"""(R50 N38) 정본 SwUDS 가 먼저다 — SwDS 파티션 맵·모듈 상속·추론은 정본이 못 채운 빈칸만 채운다.

## 실측 (2026-09-15, 고치기 전)

`generate_uds_docx` 의 정본(참조 SwUDS) 채움은 **빈칸만** 채웠다(`cur` 가 ""/TBD/N/A/- 일 때). 그런데 그 앞에서
`generate_uds_source_sections(sds_partition_map=)` 이 `asil_source="sds"` 로, `_inherit_module_asil` 이 모듈 안 첫 값을
모듈 전체로 채우면 정본이 채울 빈칸이 없다. 라이브 run 2079(kjpds02_pv):

| 축 | run 2078(SwDS 미배선) | run 2079(SwDS 배선) |
|---|---|---|
| `asil_source` reference | 657 | **35** |
| `asil_source` module_inherit | 1 | 672 |
| `asil_source` sds | 0 | 356 |
| 함수 ASIL 변경 | — | **327건** — TBD→A 215 · A→QM 45 · QM→A 40 · TBD→QM 27 (R49 는 표본 5건으로 "전부 A→QM" 이라 적었다 — 전수 정정) |
| 게이트 | 통과 | `asil_trusted_fill` 미달 |

사용자 결정(2026-09-15): **정본 SwUDS 먼저, SwDS 는 빈칸만**. 소스 주석(`comment`)·정본 직독(`uds`)·정본 자신만 지키고
나머지는 정본이 덮는다(`report_gen/provenance.reference_suds_may_override`). 덮은 건 `reference_suds.safety_fields_overridden`
(이전 출처별)·`safety_conflicts_sample` 로 공시하고, 값이 같아 출처만 올린 건 `safety_fields_agreed` 로 센다.

같은 라운드: 본문 절(`requirements` 등)이 템플릿 레이아웃에 자리가 없어 빠지면 `text_sections_unplaced` 로 공시한다
(run 2078 의 SwRS 요구 절 79,572자가 침묵으로 빠졌다).
"""
from __future__ import annotations

import json

import pytest

from report_gen.provenance import (
    REFERENCE_SUDS_KEEPS,
    WEAK_SOURCES,
    reference_suds_may_override,
)

pytest.importorskip("docx")


# ==============================================================
# 1. 판정 단일 출처
# ==============================================================

class TestPredicate:
    @pytest.mark.parametrize("src", ["comment", "reference", "uds", "COMMENT", " uds "])
    def test_kept_sources_are_not_overridden(self, src):
        assert reference_suds_may_override(src) is False

    @pytest.mark.parametrize("src", [
        "sds", "srs", "hsis", "module_inherit", "inference", "default", "", None, "unknown",
        "srs_default_qm", "swcom", "rag", "call_graph", "rule", "generated_doc", "brand_new_label",
    ])
    def test_everything_else_is_overridable(self, src):
        """뮤테이션: `sds` 를 예외에 넣으면(= SwDS 를 정본 위에 두면) 실패 — run 2079 의 역전이 되살아난다."""
        assert reference_suds_may_override(src) is True

    def test_keep_set_is_exactly_the_three_named_sources(self):
        assert REFERENCE_SUDS_KEEPS == frozenset({"comment", "reference", "uds"})

    def test_every_weak_source_is_overridable(self):
        """약한 출처가 정본을 막을 수는 없다 — 두 판정이 갈리면 모듈 상속이 정본을 이긴다(run 2079 의 672건)."""
        for w in WEAK_SOURCES:
            assert reference_suds_may_override(w) is True, w


# ==============================================================
# 2. 빌더 — 정본이 실제로 덮는가 (사이드카로 본다)
# ==============================================================

@pytest.fixture(autouse=True)
def _no_repo_template(monkeypatch, tmp_path):
    """저장소 기본 템플릿(430 heading, 300초+)·고정 참조(40MB)를 끌어오지 않는다 — `test_uds_docx_gen_stats` 와 같은 이유."""
    import config
    monkeypatch.setattr(config, "resolve_uds_template_path", lambda: "", raising=False)
    monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(tmp_path / "no_such_reference.docx"), raising=False)


def _payload(*, asil, asil_source, related="", related_source="", project="HDPDM01_PDS64_RD"):
    return {
        "project_name": project,
        "function_details": {
            "SwUFn_0001": {
                "id": "SwUFn_0001", "name": "alpha", "prototype": "void alpha(void);", "description": "설명",
                "asil": asil, "asil_source": asil_source, "related": related, "related_source": related_source,
                "inputs": [], "outputs": [], "precondition": "", "globals_global": [], "globals_static": [],
                "called": "", "logic": "",
            },
        },
    }


def _run(tmp_path, monkeypatch, ref_block, payload):
    """가짜 정본(같은 프로젝트 이름) + 파싱 결과 monkeypatch → 폴백 모드로 생성 → 보강 사이드카·통계."""
    import docx

    import config
    from report_gen import docx_builder
    from report_gen.docx_builder import enriched_function_details_path, generate_uds_docx

    ref_path = tmp_path / "(HDPDM01_SUDS) Software Unit Design Specification.docx"
    docx.Document().save(str(ref_path))
    monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(ref_path), raising=False)
    monkeypatch.setattr(docx_builder, "_extract_function_info_from_docx",
                        lambda _doc: {"REF_0001": {"name": "alpha", **ref_block}})
    out = tmp_path / "uds.docx"
    stats: dict = {}
    generate_uds_docx(None, payload, str(out), stats_out=stats)
    side = json.loads(enriched_function_details_path(str(out)).read_text(encoding="utf-8"))
    return side["function_details"]["SwUFn_0001"], stats["reference_suds"]


class TestReferenceOverridesDocSources:
    def test_sds_asil_is_overridden_by_the_reference_and_counted(self, tmp_path, monkeypatch):
        """run 2079 의 한 함수: SwDS 가 QM 이라 했고 정본은 A — 정본이 이기고, 덮은 사실이 이전 출처별로 남는다."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source="sds"))
        assert (fn["asil"], fn["asil_source"]) == ("A", "reference")
        assert ref["safety_fields_overridden"] == {"sds": 1}
        assert ref["safety_fields_agreed"] == 0
        assert ref["safety_fields_applied"] == 1
        assert ref["safety_conflicts_sample"] == [
            {"id": "SwUFn_0001", "name": "alpha", "field": "asil", "prev": "QM", "prev_source": "sds", "reference": "A"},
        ]

    @pytest.mark.parametrize("label", ["SDS", " sds ", "hsis"])
    def test_previous_source_is_canonicalised_in_the_count(self, tmp_path, monkeypatch, label):
        """이전 출처는 별칭·대소문자를 접어 센다 — `hsis`(→`sds`) 와 `SDS` 가 따로 잡히면 출처별 합이 표를 설명 못 한다."""
        _fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source=label))
        assert ref["safety_fields_overridden"] == {"sds": 1}

    def test_module_inherited_asil_is_overridden_too(self, tmp_path, monkeypatch):
        """run 2079 의 672건 — 모듈 상속은 첫 함수 값이 번진 것이라 정본 앞에 설 수 없다."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source="module_inherit"))
        assert (fn["asil"], fn["asil_source"]) == ("A", "reference")
        assert ref["safety_fields_overridden"] == {"module_inherit": 1}

    def test_comment_asil_is_kept(self, tmp_path, monkeypatch):
        """대조군 — 소스 주석 `@asil` 은 c_source 권위라 정본이 덮지 않는다(`backend/services/CLAUDE.md` 우선순위)."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source="comment"))
        assert (fn["asil"], fn["asil_source"]) == ("QM", "comment")
        assert ref["safety_fields_overridden"] == {}
        assert ref["safety_fields_applied"] == 0

    def test_same_value_is_agreed_not_conflict(self, tmp_path, monkeypatch):
        """SwDS 와 정본이 같은 값이면 충돌이 아니다 — 출처만 정본으로 올라간다(신뢰 게이트는 둘 다 trusted)."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="A", asil_source="sds"))
        assert (fn["asil"], fn["asil_source"]) == ("A", "reference")
        assert ref["safety_fields_agreed"] == 1
        assert ref["safety_fields_overridden"] == {}
        assert ref["safety_conflicts_sample"] == []

    def test_blank_fill_is_not_counted_as_override(self, tmp_path, monkeypatch):
        """예전 동작(빈칸 채움)은 그대로 — `overridden` 은 **이미 있던 값을 바꾼 것**만 센다."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="TBD", asil_source="default"))
        assert (fn["asil"], fn["asil_source"]) == ("A", "reference")
        assert ref["safety_fields_applied"] == 1
        assert ref["safety_fields_overridden"] == {} and ref["safety_fields_agreed"] == 0

    def test_identity_gate_still_blocks_overrides(self, tmp_path, monkeypatch):
        """덮어쓰기가 신원 게이트 **뒤**에 있다 — 남의 프로젝트 정본이 SwDS 값을 덮으면 두 결함이 겹친다."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source="sds", project="KJPDS02_PV"))
        assert (fn["asil"], fn["asil_source"]) == ("QM", "sds")
        assert ref["safety_fields_blocked"] == 1
        assert ref["safety_fields_overridden"] == {}

    def test_invalid_reference_asil_never_overrides(self, tmp_path, monkeypatch):
        """정본 파싱이 어긋난 값(프로토타입 문자열)은 빈칸도 못 채우고 덮지도 못한다."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "void alpha( void )"}, _payload(asil="QM", asil_source="sds"))
        assert (fn["asil"], fn["asil_source"]) == ("QM", "sds")
        assert ref["invalid_asil_rejected"] == 1 and ref["safety_fields_overridden"] == {}

    def test_related_order_and_separator_do_not_count_as_conflict(self, tmp_path, monkeypatch):
        """Related ID 나열 순서·구분자가 다른 것은 충돌이 아니다 — 아니면 충돌 계수가 나열 순서를 센 값이 된다."""
        fn, ref = _run(tmp_path, monkeypatch, {"related": "SwCom_03 / SwFn_25"},
                       _payload(asil="A", asil_source="comment", related="SwFn_25, SwCom_03", related_source="sds"))
        assert fn["related_source"] == "reference"
        assert ref["safety_fields_agreed"] == 1 and ref["safety_fields_overridden"] == {}

    def test_kept_source_conflict_is_counted_not_silent(self, tmp_path, monkeypatch):
        """(리뷰 W1) 소스 주석 `@asil QM` 과 정본 `A` 가 다르면 정본은 막히지만 그 불일치는 센다 — 가장 보고 가치가 큰 신호다."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source="comment"))
        assert (fn["asil"], fn["asil_source"]) == ("QM", "comment")
        assert ref["safety_fields_kept_conflict"] == {"comment": 1}
        _fn2, ref2 = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="A", asil_source="comment"))
        assert ref2["safety_fields_kept_conflict"] == {}, "같은 값은 불일치가 아니다"

    @pytest.mark.parametrize("prev", ["ASIL A", "asil-a", "ASIL_A", "a"])
    def test_asil_notation_variants_are_agreement_not_conflict(self, tmp_path, monkeypatch, prev):
        """(리뷰 W5) SwDS 파싱 원문 `ASIL A` 와 정본 `A` 는 같은 등급이다 — 충돌로 세면 "어긋난 함수 N건" 이 거짓 주장이 된다."""
        fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil=prev, asil_source="sds"))
        assert (fn["asil"], fn["asil_source"]) == ("A", "reference")
        assert ref["safety_fields_agreed"] == 1 and ref["safety_fields_overridden"] == {}

    def test_blank_fill_is_counted_separately_from_overrides(self, tmp_path, monkeypatch):
        """(리뷰 I1) `applied` 는 이제 합계다 — 빈칸 채움 건수를 따로 남겨야 구판 run 과 비교할 수 있다."""
        _fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="", asil_source=""))
        assert (ref["safety_fields_blank_filled"], ref["safety_fields_applied"]) == (1, 1)
        _fn2, ref2 = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source="sds"))
        assert (ref2["safety_fields_blank_filled"], ref2["safety_fields_applied"]) == (0, 1)

    def test_sample_carries_the_function_name(self, tmp_path, monkeypatch):
        """(리뷰 I4) 이름 폴백 매칭이면 id 가 정본 쪽 키일 수 있다 — 이름을 같이 싣는다."""
        _fn, ref = _run(tmp_path, monkeypatch, {"asil": "A"}, _payload(asil="QM", asil_source="sds"))
        assert ref["safety_conflicts_sample"][0]["name"] == "alpha"

    def test_related_conflict_is_counted_by_previous_source(self, tmp_path, monkeypatch):
        fn, ref = _run(tmp_path, monkeypatch, {"related": "SwFn_99"},
                       _payload(asil="A", asil_source="comment", related="SwFn_25", related_source="sds"))
        assert (fn["related"], fn["related_source"]) == ("SwFn_99", "reference")
        assert ref["safety_fields_overridden"] == {"sds": 1}
        assert ref["safety_conflicts_sample"][0]["field"] == "related"


# ==============================================================
# 2b. 모듈 상속은 정본 채움 뒤에, 씨앗은 출처 우선순위로 (리뷰 C1)
# ==============================================================

def _two_function_payload(*, beta_asil="", beta_source="", alpha_asil="QM", alpha_source="sds", gamma=None):
    """같은 모듈(SwCom_01)의 alpha(정본에 있음)·beta(정본에 없음, 빈칸) [+ gamma]. 모듈 귀속은 function_table_rows 로."""
    p = _payload(asil=alpha_asil, asil_source=alpha_source)
    fd = p["function_details"]
    fd["SwUFn_0002"] = dict(fd["SwUFn_0001"], id="SwUFn_0002", name="beta", prototype="void beta(void);",
                            asil=beta_asil, asil_source=beta_source)
    # 빌더의 `fn_module_map` 은 row[1]=모듈, row[2]=함수 ID 를 읽는다(`generate_uds_docx` 의 function_table_rows 순회).
    rows = [["SwCom_01", "SwCom_01", "SwUFn_0001", "alpha"], ["SwCom_01", "SwCom_01", "SwUFn_0002", "beta"]]
    if gamma:
        fd["SwUFn_0003"] = dict(fd["SwUFn_0001"], id="SwUFn_0003", name="gamma", prototype="void gamma(void);", **gamma)
        rows.append(["SwCom_01", "SwCom_01", "SwUFn_0003", "gamma"])
    p["function_table_rows"] = rows
    return p


def _run2(tmp_path, monkeypatch, ref_block, payload):
    import docx

    import config
    from report_gen import docx_builder
    from report_gen.docx_builder import enriched_function_details_path, generate_uds_docx

    ref_path = tmp_path / "(HDPDM01_SUDS) Software Unit Design Specification.docx"
    docx.Document().save(str(ref_path))
    monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(ref_path), raising=False)
    monkeypatch.setattr(docx_builder, "_extract_function_info_from_docx",
                        lambda _doc: {"REF_0001": {"name": "alpha", **ref_block}})
    out = tmp_path / "uds2.docx"
    stats: dict = {}
    generate_uds_docx(None, payload, str(out), stats_out=stats)
    side = json.loads(enriched_function_details_path(str(out)).read_text(encoding="utf-8"))
    return side["function_details"], stats["reference_suds"]


class TestModuleInheritanceAfterReference:
    def test_sibling_inherits_the_reference_value_not_the_sds_seed(self, tmp_path, monkeypatch):
        """run 2079 의 672건: alpha(SwDS=QM, 정본=A)·beta(빈칸). 예전엔 beta 가 QM(module_inherit) 로 굳었다 — 정본이 못 닿는 함수에
        정본이 틀렸다고 판정한 값이 남는 under-classification. 이제 상속이 정본 뒤라 beta 는 A 를 물려받는다."""
        fd, ref = _run2(tmp_path, monkeypatch, {"asil": "A"}, _two_function_payload())
        assert (fd["SwUFn_0001"]["asil"], fd["SwUFn_0001"]["asil_source"]) == ("A", "reference")
        assert (fd["SwUFn_0002"]["asil"], fd["SwUFn_0002"]["asil_source"]) == ("A", "module_inherit")
        assert ref["module_inherit_by_seed_source"] == {"reference": 1}

    def test_seed_prefers_reference_over_an_earlier_sds_sibling(self, tmp_path, monkeypatch):
        """dict 순서상 SwDS 함수가 먼저 와도 씨앗은 정본이다 — gamma(SwDS=QM, 정본에 없음)가 첫 함수여도 beta 는 A."""
        p = _two_function_payload(gamma={"asil": "QM", "asil_source": "sds"})
        # gamma 를 dict 맨 앞으로
        fd0 = p["function_details"]
        p["function_details"] = {"SwUFn_0003": fd0["SwUFn_0003"], "SwUFn_0001": fd0["SwUFn_0001"], "SwUFn_0002": fd0["SwUFn_0002"]}
        fd, ref = _run2(tmp_path, monkeypatch, {"asil": "A"}, p)
        assert fd["SwUFn_0002"]["asil"] == "A" and fd["SwUFn_0002"]["asil_source"] == "module_inherit"
        assert fd["SwUFn_0003"]["asil"] == "QM" and fd["SwUFn_0003"]["asil_source"] == "sds", "정본에 없는 SwDS 값은 그대로(빈칸 채움)"
        assert ref["module_inherit_by_seed_source"] == {"reference": 1}

    def test_sds_seed_still_fills_when_no_reference_sibling(self, tmp_path, monkeypatch):
        """대조군 — 정본이 그 모듈에 아무 함수도 없으면 SwDS 씨앗으로 채운다(사용자 결정 "SwDS 는 빈칸만"). 씨앗 출처가 남는다."""
        fd, ref = _run2(tmp_path, monkeypatch, {"related": "SwFn_99"}, _two_function_payload())
        assert (fd["SwUFn_0002"]["asil"], fd["SwUFn_0002"]["asil_source"]) == ("QM", "module_inherit")
        assert ref["module_inherit_by_seed_source"] == {"sds": 1}


# ==============================================================
# 3. 본문 절이 템플릿에 자리가 없으면 공시한다
# ==============================================================

def _make_template(path, level1_headings):
    import docx
    d = docx.Document()
    for h in level1_headings:
        d.add_heading(h, level=1)
    d.add_heading("SwUFn_0001: alpha", level=4)
    d.save(str(path))
    return str(path)


def _template_payload(**text):
    p = _payload(asil="A", asil_source="comment")
    p.update({"overview": "", "requirements": "", "interfaces": "", "uds_frames": "", "notes": ""})
    p.update(text)
    return p


def _gen_stats(tmp_path, template, payload):
    from report_gen.docx_builder import gen_stats_path, generate_uds_docx
    out = tmp_path / "uds.docx"
    generate_uds_docx(template, payload, str(out), stats_out={})
    return json.loads(gen_stats_path(str(out)).read_text(encoding="utf-8"))


class TestTextSectionsUnplaced:
    def test_requirements_without_a_heading_is_reported_with_its_length(self, tmp_path):
        """정본 레이아웃(Introduction·Software Unit Design 만) + SwRS 요구 절 → 실리지 않았다는 사실과 글자수."""
        tpl = _make_template(tmp_path / "t.docx", ["Introduction", "Software Unit Design"])
        side = _gen_stats(tmp_path, tpl, _template_payload(requirements="- SwEI_01 battery power source\n" * 40))
        unplaced = side["text_sections_unplaced"]
        assert set(unplaced) == {"requirements"}
        assert isinstance(unplaced["requirements"], int) and unplaced["requirements"] > 0

    def test_requirements_with_a_heading_is_not_reported(self, tmp_path):
        tpl = _make_template(tmp_path / "t.docx", ["Introduction", "Requirements", "Software Unit Design"])
        side = _gen_stats(tmp_path, tpl, _template_payload(requirements="- SwEI_01 battery power source"))
        assert side["text_sections_unplaced"] == {}

    def test_empty_sections_are_not_unplaced(self, tmp_path):
        """비어 있는 절은 '자리가 없어 빠진 것' 이 아니다 — 없는 입력을 결함으로 세면 매 run 오탐이다."""
        tpl = _make_template(tmp_path / "t.docx", ["Introduction"])
        side = _gen_stats(tmp_path, tpl, _template_payload())
        assert side["text_sections_unplaced"] == {}

    def test_placeholder_template_reports_every_section_unplaced(self, tmp_path):
        """(리뷰 W3) 토큰 치환 템플릿은 본문 절을 싣는 자리 자체가 없다 — 같은 키로 전부 미배치라고 말한다."""
        import docx
        d = docx.Document()
        d.add_paragraph("Owner: {{custom_owner_token}}")   # 빌더가 아는 토큰은 먼저 치환돼 사라지므로 미지 토큰으로 판정을 유도
        tpl = tmp_path / "ph.docx"
        d.save(str(tpl))
        side = _gen_stats(tmp_path, str(tpl), _template_payload(requirements="- SwEI_01", notes="n"))
        assert side["mode"] == "placeholder_substitution"
        assert set(side["text_sections_unplaced"]) == {"requirements", "notes"}

    def test_every_text_section_key_is_tracked(self, tmp_path):
        tpl = _make_template(tmp_path / "t.docx", ["Introduction"])
        side = _gen_stats(tmp_path, tpl, _template_payload(overview="o", requirements="r", interfaces="i", uds_frames="u", notes="n"))
        assert set(side["text_sections_unplaced"]) == {"overview", "requirements", "interfaces", "uds_frames", "notes"}
