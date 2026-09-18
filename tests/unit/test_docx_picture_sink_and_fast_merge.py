# tests/unit/test_docx_picture_sink_and_fast_merge.py
"""(R53 N27-c) DOCX 단계에서 그림마다·병합마다 문서 전체를 훑던 python-docx 내부 비용 — **같은 결과**를 내며 걷어낸다.

## 왜 이 파일이 있나

R47-j 뒤 남은 DOCX 자식 572초(라이브 프로파일)의 절반이 python-docx 1.2.0 내부였다:
`Run.add_picture` 아래 26.3% — `StoryPart.next_id` 가 그림마다 문서 전체에 `//@id` xpath, `ImageParts._get_by_sha1` 가
그림마다 기존 이미지 파트 **전부**를 다시 sha1(캐시 없음) — 그리고 `CT_Tc.merge` 아래 24.7% — 병합마다 `_tbl`/`_tr`/sibling
xpath. 빌더는 `_PictureSink`(그림) 와 `_hmerge_fast`(같은 행 가로 병합) 로 **같은 XML** 을 낸다.

## 무엇을 고정하나

1. **결과 동일** — 두 경로는 python-docx 원판과 document.xml · rels · 이미지 파트 · 저장 바이트가 같다(고정 시드 무작위 포함).
   빠른 경로가 감당하지 않는 모양(다른 행 · 내용 있는 셀 · vMerge · b 가 왼쪽)은 원판으로 떨어지고 그때도 같다.
2. **접근 경로** — 싱크는 `next_id` 0회 · 파트당 sha1 1회 · `//@id` 스캔 1회, 빠른 병합은 `CT_Tc.merge` 0회 · xpath 0회
   (결과가 같은 두 구현은 관측량이 **횟수**여야 갈린다 — R47-h/R47-j 교훈).
3. **배선** — `generate_uds_docx` 의 그림 삽입·원본 블록 되붙임은 두 분기 모두 싱크를 지나고, 함수 정보 표 병합은 `_merge_cells` 뿐이다.
"""
from __future__ import annotations

import copy
import logging
import random
import re
import struct

import pytest

from tests.unit._source_probe import source_of

docx = pytest.importorskip("docx", reason="python-docx 없음")

from docx.exceptions import InvalidSpanError  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.oxml.table import CT_Tc  # noqa: E402
from docx.oxml.xmlchemy import BaseOxmlElement  # noqa: E402
from docx.package import ImageParts  # noqa: E402
from docx.parts.image import ImagePart  # noqa: E402
from docx.parts.story import StoryPart  # noqa: E402
from docx.shared import Inches  # noqa: E402
from docx.table import _Cell  # noqa: E402

from report_gen import docx_builder  # noqa: E402
from report_gen.atomic_io import normalize_zip_member_times  # noqa: E402
from report_gen.docx_builder import (  # noqa: E402
    FN_ROW_FULL,
    FN_ROW_PAIR,
    _append_body_block,
    _hmerge_fast,
    _insert_logic_image_in_table,
    _merge_cells,
    _merge_function_info_table,
    _PictureSink,
    _save_docx,
)

_PNG_1x1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da6364f8cfc000000201"
    "0100c9fe92ef0000000049454e44ae426082"
)


def _png_variant(k: int) -> bytes:
    """IDAT 데이터 한 바이트만 다른 PNG — python-docx 는 청크 길이로만 걷고 CRC 를 보지 않아 헤더 해석은 그대로다(다른 sha1)."""
    i = _PNG_1x1.index(b"IDAT") + 4 + 6
    return _PNG_1x1[:i] + bytes([k & 0xFF]) + _PNG_1x1[i + 1:]


def _gif(w: int = 4, h: int = 3) -> bytes:
    return b"GIF89a" + struct.pack("<HH", w, h) + bytes(7)


@pytest.fixture
def images(tmp_path):
    files = {
        "a.png": _PNG_1x1, "a_dup.png": _PNG_1x1, "b.png": _png_variant(0x11), "c.png": _png_variant(0x22),
        "d.png": _png_variant(0x33), "g.gif": _gif(),
    }
    out = {}
    for name, blob in files.items():
        p = tmp_path / name
        p.write_bytes(blob)
        out[name] = str(p)
    return out


def _doc():
    return docx.Document()


def _snapshot(doc):
    parts = doc.part.package.image_parts._image_parts
    rels = sorted(
        (rId, rel.reltype, str(rel.target_ref if rel.is_external else rel.target_part.partname))
        for rId, rel in doc.part.rels.items()
    )
    return doc.element.xml, rels, [(str(p.partname), p.blob) for p in parts]


def _saved(doc, path):
    doc.save(str(path))
    assert normalize_zip_member_times(path)
    return path.read_bytes()


def _assert_same_docs(a, b, tmp_path, tag="x"):
    assert _snapshot(a) == _snapshot(b)
    assert _saved(a, tmp_path / f"{tag}_a.docx") == _saved(b, tmp_path / f"{tag}_b.docx")


def _template_with_images(tmp_path, images):
    """기존 이미지 파트가 있는 템플릿 — a.png · g.gif 에 더해 a.png 와 **같은 바이트의 두 번째 파트**(원판은 목록의 앞 것을 고른다)."""
    from docx.image.image import Image
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    d = _doc()
    d.add_picture(images["a.png"], width=Inches(1))
    d.add_picture(images["g.gif"], width=Inches(1))
    dup = d.part.package.image_parts._add_image_part(Image.from_file(images["a_dup.png"]))
    d.part.relate_to(dup, RT.IMAGE)                      # 관계가 있어야 저장·재적재된다
    p = tmp_path / "tpl.docx"
    d.save(str(p))
    return p


def _set_docpr_ids(el, value: int) -> int:
    n = 0
    for d in el.xpath(".//wp:docPr"):
        d.set("id", str(value))
        n += 1
    return n


# ── 1. 그림 싱크 == Run.add_picture ───────────────────────────────────────────

class TestPictureSinkMatchesRunAddPicture:
    def test_distinct_images_widths_and_heights(self, images, tmp_path):
        a, b = _doc(), _doc()
        sink = _PictureSink(b)
        for name, kw in (("a.png", {"width": Inches(5.2)}), ("b.png", {}), ("g.gif", {"height": Inches(1)}),
                         ("c.png", {"width": Inches(2), "height": Inches(3)})):
            a.add_paragraph().add_run().add_picture(images[name], **kw)
            sink.add_picture(b.add_paragraph().add_run(), images[name], **kw)
        assert sink.pictures == 4 and sink.full_id_scans == 1
        _assert_same_docs(a, b, tmp_path, "distinct")

    def test_duplicate_bytes_reuse_the_first_part_and_its_filename(self, images, tmp_path):
        a, b = _doc(), _doc()
        sink = _PictureSink(b)
        for name in ("a.png", "a_dup.png", "a.png"):
            a.add_paragraph().add_run().add_picture(images[name])
            sink.add_picture(b.add_paragraph().add_run(), images[name])
        assert len(b.part.package.image_parts._image_parts) == 1
        # 원판은 매칭된 **기존 파트**의 이름을 docPr 에 쓴다 — 세 그림 전부 "a.png"
        names = [d.get("name") for d in b.element.xpath(".//pic:cNvPr")]
        assert names == ["a.png"] * 3
        _assert_same_docs(a, b, tmp_path, "dup")

    def test_template_with_existing_parts_including_duplicates(self, images, tmp_path):
        tpl = _template_with_images(tmp_path, images)
        a, b = docx.Document(str(tpl)), docx.Document(str(tpl))
        assert len(b.part.package.image_parts._image_parts) == 3            # 전제: a · g · a 사본
        sink = _PictureSink(b)
        for name in ("a.png", "d.png", "g.gif", "a_dup.png"):
            a.add_paragraph().add_run().add_picture(images[name], width=Inches(4))
            sink.add_picture(b.add_paragraph().add_run(), images[name], width=Inches(4))
        assert len(b.part.package.image_parts._image_parts) == 4            # d.png 만 새 파트
        _assert_same_docs(a, b, tmp_path, "tpl")

    def test_foreign_block_appended_between_pictures_raises_the_next_id(self, images, tmp_path):
        a, b = _doc(), _doc()
        sink = _PictureSink(b)
        a.add_paragraph().add_run().add_picture(images["a.png"])
        sink.add_picture(b.add_paragraph().add_run(), images["a.png"])
        for d, kw in ((a, {}), (b, {"picture_sink": sink})):
            el = copy.deepcopy(d.paragraphs[-1]._p)
            assert _set_docpr_ids(el, 777) == 1
            _append_body_block(d, el, **kw)
        a.add_paragraph().add_run().add_picture(images["b.png"])
        sink.add_picture(b.add_paragraph().add_run(), images["b.png"])
        assert [d.get("id") for d in b.element.xpath(".//wp:docPr")] == ["1", "777", "778"]
        _assert_same_docs(a, b, tmp_path, "append")

    def test_foreign_add_picture_needs_invalidate(self, images, tmp_path):
        a, b, c = _doc(), _doc(), _doc()
        sb, sc = _PictureSink(b), _PictureSink(c)
        for d, s in ((b, sb), (c, sc)):
            s.add_picture(d.add_paragraph().add_run(), images["a.png"])
            d.add_picture(images["b.png"])                 # 싱크를 거치지 않은 원판 경로
        a.add_paragraph().add_run().add_picture(images["a.png"])
        a.add_picture(images["b.png"])
        sb.invalidate()                                    # c 는 통지하지 않는다 — id 가 겹쳐야 한다(자기 점검)
        for d, s in ((b, sb), (c, sc)):
            s.add_picture(d.add_paragraph().add_run(), images["c.png"])
        a.add_paragraph().add_run().add_picture(images["c.png"])
        assert sb.full_id_scans == 2
        _assert_same_docs(a, b, tmp_path, "inval")
        assert [d.get("id") for d in c.element.xpath(".//wp:docPr")] == ["1", "2", "2"], "invalidate 없이는 id 가 겹친다(자기 점검)"

    def test_partname_numbering_reuses_gaps_like_python_docx(self, images, tmp_path):
        """원판 `_next_image_partname` 은 "빈 번호 중 최소" 를 쓴다 — image2 가 빠진 템플릿에 새 그림을 넣으면 image2 가 된다."""
        d = _doc()
        for name in ("a.png", "b.png", "c.png"):
            d.add_picture(images[name])
        rels = d.part.rels
        rid2 = [rId for rId, rel in rels.items() if str(rel.target_part.partname).endswith("image2.png")][0]
        del rels[rid2]                                   # 관계가 끊긴 파트는 저장되지 않는다 → 재적재 시 번호 2 가 빈다
        tpl = tmp_path / "gap.docx"
        d.save(str(tpl))
        a, b = docx.Document(str(tpl)), docx.Document(str(tpl))
        assert sorted(str(p.partname) for p in b.part.package.image_parts._image_parts) == [
            "/word/media/image1.png", "/word/media/image3.png"]
        sink = _PictureSink(b)
        for name in ("d.png", "g.gif", "a.png"):
            a.add_paragraph().add_run().add_picture(images[name])
            sink.add_picture(b.add_paragraph().add_run(), images[name])
        assert [str(p.partname) for p in b.part.package.image_parts._image_parts][2:] == [
            "/word/media/image2.png", "/word/media/image4.gif"]
        _assert_same_docs(a, b, tmp_path, "gap")

    def test_add_picture_paragraph_matches_document_add_picture(self, images, tmp_path):
        a, b = _doc(), _doc()
        sink = _PictureSink(b)
        a.add_picture(images["a.png"], width=Inches(5))
        sink.add_picture_paragraph(images["a.png"], width=Inches(5))
        _assert_same_docs(a, b, tmp_path, "para")

    def test_missing_file_raises_the_same_and_leaves_the_document_untouched(self, tmp_path):
        a, b = _doc(), _doc()
        sink = _PictureSink(b)
        run_a, run_b = a.add_paragraph().add_run(), b.add_paragraph().add_run()
        before = _snapshot(b)
        missing = str(tmp_path / "nope.png")
        with pytest.raises(FileNotFoundError):
            run_a.add_picture(missing)
        with pytest.raises(FileNotFoundError):
            sink.add_picture(run_b, missing)
        assert _snapshot(b) == before and sink.pictures == 0
        assert _snapshot(a) == _snapshot(b)

    @pytest.mark.parametrize("seed", range(12))
    def test_random_interleavings_match(self, seed, images, tmp_path):
        rng = random.Random(20260916 + seed)
        tpl = _template_with_images(tmp_path, images)
        a, b = docx.Document(str(tpl)), docx.Document(str(tpl))
        sink = _PictureSink(b)
        names = list(images)
        for _ in range(10):
            op = rng.choice(("pic", "pic", "pic", "foreign_pic", "append"))
            name = rng.choice(names)
            kw = rng.choice(({}, {"width": Inches(rng.randint(1, 6))}, {"height": Inches(1)}))
            if op == "pic":
                a.add_paragraph().add_run().add_picture(images[name], **kw)
                sink.add_picture(b.add_paragraph().add_run(), images[name], **kw)
            elif op == "foreign_pic":
                a.add_picture(images[name], **kw)
                b.add_picture(images[name], **kw)
                sink.invalidate()
            else:
                src_a = [i for i, p in enumerate(a.paragraphs) if p._p.xpath(".//wp:docPr")]
                if not src_a:
                    continue
                idx = src_a[-1]
                ident = rng.randint(1, 900)
                for d, k in ((a, {}), (b, {"picture_sink": sink})):
                    el = copy.deepcopy(d.paragraphs[idx]._p)
                    _set_docpr_ids(el, ident)
                    _append_body_block(d, el, **k)
        _assert_same_docs(a, b, tmp_path, f"rand{seed}")


# ── 2. 접근 경로 — 횟수가 관측량이다 ────────────────────────────────────────

@pytest.fixture
def counters(monkeypatch):
    calls = {"next_id": 0, "sha1": 0, "id_scan": 0, "merge": 0, "xpath": 0, "partname": 0}
    real_next = StoryPart.next_id.fget
    real_sha1 = ImagePart.sha1.fget
    real_xpath = BaseOxmlElement.xpath
    real_merge = CT_Tc.merge
    real_partname = ImageParts._next_image_partname

    def _partname(self, ext):
        calls["partname"] += 1
        return real_partname(self, ext)
    monkeypatch.setattr(ImageParts, "_next_image_partname", _partname)

    def _next(self):
        calls["next_id"] += 1
        return real_next(self)

    def _sha1(self):
        calls["sha1"] += 1
        return real_sha1(self)

    def _xpath(self, expr):
        calls["xpath"] += 1
        if expr == "//@id":
            calls["id_scan"] += 1
        return real_xpath(self, expr)

    def _merge(self, other):
        calls["merge"] += 1
        return real_merge(self, other)

    monkeypatch.setattr(StoryPart, "next_id", property(_next))
    monkeypatch.setattr(ImagePart, "sha1", property(_sha1))
    monkeypatch.setattr(BaseOxmlElement, "xpath", _xpath)
    monkeypatch.setattr(CT_Tc, "merge", _merge)
    return calls


class TestPictureSinkAccessPath:
    def test_sink_scans_ids_once_and_hashes_each_existing_part_once(self, counters, images, tmp_path):
        tpl = _template_with_images(tmp_path, images)          # (템플릿을 만드는 add_picture 도 세므로 여기서 0 으로)
        for k in counters:
            counters[k] = 0
        new_names = ["b.png", "c.png", "d.png", "b.png", "c.png", "d.png"]
        ref = docx.Document(str(tpl))
        for name in new_names:
            ref.add_paragraph().add_run().add_picture(images[name])
        ref_calls = dict(counters)
        for k in counters:
            counters[k] = 0
        d = docx.Document(str(tpl))
        sink = _PictureSink(d)
        for name in new_names:
            sink.add_picture(d.add_paragraph().add_run(), images[name])
        # 원판: 그림마다 next_id 1회 + `_get_by_sha1` 가 목록을 **첫 일치까지** 해시 — 기존 3 뒤에 b·c·d 가 붙으므로
        #   3, 4, 5(새 파트) 다음 4, 5, 6(각각 자기 자리까지) = 27. 파트 수에 비례해 그림마다 다시 센다.
        assert ref_calls["next_id"] == 6 and ref_calls["id_scan"] == 6 and ref_calls["sha1"] == 27, ref_calls
        assert ref_calls["partname"] == 3                    # 새 파트마다 파트 전부를 훑는 채번
        # 싱크: next_id 0회 · `//@id` 1회 · 기존 파트 3개만 각 1회(새 파트는 들어온 이미지의 해시를 그대로 쓴다) · 채번 훑기 0회
        assert (counters["next_id"], counters["id_scan"], counters["sha1"], counters["partname"]) == (0, 1, 3, 0), counters
        assert sink.pictures == 6 and len(d.part.package.image_parts._image_parts) == 6

    def test_insert_logic_image_goes_through_the_sink_when_given(self, counters, images):
        d = _doc()
        t = d.add_table(rows=4, cols=6)
        t.cell(1, 0).text = "Logic Diagram"
        sink = _PictureSink(d)
        counters["next_id"] = 0
        assert _insert_logic_image_in_table(t, 6, images["a.png"], picture_sink=sink) is True
        assert counters["next_id"] == 0 and sink.pictures == 1
        t2 = d.add_table(rows=4, cols=6)
        t2.cell(1, 0).text = "Logic Diagram"
        assert _insert_logic_image_in_table(t2, 6, images["a.png"]) is True     # 싱크 없이 → 원판 경로
        assert counters["next_id"] == 1


# ── 3. 빠른 병합 == CT_Tc.merge ──────────────────────────────────────────────

def _build(spec):
    """같은 spec 에서 같은 표를 만든다(두 구현은 **각자의 새 표**에 돈다)."""
    d = _doc()
    t = d.add_table(rows=spec["rows"], cols=spec["cols"])
    for (r, c), kind in spec["widths"].items():
        tc = t._tbl.tr_lst[r].tc_lst[c]
        tcPr = tc.tcPr
        if kind == "strip":
            tc.remove(tcPr)
        elif kind == "none":
            tcPr.remove(tcPr.tcW)
        elif kind == "pct":
            tcPr.tcW.set(qn("w:type"), "pct")
            tcPr.tcW.set(qn("w:w"), "2500")
        else:
            tcPr.tcW.set(qn("w:w"), str(kind))
    for (r, c), text in spec["texts"].items():
        t._tbl.tr_lst[r].tc_lst[c].add_p().add_r().text = text
    for (r1, c1), (r2, c2) in spec["pre"]:
        try:
            t.cell(r1, c1).merge(t.cell(r2, c2))
        except (InvalidSpanError, ValueError):      # silent-ok: 사전 병합이 안 되는 모양이면 그 단계는 건너뛴다(양쪽 동일)
            pass
    return t


def _random_spec(rng):
    rows, cols = rng.randint(2, 5), rng.randint(4, 8)
    widths = {}
    texts = {}
    for r in range(rows):
        for c in range(cols):
            k = rng.random()
            if k < 0.10:
                widths[(r, c)] = "strip"
            elif k < 0.20:
                widths[(r, c)] = "none"
            elif k < 0.30:
                widths[(r, c)] = "pct"
            elif k < 0.60:
                widths[(r, c)] = rng.randint(100, 5000)
            if rng.random() < 0.15:
                texts[(r, c)] = f"t{r}{c}"
    pre = []
    for _ in range(rng.randint(0, 2)):
        r = rng.randrange(rows)
        if rng.random() < 0.7:
            c1, c2 = rng.randrange(cols), rng.randrange(cols)
            pre.append(((r, min(c1, c2)), (r, max(c1, c2))))
        else:
            r2, c = rng.randrange(rows), rng.randrange(cols)
            pre.append(((min(r, r2), c), (max(r, r2), c)))
    return {"rows": rows, "cols": cols, "widths": widths, "texts": texts, "pre": pre}


def _apply(t, op, fast: bool):
    (r1, c1), (r2, c2) = op
    try:
        if fast:
            _merge_cells(t.cell(r1, c1), t.cell(r2, c2))
        else:
            t.cell(r1, c1).merge(t.cell(r2, c2))
    except Exception as exc:   # noqa: BLE001 — 예외 종류도 대조 대상이다
        return type(exc).__name__
    return None


class TestFastMergeMatchesPythonDocx:
    @pytest.mark.parametrize("seed", range(60))
    def test_random_shapes_and_operations(self, seed):
        rng = random.Random(53_000 + seed)
        spec = _random_spec(rng)
        for _ in range(4):
            rows, cols = spec["rows"], spec["cols"]
            r = rng.randrange(rows)
            if rng.random() < 0.8:
                c1, c2 = rng.randrange(cols), rng.randrange(cols)
                op = ((r, c1), (r, c2))                     # 같은 행 — 뒤집힌 순서도 섞인다
            else:
                op = ((r, rng.randrange(cols)), (rng.randrange(rows), rng.randrange(cols)))
            ref, new = _build(spec), _build(spec)
            assert ref._tbl.xml == new._tbl.xml            # 전제: 같은 spec → 같은 표
            e_ref, e_new = _apply(ref, op, fast=False), _apply(new, op, fast=True)
            assert e_ref == e_new, (spec, op)
            assert ref._tbl.xml == new._tbl.xml, (spec, op)

    def test_row_wrapped_in_sdt_is_left_to_python_docx_which_raises(self):
        """(리뷰 W2) `w:tbl/w:sdt/w:sdtContent/w:tr` 안의 셀 — 원판은 `tr_lst.index` 에서 ValueError. 빠른 경로가 대신 성공하면 안 된다."""
        from docx.oxml import OxmlElement

        def _wrapped():
            t = _doc().add_table(rows=2, cols=4)
            tr = t._tbl.tr_lst[1]
            sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
            tr.addprevious(sdt)
            sdt.append(content)
            content.append(tr)                                   # tr 를 sdt 안으로 옮긴다
            return t, tr
        ref_t, ref_tr = _wrapped()
        new_t, new_tr = _wrapped()
        with pytest.raises(ValueError):
            ref_tr.tc_lst[0].merge(ref_tr.tc_lst[2])
        assert _hmerge_fast(new_tr.tc_lst[0], new_tr.tc_lst[2]) is False
        with pytest.raises(ValueError):
            _merge_cells(_Cell(new_tr.tc_lst[0], new_t), _Cell(new_tr.tc_lst[2], new_t))
        assert ref_t._tbl.xml == new_t._tbl.xml

    @pytest.mark.parametrize("case", ["blank", "same_cell", "b_left_of_a", "content_in_swallowed", "vmerge", "no_tcpr", "pct_and_none"])
    def test_named_shapes(self, case, counters):
        spec = {"rows": 3, "cols": 6, "widths": {}, "texts": {}, "pre": []}
        op = ((1, 1), (1, 4))
        if case == "same_cell":
            op = ((1, 2), (1, 2))
        elif case == "b_left_of_a":
            op = ((1, 4), (1, 1))
        elif case == "content_in_swallowed":
            spec["texts"] = {(1, 3): "keep me"}
        elif case == "vmerge":
            spec["pre"] = [((0, 1), (2, 1))]
        elif case == "no_tcpr":
            spec["widths"] = {(1, 1): "strip", (1, 2): "strip"}
        elif case == "pct_and_none":
            spec["widths"] = {(1, 1): "pct", (1, 2): "none", (1, 3): 800}
        ref, new = _build(spec), _build(spec)
        counters["merge"] = 0
        assert _apply(ref, op, fast=False) == _apply(new, op, fast=True)
        assert ref._tbl.xml == new._tbl.xml
        fast_expected = case in ("blank", "same_cell", "no_tcpr", "pct_and_none")
        # 참조 쪽 1회는 항상 센다 — 빠른 경로가 맡은 모양이면 새 쪽은 원판을 부르지 않는다(1회), 아니면 원판으로 떨어진다(2회)
        assert counters["merge"] == (1 if fast_expected else 2), (case, counters["merge"])

    def test_fast_path_declines_without_touching_the_table(self):
        spec = {"rows": 2, "cols": 5, "widths": {}, "texts": {(0, 2): "x"}, "pre": []}
        t = _build(spec)
        before = t._tbl.xml
        assert _hmerge_fast(t.cell(0, 0)._tc, t.cell(0, 3)._tc) is False
        assert t._tbl.xml == before

    def test_function_info_merge_takes_the_fast_path_only(self, counters):
        t = _doc().add_table(rows=20, cols=6)
        layout = [(FN_ROW_FULL, ["[ Function Information ]"])] + [(FN_ROW_PAIR, [f"L{i}", "", f"V{i}"]) for i in range(19)]
        counters["merge"] = counters["xpath"] = 0
        _merge_function_info_table(t, 6, layout)
        assert counters["merge"] == 0, "함수 정보 표 병합이 python-docx merge 로 되돌아갔다"
        assert counters["xpath"] == 0, f"빠른 경로가 xpath 를 {counters['xpath']}회 돌렸다"
        assert [len(tr.tc_lst) for tr in t._tbl.tr_lst] == [1] + [2] * 19

    def test_function_info_merge_falls_back_when_a_cell_has_content(self, counters):
        from tests.unit.test_docx_grid_snapshot import _ref_merge
        new, ref = _doc().add_table(rows=4, cols=6), _doc().add_table(rows=4, cols=6)
        for t in (new, ref):
            t.cell(2, 3).text = "already here"
        layout = [(FN_ROW_FULL, ["h"])] + [(FN_ROW_PAIR, ["L", "", "V"])] * 3
        counters["merge"] = 0
        _merge_function_info_table(new, 6, layout)
        assert counters["merge"] == 1                       # 내용 있는 칸을 삼키는 그 한 번만 원판
        _ref_merge(ref, 6, layout)
        assert new._tbl.xml == ref._tbl.xml


# ── 4. 배선 ─────────────────────────────────────────────────────────────────

class TestWiring:
    def test_all_picture_insertions_in_generate_uds_docx_go_through_the_sink(self):
        src = source_of(docx_builder.generate_uds_docx)
        assert src.count("_PictureSink(doc)") == 2, "템플릿 분기·무템플릿 분기 각각 싱크가 있어야 한다(쌍둥이 한쪽만 아님)"
        assert not re.search(r"(?<!_pics)\.add_picture(_paragraph)?\(", src)       # (리뷰 I8) 수신자 이름과 무관하게
        calls = re.findall(r"_insert_logic_image_in_table\(([^\n]*)\)", src)
        assert calls and all("picture_sink=_pics" in c for c in calls), calls
        appends = re.findall(r"_append_body_block\(doc,([^\n]*)\)", src)
        assert len(appends) == 3 and all("picture_sink=_pics" in c for c in appends), appends

    def test_every_picture_call_in_the_module_is_the_sink_or_its_documented_fallback(self):
        """(리뷰 W3/I8) 모듈 **전체**의 `*.add_picture*(` 호출을 AST 로 센다 — 허용은 싱크 자신·`picture_sink`·`_pics` 와
        `_insert_logic_image_in_table` 의 싱크 없는 폴백 한 곳뿐. 다른 곳에 생기면 `next_id` 를 우회해 docPr id 가 겹칠 수 있다."""
        import ast
        from pathlib import Path
        tree = ast.parse(Path(docx_builder.__file__).read_text(encoding="utf-8"))
        parents = {}
        for node in ast.walk(tree):
            for child in ast.iter_child_nodes(node):
                parents[child] = node

        def enclosing(node):
            while node in parents:
                node = parents[node]
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    return node.name
            return None

        offenders = []
        seen = set()
        for node in ast.walk(tree):
            if not (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)):
                continue
            if node.func.attr not in ("add_picture", "add_picture_paragraph"):
                continue
            recv, fn = ast.unparse(node.func.value), enclosing(node)
            seen.add((recv, fn))
            ok = recv in ("_pics", "picture_sink") \
                or (recv == "self" and fn == "add_picture_paragraph") \
                or (recv == "run" and fn == "_insert_logic_image_in_table")
            if not ok:
                offenders.append((node.lineno, recv, fn))
        assert offenders == [], offenders
        assert ("run", "_insert_logic_image_in_table") in seen and ("self", "add_picture_paragraph") in seen   # 허용 목록이 실제와 맞는가

    def test_save_docx_reports_duplicate_drawing_ids(self, images, tmp_path, caplog):
        """(리뷰 W3) 싱크를 거치지 않은 그림이 id 를 겹치게 하면 `gen_stats.drawing_id_duplicates` 와 경고로 드러난다."""
        d = _doc()
        d.add_picture(images["a.png"])
        d.add_picture(images["b.png"])
        ids = d.element.xpath(".//wp:docPr")
        assert [x.get("id") for x in ids] == ["1", "2"]
        stats = {}
        assert _save_docx(d, tmp_path / "clean.docx", str(tmp_path / "clean.docx"), stats)
        assert stats["drawing_id_duplicates"] == 0
        ids[1].set("id", "1")                                    # 우회 삽입이 만들 모양
        stats = {}
        with caplog.at_level(logging.WARNING):
            assert _save_docx(d, tmp_path / "dup.docx", str(tmp_path / "dup.docx"), stats)
        assert stats["drawing_id_duplicates"] == 1
        assert any("wp:docPr/@id" in r.getMessage() for r in caplog.records)

    def test_python_docx_private_surface_the_sink_relies_on(self):
        """(리뷰 I7) 싱크·빠른 병합이 기대는 python-docx 내부 이름 — 버전이 바뀌어 사라지면 여기서 먼저 시끄럽게."""
        from docx.image.image import Image
        from docx.opc.packuri import PackURI
        from docx.oxml.shape import CT_Inline
        from docx.oxml.table import CT_Row
        from docx.package import Package
        assert docx.__version__.startswith("1."), docx.__version__
        for obj, names in ((Package, ["image_parts"]), (ImageParts, ["append"]),
                           (ImagePart, ["from_image", "image", "sha1"]), (Image, ["from_file", "sha1", "ext", "filename", "scaled_dimensions"]),
                           (PackURI, ["idx"]), (CT_Inline, ["new_pic_inline"]),
                           (CT_Tc, ["_is_empty", "_add_width_of", "grid_span", "vMerge", "tcPr", "clear_content", "add_p"]),
                           (CT_Row, ["tc_lst"]), (StoryPart, ["next_id", "relate_to"])):
            missing = [n for n in names if not (hasattr(obj, n) or n in getattr(obj, "__dict__", {}))]
            assert missing == [], (obj.__name__, missing)
        assert "_image_parts" in vars(ImageParts()), "ImageParts 인스턴스가 `_image_parts` 목록을 갖지 않는다"

    def test_function_info_merge_uses_merge_cells_only(self):
        body = source_of(_merge_function_info_table).split('"""', 2)[-1]
        code = "\n".join(ln for ln in body.splitlines() if not ln.strip().startswith("#"))
        assert ".merge(" not in code and code.count("_merge_cells(") == 4

    def test_insert_logic_image_prefers_the_sink(self):
        src = source_of(_insert_logic_image_in_table)
        assert "picture_sink.add_picture(run" in src and "run.add_picture(" in src

    def test_sink_reads_the_file_before_touching_the_package(self):
        src = source_of(_PictureSink)
        body = src[src.index("def add_picture("):]
        assert body.index("Image.from_file(") < body.index("self._index_tail()") < body.index("relate_to(")
