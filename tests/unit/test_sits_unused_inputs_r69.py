"""R69 (input loss audit #8) — SITS reads only what it uses, and its 1.4 Reference lists only documents it read.

Before R69 SITS loaded the SDS summary, the SwUDS descriptions (a first time) and the HSIS signals, logged "loaded" and
wrote none of them anywhere; its 1.4 Reference listed the SwUDS and the STP whenever a path was given — a document that
yielded nothing (wrong template, empty file) was still named as a reference, and the input-document record said "열림".
"""
from __future__ import annotations

import pytest

from generators import sits as gsits
from generators import sts as gsts
from generators.suts import input_document_usable
from report_gen.generation_disclosures import build_disclosures

A_C = b"void B(void);\nvoid A(void)\n{\n    B();\n}\n"
B_C = b"void B(void)\n{\n}\n"


def _tree(tmp_path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    (tmp_path / "a.c").write_bytes(A_C)
    (tmp_path / "b.c").write_bytes(B_C)


def _empty_docx(path):
    from docx import Document
    d = Document()
    d.add_paragraph("표가 없는 문서")
    d.save(str(path))
    return str(path)


def _run(tmp_path, monkeypatch, **kw):
    """generate_sits on a two-module tree; returns (quality report, 1.4 Reference rows)."""
    seen = {}
    real = gsits.generate_sits_xlsm

    def capture(*a, **k):
        seen["refs"] = list((k.get("front_matter") or {}).get("references") or [])
        return real(*a, **k)

    monkeypatch.setattr(gsits, "generate_sits_xlsm", capture)
    _tree(tmp_path)
    out = gsits.generate_sits(source_root=str(tmp_path), output_path=str(tmp_path / "i.xlsm"), **kw)
    q = out.get("quality_report") or {}
    assert q, f"no flows — this test would be empty: {out.get('error')}"
    return q, seen.get("refs", [])


def test_the_loads_nothing_reads_are_gone(tmp_path, monkeypatch):
    # review round 2 W4: SITS reads no SwUDS description either — the record's `description` was never read
    calls = {"hsis": 0, "sds_summary": 0, "uds_desc": 0}

    def count(name, ret):
        def f(*a, **k):
            calls[name] += 1
            return ret
        return f

    monkeypatch.setattr(gsts, "_load_hsis_signals", count("hsis", {}))
    monkeypatch.setattr(gsts, "_load_sds_summary", count("sds_summary", ""))
    monkeypatch.setattr(gsts, "_load_uds_descriptions", count("uds_desc", {}))
    hsis = tmp_path / "x_HSIS.xlsx"
    hsis.write_bytes(b"")
    _run(tmp_path, monkeypatch, hsis_path=str(hsis), sds_docx_path=_empty_docx(tmp_path / "x_SDS.docx"),
         uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    assert calls == {"hsis": 0, "sds_summary": 0, "uds_desc": 0}, calls


def test_a_swuds_that_yields_no_function_is_not_a_reference(tmp_path, monkeypatch):
    q, refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    rec = q["input_documents"]["UDS"]
    assert rec["opened"] is True and rec["functions"] == 0 and "함수 0 개" in rec["reason"]
    assert not input_document_usable("UDS", rec)
    assert all(note != "SW 상세 설계서" for _f, note in refs), refs
    item = next(i for i in build_disclosures("sits", q) if i["key"] == "sits_input_documents")
    assert "SwUDS 못 읽음" in item["value"] and item["tone"] == "warning"


def test_a_swuds_that_yields_functions_is_a_reference_and_says_how_many(tmp_path, monkeypatch):
    monkeypatch.setattr(gsits, "load_uds_swcom_map", lambda p: {"A": ["SwCom_01"], "B": ["SwCom_02"]})
    q, refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    assert q["input_documents"]["UDS"]["functions"] == 2
    assert ("x_SwUDS.docx", "SW 상세 설계서") in refs
    item = next(i for i in build_disclosures("sits", q) if i["key"] == "sits_input_documents")
    assert "SwUDS 함수 2 (표 2 중 소스와 맞음)" in item["value"]


@pytest.mark.parametrize("body, listed", [(b"", False), ("시험 환경 SwITE_01 — HIL".encode("utf-8"), True)])
def test_an_stp_is_a_reference_only_when_its_text_was_read(tmp_path, monkeypatch, body, listed):
    stp = tmp_path / "x_STP.txt"
    stp.write_bytes(body)
    q, refs = _run(tmp_path, monkeypatch, stp_path=str(stp))
    rec = q["input_documents"]["STP"]
    assert (rec["read_chars"] > 0) is listed
    assert (("x_STP.txt", "SW 테스트 계획서") in refs) is listed
    if not listed:
        assert "읽은 글 0" in rec["reason"]
    else:
        item = next(i for i in build_disclosures("sits", q) if i["key"] == "sits_input_documents")
        assert f"STP 글 {rec['read_chars']}자" in item["value"]


def _headings_docx(path):
    """Meeting minutes — headings and paragraphs, no SwUDS table (review C1: the description loader keys headings)."""
    from docx import Document
    d = Document()
    for title in ("회의 안건", "결정 사항"):
        d.add_heading(title, level=1)
        d.add_paragraph("다음 주까지 정리한다.")
    d.save(str(path))
    return str(path)


def test_a_document_with_headings_only_is_not_a_swuds_reference(tmp_path, monkeypatch):
    q, refs = _run(tmp_path, monkeypatch, uds_path=_headings_docx(tmp_path / "minutes.docx"))
    rec = q["input_documents"]["UDS"]
    assert (rec["functions"], rec["functions_listed"]) == (0, 0)
    assert all(note != "SW 상세 설계서" for _f, note in refs), refs


def test_a_description_only_swuds_is_not_a_reference(tmp_path, monkeypatch):
    # review round 2 W4: even a description keyed by a source function name feeds no SITS cell
    monkeypatch.setattr(gsts, "_load_uds_descriptions", lambda p: {"A": "A 를 부른다", "baud_rate": "통신 속도"})
    q, refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    assert q["input_documents"]["UDS"]["functions"] == 0
    assert all(note != "SW 상세 설계서" for _f, note in refs), refs


def test_a_checked_table_name_alias_counts(tmp_path, monkeypatch):
    # review round 2 I-a: the design ID is found through a confirmed table-Name alias of the source name
    monkeypatch.setattr(gsits, "load_uds_design_ids", lambda p, source_functions=None: {
        "by_name": {"A_heading_typo": "SwFn_01"}, "aliases": {"A": "A_heading_typo"}, "aliases_checked": True})
    q, _refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    assert q["input_documents"]["UDS"]["functions"] == 1


def test_an_unmatched_swuds_is_not_copied_into_the_related_id_index(tmp_path, monkeypatch):
    # review round 2 W3: a table that matched no source function stayed in the "Related_ID 확인" SwUDS index
    monkeypatch.setattr(gsits, "load_uds_related_map", lambda p: {"other_fn": ["SwCom_09"], "gone_fn": ["SwCom_10"]})
    q, _refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    assert q["input_documents"]["UDS"]["functions"] == 0
    assert q["integration_flow_coverage"]["relid_index_rows"] == 0
    monkeypatch.setattr(gsits, "load_uds_related_map", lambda p: {"a": ["SwCom_01"], "gone_fn": ["SwCom_10"]})
    q, _refs = _run(tmp_path / "matched", monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    assert q["input_documents"]["UDS"]["functions"] == 1 and q["integration_flow_coverage"]["relid_index_rows"] == 2     # a used SwUDS: whole index


def test_a_non_docx_stp_says_it_cannot_be_opened(tmp_path, monkeypatch):
    stp = tmp_path / "x_STP.docx"
    stp.write_bytes(b"not a zip")
    q, _refs = _run(tmp_path, monkeypatch, stp_path=str(stp))
    assert "docx 가 아니다" in q["input_documents"]["STP"]["reason"]


def test_a_design_id_only_swuds_is_counted(tmp_path, monkeypatch):
    monkeypatch.setattr(gsits, "load_uds_design_ids", lambda p, source_functions=None: {"by_name": {"A": "SwFn_01"}})
    q, refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    assert q["input_documents"]["UDS"]["functions"] == 1
    assert ("x_SwUDS.docx", "SW 상세 설계서") in refs


def test_one_function_in_two_tables_with_different_case_is_counted_once(tmp_path, monkeypatch):
    monkeypatch.setattr(gsits, "load_uds_swcom_map", lambda p: {"a": ["SwCom_01"]})
    monkeypatch.setattr(gsits, "load_uds_design_ids", lambda p, source_functions=None: {"by_name": {"A": "SwFn_01"}})
    q, _refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    rec = q["input_documents"]["UDS"]
    assert (rec["functions"], rec["functions_listed"]) == (1, 1)


def test_a_swuds_whose_functions_match_no_source_function_is_not_a_reference(tmp_path, monkeypatch):
    # another project's or version's SwUDS: its table names functions this source does not have
    monkeypatch.setattr(gsits, "load_uds_swcom_map", lambda p: {"other_fn": ["SwCom_09"], "gone_fn": ["SwCom_10"]})
    q, refs = _run(tmp_path, monkeypatch, uds_path=_empty_docx(tmp_path / "x_SwUDS.docx"))
    rec = q["input_documents"]["UDS"]
    assert (rec["functions"], rec["functions_listed"]) == (0, 2) and "소스 함수와 맞는 것 0 개" in rec["reason"]
    assert all(note != "SW 상세 설계서" for _f, note in refs), refs


def test_an_stp_docx_without_strategy_sections_says_so(tmp_path, monkeypatch):
    q, refs = _run(tmp_path, monkeypatch, stp_path=_headings_docx(tmp_path / "x_STP.docx"))
    rec = q["input_documents"]["STP"]
    assert rec["read_chars"] == 0 and "키워드 제목" in rec["reason"]
    item = next(i for i in build_disclosures("sits", q) if i["key"] == "sits_input_documents")
    assert "STP 못 읽음" in item["value"]


def test_counted_zero_is_unusable_and_uncounted_means_opened():
    assert input_document_usable("UDS", {"opened": True})
    assert not input_document_usable("UDS", {"opened": True, "functions": 0})
    assert input_document_usable("UDS", {"opened": True, "functions": 3})
    assert not input_document_usable("STP", {"opened": True, "read_chars": 0})
    assert not input_document_usable("SRS", {"opened": True, "requirements": 0})
    assert not input_document_usable("UDS", {"opened": False, "functions": 3})
