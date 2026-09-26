"""R68 (N72 / N45 라벨 축) — 소스 단계 ASIL·Related 출처 라벨은 **값을 준 단계**가 단다.

실측(2026-09-19, `generate_uds_source_sections` 직접 호출):

| 프로젝트 | 함수 | override 가 ASIL 을 준 함수 | 그 라벨 | 값 없는 TBD | 그 라벨 |
|---|---|---|---|---|---|
| KJPDS02 | 1,146 | 228 (A 168 · QM 60) | 전부 `inference` | 918 | 전부 `inference` |
| PDS64 | 522 | 255 (A 192 · QM 63) | 전부 `inference` | 267 | 전부 `inference` |

값 사슬은 2026-04-09(`0b7e1979`) 부터 `comment > override > SDS > TBD` 였는데 라벨 식은 `comment > SDS > inference` 그대로라
override 가 준 값이 "추론" 으로, SDS 도 값이 있으면 override 값에 "SDS" 라벨이 붙었다(값과 라벨의 출처가 다름). 값이 없는 TBD 도
"추론" 이었다 — `provenance.unrecorded_source` 규약은 자리표시자를 `default`(0.30)로 둔다.

이 라운드는 생성기의 **값 사슬을 바꾸지 않는다**(아래 `TestValueChainUnchanged`). `override` 는 `inference` 와 같은 점수(0.60)·같은
약함 판정이라 override 가 준 값의 하류 덮어쓰기·신뢰도 점수는 종전과 같다. TBD 축의 점수는 원래 라벨과 무관했다 —
`validation._effective_src` 가 자리표시자를 `default` 로 재평가한다(리뷰 W2). 하류에서 딱 하나 바뀌는 집합은 "override 값 + SDS 도 값
있음" 행이다: 옛 라벨 `sds`(강함) 가 SwDS 덮어쓰기를 막았는데 이제 약함이라 덮인다(`TestDownstreamParity` 가 의도로 못박는다 —
확인되지 않은 스냅샷보다 설계 문서가 근거다). 라이브 실측(run 2111→2114) 그 집합은 1/1,146, 최종 값 변화 0. override 파일의 프로젝트
귀속(P7)은 별도 결정.
"""
from __future__ import annotations

import itertools
import json
from pathlib import Path

import pytest

from report_gen.uds_generator import _OVERRIDE_SOURCE, _source_stage_provenance

REPO = Path(__file__).resolve().parents[2]
OVERRIDE_PATH = REPO / "docs" / "uds_function_swcom_override.json"


# ─── 1. 헬퍼 — 값을 준 단계가 라벨을 단다 ──────────────────────────────────────────

class TestSourceStageProvenance:
    def _run(self, **kw):
        base = dict(comment_asil="", comment_related="", override=None, sds_asil="", sds_related="")
        base.update(kw)
        return _source_stage_provenance(**base)

    def test_comment_wins_and_is_labelled_comment(self):
        asil, asrc, rel, rsrc = self._run(
            comment_asil="C", comment_related="SwFn_01",
            override={"asil": "A", "related": "SwCom_01"}, sds_asil="QM", sds_related="SwFn_09")
        assert (asil, asrc) == ("C", "comment")
        assert (rel, rsrc) == ("SwFn_01", "comment")

    def test_override_value_is_labelled_override_not_inference(self):
        """**대표 케이스** — 실측 228/255건. 예전엔 `inference` 였다."""
        asil, asrc, rel, rsrc = self._run(override={"asil": "A", "related": "SwCom_01, SwSTR_01"})
        assert (asil, asrc) == ("A", "override")
        assert (rel, rsrc) == ("SwCom_01, SwSTR_01", "override")
        assert asrc == _OVERRIDE_SOURCE

    def test_override_value_is_not_labelled_sds_when_sds_also_has_one(self):
        """예전엔 값은 override 것인데 라벨은 `sds` 였다 — 값과 라벨의 출처가 달랐다."""
        asil, asrc, rel, rsrc = self._run(
            override={"asil": "A", "related": "SwCom_01"}, sds_asil="QM", sds_related="SwFn_09")
        assert (asil, asrc) == ("A", "override")
        assert (rel, rsrc) == ("SwCom_01", "override")

    def test_sds_when_no_comment_and_no_override(self):
        asil, asrc, rel, rsrc = self._run(sds_asil="QM", sds_related="SwFn_09")
        assert (asil, asrc) == ("QM", "sds")
        assert (rel, rsrc) == ("SwFn_09", "sds")

    def test_nothing_is_tbd_with_default_not_inference(self):
        """값이 없으면 근거도 없다 — 실측 918건이 `inference`(0.60) 였다. `default`(0.30) 가 사실이다."""
        asil, asrc, rel, rsrc = self._run()
        assert (asil, asrc) == ("TBD", "default")
        assert (rel, rsrc) == ("TBD", "default")

    def test_override_null_asil_falls_through(self):
        """override JSON 엔 `asil: null` 항목이 1건 있다 — 빈 값은 단계를 건너뛴다(라벨도 함께)."""
        asil, asrc, *_ = self._run(override={"asil": None, "related": "SwCom_01"}, sds_asil="QM")
        assert (asil, asrc) == ("QM", "sds")
        asil, asrc, *_ = self._run(override={"asil": None})
        assert (asil, asrc) == ("TBD", "default")

    def test_non_string_override_value_is_stringified(self):
        """옛 식은 원값을 그대로 담았고 새 식은 `str()` 을 건다 — 대조군(`_legacy_values`)이 `str()` 을 걸어 구조적으로 못 보는
        차이라 여기서 못박는다(리뷰 I2). 현 JSON 은 전부 문자열이라 산출물엔 차이가 없다."""
        asil, asrc, *_ = self._run(override={"asil": 4})
        assert (asil, asrc) == ("4", "override")
        assert isinstance(asil, str)

    def test_override_not_a_dict_is_ignored(self):
        asil, asrc, rel, rsrc = self._run(override="garbage", sds_asil="B")
        assert (asil, asrc) == ("B", "sds")
        assert (rel, rsrc) == ("TBD", "default")

    def test_single_chain_no_loop_specific_switch(self):
        """(R70 N84) 사슬은 하나뿐이다 — R68 의 `override_related=False`(텍스트 폴백 루프 전용 Related 사슬)는 없어졌다.
        같은 함수가 어느 루프로 가느냐에 따라 Related 가 달라지던 두 번째 규칙이 다시 생기면 여기서 잡힌다."""
        import inspect

        params = inspect.signature(_source_stage_provenance).parameters
        assert set(params) == {"comment_asil", "comment_related", "override", "sds_asil", "sds_related"}, sorted(params)
        asil, asrc, rel, rsrc = self._run(override={"asil": "A", "related": "SwCom_01"}, sds_related="SwFn_09")
        assert (rel, rsrc) == ("SwCom_01", "override"), "Related 사슬에도 override 가 SDS 보다 앞선다"

    def test_labels_are_first_class_vocabulary(self):
        """새 라벨은 점수표·화면 라벨·입력 사슬 표 세 곳에 다 있어야 한다 — 없으면 `unknown` 으로 조용히 접힌다."""
        from backend.services import docgen_field_sources as fs

        src = (REPO / "report_gen" / "validation.py").read_text(encoding="utf-8")
        for label in ("comment", "override", "sds", "default"):
            assert f'"{label}":' in src, f"{label} 이 validation.py 어휘에 없다"
        assert "override" in fs.FIELD_SOURCES["asil"] and "override" in fs.FIELD_SOURCES["related"]
        assert fs.LABEL_TO_SOURCE["함수 override 맵(정본 역추출 스냅샷)"] == "override"


# ─── 2. 값 사슬은 그대로 — 라벨만 바뀐다 ────────────────────────────────────────────

def _legacy_values(comment_asil, comment_related, override, sds_asil, sds_related):
    """2026-04-09 ~ R67 의 AST 루프 값 식(uds_generator.py 옛 1695·1696행)을 그대로 옮긴 대조군.
    (R70 N84) 텍스트 폴백 루프도 이제 이 식을 쓴다 — 옛 폴백 식(`comment or SDS or TBD`)은 대조군에서 뺐다."""
    ovr = override if isinstance(override, dict) else {}
    asil = comment_asil or (ovr.get("asil") if ovr else "") or sds_asil or "TBD"
    related = comment_related or (ovr.get("related") if ovr else "") or sds_related or "TBD"
    return str(asil), str(related)


class TestValueChainUnchanged:
    _ASIL = ["", "C"]
    _REL = ["", "SwFn_01"]
    _OVR = [None, {}, {"asil": "A", "related": "SwCom_01"}, {"asil": None, "related": "SwCom_02"}, {"asil": "QM"}]
    _SDS_A = ["", "QM"]
    _SDS_R = ["", "SwFn_09"]

    def test_values_equal_legacy_chain_for_every_combination(self):
        """80가지 조합 — 값이 하나라도 달라지면 라벨 정직화가 아니라 내용 변경이다."""
        for ca, cr, ov, sa, sr in itertools.product(self._ASIL, self._REL, self._OVR, self._SDS_A, self._SDS_R):
            asil, _, rel, _ = _source_stage_provenance(
                comment_asil=ca, comment_related=cr, override=ov, sds_asil=sa, sds_related=sr)
            assert (asil, rel) == _legacy_values(ca, cr, ov, sa, sr), (ca, cr, ov, sa, sr)

    def test_label_names_the_stage_that_gave_the_value(self):
        """라벨 ↔ 값 결합: 라벨이 가리키는 단계의 값과 실제 값이 같아야 한다(값과 라벨의 출처 분리 금지)."""
        for ca, cr, ov, sa, sr in itertools.product(self._ASIL, self._REL, self._OVR, self._SDS_A, self._SDS_R):
            asil, asrc, rel, rsrc = _source_stage_provenance(
                comment_asil=ca, comment_related=cr, override=ov, sds_asil=sa, sds_related=sr)
            ovr = ov if isinstance(ov, dict) else {}
            expect = {"comment": (ca, cr), "override": (ovr.get("asil"), ovr.get("related")),
                      "sds": (sa, sr), "default": ("TBD", "TBD")}
            assert asil == str(expect[asrc][0]), (asrc, asil, ca, ov, sa)
            assert rel == str(expect[rsrc][1]), (rsrc, rel, cr, ov, sr)


# ─── 3. 어휘 — 덮어쓰기·신뢰·게이트 판정은 `inference` 와 같다 ─────────────────────────

class TestVocabularyParityWithInference:
    def test_override_is_a_weak_source(self):
        """약해야 SwDS(`docx_builder`)·정본이 종전대로 덮는다 — 강하게 두면 KJPDS02 스냅샷이 다른 프로젝트에 굳는다."""
        from report_gen.provenance import WEAK_SOURCES, is_weak_source

        assert "override" in WEAK_SOURCES
        assert is_weak_source("override") is True
        assert is_weak_source("OVERRIDE ") is True

    def test_reference_suds_may_override_it(self):
        from report_gen.provenance import reference_suds_may_override

        assert reference_suds_may_override("override") is True
        assert reference_suds_may_override("reference") is False, "대조군 — 정본 자신은 안 덮인다"

    def test_fill_rate_gate_does_not_trust_it(self):
        """채움률 게이트(`_normalize_field_source`)는 6개 라벨만 신뢰한다 — override 는 `inference` 로 접혀 불신, 종전과 같다."""
        from backend.helpers.common import _normalize_field_source

        assert _normalize_field_source("override") == "inference"

    def test_score_equals_inference(self):
        """점수가 같아야 신뢰도 평균이 안 움직인다(라벨만 바꾸는 라운드). 올리려면 P7 결정이 먼저다."""
        import ast

        tree = ast.parse((REPO / "report_gen" / "validation.py").read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Dict) and node.keys and all(
                isinstance(v, ast.Constant) and isinstance(v.value, (int, float)) for v in node.values
            ):
                keys = {k.value for k in node.keys if isinstance(k, ast.Constant)}
                if {"comment", "inference", "override"} <= keys:
                    table = {k.value: v.value for k, v in zip(node.keys, node.values, strict=True)}
                    assert table["override"] == table["inference"]
                    return
        raise AssertionError("src_score 표에서 override 를 찾지 못했다")

    def test_confidence_report_names_the_label_and_its_evidence(self, tmp_path):
        """신뢰도 사이드카가 새 라벨을 어휘로 알아야 한다 — 모르면 `미상(분류 불가)` 로 접히고 드리프트 카운터만 오른다.
        근거 문구는 파일 출처와 "이 프로젝트의 값인지 확인되지 않았다" 를 말해야 한다(X8 — 정본 그 자체로 읽히면 안 된다)."""
        from report_gen.validation import generate_asil_related_confidence_report

        out = tmp_path / "conf.md"
        generate_asil_related_confidence_report({"function_details": {"F": {
            "name": "f", "asil": "A", "asil_source": "override", "description": "d", "description_source": "inference",
            "related": "SwCom_01", "related_source": "override"}}}, str(out))
        text = out.read_text(encoding="utf-8")
        assert "함수 override 맵(정본 역추출 스냅샷): `1` / `1`" in text, text[:1500]
        assert "docs/uds_function_swcom_override.json" in text
        assert "확인되지 않았다" in text
        assert "미상(분류 불가): `1`" not in text

    def test_docgen_field_sources_treats_it_as_ungrounded(self):
        """저장소 파일은 사용자 입력이 아니다 — 게이트가 '확보' 로 세면 안 된다."""
        from backend.services import docgen_field_sources as fs

        assert fs.source_input("override") is None
        rows = {r["source"]: r for r in fs.chain_state("asil", {})}
        assert rows["override"]["grounded"] is False
        # 사이드카 분포 줄(불릿 제거 후)이 코드로 되돌아와야 귀속이 된다 — 안 되면 라벨 문자열이 키로 남는다.
        assert fs.parse_source_distribution(["함수 override 맵(정본 역추출 스냅샷): `12` / `40` (30.0%)"]) == {"override": 12}


# ─── 4. 생성기 통합 — 진짜 소스 트리에서 라벨이 붙는다 ─────────────────────────────────

@pytest.fixture(scope="module")
def sections(tmp_path_factory):
    """`main` 은 저장소 override 에 ASIL·Related 가 있다(2026-04-09 이래). 주석 없는 함수 하나·`@asil` 주석 함수 하나를 곁들인다."""
    ovr = json.loads(OVERRIDE_PATH.read_text(encoding="utf-8"))
    if not (ovr.get("main") or {}).get("asil"):
        pytest.skip("override JSON 에 main/asil 이 없다 — 전제가 바뀌었다")
    root = tmp_path_factory.mktemp("r68src")
    (root / "app.c").write_text(
        "#include <stdint.h>\n"
        "static uint8_t g_cnt;\n"
        "void main(void)\n{\n    g_cnt++;\n}\n"
        "void r68_no_override_fn(void)\n{\n    g_cnt = 0u;\n}\n"
        "/**\n * @brief r68 commented\n * @asil C\n */\n"
        "void r68_commented_fn(void)\n{\n    g_cnt = 1u;\n}\n",
        encoding="utf-8")
    from report_gen.uds_generator import generate_uds_source_sections

    out = generate_uds_source_sections(str(root))
    by_name = {str(i.get("name")): i for i in (out.get("function_details") or {}).values() if isinstance(i, dict)}
    return ovr, by_name


class TestGeneratedSections:
    def test_override_fed_function_is_labelled_override(self, sections):
        ovr, by_name = sections
        info = by_name["main"]
        assert info["asil"] == ovr["main"]["asil"]
        assert info["asil_source"] == "override"
        assert info["related"] == ovr["main"]["related"]
        assert info["related_source"] == "override"

    def test_function_without_any_evidence_is_default(self, sections):
        _, by_name = sections
        info = by_name["r68_no_override_fn"]
        assert (info["asil"], info["asil_source"]) == ("TBD", "default")
        assert (info["related"], info["related_source"]) == ("TBD", "default")

    def test_comment_asil_still_wins(self, sections):
        _, by_name = sections
        info = by_name["r68_commented_fn"]
        assert (info["asil"], info["asil_source"]) == ("C", "comment")

    def test_fallback_loop_uses_the_same_chain(self, tmp_path, monkeypatch):
        """AST 가 놓친 함수는 텍스트 폴백 루프로 간다(R62 기법: `parse_c_project` 결과에서 빼낸다). (R70 N84) 그 루프의
        Related 사슬에도 override 가 선다 — R68 까지는 ASIL 만 override 를 봐 같은 함수가 루프에 따라 다른 Related 를 받았다."""
        import workflow.code_parser as pkg
        from report_gen.uds_generator import generate_uds_source_sections

        ovr = json.loads(OVERRIDE_PATH.read_text(encoding="utf-8"))
        if not ((ovr.get("main") or {}).get("asil") and (ovr.get("main") or {}).get("related")):
            pytest.skip("override JSON 에 main 의 asil/related 가 없다")
        (tmp_path / "fb.c").write_text(
            "static int g;\nvoid main(void)\n{\n    g = 1;\n}\nvoid r68_seen(void)\n{\n    g = 2;\n}\n", encoding="utf-8")
        real = pkg.parse_c_project

        def _without_main(*args, **kwargs):
            res = real(*args, **kwargs)
            res["functions"] = [f for f in res["functions"] if f.get("name") != "main"]
            return res

        monkeypatch.setattr(pkg, "parse_c_project", _without_main)
        sec = generate_uds_source_sections(str(tmp_path))
        by_name = {v["name"]: v for v in sec["function_details"].values() if isinstance(v, dict)}
        info = by_name["main"]
        assert (info["asil"], info["asil_source"]) == (ovr["main"]["asil"], "override")
        assert (info["related"], info["related_source"]) == (ovr["main"]["related"], "override")

    def test_no_function_carries_inference_for_asil_or_related(self, sections):
        """소스 단계는 ASIL·Related 를 추론하지 않는다 — 그 두 축에 `inference` 가 남아 있으면 옛 식이 되살아난 것이다."""
        _, by_name = sections
        leftovers = [(n, i.get("asil_source"), i.get("related_source")) for n, i in by_name.items()
                     if "inference" in (i.get("asil_source"), i.get("related_source"))]
        assert not leftovers, leftovers


class TestCacheSchemaMoved:
    def test_schema_version_is_at_least_v23(self):
        """소스 단계 라벨은 캐시 payload 안에 있다 — 버전을 안 올리면 소스가 안 바뀐 프로젝트에서 옛 `inference` 라벨이
        캐시 수명 동안 그대로 나온다(리뷰 C1, v12·v16~v22 와 같은 실패 모드)."""
        from backend.helpers.uds import _SOURCE_SECTIONS_SCHEMA_VERSION

        assert int(_SOURCE_SECTIONS_SCHEMA_VERSION.lstrip("v")) >= 23


# ─── 5. 하류 정합 — docx_builder 의 덮어쓰기는 `inference` 때와 같다 ────────────────────

@pytest.fixture
def build(tmp_path, monkeypatch):
    import docx

    import config
    from report_gen import docx_builder

    monkeypatch.setattr(config, "resolve_uds_template_path", lambda: "", raising=False)
    monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(tmp_path / "none.docx"), raising=False)

    def _run(sds_map=None, **overrides):
        tpl = tmp_path / "t.docx"
        d = docx.Document()
        d.add_heading("Software Unit Design", level=1)
        d.add_heading("SwUFn_0001: alpha", level=4)
        d.save(str(tpl))
        info = {"id": "SwUFn_0001", "name": "alpha", "prototype": "void alpha(void);",
                "description": "", "asil": "A", "asil_source": "override",
                "related": "SwCom_01", "related_source": "override", "precondition": "",
                "inputs": [], "outputs": [], "globals_global": [], "globals_static": [],
                "called": "", "logic": ""}
        info.update(overrides)
        payload = {"project_name": "T", "overview": "o", "requirements": "r", "interfaces": "i",
                   "uds_frames": "u", "notes": "n", "function_details": {"SwUFn_0001": info}}
        if sds_map is not None:
            payload["sds_partition_map"] = sds_map
        docx_builder.generate_uds_docx(str(tpl), payload, str(tmp_path / "out.docx"))
        return info

    return _run


class TestDownstreamParity:
    def test_comment_evidence_overwrites_override(self, build):
        info = build(comment_asil="D", comment_related="SwFn_09")
        assert (info["asil"], info["asil_source"]) == ("D", "comment")
        assert (info["related"], info["related_source"]) == ("SwFn_09", "comment")

    def test_swds_overwrites_override_like_it_overwrote_inference(self, build):
        """**하류에서 유일하게 거동이 바뀌는 집합**(리뷰 W1) — override 값 + SwDS 도 값 있음.
        옛 라벨 `sds`(강함) 는 이 덮어쓰기를 막았고 옛 라벨 `inference`(약함) 는 허용했다. 이제 override 는 약함이라 덮인다 —
        확인되지 않은 스냅샷보다 설계 문서가 근거이므로 **의도**다. 되돌리려면 이 테스트를 바꿔야 한다(라이브 실측 1/1,146)."""
        sds = {"alpha": {"asil": "QM", "related": "SwFn_09", "description": ""}}
        info = build(sds_map=sds)
        assert (info["asil"], info["asil_source"]) == ("QM", "sds")
        assert (info["related"], info["related_source"]) == ("SwFn_09", "sds")
        # 대조군 — 옛 `inference` 라벨도 같은 결과였다(override = inference 급).
        ctrl = build(sds_map=sds, asil_source="inference", related_source="inference")
        assert (ctrl["asil"], ctrl["asil_source"]) == ("QM", "sds")
        # 대조군 — 강한 라벨은 안 덮인다(약함 판정이 결정한다는 것을 보인다).
        strong = build(sds_map=sds, asil_source="comment", related_source="comment")
        assert (strong["asil"], strong["asil_source"]) == ("A", "comment")

    def test_module_seed_rank_treats_override_like_inference(self):
        """모듈 상속 씨앗 순위: override 는 inference 와 같은 3 이어야 SwDS(2)·정본(1)·주석(0) 씨앗이 이긴다."""
        import ast

        src = (REPO / "report_gen" / "docx_builder.py").read_text(encoding="utf-8")
        tree = ast.parse(src)
        for node in ast.walk(tree):
            if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "_SEED_RANK" for t in node.targets):
                table = ast.literal_eval(node.value)
                assert table["override"] == 3 and table["sds"] < table["override"]
                assert "inference" not in table or table["inference"] == table["override"]
                return
        raise AssertionError("_SEED_RANK 를 찾지 못했다")

    def test_override_survives_when_nothing_better_comes(self, build):
        """덮을 근거가 없으면 값도 라벨도 남는다 — 예전 `inference` 라벨이 남던 자리에 이제 사실이 남는다."""
        info = build()
        assert (info["asil"], info["asil_source"]) == ("A", "override")
        assert (info["related"], info["related_source"]) == ("SwCom_01", "override")
