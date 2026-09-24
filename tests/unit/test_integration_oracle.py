"""(R16, 격차 ②) Interprocedural source oracle — callees with a project definition run for real (SITS expectations).

Each case is a small project of translation units built in memory; the unit oracle (no provider) is the control.
"""
from __future__ import annotations

import os

import pytest

from generators import c_project_context as cpc
from generators.c_source_oracle import Unsupported, evaluate_outputs
from generators.integration_oracle import CalleeProvider, attach_integration_evidence
from generators.test_evidence import VERIFY_PREFIX

ROOT = os.path.join(os.sep, "virt_r16")
COMMON = """#ifndef COMMON_H
#define COMMON_H
typedef unsigned char U8;
typedef unsigned int U16;
typedef signed int S16;
extern U8 g_in;
extern U8 g_out;
extern U16 g_sum;
extern S16 g_s;
U8 b_scale(U8 x);
void b_store(U8 v);
U8 b_branch(U8 v);
U8 b_loop(U8 n);
U8 b_goto(U8 v);
void b_ptr(U8 *p);
void b_ext(void);
S16 b_over(S16 v);
U8 b_rec(U8 n);
U8 a_back(U8 n);
U8 b_static_peer(void);
U8 dup(void);
#endif
"""
A = """#include "common.h"
U8 g_in;
U8 g_out;
U16 g_sum;
S16 g_s;
static U8 s_a;
static U8 s_arr[2];
void entry(U8 k) {
    U8 r = b_scale(k);
    b_store(r);
    s_a = b_branch(g_in);
    g_out = (U8)(s_a + 1U);
}
void narrow(U16 w) { g_out = b_scale((U8)w); }
void convert(U16 w) { g_out = b_scale(w); }
void forked(void) { g_out = b_branch(g_in); }
void loops(void) { g_out = b_loop(g_in); }
void jumps(void) { g_out = b_goto(g_in); g_sum = 7U; }
void escapes(void) { U8 t = 5U; s_arr[0] = 1U; b_ptr(&t); b_ptr(s_arr); g_out = t; g_in = s_arr[0]; }
void statics(void) { static U8 cnt = 0U; cnt = 3U; b_ext(); g_out = cnt; }
void overflow(void) { g_s = b_over(g_s); }
void recursion(void) { g_out = b_rec(g_in); }
U8 a_back(U8 n) { return b_rec(n); }
void peer(void) { g_out = b_static_peer(); }
void ambiguous(void) { g_out = dup(); }
void stubbed(void) { g_out = b_scale(g_in); }
"""
B = """#include "common.h"
static U8 s_b;
U8 b_scale(U8 x) { return (U8)(x * 2U); }
void b_store(U8 v) { g_sum = (U16)(g_sum + v); s_b = v; }
U8 b_branch(U8 v) { if (v > 10U) { return 3U; } return 4U; }
U8 b_loop(U8 n) { U16 i; U16 j; U16 acc = 0U; for (i = 0U; i < 4000U; i++) { for (j = 0U; j < 40U; j++) { acc++; } }
    return (U8)(acc + n); }
U8 b_goto(U8 v) { if (v) { goto done; } g_sum = 1U; done: return v; }
void b_ptr(U8 *p) { *p = 9U; }
void b_ext(void) { lib_call(); }
S16 b_over(S16 v) { return (S16)(v + 32767); }
U8 b_rec(U8 n) { if (n == 0U) { return 0U; } return a_back((U8)(n - 1U)); }
static U8 hidden(void) { return 1U; }
U8 dup(void) { return 1U; }
"""
C = """#include "common.h"
static U8 hidden(void) { return 2U; }
U8 b_static_peer(void) { return 6U; }
U8 dup(void) { return 2U; }
"""


def _project(extra=None):
    files = {os.path.join(ROOT, "common.h"): COMMON, os.path.join(ROOT, "a.c"): A,
             os.path.join(ROOT, "b.c"): B, os.path.join(ROOT, "c.c"): C, **(extra or {})}
    ctx = cpc.build_project_context(files)
    units = [p for p in files if p.endswith(".c")]
    scopes = cpc.build_scopes(ctx, units)
    return files, ctx, scopes


def _eval(name, inputs, outputs, provider=True, extra=None):
    files, ctx, scopes = _project(extra)
    a = os.path.join(ROOT, "a.c")
    unit = {"name": name, "source_text": files[a], "source_path": a, "source_text_complete": True,
            "project_scope": scopes[a]}
    if provider:
        unit["callee_provider"] = CalleeProvider(ctx, files, scopes, cpc.shared_parser())
    return evaluate_outputs(unit, [inputs], [list(outputs)])[0]


def _value(record, name):
    return (record["outputs"].get(name) or {}).get("value")


def _reason(record, name):
    return (record["outputs"].get(name) or {}).get("reason", "")


def test_callee_bodies_run_across_units():
    r = _eval("entry", {"k": 5, "g_in": 20, "g_sum": 100}, ["g_out", "g_sum", "s_b"])
    assert r["status"] == "supported"
    assert (_value(r, "g_out"), _value(r, "g_sum"), _value(r, "s_b")) == (4, 110, 10)
    assert r["interprocedural"]["inlined"] == {"b_branch": 1, "b_scale": 1, "b_store": 1}
    assert r["oracle"] == "project_context_source_interprocedural"
    assert any("interpreted with the argument values" in a for a in r["assumptions"])


def test_without_provider_callees_stay_effects_only():
    r = _eval("entry", {"k": 5, "g_in": 20, "g_sum": 100}, ["g_out", "g_sum"], provider=False)
    assert _value(r, "g_out") is None and _value(r, "g_sum") is None
    assert "interprocedural" not in r


def test_argument_converts_to_parameter_type():
    assert _value(_eval("convert", {"w": 300}, ["g_out"]), "g_out") == 88   # (U8)300 = 44, * 2
    assert _value(_eval("narrow", {"w": 200}, ["g_out"]), "g_out") == 144   # (U8)(200 * 2)


def test_callee_paths_on_unknown_input_join_to_unknown():
    r = _eval("forked", {}, ["g_out"])
    assert _value(r, "g_out") is None
    assert _reason(r, "g_out").startswith("path_dependent_in_callee:b_branch")
    assert _value(_eval("forked", {"g_in": 11}, ["g_out"]), "g_out") == 3


def test_a_global_written_differently_on_callee_paths_is_unknown():
    extra_b = B + "void b_set(U8 v) { if (v > 10U) { g_sum = 1U; } else { g_sum = 2U; } g_s = 4; }\n"
    extra_a = A + "void b_set(U8 v);\nvoid sets(void) { b_set(g_in); }\n"
    files, ctx, scopes = _project({os.path.join(ROOT, "b.c"): extra_b, os.path.join(ROOT, "a.c"): extra_a})
    unit = {"name": "sets", "source_text": extra_a, "source_path": os.path.join(ROOT, "a.c"), "source_text_complete": True,
            "project_scope": scopes[os.path.join(ROOT, "a.c")],
            "callee_provider": CalleeProvider(ctx, files, scopes, cpc.shared_parser())}
    r = evaluate_outputs(unit, [{}, {"g_in": 3}], [["g_sum", "g_s"]] * 2)
    assert _value(r[0], "g_sum") is None and _reason(r[0], "g_sum") == "path_dependent_in_callee:b_set"
    assert _value(r[0], "g_s") == 4   # both paths agree
    assert _value(r[1], "g_sum") == 2


def test_callee_over_its_budget_share_falls_back_to_effects():
    r = _eval("loops", {"g_in": 1}, ["g_out"])
    assert r["status"] == "supported"
    assert _value(r, "g_out") is None
    assert r["interprocedural"]["not_inlined"].get("execution_budget") == 1


def test_unsupported_callee_falls_back_and_caller_continues():
    r = _eval("jumps", {"g_in": 1}, ["g_out", "g_sum"])
    assert _value(r, "g_out") is None and _value(r, "g_sum") == 7
    assert r["interprocedural"]["failed"]["b_goto"] == "goto_unmodeled"


def test_pointer_writes_reach_caller_locals_and_arrays_of_another_unit():
    r = _eval("escapes", {}, ["g_out", "g_in"])
    # b_ptr's unit cannot name ``s_arr`` (static in a.c): the caller havocs it in its own terms
    assert _value(r, "g_out") is None and _value(r, "g_in") is None


def test_unknown_code_in_callee_havocs_caller_static_locals():
    r = _eval("statics", {}, ["g_out"])
    assert _value(r, "g_out") is None


def test_undefined_behaviour_proven_in_callee_refuses_the_run():
    r = _eval("overflow", {"g_s": 5}, ["g_s"])
    assert r["status"] == "unsupported" and r["reason"].startswith("undefined_behavior:")


def test_recursion_is_not_interpreted():
    r = _eval("recursion", {"g_in": 2}, ["g_out"])
    assert _value(r, "g_out") is None
    assert r["interprocedural"]["not_inlined"].get("recursion", 0) >= 1


def test_static_function_of_another_unit_is_not_the_callee():
    files, ctx, scopes = _project()
    prov = CalleeProvider(ctx, files, scopes, cpc.shared_parser())
    assert prov.definition("b_static_peer", os.path.join(ROOT, "a.c"))[4].endswith("c.c")
    with pytest.raises(Unsupported, match=r"^no_external_definition$"):
        prov.definition("hidden", os.path.join(ROOT, "a.c"))
    assert prov.definition("hidden", os.path.join(ROOT, "b.c"))[4].endswith("b.c")  # its own unit names it


def test_two_external_definitions_are_not_guessed():
    r = _eval("ambiguous", {}, ["g_out"])
    assert _value(r, "g_out") is None
    assert r["interprocedural"]["failed"]["dup"] == "definition:definition_ambiguous"


def test_same_named_static_objects_of_two_units_do_not_alias():
    extra = {os.path.join(ROOT, "d.c"): '#include "common.h"\nstatic U8 s_a;\nU8 d_get(void) { s_a = 9U; return s_a; }\n'}
    files, ctx, scopes = _project(extra)
    a = os.path.join(ROOT, "a.c")
    unit = {"name": "peer", "source_path": a, "source_text": files[a].replace("b_static_peer()", "d_get()")
            .replace("void peer(void)", "U8 d_get(void);\nvoid peer(void)"), "source_text_complete": True}
    files[a] = unit["source_text"]
    ctx = cpc.build_project_context(files)
    scopes = cpc.build_scopes(ctx, [p for p in files if p.endswith(".c")])
    unit["project_scope"] = scopes[a]
    unit["callee_provider"] = CalleeProvider(ctx, files, scopes, cpc.shared_parser())
    r = evaluate_outputs(unit, [{"s_a": 1}], [["g_out", "s_a"]])[0]
    assert r["interprocedural"]["failed"]["d_get"] == "linkage_collision:s_a"
    # interpreted, d_get's ``s_a = 9U`` would have landed in a.c's s_a; the fallback write closure is by name, so it
    # conservatively havocs it instead — never the other unit's value
    assert _value(r, "s_a") is None


def test_sequence_stub_wins_over_the_body():
    r = _eval("stubbed", {"g_in": 3, "b_scale() return": 50}, ["g_out"])
    assert _value(r, "g_out") == 50 and r.get("stubs") == ["b_scale"]
    assert "b_scale" not in r["interprocedural"]["inlined"]


def test_attach_integration_evidence_writes_values_and_discloses():
    files, ctx, _ = _project()
    a = os.path.join(ROOT, "a.c")
    details = {"f1": {"name": "entry", "source_path": a}, "f2": {"name": "dup", "source_path": os.path.join(ROOT, "b.c")},
               "f3": {"name": "dup", "source_path": os.path.join(ROOT, "c.c")}}
    subs = [{"inputs": {"k": 5, "g_in": 20, "g_sum": 100}, "expected": {"g_out": f"{VERIFY_PREFIX} x", "g_sum": "N/A"}},
            {"inputs": {"k": 5}, "expected": {"g_out": f"{VERIFY_PREFIX} x", "g_sum": "N/A"}}]
    itcs = [{"tc_id": "T1", "entry_fn": "entry", "sub_cases": subs},
            {"tc_id": "T2", "entry_fn": "dup", "sub_cases": [{"inputs": {}, "expected": {"g_out": "N/A"}}]},
            {"tc_id": "T3", "entry_fn": "nowhere", "sub_cases": [{"inputs": {}, "expected": {"g_out": "N/A"}}]}]
    stats = attach_integration_evidence(itcs, {"project_context": ctx, "source_files": files}, details)
    assert subs[0]["expected"] == {"g_out": 4, "g_sum": 110}
    assert subs[0]["expected_evidence"]["g_out"]["callees_interpreted"] == ["b_branch", "b_scale", "b_store"]
    assert str(subs[1]["expected"]["g_sum"]).startswith(VERIFY_PREFIX)  # g_sum's entry value was not set
    assert stats["status"] == "evaluated" and stats["tc_evaluated"] == 1 and stats["tc_with_derived"] == 1
    assert stats["tc_skipped"] == {"entry_definition_ambiguous": 1, "entry_definition_missing": 1}
    assert str(subs[1]["expected"]["g_out"]).startswith(VERIFY_PREFIX)  # g_in not set: b_branch's paths disagree
    assert stats["cells"] == 4 and stats["derived"] == 2 and stats["derived_assigned"] == 2
    assert stats["unknown_reasons"] == {"path_dependent_in_callee": 1, "initial_value_not_in_inputs": 1}
    assert stats["inlined_calls"] == 6
    assert str(itcs[1]["sub_cases"][0]["expected"]["g_out"]) == "N/A"   # untouched when not evaluated


def test_attach_without_project_context_says_so():
    assert attach_integration_evidence([], {}, {})["status"] == "no_project_context"


HEAVY = """#include "common.h"
U8 g_in;
U8 g_out;
U16 g_sum;
S16 g_s;
void heavy(void) { U16 i; U16 j; for (i = 0U; i < 4000U; i++) { for (j = 0U; j < 4000U; j++) { g_sum++; } } g_out = 1U; }
"""


def test_a_flow_over_the_step_budget_in_its_first_sub_cases_is_not_run_further():
    a = os.path.join(ROOT, "a.c")
    files, ctx, _ = _project({a: HEAVY})
    subs = [{"inputs": {"g_sum": k}, "expected": {"g_out": "N/A"}} for k in range(4)]
    itcs = [{"tc_id": "T1", "entry_fn": "heavy", "sub_cases": subs}]
    stats = attach_integration_evidence(itcs, {"project_context": ctx, "source_files": files},
                                        {"f": {"name": "heavy", "source_path": a}})
    assert stats["tc_budget_cut"] == 1
    assert [s["expected_evidence"]["g_out"]["reason"] for s in subs[:2]] == ["execution_budget", "execution_budget"]
    assert all(s["expected_evidence"]["g_out"]["reason"] == "execution_budget:not_run_after_2_sub_cases_exceeded_it"
               for s in subs[2:])
    assert all(str(s["expected"]["g_out"]).startswith(VERIFY_PREFIX) for s in subs)


def test_evidence_sheet_lists_every_cell_with_its_basis_or_reason():
    import openpyxl

    from generators.integration_oracle import TEST_EVIDENCE_HEADERS, write_integration_evidence_sheet
    files, ctx, _ = _project()
    a = os.path.join(ROOT, "a.c")
    subs = [{"case_num": 1, "inputs": {"k": 5, "g_in": 20, "g_sum": 100}, "expected": {"g_out": "N/A", "g_sum": "N/A"}},
            {"case_num": 2, "inputs": {"k": 5}, "expected": {"g_out": "N/A"}}]
    itcs = [{"tc_id": "T1", "entry_fn": "entry", "sub_cases": subs}]
    attach_integration_evidence(itcs, {"project_context": ctx, "source_files": files},
                                {"f": {"name": "entry", "source_path": a}})
    wb = openpyxl.Workbook()
    assert write_integration_evidence_sheet(wb, itcs) == 3
    rows = list(wb["Test Evidence"].iter_rows(values_only=True))
    assert list(rows[0]) == TEST_EVIDENCE_HEADERS
    first = dict(zip(TEST_EVIDENCE_HEADERS, rows[1], strict=True))
    assert (first["Observable"], first["Expected"], first["Status"], first["Basis"]) == ("g_out", 4, "derived", "assigned")
    assert first["Callees interpreted"] == "b_branch, b_scale, b_store"
    last = dict(zip(TEST_EVIDENCE_HEADERS, rows[3], strict=True))
    assert last["Status"] == "unknown" and last["Reason"].startswith("path_dependent_in_callee")
    empty = openpyxl.Workbook()
    assert write_integration_evidence_sheet(empty, [{"tc_id": "T", "sub_cases": [{"expected": {"x": "N/A"}}]}]) == 0
    assert "Test Evidence" not in empty.sheetnames   # nothing evaluated: no sheet claiming evidence


def test_quality_report_and_disclosure_carry_the_block():
    from generators.sits import generate_sits_quality_report
    from report_gen.generation_disclosures import build_disclosures
    block = {"status": "evaluated", "tc_total": 120, "tc_with_derived": 63, "cells": 48938, "derived": 10479,
             "derived_assigned": 5511, "derived_unchanged_input": 4968, "tc_skipped": {},
             "unknown_reasons": {"observable_form_unmodeled": 8022}, "unknown_reason_kinds": 1,
             "not_inlined": {"path_budget": 7021}, "tc_budget_cut": 1}
    qr = generate_sits_quality_report([], 0, flow_stats={"integration_oracle": block})
    assert qr["integration_oracle"] is block
    item = {i["key"]: i for i in build_disclosures("sits", qr)}["sits_integration_oracle"]
    assert item["value"] == "10479 / 48938칸 (흐름이 쓴 값 5511) · 값이 있는 TC 63/120"
    assert "독립 경로(clang 등)로 대조하지 않았다" in item["note"]
    assert "흐름이 쓴 값 5511" in item["note"] and "입력을 그대로 둔 값 4968" in item["note"]
    assert "path_budget 7021" in item["note"] and "흐름 1개는 나머지 sub-case 를 돌리지 않았다" in item["note"]
    assert "실행 결과가 아니다" in item["note"] and item["tone"] == "info"
    failed = {i["key"]: i for i in build_disclosures(
        "sits", {"integration_oracle": {"status": "error:RuntimeError", "error": "boom"}})}["sits_integration_oracle"]
    assert failed["value"] == "도출 실패" and failed["tone"] == "warning" and "boom" in failed["note"]
    none = {i["key"]: i for i in build_disclosures(
        "sits", {"integration_oracle": {"status": "no_project_context"}})}["sits_integration_oracle"]
    assert none["value"] == "도출 안 함" and "[검증 필요]" in none["note"]
    assert "sits_integration_oracle" not in {i["key"] for i in build_disclosures("sits", {})}   # old outputs: no item


def test_callee_static_locals_do_not_survive_unknown_code():
    b_static = B + "U8 toggle(U8 set) { static U8 v; if (set) { v = 7U; return 0U; } return v; }\n"
    a = A + ("U8 toggle(U8 set);\nvoid keep(void) { (void)toggle(1U); g_sum = toggle(0U); }\n"
             "void lose(void) { (void)toggle(1U); b_ext(); g_sum = toggle(0U); }\n")
    files, ctx, scopes = _project({os.path.join(ROOT, "b.c"): b_static, os.path.join(ROOT, "a.c"): a})
    unit = {"source_text": a, "source_path": os.path.join(ROOT, "a.c"), "source_text_complete": True,
            "project_scope": scopes[os.path.join(ROOT, "a.c")],
            "callee_provider": CalleeProvider(ctx, files, scopes, cpc.shared_parser())}
    # the first call leaves the static at 7 and the second reads it back
    assert _value(evaluate_outputs({**unit, "name": "keep"}, [{}], [["g_sum"]])[0], "g_sum") == 7
    # unknown code in between (``lib_call`` inside b_ext) may call toggle(1U) or toggle(0U): the static is unknown
    r = evaluate_outputs({**unit, "name": "lose"}, [{}], [["g_sum"]])[0]
    assert _value(r, "g_sum") is None and r["interprocedural"]["inlined"].get("toggle") == 2


def test_separate_builds_link_within_their_own_root():
    app, boot = os.path.join(ROOT, "app"), os.path.join(ROOT, "boot")
    hdr = "typedef unsigned char U8;\ntypedef unsigned int U16;\nextern U8 g_v;\nU8 drv(void);\n"
    files = {os.path.join(app, "m.h"): hdr, os.path.join(boot, "m.h"): hdr,
             os.path.join(app, "main.c"): '#include "m.h"\nU8 g_v;\nvoid run(void) { g_v = drv(); }\n',
             os.path.join(app, "drv.c"): '#include "m.h"\nU8 drv(void) { return 1U; }\n',
             os.path.join(boot, "drv.c"): '#include "m.h"\nU8 drv(void) { return 2U; }\n'}
    ctx = cpc.build_project_context(files, roots=[app, boot])
    units = [p for p in files if p.endswith(".c")]
    scopes = cpc.build_scopes(ctx, units)
    main = os.path.join(app, "main.c")
    unit = {"name": "run", "source_text": files[main], "source_path": main, "source_text_complete": True,
            "project_scope": scopes[main], "callee_provider": CalleeProvider(ctx, files, scopes, cpc.shared_parser())}
    assert _value(evaluate_outputs(unit, [{}], [["g_v"]])[0], "g_v") == 1
    alone = CalleeProvider(cpc.build_project_context(files), files, scopes, cpc.shared_parser())
    with pytest.raises(Unsupported, match="definition_ambiguous"):
        alone.definition("drv", main)   # without the roots the two builds are one: never guessed
