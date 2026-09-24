"""R6 (P3/G4) — STS requirement boundary TCs from the requirement text (`generators/sts_requirement_tc.py`).

The verdict of every step comes from the requirement sentence alone — the whole line on that subject, joined as the
line joins it; the step is one unit of the written precision.
"""
from __future__ import annotations

from collections import Counter

import openpyxl

from generators.sts import _build_tc_dict, _classify_steps, _make_tc_id
from generators.sts_requirement_tc import (
    ACTION_PREFIX,
    BASIS_MARK,
    REQUIREMENT_EVIDENCE_HEADERS,
    append_requirement_boundary_tcs,
    boundary_steps,
    write_requirement_evidence_sheet,
)


def _req(desc, rid="SwTSR_0209", verification=""):
    return {"id": rid, "name": "n", "description": desc, "verification": verification, "asil": "B"}


def _steps(desc, **kw):
    return [s for g in boundary_steps(_req(desc, **kw)) for s in g["steps"]]


def _point(step):
    return step["action"][len(ACTION_PREFIX):].split(BASIS_MARK)[0].split(" = ", 1)[1]


def _acts(step):
    return "불성립" not in step["expected"]


def test_a_threshold_gets_the_points_either_side_and_on_it_with_the_requirements_verdict():
    (g,) = boundary_steps(_req("- High 배터리 전압인가 u16g_ApiIn_Vsup: 1604 초과"))
    assert [(_point(s), _acts(s)) for s in g["steps"]] == [("1603", False), ("1604", False), ("1605", True)]
    assert all(s["action"].startswith(ACTION_PREFIX + "u16g_ApiIn_Vsup = ") for s in g["steps"])
    e = g["evidence"]
    assert (e["op"], e["value"], e["step"], e["step_basis"]) == (">", "1604", "1", "written_precision")
    assert "1604 초과" in g["steps"][0]["expected"]   # the condition is stated in the requirement's own words


def test_the_step_is_the_precision_the_value_is_written_with():
    assert [_point(s) for s in _steps("- Battery 전압이 8.5V미만인 경우 저전압")] == ["8.4V", "8.5V", "8.6V"]
    assert [_point(s) for s in _steps("- u16s_X 가 16.04V 이상")] == ["16.03V", "16.04V", "16.05V"]
    # (review C1) ``4.00`` is written to two decimals — the step is 0.01, never 1
    assert [_point(s) for s in _steps("- u16s_Y 가 4.00V 이하")] == ["3.99V", "4.00V", "4.01V"]
    # a hex value keeps its base
    assert [_point(s) for s in _steps("- Stack Pointer값이 0x1FF0이상이면")] == ["0x1FEF", "0x1FF0", "0x1FF1"]


def test_an_or_line_on_one_subject_is_judged_as_a_whole():
    # (review C1) SwTSR_0101: 5 V is above 4.00 V but meets the other condition (≥ 4.74 V): the requirement acts
    groups = boundary_steps(_req("- Position Sensor의 Main 전원(u16g_ApiIn_MagnetLevel)이 4.00V 이하 또는 4.74V 이상인 "
                                 "상태로 특정 시간(u16s_MAGNET_ERR_TM : 100ms) 이상 유지 시", "SwTSR_0101"))
    low = next(g for g in groups if g["evidence"]["value"] == "4.00")
    assert [(_point(s), _acts(s)) for s in low["steps"]] == [("3.99V", True), ("4.00V", True), ("4.01V", False)]
    high = next(g for g in groups if g["evidence"]["value"] == "4.74")
    assert [(_point(s), _acts(s)) for s in high["steps"]] == [("4.73V", False), ("4.74V", True), ("4.75V", True)]
    assert "4.00V 이하 또는 4.74V 이상" in low["steps"][0]["expected"]


def test_an_and_line_on_one_subject_needs_every_condition():
    groups = boundary_steps(_req("- u16s_V 가 9V 이상 그리고 16V 이하"))
    nine = next(g for g in groups if g["evidence"]["value"] == "9")
    assert [(_point(s), _acts(s)) for s in nine["steps"]] == [("8V", False), ("9V", True), ("10V", True)]


def test_conditions_that_cannot_be_stepped_are_counted_with_their_reason():
    stats = Counter()
    groups = boundary_steps(_req("2. 500ms 초과 저전압 상황 동작 확인\nu16g_ApiIn_Vsup\t850 ~ 1,604\n"
                                 "( u16g_ApiIn_Vsup < u16s_BATT_ERR_LOWER_LIMIT )\n"
                                 "- 전압이 8.5V 미만이 아닌 경우 정상\n"          # (review C2) negated
                                 "- Application 안정화 시간 : 100ms 이내\n"      # (review C3) a response constraint
                                 "- u16s_V 가 9V 이상 16V 이하\n"), stats)       # one subject, joining unstated
    assert groups == []
    assert stats["skipped:kind_range"] == 1 and stats["skipped:kind_symbolic"] == 1
    assert stats["skipped:negated_condition"] == 1 and stats["skipped:response_constraint"] == 1
    assert stats["skipped:same_subject_combination_unstated"] == 2 and stats["facts_used"] == 0


def test_a_sentence_repeated_in_the_verification_criteria_is_written_once():
    stats = Counter()
    groups = boundary_steps(_req("- u16s_A: 10 초과", verification="- u16s_A: 10 초과"), stats)
    assert len(groups) == 1 and stats["skipped:duplicate_fact"] == 1
    # the same comparison in another sentence (another context: another action) is its own fact (review r2 Info)
    groups = boundary_steps(_req("- u16s_A: 10 초과", verification="1. u16s_A: 10 초과 시 경고 동작 확인"))
    assert len(groups) == 2


def test_a_hold_time_gets_held_points_around_it():
    (g,) = boundary_steps(_req("- LIN 통신 이상으로 특정시간(300ms)초과하여 유지되는 경우"))
    assert [s["action"].split(BASIS_MARK)[0] for s in g["steps"]] == [
        ACTION_PREFIX + "조건 지속 시간 = 299ms", ACTION_PREFIX + "조건 지속 시간 = 300ms",
        ACTION_PREFIX + "조건 지속 시간 = 301ms"]
    assert [_acts(s) for s in g["steps"]] == [False, False, True]


def test_a_hold_time_on_a_phrase_says_what_lasts():
    # (review r2 W6) ``15초이상 끊길`` — the step must not read "keep the communication for 14 s"
    (g,) = boundary_steps(_req("- LIN 통신이 15초이상 끊길 경우 오류"))
    assert g["steps"][0]["action"].split(BASIS_MARK)[0] == ACTION_PREFIX + "LIN 통신 지속 시간 = 14초 (끊길 상태)"


def test_a_parameter_threshold_sets_the_quantity_it_is_compared_with():
    # (review r2 C-B) Param_X is the reference constant; the opening angle is what the step sets
    (g,) = boundary_steps(_req("3. Param_AntipinchReverseActAngle( 3도 ) 초과한 열림각에서 끼임 발생한 경우", "SwTR_0202"))
    head = [s["action"].split(BASIS_MARK)[0] for s in g["steps"]]
    assert head[0] == ACTION_PREFIX + "열림각 = 2도 (기준 Param_AntipinchReverseActAngle = 3도)"
    assert g["evidence"]["reference_constant"] == "Param_AntipinchReverseActAngle"
    stats = Counter()
    assert boundary_steps(_req("- Param_X( 3도 ) 초과"), stats) == []
    assert stats["skipped:monitored_quantity_unknown"] == 1


def test_a_named_hold_time_constant_is_the_reference_not_what_is_held():
    groups = boundary_steps(_req("- Position Sensor의 Main 전원(u16g_ApiIn_MagnetLevel)이 4.00V 이하 또는 4.74V 이상인 "
                                 "상태로 특정 시간(u16s_MAGNET_ERR_TM : 100ms) 이상 유지 시", "SwTSR_0101"))
    hold = next(g for g in groups if g["evidence"]["kind"] == "duration")
    assert hold["steps"][0]["action"].split(BASIS_MARK)[0] == \
        ACTION_PREFIX + "조건 지속 시간 = 99ms (기준 u16s_MAGNET_ERR_TM = 100ms)"
    # (review r2 C-A) the hold time qualifies the voltage condition: while stepping it, the condition holds (AND)
    assert "성립 상태로 둔다" in hold["evidence"]["combination_note"]
    low = next(g for g in groups if g["evidence"]["value"] == "4.00")
    assert low["evidence"]["combination_note"] == \
        "다른 조건 [조건 지속 시간 100ms 이상 (기준 u16s_MAGNET_ERR_TM)]: 성립 상태로 둔다"
    # (review r3 W2) the voltage conditions are one condition for the hold time's note, not two contradictory ones
    assert hold["evidence"]["combination_note"] == \
        "다른 조건 [u16g_ApiIn_MagnetLevel 4.00V 이하 또는 u16g_ApiIn_MagnetLevel 4.74V 이상]: 성립 상태로 둔다"


def test_a_comma_between_conditions_is_not_read_as_a_join():
    # ``P || R, Pos ≥ 50``: whether Pos joins by and/or is not written — the tester is told, nothing is assumed
    groups = boundary_steps(_req("- 오른쪽 도어: DoorPitch >= 8 || DoorRoll <= -4, 도어 Positon 50도 이상", "SwTR_0201"))
    pos = next(g for g in groups if g["evidence"]["value"] == "50")
    assert "결합 조건 확인 필요" in pos["evidence"]["combination_note"]
    pitch = next(g for g in groups if g["evidence"]["value"] == "8")
    assert "다른 조건 [DoorRoll -4 이하]: 불성립 상태로 둔다" in pitch["evidence"]["combination_note"]


def test_an_inequality_sibling_takes_part_in_the_verdict():
    # (review r2 W4) ``u16g_A > 10 && u16g_A != 11``: at 11 the line does not hold
    (g, _ne) = boundary_steps(_req("- u16g_A > 10 && u16g_A != 11")) + [None]
    assert [(_point(s), _acts(s)) for s in g["steps"]] == [("9", False), ("10", False), ("11", False)]


def test_other_subjects_are_held_false_on_or_and_true_on_and_in_the_precondition():
    reqs = [_req("- u16s_A: 10 초과 또는 u16s_B: 20 초과")]
    tcs = []
    append_requirement_boundary_tcs(tcs, reqs, _build, _make_tc_id, _classify_steps)
    assert "다른 조건 [u16s_B 20 초과]: 불성립 상태로 둔다" in tcs[0]["precondition"]
    reqs = [_req("- u16s_A: 10 초과 && u16s_B: 20 초과")]
    tcs = []
    append_requirement_boundary_tcs(tcs, reqs, _build, _make_tc_id, _classify_steps)
    assert "다른 조건 [u16s_B 20 초과]: 성립 상태로 둔다" in tcs[0]["precondition"]
    # (review r2 W7) no guessed "mid" initial value for the stepped subject
    assert "입력:" not in tcs[0]["precondition"]


def _build(**kw):
    # the same as `generate_sts` wires it: the step sets the value, no guessed initial input
    return _build_tc_dict(test_env="SwTE_01", is_safety=False, derive_inputs=False, **kw)


def test_boundary_tcs_number_on_after_the_requirements_own_tcs_and_are_labelled_by_their_steps():
    reqs = [_req("- u16g_ApiIn_Vsup: 1604 초과", "SwTSR_0209"), _req("설명만 있음", "SwTR_0101")]
    existing = [_build(tc_id=_make_tc_id("SwTSR_0209", 1), req=reqs[0], steps=[{"action": "a", "expected": "b"}],
                       test_method="RBT", gen_method="AOR"),
                _build(tc_id=_make_tc_id("SwTR_0101", 1), req=reqs[1], steps=[{"action": "a", "expected": "b"}],
                       test_method="RBT", gen_method="AOR")]
    tcs = list(existing)
    stats = append_requirement_boundary_tcs(tcs, reqs, _build, _make_tc_id, _classify_steps)
    assert [tc["id"] for tc in tcs] == ["SwTC_SwTSR_0209_01", "SwTC_SwTSR_0209_02", "SwTC_SwTR_0101_01"]
    added = tcs[1]
    # (review W1) BAA comes from the step wording the classifier (and the validator's label audit) reads
    assert added["requirement_boundary"] and added["gen_method"] == "BAA" and len(added["steps"]) == 3
    assert _classify_steps(added["steps"])[1] == "BAA"
    assert stats["tcs"] == 1 and stats["steps"] == 3


def test_a_facts_points_never_split_across_tcs():
    req = _req("- u16s_A: 10 초과\n- u16s_B: 20 초과\n- u16s_C: 30 초과")
    tcs = []
    append_requirement_boundary_tcs(tcs, [req], _build, _make_tc_id, _classify_steps, max_steps=4)
    assert [len(tc["steps"]) for tc in tcs] == [3, 3, 3]
    assert [[e["signal"] for e in tc["requirement_evidence"]] for tc in tcs] == [["u16s_A"], ["u16s_B"], ["u16s_C"]]
    tcs = []
    append_requirement_boundary_tcs(tcs, [req], _build, _make_tc_id, _classify_steps, max_steps=6)
    assert [len(tc["steps"]) for tc in tcs] == [6, 3]


def test_the_evidence_sheet_names_the_sentence_each_tc_comes_from_and_replaces_an_old_one():
    tcs = []
    append_requirement_boundary_tcs(tcs, [_req("- u16g_ApiIn_Vsup: 1604 초과")], _build, _make_tc_id,
                                    _classify_steps)
    wb = openpyxl.Workbook()
    wb.create_sheet("Requirement Evidence")   # a template made from an earlier output
    assert write_requirement_evidence_sheet(wb, tcs) == 1
    assert wb.sheetnames.count("Requirement Evidence") == 1 and "Requirement Evidence1" not in wb.sheetnames
    rows = list(wb["Requirement Evidence"].iter_rows(values_only=True))
    assert list(rows[0]) == REQUIREMENT_EVIDENCE_HEADERS
    assert rows[1][:8] == ("SwTC_SwTSR_0209_01", "SwTSR_0209", "threshold", "u16g_ApiIn_Vsup", "—", ">", "1604", "—")
    assert rows[1][10] == "1603(F), 1604(F), 1605(T)"


def test_no_boundary_tc_writes_no_evidence_sheet_and_drops_a_template_one():
    wb = openpyxl.Workbook()
    assert write_requirement_evidence_sheet(wb, []) == 0 and "Requirement Evidence" not in wb.sheetnames
    wb.create_sheet("Requirement Evidence")
    assert write_requirement_evidence_sheet(wb, []) == 0 and "Requirement Evidence" not in wb.sheetnames


def test_generate_sts_writes_boundary_tcs_and_their_sheet_only_when_enabled(tmp_path):
    from generators.sts import generate_sts
    reqs = ["SwTSR_0209: u16g_ApiIn_Vsup: 1604 초과 시 고전압 고장을 검출한다."]
    on = generate_sts(reqs, {}, str(tmp_path / "on.xlsm"), project_config={"project_id": "T"})
    off = generate_sts(reqs, {}, str(tmp_path / "off.xlsm"),
                       project_config={"project_id": "T", "requirement_boundary_tcs": False})
    assert on["test_case_count"] == off["test_case_count"] + 1
    assert on["quality_report"]["generation_stats"]["requirement_boundary"]["tcs"] == 1
    assert "requirement_boundary" not in off["quality_report"]["generation_stats"]
    wb_on = openpyxl.load_workbook(on["output_path"], read_only=True)
    wb_off = openpyxl.load_workbook(off["output_path"], read_only=True)
    try:
        assert "Requirement Evidence" in wb_on.sheetnames and "Requirement Evidence" not in wb_off.sheetnames
    finally:
        wb_on.close()
        wb_off.close()
    assert not [w for w in (on.get("validation") or {}).get("warnings") or [] if "BAA" in w]
    wb = openpyxl.load_workbook(on["output_path"], read_only=True)
    try:
        ws = wb["3.SW Test Spec"]
        rows = [r for r in ws.iter_rows(values_only=True) if r and len(r) > 10 and r[10]
                and str(r[10]).startswith("입력 설정 (요구 경계)")]
        assert rows and all("입력:" not in str(r[9] or "") for r in rows)   # (review r2 W7) wiring passes it
    finally:
        wb.close()
