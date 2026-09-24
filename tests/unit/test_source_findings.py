"""R10 — source findings: undefined behaviour the oracle proved while deriving expected values becomes a sheet of
defect candidates (one per function and kind), never a silent "unknown"."""
from __future__ import annotations

import openpyxl

from generators.source_findings import (
    FINDINGS_HEADERS,
    FINDINGS_SHEET,
    NOT_REPRODUCED_HERE,
    collect_findings,
    summarize_findings,
    write_source_findings_sheet,
)
from report_gen.generation_disclosures import build_disclosures


def _seq(n, inputs, evidence):
    return {"seq_num": n, "inputs": inputs, "expected_evidence": evidence}


UNITS = [{"fid": "SwUFn_1", "name": "f"}, {"fid": "SwUFn_2", "name": "g"}]
SEQS = {
    "SwUFn_1": [
        _seq(1, {"a": 32767}, {"x": {"status": "unknown", "reason": "undefined_behavior:signed_overflow",
                                     "source_path": "f.c", "source_hash": "h1"},
                               "y": {"status": "unknown", "reason": "undefined_behavior:signed_overflow"}}),
        _seq(2, {"a": 0}, {"x": {"status": "unknown", "reason": "scope_x;undefined_behavior:division_by_zero"},
                           "y": {"status": "derived", "reason": ""}}),
    ],
    "SwUFn_2": [_seq(1, {"b": 1}, {"z": {"status": "unknown", "reason": "callee_pointer_write"}})],
}


def test_findings_group_by_function_and_kind_with_an_example():
    findings = collect_findings(UNITS, SEQS, {"SwUFn_1": "SwUTC_SwUFn_1"})
    assert [(f["function"], f["kind"], f["occurrences"]) for f in findings] == [
        ("f", "division_by_zero", 1), ("f", "signed_overflow", 2)]
    over = findings[1]
    assert over["observables"] == ["x", "y"] and over["example_inputs"] == {"a": 32767}
    assert (over["example_tc"], over["example_sequence"], over["source_path"]) == ("SwUTC_SwUFn_1", 1, "f.c")
    # (review W4) the measure is sequences (a UB refuses the whole run); blocked slots are reported beside it
    assert summarize_findings(findings) == {"findings": 2, "by_kind": {"division_by_zero": 1, "signed_overflow": 1},
                                            "functions": 1, "sequences": 2, "blocked_expected_slots": 3,
                                            "reproduction": "not_run"}


def test_an_unknown_for_another_reason_is_not_a_finding():
    assert collect_findings(UNITS[1:], SEQS, {}) == []


def test_the_sheet_says_what_a_finding_means_and_that_it_is_not_reproduced_here():
    wb = openpyxl.Workbook()
    wb.create_sheet(FINDINGS_SHEET)   # a template made from an earlier output
    assert write_source_findings_sheet(wb, collect_findings(UNITS, SEQS, {})) == 2
    assert wb.sheetnames.count(FINDINGS_SHEET) == 1
    rows = list(wb[FINDINGS_SHEET].iter_rows(values_only=True))
    assert list(rows[0]) == FINDINGS_HEADERS
    assert rows[2][2] == "signed_overflow" and "결함 후보" in str(rows[2][11]) and rows[2][12] == NOT_REPRODUCED_HERE
    assert (rows[2][3], rows[2][4]) == (1, 2)   # one sequence, two blocked expected slots
    wb2 = openpyxl.Workbook()
    wb2.create_sheet(FINDINGS_SHEET)
    assert write_source_findings_sheet(wb2, []) == 0 and FINDINGS_SHEET not in wb2.sheetnames


def test_the_suts_disclosure_names_the_findings():
    items = {i["key"]: i for i in build_disclosures("suts", {"source_findings": summarize_findings(
        collect_findings(UNITS, SEQS, {}))})}
    item = items["suts_source_findings"]
    assert item["value"] == "소견 2 (함수 1 · 시퀀스 2)" and "signed_overflow 1" in item["note"]
    assert "적어도 한 실행 경로" in item["note"] and "증명" not in item["note"]   # (review W5)
    assert "suts_source_findings" not in {i["key"] for i in build_disclosures("suts", {})}



# ── review round 2 (W3, W7): the reproduction script ──────────────────────────────

def _script():
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
    import source_findings as sf
    return sf


def test_reproduce_sorts_every_finding_into_one_of_four_answers(monkeypatch):
    sf = _script()
    import source_oracle_clang_check as socc
    findings = [{"source_path": "f.c", "function": "f", "kind": k, "example": {"inputs": {}, "observable": "x"}}
                for k in ("signed_overflow", "signed_overflow", "division_by_zero", "signed_overflow", "division_by_zero")]
    reports = iter([
        {"eval_error": 1, "checked": 1, "eval_errors": [{"clang": "signed integer overflow: 32767 + 1"}]},   # reproduced
        {"eval_error": 1, "checked": 1, "eval_errors": [{"clang": "division by zero"}]},              # another error
        {"eval_error": 0, "checked": 1},                                                               # every fill a value
        {"eval_error": 0, "checked": 0, "unchecked_reasons": {"constexpr_limit": 1}},                  # clang could not
    ])
    monkeypatch.setattr(sf, "_units", lambda fs, roots: {("f.c", "f"): {"name": "f"}} if fs else {})
    def _check(claims, clang, ub_probe):
        assert ub_probe is True   # (round 3 I6) the placeholder value is only sound in UB-probe mode
        return next(reports)
    monkeypatch.setattr(socc, "check_claims", _check)
    findings[4]["source_path"] = "gone.c"
    verdicts = sf.reproduce(findings, [])
    assert [f["reproduction"] for f in findings] == ["reproduced", "other_evaluation_error", "not_reproduced",
                                                     "unchecked:constexpr_limit", "unchecked:source_not_found"]
    assert verdicts == {"reproduced": 1, "other_evaluation_error": 1, "not_reproduced": 1, "unchecked": 2}


def test_a_fill_clang_could_not_evaluate_is_never_not_reproduced():
    """(W3) fill 0 hits a constexpr limit, fills 90/201 give a value ≠ the placeholder: probing, that is unchecked —
    "not reproduced" needs every fill evaluated."""
    import source_oracle_clang_check as socc
    assert socc._PROBE_RANK["constexpr_limit"] > socc._PROBE_RANK["mismatch"]
    assert socc._PROBE_RANK["eval_error"] > socc._PROBE_RANK["constexpr_limit"]
    assert socc._CHECK_RANK["mismatch"] > socc._CHECK_RANK["eval_error"] > socc._CHECK_RANK["constexpr_limit"]


def test_an_out_of_bounds_index_has_a_meaning_and_a_clang_message():
    sf = _script()
    assert "배열" in MEANING_OOB and sf._KIND_MESSAGE["array_index_out_of_bounds"].search(
        "read of dereferenced one-past-the-end pointer")


from generators.source_findings import MEANING as _MEANING  # noqa: E402

MEANING_OOB = _MEANING["array_index_out_of_bounds"]


def test_run_tu_probe_answers_unchecked_when_one_fill_could_not_be_evaluated(monkeypatch, tmp_path):
    """(W3, behaviour) fill 0: a harness limit; fill 90: a value that is not the placeholder. Probing → constexpr_limit
    (unchecked), never mismatch ("every fill evaluated"); checking → mismatch as before."""
    import subprocess
    import types

    import source_oracle_clang_check as socc
    out = "\n".join([
        "u.cpp:10:1: error: static assertion expression is not an integral constant expression "
        "(__oracle_check_0_0_0)",
        "u.cpp:10:1: note: cast that performs the conversions of a reinterpret_cast is not allowed",
        "u.cpp:11:1: error: static assertion failed due to requirement '__oracle_check_0_0_1'",
        "u.cpp:11:1: note: expression evaluates to '2 == 0'",
        "u.cpp:99:1: error: __oracle_canary"])
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: types.SimpleNamespace(stdout=out, stderr="", returncode=1))
    checks = {10: (0, "g_o", 0, 0, "0_0_0"), 11: (0, "g_o", 0, 1, "0_0_1"), 12: (0, "g_o", 0, 2, "0_0_2")}
    probe, _d, _r = socc._run_tu("", checks, str(tmp_path / "u.cpp"), "clang", "msp430", 10, ub_probe=True)
    plain, _d, _r = socc._run_tu("", checks, str(tmp_path / "u.cpp"), "clang", "msp430", 10)
    assert probe == {(0, "g_o"): "constexpr_limit"} and plain == {(0, "g_o"): "mismatch"}
