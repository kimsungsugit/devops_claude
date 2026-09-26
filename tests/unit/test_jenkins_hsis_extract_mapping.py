"""Regression: /api/jenkins/hsis/extract-mapping 시스템 기준(SyRS) 커버리지.

HSIS(HW-SW 인터페이스)는 시스템레벨 문서라 SW 요구(SwRS 68)가 아니라 시스템 요구(SyRS)를
참조한다. 매트릭스가 Sy→Sw 평탄화로 SW 68에 조인하면 시스템-only 요구(SyTR_0802 등)가
침묵 탈락해 밴드가 6/68로 과소 표시됐다(실측 근본원인, SyRS 기준은 21/21). 이 커밋은
엔드포인트가 syrs_path를 받아 원본 Sy* 참조를 SyRS에 조인한 `system_basis` 진단을 병행
반환하는 것(기존 SW 필드 불변)의 회귀 가드. 파서(parse_hsis_signals)는 모킹해 분리 검증.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from docx import Document


@pytest.fixture(autouse=True)
def _local_mode(monkeypatch: pytest.MonkeyPatch):
    from backend.services import file_resolver as fr
    monkeypatch.setattr(fr, "_resolver", fr.LocalFileResolver())


def _syrs_doc(tmp_path: Path, tokens: str) -> str:
    p = tmp_path / "syrs.docx"
    d = Document()
    d.add_paragraph(tokens)
    d.save(str(p))
    return str(p)


def _empty_xlsx(tmp_path: Path) -> str:
    p = tmp_path / "hsis.xlsx"
    openpyxl.Workbook().save(str(p))
    return str(p)


def _mock_signals(monkeypatch: pytest.MonkeyPatch, signals: list) -> None:
    """엔드포인트 내부 `from generators.sts import parse_hsis_signals`를 대체."""
    import generators.sts as gsts
    monkeypatch.setattr(gsts, "parse_hsis_signals",
                        lambda _p: {"signals": signals, "sw_var_names": []})


def test_load_syrs_req_set(tmp_path: Path) -> None:
    """SyRS docx에서 Sy* 요구 ID 집합(대문자) 로드."""
    from backend.routers.jenkins import _load_syrs_req_set
    from backend.services.file_resolver import LocalFileResolver
    p = _syrs_doc(tmp_path, "SyTR_0101 SyEIF_02 무관텍스트 SyTSR_0110")
    s = _load_syrs_req_set(p, LocalFileResolver())
    assert s == {"SYTR_0101", "SYEIF_02", "SYTSR_0110"}


def test_system_basis_joins_syrs_not_sw(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """system_basis가 원본 Sy* 참조를 SyRS에 조인 — SW 평탄화 손실 없이 참 커버리지.

    mutation: sys_refs 수집(Sy* 원본 보존)을 제거하면 joined=0이 되어 이 assert 실패.
    """
    from backend.routers.jenkins import jenkins_hsis_extract_mapping
    _mock_signals(monkeypatch, [
        {"id": "HSI_01", "related_id": "SyEIF_01"},
        {"id": "HSI_02", "related_id": "SyEIF_02"},
        {"id": "HSI_03", "related_id": "SyTR_0507"},   # SW SRS엔 없지만 SyRS엔 있음
        {"id": "HSI_04", "related_id": "SyABC_9999"},  # SyRS에도 없음(unmatched)
    ])
    syrs = _syrs_doc(tmp_path, "SyEIF_01 SyEIF_02 SyEIF_03 SyEIF_06 SyTR_0507 SyTSR_0109")
    res = jenkins_hsis_extract_mapping(
        {"hsis_path": _empty_xlsx(tmp_path), "syrs_path": syrs})
    assert res["ok"] is True
    sb = res["system_basis"]
    assert sb["refs_total"] == 4
    assert sb["joined"] == 3                              # SyEIF_01/02 + SyTR_0507
    assert sb["unmatched"] == ["SYABC_9999"]              # SyRS 밖 표면화(침묵 금지)
    # 시스템 인터페이스요구(SyEIF): SyRS 4개 중 2 커버, 2 미커버
    assert sb["interface_total"] == 4
    assert sb["interface_covered"] == 2
    assert sb["interface_missing"] == ["SYEIF_03", "SYEIF_06"]


def test_no_syrs_path_no_system_basis(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """하위호환: syrs_path 미제공 시 system_basis 없음, 기존 SW 필드 유지."""
    from backend.routers.jenkins import jenkins_hsis_extract_mapping
    _mock_signals(monkeypatch, [{"id": "HSI_01", "related_id": "SyEIF_01"}])
    res = jenkins_hsis_extract_mapping({"hsis_path": _empty_xlsx(tmp_path)})
    assert res["ok"] is True
    assert "system_basis" not in res
    assert "hsis_pairs" in res and "total_signals" in res  # 기존 계약 불변


def test_bad_syrs_degrades_not_crash(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """비-docx/손상 syrs_path → 엔드포인트 크래시 아닌 degrade, hsis_pairs(주 산출물) 보존.

    deep-review Warning#2: docx 파싱 실패(PackageNotFoundError는 OSError 아님)가 uncaught면
    500으로 hsis_pairs까지 파괴돼 "additive·하위호환" 계약이 깨진다. 광역 except로 degrade.
    """
    from backend.routers.jenkins import jenkins_hsis_extract_mapping
    _mock_signals(monkeypatch, [{"id": "HSI_01", "related_id": "SyEIF_01"}])
    bad = tmp_path / "not_a.docx"
    bad.write_text("this is plainly not a docx package", encoding="utf-8")
    res = jenkins_hsis_extract_mapping(
        {"hsis_path": _empty_xlsx(tmp_path), "syrs_path": str(bad)})
    assert res["ok"] is True                    # 크래시 안 함
    assert "hsis_pairs" in res                  # 주 산출물 보존
    assert "system_basis" not in res            # 진단은 degrade
    assert res.get("system_basis_error")        # fail-loud (침묵 아님)
    assert "/" not in res["system_basis_error"] and "\\" not in res["system_basis_error"]  # 경로 미노출


def test_system_basis_all_covered_no_missing(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """인터페이스요구 전부 커버 시 interface_missing 빈 배열(오경보 없음)."""
    from backend.routers.jenkins import jenkins_hsis_extract_mapping
    _mock_signals(monkeypatch, [
        {"id": "HSI_01", "related_id": "SyEIF_01"},
        {"id": "HSI_02", "related_id": "SyEIF_02"},
    ])
    syrs = _syrs_doc(tmp_path, "SyEIF_01 SyEIF_02")
    res = jenkins_hsis_extract_mapping(
        {"hsis_path": _empty_xlsx(tmp_path), "syrs_path": syrs})
    sb = res["system_basis"]
    assert sb["interface_covered"] == 2 and sb["interface_total"] == 2
    assert sb["interface_missing"] == []
    assert sb["unmatched"] == []
