"""TD-P0: independent capability benchmark detector tests (synthetic QM)."""
import pytest

from scripts.test_generation_semantic_baseline import check_rows, independent_conditions


def test_td_p0_or_all_true_toggle_is_not_mcdc():
    vectors = [(True, True), (False, True), (True, False)]
    assert independent_conditions(vectors, lambda v: v[0] or v[1]) == set()


def test_td_p0_or_false_baseline_has_both_independent_effects():
    vectors = [(False, False), (True, False), (False, True)]
    assert independent_conditions(vectors, lambda v: v[0] or v[1]) == {0, 1}


def test_td_p0_changing_two_conditions_is_not_independence():
    assert independent_conditions([(False, False), (True, True)], all) == set()


def test_td_p0_affine_detector_rejects_type_midpoint():
    rows = [{"inputs": {"x": 127}, "expected": {"y": 32767}}]
    assert check_rows(rows, lambda x: 2*x+3)


def test_td_p0_affine_detector_accepts_exact_calculation():
    rows = [{"inputs": {"x": 127}, "expected": {"y": 257}}]
    assert not check_rows(rows, lambda x: 2*x+3)


def test_td_p0_missing_behavior_requires_visible_review_marker():
    assert check_rows([{"inputs": {"x": 0}, "expected": {"y": 0}}])
    assert check_rows([{"inputs": {"x": 0}, "expected": {}}])
    assert not check_rows([{"inputs": {"x": 0}, "expected": {"y": "[검증 필요] no behavior"}}])


def test_td_p1_suts_affine_expected_values_have_source_evidence():
    # TD-P1-ORACLE: the oracle is independent arithmetic, not the evaluator under test.
    from generators.suts import generate_sequences

    unit = {"name": "affine", "input_vars": ["x"], "output_vars": ["return"],
            "param_types": {"x": "uint8_t", "return": "uint16_t"}, "logic_flow": [],
            "source_text": "#include <stdint.h>\nuint16_t affine(uint8_t x) { return 2 * x + 3; }"}
    rows = generate_sequences(unit, max_seq=6)
    assert rows
    for row in rows:
        value = row["expected"]["return"]
        evidence = row["expected_evidence"]["return"]
        assert evidence["execution_status"] == "not_run"
        x = row["inputs"]["x"]
        if 0 <= x <= 255:
            assert value == 2*x+3
            assert evidence["status"] == "derived"
            assert evidence["oracle_kind"] == "source"
            assert evidence["source_hash"]
        else:
            assert isinstance(value, str) and value.startswith("[검증 필요]")
            assert evidence["status"] == "unknown"


def test_td_p1_sits_type_and_call_chain_never_prove_output_behavior():
    # TD-P1-UNKNOWN: includes boundary, combinations, and error propagation rows.
    from generators.sits import _generate_sub_cases

    flow = {"input_vars": ["x", "z"], "expected_vars": ["y"],
            "input_raws": ["U8 x", "U8 z"], "expected_raws": ["U16 y"],
            "call_chain": "Source -> Consumer"}
    rows = _generate_sub_cases(flow, max_cases=20)
    assert len(rows) > 7
    for row in rows:
        assert isinstance(row["expected"]["y"], str)
        assert row["expected"]["y"].startswith("[검증 필요]")
        evidence = row["expected_evidence"]["y"]
        assert evidence["status"] == "unknown"
        assert evidence["oracle_kind"] == "none"
        assert evidence["reason"]
        assert evidence["execution_status"] == "not_run"


@pytest.mark.parametrize("x, expected", [(0, 3), (10, 23), (11, 7), (255, 7)])
def test_td_p1_source_branches_and_local_calculation(x, expected):
    from generators.c_test_semantics import evaluate_function

    source = "#include <stdint.h>\nuint16_t f(uint8_t x) { uint16_t y = 2*x+3; if (x > 10) { y = 7; } return y; }"
    result = evaluate_function(source, "f", {"x": x})
    assert result["status"] == "supported"
    assert result["value"] == expected
    assert result["execution_status"] == "not_run"


@pytest.mark.parametrize("body", [
    "return helper(x);",
    "return (uint8_t)x;",
    "while (x > 0) { x = x - 1; } return x;",
    "if (x > 0) { return helper(x); } return 0;",
])
def test_td_p1_unsupported_source_is_explicit_even_on_unvisited_branch(body):
    from generators.c_test_semantics import evaluate_function

    result = evaluate_function("#include <stdint.h>\nuint8_t f(uint8_t x) { " + body + " }", "f", {"x": 0})
    assert result["status"] == "unsupported"
    assert result["reason"]


@pytest.mark.parametrize("body", [
    "{ uint16_t local = 7; } return local;",
    "uint16_t local = 7; { uint16_t local = 8; } return local;",
])
def test_td_p1_block_scope_must_not_leak_or_silently_shadow(body):
    from generators.c_test_semantics import evaluate_function

    result = evaluate_function("#include <stdint.h>\nuint16_t f(uint8_t x) { " + body + " }", "f", {"x": 0})
    assert result["status"] == "unsupported"


def test_td_p1_unsigned_unary_promotion_requires_target_semantics():
    from generators.c_test_semantics import evaluate_function

    result = evaluate_function("#include <stdint.h>\nint f(uint16_t x) { return -x; }", "f", {"x": 1})
    assert result["status"] == "unsupported"


def test_td_p1_custom_typedef_must_not_inherit_builtin_width_by_name():
    from generators.c_test_semantics import evaluate_function

    source = "typedef unsigned long uint8_t; uint8_t f(uint8_t x) { return x; }"
    result = evaluate_function(source, "f", {"x": 1})
    assert result["status"] == "unsupported"


def test_td_p1_explicit_void_parameter_constant_return():
    from generators.c_test_semantics import evaluate_function

    result = evaluate_function("int f(void) { return 7; }", "f", {})
    assert result["status"] == "supported"
    assert result["value"] == 7


def test_td_p1_real_source_to_export_preserves_computed_oracle_and_provenance(tmp_path):
    from hashlib import sha256

    import openpyxl

    from generators.suts import attach_unit_sources, collect_unit_functions, generate_sequences, generate_suts_xlsm
    from report_gen.uds_generator import generate_uds_source_sections

    source = "#include <stdint.h>\nuint16_t affine(uint8_t x) { return 2 * x + 3; }\n"
    (tmp_path / "affine.c").write_text(source, encoding="utf-8")
    sections = generate_uds_source_sections(str(tmp_path), preprocess=False)
    details = sections["function_details"]
    detail = next(d for d in details.values() if d["name"] == "affine")
    # (R80) 원문은 파일당 한 번(`source_files`) — 생성 경로가 `attach_unit_sources` 로 unit 에 붙인다.
    assert sections["source_files"][detail["source_path"]] == source
    units = collect_unit_functions(details, sds_map={}, uds_io_map={})
    attach_unit_sources(units, sections["source_files"])
    unit = next(u for u in units if u["name"] == "affine")
    assert unit["source_text"] == source
    assert "return" in unit["output_vars"]
    rows = generate_sequences(unit, max_seq=6)
    assert not check_rows(rows, lambda x: 2*x+3, observable="return")
    output = tmp_path / "semantic_evidence.xlsx"
    generate_suts_xlsm(None, [unit], {unit["fid"]: rows}, str(output), {"project_id": "QM_SEMANTIC"})
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    try:
        records = list(workbook["Test Evidence"].values)
        data = [dict(zip(records[0], row, strict=False)) for row in records[1:]]
        assert len(data) == len(rows)
        assert all(row["Execution"] == "not_run" for row in data)
        computed = [row for row in data if row["Status"] == "derived"]
        assert computed
        assert all(row["Oracle"] == "source" for row in computed)
        assert all(row["Source SHA256"] == sha256(source.encode()).hexdigest() for row in computed)
        assert sorted(row["Expected"] for row in computed) == [3, 3, 257, 513]
    finally:
        workbook.close()
    from tools.export_suts_vectorcast import build_vectorcast_model

    model = build_vectorcast_model(str(output))
    cases = model["units"][0]["test_cases"]
    assert len(cases) == len(rows)
    derived = [case for case in cases if case["expected_evidence"]["return"]["status"] == "derived"]
    assert len(derived) == 4
    assert sorted(case["expected"]["return"] for case in derived) == [3, 3, 257, 513]
    for case in derived:
        evidence = case["expected_evidence"]["return"]
        assert evidence["source_hash"] == sha256(source.encode()).hexdigest()
        assert evidence["oracle_kind"] == "source"
        assert evidence["execution_status"] == "not_run"
        assert evidence["requirement_verified"] is False


@pytest.fixture
def exported_semantic_workbook(tmp_path):
    from generators.suts import generate_sequences, generate_suts_xlsm

    unit = {"fid": "SwUFn_0001", "name": "affine", "input_vars": ["x"],
            "output_vars": ["return", "pt_Out"],
            "param_types": {"x": "uint8_t", "return": "uint16_t", "pt_Out": "Thing *"},
            "logic_flow": [], "source_text": "#include <stdint.h>\nuint16_t affine(uint8_t x) { return 2*x+3; }"}
    rows = generate_sequences(unit, max_seq=6)
    output = tmp_path / "export.xlsx"
    generate_suts_xlsm(None, [unit], {unit["fid"]: rows}, str(output), {"project_id": "QM"})
    return output


@pytest.mark.parametrize("mutation, reason", [
    ("stale", "stale_expected_evidence"),
    ("duplicate", "ambiguous_expected_evidence"),
    ("missing", "missing_expected_evidence"),
    ("input", "stale_input_evidence"),
    ("function", "stale_function_evidence"),
])
def test_td_p1_export_rejects_stale_duplicate_or_missing_provenance(exported_semantic_workbook, mutation, reason):
    import openpyxl

    from tools.export_suts_vectorcast import build_vectorcast_model

    path = exported_semantic_workbook
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Test Evidence"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    selected = next(row for row in range(2, sheet.max_row+1)
                    if sheet.cell(row, headers["Status"]).value == "derived")
    sequence = sheet.cell(selected, headers["Sequence"]).value
    if mutation == "stale":
        sheet.cell(selected, headers["Expected"]).value = 9999
    elif mutation == "duplicate":
        sheet.append([cell.value for cell in sheet[selected]])
    elif mutation == "input":
        from generators.suts import _DATA_START_ROW, _INPUT_COL_START

        workbook["2.SW Unit Test Spec"].cell(_DATA_START_ROW + sequence, _INPUT_COL_START).value = 42
    elif mutation == "function":
        from generators.suts import _COL_UNIT, _DATA_START_ROW

        workbook["2.SW Unit Test Spec"].cell(_DATA_START_ROW, _COL_UNIT).value = "different_function"
    else:
        sheet.delete_rows(selected)
    workbook.save(path)
    workbook.close()
    model = build_vectorcast_model(str(path))
    case = next(case for case in model["units"][0]["test_cases"] if case["sequence_no"] == sequence)
    assert case["expected_evidence"]["return"]["status"] == "conflict"
    assert case["expected_evidence"]["return"]["reason"] == reason
    assert case["expected"]["return"]["verification_required"] is True
    assert reason in {warning["code"] for warning in model["export_warnings"]}


def test_td_p1_export_preserves_blank_unknown_observable_evidence(exported_semantic_workbook):
    from tools.export_suts_vectorcast import build_vectorcast_model

    model = build_vectorcast_model(str(exported_semantic_workbook))
    cases = model["units"][0]["test_cases"]
    assert cases
    for case in cases:
        assert "pt_Out" not in case["expected"]
        evidence = case["expected_evidence"]["pt_Out"]
        assert evidence["status"] == "unknown"
        assert evidence["execution_status"] == "not_run"


def test_td_p1_export_unproven_numeric_expectation_requires_verification(exported_semantic_workbook):
    import openpyxl

    from tools.export_suts_vectorcast import build_vectorcast_model

    path = exported_semantic_workbook
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Test Evidence"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    selected = next(row for row in range(2, sheet.max_row+1)
                    if sheet.cell(row, headers["Status"]).value == "derived")
    sequence = sheet.cell(selected, headers["Sequence"]).value
    sheet.cell(selected, headers["Status"]).value = "unknown"
    sheet.cell(selected, headers["Oracle"]).value = "none"
    workbook.save(path)
    workbook.close()
    model = build_vectorcast_model(str(path))
    case = next(case for case in model["units"][0]["test_cases"] if case["sequence_no"] == sequence)
    assert case["expected_evidence"]["return"]["status"] == "unknown"
    assert case["expected"]["return"]["verification_required"] is True
    assert "unproven_expected_evidence" in {warning["code"] for warning in model["export_warnings"]}
