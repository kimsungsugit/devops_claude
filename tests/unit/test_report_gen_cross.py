# tests/unit/test_report_gen_cross.py
"""Verify report_gen modules import correctly and cross-references resolve."""

from __future__ import annotations

import importlib
import pytest


REPORT_GEN_MODULES = [
    "report_gen",
    "report_gen.source_parser",
    "report_gen.function_analyzer",
    "report_gen.requirements",
    "report_gen.uds_text",
    "report_gen.uds_generator",
    "report_gen.docx_builder",
    "report_gen.validation",
    "report_gen.utils",
]


class TestReportGenImports:
    @pytest.mark.parametrize("module_name", REPORT_GEN_MODULES)
    def test_module_imports(self, module_name):
        mod = importlib.import_module(module_name)
        assert mod is not None

    def test_shim_backward_compat(self):
        """The root report_generator.py shim should re-export key functions."""
        import report_generator

        assert hasattr(report_generator, "generate_uds_docx")
        assert hasattr(report_generator, "validate_uds_docx_structure")
        assert hasattr(report_generator, "generate_uds_preview_html")
        assert hasattr(report_generator, "generate_uds_source_sections")
        assert callable(report_generator.generate_uds_docx)

    def test_generators_shim_backward_compat(self):
        """The root sts_generator.py / suts_generator.py shims work."""
        import sts_generator
        import suts_generator

        assert hasattr(sts_generator, "generate_sts")
        assert hasattr(suts_generator, "generate_suts")
        assert callable(sts_generator.generate_sts)
        assert callable(suts_generator.generate_suts)

    def test_cross_module_function_access(self):
        """Functions that depend on cross-module imports resolve."""
        from report_gen.utils import _safe_dict, _safe_list, _fmt_bool

        assert _safe_dict(None) == {}
        assert _safe_dict({"a": 1}) == {"a": 1}
        assert _safe_list(None) == []
        assert _safe_list([1, 2]) == [1, 2]
        assert _fmt_bool(True) in ("True", "true", "Yes", "yes", "O", "TRUE", "YES")
        assert _fmt_bool(False) in ("False", "false", "No", "no", "X", "FALSE", "NO")

    def test_constants_available(self):
        """report.constants should be importable."""
        from report.constants import UDS_RULES, DEFAULT_TYPE_RANGES

        assert isinstance(UDS_RULES, (dict, list))
        assert isinstance(DEFAULT_TYPE_RANGES, dict)

    def test_c_parsing_available(self):
        """report.c_parsing functions should be importable."""
        from report.c_parsing import _strip_c_comments, _extract_c_prototypes

        assert callable(_strip_c_comments)
        assert callable(_extract_c_prototypes)

    def test_strip_c_comments(self):
        from report.c_parsing import _strip_c_comments

        code = "int x = 1; /* comment */ int y = 2;"
        result = _strip_c_comments(code)
        assert "comment" not in result
        assert "int x" in result
        assert "int y" in result

    def test_extract_c_prototypes(self):
        from report.c_parsing import _extract_c_prototypes

        code = "void foo(int x);\nint bar(void);\n"
        protos = _extract_c_prototypes(code)
        assert isinstance(protos, list)


class TestConfigFunctions:
    def test_resolve_oai_api_keys(self):
        """config.resolve_oai_api_keys replaces ENV: placeholders."""
        import os
        from config import resolve_oai_api_keys

        os.environ["_TEST_API_KEY"] = "test-secret-123"
        try:
            configs = [
                {"model": "test-model", "api_key": "ENV:_TEST_API_KEY"},
                {"model": "other", "api_key": "hardcoded-value"},
            ]
            resolved = resolve_oai_api_keys(configs)
            assert resolved[0]["api_key"] == "test-secret-123"
            assert resolved[1]["api_key"] == "hardcoded-value"
        finally:
            del os.environ["_TEST_API_KEY"]

    def test_load_oai_config_list_missing_file(self):
        from config import load_oai_config_list

        result = load_oai_config_list("/nonexistent/path/config.json")
        assert result == []


class TestValidationHelpers:
    def test_valid_call_names_filters_type_tokens(self):
        from report_gen.validation import _valid_call_names

        result = _valid_call_names(["U8", "U16", "if", "MotorPwm_SetRatio", "lin_checksum"])
        assert result == ["MotorPwm_SetRatio", "lin_checksum"]

    def test_has_doc_output_slot_only_marks_mutable_out_params(self):
        from report_gen.validation import _has_doc_output_slot

        assert _has_doc_output_slot("Foo(void)") is False
        assert _has_doc_output_slot("Foo(const uint8 *src)") is False
        assert _has_doc_output_slot("Foo(uint8 *dst)") is True
        assert _has_doc_output_slot("Foo(uint8 buf[8])") is True
        assert _has_doc_output_slot("Foo(void (*cb)(void))") is False

    def test_confidence_report_includes_evidence_mix_sections(self, tmp_path):
        from report_gen.validation import generate_asil_related_confidence_report

        payload = {
            "function_details_by_name": {
                "MotorDispatcher": {
                    "id": "SwUFn_0198",
                    "name": "MotorDispatcher",
                    "description": "Motor control function",
                    "description_source": "sds",
                    "description_source_detail": "hsis+sds_match",
                    "description_evidence_sources": ["sds", "code", "hsis"],
                    "asil": "A",
                    "asil_source": "sds",
                    "related": "SwTR_0608",
                    "related_source": "sds",
                    "related_source_detail": "hsis+sds",
                    "related_evidence_sources": ["sds", "hsis"],
                }
            }
        }

        out = tmp_path / "confidence.md"
        generate_asil_related_confidence_report(payload, str(out))
        text = out.read_text(encoding="utf-8")

        assert "## Description Evidence Mix" in text
        assert "## Related ID Evidence Mix" in text
        assert "## Description Source Detail" in text
        assert "## Related ID Source Detail" in text
        assert "## Operating Judgment" in text
        assert "Canonical policy: `doc-first` with `code/HSIS` as supporting evidence." in text

    def test_confidence_report_keeps_payload_source_when_docx_exists(self, tmp_path, monkeypatch):
        from report_gen.validation import generate_asil_related_confidence_report

        fake_docx = tmp_path / "generated.docx"
        fake_docx.write_bytes(b"stub")

        payload = {
            "function_details_by_name": {
                "MotorDispatcher": {
                    "id": "SwUFn_0198",
                    "name": "MotorDispatcher",
                    "description": "Motor control function",
                    "description_source": "sds",
                    "comment_description": "Comment text that should not override canonical payload source",
                    "asil": "A",
                    "asil_source": "sds",
                    "related": "SwTR_0608",
                    "related_source": "sds",
                }
            }
        }

        class _FakeDoc:
            pass

        monkeypatch.setattr("docx.Document", lambda *_args, **_kwargs: _FakeDoc())
        monkeypatch.setattr(
            "report_gen.validation._extract_function_info_from_docx",
            lambda _doc: {
                "SwUFn_0198": {
                    "id": "SwUFn_0198",
                    "name": "MotorDispatcher",
                    "description": "Reference description",
                    "related": "SwTR_9999",
                }
            },
        )

        out = tmp_path / "confidence_keep_payload.md"
        generate_asil_related_confidence_report(payload, str(out), str(fake_docx))
        text = out.read_text(encoding="utf-8")

        assert "- SDS: `1` / `1` (100.0%)" in text

    def test_remove_docx_paragraphs_keeps_na_in_function_info_tables(self):
        import docx
        from report_gen.docx_builder import _remove_docx_paragraphs

        doc = docx.Document()
        table = doc.add_table(rows=2, cols=3)
        for cell in table.rows[0].cells:
            cell.text = "[ Function Information ]"
        table.cell(1, 0).text = "Called Function"
        table.cell(1, 2).text = "N/A"

        _remove_docx_paragraphs(doc, ["N/A"])

        assert table.cell(1, 2).text == "N/A"


class TestRequirementsEnrichment:
    def test_related_and_prototype_assist_sds_match(self, monkeypatch):
        from report_gen.requirements import enrich_function_details_with_docs

        monkeypatch.setattr(
            "report_gen.requirements._build_req_map_from_doc_paths",
            lambda paths: {
                "swtr_0608": {"asil": "A", "related": "SwTR_0608"},
            },
        )
        monkeypatch.setattr(
            "report_gen.requirements._extract_sds_partition_map",
            lambda path: {
                "g_ap_motorctrl_func": {
                    "asil": "A",
                    "related": "SwTR_0608",
                    "description": "Motor control function",
                },
                "swcom_01": {
                    "asil": "QM",
                    "related": "SwTR_0001",
                    "description": "Fallback swcom row",
                },
            },
        )

        details = {
            "SwUFn_0198": {
                "id": "SwUFn_0198",
                "name": "MotorDispatcher",
                "module_name": "UnknownModule",
                "prototype": "void g_Ap_MotorCtrl_Func(void)",
                "related": "SwTR_0608",
                "asil": "TBD",
                "description": "",
            }
        }

        enrich_function_details_with_docs(
            details,
            [["SwCom_01", "", "SwUFn_0198", "MotorDispatcher"]],
            req_doc_paths=["srs.docx"],
            sds_doc_paths=["sds.docx"],
        )

        info = details["SwUFn_0198"]
        assert info["asil"] == "A"
        assert info["sds_match_key"] == "g_ap_motorctrl_func"
        assert info["sds_match_mode"] == "related_prototype"
        assert info["sds_match_scope"] == "function"
        assert info["mapping_confidence"] >= 0.8
