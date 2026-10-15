"""R88 — a reference SITS used as the template keeps none of its own test spec / integration strategy sheets.

The generator writes `3.SW Integration Test Spec` and `2.SW Integration Strategy` (the KJPDS02 reference's names). The
HDPDM01 reference names them `4.SW Integration Test Spec` · `3-1-1. / 3-1-2. SW Integration Strategy` ·
`3-2. Sw Integration Strategy`: with it as the template, the generated workbook still carried the reference's 77 test
cases in `4.SW…` beside ours (found reading the generated workbook back, 2026-10-15)."""
from __future__ import annotations

import openpyxl

from generators.sits import _drop_template_spec_sheets, generate_sits_xlsm
from tests.unit.test_sits_front_matter import _itcs


def _template(path, names):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    for name in names:
        ws = wb.create_sheet(name)
        ws["B7"] = "SwITC_OLD_REFERENCE"
    wb.create_sheet("Notes")["A1"] = "keep me"
    wb.save(path)
    return path


def _cells(ws):
    return {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}


def test_the_hdpdm01_style_reference_sheets_are_dropped(tmp_path):
    tpl = _template(tmp_path / "hd_tpl.xlsx", ["4.SW Integration Test Spec", "3-1-1. SW Integration Strategy",
                                              "3-1-2. SW Integration Strategy", "3-2. Sw Integration Strategy"])
    out = tmp_path / "out.xlsx"
    generate_sits_xlsm(str(tpl), _itcs(), str(out))
    wb = openpyxl.load_workbook(str(out))
    for gone in ("4.SW Integration Test Spec", "3-1-1. SW Integration Strategy", "3-1-2. SW Integration Strategy",
                 "3-2. Sw Integration Strategy"):
        assert gone not in wb.sheetnames, gone
    assert "3.SW Integration Test Spec" in wb.sheetnames and "Notes" in wb.sheetnames and "Cover" in wb.sheetnames
    assert not any("SwITC_OLD_REFERENCE" in _cells(wb[s]) for s in wb.sheetnames), "the reference's old test stayed"
    assert "SwITC_01" in _cells(wb["3.SW Integration Test Spec"])


def test_the_kjpds02_style_names_are_ours_and_kept_in_place(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.create_sheet("2.SW Integration Strategy")["A1"] = "Software Integration Strategy"
    wb.create_sheet("3.SW Integration Test Spec")["B7"] = "SwITC_OLD_REFERENCE"
    assert _drop_template_spec_sheets(wb) == []   # same names as ours: written over below, not dropped here
    tpl = tmp_path / "pv_tpl.xlsx"
    wb.save(tpl)
    out = tmp_path / "out.xlsx"
    generate_sits_xlsm(str(tpl), _itcs(), str(out))
    got = openpyxl.load_workbook(str(out))
    assert "2.SW Integration Strategy" in got.sheetnames and "SwITC_OLD_REFERENCE" not in _cells(got["3.SW Integration Test Spec"])


def test_the_dropped_sheets_are_recorded(tmp_path):
    tpl = _template(tmp_path / "hd_tpl.xlsx", ["4.SW Integration Test Spec", "3-2. Sw Integration Strategy"])
    stats: dict = {}
    generate_sits_xlsm(str(tpl), _itcs(), str(tmp_path / "out.xlsx"), strategy_context={"stats_out": stats})
    assert stats.get("template_sheets_dropped") == ["4.SW Integration Test Spec", "3-2. Sw Integration Strategy"]


def test_only_spec_and_strategy_sheets_match():
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    for name in ("Related_ID 확인", "Test Evidence", "Interface Evidence", "1.Introduction", "2.Test Environment",
                 "4.SW Integration Test Spec"):
        wb.create_sheet(name)
    assert _drop_template_spec_sheets(wb) == ["4.SW Integration Test Spec"]
    assert "Related_ID 확인" in wb.sheetnames and "2.Test Environment" in wb.sheetnames


def test_our_named_strategy_sheet_keeps_none_of_the_reference_blocks(tmp_path):
    # (리뷰 R88 W2) the KJPDS02 reference's `2.SW Integration Strategy` is reused by name — its old blocks must not stay
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    ws = wb.create_sheet("2.SW Integration Strategy")
    ws["A1"] = "Software Integration Strategy"
    for r in range(6, 60):
        ws.cell(row=r, column=1, value=f"SwIT_OLD_{r}")
        ws.cell(row=r, column=7, value="old leaf")
    ws.merge_cells("A40:C41")
    tpl = tmp_path / "pv_tpl.xlsx"
    wb.save(tpl)
    out = tmp_path / "out.xlsx"
    generate_sits_xlsm(str(tpl), _itcs(), str(out), strategy_context={
        "calls_map": {}, "file_of": {}, "depth_of": {}, "uds_related_map": {}, "design_ids": set(), "stats_out": {}},
        flows=[{"entry_fn": "Ap_Door_Run"}])
    got = openpyxl.load_workbook(str(out))["2.SW Integration Strategy"]
    blob = _cells(got)
    assert "Software Integration Strategy" in blob, "the heading rows stay"
    assert not any(v.startswith("SwIT_OLD_") for v in blob) and "old leaf" not in blob


def test_a_dropped_active_sheet_does_not_leave_the_workbook_opening_elsewhere(tmp_path):
    # (리뷰 R88 I1) the HDPDM01 reference opens on `4.SW…`, the sheet dropped
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    for name in ("1.Introduction", "4.SW Integration Test Spec", "Notes"):
        wb.create_sheet(name)
    wb.active = 2
    assert _drop_template_spec_sheets(wb) == ["4.SW Integration Test Spec"]
    assert wb.active.title == "Cover"


def test_the_drop_reaches_the_quality_report_and_the_disclosure(tmp_path):
    # (리뷰 R88 W1) reporting it is not reaching the report — run the whole generation
    from generators.sits import generate_sits
    from report_gen.generation_disclosures import build_disclosures

    src = tmp_path / "src"
    src.mkdir()
    (src / "a.c").write_text("void b(void);\nvoid a(void) { b(); }\n", encoding="utf-8")
    (src / "b.c").write_text("int g;\nvoid b(void) { g = 1; }\n", encoding="utf-8")
    tpl = _template(tmp_path / "hd_tpl.xlsx", ["4.SW Integration Test Spec", "3-2. Sw Integration Strategy"])
    res = generate_sits(str(src), str(tmp_path / "out.xlsx"), template_path=str(tpl))
    qr = res["quality_report"]
    assert qr["template_sheets_dropped"] == ["4.SW Integration Test Spec", "3-2. Sw Integration Strategy"]
    item = {i["key"]: i for i in build_disclosures("sits", qr)}["sits_template_sheets_dropped"]
    assert item["value"] == "2개" and "4.SW Integration Test Spec" in item["note"]
