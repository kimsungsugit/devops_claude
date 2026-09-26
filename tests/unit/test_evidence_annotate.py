"""R11 — every derived expected value gets an independent-check label in a copy of the SUTS; rows that are not derived
say so; the input workbook is never overwritten; the column header says it is not an execution."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

import evidence_annotate as ea  # noqa: E402

HEADER = ["Function ID", "Function", "Sequence", "Observable", "Expected", "Status", "Oracle", "Execution",
          "Source SHA256", "Source path", "Reason", "Test Case ID", "Source hash scope", "Inputs JSON", "Basis",
          "Assumptions"]


def _book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Evidence"
    ws.append(HEADER)
    ws.append(["F1", "f", 1, "g_o", 2, "derived", "source", "not_run", "h", "a.c", "", "TC_f", "", "{}", "assigned", ""])
    ws.append(["F1", "f", 2, "g_o", 3, "derived", "source", "not_run", "h", "a.c", "", "TC_f", "", "{}", "assigned", ""])
    ws.append(["F1", "f", 3, "g_p", "[검증 필요] x", "unknown", "none", "not_run", "", "a.c", "x", "TC_f", "", "{}", "", ""])
    wb.save(path)
    return str(path)


def test_rows_get_the_verdict_of_their_claim_and_non_derived_rows_a_dash(tmp_path):
    src = _book(tmp_path / "s.xlsx")
    labels = {("a.c", "f", "TC_f", 1, "g_o"): "clang_constexpr_agree",
              ("a.c", "f", "TC_f", 2, "g_o"): "unchecked:constexpr_limit"}
    counts = ea.annotate(src, str(tmp_path / "out.xlsx"), labels)
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx")["Test Evidence"]
    col = [c.value for c in ws[1]].index(ea.COLUMN) + 1
    assert [ws.cell(r, col).value for r in range(2, 5)] == ["clang_constexpr_agree", "unchecked:constexpr_limit", "—"]
    assert counts == {"clang_constexpr_agree": 1, "unchecked": 1, "—": 1}
    assert "not an execution" in ea.COLUMN
    assert [ws.cell(r, 8).value for r in range(2, 5)] == ["not_run"] * 3   # Execution is untouched


def test_the_verdicts_come_from_the_check_per_output(monkeypatch):
    import source_oracle_clang_check as socc
    claims = [{"key": ("a.c", "f", "TC_f", 1), "outputs": {"g_o": 2}}, {"key": ("a.c", "f", "TC_f", 2), "outputs": {"g_o": 3}}]

    def check(cs, clang):
        cs[0]["verdicts"] = {"g_o": "agree_with_stubbed_callees"}
        cs[1]["verdicts"] = {"g_o": "mismatch"}
        return {"claims": 2, "checked": 2, "agree": 1, "mismatch": 1, "eval_error": 0, "unchecked": 0}
    monkeypatch.setattr(socc, "_claims_from_xlsm", lambda x, r: (claims, {}))
    monkeypatch.setattr(socc, "check_claims", check)
    monkeypatch.setattr(ea, "_derived_outputs", lambda x: {})
    labels, report = ea.verdicts("x.xlsx", "root")
    assert labels == {("a.c", "f", "TC_f", 1, "g_o"): "clang_constexpr_agree_stubbed_callees",
                      ("a.c", "f", "TC_f", 2, "g_o"): "clang_mismatch"}


def test_the_input_workbook_is_never_the_output(tmp_path):
    src = _book(tmp_path / "s.xlsx")
    with pytest.raises(SystemExit, match="refusing to overwrite"):
        ea.main(["--xlsm", src, "--source-root", str(tmp_path), "--out", src])



# ── review round 2 ──────────────────────────────────────────────

def test_the_output_guard_refuses_a_hard_link_a_source_tree_and_another_kind(tmp_path):
    import os
    src = _book(tmp_path / "s.xlsm")
    tree = tmp_path / "src"
    tree.mkdir()
    link = tmp_path / "link.xlsm"
    os.link(src, link)
    for out, why in ((str(link), "refusing to overwrite"), (str(tree / "c.xlsm"), "inside read-only source"),
                     (str(tmp_path / "c.xlsx"), "must keep the input's kind")):
        with pytest.raises(SystemExit, match=why):
            ea.guard_output(out, src, str(tree))
    ea.guard_output(str(tmp_path / "c.xlsm"), src, str(tree))   # a sibling copy is fine


def test_the_copy_carries_a_legend():
    labels = [row[0] for row in ea.LEGEND]
    assert "clang_constexpr_agree_stubbed_callees" in labels and any("타깃 실행이 아니다" in r[1] for r in ea.LEGEND)


def test_a_row_from_a_changed_source_says_so(monkeypatch):
    """(W4) the check skipped a sequence whose file hash is not the current text: its rows are not "not in check"."""
    import source_oracle_clang_check as socc
    monkeypatch.setattr(socc, "_claims_from_xlsm", lambda x, r: ([], {"skipped_sequences": {
        ("a.c", "f", "TC_f", 1): "source_changed"}, "skipped_reasons": {"source_changed": 1}}))
    monkeypatch.setattr(socc, "check_claims", lambda cs, clang: {"claims": 0, "checked": 0, "agree": 0,
                                                                  "mismatch": 0, "eval_error": 0, "unchecked": 0})
    monkeypatch.setattr(ea, "_derived_outputs", lambda x: {("a.c", "f", "TC_f", 1): ["g_o", "g_p"]})
    labels, report = ea.verdicts("x.xlsx", "root")
    assert labels == {("a.c", "f", "TC_f", 1, "g_o"): "unchecked:source_changed",
                      ("a.c", "f", "TC_f", 1, "g_p"): "unchecked:source_changed"}
    assert report["skipped_reasons"] == {"source_changed": 1}


# ── review round 3 ──────────────────────────────────────────────

def test_the_claims_compare_the_recorded_hash_with_the_real_text(tmp_path):
    """(W-b) not a mock: a CRLF file is the text its LF hash names, another hash is ``source_changed`` and a path
    outside the tree is ``source_missing``."""
    import hashlib

    import source_oracle_clang_check as socc
    tree = tmp_path / "src"
    tree.mkdir()
    lf = "unsigned char g_o;\nvoid f(void) { g_o = 2U; }\n"
    (tree / "a.c").write_bytes(lf.replace("\n", "\r\n").encode("utf-8"))
    path = str((tree / "a.c").resolve())
    sha = hashlib.sha256(lf.encode()).hexdigest()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Evidence"
    ws.append(HEADER)
    for seq, file, digest in ((1, path, sha), (2, path, "0" * 64), (3, str(tmp_path / "gone.c"), sha)):
        ws.append(["F1", "f", seq, "g_o", 2, "derived", "source", "not_run", digest, file, "", "TC_f", "", "{}",
                   "assigned", ""])
    ws.append(["F1", "f", 4, "g_o", "[검증 필요] x", "unknown", "none", "not_run", "0" * 64, path, "x", "TC_f", "", "{}",
               "", ""])
    wb.save(tmp_path / "s.xlsx")
    claims, meta = socc._claims_from_xlsm(str(tmp_path / "s.xlsx"), str(tree))
    assert [c["key"] for c in claims] == [(path, "f", "TC_f", 1)]
    assert meta["skipped_sequences"] == {(path, "f", "TC_f", 2): "source_changed",
                                         (str(tmp_path / "gone.c"), "f", "TC_f", 3): "source_missing"}
    assert meta["derived_sequences"] == 3   # the unverified row is not a claim


def test_the_output_folder_must_exist_before_the_check_runs(tmp_path, monkeypatch):
    """(I5) a missing folder is refused up front — not after the whole clang run, at save time."""
    src = _book(tmp_path / "s.xlsm")
    monkeypatch.setattr(ea, "verdicts", lambda *a: pytest.fail("the check must not run"))
    with pytest.raises(SystemExit, match="folder does not exist"):
        ea.main(["--xlsm", src, "--source-root", str(tmp_path / "src"), "--out", str(tmp_path / "no" / "c.xlsm")])
