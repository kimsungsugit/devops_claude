"""R67 N79 — STS 의 Test Method / Gen. Method 라벨은 스텝이 증명해야 한다.

회귀 대상(실측 KJPDS02_PV 2026-09-14, TC 294):
  - 라벨이 요구 종류·함수 입력 타입에서 미리 정해져 스텝과 무관했다. BAA 51건 중
    경계값 스텝(경계 최솟값/최댓값·유효 범위 초과의 숫자 입력)이 있는 것 6건, AOR 인데
    경계값 스텝이 있는 것 27건, FIT 213건 중 고장 주입 스텝이 없는 것 202건
    (`validate_sts_xlsm` 의 `label_audit` 로 잰 값 — 모듈 주석과 같은 수치).
  - 2026-08-11 어휘 정규화(RVW→RBT) 뒤 리뷰 전용 TC 가 실행 시험으로 세어져
    executable_pct 가 항상 100 이었다(품질 DB: 08-11 이전 79.2/리뷰전용 15 → 이후 100/0).
"""
from __future__ import annotations

import openpyxl
import pytest

from generators.sts import (
    _COL_HEADERS,
    _HEADER_ROW,
    _SPEC_SHEET_NAME,
    STS_COL,
    _classify_steps,
    _relabel_from_steps,
    generate_quality_report,
    generate_sts_xlsm,
    generate_test_cases,
    generate_traceability_matrix,
    validate_sts_xlsm,
)


def _s(action: str, expected: str = "ok") -> dict:
    return {"action": action, "expected": expected}


# ── 1. 분류기 ────────────────────────────────────────────────────────────────

class TestClassifySteps:
    def test_numeric_boundary_input_is_baa(self):
        steps = [_s("입력 설정 (경계 최솟값): x=0, y=-128"), _s("f() 호출"), _s("경계값 출력 확인")]
        assert _classify_steps(steps) == ("RBT", "BAA", False)

    def test_boundary_label_without_numbers_is_not_baa(self):
        """타입을 몰라 값을 못 만든 자리("경계 최솟값: param")는 경계값 분석이 아니다."""
        steps = [_s("입력 설정 (경계 최솟값): [IN] void *p"), _s("f() 호출")]
        assert _classify_steps(steps)[1] == "AOR"

    def test_out_of_range_numeric_input_is_fit_and_baa(self):
        """유효 범위 밖 값은 고장 주입이자 경계값 분석 — SwUTS 의 BV_MAX_INV→FI/ABV 와 같은 규칙."""
        steps = [_s("입력 설정 (유효 범위 초과): x=256"), _s("f() 호출"), _s("에러 처리 / 포화 출력 확인")]
        assert _classify_steps(steps) == ("FIT", "BAA", False)

    def test_error_condition_step_is_fit(self):
        steps = [_s("에러 조건 설정: ( ptr == NULL )"), _s("에러 처리 결과 확인")]
        assert _classify_steps(steps) == ("FIT", "AOR", False)

    @pytest.mark.parametrize("action", [
        "조건 충족 설정: ( a > 0 )",
        "조건 미충족 설정: NOT ( a > 0 )",
        "else-if 조건 설정: ( a == 2 )",
    ])
    def test_branch_partition_is_eca(self, action):
        assert _classify_steps([_s(action, "조건 분기 → True 경로 진입")]) == ("RBT", "ECA", False)

    def test_switch_case_is_eca_via_expected(self):
        assert _classify_steps([_s("mode = 3 설정", "switch 분기 → case 3 진입")])[1] == "ECA"

    def test_assign_step_is_not_partition(self):
        """`var = val 설정 확인`(대입)은 분기가 아니다 — switch case 와 글자가 비슷하다."""
        assert _classify_steps([_s("cnt = 0 설정 확인", "cnt 값 변경 정상")])[1] == "AOR"

    def test_boundary_wins_over_partition(self):
        steps = [_s("조건 충족 설정: ( a > 0 )"), _s("입력 설정 (경계 최댓값): a=255")]
        assert _classify_steps(steps)[1] == "BAA"

    def test_review_steps_are_flagged(self):
        steps = [_s("소스 코드에서 해당 요구사항 구현부 확인"), _s("요구사항 내용 리뷰: 전원 관리")]
        assert _classify_steps(steps) == ("RBT", "AOR", True)

    def test_plain_call_is_rbt_aor(self):
        steps = [_s("f() 호출"), _s("함수 반환값 확인"), _s("출력/반환값 확인")]
        assert _classify_steps(steps) == ("RBT", "AOR", False)

    def test_empty_and_malformed_steps(self):
        assert _classify_steps([]) == ("RBT", "AOR", False)
        assert _classify_steps([None, {}, {"action": None}])[0] == "RBT"


# ── 2. 생성 경로 — 라벨은 최종 스텝을 따른다 ─────────────────────────────────

class TestGeneratedLabelsFollowSteps:
    @staticmethod
    def _run(fd, req=None, cfg=None):
        req = req or {"id": "SwTR_0001", "req_type": "TSR", "asil": "B"}
        return generate_test_cases([req], fd, {req["id"]: list(fd)}, cfg or {"max_tc_per_req": 9})

    def test_typed_inputs_with_branch_flow_are_not_baa(self):
        """옛 휴리스틱: `u8` 입력이 있으면 BAA, TSR 이면 FIT. 스텝엔 경계값도 고장도 없다."""
        fd = {"F1": {"id": "F1", "name": "fn", "inputs": ["[IN] u8 mode"], "output": "u8",
                     "logic_flow": [{"type": "if", "condition": "( mode == 1 )",
                                     "true_body": [{"type": "call", "name": "sub"}],
                                     "false_body": []}]}}
        tcs = self._run(fd)
        assert tcs, "TC 가 있어야 한다"
        for tc in tcs:
            m, g, _ = _classify_steps(tc["steps"])
            assert (tc["test_method"], tc["gen_method"]) == (m, g), tc
        assert {tc["gen_method"] for tc in tcs} == {"ECA"}
        assert {tc["test_method"] for tc in tcs} == {"RBT"}

    def test_guard_branch_produces_fit_only_on_error_path_tc(self):
        fd = {"F1": {"id": "F1", "name": "fn", "inputs": [], "output": "void",
                     "logic_flow": [{"type": "if", "condition": "( ptr == NULL )",
                                     "true_body": [{"type": "return", "value": "E_NOT_OK"}],
                                     "false_body": [{"type": "call", "name": "work"}]}]}}
        tcs = self._run(fd, req={"id": "SwTR_0002", "req_type": "TR"})
        by_method = {tc["test_method"] for tc in tcs}
        assert by_method == {"RBT", "FIT"}, [(t["test_method"], t["steps"]) for t in tcs]
        fit = [tc for tc in tcs if tc["test_method"] == "FIT"]
        assert all(any(s["action"].startswith("에러 조건 설정:") for s in tc["steps"]) for tc in fit)

    def test_no_flow_simple_steps_get_content_labels(self):
        """logic_flow 가 없는 함수: TC1 정상 → RBT/AOR, TC2 경계 → RBT/BAA, TC3 초과 → FIT/BAA."""
        fd = {"F1": {"id": "F1", "name": "fn", "inputs": ["[IN] U8 level"], "output": "U8",
                     "logic_flow": []}}
        tcs = self._run(fd, req={"id": "SwTR_0003", "req_type": "TR"})
        labels = [(tc["test_method"], tc["gen_method"]) for tc in tcs]
        assert labels == [("RBT", "AOR"), ("RBT", "BAA"), ("FIT", "BAA")], labels

    def test_review_only_tc_is_flagged_and_labelled_rbt_aor(self):
        req = {"id": "SwNTR_0201", "req_type": "NTR", "description": "문서화 요구"}
        tcs = generate_test_cases([req], {}, {})
        assert len(tcs) == 1
        assert tcs[0]["review_only"] is True
        assert (tcs[0]["test_method"], tcs[0]["gen_method"]) == ("RBT", "AOR")

    def test_function_tc_is_not_review_only(self):
        fd = {"F1": {"id": "F1", "name": "fn", "inputs": [], "output": "void", "logic_flow": []}}
        tcs = self._run(fd, req={"id": "SwTR_0004", "req_type": "TR"})
        assert all(tc["review_only"] is False for tc in tcs)


# ── 3. 스텝이 바뀌면 라벨도 다시 읽는다(AI 보강 경로) ─────────────────────────

class TestRelabel:
    def test_relabel_follows_replaced_steps(self):
        tcs = [
            {"id": "A", "test_method": "FIT", "gen_method": "BAA", "steps": [_s("f() 호출")]},
            {"id": "B", "test_method": "RBT", "gen_method": "AOR",
             "steps": [_s("입력 설정 (경계 최댓값): a=255")]},
        ]
        _relabel_from_steps(tcs)
        assert (tcs[0]["test_method"], tcs[0]["gen_method"], tcs[0]["review_only"]) == ("RBT", "AOR", False)
        assert (tcs[1]["test_method"], tcs[1]["gen_method"]) == ("RBT", "BAA")

    def test_relabel_is_idempotent(self):
        tcs = [{"id": "A", "steps": [_s("에러 조건 설정: x"), _s("입력 설정 (경계 최솟값): a=0")]}]
        _relabel_from_steps(tcs)
        once = dict(tcs[0])
        _relabel_from_steps(tcs)
        assert tcs[0] == once

    class _Stop(RuntimeError):
        def __init__(self, tcs):
            super().__init__("stop")
            self.tcs = tcs

    def test_review_only_is_downward_only(self):
        """리뷰 TC 는 함수가 없어 생긴 것 — 스텝 문구가 바뀌어도 실행 시험이 되지 않는다(리뷰 W1)."""
        tcs = [{"id": "A", "review_only": True, "steps": [_s("입력 설정 (경계 최댓값): a=255")]}]
        _relabel_from_steps(tcs)
        assert (tcs[0]["gen_method"], tcs[0]["review_only"]) == ("BAA", True)

    def test_ai_enhancement_skips_review_only_tcs(self, monkeypatch):
        """리뷰 TC 는 2~3 스텝이라 예전엔 항상 AI 후보 1순위였고, 자유 문구로 갈아 끼워지면
        리뷰 표지가 사라졌다. 후보에서 빠져야 하고 스텝은 원본 그대로여야 한다."""
        import json
        import sys

        import generators.sts as gsts

        monkeypatch.setattr(gsts, "_sts_ai_call_with_retry", lambda *a, **k: json.dumps({
            "steps": [{"action": "임의 문구", "expected": "임의 기대"}]}))
        monkeypatch.setitem(sys.modules, "workflow.ai", type(sys)("workflow.ai"))
        sys.modules["workflow.ai"].agent_call = lambda *a, **k: ""
        review = {"id": "R", "srs_id": "SwNTR_0201", "review_only": True,
                  "steps": [_s("소스 코드에서 해당 요구사항 구현부 확인"), _s("요구사항 내용 리뷰: x")]}
        fn = {"id": "F", "srs_id": "SwTR_0001", "review_only": False, "steps": [_s("f() 호출")]}
        gsts.enhance_test_cases_with_ai([review, fn], {}, {"model": "x"})
        assert review["steps"][0]["action"] == "소스 코드에서 해당 요구사항 구현부 확인"
        assert fn["steps"][0]["action"] == "임의 문구", "대조군: 함수 TC 는 보강돼야 한다"
        _relabel_from_steps([review, fn])
        cov = generate_traceability_matrix([review, fn], [{"id": "SwNTR_0201"}, {"id": "SwTR_0001"}])["coverage"]
        assert cov["executable_pct"] == 50.0 and cov["review_only_reqs"] == ["SwNTR_0201"]

    def test_generate_sts_relabels_after_ai_enhancement(self, tmp_path, monkeypatch):
        """배선 확인 — 헬퍼 단독 테스트는 호출부가 안 부르는 것을 못 본다.

        AI 가 리뷰 TC 의 스텝을 경계값 스텝으로 갈아 끼우면(이 테스트는 후보 필터를
        우회해 강제로 바꾼다), 매트릭스로 넘어가는 TC 의 라벨은 BAA 여야 하고
        review_only 는 **그대로 True** 여야 한다(하향 전용).
        """
        import generators.sts as gsts

        def _fake_ai(test_cases, *a, **k):
            for tc in test_cases:
                tc["steps"] = [_s("입력 설정 (경계 최댓값): a=255"), _s("f() 호출")]
            return test_cases

        def _capture(test_cases, reqs):
            raise self._Stop(test_cases)

        monkeypatch.setattr(gsts, "enhance_test_cases_with_ai", _fake_ai)
        monkeypatch.setattr(gsts, "generate_traceability_matrix", _capture)
        with pytest.raises(self._Stop) as exc:
            gsts.generate_sts(
                requirements_text=["SwNTR_0201: 문서화 요구"],
                function_details={},
                output_path=str(tmp_path / "o.xlsm"),
                ai_config={"model": "x"},
            )
        tcs = exc.value.tcs
        assert tcs, "TC 가 있어야 한다"
        assert {(t["test_method"], t["gen_method"], t["review_only"]) for t in tcs} == {("RBT", "BAA", True)}


# ── 4. 실행 시험 축은 review_only 플래그를 본다 ──────────────────────────────

class TestExecutableAxisUsesFlag:
    def test_review_only_flag_excludes_from_executable(self):
        tcs = [
            {"id": "T1", "srs_id": "R1", "test_method": "RBT", "review_only": True,
             "steps": [_s("소스 코드에서 해당 요구사항 구현부 확인")]},
            {"id": "T2", "srs_id": "R2", "test_method": "RBT", "review_only": False,
             "steps": [_s("f() 호출")]},
        ]
        cov = generate_traceability_matrix(tcs, [{"id": "R1"}, {"id": "R2"}])["coverage"]
        assert cov["pct"] == 100.0
        assert cov["executable_pct"] == 50.0
        assert cov["review_only_reqs"] == ["R1"]

    def test_legacy_rvw_label_still_excluded(self):
        """플래그가 없는 옛 TC dict: RVW 라벨만으로도 계속 비실행으로 센다(호환)."""
        tcs = [{"id": "T1", "srs_id": "R1", "test_method": "RVW", "steps": []}]
        cov = generate_traceability_matrix(tcs, [{"id": "R1"}])["coverage"]
        assert cov["executable_pct"] == 0.0

    def test_end_to_end_review_only_requirement_is_reported(self):
        """생성 → 매트릭스 → 품질 리포트: 리뷰 전용 요구가 다시 경고에 나와야 한다."""
        reqs = [{"id": "SwNTR_0201", "req_type": "NTR", "description": "문서화"},
                {"id": "SwTR_0001", "req_type": "TR"}]
        fd = {"F1": {"id": "F1", "name": "fn", "inputs": [], "output": "void", "logic_flow": []}}
        tcs = generate_test_cases(reqs, fd, {"SwTR_0001": ["F1"]})
        trace = generate_traceability_matrix(tcs, reqs)
        assert trace["coverage"]["review_only_reqs"] == ["SwNTR_0201"]
        assert trace["coverage"]["executable_pct"] == 50.0
        qr = generate_quality_report(tcs, trace)
        assert any("RVW" in w and "SwNTR_0201" in w for w in qr["coverage_warnings"]), qr["coverage_warnings"]


# ── 5. 검증기 — 산출물에서 라벨 셀과 스텝 셀을 대조한다 ───────────────────────

def _write_sheet(path, blocks):
    """blocks: [(tc_id, method, gen, [(action, expected), ...])]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _SPEC_SHEET_NAME
    for col, label in enumerate(_COL_HEADERS, 1):
        ws.cell(row=_HEADER_ROW, column=col, value=label)
    r = _HEADER_ROW + 1
    for tc_id, method, gen, steps in blocks:
        for i, (a, e) in enumerate(steps):
            if i == 0:
                ws.cell(row=r, column=STS_COL["tc_id"], value=tc_id)
                ws.cell(row=r, column=STS_COL["title"], value=tc_id)
                ws.cell(row=r, column=STS_COL["test_method"], value=method)
                ws.cell(row=r, column=STS_COL["gen_method"], value=gen)
                ws.cell(row=r, column=STS_COL["srs"], value="SwTR_0001")
            ws.cell(row=r, column=STS_COL["action"], value=a)
            ws.cell(row=r, column=STS_COL["expected"], value=e)
            r += 1
    wb.save(str(path))
    return str(path)


class TestValidatorLabelAudit:
    def test_mismatches_are_counted_and_warned(self, tmp_path):
        p = _write_sheet(tmp_path / "sts.xlsx", [
            ("T1", "FIT", "BAA", [("f() 호출", "ok"), ("출력 확인", "ok")]),          # BAA·FIT 근거 없음
            ("T2", "RBT", "AOR", [("입력 설정 (경계 최솟값): a=0", "ok")]),           # 경계값인데 AOR
            ("T3", "RBT", "ECA", [("f() 호출", "ok")]),                              # ECA 근거 없음
            ("T4", "RBT", "AOR", [("에러 조건 설정: x", "에러 처리 경로 진입")]),    # 고장인데 RBT
            ("T5", "FIT", "BAA", [("입력 설정 (유효 범위 초과): a=256", "ok")]),     # 정합
            ("T6", "RBT", "ECA", [("mode = 1 설정", "switch 분기 → case 1 진입")]),  # 정합(expected)
            ("T7", "RBT", "AOR", [("조건 충족 설정: ( a > 0 )", "조건 분기 → True 경로 진입")]),  # 분기인데 AOR
            ("T8", "RBT", "BAA", [("조건 충족 설정: ( a > 0 )", "ok"),
                                  ("입력 설정 (경계 최댓값): a=255", "ok")]),           # 분기+경계 → BAA 가 맞다
        ])
        v = validate_sts_xlsm(p)
        a = v["stats"]["label_audit"]
        assert a == {
            "baa_total": 3, "baa_without_boundary": 1,
            "eca_total": 2, "eca_without_partition": 1,
            "fit_total": 2, "fit_without_fault": 1,
            "boundary_without_baa": 1, "partition_without_eca": 1, "fault_without_fit": 1,
        }, a
        hit = [w for w in v["warnings"] if w.startswith("라벨↔스텝 불일치")]
        assert len(hit) == 1
        assert "BAA 인데 경계값 스텝 없음 1건" in hit[0] and "고장 주입 스텝이 있는데 FIT 아님 1건" in hit[0]
        assert "분기 스텝이 있는데 ECA 아님 1건" in hit[0]

    def test_consistent_sheet_has_no_warning(self, tmp_path):
        p = _write_sheet(tmp_path / "ok.xlsx", [
            ("T1", "RBT", "AOR", [("f() 호출", "ok")]),
            ("T2", "FIT", "BAA", [("입력 설정 (유효 범위 초과): a=256", "ok")]),
        ])
        v = validate_sts_xlsm(p)
        assert not [w for w in v["warnings"] if "불일치" in w]
        assert sum(n for k, n in v["stats"]["label_audit"].items() if "without" in k) == 0

    def test_writer_to_validator_roundtrip_has_zero_mismatch(self, tmp_path):
        """생성기가 붙인 라벨이 라이터를 거쳐 파일에 남고, 검증기가 그 파일에서 0건 불일치를 봐야 한다."""
        reqs = [{"id": "SwTR_0001", "req_type": "TR"}, {"id": "SwNTR_0201", "req_type": "NTR"}]
        fd = {
            "F1": {"id": "F1", "name": "fa", "inputs": ["[IN] U8 level"], "output": "U8", "logic_flow": []},
            "F2": {"id": "F2", "name": "fb", "inputs": [], "output": "void",
                   "logic_flow": [{"type": "if", "condition": "( p == NULL )",
                                   "true_body": [{"type": "return", "value": "1"}], "false_body": []}]},
        }
        tcs = generate_test_cases(reqs, fd, {"SwTR_0001": ["F1", "F2"]}, {"max_tc_per_req": 9})
        assert {tc["gen_method"] for tc in tcs} >= {"AOR", "BAA", "ECA"}
        out = generate_sts_xlsm(None, tcs, {"coverage": {}}, str(tmp_path / "rt.xlsx"), {"project_id": "T"})
        v = validate_sts_xlsm(out)
        a = v["stats"]["label_audit"]
        assert v["stats"]["tc_count"] == len(tcs)
        assert a["baa_total"] >= 1 and a["fit_total"] >= 1 and a["eca_total"] >= 1, a
        assert sum(n for k, n in a.items() if "without" in k) == 0, a
