"""R3 (P0) — offline alignment of a reference SUTS with the current source (`scripts/reference_alignment.py`).

No SVN access: the reference's vectors are re-evaluated on the current source by the source oracle, slot by slot.
The workbook here uses the HDPDM01 reference layout (header rows 5-6, TC ID in C, sequence number in M, inputs from N,
expected results from BK, Related ID in ES) — the same layout `tools/export_suts_vectorcast` parses.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

import reference_alignment  # noqa: E402
from reference_alignment import align, reference_value  # noqa: E402

COMMON = """#ifndef COMMON_H
#define COMMON_H
typedef unsigned char U8;
typedef unsigned int U16;
typedef signed int S16;
#endif
"""
UNIT = """#include "common.h"
U8 g_i;
U8 g_o;
U8 g_x;
void f(void) { g_o = (U8)(g_i + 1U); }
void h(void) { g_o = (U8)(g_x * 2U); }
void k(void) { g_o = g_i; }
void e(void) { g_x = 1U; }
"""


def _reference(path: Path, blocks):
    """blocks: [(tc_id, signature, inputs, outputs, rows)] with rows = [(seq, {input: v}, {output: v})]."""
    from openpyxl.utils import column_index_from_string as ci
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2.SW Unit Test Spec"
    ws.cell(5, ci("C"), "Test Case")
    ws.cell(5, ci("N"), "Input")
    ws.cell(5, ci("BK"), "Expected Result")
    ws.cell(5, ci("ES"), "Related ID")
    ws.cell(6, ci("C"), "TC ID")
    row = 7
    for tc_id, signature, inputs, outputs, rows in blocks:
        ws.cell(row, ci("C"), tc_id)
        ws.cell(row, ci("D"), signature)
        for j, name in enumerate(inputs):
            ws.cell(row, ci("N") + j, name)
        for j, name in enumerate(outputs):
            ws.cell(row, ci("BK") + j, name)
        for seq, ins, outs in rows:
            row += 1
            ws.cell(row, ci("M"), seq)
            for j, name in enumerate(inputs):
                if name in ins:
                    ws.cell(row, ci("N") + j, ins[name])
            for j, name in enumerate(outputs):
                if name in outs:
                    ws.cell(row, ci("BK") + j, outs[name])
        row += 1
    wb.save(path)


@pytest.fixture()
def tree(tmp_path):
    root = tmp_path / "src"
    root.mkdir()
    (root / "common.h").write_text(COMMON, encoding="utf-8", newline="\n")
    (root / "unit.c").write_text(UNIT, encoding="utf-8", newline="\n")
    return root


@pytest.mark.parametrize("cell, value", [
    (5, 5), ("0x1", 1), ("0xFFU", 255), ("-1", -1), ("12UL", 12), (3.0, 3),
    ("-", None), ("N/A", None), (2.5, None), ("REG.Bits.P3", None), (True, None),
])
def test_reference_cells_are_read_as_integers_only_when_they_state_one(cell, value):
    assert reference_value(cell) == value


def test_agreeing_disagreeing_and_incomparable_slots_are_kept_apart(tmp_path, tree):
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [
        ("SwUTC_1", "void f( void )", ["g_i"], ["g_o", "REG.Bits.P3"], [
            (1, {"g_i": 1}, {"g_o": 2, "REG.Bits.P3": 1}),        # agree · register field not comparable
            (2, {"g_i": -1}, {"g_o": "0x0"}),                      # -1 into U8 holds 255: 255 + 1 → (U8)0 — agree
            (3, {"g_i": 5}, {"g_o": 7}),                           # the source computes 6 — disagree
        ]),
        ("SwUTC_2", "void h( void )", ["g_i"], ["g_o"], [
            (1, {"g_i": 1}, {"g_o": 4}),                           # g_x is not set by the reference — unknown
        ]),
        ("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9})]),
        ("SwUTC_4", "void gone( void )", [], ["g_o"], [(1, {}, {"g_o": 0})]),
    ])
    report = align(str(ref), [tree])
    by = {f["function"]: f for f in report["functions"]}
    assert by["f"]["status"] == "divergent"
    assert by["f"]["slots"] == {"agree": 2, "not_comparable": 1, "disagree": 1}
    (d,) = by["f"]["disagreements"]
    assert (d["sequence"], d["reference"], d["source"]) == (3, 7, 6)
    # the -1 was converted — and says what into
    assert by["f"]["input_conversions"] == [{"input": "g_i", "reference": -1, "held": 255, "count": 1}]
    assert by["f"]["agree_basis"] == {"assigned": 2}
    assert by["h"]["status"] == "unknown" and "initial_value_not_in_inputs" in by["h"]["unknown_reasons"]
    assert by["k"]["status"] == "aligned"
    assert by["gone"]["status"] == "not_in_source"
    assert report["summary"]["aligned_functions"] == ["k"] and report["summary"]["divergent_functions"] == ["f"]


def test_an_agreement_on_an_input_the_function_never_writes_is_echo_only(tmp_path, tree):
    ref = tmp_path / "ref.xlsm"
    # e() writes g_x only; the reference states g_o, which e() leaves as the input set it — nothing about the code
    _reference(ref, [("SwUTC_5", "void e( void )", ["g_o"], ["g_o"], [(1, {"g_o": 4}, {"g_o": 4})])])
    (rec,) = align(str(ref), [tree])["functions"]
    assert rec["status"] == "echo_only" and rec["agree_basis"] == {"unchanged_input": 1}


def test_a_unit_without_expected_values_is_not_an_unknown_without_reason(tmp_path, tree):
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {})])])
    report = align(str(ref), [tree])
    assert report["functions"][0]["status"] == "no_reference_expectation"
    assert report["summary"]["status"] == {"no_reference_expectation": 1}


def test_a_symbolic_reference_value_the_oracle_computes_is_not_a_disagreement(tmp_path, tree, monkeypatch):
    # the oracle determines g_o but the reference cell is text: not comparable, never ``None != value`` (review W6)
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": "see note"})])])
    (rec,) = align(str(ref), [tree])["functions"]
    assert rec["slots"] == {"not_comparable": 1} and rec["unknown_reasons"] == {"reference_value_not_integer": 1}
    assert rec["status"] == "unknown"


def _twin(tmp_path, body):
    boot = tmp_path / "boot"
    boot.mkdir()
    (boot / "common.h").write_text(COMMON, encoding="utf-8", newline="\n")
    (boot / "boot.c").write_text('#include "common.h"\nU8 g_i;\nU8 g_o;\nU8 g_x;\n' + body, encoding="utf-8",
                                 newline="\n")
    return boot


def test_a_definition_with_nothing_comparable_cannot_claim_the_unit_and_hide_a_disagreement(tmp_path, tree):
    # review C1 scenario A: boot's k() is not comparable (reads g_x, which the reference never sets); the app's k()
    # agrees on one slot and disagrees on another — the unit is divergent, not "unknown" through the silent twin
    boot = _twin(tmp_path, "void k(void) { g_o = g_x; }\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9}),
                                                                         (2, {"g_i": 3}, {"g_o": 4})])])
    (rec,) = align(str(ref), [tree, boot])["functions"]
    assert rec["status"] == "divergent" and rec["matched_definition"].endswith("unit.c")
    assert {c["path"][-6:]: (c["agree"], c["disagree"]) for c in rec["candidates"]} == {"unit.c": (1, 1),
                                                                                         "boot.c": (0, 0)}


def test_the_definition_with_more_agreements_wins_and_keeps_its_disagreement(tmp_path, tree):
    # review C1 scenario B: boot agrees once, the app agrees twice and disagrees once — the app owns the unit
    boot = _twin(tmp_path, "void k(void) { g_o = (U8)(g_i + 100U); }\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [
        (1, {"g_i": 1}, {"g_o": 1}), (2, {"g_i": 2}, {"g_o": 2}), (3, {"g_i": 0}, {"g_o": 100})])])
    (rec,) = align(str(ref), [tree, boot])["functions"]
    assert rec["matched_definition"].endswith("unit.c") and rec["status"] == "divergent"
    assert rec["slots"] == {"agree": 2, "disagree": 1}


def test_two_definitions_with_the_same_evidence_are_ambiguous_not_a_silent_pick(tmp_path, tree):
    boot = _twin(tmp_path, "void k(void) { g_o = g_i; }\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9})])])
    report = align(str(ref), [tree, boot])
    (rec,) = report["functions"]
    assert rec["status"] == "ambiguous_definition" and rec["matched_definition"] == ""
    assert report["summary"]["aligned_functions"] == []


def test_the_report_is_never_written_into_a_read_only_input(tmp_path, tree):
    ref = tmp_path / "docs" / "ref.xlsm"
    ref.parent.mkdir()
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9})])])
    for out in (tree / "out.json", ref.parent / "out.json"):
        with pytest.raises(SystemExit, match="read-only input"):
            reference_alignment.main(["--reference", str(ref), "--source-root", str(tree), "--out", str(out)])
        assert not out.exists()
    assert reference_alignment.main(["--reference", str(ref), "--source-root", str(tree),
                                     "--out", str(tmp_path / "out.json")]) == 0


def test_a_name_defined_twice_is_matched_to_the_definition_whose_values_the_reference_states(tmp_path, tree):
    boot = tmp_path / "boot"
    boot.mkdir()
    (boot / "common.h").write_text(COMMON, encoding="utf-8", newline="\n")
    (boot / "boot.c").write_text('#include "common.h"\nU8 g_i;\nU8 g_o;\nvoid k(void) { g_o = (U8)(g_i + 100U); }\n',
                                 encoding="utf-8", newline="\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 1}, {"g_o": 101})])])
    report = align(str(ref), [tree, boot])
    (rec,) = report["functions"]
    assert rec["status"] == "aligned" and rec["matched_definition"].endswith("boot.c")
    assert len(rec["defined_in"]) == 2 and len(rec["candidates"]) == 2


def test_clang_confirms_that_a_disagreement_is_about_the_source(tmp_path, tree):
    if shutil.which("clang") is None:
        pytest.skip("clang not installed")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_1", "void f( void )", ["g_i"], ["g_o"], [(3, {"g_i": 5}, {"g_o": 7})])])
    report = align(str(ref), [tree], clang=True)
    (d,) = report["functions"][0]["disagreements"]
    assert d["clang"] == "confirms_source_value"
    assert report["summary"]["clang"] == {"confirms_source_value": 1}


# ── review round 2 ──────────────────────────────────────────────────────────


def test_equal_agreement_never_lets_fewer_disagreements_pick_the_definition(tmp_path, tree):
    # W-A: boot agrees on seq 1 and cannot compare seq 2 (reads g_x); the app agrees on 1 and disagrees on 2
    boot = _twin(tmp_path, "void k(void) { g_o = (U8)(g_x + g_i - g_i); }\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i", "g_x"], ["g_o"], [(1, {"g_i": 9, "g_x": 9}, {"g_o": 9}),
                                                                             (2, {"g_i": 3}, {"g_o": 4})])])
    report = align(str(ref), [tree, boot])
    (rec,) = report["functions"]
    assert rec["status"] == "ambiguous_definition" and rec["matched_definition"] == ""
    assert [d["definition"][-6:] for d in rec["disagreements"]] == ["unit.c"]
    assert report["summary"]["aligned_functions"] == [] and report["summary"]["units_without_slot_totals"] == 1


def test_equal_agreement_where_every_definition_disagrees_is_divergent(tmp_path, tree):
    # W-B: the same body twice — whichever the reference belongs to, it differs from the source
    boot = _twin(tmp_path, "void k(void) { g_o = g_i; }\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9}),
                                                                         (2, {"g_i": 3}, {"g_o": 4})])])
    report = align(str(ref), [tree, boot])
    (rec,) = report["functions"]
    assert rec["status"] == "divergent" and len(rec["disagreements"]) == 2
    assert report["summary"]["divergent_functions"] == ["k"]


def test_definitions_with_nothing_comparable_are_unknown_not_ambiguous(tmp_path, tree):
    boot = _twin(tmp_path, "void k(void) { g_o = g_x; }\n")
    (tree / "unit.c").write_text(UNIT.replace("void k(void) { g_o = g_i; }", "void k(void) { g_o = g_x; }"),
                                 encoding="utf-8", newline="\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9})])])
    (rec,) = align(str(ref), [tree, boot])["functions"]
    assert rec["status"] == "unknown"
    assert set(rec["unknown_reasons_by_definition"]) == set(rec["defined_in"])
    assert all(r == {"initial_value_not_in_inputs": 1} for r in rec["unknown_reasons_by_definition"].values())


def test_a_symbolic_expected_value_is_counted_within_its_basis_not_as_another(tmp_path, tree):
    (tree / "unit.c").write_text(UNIT + "#define K_NINE 9U\n", encoding="utf-8", newline="\n")
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": "K_NINE"})])])
    report = align(str(ref), [tree])
    (rec,) = report["functions"]
    assert rec["status"] == "aligned" and rec["agree_basis"] == {"assigned": 1}
    assert rec["agree_via_reference_symbol"] == 1 and report["summary"]["agree_via_reference_symbol"] == 1


def test_the_unit_record_keys_by_tc_id_and_says_the_evidence_strength(tmp_path, tree):
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_1", "void f( void )", ["g_i"], ["g_o", "REG.Bits.P3"], [
        (1, {"g_i": 1}, {"g_o": 2, "REG.Bits.P3": 1}), (3, {"g_i": 5}, {"g_o": 7})])])
    (rec,) = align(str(ref), [tree])["functions"]
    assert rec["tc_id"].startswith("SwUTC_1") and rec["related_id"] != rec["tc_id"]
    assert rec["evidence_strength"] == round(1 / 3, 4)
    (d,) = rec["disagreements"]
    assert d["input_conversions"] == [] and d["inputs"] == {"g_i": 5}


def test_a_definition_that_raises_is_kept_as_an_error_and_fails_the_run(tmp_path, tree, monkeypatch):
    ref = tmp_path / "docs" / "ref.xlsm"
    ref.parent.mkdir()
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9})]),
                     ("SwUTC_1", "void f( void )", ["g_i"], ["g_o"], [(1, {"g_i": 1}, {"g_o": 2})])])
    real = reference_alignment._param_types

    def boom(unit):
        if unit["name"] == "k":
            raise RuntimeError("model defect")
        return real(unit)
    monkeypatch.setattr(reference_alignment, "_param_types", boom)
    report = align(str(ref), [tree])
    by = {f["function"]: f for f in report["functions"]}
    assert by["k"]["status"] == "error:RuntimeError" and "model defect" in by["k"]["failed_definitions"][0]["error"]
    assert by["f"]["status"] == "aligned"  # the run went on
    assert reference_alignment.main(["--reference", str(ref), "--source-root", str(tree),
                                     "--out", str(tmp_path / "o.json")]) == 1


def test_a_clang_agreement_through_stubbed_callees_is_not_an_independent_confirmation(monkeypatch):
    import source_oracle_clang_check as socc
    reports = iter([{"agree": 1, "agree_with_stubbed_callees": 1, "mismatch": 0, "eval_error": 0, "unchecked_reasons": {}},
                    {"agree": 1, "agree_with_stubbed_callees": 0, "mismatch": 0, "eval_error": 0, "unchecked_reasons": {}},
                    {"agree": 0, "mismatch": 1, "eval_error": 0, "unchecked_reasons": {}}])
    monkeypatch.setattr(socc, "check_claims", lambda claims, clang: next(reports))
    entries = [{}, {}, {}]
    verdicts = reference_alignment._clang_confirm([{"_entry": e} for e in entries], "clang")
    assert [e["clang"] for e in entries] == ["confirms_with_stubbed_callees", "confirms_source_value",
                                            "contradicts_oracle"]
    assert verdicts == {"confirms_with_stubbed_callees": 1, "confirms_source_value": 1, "contradicts_oracle": 1}


def test_the_report_may_not_replace_an_input_nor_a_source_root_itself(tmp_path, tree):
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [(1, {"g_i": 9}, {"g_o": 9})])])
    with pytest.raises(SystemExit, match="read-only input"):
        reference_alignment._guard_output(tree, str(ref), [tree])
    with pytest.raises(SystemExit, match="input of this run"):
        reference_alignment._guard_output(tmp_path / "gen.xlsm", str(tmp_path / "x" / "ref.xlsm"), [tree],
                                          inputs=(str(tmp_path / "gen.xlsm"),))


def test_a_unit_without_sequences_says_so(tmp_path, tree):
    ref = tmp_path / "ref.xlsm"
    _reference(ref, [("SwUTC_3", "void k( void )", ["g_i"], ["g_o"], [])])
    report = align(str(ref), [tree])
    assert [f["status"] for f in report["functions"]] == ["no_sequences"]
