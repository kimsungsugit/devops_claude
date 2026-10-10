"""R46 — a review item's missing part: looked for in the other documents first, and what filling it would improve.

2026-09-30 user direction: "부족한 건 클라우디움에서 파일들을 스캔하고 뽑아내봐. 그러고도 없으면 부족한 부분을 게이트나 결과에
표시해서 채워지면 개선된다는 것도 안내해야 해". Two parts:

* **evidence** (`review_evidence_corpus` / `attach_review_evidence`): the other sentences of the given SRS / SyRS / SyDS
  (all blocks, and a SyDS element's ``Input Information``) that state the item's value in the same unit **in a form the
  generator steps** — quoted as candidates (the same number may be another quantity), never used.
* **if filled** (`fill_guidance`): what writing the gap into the requirement produces — **tried in the item's own
  sentence** (review W1: a name written in front of the value is not enough in ``… 10ms 이상의 신호를 주입``, and in
  ``Pull-Up 전원 : 2.25V 이상 ~ Pull-Up 전원 2.75V 이하`` it even cost 2.25 V its step). A rewrite (a negation, a
  deadline, a label) gets an example that keeps the sentence's meaning (review W2).
"""
from __future__ import annotations

import re
from collections import Counter

import docx
import openpyxl
import pytest

from generators.sts import _build_tc_dict, _classify_steps, _make_tc_id
from generators.sts_requirement_tc import (
    REQUIREMENT_REVIEW_HEADERS,
    REQUIREMENT_REVIEW_SHEET,
    REVIEW_TEXT,
    append_requirement_boundary_tcs,
    boundary_steps,
    if_filled,
    parse_system_requirement_docx,
    traced_system_facts,
    write_requirement_review_sheet,
)


def _req(desc, rid="SwTR_0102", related="", verification=""):
    return {"id": rid, "name": "n", "description": desc, "verification": verification, "asil": "B",
            "related_id": related}


def _append(requirements, system=None):
    tcs, out = [], []

    def build(**kw):
        return _build_tc_dict(test_env="HILS", derive_inputs=False, is_safety=False, **kw)
    stats = append_requirement_boundary_tcs(tcs, requirements, build, _make_tc_id, _classify_steps, system=system,
                                            review_out=out)
    return stats, out, tcs


def _held(out):
    return [r for r in out if r["kind"] == "held_back"]


def _review(sentence):
    review = []
    boundary_steps(_req(sentence), Counter(), None, review)
    return review


def _stepped_values(sentence):
    return sorted(g["evidence"]["value"] for g in boundary_steps(_req(sentence)))


# ── evidence ──────────────────────────────────────────────────────────────────────────────────────────────────────

BATTERY = {"SyEI_01": {"doc": "SyDS", "name": "Battery Power", "fields": {"Range": "9 ~ 16V"}},
           # (R47) Input Information is a text field now — read as evidence as before, and stepped under a citing SRS
           "SyEL_05": {"doc": "SyDS", "name": "Power Supply",
                       "fields": {"Description": "전원 공급 블록",
                                  "Input Information": "BAT 전압(9V 이상 16.0V 이하), 5V, Buzzer chk"}}}


def test_a_range_cell_finds_the_same_range_written_as_two_comparisons_elsewhere():
    """HDPDM01 / KJPDS02: SyDS ``Range: 9 ~ 16V`` (block Battery Power) — a SyDS element's Input Information writes
    ``BAT 전압(9V 이상 16.0V 이하)``. ``16.0`` and ``16`` are one value."""
    stats, out, _ = _append([_req("- 동작", related="SyEI_01")], BATTERY)
    (item,) = _held(out)
    assert (item["reason"], item["fact_kind"], item["lo"], item["hi"], item["unit"]) == \
        ("no_subject_in_value_field", "range", "9", "16", "V")
    (cand,) = item["evidence"]
    assert cand["where"] == "SyDS SyEL_05 · Input Information"
    assert cand["condition"] == "BAT 전압 9V 이상 그리고 16.0V 이하"
    assert (cand["related"], cand["inclusion"], cand["weak_subject"], item["evidence_total"]) == \
        (False, "range", True, 1)
    assert (stats["review_evidence_searched"], stats["review_evidence_found"], stats["review_evidence_related"],
            stats["review_evidence_documents"]) == (1, 1, 0, ["SRS", "SyDS"])


# (R47) Input Information is no longer evidence only: it is stepped under a citing SRS requirement — the two R46
#   tests that fixed it apart (never a step · parsed apart) moved to tests/unit/test_sts_input_information_r47.py


def test_a_value_without_subject_finds_the_named_sentence_of_its_own_requirement_first():
    """HDPDM01 ``SwTSR_0101``: SyDS SySM_04 ``4.85V 미만`` names nothing; the requirement's own SRS text names
    ``u16g_ApiIn_HallSnsrLevel`` at 4.85 V written ``이하`` — the other inclusion (R33's word), tagged so."""
    system = {"SySM_04": {"doc": "SyDS", "name": "Sensor Monitor", "fields": {"Description": "- 4.85V 미만 시 정지"}}}
    reqs = [_req("- u16g_Other 이 4.85V 미만 시 경고", rid="SwTR_0201"),     # unrelated, same comparison
            _req("- u16g_ApiIn_HallSnsrLevel 이 4.85V 이하 또는 5.14V 이상인 상태", rid="SwTSR_0101", related="SySM_04")]
    _stats, out, _ = _append(reqs, system)
    (item,) = [r for r in _held(out) if r["source"].startswith("SyDS")]
    first, second = item["evidence"]
    assert (first["where"], first["related"], first["inclusion"]) == ("SRS SwTSR_0101", True, "flip")
    assert first["condition"].startswith("u16g_ApiIn_HallSnsrLevel 4.85V 이하")
    assert (second["where"], second["related"], second["inclusion"]) == ("SRS SwTR_0201", False, "same")


def test_the_other_side_is_told_apart_from_the_other_inclusion():
    """(review W3) ``3deg 이하`` ↔ ``Door Angle 3deg 초과`` is the other side (a complement or another condition), not
    the R33 "경계 포함 불일치" ``이하`` ↔ ``미만``."""
    reqs = [_req("- 3deg 이하 시 반전"),
            _req("- Door Angle 3deg 초과에서 끼임", rid="SwTR_0301"), _req("- Door Angle 3deg 미만에서 끼임", rid="SwTR_0302")]
    _stats, out, _ = _append(reqs)
    (item,) = _held(out)
    assert [(c["condition"], c["inclusion"]) for c in item["evidence"]] == [
        ("Door Angle 3deg 미만", "flip"), ("Door Angle 3deg 초과", "opposite")]
    wb = openpyxl.Workbook()
    write_requirement_review_sheet(wb, out)
    cell = list(wb[REQUIREMENT_REVIEW_SHEET].iter_rows(values_only=True))[1][7]
    assert "3deg 미만 [포함 다름]" in cell and "3deg 초과 [방향 반대]" in cell


def test_another_comparison_of_the_items_own_sentence_is_evidence():
    """(review W4) ``u16g_A 가 4.85V 미만 시 경고, 4.85V 미만 유지 시 정지``: the second value names nothing — the
    first, in the same sentence, is the strongest candidate (the R46 first cut excluded the whole sentence)."""
    _stats, out, _ = _append([_req("- u16g_A 가 4.85V 미만 시 경고, 4.85V 미만 유지 시 정지")])
    (item,) = _held(out)
    assert [(c["where"], c["condition"], c["related"]) for c in item["evidence"]] == [
        ("SRS SwTR_0102", "u16g_A 4.85V 미만", True)]


def test_one_quantity_in_two_units_is_one_value():
    """(review I1) ``3초`` and ``3s``, ``500ms`` and ``0.5s`` are the same number of seconds."""
    _stats, out, _ = _append([_req("- 3초 이상 끊길 경우 정지"), _req("- LIN 통신이 3s 이상 끊길 경우 정지", rid="SwTR_0200"),
                              _req("- 500ms 초과 시 정지", rid="SwTR_0300"), _req("- u16g_T 가 0.5s 초과 시 정지", rid="SwTR_0400")])
    by_fact = {r["fact"]: r for r in _held(out)}
    assert [c["where"] for c in by_fact["3초 이상"]["evidence"]] == ["SRS SwTR_0200"]
    assert [c["where"] for c in by_fact["500ms 초과"]["evidence"]] == ["SRS SwTR_0400"]


def test_an_unnamed_hold_time_is_evidence_only_within_the_items_own_requirement():
    reqs = [_req("- 300ms 이상 단절된 경우 전송", rid="SwTSR_0201"),
            _req("- 전원 이상인 상태로 특정시간(300ms) 이상 유지 시 정지", rid="SwTSR_0101"),       # unrelated, unnamed
            _req("- LIN 통신이 300ms 이상 끊길 경우 정지", rid="SwTSR_0207")]                       # unrelated, named
    _stats, out, _ = _append(reqs)
    (item,) = _held(out)
    assert [c["where"] for c in item["evidence"]] == ["SRS SwTSR_0207"]
    _stats, out, _ = _append([reqs[0], {**reqs[1], "id": "SwTSR_0201"}, reqs[2]])                   # now its own
    (item,) = _held(out)
    assert [(c["where"], c["related"]) for c in item["evidence"]] == [("SRS SwTSR_0201", True),
                                                                      ("SRS SwTSR_0207", False)]


def test_a_value_without_a_unit_is_not_searched_single_or_range():
    """(review I2) ``3`` and ``0 ~ 65535`` (the domain of every 16-bit variable) match everywhere."""
    _stats, out, _ = _append([_req("- u16g_A 와 3 초과 여부를 5 초과 조건과 비교"),
                              _req("u16g_B\t0 ~ 65535", rid="SwTR_0200"), _req("- u16g_C 가 3 초과", rid="SwTR_0300")])
    unitless = [r for r in _held(out) if not r.get("unit")]
    assert unitless and all(r["evidence"] is None for r in unitless)


def test_candidates_are_capped_and_counted():
    reqs = [_req("- 8.5V 미만 시 정지", rid="SwTR_0100")] + \
           [_req(f"- u16g_V{i} 가 8.5V 미만 시 경고", rid=f"SwTR_02{i:02d}") for i in range(6)]
    _stats, out, _ = _append(reqs)
    (item,) = _held(out)
    assert len(item["evidence"]) == 3 and item["evidence_total"] == 6


# ── if filled: tried in the item's own sentence ───────────────────────────────────────────────────────────────────

_NAMES = {"u16g_<신호>": "u16g_Drive", "<신호>": "구동전압", "<감시량>": "열림각"}


def _write_in(example):
    """The sentence the 'If Filled' example shows, with the placeholders replaced by names the test picks (not the
    ones `fill_guidance` tried)."""
    for placeholder, name in _NAMES.items():
        example = example.replace(placeholder, name)
    assert "<" not in example and "…" not in example, example
    return example


# (sentence, the item's fact, the values the filled sentence must step)
_IN_SENTENCE = [
    ("- 4.85V 미만 시 정지", "4.85V 미만", ["4.85"]),
    # KJPDS02_PV SySM_06 (review W1): a Korean name is overwritten by ``의 신호를 주입`` — a C identifier works
    ("- Watchdog Input 을 10ms 이상의 신호를 주입", "10ms 이상", ["10"]),
    # HDPDM01 / KJPDS02 SyDB_0104 (review r2 W7): a one-word name was read with ``따라`` — two words are not
    ("- 위치에 따라 3deg 이하이면 반전", "3deg 이하", ["3"]),
    # a lower and an upper bound: the value before names the quantity — join the two, do not name it again; neither
    #   names it — the name once, and the join
    ("- Sensor 전압 : 4.5V 이상 / 5.1V 이하", "5.1V 이하", ["4.5", "5.1"]),
    ("- 4.5V 이상 / 5.1V 이하 시 정상", "5.1V 이하", ["4.5", "5.1"]),
    # SyII_42 (review W1): 2.25 V must keep its step
    ("Pull-Up 전원 : 2.25V 이상 ~ 2.75V 이하", "2.75V 이하", ["2.25", "2.75"]),
    ("- u16s_V 가 9V 이상 16V 이상", "9V 이상", ["16", "9"]),
    ("- Param_X( 3도 ) 초과 시 반전", "Param_X( 3도 ) 초과", ["3"]),
    ("- 정상 작동 전압: 9.00V ~ 16.00V", "9.00V ~ 16.00V", ["16.00", "9.00"]),
    ("- 입력전원 9~16V 범위", "9~16V", ["16", "9"]),
    # (R47) ``RPM`` is a unit now; a word the extractor does not know as one (``Nm``) still follows the range
    ("- 회전수 0 ~ 5000 RPM", "0 ~ 5000 RPM", ["0", "5000"]),
    ("- 토크 0 ~ 50 Nm", "0 ~ 50 ", ["0", "50"]),
]


@pytest.mark.parametrize("sentence, fact, values", _IN_SENTENCE)
def test_the_fill_is_written_into_the_sentence_and_then_steps(sentence, fact, values):
    (item,) = [r for r in _review(sentence) if r["fact"] == fact]
    assert item["fill_check"] == "verified", item["if_filled"]
    assert "u16g_" not in item["if_filled"] or "10ms" in fact           # a C name only where Korean is not read
    example = re.search(r"예: '([^']+)'", item["if_filled"]).group(1)
    assert _stepped_values(_write_in(example)) == values



def test_a_range_example_keeps_the_space_before_what_follows():
    """(review r2 I-a) ``0 ~ 50 Nm`` (R47: was ``RPM``, a unit now): the range's match ends in the space — copied, the
    example must keep it."""
    (item,) = _review("- 토크 0 ~ 50 Nm")
    assert "0 이상이고 50 이하 Nm" in item["if_filled"]

def test_naming_the_value_again_would_have_cost_the_other_its_step():
    """(review W1) the tempting fill — the label's name written in front of 2.75 V — makes both values an unstated
    join; the guidance says to join them instead, and says why."""
    assert _stepped_values("Pull-Up 전원 : 2.25V 이상 ~ Pull-Up 전원 2.75V 이하") == []
    (item,) = _review("Pull-Up 전원 : 2.25V 이상 ~ 2.75V 이하")
    assert "이고 2.75V 이하" in item["if_filled"] and "신호 이름만 다시 적으면 결합 미기재로 두 값이 모두 보류된다" in item["if_filled"]


def test_a_fill_that_costs_another_value_its_step_or_names_nothing_is_refused():
    """The two checks of a tried fill: every value that stepped before still steps, and the written name is what the
    sentence then reads as the subject (``미만/대상신호`` is not ``대상신호``)."""
    from decimal import Decimal

    from generators.sts_requirement_tc import _try_fill
    v225, v275 = ("V", Decimal("2.25"), ">="), ("V", Decimal("2.75"), "<=")
    assert _try_fill("Pull-Up 전원 : 2.25V 이상 ~ Pull-Up 전원 2.75V 이하", [], Counter({v225: 1})) == \
        "같은 줄의 2.25V 이상 스텝이 사라짐"
    assert _try_fill("Pull-Up 전원 : 2.25V 이상이고 2.75V 이하", [(v275, "Pull-Up 전원")], Counter({v225: 1})) == ""
    # (review r2 I-c) counted: two equal comparisons before, one after, is a loss
    assert _try_fill("- u16g_A 가 3V 이상 시 경고", [], Counter({("V", Decimal("3"), ">="): 2})).endswith("스텝이 사라짐")
    why = _try_fill("- Port 를 4.85V 미만/채움 신호 5.15V 초과로", [(("V", Decimal("5.15"), ">"), "채움 신호")], Counter())
    assert why.startswith("적은 이름이 주어로 읽히지 않음(주어로 '미만/채움 신호'")


@pytest.mark.parametrize("sentence, fact", [
    ("- Pull-up Port를 4.85V 미만/5.15V초과로 Setting", "5.15V초과"),       # SySM_04: the other value names it
    ("- Port 를 4.85V 미만/5.15V 초과로 높여 Setting", "5.15V 초과"),        # SySM_04 Input 1: neither names it
])
def test_a_lower_bound_above_an_upper_one_is_a_meaning_to_decide(sentence, fact):
    """(review r2 W8, as R45 C1) outside a band or an entry and a return — the shape does not tell: the guidance says
    to decide, and shows the band reading only as tried in the sentence."""
    (item,) = [r for r in _review(sentence) if r["fact"] == fact]
    assert item["fill_check"] == "rewrite"
    assert "대역 밖(이거나)인지, 진입·복귀 같은 두 조건(히스테리시스" in item["if_filled"]
    example = re.search(r"대역 밖이면 예: '([^']+)'", item["if_filled"]).group(1)
    assert _stepped_values(_write_in(example)) == ["4.85", "5.15"]


def test_a_hysteresis_is_never_given_a_join():
    """(review r2 W8) ``8.5V 이하, 9.0V 이상 복귀`` — an entry and a return: no ``이고`` (never true) nor ``이거나``."""
    items = _review("- 저전압 판정: u16g_V 가 8.5V 이하, 9.0V 이상 복귀")
    assert items and all(r["fill_check"] == "rewrite" and "히스테리시스" in r["if_filled"]
                         and "이고" not in r["if_filled"].split("—")[0] for r in items)


def test_a_korean_name_that_the_sentence_does_not_read_is_said_so():
    (item,) = _review("Input : f/w 변경을 통해 Watchdog Input 을 10ms 이상의 신호를 주입")
    assert "u16g_<신호> 10ms 이상" in item["if_filled"] and "한국어 이름은 이 문장에서 주어로 읽히지 않는다" in item["if_filled"]


def test_what_a_name_alone_cannot_settle_says_what_else_must():
    (item,) = _review("- 8.5V 미만이 아닌 경우 정상")
    assert item["fill_check"] == "needs_more"
    assert "신호 이름을 적어도 스텝되지 않는다" in item["if_filled"] and "부정" in item["if_filled"]


# rewrites: the example keeps the meaning (review W2) — and its form steps with that meaning
@pytest.mark.parametrize("sentence, fact, op", [
    ("- u16g_V 가 8.5V 미만이 아닌 경우 정상", "8.5V 미만", ">="),                          # the complement
    ("- u16g_V 가 8.5V 이하가 아닐 경우 정상", "8.5V 이하", ">"),                           # (review r2 W5)
    ("- u16g_V 가 9V 이상을 벗어나면 정지", "9V 이상", "<"),
    ("- u16g_V 가 8.5V 이하가 되지 않으면 정상", "8.5V 이하", ">"),
    ("- u16g_T 가 300ms 이상 유지하지 않으면 경고", "300ms 이상 유지", "<"),                     # (review r3 W9)
    ("- LIN Signal 이 300ms 이상 수신되지 않는 경우 정지", "300ms 이상", ">="),             # the negation in the name
    # (review r2 W6) ``이내`` is ``이하``: the non-event outlasts it; the condition lasts at most that long
    ("- 20ms 이내에 Message Flag가 Set/Clear가 되지 않으면 LIN 통신 Error", "20ms 이내", ">"),
    ("- 저전압이 500ms 이내에 9V 이상으로 복귀할 경우 남은 동작을 수행한다", "500ms 이내", "<="),
    # (review r4 W10) every ending of a negated clause
    ("- 20ms 이내에 Message Flag가 변경되지 않은 경우 LIN 통신 Error", "20ms 이내", ">"),
    ("- 20ms 이내에 응답하지 않고 대기하면 Error", "20ms 이내", ">"),
    ("- 100ms 이내에 응답을 받지 못하는 경우 재전송", "100ms 이내", ">"),
    ("- 100ms 이내에 정상 상태가 아닐 경우 고장", "100ms 이내", ">"),
    # (review r3 W9) the adverb ``없이`` is no negative clause
    ("- 저전압이 500ms 이내에 이상 없이 9V 이상으로 복귀할 경우 남은 동작", "500ms 이내", "<="),
])
def test_a_rewrite_example_keeps_the_sentences_meaning(sentence, fact, op):
    (item,) = [r for r in _review(sentence) if r["fact"] == fact]
    assert item["fill_check"] == "rewrite"
    example = re.search(r"'([^']*<(?:신호|조건)>[^']*)'", item["if_filled"]).group(1)
    written = example.replace("<신호>", "구동전압").replace("<조건>", "구동전압")
    groups = boundary_steps(_req(f"- {written} 시 동작"))
    assert [(g["evidence"]["op"]) for g in groups] == [op], (written, groups)



def test_a_negation_the_guidance_cannot_place_gets_no_comparison_example():
    (item,) = _review("- u16g_V 가 8.5V 이하인 상태가 없으면 정상")
    assert item["reason"] == "negated_condition" and item["fill_check"] == "rewrite"
    assert "'<신호>" not in item["if_filled"] and "정해 부정 없이 적으면" in item["if_filled"]


@pytest.mark.parametrize("sentence", ["- u16s_V 가 9V 미만 9V 이상", "- u8g_X 가 10 초과 11 미만"])
def test_both_sides_of_one_value_or_an_empty_range_is_a_meaning_to_decide(sentence):
    """(review r3 I-g, as R45 I1) *and* never holds, *or* always does — no join is offered as a fill."""
    items = _review(sentence)
    assert items and all(r["fill_check"] == "rewrite" and "늘 거짓" in r["if_filled"] for r in items)


def test_a_reason_a_filled_sentence_is_still_held_for_is_said_in_words():
    """(review r3 I-h) not the code: ``u8g_X 45000 이상`` is outside the width the name declares."""
    from decimal import Decimal

    from generators.sts_requirement_tc import _try_fill
    why = _try_fill("- u8g_X 가 45000 이상 시 정지", [(("", Decimal(45000), ">="), None)], Counter())
    assert why == "값이 신호 이름이 선언한 폭 밖"

def test_every_review_reason_gets_guidance():
    """(review I10) a new review reason must not leave the column empty."""
    samples = {
        "no_subject_for_value": "- 4.85V 미만 시 정지",
        "subject_unclear": "- Port를 4.85V 미만/5.15V초과로 Setting",
        "kind_range": "- 입력전원 9~16V 범위",
        "parenthesis_labels_may_share_a_quantity": "- A조건(0.8m/s 미만) 또는 B조건(0.8m/s 이상, 2.0m/s 이상) 판단",
        "same_subject_combination_unstated": "- u16s_V 가 9V 이상 16V 이상",
        "negated_condition": "- u16g_V 가 8.5V 미만이 아닌 경우 정상",
        "monitored_quantity_unknown": "- Param_X( 3도 ) 초과 시 반전",
        "deadline_or_window": "- 20ms 이내에 Message Flag가 Set/Clear가 되지 않으면 LIN 통신 Error",
    }
    seen = {}
    for sentence in samples.values():
        for r in _review(sentence):
            seen.setdefault(r["reason"], r)
    system = {"SyII_42": {"doc": "SyDS", "name": "Sensor chk", "fields": {"Range": "2.75V 이하"}}}
    value_field = []
    boundary_steps(_req("- 동작", related="SyII_42"), Counter(), system, value_field)
    seen.setdefault(value_field[0]["reason"], value_field[0])
    assert set(seen) == set(REVIEW_TEXT)
    for reason, item in seen.items():
        assert item["fill_check"] in {"verified", "needs_more", "rewrite"} and item["if_filled"] not in {"", "—"}, reason


def test_readings_and_conflicts_get_their_own_guidance():
    assert "확인만" in if_filled({"kind": "read", "reason": "read_as_range"})
    assert "통일하면" in if_filled({"kind": "inclusion_conflict"})


def test_a_guidance_failure_never_costs_the_item(monkeypatch):
    import generators.sts_requirement_tc as rtc

    def boom(*a, **kw):
        raise RuntimeError("fill")
    monkeypatch.setattr(rtc, "fill_guidance", boom)
    (item,) = _review("- 4.85V 미만 시 정지")
    assert item["fill_check"] == "unchecked" and "안내를 만들지 못함" in item["if_filled"]


# ── the document, the report, the web, the gate ───────────────────────────────────────────────────────────────────

def test_the_review_sheet_carries_if_filled_and_the_evidence():
    stats, out, _ = _append([_req("- 동작", related="SyEI_01"), _req("- u16g_X 가 3 초과", rid="SwTR_0800"),
                             _req("u16g_B\t0 ~ 65535", rid="SwTR_0900")], BATTERY)
    out.append({"kind": "inclusion_conflict", "srs_id": "SwTSR_0104", "role": "condition", "same_subject": True,
                "a": {"source": "SRS", "subject": "저전압", "text": "8.5V 이하", "line": "a"},
                "b": {"source": "SRS", "subject": "저전압", "text": "8.5V 미만", "line": "b"}})
    wb = openpyxl.Workbook()
    assert write_requirement_review_sheet(wb, out) == len(out)
    rows = list(wb[REQUIREMENT_REVIEW_SHEET].iter_rows(values_only=True))
    head = list(rows[0])
    assert head == REQUIREMENT_REVIEW_HEADERS and head[6:8] == ["If Filled", "Evidence (Other Sentences)"]
    by_fact = {r[3]: r for r in rows[1:]}
    battery = by_fact["9 ~ 16V"]
    assert "두 비교로" in battery[6] and "이 문장에 넣어 확인" in battery[6]
    assert battery[7].startswith("SyDS SyEL_05 · Input Information: BAT 전압 9V 이상 그리고 16.0V 이하 [괄호 앞 명사]")
    assert "같은 양인지 확인" in battery[7]
    assert by_fact["0 ~ 65535"][7].startswith("탐색 안 함")
    assert "통일하면" in rows[-1][6] and rows[-1][7] == "—"
    assert stats["review_fill_checks"] == {"verified": 2}


def test_an_item_with_nothing_found_says_the_words_must_be_decided():
    stats, out, _ = _append([_req("- 8.5V 이하 시 정지")])
    wb = openpyxl.Workbook()
    write_requirement_review_sheet(wb, out)
    row = list(wb[REQUIREMENT_REVIEW_SHEET].iter_rows(values_only=True))[1]
    assert row[7].startswith("준 문서의 다른 문장에") and "적을 말부터 요구 문서에서 정해야 한다" in row[7]
    assert (stats["review_evidence_searched"], stats["review_evidence_found"]) == (1, 0)


def test_a_failed_evidence_search_changes_no_candidate_and_says_so_in_the_sheet(monkeypatch):
    import generators.sts_requirement_tc as rtc

    def boom(*a, **kw):
        raise RuntimeError("corpus")
    monkeypatch.setattr(rtc, "review_evidence_corpus", boom)
    stats, out, tcs = _append([_req("- 8.5V 이하 시 정지"), _req("- u16g_V 가 9V 초과", rid="SwTR_0200")])
    assert stats["review_evidence_error"] == "RuntimeError: corpus"
    assert all("evidence" not in r for r in out) and "review_evidence_searched" not in stats
    assert len(tcs) == 1                                                        # the boundary TCs are still there
    wb = openpyxl.Workbook()
    write_requirement_review_sheet(wb, out)
    row = list(wb[REQUIREMENT_REVIEW_SHEET].iter_rows(values_only=True))[1]
    assert row[7] == "근거 탐색 실패 — RuntimeError: corpus" and row[6]      # (review I4) not "—"; guidance stays


def _disclosure(rb):
    from report_gen.generation_disclosures import build_disclosures
    items = {i["key"]: i for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": rb}})}
    return items.get("sts_requirement_review")


_BASE = {"tcs": 1, "steps": 3, "facts_used": 1, "facts": 2, "review_item_count": 5,
         "review_by_reason": {"kind_range": 5}, "review_items": []}


def test_the_disclosure_says_what_was_found_how_the_guidance_was_checked_and_what_filling_improves():
    old = _disclosure(_BASE)                                  # an output from before R46: nothing said about evidence
    assert "근거 후보" not in old["value"] and "If Filled" not in old["note"]
    rb = {**_BASE, "review_evidence_searched": 4, "review_evidence_found": 3, "review_evidence_related": 1,
          "review_evidence_documents": ["SRS", "SyRS", "SyDS"],
          "review_fill_checks": {"needs_more": 1, "rewrite": 1, "verified": 3}}
    item = _disclosure(rb)
    assert item["value"] == "5건 · 다른 문장에 근거 후보 3/4"
    assert "이 문장에 넣어 확인 3 · 더 풀 것이 남음 1 · 다시 쓰기 예시 1" in item["note"]
    assert "준 문서(SRS·SyRS·SyDS)" in item["note"] and "같은 요구·인용 블록 1건" in item["note"]
    assert "후보 없는 1건은 적을 말부터 요구 문서에서 정해야 한다" in item["note"]
    tail = item["note"][-120:]                                # (review I13) the gate board keeps the last 120
    assert "채우면: 'If Filled' 칸대로 요구 문서에 적고 다시 생성하면 그 값의 경계 스텝이 생긴다(더 풀 것이 남은 항목은 그것까지)." in tail
    failed = _disclosure({**_BASE, "review_evidence_error": "RuntimeError: corpus"})
    assert "근거 후보 탐색은 실패했다 — RuntimeError: corpus" in failed["note"]


def test_nothing_searched_is_not_written_as_zero_found():
    """(review I3) no held-back item with a unit: no "0/0", no document list that was not read."""
    item = _disclosure({**_BASE, "review_evidence_searched": 0, "review_evidence_found": 0,
                        "review_evidence_related": 0, "review_evidence_documents": [],
                        "review_fill_checks": {"rewrite": 2}})
    assert item["value"] == "5건" and "찾지 않았다" in item["note"] and "0/0" not in item["note"]


def test_the_gate_board_says_what_to_do_with_review_items():
    from backend.services.issue_explainer import rule_action
    from workflow.quality.issues import _from_generation_disclosures
    rb = {**_BASE, "review_item_count": 57, "review_evidence_searched": 55, "review_evidence_found": 18,
          "review_evidence_related": 10, "review_evidence_documents": ["SRS", "SyRS", "SyDS"],
          "review_fill_checks": {"verified": 40, "needs_more": 5, "rewrite": 10},
          "review_items": [{"srs_ids": ["SwTSR_0101"], "source": "SyDS SySM_04 · Verification criteria",
                            "fact": "4.85V 미만", "reason": "no_subject_for_value"}] * 5}
    from report_gen.generation_disclosures import build_disclosures
    items = build_disclosures("sts", {"generation_stats": {"requirement_boundary": rb}})
    (issue,) = [i for i in _from_generation_disclosures({"present": True, "items": items})
                if i["code"] == "disclosure:sts_requirement_review"]
    assert "근거 후보 18/55" in issue["message"] and "그 값의 경계 스텝이 생긴다" in issue["message"]
    action = rule_action(issue["code"])
    assert "If Filled" in action and "이 문장에 넣어 확인" in action and "후보를 쓰지 않는다" in action


def test_the_readiness_gate_says_what_the_system_documents_add():
    from backend.services.docgen_requirements import IN_SYDS, IN_SYRS, requirements_for
    optional = requirements_for("sts")["optional"]
    assert "근거 후보" in optional[IN_SYRS] and "Input Information" in optional[IN_SYDS]
