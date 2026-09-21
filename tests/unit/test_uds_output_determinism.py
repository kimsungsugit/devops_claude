"""R47-k N27-e — 같은 입력의 두 run 이 다른 문서를 냈다.

실측(2026-09-14, run 2072/2073 — 같은 요청, 백엔드 재기동 사이): `.payload.json` function_details 1,157 중
132 항목이 달랐는데 **전부 `(idx: …)` 첨자 나열의 순서만** 달랐다(`(idx: 7, 2, 1, …)` ↔ `(idx: 0, 1, 4, …)`).
근본은 둘이다.

1. `function_analyzer._scan_name_usage` 는 첨자를 `set` 에 모으고, 소비처(`uds_generator`)는 바로 윗줄의
   `members` 는 `sorted()` 하면서 첨자는 set 그대로 순회했다(쌍둥이 한쪽). 문자열 해시 시드는 프로세스마다
   달라 순서가 **재기동마다** 바뀐다. → `_observed_index_values` 단일 출처(정규화→중복 제거→정렬).
2. python-docx 는 zip 멤버 2,008개 전부에 **저장 시각**을 박는다. 내용이 같아도 `output_sha256` 은 늘 달랐다
   (두 run 의 zip: 내용이 다른 멤버 1 · 시각이 다른 멤버 2,008). → `normalize_zip_member_times` 로 시각 고정.

가드는 세 층이다: 헬퍼 단위 · **PYTHONHASHSEED 를 바꾼 서브프로세스**(실제 비결정 축) · 빌더 산출물 바이트.
"""
from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import docx
import pytest

from report_gen import docx_builder
from report_gen.atomic_io import ZIP_FIXED_DATE_TIME, normalize_zip_member_times
from report_gen.function_analyzer import _observed_index_values

ROOT = Path(__file__).resolve().parents[2]


# --------------------------------------------------------------------------------------------
# 1. 헬퍼 — 입력 순서와 무관하게 같은 목록
# --------------------------------------------------------------------------------------------
class TestObservedIndexValues:
    _NUMS = ["7", "2", "1", "3", "4", "6", "0", "8", "5"]     # run 2072 가 실제로 낸 순서
    _WANT = ["0", "1", "2", "3", "4", "5", "6", "7", "8"]

    def test_order_does_not_depend_on_input_iteration(self):
        assert _observed_index_values(self._NUMS, {}) == self._WANT
        assert _observed_index_values(list(reversed(self._NUMS)), {}) == self._WANT
        assert _observed_index_values(set(self._NUMS), {}) == self._WANT

    def test_numbers_by_value_first_then_symbols_naturally(self):
        """`10` 이 `2` 뒤에(사전순이 아니라 값), 기호는 자연 정렬(`Idx2` < `Idx10`)."""
        got = _observed_index_values(["u8t_Idx10", "10", "u8t_Idx2", "2", "u8t_Cur", "0"], {})
        assert got == ["0", "2", "10", "u8t_Cur", "u8t_Idx2", "u8t_Idx10"]

    def test_duplicates_collapse_after_normalization(self):
        """`2U` · `( 2U )` · 매크로 `IDX_A(=2)` 는 같은 첨자다 — 정규화 **뒤에** 중복을 지운다."""
        assert _observed_index_values(["( 2U )", "2U", "IDX_A", "2"], {"IDX_A": "2"}) == ["2"]

    @pytest.mark.parametrize("raw", [None, set(), [], ["", "  "]])
    def test_empty_inputs_give_empty_list(self, raw):
        assert _observed_index_values(raw, {}) == []

    def test_symbolic_and_numeric_mixed_matches_live_shape(self):
        """run 2072 의 `u8s_TempLut (idx: u8t_CurIdx, 0, 4, u8t_NextIdx)` 가 정렬 뒤 어떻게 보이는가."""
        got = _observed_index_values({"u8t_CurIdx", "0", "4", "u8t_NextIdx"}, {})
        assert got == ["0", "4", "u8t_CurIdx", "u8t_NextIdx"]

    def test_natural_key_ties_are_broken_by_the_raw_text(self):
        """(리뷰 C1) `u8t_Ch1` 과 `u8t_Ch01` 은 자연 키가 같다 — 타이브레이커가 없으면 안정 정렬이 입력(=set 순회) 순서를
        그대로 내보내 고치려던 비결정이 그 자리에 남는다. 두 입력 순서 모두 같은 답이어야 한다."""
        a = _observed_index_values(["u8t_Ch1", "u8t_Ch01", "u8t_Ch001"], {})
        b = _observed_index_values(["u8t_Ch001", "u8t_Ch1", "u8t_Ch01"], {})
        assert a == b == ["u8t_Ch001", "u8t_Ch01", "u8t_Ch1"]


# --------------------------------------------------------------------------------------------
# 2. 프로세스 가드 — 해시 시드가 달라도 function_details 가 같다
# --------------------------------------------------------------------------------------------
_C_SRC = """
typedef unsigned char U8;
static U8 u8s_Lut[8];
U8 g_Tbl[16];
static U8 u8t_Cur;
static U8 u8t_Next;

U8 Sample(const U8 * Data, U8 * Out, U8 Sel)
{
    U8 v;
    v = Data[7] + Data[2] + Data[1] + Data[3] + Data[4] + Data[6] + Data[0] + Data[8] + Data[5];
    Out[3] = v; Out[1] = v; Out[10] = v; Out[2] = v;
    u8s_Lut[u8t_Cur] = Data[0];
    u8s_Lut[0] = 1U;
    u8s_Lut[4] = 2U;
    u8s_Lut[u8t_Next] = 3U;
    g_Tbl[2] = v;
    g_Tbl[0] = v;
    g_Tbl[1] = v;
    g_Tbl[10] = v;
    return v;
}
"""

_CHILD = """
import json, sys
sys.path.insert(0, sys.argv[1])
# 이 자식 프로세스는 **conftest 의 격리를 못 받는다**(`_default_local_resolver` 는 같은
# 프로세스의 fixture 다). 그래서 `config/file_mode.json` 이 cloudium 인 머신에서는
# `generate_uds_source_sections` 의 경로 판정이 worker(127.0.0.1)로 나가고, 워커가 안 떠
# 있으면 `PermissionError` 로 죽는다 — **같은 코드가 머신에 따라 통과/실패**한다.
# 이 시험이 재는 것은 해시 시드 독립성이지 파일 모드가 아니므로 여기서 로컬로 고정한다.
# (2026-09-21: 이 머신에서 재현 — HEAD 에 `config/file_mode.json` 만 복사해도 같다)
try:
    from backend.services import file_resolver as _fr
    _fr._resolver = _fr.LocalFileResolver()
except Exception:   # backend 없는 환경 — 그때는 원래 로컬이다
    pass
from report_gen.uds_generator import generate_uds_source_sections
res = generate_uds_source_sections(sys.argv[2], preprocess=False)
# (리뷰 W6) 네 필드만 비교하면 `global_data`(c_parser 의 set 순회) 같은 다른 축을 못 본다 — **전체**를 내보낸다.
print(json.dumps(res, sort_keys=True, ensure_ascii=False, default=str))
"""

_FIELDS = ("inputs", "outputs", "globals_static", "globals_global")


class TestSourceSectionsAreHashSeedIndependent:
    @pytest.fixture(scope="class")
    def outputs(self, tmp_path_factory):
        td = tmp_path_factory.mktemp("seedsrc")
        (td / "sample.c").write_text(_C_SRC, encoding="utf-8")
        child = td / "child.py"
        child.write_text(_CHILD, encoding="utf-8")
        outs = {}
        for seed in ("0", "1", "2", "3"):
            env = dict(os.environ, PYTHONHASHSEED=seed)
            r = subprocess.run([sys.executable, str(child), str(ROOT), str(td)], capture_output=True, text=True,
                               encoding="utf-8", errors="replace", env=env, cwd=str(ROOT), timeout=180)
            assert r.returncode == 0, r.stderr[-800:]
            outs[seed] = r.stdout.strip().splitlines()[-1]
        return outs

    def test_every_seed_prints_the_same_whole_payload(self, outputs):
        """"허용한 것 말고 전부 같다" — function_details 뿐 아니라 `global_data`·`globals` 등 payload 전체가 시드에 독립."""
        distinct = set(outputs.values())
        assert len(distinct) == 1, "\n---\n".join(v[:600] for v in outputs.values())

    def test_the_lists_are_sorted_not_merely_stable(self, outputs):
        """네 필드 전부에 `(idx:)` 가 있고, 그 순서가 정렬 규약(숫자 값 → 기호)이다."""
        fd = json.loads(next(iter(outputs.values())))["function_details"]
        assert len(fd) == 1
        entry = next(iter(fd.values()))
        for f in _FIELDS:
            assert any("(idx: " in s for s in entry[f]), (f, entry[f])
        assert any("(idx: 0, 1, 2, 3, 4, 5, 6, 7, 8)" in s for s in entry["inputs"]), entry["inputs"]
        assert any("(idx: 1, 2, 3, 10)" in s for s in entry["outputs"]), entry["outputs"]
        assert any("(idx: 0, 4, u8t_Cur, u8t_Next)" in s for s in entry["globals_static"]), entry["globals_static"]
        assert any("(idx: 0, 1, 2, 10)" in s for s in entry["globals_global"]), entry["globals_global"]


# --------------------------------------------------------------------------------------------
# 3. zip 멤버 시각 고정
# --------------------------------------------------------------------------------------------
def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _write_zip(p: Path, stamps: dict) -> None:
    """멤버별로 다른 시각을 박은 zip. 내용은 멤버 이름으로 고정."""
    with zipfile.ZipFile(p, "w") as z:
        for name, dt in stamps.items():
            zi = zipfile.ZipInfo(name, date_time=dt)
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, f"payload-of-{name}".encode() * 50)


class TestNormalizeZipMemberTimes:
    def test_every_member_gets_the_fixed_time_and_content_order_is_kept(self, tmp_path):
        p = tmp_path / "a.docx"
        _write_zip(p, {"[Content_Types].xml": (2026, 9, 14, 9, 3, 36), "word/document.xml": (2026, 9, 14, 9, 50, 54),
                       "word/media/image1.png": (2020, 1, 2, 3, 4, 6)})
        before = [(i.filename, i.CRC, i.file_size) for i in zipfile.ZipFile(p).infolist()]
        assert normalize_zip_member_times(p) is True
        infos = zipfile.ZipFile(p).infolist()
        assert [(i.filename, i.CRC, i.file_size) for i in infos] == before
        assert {i.date_time for i in infos} == {ZIP_FIXED_DATE_TIME}
        assert zipfile.ZipFile(p).testzip() is None

    def test_same_content_with_different_stamps_becomes_byte_identical(self, tmp_path):
        """N27-e 의 핵심 주장 — 내용이 같으면 바이트가 같다."""
        a, b = tmp_path / "a.docx", tmp_path / "b.docx"
        _write_zip(a, {"x.xml": (2026, 9, 14, 9, 3, 36), "y.xml": (2026, 9, 14, 9, 3, 36)})
        _write_zip(b, {"x.xml": (2026, 9, 14, 9, 50, 54), "y.xml": (2021, 5, 6, 7, 8, 10)})
        assert _sha(a) != _sha(b)
        assert normalize_zip_member_times(a) and normalize_zip_member_times(b)
        assert _sha(a) == _sha(b)

    def test_not_a_zip_is_left_untouched_and_reported(self, tmp_path, caplog):
        p = tmp_path / "broken.docx"
        p.write_bytes(b"not a zip at all")
        import logging
        with caplog.at_level(logging.WARNING):
            assert normalize_zip_member_times(p) is False
        assert p.read_bytes() == b"not a zip at all"
        assert [f.name for f in tmp_path.iterdir()] == ["broken.docx"], "임시 파일이 남았다"
        assert any("정규화 실패" in r.getMessage() for r in caplog.records)

    def test_missing_file_is_false_not_an_exception(self, tmp_path):
        assert normalize_zip_member_times(tmp_path / "nope.docx") is False
        assert list(tmp_path.iterdir()) == []

    def test_archive_and_member_comments_survive(self, tmp_path):
        """(리뷰 W4) 재작성이 주석을 버리면 검증이 못 잡던 침묵 손실 — 이제 복사하고 검증한다."""
        p = tmp_path / "c.docx"
        with zipfile.ZipFile(p, "w") as z:
            zi = zipfile.ZipInfo("x.xml", date_time=(2026, 9, 14, 9, 3, 36))
            zi.comment = b"member-comment"
            z.writestr(zi, b"x" * 100)
            z.comment = b"ARCHIVE-COMMENT"
        assert normalize_zip_member_times(p) is True
        with zipfile.ZipFile(p) as z:
            assert z.comment == b"ARCHIVE-COMMENT"
            assert z.infolist()[0].comment == b"member-comment"
            assert z.infolist()[0].date_time == ZIP_FIXED_DATE_TIME

    def test_members_with_extra_fields_are_refused_not_silently_stripped(self, tmp_path):
        """extra(UT/NTFS 확장 시각 등)는 복사하면 정규화가 무력화되고 버리면 손실이다 — 손대지 않고 False."""
        p = tmp_path / "e.docx"
        with zipfile.ZipFile(p, "w") as z:
            zi = zipfile.ZipInfo("x.xml", date_time=(2026, 9, 14, 9, 3, 36))
            zi.extra = b"\x55\x54\x05\x00\x01\x00\x00\x00\x00"   # 0x5455 UT
            z.writestr(zi, b"x" * 100)
        before = p.read_bytes()
        assert normalize_zip_member_times(p) is False
        assert p.read_bytes() == before

    def test_unexpected_exception_types_still_leave_the_original_and_return_false(self, tmp_path, monkeypatch):
        """(리뷰 W1) `RuntimeError`(암호화 멤버)·`zlib.error` 는 BadZipFile/OSError/ValueError 가 아니다 — 새어 나가면
        `generate_uds_docx` 를 뚫고 25분짜리 빌드가 'failed' 로 기록된다. 무엇이 나도 False + 원본 유지여야 한다."""
        p = tmp_path / "a.docx"
        _write_zip(p, {"x.xml": (2026, 9, 14, 9, 3, 36)})
        before = p.read_bytes()

        def _boom(self, *a, **k):
            raise RuntimeError("File 'x.xml' is encrypted, password required")

        monkeypatch.setattr(zipfile.ZipFile, "read", _boom)
        assert normalize_zip_member_times(p) is False
        monkeypatch.undo()
        assert p.read_bytes() == before
        assert [f.name for f in tmp_path.iterdir()] == ["a.docx"]

    def test_a_corrupted_rewrite_never_replaces_the_original(self, tmp_path, monkeypatch):
        """검증은 장식이 아니다 — 다시 쓴 멤버가 원본과 다르면(CRC) 원본을 덮지 않고 False.
        `ZipFile.read` 를 망가뜨려 재작성이 틀리게 만든다(뮤테이션 M10: 검증을 `pass` 로 바꾸면 원본이 사라진다)."""
        p = tmp_path / "a.docx"
        _write_zip(p, {"x.xml": (2026, 9, 14, 9, 3, 36), "y.xml": (2026, 9, 14, 9, 3, 36)})
        before = p.read_bytes()
        monkeypatch.setattr(zipfile.ZipFile, "read", lambda self, name, *a, **k: b"CORRUPT")
        assert normalize_zip_member_times(p) is False
        monkeypatch.undo()
        assert p.read_bytes() == before
        assert [f.name for f in tmp_path.iterdir()] == ["a.docx"]


# --------------------------------------------------------------------------------------------
# 4. 빌더 — 실제 generate_uds_docx 산출물이 재현 가능하다
# --------------------------------------------------------------------------------------------
@pytest.fixture
def build(tmp_path, monkeypatch):
    import config

    monkeypatch.setattr(config, "resolve_uds_template_path", lambda: "", raising=False)
    monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(tmp_path / "none.docx"), raising=False)
    tpl = tmp_path / "t.docx"
    d = docx.Document()
    d.add_heading("Software Unit Design", level=1)
    d.add_heading("SwUFn_0001: alpha", level=4)
    d.save(str(tpl))
    info = {"id": "SwUFn_0001", "name": "alpha", "prototype": "void alpha(const U8 * Data);",
            "description": "d", "asil": "", "related": "", "precondition": "",
            "inputs": ["[IN] const U8 * Data (idx: 0, 1, 2) (range: 0x00000000 ~ 0xFFFFFFFF)"],
            "outputs": [], "globals_global": [], "globals_static": [], "called": "", "logic": ""}
    payload = {"project_name": "T", "overview": "o", "requirements": "r", "interfaces": "i",
               "uds_frames": "u", "notes": "n", "function_details": {"SwUFn_0001": dict(info)}}

    def _run(sub: str) -> Path:
        out = tmp_path / sub / "out.docx"
        out.parent.mkdir(parents=True, exist_ok=True)
        docx_builder.generate_uds_docx(str(tpl), json.loads(json.dumps(payload)), str(out))
        return out

    return _run


class TestBuilderOutputIsReproducible:
    def test_members_carry_the_fixed_time(self, build):
        out = build("a")
        infos = zipfile.ZipFile(out).infolist()
        assert infos and {i.date_time for i in infos} == {ZIP_FIXED_DATE_TIME}

    def test_two_builds_of_the_same_payload_are_byte_identical(self, build):
        a, b = build("a"), build("b")
        assert _sha(a) == _sha(b)

    def test_every_save_path_goes_through_the_normalizer(self, build, monkeypatch):
        calls = []
        real = docx_builder.normalize_zip_member_times
        monkeypatch.setattr(docx_builder, "normalize_zip_member_times", lambda p, **kw: (calls.append(Path(p)), real(p, **kw))[1])
        out = build("a")
        assert calls == [out]

    def test_gen_stats_sidecar_records_the_normalization_outcome(self, build, monkeypatch):
        """(리뷰 W3) 빌더는 로깅 없는 서브프로세스에서 돌아 INFO/WARNING 이 안 남는다 — 사이드카가 유일한 관측 자리."""
        out = build("a")
        stats = json.loads(docx_builder.gen_stats_path(str(out)).read_text(encoding="utf-8"))
        assert stats["zip_time_normalized"] is True
        assert isinstance(stats["zip_time_normalize_sec"], float)
        monkeypatch.setattr(docx_builder, "normalize_zip_member_times", lambda p, **kw: False)
        out2 = build("b")
        stats2 = json.loads(docx_builder.gen_stats_path(str(out2)).read_text(encoding="utf-8"))
        assert stats2["zip_time_normalized"] is False, "실패가 사이드카에 안 남으면 나중에 해시가 다를 때 '입력이 달랐나' 로 오독한다"
        assert out2.exists() and out2.stat().st_size > 0, "정규화 실패가 산출물을 지우면 안 된다(fail-open)"

    def test_docx_builder_has_a_single_save_site(self):
        """세 종결 분기(토큰 템플릿·구조 복제·무템플릿)가 전부 `_save_docx` 를 지난다 — 하나라도 `doc.save` 로 돌아가면 그 경로만 시각이 남는다."""
        src = (ROOT / "report_gen" / "docx_builder.py").read_text(encoding="utf-8")
        body = "\n".join(ln for ln in src.splitlines() if not ln.lstrip().startswith("#"))
        assert body.count("doc.save(") == 1
        assert body.count("_save_docx(doc, out, output_path, ") == 3
        assert body.count("_write_gen_stats(output_path, ") == 1, "gen_stats 는 _save_docx 가 저장 뒤에 한 번 쓴다"


# --------------------------------------------------------------------------------------------
# 5. 캐시 — 파서 출력 모양이 바뀌었으니 소스 섹션 캐시 버전이 올라가 있어야 한다
# --------------------------------------------------------------------------------------------
def test_source_sections_cache_version_was_bumped_past_v15():
    from backend.helpers.uds import _SOURCE_SECTIONS_SCHEMA_VERSION as ver

    assert ver.startswith("v") and int(ver[1:]) >= 16, ver
