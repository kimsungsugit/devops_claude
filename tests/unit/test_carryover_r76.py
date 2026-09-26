"""R76 — 이월 일괄 정리(N82 · N86 · N87 · N91 · N93~N99 · N101~N106)의 가드.

각 항목의 **관측량**을 단언한다(산출 파일의 셀 · 리포트의 수 · 핸들러가 생성기에 넘긴 인자). 실측 근거는 계획서 R76 절.
"""
from __future__ import annotations

import inspect
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]


# ── N102: STS 안전 표시의 극성 ────────────────────────────────────────────────
class TestStsSafetyPolarity:
    @staticmethod
    def _tc(i: int, mark: str):
        return {"id": f"SwTC_{i:03d}", "title": "t", "safety_related": mark, "test_environment": "SwTE_01",
                "test_method": "RBT", "gen_method": "AOR", "fs_req": "", "description": "d", "precondition": "p",
                "srs_id": "SwTR_0001", "steps": [{"action": "a1", "expected": "e1"}, {"action": "a2", "expected": "e2"}]}

    def test_quality_report_counts_the_safety_rows_not_the_qm_rows(self):
        from generators.sts import generate_quality_report

        tcs = [self._tc(1, "O"), self._tc(2, "O"), self._tc(3, "X"), self._tc(4, "X"), self._tc(5, "X"), self._tc(6, "")]
        assert generate_quality_report(tcs, {"coverage": {}})["safety_test_cases"] == 2

    def test_writer_highlights_the_safety_rows_not_the_qm_rows(self, tmp_path):
        from openpyxl import load_workbook

        from generators.sts import _HEADER_ROW, STS_COL, generate_sts_xlsm

        out = generate_sts_xlsm(None, [self._tc(1, "O"), self._tc(2, "X"), self._tc(3, "")],
                                {"matrix": [], "coverage": {}}, str(tmp_path / "p.xlsx"), {})
        wb = load_workbook(out)
        ws = next(wb[n] for n in wb.sheetnames if "Test Spec" in n)
        fills = {}
        for r in range(_HEADER_ROW + 1, ws.max_row + 1):
            mark = ws.cell(row=r, column=STS_COL["safety_related"]).value
            if mark is not None or r == _HEADER_ROW + 1:
                cur = mark
            fills.setdefault(cur, set()).add(str(ws.cell(row=r, column=STS_COL["action"]).fill.start_color.rgb)[-6:])
        assert fills["O"] == {"FFFFCC"}, fills
        assert "FFFFCC" not in fills["X"], fills

    def test_marks_come_from_one_module(self):
        from generators import safety_marks as sm

        assert (sm.resolve_safety_related("ASIL B"), sm.resolve_safety_related("QM"), sm.resolve_safety_related("TBD")) == (
            sm.SAFETY_RELATED_MARK, sm.NON_SAFETY_MARK, "")


# ── N82: flow 없는 함수의 TC 도 스텝 상한을 받는다 ────────────────────────────
class TestStsNoFlowStepCap:
    INFO = {"name": "f", "prototype": "U8 f(U8 a, U8 b)", "inputs": ["[IN] U8 a", "[IN] U8 b"], "outputs": ["[OUT] U8 ret"]}

    def test_premise_no_flow_function_has_a_long_boundary_tc(self):
        from generators.sts import _generate_simple_steps

        tcs = _generate_simple_steps(dict(self.INFO))
        # [정상, 경계값, 범위 초과] — 둘째가 경계값 TC 이고 4 스텝을 넘는다(이 전제가 깨지면 아래 두 테스트는 아무것도 안 잰다)
        assert len(tcs) == 3 and len(tcs[1]) > 4 and tcs[1][0]["action"].startswith("입력 설정 (경계"), [len(t) for t in tcs]

    def test_steps_are_cut_to_the_cap(self):
        from generators.sts import _generate_steps_from_flow

        tcs = _generate_steps_from_flow([], dict(self.INFO), max_steps=4)
        assert tcs and max(len(t) for t in tcs) <= 4

    def test_a_cap_below_four_drops_the_boundary_tc_instead_of_halving_it(self):
        from generators.sts import _classify_steps, _generate_steps_from_flow

        tcs = _generate_steps_from_flow([], dict(self.INFO), max_steps=3)
        assert len(tcs) == 2 and all(len(t) <= 3 for t in tcs)           # 정상 + 범위 초과 — 경계값 TC 만 빠진다
        assert not any(st["action"].startswith("입력 설정 (경계") for t in tcs for st in t)
        assert _classify_steps(tcs[0])[1] != "BAA"

    def test_default_cap_changes_nothing(self):
        from generators.sts import _generate_simple_steps, _generate_steps_from_flow

        assert _generate_steps_from_flow([], dict(self.INFO)) == _generate_simple_steps(dict(self.INFO))


# ── N99: 두 번째 HSIS 파서도 묶음 머리 아래의 Name 열을 읽는다 ────────────────
class TestSecondHsisParserGroupedHeader:
    @staticmethod
    def _book(path, sub_header=True, gap_row=False):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "2.HSIS"
        ws.append(["HW SW Interface"])
        #        0      1              2       3    4(group=Type 열)  5              6
        ws.append(["ID", "Signal Name", "Direction", "", "SW Variable", "", "Related ID"])
        if gap_row:
            ws.append([None] * 7)
        if sub_header:
            ws.append(["", "", "", "Name", "Type", "Value Range", ""])
        ws.append(["HSI_01", "VSUP", "IN", "u16g_DrvIn_Vsup", "U16", "0x0000U ~ 0xFFFFU", "SyTR_0401"])
        ws.append(["HSI_02", "TEMP", "IN", "s16g_DrvIn_Temp", "S16", "-40 ~ 150", "SyTR_0402"])
        wb.save(path)
        return str(path)

    def test_names_not_types(self, tmp_path):
        from generators.sts import parse_hsis_signals

        got = parse_hsis_signals(self._book(tmp_path / "h.xlsx"))
        assert [s["sw_var_name"] for s in got["signals"]] == ["u16g_DrvIn_Vsup", "s16g_DrvIn_Temp"]
        assert got["sw_var_names"] == ["u16g_DrvIn_Vsup", "s16g_DrvIn_Temp"]
        assert {s["column_layout"] for s in got["signals"]} == {"header"}

    def test_blank_row_between_group_and_sub_header(self, tmp_path):
        from generators.sts import parse_hsis_signals

        got = parse_hsis_signals(self._book(tmp_path / "g.xlsx", gap_row=True))
        assert [s["sw_var_name"] for s in got["signals"]] == ["u16g_DrvIn_Vsup", "s16g_DrvIn_Temp"]

    def test_no_sub_header_is_disclosed_as_fallback(self, tmp_path):
        from generators.sts import parse_hsis_signals

        got = parse_hsis_signals(self._book(tmp_path / "n.xlsx", sub_header=False))
        assert {s["column_layout"] for s in got["signals"]} == {"group-fallback"}
        assert len(got["signals"]) == 2          # 데이터 행을 머리행으로 삼키지 않는다

    def test_a_data_row_is_never_taken_for_the_sub_header(self, tmp_path):
        """하위 머리행이 없는 문서에서, 둘째 데이터 행의 변수 이름이 하필 `Name` 이어도 머리행으로 삼지 않는다 —
        첫 데이터 행을 본 순간 기다림이 끝난다(`_load_hsis_signals` 와 같은 규칙)."""
        from openpyxl import Workbook

        from generators.sts import parse_hsis_signals

        wb = Workbook()
        ws = wb.active
        ws.title = "2.HSIS"
        ws.append(["ID", "Signal Name", "Direction", "", "SW Variable", "", "Related ID"])
        ws.append(["HSI_01", "VSUP", "IN", "u16g_Vsup", "U16", "", "SyTR_0401"])
        ws.append(["HSI_02", "TEMP", "IN", "Name", "S16", "", "SyTR_0402"])
        wb.save(tmp_path / "d.xlsx")
        got = parse_hsis_signals(str(tmp_path / "d.xlsx"))
        assert {s["column_layout"] for s in got["signals"]} == {"group-fallback"}
        assert [s["sw_var_name"] for s in got["signals"]] == ["U16", "S16"]      # 묶음 열(=Type) — 틀린 값이지만 **공시된** 폴백

    def test_both_parsers_agree(self, tmp_path):
        from generators.sts import _load_hsis_signals, parse_hsis_signals

        p = self._book(tmp_path / "b.xlsx")
        a = {s["id"]: s["sw_var_name"] for s in _load_hsis_signals(p)["signals"]}
        b = {s["id"]: s["sw_var_name"] for s in parse_hsis_signals(p)["signals"]}
        assert a == b and len(a) == 2

    @pytest.mark.parametrize("low, group, want", [
        (["", "", "", "name", "type", "value range", ""], 4, (3, 4, 5)),          # KJPDS02: 라벨이 Type 열 위 → 바로 왼쪽
        (["", "", "name", "name", "type", "value range"], 3, (3, 4, 5)),          # 표준: 라벨이 Name 열 위 — 왼쪽의 남의 Name 을 집지 않는다
        (["name", "", "", "", "", "name", "type"], 4, (5, 6, -1)),                # 오른쪽에만 있을 때
        (["name", "", "", "type", "", "", "", "", "", "name"], 3, (-1, -1, -1)),  # 구간 밖 Name 은 남의 것
    ])
    def test_sub_header_name_search_order(self, low, group, want):
        from generators.sts import _hsis_sw_variable_subcols

        assert _hsis_sw_variable_subcols(low, group) == want

    def test_both_parsers_wait_the_same_number_of_rows(self):
        from generators import sts
        from tests.unit._source_probe import source_of

        assert sts._HSIS_SUB_HEADER_WAIT_ROWS == 3
        assert "_HSIS_SUB_HEADER_WAIT_ROWS" in source_of(sts._load_hsis_signals)
        assert "_HSIS_SUB_HEADER_WAIT_ROWS" in source_of(sts.parse_hsis_signals)

    def test_flat_header_layout_unchanged(self, tmp_path):
        from openpyxl import Workbook

        from generators.sts import parse_hsis_signals

        wb = Workbook()
        ws = wb.active
        ws.title = "2.HSIS"
        ws.append(["ID", "Signal Name", "SW Variable Name", "Related ID"])
        ws.append(["HSI_01", "VSUP", "u16g_Vsup", "SwTR_0101"])
        wb.save(tmp_path / "f.xlsx")
        got = parse_hsis_signals(str(tmp_path / "f.xlsx"))
        assert [s["sw_var_name"] for s in got["signals"]] == ["u16g_Vsup"]


# ── N106: STS 요구-함수 매핑 — 토큰 한가운데 일치 금지 · 경로 공시 ────────────
class TestStsMappingTokenBoundary:
    @pytest.mark.parametrize("raw, short, want", [
        ("s_UDS_SlipDetect_ReadCustomPayload", "adc", False),      # Re[adC]ustom
        ("design guideline v1", "Lin", False),                     # guide[lin]e
        ("reproglinsendtask", "LinSend", False),                   # ⚠ 소문자로 접힌 이름엔 경계가 없다 — 그래서 호출부가 이 방향엔 안 쓴다
        ("g_uds_lincomp_reset", "UDS Lin Comp", True),             # 시작은 토큰 경계, 끝은 아님 — 살린다
        ("u16s_DiagCheck_HB_Helper2", "s_DiagCheck_HB", True),     # 헝가리안 접두의 범위 글자는 토큰 시작
        ("u8g_Mode", "g_Mode", True),
        ("status_Mode", "s_Mode", False),                          # 접두가 아닌 첫 토큰은 쪼개지 않는다
        ("s_Sleep_DisableInterrupts", "eint", False),              # Disabl[eInt]errupts
        ("LinRx", "lin", True),
        ("SCI_T_Bytes", "sci", True),
        ("s_UDS_WDBI_UsOptMain", "main", True),
        ("abc", "", False),
    ])
    def test_rule(self, raw, short, want):
        from report_gen.requirements import contains_at_token_start, normalize_sds_key

        assert contains_at_token_start(raw, normalize_sds_key(raw), normalize_sds_key(short)) is want

    def test_second_occurrence_at_a_boundary_counts(self):
        from report_gen.requirements import contains_at_token_start, normalize_sds_key

        raw = "readc_adc_value"          # 첫 `adc` 는 토큰 한가운데(re[adc]), 둘째는 토큰 시작
        assert contains_at_token_start(raw, normalize_sds_key(raw), "adc") is True

    def test_function_name_inside_a_lowercased_key_keeps_the_old_rule(self):
        """키는 소문자로 접혀 camelCase 경계가 없다 — 그 방향에 토큰 규칙을 걸면 옳은 링크가 끊긴다(라이브 실측 6 함수)."""
        from generators.sts import map_requirements_to_functions

        reqs = [{"id": "SwTR_0106"}]
        got = map_requirements_to_functions(reqs, {"f1": {"name": "LinRecv", "module_name": ""}},
                                            sds_map={"reproglinrecvtask": {"related": "SwTR_0106"}})
        assert got["SwTR_0106"] == ["f1"]

    def test_hungarian_prefixed_helper_still_finds_its_partition(self):
        from generators.sts import map_requirements_to_functions

        reqs = [{"id": "SwTSR_0101"}]
        got = map_requirements_to_functions(reqs, {"f1": {"name": "u16s_DiagCheck_HB_Helper2", "module_name": ""}},
                                            sds_map={"s_diagcheck_hb": {"related": "SwTSR_0101"}})
        assert got["SwTSR_0101"] == ["f1"]

    def test_mid_token_partition_no_longer_links_the_function(self):
        from generators.sts import map_requirements_to_functions

        reqs = [{"id": "SwTR_0101"}, {"id": "SwTR_0202"}]
        details = {"f1": {"name": "s_UDS_ReadCustomPayload", "module_name": ""},
                   "f2": {"name": "ADC_Measure", "module_name": ""}}
        sds = {"adc": {"related": "SwTR_0101"}}
        got = map_requirements_to_functions(reqs, details, sds_map=sds)
        assert got["SwTR_0101"] == ["f2"]

    def test_paths_are_disclosed(self):
        from generators.sts import map_requirements_to_functions

        reqs = [{"id": "SwTR_0101"}, {"id": "SwTR_0202"}]
        details = {"f1": {"name": "own", "related": "SwTR_0101"},
                   "f2": {"name": "ADC_Measure"},            # sds_substring (adc ⊂ ADC_Measure, 토큰 시작)
                   "f3": {"name": "exact_fn"},               # sds_exact
                   "f4": {"name": "nothing"}}
        sds = {"adc": {"related": "SwTR_0202"}, "exact_fn": {"related": "SwTR_0202"}}
        stats: dict = {}
        map_requirements_to_functions(reqs, details, sds_map=sds, stats_out=stats)
        m = stats["requirement_function_mapping"]
        assert m["functions_by_path"] == {"own_related": 1, "sds_substring": 1, "sds_exact": 1}
        assert (m["functions_total"], m["design_id_bridge_gained"], m["design_id_bridge_only"], m["unlinked"]) == (4, 0, 0, 1)
        assert m["links"] == 3 and m["functions_per_requirement"] == {"median": 2, "max": 2}

    def test_median_is_not_the_max(self):
        from generators.sts import map_requirements_to_functions

        reqs = [{"id": "SwTR_0001"}, {"id": "SwTR_0002"}, {"id": "SwTR_0003"}]
        details = {"a": {"name": "a", "related": "SwTR_0001, SwTR_0002, SwTR_0003"},
                   "b": {"name": "b", "related": "SwTR_0002, SwTR_0003"},
                   "c": {"name": "c", "related": "SwTR_0003"}}
        stats: dict = {}
        map_requirements_to_functions(reqs, details, sds_map={}, stats_out=stats)
        assert stats["requirement_function_mapping"]["functions_per_requirement"] == {"median": 2, "max": 3}

    def test_the_stats_reach_the_quality_report(self, tmp_path):
        """헬퍼 단독 테스트는 호출부가 값을 버리는 것을 못 본다 — `generate_sts` 의 리포트에서 읽는다."""
        import generators.sts as gsts

        out = gsts.generate_sts(requirements_text=["SwTR_0606: 차속에 따른 도어 open 방지"],
                                function_details={"f1": {"id": "f1", "name": "foo", "module_name": "", "related": "SwTR_0606"}},
                                output_path=str(tmp_path / "out.xlsx"))
        m = out["quality_report"]["generation_stats"]["requirement_function_mapping"]
        assert (m["functions_total"], m["functions_by_path"]["own_related"], m["unlinked"]) == (1, 1, 0)

    def test_unchecked_direction_has_its_own_label(self):
        from generators.sts import map_requirements_to_functions

        stats: dict = {}
        map_requirements_to_functions([{"id": "SwTR_0106"}], {"f1": {"name": "LinRecv"}},
                                      sds_map={"reproglinrecvtask": {"related": "SwTR_0106"}}, stats_out=stats)
        assert stats["requirement_function_mapping"]["functions_by_path"] == {"own_related": 0, "sds_name_in_key": 1}

    def test_a_function_saved_by_the_bridge_is_not_also_unlinked(self):
        """(리뷰 W2) 못 붙은 함수는 세 티어가 끝난 뒤에 센다 — 칸의 합이 함수 수다."""
        from generators.sts import map_requirements_to_functions

        reqs = [{"id": "SwTR_0101"}]
        details = {"f1": {"name": "bridged_fn"}, "f2": {"name": "lost_fn"}, "f3": {"name": "own", "related": "SwTR_0101"}}
        sds = {"swfn_01": {"related": "SwTR_0101"}}
        stats: dict = {}
        got = map_requirements_to_functions(reqs, details, sds_map=sds, uds_design_ids={"bridged_fn": ["SwFn_01"]}, stats_out=stats)
        assert sorted(got["SwTR_0101"]) == ["f1", "f3"]
        m = stats["requirement_function_mapping"]
        assert (m["design_id_bridge_gained"], m["design_id_bridge_only"], m["unlinked"]) == (1, 1, 1)
        assert sum(m["functions_by_path"].values()) + m["design_id_bridge_only"] + m["unlinked"] == m["functions_total"] == 3

    def test_suts_asil_fallback_is_deliberately_left_alone(self):
        """SUTS 의 ASIL 폴백은 옛 부분문자열 규칙 그대로다(다른 가드가 묶는 결정 — 여기선 '안 바꿨다' 만 본다)."""
        from generators.suts import _resolve_unit_asil

        got, ev = _resolve_unit_asil({"module_name": "ReadCustom"}, {"adc": {"asil": "B"}})
        assert (got, ev.startswith("sds-fuzzy")) == ("B", True)


# ── N98 · N95 · N97 · N101 · N104: SUTS ───────────────────────────────────────
class TestSutsPointerAddressRange:
    @staticmethod
    def _unit():
        return {"name": "f", "fid": "SwUFn_001", "input_vars": ["Values", "n", "k"], "output_vars": [],
                "param_types": {"Values": "U16 *", "n": "U8", "k": "U8"},
                "uds_param_info": {"Values": {"range": [0, 0xFFFFFFFF]}, "n": {"range": [0, 10]}, "k": {"range": [0, 999]}}}

    def test_pointer_address_range_is_not_a_document_error(self):
        from generators.suts import generate_sequences

        u = self._unit()
        generate_sequences(u, type_cache={})
        assert u["pointer_address_ranges"] == ["Values(SwUDS 0~4294967295 vs uint16_t)"]
        assert u["range_conflicts"] == ["k(SwUDS 0~999 vs uint8_t)"]          # 스칼라의 상충은 그대로 문서 오류 후보
        assert u["bounds_source"] == {"Values": "type", "n": "uds_range", "k": "type"}   # 값 처리는 안 바뀐다

    def test_array_declaration_is_an_address_slot_too(self):
        from generators.suts import generate_sequences

        u = self._unit()
        u["param_types"]["Values"] = "U16 []"
        generate_sequences(u, type_cache={})
        assert u["pointer_address_ranges"] == ["Values(SwUDS 0~4294967295 vs uint16_t)"] and u["bounds_source"]["Values"] == "type"

    def test_a_range_that_fits_the_pointee_is_still_used(self):
        from generators.suts import generate_sequences

        u = self._unit()
        u["uds_param_info"]["Values"] = {"range": [0, 1000]}
        generate_sequences(u, type_cache={})
        assert u["bounds_source"]["Values"] == "uds_range" and u["pointer_address_ranges"] == []

    def test_report_counts_the_two_axes_apart(self):
        from generators.suts import generate_sequences, generate_suts_quality_report

        u = self._unit()
        seqs = {u["fid"]: generate_sequences(u, type_cache={})}
        q = generate_suts_quality_report([u], seqs, 1)
        assert (q["range_conflict_count"], q["pointer_address_range_count"]) == (1, 1)
        assert q["pointer_address_ranges"] == ["Values(SwUDS 0~4294967295 vs uint16_t)"]


class TestSutsDisclosures:
    def test_verify_needed_slots_are_counted(self):
        from generators.suts import _VERIFY_NEEDED_PREFIX, generate_suts_quality_report

        u = {"name": "f", "fid": "F1", "input_vars": ["a"], "output_vars": ["o"]}
        seqs = {"F1": [{"seq_num": 1, "inputs": {"a": 1}, "expected": {"o": f"{_VERIFY_NEEDED_PREFIX} 256", "p": 3}},
                       {"seq_num": 2, "inputs": {"a": 2}, "expected": {"o": f"{_VERIFY_NEEDED_PREFIX} -1"}}]}
        assert generate_suts_quality_report([u], seqs, 1)["verify_needed_expected_slots"] == 2

    def test_the_generator_writes_the_same_marker_it_counts(self):
        from generators import suts
        from tests.unit._source_probe import source_of

        src = source_of(suts)
        assert src.count('f"{_VERIFY_NEEDED_PREFIX} {raw}"') == 2 and 'f"[검증 필요]' not in src

    def test_requirement_link_evidence_distribution(self):
        from generators.suts import generate_suts_quality_report

        units = [{"name": "a", "fid": "A", "srs_req_ids": "SwTR_1", "srs_req_link": "sts_mapping"},
                 {"name": "b", "fid": "B", "srs_req_ids": "SwTR_2", "srs_req_link": "hsis"},
                 {"name": "c", "fid": "C", "srs_req_ids": "SwTR_3"},
                 {"name": "d", "fid": "D", "srs_req_ids": "", "srs_req_link": "sts_mapping"}]
        q = generate_suts_quality_report(units, {}, 4)
        assert q["srs_req_link_distribution"] == {"sts_mapping": 1, "hsis": 1, "unrecorded": 1}

    def _run(self, tmp_path, monkeypatch, **kw):
        from generators import suts as suts_mod

        (tmp_path / "a.c").write_text("unsigned char f(unsigned char a) { if (a) { return 1; } return 0; }\n", encoding="utf-8")
        return suts_mod.generate_suts(source_root=str(tmp_path), output_path=str(tmp_path / "u.xlsm"), scope="source", **kw)

    def test_caps_are_disclosed_like_sits(self, tmp_path, monkeypatch):
        ref = self._run(tmp_path, monkeypatch, max_sequences=9)["quality_report"]
        assert (ref["caps_requested"], ref["caps_effective"]) == ({"max_sequences": 9}, {"max_sequences": 9})
        ext = self._run(tmp_path, monkeypatch, max_sequences=9, tc_profile="extended")["quality_report"]
        assert (ext["caps_requested"], ext["caps_effective"]) == ({"max_sequences": 9}, {"max_sequences": None})

    def test_enrichment_failure_reaches_the_report_and_the_warnings(self, tmp_path, monkeypatch):
        from generators import sts as sts_mod

        srs = tmp_path / "srs.docx"
        srs.write_bytes(b"x")

        def _boom(_p):
            raise RuntimeError("SRS 표가 깨졌다")

        monkeypatch.setattr(sts_mod, "parse_srs_docx_tables", _boom)
        out = self._run(tmp_path, monkeypatch, srs_docx_path=str(srs))
        errs = out["quality_report"]["enrichment_errors"]
        assert errs == [{"stage": "srs_req_ids", "error": "RuntimeError: SRS 표가 깨졌다"}]
        assert any("srs_req_ids" in w and "SRS 표가 깨졌다" in w for w in out["validation"]["warnings"])

    def test_no_failure_means_an_empty_list_not_a_missing_key(self, tmp_path, monkeypatch):
        assert self._run(tmp_path, monkeypatch)["quality_report"]["enrichment_errors"] == []


# ── N103: SUTS·SITS 라이터의 데이터 영역 병합은 기존 병합을 훑지 않는다 ────────
def _assert_disjoint(ws) -> None:
    """(리뷰 G7) `merge_fresh` 가 건너뛴 검사의 **실효** — 병합 범위가 서로 겹치지 않는다(겹치면 Excel 이 "복구" 를 묻는다)."""
    cells: set = set()
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                assert (r, c) not in cells, f"overlapping merge at {(r, c)} in {rng}"
                cells.add((r, c))


class TestWritersDoNotScanExistingMerges:
    @staticmethod
    def _count(monkeypatch):
        from openpyxl.worksheet.cell_range import MultiCellRange

        calls = {"n": 0}
        real = MultiCellRange.add

        def _counting(self, coord):
            calls["n"] += 1
            return real(self, coord)

        monkeypatch.setattr(MultiCellRange, "add", _counting)
        return calls

    def test_suts(self, tmp_path, monkeypatch):
        from openpyxl import load_workbook

        from generators.suts import generate_suts_xlsm

        calls = self._count(monkeypatch)

        def _mk(n):
            units = [{"fid": f"F{i}", "name": f"f{i}", "component": "SwCom_01", "input_vars": ["a"], "output_vars": [], "asil": "QM"}
                     for i in range(n)]
            seqs = {u["fid"]: [{"seq_num": 1, "strategy": "BV_MIN", "inputs": {"a": 0}, "expected": {}},
                               {"seq_num": 2, "strategy": "BV_MID", "inputs": {"a": 1}, "expected": {}},
                               {"seq_num": 3, "strategy": "BV_MAX_INV", "inputs": {"a": 256}, "expected": {}}] for u in units}
            return units, seqs

        per = {}
        for n in (1, 6):
            calls["n"] = 0
            out = generate_suts_xlsm(None, *_mk(n), str(tmp_path / f"u{n}.xlsx"), {"project_id": "T"})
            per[n] = calls["n"]
        assert per[1] == per[6], per
        ws = load_workbook(out)["2.SW Unit Test Spec"]
        data_merges = [r for r in ws.merged_cells.ranges if r.min_row >= 5 and r.max_row > r.min_row]
        assert len(data_merges) >= 6 * 5                       # TC 메타 5열 × 6 TC — 병합 자체는 그대로 있다
        _assert_disjoint(ws)

    def test_sits(self, tmp_path, monkeypatch):
        from openpyxl import load_workbook

        from generators.sits import _SPEC_SHEET_NAME, generate_sits_xlsm
        from tests.unit.test_generators_sits import _round_trip_itcs

        calls = self._count(monkeypatch)
        per = {}
        for n in (1, 6):
            calls["n"] = 0
            out = generate_sits_xlsm(None, _round_trip_itcs(n_tc=n, n_sub=3), str(tmp_path / f"i{n}.xlsx"))
            per[n] = calls["n"]
        assert per[1] == per[6], per
        ws = load_workbook(out)[_SPEC_SHEET_NAME]
        assert len([r for r in ws.merged_cells.ranges if r.min_row >= 5 and r.max_row > r.min_row]) >= 6 * 5
        _assert_disjoint(ws)


# ── N87: SITS 방법 분포는 문서에 쓰는 값 ───────────────────────────────────────
class TestSitsMethodDistributionIsWhatTheDocumentSays:
    @staticmethod
    def _itc(i, internal, method=None):
        d = {"tc_id": f"SwITC_{i:02d}", "gen_method": internal, "sub_cases": [{}], "related_ids": [], "input_vars": ["a"]}
        if method:
            d["test_method"] = method
        return d

    def test_internal_labels_do_not_inflate_the_distribution(self):
        from generators.sits import generate_sits_quality_report

        itcs = [self._itc(1, "AOR, ABV"), self._itc(2, "ABV, AEC"), self._itc(3, "ABV")]
        assert generate_sits_quality_report(itcs)["gen_method_distribution"] == {"AOR": 3, "AEC": 3}

    def test_fault_injection_tc_adds_the_boundary_code(self):
        from generators.sits import _SITS_METHOD_FAULT, generate_sits_quality_report

        itcs = [self._itc(1, "ABV"), self._itc(2, "ABV", method=_SITS_METHOD_FAULT)]
        assert generate_sits_quality_report(itcs)["gen_method_distribution"] == {"AOR": 2, "AEC": 1, "ABV": 1}

    def test_distribution_matches_the_written_cells(self, tmp_path):
        from openpyxl import load_workbook

        from generators.sits import (
            _GEN_COL,
            _SPEC_SHEET_NAME,
            _split_method_codes,
            generate_sits_quality_report,
            generate_sits_xlsm,
        )
        from tests.unit.test_generators_sits import _round_trip_itcs

        itcs = _round_trip_itcs(n_tc=3, n_sub=2)
        out = generate_sits_xlsm(None, itcs, str(tmp_path / "d.xlsx"))
        ws = load_workbook(out)[_SPEC_SHEET_NAME]
        written: dict = {}
        for r in range(5, ws.max_row + 1):
            for code in _split_method_codes(ws.cell(row=r, column=_GEN_COL).value):
                written[code] = written.get(code, 0) + 1
        assert written == generate_sits_quality_report(itcs)["gen_method_distribution"] and written

    def test_evaluator_denominator_is_the_writable_vocabulary(self):
        from generators.sits import SITS_WRITABLE_GEN_CODES
        from workflow.quality import evaluator

        assert evaluator._SITS_GEN_METHOD_VOCAB_SIZE == float(len(SITS_WRITABLE_GEN_CODES)) == 3.0

    @pytest.mark.parametrize("dist, want", [({"AOR": 5, "AEC": 5}, 66.67), ({"AOR": 5, "AEC": 4, "ABV": 1}, 100.0), ({}, 0.0)])
    def test_evaluator_value(self, dist, want):
        from workflow.quality.evaluator import evaluate_sits

        ms = evaluate_sits({"total_test_cases": 5, "gen_method_distribution": dist})
        assert next(m["value"] for m in ms if m["metric_name"] == "method_diversity_pct") == want


# ── N86: UDS 소스 단계 출처 — 자리표시자는 값이 아니다 ─────────────────────────
class TestUdsSourceStagePlaceholders:
    @pytest.mark.parametrize("placeholder", ["N/A", "-", "none", "TBD", "  n/a "])
    def test_placeholder_in_an_earlier_stage_does_not_mask_a_later_value(self, placeholder):
        from report_gen.uds_generator import _source_stage_provenance

        got = _source_stage_provenance(comment_asil=placeholder, comment_related=placeholder,
                                       override={"asil": placeholder}, sds_asil="B", sds_related="SwTR_0101")
        assert got == ("B", "sds", "SwTR_0101", "sds")

    def test_all_placeholders_is_tbd_with_default_label(self):
        from report_gen.uds_generator import _source_stage_provenance

        got = _source_stage_provenance(comment_asil="N/A", comment_related="-", override={}, sds_asil="none", sds_related="")
        assert got == ("TBD", "default", "TBD", "default")

    def test_real_values_keep_the_chain_order(self):
        from report_gen.uds_generator import _source_stage_provenance

        got = _source_stage_provenance(comment_asil="QM", comment_related="", override={"related": "SwTR_1"},
                                       sds_asil="B", sds_related="SwTR_2")
        assert got == ("QM", "comment", "SwTR_1", "override")


# ── N96: UDS Value Range — enum 의 닫힌 값 집합 ───────────────────────────────
class TestUdsEnumValueRange:
    def test_small_domain_is_listed(self):
        from report_gen.function_analyzer import _param_value_range

        assert _param_value_range({"type": "Mode_t", "value_domain": {"values": [2, 0, 1, 1]}}) == "0, 1, 2 (enum 열거자)"

    def test_large_domain_is_a_span_with_a_count(self):
        from report_gen.function_analyzer import _ENUM_RANGE_LIST_MAX, _param_value_range

        n = _ENUM_RANGE_LIST_MAX + 1
        assert _param_value_range({"type": "E", "value_domain": {"values": list(range(n))}}) == f"0 ~ {n - 1} (enum 열거자 {n}개)"
        assert "~" not in _param_value_range({"type": "E", "value_domain": {"values": list(range(_ENUM_RANGE_LIST_MAX))}})

    def test_junk_range_falls_to_the_domain_and_a_real_range_wins(self):
        from report_gen.function_analyzer import _param_value_range

        dom = {"values": [0, 1]}
        assert _param_value_range({"type": "E", "range": "}", "value_domain": dom}) == "0, 1 (enum 열거자)"
        assert _param_value_range({"type": "E", "range": "0 ~ 5", "value_domain": dom}) == "0 ~ 5"

    @pytest.mark.parametrize("dom", [None, {}, {"values": []}, {"values": ["x"]}])
    def test_no_domain_stays_na(self, dom):
        from report_gen.function_analyzer import _param_value_range

        assert _param_value_range({"type": "E", "value_domain": dom}) == "N/A"

    def test_component_global_table_uses_the_same_text(self):
        from report_gen.utils import _build_global_rows

        hdr = ["Name", "Type", "Value Range", "Reset Value", "Description"]
        rows = _build_global_rows(["g_Mode", "g_U8"], {"g_Mode": {"type": "Mode_t", "value_domain": {"values": [0, 1]}},
                                                        "g_U8": {"type": "U8", "range": "0 ~ 255"}}, hdr, with_labels=False)
        assert [r[2] for r in rows] == ["0, 1 (enum 열거자)", "0 ~ 255"]


# ── N94: 영향도 초안의 타입 출처 라벨 ──────────────────────────────────────────
class TestImpactDraftTypedefLabel:
    GIM = {"g_Speed": {"type": "Speed_t", "base_type": "uint16_t"},      # 별칭 → 풀어서 앎
           "g_Cnt": {"type": "U16", "base_type": "unsigned int"},        # 선언이 이미 아는 이름 — 풀린 쪽을 쓰지 않는다
           "g_Raw": {"type": "uint8_t"}}

    def test_resolved_set_matches_what_the_type_map_picked(self):
        from generators.suts import _gim_to_type_map, _gim_typedef_resolved

        assert _gim_typedef_resolved(self.GIM) == {"g_Speed"}
        assert _gim_to_type_map(self.GIM)["g_Speed"] == "uint16_t" and _gim_to_type_map(self.GIM)["g_Cnt"] == "U16"

    def test_label(self):
        from generators.suts import _gim_to_type_map, _gim_typedef_resolved
        from workflow.impact_doc_draft import build_var_types

        vt = build_var_types(["g_Speed", "g_Cnt", "g_Raw[2]"], _gim_to_type_map(self.GIM),
                             typedef_resolved=_gim_typedef_resolved(self.GIM))
        assert {k: v["source"] for k, v in vt.items()} == {"g_Speed": "globals_map_typedef", "g_Cnt": "globals_map", "g_Raw": "globals_map"}
        assert vt["g_Speed"]["type"] == "uint16_t"

    def test_without_the_set_the_old_label_stays(self):
        from workflow.impact_doc_draft import resolve_var_type

        assert resolve_var_type("g_Speed", {"g_Speed": "uint16_t"}) == {"type": "uint16_t", "source": "globals_map"}

    def test_orchestrator_passes_the_set(self):
        src = (ROOT / "workflow/impact_orchestrator.py").read_text(encoding="utf-8")
        assert "build_var_types(_vars, _local_tc, typedef_resolved=_td_resolved)" in src
        assert "_td_resolved = _gim_typedef_resolved(gim)" in src


# ── N93: 영향도 초안의 STS 미리보기 — 경계값 TC 가 절단에 먼저 잘리지 않는다 ───
class TestImpactDraftKeepsTheBoundaryTc:
    def _proposal(self, monkeypatch, n_branch, boundary=True, **kw):
        from generators import sts as gsts
        from tests.unit.test_impact_doc_content import _proposal_sections, _stub_generators
        from workflow.impact_orchestrator import _build_doc_proposal

        tcs = [[{"action": f"branch {i}", "expected": "ok"}] for i in range(n_branch)]
        if boundary:
            tcs.append([{"action": "입력 설정 (경계 최솟값): a = 0", "expected": "ok"}])
        _stub_generators(monkeypatch)

        def _gen(lf, info, stats=None, keep_boundary=False, **_k):
            assert keep_boundary is True
            if stats is not None:
                stats["boundary_appended"] = boundary
                stats["boundary_index"] = (len(tcs) - 1) if boundary else None
                stats["boundary_beyond_function_cap"] = bool(boundary and n_branch >= 5)
            return [list(t) for t in tcs]

        monkeypatch.setattr(gsts, "_generate_steps_from_flow", _gen)
        fd = {"f1": {"name": "s_foo", "prototype": "U16 s_foo(U16 x)", "logic_flow": [], "calls_list": []}}
        return _build_doc_proposal(_proposal_sections(fd), {"s_foo"}, **kw)

    def test_boundary_tc_takes_the_last_slot_when_the_cap_cuts(self, monkeypatch):
        out = self._proposal(monkeypatch, n_branch=9)
        shown, meta = out["sts"]["s_foo"], out["sts_meta"]["s_foo"]
        assert len(shown) == 6 and shown[-1][0]["action"].startswith("입력 설정 (경계")
        assert [t[0]["action"] for t in shown[:5]] == [f"branch {i}" for i in range(5)]
        assert (meta["gen_total"], meta["gen_truncated"], meta["boundary_tc_index"]) == (10, True, 5)
        assert meta["boundary_tc_extended_only"] is True            # 함수당 상한 밖 — 기본 프로파일 문서엔 없다(리뷰 W3)

    def test_no_cut_no_reordering(self, monkeypatch):
        out = self._proposal(monkeypatch, n_branch=2)
        assert [t[0]["action"] for t in out["sts"]["s_foo"]][:2] == ["branch 0", "branch 1"]
        assert out["sts_meta"]["s_foo"]["boundary_tc_index"] == 2
        assert out["sts_meta"]["s_foo"]["boundary_tc_extended_only"] is False

    def test_no_boundary_tc_means_none_and_a_plain_head_cut(self, monkeypatch):
        out = self._proposal(monkeypatch, n_branch=9, boundary=False)
        assert out["sts_meta"]["s_foo"]["boundary_tc_index"] is None
        assert out["sts_meta"]["s_foo"]["boundary_tc_extended_only"] is False
        assert [t[0]["action"] for t in out["sts"]["s_foo"]] == [f"branch {i}" for i in range(6)]


class TestBoundaryIndexFromTheRealGenerator:
    """(리뷰 G4·G5·I5) 위 가드는 생성기를 스텁으로 갈아 끼운다 — 여기선 **실제 생성기**가 자리를 옳게 말하는지 본다."""
    INFO = {"name": "f", "prototype": "U8 f(U8 a, U8 b)", "inputs": ["[IN] U8 a", "[IN] U8 b"], "outputs": ["[OUT] U8 ret"]}

    @staticmethod
    def _flow(n):
        return [{"type": "if", "condition": f"a == {i}", "then": [{"type": "action", "text": f"x = {i}"}], "else": []} for i in range(n)]

    def test_no_flow_function_reports_the_second_slot(self):
        from generators.sts import _generate_steps_from_flow

        st: dict = {}
        tcs = _generate_steps_from_flow([], dict(self.INFO), stats=st, keep_boundary=True)
        assert st["boundary_index"] == 1 and tcs[1][0]["action"].startswith("입력 설정 (경계")
        assert st["boundary_appended"] is False                     # 요구당 상한 계수의 뜻은 그대로

    def test_no_flow_function_without_a_boundary_tc_reports_none(self):
        from generators.sts import _generate_steps_from_flow

        st: dict = {}
        _generate_steps_from_flow([], dict(self.INFO), max_steps=3, stats=st)
        assert st["boundary_index"] is None

    def test_flow_function_index_and_beyond_cap_flag(self):
        from generators.sts import _generate_steps_from_flow

        st: dict = {}
        kept = _generate_steps_from_flow(self._flow(8), dict(self.INFO), max_tc=3, stats=st, keep_boundary=True)
        bi = st["boundary_index"]
        assert kept[bi][0]["action"].startswith("입력 설정 (경계") and bi == len(kept) - 1
        assert st["boundary_beyond_function_cap"] is True and len(kept) == 4
        # 같은 함수를 **문서가 쓰는 호출**(기본 프로파일)로 부르면 그 TC 는 없다 — 초안이 `extended_only` 라 말해야 하는 이유
        st2: dict = {}
        doc = _generate_steps_from_flow(self._flow(8), dict(self.INFO), max_tc=3, stats=st2)
        assert len(doc) == 3 and st2["boundary_index"] is None
        assert not any(t[0]["action"].startswith("입력 설정 (경계") for t in doc)

    def test_within_the_cap_is_not_extended_only(self):
        from generators.sts import _generate_steps_from_flow

        st: dict = {}
        _generate_steps_from_flow(self._flow(1), dict(self.INFO), max_tc=5, stats=st, keep_boundary=True)
        assert st["boundary_index"] is not None and st["boundary_beyond_function_cap"] is False


# ── N91 · N105: 로컬 핸들러 ───────────────────────────────────────────────────
_LOCAL_DOC_HANDLERS = ("local_sts_generate", "local_sts_generate_stream", "local_sts_generate_async",
                       "local_suts_generate", "local_suts_generate_stream", "local_suts_generate_async",
                       "local_sits_generate", "local_sits_generate_stream", "local_sits_generate_async")


def _handler_source(name: str) -> str:
    from backend.routers import local
    from tests.unit._source_probe import source_of

    return source_of(getattr(local, name))


class TestLocalHandlers:
    @pytest.mark.parametrize("name", _LOCAL_DOC_HANDLERS)
    def test_every_handler_takes_the_profile_as_an_optional_form_field(self, name):
        from backend.routers import local

        p = inspect.signature(getattr(local, name)).parameters["tc_profile"]
        assert getattr(p.default, "default", None) == ""           # Form("") — 안 보내면 정본 규모

    @pytest.mark.parametrize("name", _LOCAL_DOC_HANDLERS)
    def test_every_handler_forwards_it(self, name):
        src = _handler_source(name)
        forwarded = '"tc_profile": tc_profile,' if "_sts_" in name else "tc_profile=tc_profile,"
        assert src.count(forwarded) == 1, name

    @pytest.mark.parametrize("name", [n for n in _LOCAL_DOC_HANDLERS if "_suts_" in n])
    def test_local_suts_never_discovers_a_foreign_hsis(self, name):
        src = _handler_source(name)
        assert src.count('_no_discovery, label="HSIS")') == 1, name
        assert "_discover_hsis_path" not in src

    # (R77 N111) STS·SITS 의 자동 탐색도 끊었다 — 9곳 전부를 `test_followups_r77.py::TestNoForeignHsis` 가 묶는다.

    def test_unspecified_hsis_resolves_to_none_and_a_bad_path_still_warns(self, caplog):
        from backend.routers.local import _doc_or_discovered, _no_discovery

        assert _doc_or_discovered(None, "", _no_discovery, label="HSIS") is None
        with caplog.at_level("WARNING"):
            assert _doc_or_discovered(None, "U:/x/hsis.xlsx", _no_discovery, label="HSIS") is None
        assert any("해석하지 못해" in r.getMessage() for r in caplog.records)
        assert _doc_or_discovered("C:/given.xlsx", "C:/given.xlsx", _no_discovery, label="HSIS") == "C:/given.xlsx"
