"""R60 — MC/DC pairs whose two rows differ in an output the function writes (`generators/observable_mcdc.py`,
extended SUTS only).

``gate`` hides its outer decision behind an inner one: the designed pair of ``(g_req == 0U) && (g_en == 1U)`` leaves
``g_inner`` at its default (0), so the inner ``if`` is false on both rows and ``g_out`` is 0 on both — a wrong operator in
the outer conditions would pass. The extended profile keeps the pair and adds two rows with the same decision inputs and
``g_inner`` taken from another row of the test case, on which ``g_out`` differs (7 / 0).
``step`` is the state-machine shape (review R60 C1): the decision reads ``g_state``, which the function writes only behind
the inner condition — rows that differ in ``g_state`` but never write it only echo their inputs; that is no difference."""
from __future__ import annotations

import os

import openpyxl
from openpyxl.styles import Border, Font, PatternFill

from generators import c_project_context as cpc
from generators import observable_mcdc as om
from generators import suts
from generators.suts import generate_sequences, summarize_mcdc_design

ROOT = os.path.join(os.sep, "proj")
COMMON = """#ifndef COMMON_H
#define COMMON_H
typedef unsigned char U8;
typedef unsigned int U16;
typedef signed int S16;
#endif
"""
UNIT = """#include "common.h"
U8 g_req;
U8 g_en;
U8 g_inner;
U8 g_out;
U8 g_state;
U8 g_x;
void gate(void) { g_out = 0U; if ((g_req == 0U) && (g_en == 1U)) { if (g_inner > 5U) { g_out = 7U; } } }
void flat(void) { if ((g_req == 0U) && (g_en == 1U)) { g_out = 7U; } else { g_out = 0U; } }
void step(void) { if ((g_state == 0U) && (g_en == 1U)) { if (g_inner > 5U) { g_state = 2U; } } }
void local(void) { U8 t = (U8)(g_x + 1U); g_out = 0U; if ((t > 5U) && (g_en == 1U)) { if (g_inner > 5U) { g_out = 7U; } } }
U8 g_sw;
U8 g_old;
void latch(void) { g_out = 0U; if ((g_sw == 1U) && (g_old == 0U)) { if (g_inner > 5U) { g_out = 7U; } } g_old = g_sw; }
void flag(void) { if ((g_req == 0U) && (g_en == 1U)) { g_out = 1U; } else { g_out = 0U; } }
U8 g_mode;
void branchy(void)
{
    g_out = 0U;
    if (g_mode == 1U) { g_out = (U8)(g_en + 3U); }
    else { if ((g_req == 0U) && (g_en == 1U)) { if (g_inner == 200U) { g_out = 7U; } } }
}
void elsewhere(void)
{
    if (g_mode == 0U) { g_out = (U8)(g_en + 3U); }
    else { if ((g_req == 0U) && (g_en == 1U)) { g_out = 7U; } else { g_out = 0U; } }
}
"""


def _unit(name="gate", inputs=("g_req", "g_en", "g_inner"), outputs=("g_out",)):
    files = {os.path.join(ROOT, "common.h"): COMMON, os.path.join(ROOT, "unit.c"): UNIT}
    context = cpc.build_project_context(files)
    path = os.path.join(ROOT, "unit.c")
    return {"name": name, "fid": "F1", "source_text": UNIT, "source_path": path, "source_text_complete": True,
            "project_scope": cpc.build_scopes(context, [path])[path], "input_vars": list(inputs),
            "output_vars": list(outputs), "param_types": {v: "U8" for v in (*inputs, *outputs)}}


def _ext(unit):
    return generate_sequences(unit, None, type_cache={}, extended=True)


def _outer(unit, var="g_req"):
    return next(d for d in unit["mcdc_design"]["decisions"] if var in d["expression"])


def _value(seq, var="g_out"):
    ev = (seq.get("expected_evidence") or {}).get(var) or {}
    return seq["expected"].get(var) if ev.get("status") == "derived" else None


def _obs_rows(seqs):
    return [s for s in seqs if s["strategy"].startswith(suts.OBS_MCDC_PREFIX)]


def test_a_masked_outer_pair_gets_two_rows_whose_outputs_differ():
    unit = _unit()
    seqs = _ext(unit)
    by_num = {s["seq_num"]: s for s in seqs}
    outer = _outer(unit)
    marks = [p["observable"] for p in outer["pairs"]]
    assert {m["status"] for m in marks} == {"searched"}
    for pair, mark in zip(outer["pairs"], marks, strict=True):
        a, b = by_num[mark["seq_a"]], by_num[mark["seq_b"]]
        assert a["strategy"].startswith(suts.OBS_MCDC_PREFIX) and b["strategy"].startswith(suts.OBS_MCDC_PREFIX)
        # the decision's inputs keep the pair's values; only what the decision expression does not read moved
        for var in outer["variables"]:
            assert int(a["inputs"][var]) == pair["inputs_a"][var] and int(b["inputs"][var]) == pair["inputs_b"][var]
        assert mark["overlay_inputs"] == ["g_inner"] and mark["evaluation"] == "expression"
        assert {_value(a), _value(b)} == {7, 0}
        assert a["mcdc_observable"]["pair_id"] == pair["pair_id"] and a["tc_profile"] == suts.TC_PROFILE_EXTENDED
        # an extended-profile, MC/DC-technique row by its name (R75 rule; review R60 W1)
        assert suts.is_extended_strategy(a["strategy"]) and suts.is_extended_strategy(b["strategy"])
        assert suts.resolve_seq_gen_method(a["strategy"]) == suts._GEN_EQUIV
        assert "결정식이 읽지 않는 입력" in a["description"] and "경로 주장 아님" in a["description"]
    # the designed pair's own rows really were masked
    for pair in outer["pairs"]:
        assert _value(by_num[pair["seq_a"]]) == _value(by_num[pair["seq_b"]]) == 0


def test_a_pair_whose_rows_already_differ_is_marked_and_adds_nothing():
    unit = _unit("flat", ("g_req", "g_en"))
    seqs = _ext(unit)
    (d,) = unit["mcdc_design"]["decisions"]
    assert {p["observable"]["status"] for p in d["pairs"]} == {"observable"}
    assert all(p["observable"]["seq_a"] == p["seq_a"] for p in d["pairs"])
    assert not _obs_rows(seqs)
    assert unit["observable_mcdc"]["rows"] == 0 and unit["observable_mcdc"]["observable"] == 2


def test_an_output_that_only_echoes_its_input_is_no_difference():
    # review R60 C1: the g_state pair's rows differ in g_state (0 / 1) but neither run writes it — searched, and the
    #   found rows differ because one run writes g_state = 2
    unit = _unit("step", ("g_state", "g_en", "g_inner"), ("g_state",))
    seqs = _ext(unit)
    by_num = {s["seq_num"]: s for s in seqs}
    d = _outer(unit, "g_state")
    pair = next(p for p in d["pairs"] if p["condition_id"] == "C1")
    assert pair["observable"]["status"] == "searched"
    a, b = by_num[pair["observable"]["seq_a"]], by_num[pair["observable"]["seq_b"]]
    bases = {s["expected_evidence"]["g_state"]["basis"] for s in (a, b)}
    assert "assigned" in bases and {_value(a, "g_state"), _value(b, "g_state")} >= {2}


def test_a_decision_designed_on_the_model_is_checked_again_for_the_new_rows(monkeypatch):
    # `t` is a local: the decision is designed on the modeled run; the overlay may move inputs the decision reads
    #   through `t` only if the model still evaluates it as the pair claims
    unit = _unit("local", ("g_x", "g_en", "g_inner"))
    _ext(unit)
    d = _outer(unit, "t > 5U")
    assert d.get("evaluation") == "source_path"
    statuses = {p["observable"]["status"] for p in d["pairs"]}
    assert "searched" in statuses
    assert all(p["observable"].get("evaluation") == "source_path" for p in d["pairs"]
               if p["observable"]["status"] == "searched")
    # (review R60 W7 M2 · round 2) the re-check refusing every candidate: no searched pair is claimed — patched only
    #   inside the observable search, so the design keeps its retained pairs (the precondition is asserted)
    import generators.c_source_oracle as cso
    real_find = om.find_observable_pairs

    def refusing(*a, **k):
        real_observe = cso.observe_decisions
        cso.observe_decisions = lambda unit, vectors, *rest, **kw: [{"status": "unsupported"}] * len(vectors)
        try:
            return real_find(*a, **k)
        finally:
            cso.observe_decisions = real_observe
    monkeypatch.setattr(om, "find_observable_pairs", refusing)
    again = _unit("local", ("g_x", "g_en", "g_inner"))
    _ext(again)
    pairs = _outer(again, "t > 5U")["pairs"]
    assert pairs and all(p["retained_status"] == "retained" for p in pairs)
    statuses = {p["observable"]["status"] for p in pairs}
    assert "searched" not in statuses and "recheck_refused" in statuses


def test_the_decision_inputs_are_kept_out_of_the_overlay():
    # review R60 round 2 (mutant M1 `fixed = set()`): a base row that would also close the gate (g_en = 0) — only
    #   g_inner may move, and the pair is found; moving g_en too would close the gate on both rows
    unit = _unit()
    seqs = _ext(unit)
    outer = _outer(unit)
    pair = next(p for p in outer["pairs"] if p["condition_id"] == "C1")
    by_num = {s["seq_num"]: s for s in seqs}
    a_inputs = {k: int(v) for k, v in by_num[pair["seq_a"]]["inputs"].items()}
    for p in outer["pairs"]:
        p.pop("observable", None)
    base = [{"seq_num": 999, "inputs": {**a_inputs, "g_en": 0, "g_inner": 9}}]
    domains = {"g_req": (0, 255), "g_en": (0, 255), "g_inner": (0, 255)}
    found = om.find_observable_pairs(unit, seqs, ["g_out"], domains, base)
    mark = next(m for p, m in found["marks"] if p is pair)
    assert mark["status"] == "searched" and mark["overlay_inputs"] == ["g_inner"]


def test_values_outside_the_domain_are_never_used_and_masked_is_not_underived():
    # review R60 round 2 (mutants M5 `inside() -> True` · M6 masked/underived swapped): g_inner may only move inside
    #   (0, 3) — the inner condition (> 5) never opens, every candidate derives g_out equal on both rows: masked
    unit = _unit()
    seqs = _ext(unit)
    outer = _outer(unit)
    by_num = {s["seq_num"]: s for s in seqs}
    pair = next(p for p in outer["pairs"] if p["condition_id"] == "C1")
    a_inputs = {k: int(v) for k, v in by_num[pair["seq_a"]]["inputs"].items()}
    # a base outside the domain (g_inner 9 would open the inner write) comes first, one inside it second
    bases = [{"seq_num": 998, "inputs": {**a_inputs, "g_inner": 9}}, {"seq_num": 997, "inputs": {**a_inputs, "g_inner": 2}}]
    found = om.find_observable_pairs(unit, seqs, ["g_out"], {"g_inner": (0, 3)}, bases)
    marks = {id(p): m for p, m in found["marks"]}
    assert marks[id(pair)]["status"] == "masked" and marks[id(pair)]["candidates_tried"] == 1
    # no value for the overlay at all: no_candidate (neither masked nor underived)
    empty = om.find_observable_pairs(unit, seqs, ["g_out"], {"g_inner": (0, 3)}, [])
    assert {m["status"] for p, m in empty["marks"] if p is pair} == {"no_candidate"}
    # an output no row derives: nothing to compare — underived
    none = om.find_observable_pairs(unit, seqs, ["g_nope"], {"g_inner": (0, 255)}, bases)
    marks = {id(p): m for p, m in none["marks"]}
    assert marks[id(pair)]["status"] == "underived"


def test_a_latch_copying_the_moved_input_is_no_difference():
    # review R60 round 2 W1: `g_old = g_sw` differs whenever g_sw does, whatever the decision did — the g_sw pair is
    #   searched (g_inner opens the inner write), not taken as observable
    unit = _unit("latch", ("g_sw", "g_old", "g_inner"), ("g_out", "g_old"))
    _ext(unit)
    d = _outer(unit, "g_sw")
    pair = next(p for p in d["pairs"] if p["condition_id"] == "C1")
    assert pair["observable"]["status"] == "searched"
    assert "g_out" in pair["observable"]["outputs_changed"] and "g_old" not in pair["observable"]["outputs_changed"]


def test_a_pair_within_one_row_and_a_unit_without_outputs_are_not_checked():
    unit = _unit()
    seqs = _ext(unit)
    outer = _outer(unit)
    pair = outer["pairs"][0]
    pair["seq_b"] = pair["seq_a"]                    # two evaluations of one run (a loop) — review R60 W4
    found = om.find_observable_pairs(unit, seqs, ["g_out"], {}, seqs)
    mark = next(m for p, m in found["marks"] if p is pair)
    assert mark == {"status": "not_checked", "reason": "pair_within_one_row"}
    none = om.find_observable_pairs(unit, seqs, [], {}, seqs)
    assert {m["reason"] for _p, m in none["marks"]} == {"no_outputs"}
    assert none["report"]["not_checked"] == none["report"]["pairs"] == len(none["marks"]) > 0
    # (review R60 round 2 I3) a pair row that is not in the document is not checked
    pair["seq_a"] = pair["seq_b"] = None
    pair["seq_a"] = 9999
    found = om.find_observable_pairs(unit, seqs, ["g_out"], {}, seqs)
    assert next(m for p, m in found["marks"] if p is pair) == {"status": "not_checked", "reason": "rows_or_variables_unknown"}


def test_the_rows_before_are_the_rows_without_the_search(monkeypatch):
    unit = _unit()
    seqs = _ext(unit)
    monkeypatch.setattr(suts, "_append_observable_mcdc_rows", lambda *a, **k: None)
    plain = _ext(_unit())
    assert [(s["strategy"], s["inputs"], s["expected"]) for s in seqs[:len(plain)]] == \
        [(s["strategy"], s["inputs"], s["expected"]) for s in plain]
    assert all(s["strategy"].startswith(suts.OBS_MCDC_PREFIX) for s in seqs[len(plain):])


def test_the_reference_profile_has_no_observable_search():
    unit = _unit()
    seqs = generate_sequences(unit, None, type_cache={})
    assert "observable_mcdc" not in unit
    assert not any("observable" in p for d in unit["mcdc_design"]["decisions"] for p in d.get("pairs") or [])
    assert not _obs_rows(seqs)


def test_a_failing_search_leaves_no_rows_and_no_marks(monkeypatch):
    def fail(*_a, **_k):
        raise RuntimeError("down")
    monkeypatch.setattr(om, "find_observable_pairs", fail)
    unit = _unit()
    seqs = _ext(unit)
    assert unit["observable_mcdc"] == {"error": "RuntimeError", "rows": 0}
    assert not _obs_rows(seqs)
    assert not any("observable" in p for d in unit["mcdc_design"]["decisions"] for p in d.get("pairs") or [])
    assert summarize_mcdc_design([unit])["observable_errors"] == 1


def test_a_failure_after_rows_were_placed_removes_them(monkeypatch):
    # review R60 W7 M7: the rows are already appended when the label fails
    def fail(*_a, **_k):
        raise RuntimeError("label")
    monkeypatch.setattr(suts, "_observable_row_label", fail)
    unit = _unit()
    seqs = _ext(unit)
    assert unit["observable_mcdc"]["error"] == "RuntimeError" and not _obs_rows(seqs)
    assert not any("observable" in p for d in unit["mcdc_design"]["decisions"] for p in d.get("pairs") or [])


def test_the_evaluation_budget_per_function_is_counted_per_pair(monkeypatch):
    real = om.find_observable_pairs

    def small(*a, **k):
        k["max_evaluations"] = 3
        return real(*a, **k)
    monkeypatch.setattr(om, "find_observable_pairs", small)
    unit = _unit()
    _ext(unit)
    statuses = [p["observable"]["status"] for d in unit["mcdc_design"]["decisions"] for p in d.get("pairs") or []
                if "observable" in p]
    assert "budget_exhausted" in statuses
    assert unit["observable_mcdc"]["evaluation_budget_exhausted"] is True
    assert unit["observable_mcdc"]["budget_exhausted"] == statuses.count("budget_exhausted")


def test_evidence_that_does_not_confirm_downgrades_the_pair_in_the_pipeline(monkeypatch):
    # review R60 W6/W7: after the evidence of the new rows, a pair whose rows do not differ is downgraded and its new
    #   rows no longer claim a pair
    real = suts.apply_sequence_evidence

    def flatten(unit, seqs):
        out = real(unit, seqs)
        for s in out:
            if str(s.get("strategy") or "").startswith(suts.OBS_MCDC_PREFIX):
                s["expected"]["g_out"] = 0
        return out
    monkeypatch.setattr(suts, "apply_sequence_evidence", flatten)
    unit = _unit()
    seqs = _ext(unit)
    outer = _outer(unit)
    assert {p["observable"]["status"] for p in outer["pairs"]} == {"evidence_mismatch"}
    assert unit["observable_mcdc"]["evidence_mismatch"] >= 2
    assert all("근거 불일치" in s["description"].split("\n")[0] for s in _obs_rows(seqs))
    s = summarize_mcdc_design([unit])
    assert s["observable_evidence_mismatch"] >= 2
    # (review R60 round 4 W1) the breakdown of searched pairs adds up to the searched count after downgrades
    assert sum(s["observable_searched_own_rows"].values()) == s["observable_searched"]


def test_summary_sheet_and_disclosure():
    from report_gen import generation_disclosures as gd
    from report_gen.generation_disclosures import build_disclosures
    unit = _unit()
    seqs = _ext(unit)
    s = summarize_mcdc_design([unit])
    # two outer pairs + the inner pair (its rows had the gate closed: g_en from another row opens it; both members
    #   already exist as rows — no row added for it)
    assert s["observable_searched"] == 3 and s["observable_rows"] == unit["observable_mcdc"]["rows"] == 3
    assert s["observable_evaluations"] > 0
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": s}) if i["key"] == "suts_mcdc_design"]
    for text in ("'쓴 출력' 에서 갈리는지", "탐색으로 찾은 쌍 3", "더한 행 3", "결정 탓으로 돌리지 않는다", "Observable",
                 "원인은 가르지 않는다", "래치처럼 입력을 복사한 출력"):
        assert text in item["note"], text
    assert (gd.MAX_OBSERVABLE_EVALUATIONS, gd.MAX_OBSERVABLE_CANDIDATES) == (om.MAX_EVALUATIONS, om.MAX_CANDIDATES)
    wb = openpyxl.Workbook()
    suts._write_mcdc_design_sheet(wb, [unit], {"F1": seqs}, {"F1": "TC_1"}, Border(), PatternFill(), Font(), Font())
    values = list(wb["MCDC Design"].iter_rows(values_only=True))
    assert list(values[0][-2:]) == ["Observable", "Observable Sequences"]
    cells = {(r[2], r[3]): (r[-2], r[-1]) for r in values[1:]}
    outer = _outer(unit)
    for p in outer["pairs"]:
        assert cells[(outer["decision_id"], p["condition_id"])] == (
            f"searched [own:{p['observable']['own_rows']}]", f"{p['observable']['seq_a']},{p['observable']['seq_b']}")


def test_a_flag_equal_to_its_moved_input_is_reported_as_copy_only():
    # review R60 round 3 W1: g_out = 1/0 while g_en = 1/0 — the copy rule cannot tell this flag from a latch; the pair is
    #   not counted observable, and says why
    unit = _unit("flag", ("g_req", "g_en"))
    _ext(unit)
    (d,) = unit["mcdc_design"]["decisions"]
    c2 = next(p for p in d["pairs"] if p["condition_id"] == "C2")
    assert c2["observable"]["status"] == "copy_only" and c2["observable"]["own_rows"] == "copy_only"
    assert summarize_mcdc_design([unit])["observable_copy_only"] >= 1


def test_an_expression_designed_pair_is_only_claimed_where_the_rows_reach_the_decision():
    # review R60 round 3 W2: a base row with g_mode = 1 makes the outputs differ (g_en + 3) without the second decision
    #   ever being evaluated — the model re-check refuses it; only rows that take the decision count
    unit = _unit("branchy", ("g_mode", "g_req", "g_en", "g_inner"))
    seqs = _ext(unit)
    by_num = {s["seq_num"]: s for s in seqs}
    d = _outer(unit, "g_req")
    assert d.get("evaluation") != "source_path" and d.get("reach_spec")
    for p in d["pairs"]:
        mark = p["observable"]
        if mark["status"] in ("observable", "searched"):
            for n in (mark["seq_a"], mark["seq_b"]):
                assert int(by_num[n]["inputs"]["g_mode"]) != 1


def test_the_echo_guard_holds_where_the_input_cell_is_not_a_plain_integer():
    # review R60 round 3: `"1U"` is no integer to `_as_int`, so the copy rule cannot fire — the echo rule still keeps an
    #   unwritten output from counting
    ev = {"g_x": {"status": "derived", "basis": "unchanged_input"}}
    a = {"inputs": {"g_x": "0U"}, "expected": {"g_x": 0}, "expected_evidence": ev}
    b = {"inputs": {"g_x": "1U"}, "expected": {"g_x": 1}, "expected_evidence": ev}
    assert om.rows_differ(a, b) == []
    ev2 = {"g_x": {"status": "derived", "basis": "assigned"}}
    assert om.rows_differ({**a, "expected_evidence": ev2}, {**b, "expected_evidence": ev2}) == ["g_x"]


def test_one_key_for_the_same_row():
    # review R60 round 3 W3: a formatted cell and its integer are the same row
    assert om.vector_key({"g_b": "0x1", "g_a": 3}) == om.vector_key({"g_a": 3, "g_b": 1}) == (("g_a", 3), ("g_b", 1))
    assert om.vector_key({"g_a": "x"}) is None


def test_a_confirmation_failure_is_counted_and_disclosed(monkeypatch):
    # review R60 round 3 W4
    def fail(*_a, **_k):
        raise RuntimeError("confirm")
    monkeypatch.setattr(suts, "_confirm_observable_rows", fail)
    unit = _unit()
    _ext(unit)
    assert unit["observable_mcdc"]["confirm_error"] == "RuntimeError"
    assert summarize_mcdc_design([unit])["observable_errors"] == 1


def test_the_sheet_says_why_a_pair_was_not_checked():
    assert suts._observable_cells({"status": "not_checked", "reason": "pair_within_one_row"}) == \
        ["not_checked:pair_within_one_row", ""]
    assert suts._observable_cells({"status": "searched", "seq_a": 3, "seq_b": 4}) == ["searched", "3,4"]
    assert suts._observable_cells({"status": "masked", "own_rows": "equal"}) == ["masked [own:equal]", ""]


def test_rows_that_differ_without_taking_the_decision_are_not_observable():
    # review R60 round 3 W2: the designed rows leave g_mode at 0 — the decision sits in the else branch, so the rows
    #   never evaluate it, yet g_out = g_en + 3 differs (4 / 3); that is not this decision's pair showing
    unit = _unit("elsewhere", ("g_mode", "g_req", "g_en"))
    _ext(unit)
    d = _outer(unit, "g_req")
    c2 = next(p for p in d["pairs"] if p["condition_id"] == "C2")
    assert c2["observable"]["status"] != "observable" and c2["observable"].get("own_rows") == "unreached"
