"""TD-P2: independently evaluate generated MC/DC inputs (synthetic QM).

These assertions do not equate source-level design pairs with target execution
coverage. Predicates are handwritten from each test requirement, independent of
the production condition parser and its reported coverage labels.
"""
from itertools import combinations

import pytest

from generators.suts import generate_sequences


def _integer(value):
    return int(value, 0) if isinstance(value, str) else value


def _generate(condition, names, **extra):
    unit = {"name": "decision", "input_vars": names, "output_vars": [],
            "param_types": dict.fromkeys(names, "uint8_t"),
            "logic_flow": [{"type": "if", "condition": condition}], **extra}
    rows = generate_sequences(unit, max_seq=None, type_cache={})
    return unit, [row for row in rows if row["strategy"].startswith("MCDC")]


def _covered(rows, atoms, decision):
    vectors = [tuple(atom({k: _integer(v) for k, v in row["inputs"].items()}) for atom in atoms)
               for row in rows]
    covered = set()
    for left, right in combinations(vectors, 2):
        changed = [i for i in range(len(atoms)) if left[i] != right[i]]
        if len(changed) == 1 and decision(left) != decision(right):
            covered.add(changed[0])
    return covered


@pytest.mark.parametrize("condition, names, atoms, decision", [
    ("a > 10 || b > 20", ["a", "b"],
     [lambda x: x["a"] > 10, lambda x: x["b"] > 20], lambda v: v[0] or v[1]),
    ("a > 10 && b > 20", ["a", "b"],
     [lambda x: x["a"] > 10, lambda x: x["b"] > 20], lambda v: v[0] and v[1]),
    ("!(a > 10) || b > 20", ["a", "b"],
     [lambda x: x["a"] > 10, lambda x: x["b"] > 20], lambda v: not v[0] or v[1]),
    ("(a > 10 && b > 20) || c == 3", ["a", "b", "c"],
     [lambda x: x["a"] > 10, lambda x: x["b"] > 20, lambda x: x["c"] == 3],
     lambda v: (v[0] and v[1]) or v[2]),
    ("a > 10 && a < 20", ["a"],
     [lambda x: x["a"] > 10, lambda x: x["a"] < 20], lambda v: v[0] and v[1]),
    ("a == 0 || b != 255", ["a", "b"],
     [lambda x: x["a"] == 0, lambda x: x["b"] != 255], lambda v: v[0] or v[1]),
])
def test_td_p2_generated_pairs_demonstrate_independent_effect(condition, names, atoms, decision):
    _, rows = _generate(condition, names)
    assert rows, "No concrete MC/DC candidate rows produced"
    assert all(0 <= _integer(value) <= 255 for row in rows for value in row["inputs"].values())
    assert _covered(rows, atoms, decision) == set(range(len(atoms)))


def test_td_p2_mcdc_candidates_stay_in_authoritative_design_domain():
    _, rows = _generate("a >= 10 && a <= 20", ["a"],
                        uds_param_info={"a": {"type": "uint8_t", "range": [10, 20]}})
    assert all(10 <= _integer(row["inputs"]["a"]) <= 20 for row in rows)
    assert _covered(rows, [lambda x: x["a"] >= 10, lambda x: x["a"] <= 20], all) == set()


def test_td_p2_bare_boolean_conditions_and_not_have_real_pairs():
    _, rows = _generate("a || !b", ["a", "b"], param_types={"a": "bool", "b": "bool"})
    assert all(_integer(value) in (0, 1) for row in rows for value in row["inputs"].values())
    assert _covered(rows, [lambda x: bool(x["a"]), lambda x: bool(x["b"])],
                    lambda v: v[0] or not v[1]) == {0, 1}


def test_td_p2_repeated_atoms_cannot_claim_unique_cause_independence():
    unit, rows = _generate("a > 10 || a > 10", ["a"])
    assert _covered(rows, [lambda x: x["a"] > 10, lambda x: x["a"] > 10], any) == set()
    report = unit["mcdc_design"]
    assert len(report["decisions"][0]["conditions"]) == 2
    assert not report["decisions"][0]["pairs"]


def test_td_p2_cap_never_claims_a_pair_without_both_emitted_rows():
    unit = {"name": "f", "input_vars": ["a", "b"], "output_vars": [],
            "param_types": {"a": "uint8_t", "b": "uint8_t"},
            "logic_flow": [{"type": "if", "condition": "a > 10 || b > 20"}]}
    rows = generate_sequences(unit, max_seq=2, type_cache={})
    assert len(rows) <= 2
    by_number = {row["seq_num"]: row for row in rows}
    report = unit["mcdc_design"]
    for decision in report["decisions"]:
        assert decision["execution_status"] == "not_run"
        for pair in decision["pairs"]:
            if pair["retained_status"] == "retained":
                assert pair["seq_a"] in by_number and pair["seq_b"] in by_number
                actual_rows = [by_number[pair["seq_a"]], by_number[pair["seq_b"]]]
                assert _covered(actual_rows, [lambda x: x["a"] > 10, lambda x: x["b"] > 20], any)
            else:
                assert pair["retained_status"] == "truncated"


def test_td_p2_retained_pair_metadata_matches_actual_rows_and_independent_predicates():
    unit, rows = _generate("a > 10 || b > 20", ["a", "b"])
    by_number = {row["seq_num"]: row for row in rows}
    decision = unit["mcdc_design"]["decisions"][0]
    assert len(decision["conditions"]) == 2
    assert decision["reachability"] == "unverified"
    assert decision["execution_status"] == "not_run"
    assert decision["pairs"]
    for pair in decision["pairs"]:
        assert pair["retained_status"] == "retained"
        left, right = by_number[pair["seq_a"]], by_number[pair["seq_b"]]
        assert left["inputs"] == pair["inputs_a"]
        assert right["inputs"] == pair["inputs_b"]
        assert _covered([left, right], [lambda x: x["a"] > 10, lambda x: x["b"] > 20], any)


def test_td_p2_unsatisfiable_atomic_condition_stays_uncovered():
    unit, rows = _generate("a > 255 || b > 20", ["a", "b"])
    assert all(0 <= _integer(value) <= 255 for row in rows for value in row["inputs"].values())
    assert _covered(rows, [lambda x: x["a"] > 255, lambda x: x["b"] > 20], any) == {1}
    decision = unit["mcdc_design"]["decisions"][0]
    impossible_id = decision["conditions"][0]["condition_id"]
    assert all(pair["condition_id"] != impossible_id for pair in decision["pairs"])


def test_td_p2_search_cap_does_not_report_exhaustive_search():
    from generators.mcdc_design import build_mcdc_design

    unit = {"name": "f", "input_vars": ["a", "b"], "output_vars": [],
            "param_types": {"a": "uint8_t", "b": "uint8_t"},
            "logic_flow": [{"type": "if", "condition": "a > 10 || b > 20"}]}
    report = build_mcdc_design(unit, max_candidates=1)
    decision = report["decisions"][0]
    assert decision["search_complete"] is False
    assert not decision["pairs"]


def test_td_p2_unknown_pointer_domain_cannot_supply_fabricated_inputs():
    unit, rows = _generate("a > 10 || b > 20", ["a", "b"],
                          param_types={"a": "Opaque *", "b": "uint8_t"})
    assert not rows
    assert not unit["mcdc_design"]["decisions"][0]["pairs"]


@pytest.fixture
def mcdc_workbook(tmp_path):
    from generators.suts import generate_suts_xlsm

    unit = {"fid": "SwUFn_0001", "name": "f", "input_vars": ["a", "b"], "output_vars": [],
            "param_types": {"a": "uint8_t", "b": "uint8_t"},
            "logic_flow": [{"type": "if", "condition": "a > 10 || b > 20"}]}
    rows = generate_sequences(unit, max_seq=None, type_cache={})
    path = tmp_path / "mcdc.xlsx"
    generate_suts_xlsm(None, [unit], {unit["fid"]: rows}, str(path), {"project_id": "QM"})
    return path


def test_td_p2_exported_design_requires_source_revalidation_not_execution_claim(mcdc_workbook):
    from tools.export_suts_vectorcast import build_vectorcast_model

    model = build_vectorcast_model(str(mcdc_workbook))
    evidence = model["units"][0]["mcdc_design_evidence"]
    assert evidence
    assert all(item["verification_status"] == "source_revalidation_required" for item in evidence)
    assert all(item["execution_status"] == "not_run" for item in evidence)
    assert all(item["executed_mcdc_coverage"] is None for item in evidence)


@pytest.mark.parametrize("mutation", ["input", "execution"])
def test_td_p2_import_cannot_trust_edited_input_or_forged_execution(mcdc_workbook, mutation):
    import openpyxl

    from tools.export_suts_vectorcast import build_vectorcast_model

    wb = openpyxl.load_workbook(mcdc_workbook)
    sheet = wb["MCDC Design"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    selected = next(row for row in range(2, sheet.max_row+1) if sheet.cell(row, headers["Pair ID"]).value)
    if mutation == "input":
        sheet.cell(selected, headers["Inputs A JSON"]).value = '{"a":999,"b":999}'
    else:
        execution_col = headers.get("Execution", sheet.max_column+1)
        sheet.cell(1, execution_col).value = "Execution"
        sheet.cell(selected, execution_col).value = "100% MC/DC executed"
    wb.save(mcdc_workbook)
    wb.close()
    evidence = build_vectorcast_model(str(mcdc_workbook))["units"][0]["mcdc_design_evidence"]
    assert evidence
    assert all(item["execution_status"] == "not_run" and item["executed_mcdc_coverage"] is None for item in evidence)
    if mutation == "input":
        assert any(item["verification_status"] == "conflict" and item["reason"] == "stale_mcdc_pair_input" for item in evidence)


@pytest.fixture(scope="module")
def compiled_decision_probe(tmp_path_factory):
    import shutil
    import subprocess

    compiler = shutil.which("gcc") or shutil.which("clang")
    if compiler is None:
        pytest.skip("C compiler unavailable: compiled QM decision probe not executed")
    directory = tmp_path_factory.mktemp("mcdc_c_probe")
    source = directory / "probe.c"
    executable = directory / "probe.exe"
    source.write_text("""#include <stdio.h>
static unsigned int observed;
static int atom(unsigned int bit, int truth) { observed |= 1u << bit; return truth; }
int main(void) {
    unsigned int mode, a, b, c;
    while (scanf("%u %u %u %u", &mode, &a, &b, &c) == 4) {
        observed = 0;
        int result = mode == 0 ? (atom(0, a > 10) || atom(1, b > 20))
            : ((atom(0, a > 10) && atom(1, b > 20)) || atom(2, c == 3));
        printf("%d %u\\n", result, observed);
    }
    return 0;
}
""", encoding="utf-8")
    subprocess.run([compiler, "-std=c99", "-O0", str(source), "-o", str(executable)],
                   check=True, capture_output=True, text=True, timeout=30)
    return executable


@pytest.mark.parametrize("mode, condition, names", [
    (0, "a > 10 || b > 20", ["a", "b"]),
    (1, "(a > 10 && b > 20) || c == 3", ["a", "b", "c"]),
])
def test_td_p2_compiled_c_confirms_generated_decision_and_short_circuit_observation(
        compiled_decision_probe, mode, condition, names):
    import subprocess

    unit, rows = _generate(condition, names)
    by_number = {row["seq_num"]: row for row in rows}
    pairs = unit["mcdc_design"]["decisions"][0]["pairs"]
    assert pairs
    inputs, expected = [], []
    for pair in pairs:
        for side in ("a", "b"):
            values = {name: _integer(value) for name, value in by_number[pair[f"seq_{side}"]]["inputs"].items()}
            inputs.append(f"{mode} {values['a']} {values['b']} {values.get('c', 0)}")
            expected.append((pair[f"decision_{side}"], pair[f"observed_{side}"]))
    process = subprocess.run([str(compiled_decision_probe)], input="\n".join(inputs)+"\n",
                             check=True, capture_output=True, text=True, timeout=10)
    actual = [tuple(map(int, line.split())) for line in process.stdout.splitlines()]
    assert len(actual) == len(expected)
    for (result, mask), (decision, observed) in zip(actual, expected, strict=True):
        assert bool(result) == bool(decision)
        assert [bool(mask & (1 << i)) for i in range(len(names))] == observed


def test_td_p2_duplicate_actual_sequence_cannot_select_arbitrary_pair_member(mcdc_workbook):
    import openpyxl

    from generators.suts import _DATA_START_ROW, _SEQ_COL
    from tools.export_suts_vectorcast import build_vectorcast_model

    wb = openpyxl.load_workbook(mcdc_workbook)
    evidence = wb["MCDC Design"]
    headers = {cell.value: cell.column for cell in evidence[1]}
    selected = next(row for row in range(2, evidence.max_row+1) if evidence.cell(row, headers["Pair ID"]).value)
    sequence = evidence.cell(selected, headers["Sequence A"]).value
    sheet = wb["2.SW Unit Test Spec"]
    source_row = next(row for row in range(_DATA_START_ROW+1, sheet.max_row+1)
                      if sheet.cell(row, _SEQ_COL).value == sequence)
    sheet.append([cell.value for cell in sheet[source_row]])
    wb.save(mcdc_workbook)
    wb.close()
    evidence = build_vectorcast_model(str(mcdc_workbook))["units"][0]["mcdc_design_evidence"]
    assert any(item["verification_status"] == "conflict" and item["reason"] == "ambiguous_mcdc_sequence" for item in evidence)


# ── R80: 배선 후 계약 — 분모 유지·절단 공시·탐색 규모 ──────────────────────────

def test_r80_long_and_chain_is_fully_designed_within_budget():
    """조건 8개 AND 사슬 — 전 곱(5^8)을 도는 탐색은 예산 4096 안에서 첫 변수를 한 번도 못 바꿔 쌍 0 이었다.

    변수로 이어진 원자끼리만 묶어 찾으면(성분 탐색) 조건마다 독립 영향 쌍이 나온다. 쌍은 독립 술어로 다시 확인한다.
    """
    from generators.mcdc_design import build_mcdc_design

    names = [f"v{i}" for i in range(8)]
    unit = {"name": "f", "input_vars": names, "output_vars": [], "param_types": dict.fromkeys(names, "uint8_t"),
            "logic_flow": [{"type": "if", "condition": " && ".join(f"v{i} > {i}" for i in range(8))}]}
    decision = build_mcdc_design(unit)["decisions"][0]
    assert decision["status"] == "designed", decision["reason"]
    atoms = [lambda x, i=i: x[f"v{i}"] > i for i in range(8)]
    for pair in decision["pairs"]:
        i = int(pair["condition_id"][1:]) - 1
        a, b = pair["inputs_a"], pair["inputs_b"]
        assert atoms[i](a) != atoms[i](b)
        assert all(atoms[j](a) == atoms[j](b) for j in range(8) if j != i)
        assert all(atoms[j](a) for j in range(8)) != all(atoms[j](b) for j in range(8))


def test_r80_sheet_keeps_unsupported_decisions_in_the_denominator(tmp_path):
    import openpyxl

    from generators.suts import generate_suts_xlsm

    unit = {"fid": "SwUFn_0002", "name": "g", "input_vars": ["a", "b"], "output_vars": [],
            "param_types": {"a": "uint8_t", "b": "uint8_t"},
            "logic_flow": [{"type": "if", "condition": "a > 10 || b > 20"},
                           {"type": "if", "condition": "a > LIMIT"}]}
    rows = generate_sequences(unit, max_seq=None, type_cache={})
    path = tmp_path / "den.xlsx"
    generate_suts_xlsm(None, [unit], {unit["fid"]: rows}, str(path), {"project_id": "QM"})
    wb = openpyxl.load_workbook(path)
    ws = wb["MCDC Design"]
    headers = [c.value for c in ws[1]]
    records = [dict(zip(headers, r, strict=False)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    unsupported = [r for r in records if r["Decision ID"] == "D2"]
    assert len(unsupported) == 1 and not unsupported[0]["Pair ID"]
    assert unsupported[0]["Decision Status"] == "unsupported"
    assert unsupported[0]["Reason"].startswith("missing_declared_domain")
    assert {r["Execution"] for r in records} == {"not_run"}
    assert {r["Reachability"] for r in records} == {"unverified"}


def test_r80_truncated_pairs_are_disclosed_not_exported_as_conflict_or_coverage(tmp_path):
    from generators.suts import generate_suts_xlsm, summarize_mcdc_design
    from tools.export_suts_vectorcast import build_vectorcast_model

    names = ["a", "b", "c"]
    unit = {"fid": "SwUFn_0003", "name": "h", "input_vars": names, "output_vars": [],
            "param_types": dict.fromkeys(names, "uint8_t"),
            "logic_flow": [{"type": "if", "condition": "(a > 10 && b > 20) || c == 3"}]}
    rows = generate_sequences(unit, max_seq=7, type_cache={})
    summary = summarize_mcdc_design([unit])
    assert summary["decisions"] == 1 and summary["conditions_paired"] == 3
    assert summary["truncated_pairs"] >= 1, "상한 7 이면 BV 6 뒤 MC/DC 행 1개만 남아 쌍이 잘려야 한다"
    assert summary["retained_pairs"] + summary["truncated_pairs"] + summary["invalidated_pairs"] == 3
    path = tmp_path / "trunc.xlsx"
    generate_suts_xlsm(None, [unit], {unit["fid"]: rows}, str(path), {"project_id": "QM"})
    evidence = build_vectorcast_model(str(path))["units"][0]["mcdc_design_evidence"]
    truncated = [e for e in evidence if e["verification_status"] == "not_retained"]
    assert truncated and all(e["reason"] == "mcdc_pair_truncated" for e in truncated)
    assert not any(e["verification_status"] == "conflict" for e in evidence)
    assert all(e["executed_mcdc_coverage"] is None for e in evidence)


# ── R80 리뷰 후속: C1·W1~W5 · 캐시 스키마 · 손실 흐름 텍스트 ─────────────────────

def test_r80_c1_row_label_never_claims_a_pair_whose_partner_row_was_cut():
    """(리뷰 C1) 라벨은 finalize 뒤에 쓴다 — 짝 행이 잘린 벡터가 "독립 영향 쌍" 이라고 적히면 MCDC Design 시트와 모순."""
    names = ["a", "b", "c"]
    unit = {"name": "h", "input_vars": names, "output_vars": [], "param_types": dict.fromkeys(names, "uint8_t"),
            "logic_flow": [{"type": "if", "condition": "(a > 10 && b > 20) || c == 3"}]}
    full = generate_sequences(dict(unit), max_seq=None, type_cache={})
    head = sum(1 for r in full if not r["strategy"].startswith("MCDC_"))
    rows = generate_sequences(unit, max_seq=head + 2, type_cache={})   # MC/DC 두 행만 남긴다 → 일부 쌍은 짝을 잃는다
    mcdc = [r for r in rows if r["strategy"].startswith("MCDC_")]
    assert len(mcdc) == 2
    assert any(not r["mcdc_design"] for r in mcdc) or all(r["mcdc_design"] for r in mcdc)
    for row in mcdc:
        first = row["description"].splitlines()[0]
        if row["mcdc_design"]:
            assert first.startswith("MC/DC 독립 영향 쌍")
        else:
            assert "독립 영향 쌍:" not in first and "미성립" in first, first


def test_r80_w1_source_parameter_type_wins_over_a_same_named_global():
    """(리뷰 W1) 소스가 `int8_t mode` 라고 선언했으면 호출자 도메인(같은 이름 전역의 U16)이 이기지 못한다."""
    from generators.mcdc_design import build_mcdc_design

    src = "#include <stdint.h>\nint8_t f(int8_t mode) { if (mode > 100) { return 1; } return 0; }\n"
    unit = {"name": "f", "input_vars": ["mode"], "source_text": src, "source_text_complete": True}
    report = build_mcdc_design(unit, declared_domains={
        "mode": {"min": 0, "max": 65535, "type": "uint16_t", "source": "declared_type", "origin": "global"}})
    assert report["domains"]["mode"]["max"] == 127
    assert all(-128 <= v["mode"] <= 127 for v in report["selected_inputs"])


def test_r80_w2_pointer_parameter_gets_no_integer_domain():
    names = ["pv"]
    unit = {"name": "f", "input_vars": names, "output_vars": [], "param_types": {"pv": "U16 *"},
            "logic_flow": [{"type": "if", "condition": "pv"}]}
    rows = generate_sequences(unit, max_seq=None, type_cache={})
    assert not [r for r in rows if r["strategy"].startswith("MCDC_")]
    assert unit["mcdc_design"]["decisions"][0]["status"] == "unsupported"


def test_r80_w3_enum_values_follow_the_design_range():
    from generators.mcdc_design import build_mcdc_design

    unit = {"name": "f", "input_vars": ["m"], "logic_flow": [{"type": "if", "condition": "m > 2"}],
            "uds_param_info": {"m": {"range": [2, 3]}}}
    report = build_mcdc_design(unit, declared_domains={
        "m": {"min": 0, "max": 3, "values": [0, 1, 2, 3], "type": "enum", "source": "enum_declaration",
              "origin": "parameter"}})
    assert report["domains"]["m"]["values"] == [2, 3]
    assert {v["m"] for v in report["selected_inputs"]} <= {2, 3}
    assert report["decisions"][0]["status"] == "designed"


def test_r80_w4_unused_enum_input_defaults_to_an_enumerator():
    from generators.mcdc_design import build_mcdc_design

    unit = {"name": "f", "input_vars": ["a", "e"], "param_types": {"a": "uint8_t"},
            "logic_flow": [{"type": "if", "condition": "a > 1"}]}
    report = build_mcdc_design(unit, declared_domains={
        "e": {"min": -2, "max": 5, "values": [-2, 5], "type": "enum", "source": "enum_declaration",
              "origin": "parameter"}})
    assert report["selected_inputs"] and {v["e"] for v in report["selected_inputs"]} <= {-2, 5}


def test_r80_w5_source_parameter_outside_unit_inputs_is_unsupported_not_truncated():
    from generators.mcdc_design import build_mcdc_design

    src = "#include <stdint.h>\nuint8_t f(uint8_t a, uint8_t b) { if (a > 1 || b > 2) { return 1; } return 0; }\n"
    unit = {"name": "f", "input_vars": ["a"], "source_text": src, "source_text_complete": True}
    decision = build_mcdc_design(unit)["decisions"][0]
    assert decision["status"] == "unsupported"
    assert decision["reason"].startswith("decision_variable_not_in_unit_inputs")


def test_r80_lossy_flow_rendering_is_reported_as_such():
    """logic_flow 는 `&&` 를 AND 로 적고 긴 식을 `...` 로 자른다 — 그 이유가 증상(call_expression)에 가려지지 않는다."""
    from generators.mcdc_design import build_mcdc_design

    for text in ("( a == 1 ) AND ( b == 2 )", "( a == 1 ) OR (...)"):
        unit = {"name": "f", "input_vars": ["a", "b"], "param_types": {"a": "uint8_t", "b": "uint8_t"},
                "logic_flow": [{"type": "if", "condition": text}]}
        assert build_mcdc_design(unit)["decisions"][0]["reason"] == "lossy_condition_text"
    # 단어 경계 — 식별자 안의 AND/OR 는 손실 표시가 아니다.
    unit = {"name": "f", "input_vars": ["ERROR_FLAG"], "param_types": {"ERROR_FLAG": "uint8_t"},
            "logic_flow": [{"type": "if", "condition": "ERROR_FLAG > 0"}]}
    assert build_mcdc_design(unit)["decisions"][0]["status"] == "designed"


def test_r80_design_range_conflict_keeps_decisions_and_adds_no_phantom():
    from generators.mcdc_design import build_mcdc_design

    unit = {"name": "f", "input_vars": ["a"], "param_types": {"a": "uint8_t"},
            "uds_param_info": {"a": {"range": [0, 1000]}},
            "logic_flow": [{"type": "if", "condition": "a > 10"}]}
    report = build_mcdc_design(unit)
    assert report["domains"]["a"]["design_range_conflict"]["max"] == 1000
    assert report["domains"]["a"]["max"] == 255
    assert report["decisions"][0]["status"] == "designed"
    empty = build_mcdc_design({"name": "g", "input_vars": ["a"], "param_types": {"a": "uint8_t"},
                               "uds_param_info": {"a": {"range": [0, 1000]}}, "logic_flow": []})
    assert empty["decisions"] == []


def test_r80_source_sections_cache_schema_carries_source_text(tmp_path):
    """P1 이 함수 레코드에 원문을 추가하고 캐시 버전을 안 올려 실 생성이 v25 캐시를 히트했다 — 원문이 도달하지 않았다."""
    from backend.helpers.uds import _SOURCE_SECTIONS_SCHEMA_VERSION
    from report_gen.uds_generator import generate_uds_source_sections

    assert int(_SOURCE_SECTIONS_SCHEMA_VERSION.lstrip("v")) >= 26
    (tmp_path / "m.c").write_text("#include <stdint.h>\nuint8_t f(uint8_t a) { return a; }\n", encoding="utf-8")
    sections = generate_uds_source_sections(str(tmp_path))
    rec = next(v for v in sections["function_details"].values() if v.get("name") == "f")
    assert rec["source_text_complete"] is True
    # (리뷰 N2) 원문은 파일당 한 번 — 함수 레코드에 전문을 싣지 않는다(캐시 14배·사이드카에 소스 사본).
    assert "source_text" not in rec
    assert "uint8_t f(uint8_t a)" in sections["source_files"][rec["source_path"]]

    from generators.suts import attach_unit_sources
    units = [{"name": "f", "source_path": rec["source_path"]}, {"name": "g", "source_path": "missing.c"}]
    assert attach_unit_sources(units, sections["source_files"]) == 1
    assert units[0]["source_text_complete"] is True and "return a" in units[0]["source_text"]
    assert units[1]["source_unavailable_reason"] == "source_file_not_in_source_stage"


def test_r80_n1_empty_input_list_never_falls_back_to_source_parameters():
    from generators.mcdc_design import build_mcdc_design

    src = "#include <stdint.h>\nuint8_t f(uint8_t a, uint8_t b) { if (a > 1 || b > 2) { return 1; } return 0; }\n"
    report = build_mcdc_design({"name": "f", "input_vars": [], "source_text": src, "source_text_complete": True})
    assert report["selected_inputs"] == []
    assert report["decisions"][0]["reason"].startswith("decision_variable_not_in_unit_inputs")


def test_r80_n3_unenumerated_functions_are_not_counted_as_decisions():
    from generators.mcdc_design import build_mcdc_design
    from generators.suts import summarize_mcdc_design

    no_decision = {"name": "f", "input_vars": [], "source_text": '#include "proj.h"\nvoid f(void) { }\n'}
    unparsable = {"name": "missing_fn", "input_vars": [], "source_text": "void other(void) { }\n"}
    for unit in (no_decision, unparsable):
        unit["mcdc_design"] = build_mcdc_design(unit)
    assert no_decision["mcdc_design"]["decisions"] == []
    assert unparsable["mcdc_design"]["decisions"][0]["status"] == "unenumerated"
    summary = summarize_mcdc_design([no_decision, unparsable])
    assert summary["decisions"] == 0 and summary["unenumerated_functions"] == 1
    assert summary["unenumerated_reasons"] == {"function_identity_missing_or_ambiguous": 1}


def test_r80_w_b_enumeration_exception_keeps_the_function_visible(monkeypatch):
    """(리뷰 R3 W-B) 결정 열거 중 예외가 나도 함수가 분모에서 사라지지 않는다 — unenumerated 로 남는다."""
    import generators.mcdc_design as md

    def boom(unit, declared_domains=None):
        raise RecursionError("deep")
    monkeypatch.setattr(md, "_source_decisions", boom)
    report = md.build_mcdc_design({"name": "f", "input_vars": []})
    assert [d["status"] for d in report["decisions"]] == ["unenumerated"]
    assert report["decisions"][0]["reason"] == "source_exception:RecursionError"
    assert report["summary"]["total_decisions"] == 0 and report["summary"]["unenumerated"] == 1


def test_r80_unit_without_attached_source_uses_flow_path_not_incomplete_marker():
    from generators.mcdc_design import build_mcdc_design

    unit = {"name": "f", "input_vars": ["a"], "param_types": {"a": "uint8_t"}, "source_text": "",
            "source_text_complete": False, "logic_flow": [{"type": "if", "condition": "a > 1"}]}
    decision = build_mcdc_design(unit)["decisions"][0]
    assert decision["source_kind"] == "logic_flow" and decision["status"] == "designed"
