"""R51 — measurement requirements: the listed inputs, the output that reports them, the conversion, a judgment tolerance.

HDPDM01 / KJPDS02 ``SwEI_01`` (Battery Power Source) verification criteria, lost until now to a label spelling
(``Verification Criteria`` — the parser read only ``Verification criteria``):

    입력 전원 인가 (9.0, 12.0, 16.0)
    Canoe로 Output_Battery Voltage(9.0, 12.0, 16.0)값 확인
    (0x0~0xFF, offset : 4V, Resolution : 0x05V, 4V ~ 16.75V)

The output reports the input: at each listed value the expected code is (v − 4) / 0.05 (0x64 · 0xA0 · 0xF0), judged
within the HW measurement accuracy of the input's monitor path (R49/R50, reused) plus one resolution step (the rounding
is not written). ``0x05V`` read with the range and the codes is 0.05 V — kept, and a review item says so.
"""
from __future__ import annotations

from collections import Counter

import docx
import openpyxl

from generators.sts import _build_tc_dict, _classify_steps, _make_tc_id, parse_srs_docx_tables
from generators.sts_requirement_tc import (
    MEASURE_PREFIX,
    append_requirement_boundary_tcs,
    measurement_groups,
    measurement_spec,
    write_requirement_review_sheet,
)

HD_VC = ("입력 전원 인가 (9.0, 12.0, 16.0)\nCanoe로 Output_Battery Voltage(9.0, 12.0, 16.0)값 확인\n"
         "(0x0~0xFF, offset : 4V, Resolution : 0x05V, 4V ~ 16.75V)")
PV_VC = HD_VC.replace("(9.0, 12.0, 16.0)\nCanoe", "(9.0V, 12.0V, 16.0V)\nCanoe")
BATTERY = {"name": "Battery Voltage Monitor", "related": ["SyEI_01", "SySM_04"],
           "tolerances": [{"text": "허용 오차: ±3%", "value": "3", "unit": "%", "measured_units": ["V"]}],
           "ranges": [{"lo": "9", "hi": "16", "unit": "V"}], "nominals": [], "monitor_node": False,
           "full_text": "Battery Voltage Monitor", "scales": []}


def _req(vc=HD_VC, related="SyEI_01, SyII_03", rid="SwEI_01"):
    return {"id": rid, "name": "Battery Power Source", "related_id": related, "asil": "",
            "description": "Battery 전압 ADC 값은 u16g_ApiIn_Vsup 16bit변수를 입력 받아 내부 계산에 의해 8bit로 출력", "verification": vc}


def test_the_parser_reads_the_verification_criteria_whatever_its_capitals(tmp_path):
    d = docx.Document()
    for rows in ([("ID", "SwEI_01"), ("Name", "Battery Power Source"), ("Description", "d"), ("Verification Criteria", HD_VC)],
                 [("ID", "SwTR_0101"), ("Name", "n"), ("Description", "d"), ("Verification criteria", "v")]):
        t = d.add_table(rows=len(rows), cols=2)
        for row, (k, v) in zip(t.rows, rows, strict=True):
            row.cells[0].text, row.cells[1].text = k, v
    d.save(tmp_path / "srs.docx")
    reqs = {r["id"]: r for r in parse_srs_docx_tables(str(tmp_path / "srs.docx"))}
    assert reqs["SwEI_01"]["verification"] == HD_VC and reqs["SwTR_0101"]["verification"] == "v"


def test_the_spec_is_read_whole_or_not_at_all():
    spec = measurement_spec(_req())
    assert {k: spec[k] for k in ("input", "output", "unit", "values", "code_low", "code_high", "low", "high", "resolution",
                                 "resolution_written")} == {
        "input": "입력 전원", "output": "Output_Battery Voltage", "unit": "V", "values": ["9", "12", "16"],
        "code_low": 0, "code_high": 255, "low": "4", "high": "16.75", "resolution": "0.05", "resolution_written": "0x05V"}
    assert measurement_spec(_req(PV_VC))["values"] == ["9", "12", "16"]
    # the output checked at other values · an offset that is not the range's low end · no conversion: no spec
    assert measurement_spec(_req(HD_VC.replace("Voltage(9.0, 12.0, 16.0)", "Voltage(9.0, 12.0)"))) is None
    assert measurement_spec(_req(HD_VC.replace("offset : 4V", "offset : 3V"))) is None
    assert measurement_spec(_req(HD_VC.split("\n(0x0")[0])) is None
    assert measurement_spec(_req(HD_VC.replace("16.0)\nCanoe", "17.0)\nCanoe").replace("12.0, 16.0)값", "12.0, 17.0)값"))) is None
    # a resolution written as the range gives it: nothing to review
    assert measurement_spec(_req(HD_VC.replace("0x05V", "0.05V")))["resolution_written"] is None


def test_each_listed_value_is_a_step_with_its_code_and_band():
    stats, review = Counter(), []
    (g,) = measurement_groups(_req(), stats, review=review)
    actions = [s["action"] for s in g["steps"]]
    assert actions[0].startswith(MEASURE_PREFIX + "입력 전원 = 9V — 근거: ")
    expected = [s["expected"] for s in g["steps"]]
    # no HW requirement given: the band is one resolution step — and the review item says the HW accuracy is missing
    assert expected[0].startswith("Output_Battery Voltage = 9V(코드 0x64) — 판정 범위 8.95V ~ 9.05V(코드 0x63 ~ 0x65)")
    assert "(코드 0xA0)" in expected[1] and "(코드 0xF0)" in expected[2]
    # (review W4) no HW requirements given is not 'none linked'
    assert [r["reason"] for r in review] == ["measurement_resolution_written_otherwise", "measurement_hw_not_given"]
    assert "HW 측정 정확도 없음(HW 요구사항서 미입력) + 분해능 0.05V" in expected[0] and "hw_tolerance" not in g["evidence"]
    stats, review = Counter(), []
    (g,) = measurement_groups(_req(), stats, hw={}, review=review)
    assert review[-1]["reason"] == "measurement_tolerance_unlinked" and g["evidence"]["hw_tolerance"] == []
    assert "HW 측정 정확도 없음(이어진 HW 블록 없음)" in g["steps"][0]["expected"]
    assert "HSIS" not in review[-1]["if_filled"]               # (review r2 W-A)
    assert (stats["measurement_tcs"], stats["measurement_steps"], stats["measurement_hw_linked"]) == (1, 3, 0)
    assert g["evidence"]["kind"] == "measurement" and all(p["holds"] is None for p in g["evidence"]["points"])


def test_the_band_takes_the_hw_accuracy_of_the_monitor_path():
    """HwTSR_0204 (Battery Voltage Monitor ±3%) cites SyEI_01, as SwEI_01 does: at 9 V ±0.27 + 0.05 → 8.68 ~ 9.32 V,
    codes 0x5E ~ 0x6A (a code is in the band only when its value is)."""
    stats, review = Counter(), []
    (g,) = measurement_groups(_req(), stats, hw={"HwTSR_0204": BATTERY}, review=review)
    assert "판정 범위 8.68V ~ 9.32V(코드 0x5E ~ 0x6A): HW 측정 ±0.27V(HwTSR_0204) + 분해능 0.05V" in g["steps"][0]["expected"]
    assert "판정 범위 15.47V ~ 16.53V" in g["steps"][2]["expected"]
    assert [r["reason"] for r in review] == ["measurement_resolution_written_otherwise"]
    assert stats["measurement_hw_linked"] == 1
    # two blocks left (a hub): the largest, said undecided — and a review item
    other = {**BATTERY, "name": "Supply Monitor", "tolerances": [{**BATTERY["tolerances"][0], "text": "허용 오차: ±5%",
                                                                   "value": "5"}]}
    stats, review = Counter(), []
    (g,) = measurement_groups(_req(), stats, hw={"HwTSR_0204": BATTERY, "HwTSR_0205": other}, review=review)
    assert "HW 측정 ±0.45V(HwTSR_0204 · HwTSR_0205 — 감시 경로 미정, 가장 큰 값)" in g["steps"][0]["expected"]
    assert "measurement_path_undecided" in [r["reason"] for r in review]


def test_the_measurement_tc_is_its_own_requirement_based_tc():
    """Not a boundary analysis (BAA): the requirement's own listed values — AOR, executed (RBT), one TC of its own."""
    tcs = []

    def build(**kw):
        return _build_tc_dict(test_env="HILS", derive_inputs=False, is_safety=False, **kw)
    review = []
    stats = append_requirement_boundary_tcs(tcs, [_req()], build, _make_tc_id, _classify_steps, review_out=review)
    (tc,) = [t for t in tcs if t["requirement_evidence"][0]["kind"] == "measurement"]
    assert (tc["test_method"], tc["gen_method"], len(tc["steps"])) == ("RBT", "AOR", 3)
    assert stats["measurement_tcs"] == 1
    wb = openpyxl.Workbook()
    write_requirement_review_sheet(wb, review)
    rows = list(wb["Requirement Review"].iter_rows(values_only=True))[1:]
    reasons = [r[4] for r in rows if r[1] == "SwEI_01"]
    assert any(r.startswith("분해능 표기가 범위·코드와 맞지 않음") for r in reasons)


def test_the_disclosure_says_what_the_band_holds():
    """The board learns which measurement TCs judge with the resolution alone — and what linking the HW block gives."""
    from backend.services.issue_explainer import _RULE_ACTION
    from report_gen.generation_disclosures import build_disclosures
    items = {i["key"]: i for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
        "tcs": 1, "steps": 3, "facts": 1, "facts_used": 1, "measurement_tcs": 1, "measurement_steps": 3,
        "measurement_hw_linked": 0}}})}
    item = items["sts_measurement"]
    assert (item["value"], item["tone"]) == ("측정 TC 1 · 스텝 3", "warning")
    # (review r2 W-A) the HSIS links nothing by itself — no promise of it
    assert "판정 범위가 분해능뿐" in item["note"] and "Related ID 에 적으면" in item["note"] and "HSIS" not in item["note"]
    assert "disclosure:sts_measurement" in _RULE_ACTION
    none = {i["key"] for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
        "tcs": 1, "steps": 3, "facts": 1, "facts_used": 1}}})}
    assert "sts_measurement" not in none


def test_a_pass_criterion_of_a_non_functional_requirement_is_measured_not_set():
    """``SwNTR_0301`` verification (read since the parser takes ``Verification Criteria``): ``Load(Average)는 70% 이하.``
    is what a test measures — CPU load cannot be set. A '측정' step per line, no boundary points (no discrimination
    claimed); a SW variable or a functional requirement keeps R33's reading (a test input)."""
    from generators.sts_requirement_tc import ACCEPT_PREFIX, boundary_steps
    ntr = {"id": "SwNTR_0301", "name": "CPU 관리", "description": "CPU 부하 관리", "asil": "", "related_id": "",
           "verification": "1) CPU 부하율\n- Load(Average)는 70% 이하.\n- Load(Peak)는 90% 이하.\n"
                           "3) 메모리 점유율\n- DDM, ADM : 88% 이하.\n- 그 외 : Proto 60% 이하, Master 80% 이하."}
    stats = Counter()
    groups = boundary_steps(ntr, stats)
    crit = [g for g in groups if g["evidence"]["kind"] == "acceptance"]
    steps = [st for g in crit for st in g["steps"]]
    assert all(st["action"].startswith(ACCEPT_PREFIX) for st in steps) and all(g["evidence"]["points"] == [] for g in crit)
    assert not [x for x in groups if x["evidence"]["kind"] != "acceptance"]          # nothing stepped as an input
    assert "Load(Average)는 70% 이하." in steps[0]["action"]
    assert "원문 판정 기준을 만족 — 'CPU 부하율 — Load(Average)는 70% 이하.'" in steps[0]["expected"]
    # (review W5) the clause as written · (r2 W-B) under the heading that says what is measured
    assert any("'메모리 점유율 — DDM, ADM : 88% 이하.'" in st["expected"] for st in steps)
    assert stats["skipped:acceptance_criterion"] == stats["acceptance_criteria"] >= 4 and len(crit) == stats["acceptance_steps"]
    # a SW variable: a test input (SIL sets it); a functional requirement's bare line: a test input too
    assert not [x for x in boundary_steps(dict(ntr, verification="- u16s_C: 30 초과")) if x["evidence"]["kind"] == "acceptance"]
    assert not [x for x in boundary_steps(dict(ntr, id="SwTR_0101")) if x["evidence"]["kind"] == "acceptance"]
    # its own TC, a requirement-based one (AOR / RBT)
    tcs = []

    def build(**kw):
        return _build_tc_dict(test_env="HILS", derive_inputs=False, is_safety=False, **kw)
    append_requirement_boundary_tcs(tcs, [ntr], build, _make_tc_id, _classify_steps)
    (tc,) = [t for t in tcs if t["requirement_evidence"][0]["kind"] == "acceptance"]     # the criteria, one TC
    assert (tc["test_method"], tc["gen_method"]) == ("RBT", "AOR") and len(tc["steps"]) == len(tc["requirement_evidence"])
    wb = openpyxl.Workbook()
    from generators.sts_requirement_tc import write_requirement_evidence_sheet
    write_requirement_evidence_sheet(wb, [tc])
    assert list(wb["Requirement Evidence"].iter_rows(values_only=True))[1][10] in ("", None)   # no points


def test_the_cross_scorer_takes_no_point_from_a_measurement_or_a_criterion(tmp_path):
    """A measurement step (``입력 전원 = 9V``) and a pass criterion judge no threshold: no point, no region — their
    ``= 9V`` read as a named point credited 'optimistic' kills (HDPDM01 103 → 121)."""
    from generators.sts_requirement_tc import ACCEPT_PREFIX, MEASURE_PREFIX
    from scripts.requirement_oracle_eval import read_sts
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3.SW Test Spec"
    for _ in range(3):
        ws.append([None] * 13)
    ws.append([None] * 10 + [MEASURE_PREFIX + "입력 전원 = 9V — 근거: x", "Output = 9V(코드 0x64)", "SwEI_01"])
    ws.append([None] * 10 + [ACCEPT_PREFIX + "- Load(Average)는 70% 이하.", "측정값이 [70% 이하] 를 만족", "SwNTR_0301"])
    ws.append([None] * 10 + ["입력 설정 (요구 경계): u16g_X = 9V", "조건 [u16g_X 9V 이상] 성립 → 동작", "SwTR_0101"])
    wb.save(tmp_path / "g.xlsx")
    got = read_sts(str(tmp_path / "g.xlsx"))
    assert not got.get("SwEI_01", {}).get("points") and not got.get("SwNTR_0301", {}).get("points")
    assert len(got["SwTR_0101"]["points"]) == 1                     # a boundary step still counts


def test_the_measurement_is_its_own_tc_and_its_points_carry_no_verdict():
    """A boundary fact of the same requirement never shares the measurement's TC (one TC says one kind of thing); the
    sheet writes the listed values with no verdict as '—', never 'F'."""
    from generators.sts_requirement_tc import write_requirement_evidence_sheet
    tcs = []

    def build(**kw):
        return _build_tc_dict(test_env="HILS", derive_inputs=False, is_safety=False, **kw)
    req = _req()
    req["description"] += "\n- u16g_ApiIn_Vsup 가 8.50V 미만 시 저전압"
    append_requirement_boundary_tcs(tcs, [req], build, _make_tc_id, _classify_steps)
    kinds = [[e["kind"] for e in tc["requirement_evidence"]] for tc in tcs]
    assert ["measurement"] in kinds and all(k == ["measurement"] or "measurement" not in k for k in kinds)
    wb = openpyxl.Workbook()
    write_requirement_evidence_sheet(wb, [t for t in tcs if t["requirement_evidence"][0]["kind"] == "measurement"])
    assert list(wb["Requirement Evidence"].iter_rows(values_only=True))[1][10] == "9(—), 12(—), 16(—)"


def test_the_conversion_line_is_the_measurements_not_a_review_item():
    """(review W1) ``0x0~0xFF`` and ``4V ~ 16.75V`` of the conversion line were subjectless review items whose advice
    (write a subject in) broke the measurement spec — the line is the measurement TC's."""
    from generators.sts_requirement_tc import boundary_steps
    stats, review = Counter(), []
    boundary_steps(_req(), stats, review=review)
    assert not [r for r in review if "0x0~0xFF" in r["line"]] and stats["skipped:measurement_conversion"] >= 1


def test_a_criterion_line_without_a_subject_is_still_a_criterion():
    """(review W2) ``2) Slack Time`` / ``- Task 할당 시간의 10% 이상.``: the subject is the line above — measured, so no
    review item advising a subject 'for a boundary step'."""
    from generators.sts_requirement_tc import ACCEPT_PREFIX, boundary_steps
    ntr = {"id": "SwNTR_0301", "name": "CPU 관리", "description": "CPU 부하 관리", "asil": "", "related_id": "",
           "verification": "2) Slack Time\n- Task 할당 시간의 10% 이상."}
    stats, review = Counter(), []
    groups = boundary_steps(ntr, stats, review=review)
    assert [s["action"].startswith(ACCEPT_PREFIX) for g in groups for s in g["steps"]] == [True] and not review


def test_a_linked_block_on_another_scale_is_said_so():
    """(review W6) the block is linked but its tolerance is a monitor node's (a divider unknown): not 'no accuracy'."""
    node = {"name": "Battery Monitor Node", "related": ["SyEI_01"],
            "tolerances": [{"text": "허용 오차: ±0.1V", "value": "0.1", "unit": "V", "measured_units": []}],
            "ranges": [], "nominals": [{"value": "2.5", "unit": "V"}], "monitor_node": True, "full_text": "", "scales": []}
    stats, review = Counter(), []
    (g,) = measurement_groups(_req(), stats, hw={"HwX_01": node}, review=review)
    assert "HW 측정 정확도 척도 불명(HwX_01 — 분압비 확인) + 분해능 0.05V" in g["steps"][0]["expected"]
    assert review[-1]["reason"] == "measurement_scale_unknown" and stats["measurement_scale_unknown"] == 1


def test_an_exact_quotient_gives_the_codes_and_a_resolution_written_to_its_precision_is_right():
    """(review I1) 10 bits over 0~5V: 5/1023 written ``0.004888V`` is right to its precision; 2.5V is code 0x200 from
    the exact quotient (the 6-decimal rounding gave 0x1FF)."""
    vc = ("입력 전원 인가 (2.5, 4.0)\nCanoe로 Output_Sense Voltage(2.5, 4.0)값 확인\n"
          "(0x0~0x3FF, offset : 0V, Resolution : 0.004888V, 0V ~ 5V)")
    spec = measurement_spec(_req(vc))
    assert spec["resolution_written"] is None and spec["resolution"] == "0.004888"
    stats = Counter()
    (g,) = measurement_groups(_req(vc), stats)
    assert "(코드 0x200)" in g["steps"][0]["expected"]


def test_each_criterion_row_says_its_heading_and_clause():
    """(review r2 W-B) one evidence row per criterion line: what is measured (the heading above it) and the clause —
    no '60% 이하 · 80% 이하' summed into one cell, no row that points at the first line only."""
    from generators.sts_requirement_tc import write_requirement_evidence_sheet
    ntr = {"id": "SwNTR_0301", "name": "CPU 관리", "description": "CPU 부하 관리", "asil": "", "related_id": "",
           "verification": "2) Slack Time\n- Task 할당 시간의 10% 이상.\n3) 메모리 점유율\n- 그 외 : Proto 60% 이하, Master 80% 이하."}
    tcs = []

    def build(**kw):
        return _build_tc_dict(test_env="HILS", derive_inputs=False, is_safety=False, **kw)
    stats = append_requirement_boundary_tcs(tcs, [ntr], build, _make_tc_id, _classify_steps)
    wb = openpyxl.Workbook()
    write_requirement_evidence_sheet(wb, tcs)
    rows = list(wb["Requirement Evidence"].iter_rows(values_only=True))[1:]
    assert [(r[3], r[6]) for r in rows] == [("Slack Time", "Task 할당 시간의 10% 이상."),
                                            ("메모리 점유율", "그 외 : Proto 60% 이하, Master 80% 이하.")]
    from generators.sts_requirement_tc import REQUIREMENT_EVIDENCE_HEADERS
    hw_col = REQUIREMENT_EVIDENCE_HEADERS.index("HW Tolerance (Quoted)")
    assert rows[0][hw_col] == "— (판정 기준 측정 — 해당 없음)" and rows[0][-1] == "—"   # (review r2 I-4) · (R54) AI 열
    # (review r2 W-D) the criteria TC is not a boundary TC
    assert (stats.get("tcs", 0), stats["acceptance_tcs"]) == (0, 1)


def test_a_variable_with_a_particle_stays_a_test_input():
    """(review r2 W-C) ``u16s_C는`` — \\b failed before Hangul: the variable went unseen and the line became a criterion."""
    from generators.sts_requirement_tc import boundary_steps
    ntr = {"id": "SwNTR_0301", "name": "n", "description": "", "asil": "", "related_id": "",
           "verification": "- u16s_C는 30 초과."}
    assert not [g for g in boundary_steps(ntr) if g["evidence"]["kind"] == "acceptance"]


def test_a_fact_beside_the_conversion_is_read_as_any_other():
    """(review r2 W-E) only the facts inside the conversion the measurement reads are its own."""
    from generators.sts_requirement_tc import boundary_steps
    stats = Counter()
    vc = HD_VC.replace("4V ~ 16.75V)", "4V ~ 16.75V) u16g_ApiIn_Vsup 8.5V 미만 시 저전압")
    groups = boundary_steps(_req(vc), stats)
    assert any(g["evidence"]["kind"] == "threshold" and g["evidence"]["value"] in ("8.5", 8.5) for g in groups)


def test_a_review_item_on_a_criterion_line_promises_no_step():
    """(review r2 W-F) ``- CPU Load 60 ~ 80%.`` of a non-functional requirement: written as two comparisons it is a
    criterion again — the 'If Filled' never promises a boundary step."""
    from generators.sts_requirement_tc import boundary_steps
    review = []
    boundary_steps({"id": "SwNTR_0301", "name": "n", "description": "", "asil": "", "related_id": "",
                    "verification": "- CPU Load 60 ~ 80%."}, Counter(), review=review)
    assert review and all(r["fill_check"] == "not_applicable" and "경계 스텝은 만들지 않는다" in r["if_filled"]
                          for r in review)


def test_a_written_resolution_twice_the_range_one_is_flagged():
    """(review r2 I-1) ``0.1V`` for 0.05 V passed the 'written precision' test — within 1 % only."""
    assert measurement_spec(_req(HD_VC.replace("0x05V", "0.1V")))["resolution_written"] == "0.1V"
    assert measurement_spec(_req(HD_VC.replace("0x05V", "0.05V")))["resolution_written"] is None


def test_what_the_criteria_and_the_conversion_use_is_no_unused_fact():
    """(review r2 W-D) the 'unused facts' of the boundary disclosure leave out what the criteria TC and the measurement
    TC use — the same seven criteria were listed as unused next to the item saying they were used."""
    from report_gen.generation_disclosures import build_disclosures
    items = {i["key"]: i for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
        "tcs": 1, "steps": 3, "facts": 10, "facts_used": 1, "skipped:acceptance_criterion": 7,
        "skipped:measurement_conversion": 2, "skipped:kind_range": 1}}})}
    note = items["sts_requirement_boundary"]["note"]
    assert "acceptance_criterion" not in note and "measurement_conversion" not in note and "판정 기준" not in note
    assert "범위 1" in note


def test_a_minus_sign_is_no_bullet():
    """(review r3 W-1) ``- -5% 이상.`` quoted ``5% 이상.`` — a criterion read wrongly (a −3 % measurement would fail)."""
    from generators.sts_requirement_tc import boundary_steps
    groups = boundary_steps({"id": "SwNTR_0301", "name": "n", "description": "", "asil": "", "related_id": "",
                             "verification": "1) 전류 오차\n- -5% 이상.\n- -5% 이상, 5% 이하."})
    expected = [st["expected"] for g in groups for st in g["steps"]]
    assert any("'전류 오차 — -5% 이상.'" in e for e in expected) and any("-5% 이상, 5% 이하." in e for e in expected)


def test_one_sentence_under_two_headings_is_two_criteria():
    """(review r3 W-2) keyed by the sentence's text, ``ROM 사용률``'s criterion merged into ``RAM 사용률``'s."""
    from generators.sts_requirement_tc import boundary_steps
    stats = Counter()
    groups = boundary_steps({"id": "SwNTR_0301", "name": "n", "description": "", "asil": "", "related_id": "",
                             "verification": "1) RAM 사용률\n- 사용률 70% 이하.\n2) ROM 사용률\n- 사용률 70% 이하."}, stats)
    actions = [st["action"] for g in groups for st in g["steps"]]
    assert len(actions) == 2 and "RAM 사용률" in actions[0] and "ROM 사용률" in actions[1]
    assert stats["acceptance_steps"] == 2


def test_a_cited_blocks_criteria_are_no_unused_traced_facts():
    """(review r3 W-3) the traced path listed ``acceptance_criterion`` (an internal key) as an unused fact."""
    from report_gen.generation_disclosures import build_disclosures
    items = build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
        "tcs": 1, "steps": 3, "facts": 3, "facts_used": 1, "system_documents": [{"label": "SyDS", "file": "s.docx",
                                                                                  "blocks": 3}],
        "traced:facts": 4, "traced:facts_used": 1, "traced:tcs": 1, "traced:steps": 3,
        "traced:skipped:acceptance_criterion": 2, "traced:skipped:kind_range": 1}}})
    text = " ".join(str(i.get("note")) for i in items)
    assert "acceptance_criterion" not in text


def test_the_conversion_values_are_said_used():
    """(review r3 Info-a) the conversion line's two values counted in the facts were explained nowhere."""
    from report_gen.generation_disclosures import build_disclosures
    items = {i["key"]: i for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
        "tcs": 1, "steps": 3, "facts": 3, "facts_used": 1, "measurement_tcs": 1, "measurement_steps": 3,
        "measurement_hw_linked": 1, "skipped:measurement_conversion": 2}}})}
    assert "변환식에서 읽힌 값 2 개" in items["sts_measurement"]["note"]


def test_a_criterion_line_starting_with_a_minus_sign_takes_the_heading():
    """(review r3 W-1, follow-up) ``-5% 이상.`` names no subject — it is a sub-line of its heading, and the heading search
    for the next line passes over it (it is no heading either)."""
    from generators.sts_requirement_tc import _criterion_heading
    text = "1) 전류 오차\n-5% 이상.\n- 5% 이하."
    assert _criterion_heading(text, text.index("-5%")) == "전류 오차"
    assert _criterion_heading(text, text.index("- 5%")) == "전류 오차"
    assert _criterion_heading(text, 0) == "" and _criterion_heading(text, 999) == ""


def test_a_cited_blocks_criteria_carry_no_traced_boundary_facts():
    """(review r3 W-3) 'the blocks carrying the traced facts' counted a cited block's pass criteria (no boundary)."""
    from generators.sts_requirement_tc import append_requirement_boundary_tcs
    from tests.unit.test_sts_requirement_tc_r29 import _build, _classify_steps, _make_tc_id, _req
    system = {"SyNT_01": {"doc": "SyDS", "fields": {"Verification criteria": "1) CPU 부하율\n- Load는 70% 이하."}}}
    tcs: list = []
    stats = append_requirement_boundary_tcs(tcs, [_req(desc="", rid="SwNTR_0301", related="SyNT_01")], _build,
                                            _make_tc_id, _classify_steps, system=system)
    assert stats.get("acceptance_tcs") == 1 and "traced_facts_by_block" not in stats
    # (review r4 W-A) a cited block's line is found at its own place in that block's field too
    assert any("CPU 부하율 — Load는 70% 이하." in str(st) for tc in tcs for st in tc["steps"])


def test_a_criterion_review_item_points_to_no_other_sentence():
    """(review r3 Info-f) a pass criterion is measured — moving its value into another sentence is no fix."""
    from generators.sts_requirement_tc import _evidence_text
    assert _evidence_text({"kind": "held_back", "fill_check": "not_applicable", "evidence": []}).startswith(
        "— (측정 요구 · 판정 기준 항목")


def test_a_repeated_criterion_is_headed_by_its_own_place():
    """(review r4 W-A) the repeat was counted by its text and found by a spacing-free or partial match — ``사용률  70%``
    (two spaces) then ``사용률 70%`` were both headed ``RAM``, and a longer row holding the sentence was taken for it."""
    from generators.sts_requirement_tc import boundary_steps
    for verification in ("1) RAM 사용률\n- 사용률  70% 이하.\n2) ROM 사용률\n- 사용률 70% 이하.",
                         "1) A 부하\n- 피크 - 사용률 70% 이하.\n2) B 부하\n- 사용률 70% 이하.\n3) C 부하\n- 사용률 70% 이하."):
        groups = boundary_steps({"id": "SwNTR_0301", "name": "n", "description": "", "asil": "", "related_id": "",
                                 "verification": verification})
        heads = [st["action"].split(": ", 1)[1].split(" — ")[0] for g in groups for st in g["steps"]]
        want = [r.split(") ", 1)[1] for r in verification.split("\n") if r[:1].isdigit()]
        assert heads == want, (verification, heads)


def test_a_criterion_review_item_is_not_searched():
    """(review r4 W-B) a criterion with a unit was searched and counted ('1 of 1 found') while its cell said '탐색 대상
    아님' — the board and the sheet told two stories."""
    from generators.sts_requirement_tc import attach_review_evidence, boundary_steps
    req = {"id": "SwNTR_0301", "name": "n", "description": "부하는 80% 이하로 유지한다.", "asil": "", "related_id": "",
           "verification": "- CPU Load 60 ~ 80%."}
    review: list = []
    boundary_steps(req, None, None, review)
    assert [r["fill_check"] for r in review] == ["not_applicable"]
    stats = attach_review_evidence(review, [req], None)
    assert stats["review_evidence_searched"] == 0 and review[0]["evidence"] is None
    assert stats["review_fill_checks"] == {"not_applicable": 1}


def test_a_cited_blocks_repeated_criterion_is_headed_by_its_own_place():
    """(review r5 I-4) the traced path too: one sentence under two headings of a cited block's criteria."""
    from generators.sts_requirement_tc import append_requirement_boundary_tcs
    from tests.unit.test_sts_requirement_tc_r29 import _build, _classify_steps, _make_tc_id, _req
    system = {"SyNT_01": {"doc": "SyDS", "fields": {
        "Verification criteria": "1) RAM 사용률\n- 사용률 70% 이하.\n2) ROM 사용률\n- 사용률 70% 이하."}}}
    tcs: list = []
    append_requirement_boundary_tcs(tcs, [_req(desc="", rid="SwNTR_0301", related="SyNT_01")], _build, _make_tc_id,
                                    _classify_steps, system=system)
    said = " ".join(str(st) for tc in tcs for st in tc["steps"])
    assert "RAM 사용률 — 사용률 70% 이하." in said and "ROM 사용률 — 사용률 70% 이하." in said


def test_only_criteria_held_back_is_no_unit_less_search():
    """(review r5 I-2) with only criteria held back the board said "no item with a unit" — they have one (%)."""
    from report_gen.generation_disclosures import build_disclosures
    items = build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
        "tcs": 1, "steps": 1, "facts": 1, "facts_used": 1, "review_item_count": 1,
        "review_by_reason": {"kind_range": 1}, "review_evidence_searched": 0, "review_evidence_found": 0,
        "review_fill_checks": {"not_applicable": 1}}}})
    text = " ".join(str(i.get("note")) for i in items)
    assert "판정 기준 항목은 탐색 대상 아님" in text and "단위 있는 값의 보류 항목이 없어" not in text

