# tests/unit/test_docx_grid_snapshot.py
"""(R47-j N27-b) 표 셀 접근은 **표당 그리드 1회** — `table.cell()` 을 셀마다 부르면 셀 수의 제곱이다.

## 왜 이 파일이 있나

`_fill_function_info_table` 은 이 함정을 이미 고쳐 두었다(`test_docx_function_info_fill.py`). 그런데 같은 파일의 형제 셋
(`_add_blank_table`·`_merge_function_info_table`·`_insert_logic_image_in_table`)은 그대로 `table.cell()` 을 셀마다 불렀고,
라이브 프로파일(kjpds02_pv · 1,157함수 · DOCX 단계 1,486초 단독)에서 그 API 아래가 DOCX 시간의 **55%** 였다
(`_add_blank_table` 38.6% · `_merge_function_info_table` 12.7%).

## 무엇을 고정하나

1. **등가성** — 스냅샷 `_grid_cells(table)[c + r*cols]` 는 그 시점의 `table.cell(r, c)` 와 같은 `<w:tc>` 다(균일·가로 병합·세로 병합).
2. **결과 동일** — 세 함수의 새 구현은 `table.cell()` 을 쓰던 참조 구현과 **XML 바이트가 같다**(참조 구현은 여기 그대로 둔다).
3. **접근 경로** — 세 함수는 `Table.cell` 을 0회, `Table._cells` 를 표당 1회만 밟는다(등가성만 재면 되돌려도 통과한다 —
   R47-h 의 교훈: 결과가 같은 두 구현은 관측량이 **횟수**여야 갈린다).
"""
from __future__ import annotations

import copy

import pytest

from tests.unit._source_probe import source_of

docx = pytest.importorskip("docx", reason="python-docx 없음")

from docx.table import Table  # noqa: E402

from report_gen.docx_builder import (  # noqa: E402
    FN_ROW_FULL,
    FN_ROW_GRID,
    FN_ROW_PAIR,
    PARAM_GRID_COLS,
    _add_blank_table,
    _grid_cells,
    _insert_logic_image_in_table,
    _merge_function_info_table,
)

_PNG_1x1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da6364f8cfc000000201"
    "0100c9fe92ef0000000049454e44ae426082"
)


def _doc():
    return docx.Document()


def _layout(n_pairs: int, grid_rows: int = 0):
    lay = [(FN_ROW_FULL, ["[ Function Information ]"])]
    lay += [(FN_ROW_PAIR, [f"L{i}", "", f"V{i}"]) for i in range(n_pairs)]
    if grid_rows:
        lay.append((FN_ROW_FULL, ["[ Input Parameters ]"]))
        lay += [(FN_ROW_GRID, [str(i), f"p{i}", "U8", "0~255", "0", "d"]) for i in range(grid_rows)]
    return lay


# ── 참조 구현(변경 전 코드 그대로 — `table.cell()` 사용) ─────────────────────

def _ref_add_blank_table(doc, rows, cols, style=None, header_rows=None, data_rows=None):
    if rows <= 0 or cols <= 0:
        return None
    table = doc.add_table(rows=rows, cols=cols)
    try:
        if style:
            table.style = style
    except Exception:  # silent-ok: 변경 전 코드 그대로(참조 구현)
        pass
    row_offset = 0
    if header_rows:
        for r_idx, row in enumerate(header_rows):
            if r_idx >= rows:
                break
            for c_idx, val in enumerate(row[:cols]):
                table.cell(r_idx, c_idx).text = val or ""
        row_offset = min(len(header_rows), rows)
    if data_rows:
        max_rows = rows - row_offset
        for r_idx, row in enumerate(data_rows[:max_rows]):
            for c_idx, val in enumerate(row[:cols]):
                table.cell(row_offset + r_idx, c_idx).text = str(val) if val is not None else ""
    for r_idx in range(row_offset, rows):
        for c in table.rows[r_idx].cells:
            if c.text is None:
                c.text = ""
    return table


def _ref_merge(table, cols, layout=None):
    if not table or cols < 4:
        return
    try:                                        # (리뷰 W2) 원본에도 있던 try/except — 이게 있어야 ragged 표를 넣을 수 있다
        kinds = [str(k) for k, _ in (layout or [])]
        for r_idx in range(len(table.rows)):
            kind = kinds[r_idx] if r_idx < len(kinds) else (FN_ROW_FULL if r_idx == 0 else FN_ROW_PAIR)
            if kind == FN_ROW_GRID:
                if cols > PARAM_GRID_COLS:
                    table.cell(r_idx, PARAM_GRID_COLS - 1).merge(table.cell(r_idx, cols - 1))
                continue
            if kind == FN_ROW_FULL:
                table.cell(r_idx, 0).merge(table.cell(r_idx, cols - 1))
                continue
            if cols >= 2:
                table.cell(r_idx, 0).merge(table.cell(r_idx, 1))
            if cols >= 4:
                table.cell(r_idx, 2).merge(table.cell(r_idx, cols - 1))
    except Exception:  # silent-ok: 변경 전 코드 그대로(참조 구현)
        pass


_RAGGED_WIDTHS = [1, 2, 2, 2, 2, 2, 2, 1, 6, 6]     # 정본/템플릿 실측 행 폭(tblGrid=6) — 리뷰 W1 재현 모양


def _ragged_table(doc):
    t = doc.add_table(rows=len(_RAGGED_WIDTHS), cols=6)
    for tr, w in zip(t._tbl.tr_lst, _RAGGED_WIDTHS, strict=True):
        for tc in tr.tc_lst[w:]:
            tr.remove(tc)
    return t


# ── 1. 등가성 ─────────────────────────────────────────────────────────────────

class TestSnapshotEqualsTableCell:
    @staticmethod
    def _assert_same(table):
        rows, cols = len(table.rows), len(table.columns)
        grid = _grid_cells(table)
        assert len(grid) == rows * cols
        for r in range(rows):
            for c in range(cols):
                assert grid[c + r * cols]._tc is table.cell(r, c)._tc, (r, c)

    def test_uniform_table(self):
        self._assert_same(_doc().add_table(rows=5, cols=6))

    @pytest.mark.parametrize("cols", [6, 7])
    def test_horizontally_merged_function_info_table(self, cols):
        t = _doc().add_table(rows=8, cols=cols)
        _ref_merge(t, cols, _layout(3, grid_rows=3))
        self._assert_same(t)

    def test_vertically_merged_table(self):
        """vMerge 이어짐 칸은 `_cells` 가 **위 셀을 다시** 넣는다 — 스냅샷도 같은 객체를 낸다."""
        t = _doc().add_table(rows=4, cols=3)
        t.cell(0, 1).merge(t.cell(2, 1))
        assert t._tbl.xpath(".//w:vMerge"), "세로 병합이 안 만들어졌다 — 테스트 전제 확인"
        self._assert_same(t)


# ── 2. 결과 동일(참조 구현 대조) ─────────────────────────────────────────────

class TestXmlMatchesReferenceImplementation:
    @pytest.mark.parametrize("rows, hdr, data", [
        (4, [["A", "B", "C"]], [["1", None, "3"], ["x", "y", "z"]]),          # None → "" · 빈 꼬리 행
        (2, [["A", "B", "C"]], [["1", "2", "3"], ["4", "5", "6"], ["7", "8", "9"]]),  # 잘림
        (3, None, [["only", "data"]]),                                           # 머리 없음 · 짧은 행
        (30, [["h1", "h2", "h3"]], [[str(i), str(i * 2), None] for i in range(29)]),
    ], ids=["none-and-blank-rows", "truncated", "no-header", "30rows"])
    def test_add_blank_table(self, rows, hdr, data):
        new = _add_blank_table(_doc(), rows, 3, None, hdr, data)
        ref = _ref_add_blank_table(_doc(), rows, 3, None, hdr, data)
        assert new._tbl.xml == ref._tbl.xml

    def test_add_blank_table_zero_rows_is_none(self):
        assert _add_blank_table(_doc(), 0, 3) is None and _add_blank_table(_doc(), 3, 0) is None

    @pytest.mark.parametrize("cols", [4, 6, 7])
    @pytest.mark.parametrize("layout", [None, _layout(4), _layout(2, grid_rows=3)], ids=["uniform", "pairs", "pairs+grid"])
    def test_merge_function_info_table(self, cols, layout):
        n = len(layout) if layout else 6
        new = _doc().add_table(rows=n, cols=cols)
        ref = _doc().add_table(rows=n, cols=cols)
        _merge_function_info_table(new, cols, layout)
        _ref_merge(ref, cols, layout)
        assert new._tbl.xml == ref._tbl.xml

    def test_merge_on_a_ragged_table_matches_reference(self):
        """(리뷰 W1) 행 폭이 tblGrid 와 다른 표 — 스냅샷을 그대로 쓰면 3·4행의 라벨|값 2칸이 1칸으로 접혔다. 옛 경로로 갈라탄다."""
        new, ref = _ragged_table(_doc()), _ragged_table(_doc())
        _merge_function_info_table(new, 6, None)
        _ref_merge(ref, 6, None)
        assert new._tbl.xml == ref._tbl.xml
        assert [len(tr.tc_lst) for tr in new._tbl.tr_lst] == [len(tr.tc_lst) for tr in ref._tbl.tr_lst]

    def test_merge_is_idempotent_on_an_already_merged_table(self):
        """정규화 경로는 문서에서 읽은(이미 병합된) 표에 다시 돈다 — 스냅샷도 그 자리에서 무해해야 한다."""
        t = _doc().add_table(rows=5, cols=6)
        _ref_merge(t, 6, _layout(4))
        before = t._tbl.xml
        _merge_function_info_table(t, 6, _layout(4))
        assert t._tbl.xml == before

    def test_insert_logic_image_targets_the_same_cell(self, tmp_path):
        img = tmp_path / "logic.png"
        img.write_bytes(_PNG_1x1)
        a, b = _doc(), _doc()
        ta = _add_blank_table(a, 4, 6, None, None, [["x"] * 6, ["Logic Diagram", "", "", "", "", ""], ["y"] * 6, ["z"] * 6])
        tb = _ref_add_blank_table(b, 4, 6, None, None, [["x"] * 6, ["Logic Diagram", "", "", "", "", ""], ["y"] * 6, ["z"] * 6])
        assert _insert_logic_image_in_table(ta, 6, str(img)) is True
        # 참조: 변경 전 코드는 `table.cell(r_idx, min(2, cols-1))` 에 넣었다 — 같은 칸(1행 2열)에 그림이 있어야 한다.
        assert tb.cell(1, 2).text == ""            # 전제: 그 칸은 비어 있었다
        drawings = [len(c._tc.xpath(".//w:drawing")) for r in ta.rows for c in r.cells]
        assert sum(drawings) == 1
        assert ta.cell(1, 2)._tc.xpath(".//w:drawing"), "그림이 1행 2열이 아닌 칸에 들어갔다"


# ── 3. 접근 경로 — 횟수가 관측량이다 ────────────────────────────────────────

class TestAccessPathContract:
    @pytest.fixture
    def counters(self, monkeypatch):
        calls = {"cell": 0, "_cells": 0}
        real_cell = Table.cell
        real_cells = Table._cells.fget

        def _cell(self, r, c):
            calls["cell"] += 1
            return real_cell(self, r, c)

        def _cells(self):
            calls["_cells"] += 1
            return real_cells(self)

        monkeypatch.setattr(Table, "cell", _cell)
        monkeypatch.setattr(Table, "_cells", property(_cells))
        return calls

    @pytest.mark.parametrize("n", [10, 300], ids=lambda n: f"{n}행")
    def test_add_blank_table_builds_the_grid_once(self, counters, n):
        _add_blank_table(_doc(), n + 1, 6, None, [["h"] * 6], [[str(i)] * 6 for i in range(n)])
        assert counters == {"cell": 0, "_cells": 1}, counters

    @pytest.mark.parametrize("n", [5, 40], ids=lambda n: f"{n}행")
    def test_merge_builds_the_grid_once(self, counters, n):
        t = _doc().add_table(rows=n, cols=6)
        counters["cell"] = counters["_cells"] = 0
        _merge_function_info_table(t, 6, _layout(n - 1))
        assert counters == {"cell": 0, "_cells": 1}, counters

    def test_insert_logic_image_builds_the_grid_once(self, counters, tmp_path):
        img = tmp_path / "logic.png"
        img.write_bytes(_PNG_1x1)
        t = _doc().add_table(rows=6, cols=6)
        t.cell(3, 0).text = "Logic Diagram"
        counters["cell"] = counters["_cells"] = 0
        assert _insert_logic_image_in_table(t, 6, str(img)) is True
        assert counters == {"cell": 0, "_cells": 1}, counters

    def test_ragged_table_takes_the_per_call_path(self, counters, caplog):
        """비직사각 표는 스냅샷을 만들지 않고(0회) 옛 경로(`Table.cell`)로 돈다 — 그리고 그 사실을 로그로 말한다."""
        import logging
        t = _ragged_table(_doc())
        counters["cell"] = counters["_cells"] = 0
        with caplog.at_level(logging.WARNING):
            _merge_function_info_table(t, 6, None)
        # 옛 경로 = `Table.cell` 호출마다 `_cells` 재구성(둘이 같은 수) — 스냅샷 경로는 `_cells` 1회·`cell` 0회다.
        assert counters["cell"] > 0 and counters["_cells"] == counters["cell"], counters
        assert any("행 폭이 tblGrid" in r.getMessage() for r in caplog.records)

    @pytest.mark.parametrize("fn", [_add_blank_table, _insert_logic_image_in_table], ids=lambda f: f.__name__)
    def test_source_does_not_call_table_cell(self, fn):
        body = source_of(fn).split('"""', 2)[-1]
        code = "\n".join(ln for ln in body.splitlines() if not ln.strip().startswith("#"))   # 주석의 설명은 호출이 아니다
        assert "table.cell(" not in code, f"{fn.__name__} 이 table.cell() 로 되돌아갔다"

    def test_reference_doc_is_released_right_after_extraction(self):
        """정본(실측 50MB docx) DOM 을 `ref_map` 을 뽑은 뒤에도 함수 끝까지 붙들면 템플릿 사본과 두 벌이 상주한다."""
        from report_gen import docx_builder
        src = source_of(docx_builder.generate_uds_docx)
        i = src.index("ref_map = _extract_function_info_from_docx(ref_doc)")
        window = src[i:i + 700]
        # (리뷰 I3) 추출이 던져도 놓아야 한다 — finally 안의 del.
        assert "finally:" in window and "del ref_doc" in window.split("finally:", 1)[1], "ref_doc 를 finally 에서 놓지 않는다"


def test_copy_of_table_keeps_reference_semantics_of_the_probe():
    """(자기 점검) `copy.deepcopy` 한 표를 참조 구현에 쓰지 않는다 — 두 구현은 **각자의 새 표**에 돈다."""
    t = _doc().add_table(rows=2, cols=2)
    assert copy.deepcopy(t._tbl) is not t._tbl
