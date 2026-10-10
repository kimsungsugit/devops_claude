"""R49 — the HW measurement tolerance next to the boundary steps it concerns.

A boundary step's spacing is the precision the requirement writes (``8.50V`` → 0.01 V). A HW requirements specification
states how precisely the monitor path measures what the SW compares — HDPDM01 ``HwTSR_0204`` Battery Voltage Monitor
``허용 오차: ±3%`` (8.50 V → ±0.255 V). With the spacing inside the tolerance, a HIL run cannot tell the boundary's
inclusion. The tolerance is linked only through the Related IDs (a HW block citing a system ID the SRS requirement
cites), quoted — never used to move a point — and said in the TC precondition, the 'Requirement Evidence' sheet, the
report and the disclosure; without the document, the disclosure says what giving it would add.

Review round 1: a hub block (``SySM_04``, cited by every power requirement) links several HW blocks — the first by ID
was the wrong monitor path for 21 of 48 notes (C2), so more than one block means "path undecided" and the points
outside use the largest candidate; an absolute tolerance written for a 2.5 V monitor node is not one of a 4.85 V
threshold (W1 — "척도 불명", no numbers); a HwRS that could not be localised is the HW record's reason (W2).
"""
from __future__ import annotations

import docx
import openpyxl
import pytest

from generators.sts import _build_tc_dict, _classify_steps, _make_tc_id
from generators.sts_requirement_tc import (
    REQUIREMENT_EVIDENCE_HEADERS,
    REQUIREMENT_EVIDENCE_SHEET,
    append_requirement_boundary_tcs,
    boundary_steps,
    hw_tolerance_summary,
    hw_tolerances_for,
    load_hw_requirements,
    parse_hw_requirement_docx,
    write_requirement_evidence_sheet,
)

BATTERY = {"name": "Battery Voltage Monitor", "related": ["SySM_04", "SyTSR_0116"],
           "tolerances": [{"text": "허용 오차: ±3%", "value": "3", "unit": "%", "measured_units": ["V"]}],
           "ranges": [{"lo": "9", "hi": "16", "unit": "V"}], "nominals": []}
# the sensor supply (0~5 V) on its own scale
SUPPLY = {"name": "Sensor Supply Monitor", "related": ["SySM_04", "SyTSR_0114"],
          "tolerances": [{"text": "허용 오차: ±0.15V", "value": "0.15", "unit": "V", "measured_units": ["V"]}],
          "ranges": [{"lo": "0", "hi": "5", "unit": "V"}], "nominals": []}
# HDPDM01 HwTSR_0203 as written: "Monitor전압: 2.5V(±0.15V) … 허용 오차: ±0.15V" — a monitor node's scale
NODE = {"name": "Hall Sensor Block 전원 Monitor", "related": ["SySM_04", "SyTSR_0114"],
        "tolerances": [{"text": "허용 오차: ±0.15V", "value": "0.15", "unit": "V", "measured_units": []}],
        "ranges": [], "nominals": [{"value": "2.5", "unit": "V"}]}
CURRENT = {"name": "Motor Current Monitor", "related": ["SySM_03"],
           "tolerances": [{"text": "허용 오차: ± 1A", "value": "1", "unit": "A", "measured_units": ["A"]}],
           "ranges": [{"lo": "0", "hi": "20", "unit": "A"}], "nominals": []}


def _req(desc, related, rid="SwTSR_0104"):
    return {"id": rid, "name": "n", "description": desc, "verification": "", "asil": "B", "related_id": related}


def _evidence(desc, related="SySM_04, SyTSR_0116"):
    (g,) = [g for g in boundary_steps(_req(desc, related)) if g["evidence"]["kind"] == "threshold"]
    return g["evidence"]


def _docx(path, tables):
    d = docx.Document()
    for rows in tables:
        t = d.add_table(rows=len(rows), cols=2)
        for row, (k, v) in zip(t.rows, rows, strict=True):
            row.cells[0].text, row.cells[1].text = k, v
    d.save(path)
    return str(path)


def test_parsing_reads_the_monitor_accuracy_and_what_the_block_measures(tmp_path):
    path = _docx(tmp_path / "hrs.docx", [
        [("ID", "HwTSR_0204"), ("Name", "Battery Voltage Monitor"),
         ("Description", "Battery Voltage Range: 9 ~16V / 분해능: 16Bit / 허용 오차: ±3%"),
         ("Related ID", "SyTR_0703, SyTSR_0116, SySM_04")],
        [("ID", "HwTSR_0203"), ("Name", "Hall Sensor Block 전원 Monitor"),
         ("Description", "Hall Sensor Block 전원 Monitor전압: 2.5V(±0.15V) / 허용 오차: ±0.15V"), ("Related ID", "SySM_04")],
        # a value's own spec and a test's measuring error are not the monitor accuracy
        [("ID", "HwTR_0605"), ("Name", "Motor Encoder Interface"), ("Description", "High: 5(±0.15)V / Low: 0V(±0.15)"),
         ("Verification criteria", "- Pass: 입력 전원 전압(측정오차±0.1) 이내"), ("Related ID", "SyII_29")],
        [("ID", "SyTR_0101"), ("Description", "허용 오차: ±1V")],                    # not a HW block
    ])
    blocks = parse_hw_requirement_docx(path)
    assert set(blocks) == {"HwTSR_0204", "HwTSR_0203", "HwTR_0605"}
    # (R50) the full text and the block's own formula ratios are in the block too — compared in the R50 tests
    assert blocks["HwTSR_0204"]["scales"] == [] and "허용 오차: ±3%" in blocks["HwTSR_0204"]["full_text"]
    r49_keys = {k: v for k, v in blocks["HwTSR_0204"].items() if k not in ("full_text", "scales", "monitor_node")}
    assert r49_keys == {"name": "Battery Voltage Monitor", "related": ["SyTR_0703", "SyTSR_0116", "SySM_04"],
                                    "tolerances": [{"text": "허용 오차: ±3%", "value": "3", "unit": "%",
                                                    "measured_units": ["V"]}],
                                    "ranges": [{"lo": "9", "hi": "16", "unit": "V"}], "nominals": []}
    assert blocks["HwTSR_0203"]["nominals"] == [{"value": "2.5", "unit": "V"}] and blocks["HwTSR_0203"]["ranges"] == []
    assert blocks["HwTR_0605"]["tolerances"] == [] and blocks["HwTR_0605"]["nominals"] == [
        {"value": "5", "unit": "V"}, {"value": "0", "unit": "V"}]
    got, record = load_hw_requirements(path)
    assert got == blocks and record == {"file": "hrs.docx", "blocks": 3, "blocks_with_tolerance": 2}
    assert load_hw_requirements(None) == (None, None)
    none, bad = load_hw_requirements(str(tmp_path / "missing.docx"))
    assert none is None and bad["file"] == "missing.docx" and "error" in bad


def test_a_relative_tolerance_applies_to_what_the_block_measures():
    """HDPDM01 ``SwTSR_0104`` ``u16g_ApiIn_Vsup 8.50V 미만``: ±3% of 8.50 V is 0.255 V — the spacing 0.01 V is inside."""
    e = _evidence("- u16g_ApiIn_Vsup 가 8.50V 미만 시 저전압")
    (t,) = hw_tolerances_for(_req("", "SySM_04, SyTSR_0116"), e, {"HwTSR_0204": BATTERY, "HwTSR_0101": CURRENT})
    assert (t["hw_id"], t["text"], t["size"], t["unit"], t["shared"], t["inside"], t["scale_known"], t["direct"]) == \
        ("HwTSR_0204", "허용 오차: ±3%", "0.255", "V", ["SySM_04", "SyTSR_0116"], True, True, False)


def test_a_percent_is_not_applied_to_another_quantity():
    """(the probe's first reading) the Battery monitor's ±3% is not a hold time's: ``500ms`` gets nothing."""
    (g,) = [g for g in boundary_steps(_req("- 저전압 상태가 500ms 이상 유지 시 고장", "SySM_04"))]
    assert hw_tolerances_for(_req("", "SySM_04"), g["evidence"], {"HwTSR_0204": BATTERY}) == []


def test_absolute_tolerances_need_the_same_unit_and_scale():
    e = _evidence("- u16g_ApiIn_HallSnsrLevel 이 4.85V 이하 시 고장", "SySM_04, SyTSR_0114")
    (t,) = hw_tolerances_for(_req("", "SySM_04, SyTSR_0114"), e, {"HwX_01": SUPPLY})
    assert (t["size"], t["inside"], t["scale_known"]) == ("0.15", True, True)
    milli = {"HwX_01": {**SUPPLY, "tolerances": [{"text": "허용 오차: ±150mV", "value": "150", "unit": "mV",
                                                 "measured_units": []}]}}
    (t,) = hw_tolerances_for(_req("", "SySM_04"), e, milli)
    assert (t["size"], t["unit"]) == ("0.15", "V")
    # (review W1) HwTSR_0203 writes its ±0.15V for a 2.5 V monitor node: at 4.85 V the scale is unknown — no number
    (t,) = hw_tolerances_for(_req("", "SySM_04, SyTSR_0114"), e, {"HwTSR_0203": NODE})
    assert (t["scale_known"], t["size"], t["written_size"], t["inside"]) == (False, None, "0.15", None)
    # a nominal on the value's scale is the block's own scale (``5V(±0.15)`` for a 4.85 V threshold)
    (t,) = hw_tolerances_for(_req("", "SySM_04"), e, {"HwX_02": {**NODE, "nominals": [{"value": "5", "unit": "V"}]}})
    assert (t["scale_known"], t["size"]) == (True, "0.15")
    # an absolute tolerance in another unit (the motor current's ±1A) says nothing about a voltage
    assert hw_tolerances_for(_req("", "SySM_04"), e, {"HwTSR_0101": {**CURRENT, "related": ["SySM_04"]}}) == []
    # a spacing wider than the tolerance is not inside
    wide = _evidence("- u16g_ApiIn_Vsup 가 9V 미만 시 저전압")                    # written precision 1 V
    (t,) = hw_tolerances_for(_req("", "SySM_04, SyTSR_0116"), wide, {"HwTSR_0204": BATTERY})
    assert (t["size"], t["inside"]) == ("0.27", False)


def test_only_the_related_ids_link():
    e = _evidence("- u16g_ApiIn_Vsup 가 8.50V 미만 시 저전압")
    assert hw_tolerances_for(_req("", "SyTR_0101"), e, {"HwTSR_0204": BATTERY}) == []
    assert hw_tolerances_for(_req("", ""), e, {"HwTSR_0204": BATTERY}) == []


def test_a_hub_leaves_the_monitor_path_undecided_and_the_largest_candidate_bounds_the_points():
    """(review C2) ``SwTSR_0301`` ``입력전원 8.9V 미만`` cites only the hub ``SySM_04``: the Battery (±3% → ±0.267V) and the
    sensor supply (±0.15V, not on the 8.9 V scale here) both link. The first by ID was the supply — its "outside" points
    8.65 V · 9.15 V lay inside the Battery path's ±0.267 V. Now: undecided, and the points are outside every candidate."""
    e = _evidence("- 입력전원 8.9V 미만 시 저전압", "SySM_04")
    e["hw_tolerance"] = hw_tolerances_for(_req("", "SySM_04"), e, {"HwTSR_0203": SUPPLY, "HwTSR_0204": BATTERY})
    s = hw_tolerance_summary(e)
    assert (s["certain"], s["blocks"], s["largest"], s["inside"]) == (False, ["HwTSR_0203", "HwTSR_0204"], "0.267", True)
    # two candidates on known scales: the largest bounds the points, whichever comes first (4.85 V: 3% = 0.1455 < 0.15)
    e = _evidence("- u16g_ApiIn_HallSnsrLevel 이 4.85V 이하 시 고장", "SySM_04")
    e["hw_tolerance"] = hw_tolerances_for(_req("", "SySM_04"), e, {"HwTSR_0204": BATTERY, "HwTSR_0205": SUPPLY})
    assert [(t["hw_id"], t["size"]) for t in e["hw_tolerance"]] == [("HwTSR_0204", "0.1455"), ("HwTSR_0205", "0.15")]
    assert hw_tolerance_summary(e)["largest"] == "0.15"
    # (review r2 I1) the summary is of all candidates — the sheet keeps three, the largest may be the fourth
    many = [{**e["hw_tolerance"][0], "hw_id": f"HwX_0{i}", "size": "0.01", "inside": False} for i in range(3)]
    assert hw_tolerance_summary(e, many + e["hw_tolerance"])["largest"] == "0.15"


def test_the_system_block_of_the_fact_decides_the_path():
    """A traced fact of ``SyTSR_0116`` — only the Battery monitor cites that block: one path, however many share the hub."""
    system = {"SyTSR_0116": {"doc": "SyRS", "name": "n", "fields": {"Description": "- BAT 8.5V미만 시 고장"}}}
    (g,) = boundary_steps(_req("- 고장 판단", "SySM_04, SyTSR_0116"), None, system)
    g["evidence"]["hw_tolerance"] = hw_tolerances_for(_req("", "SySM_04, SyTSR_0116"), g["evidence"],
                                                      {"HwTSR_0203": SUPPLY, "HwTSR_0204": BATTERY})
    assert [(t["hw_id"], t["direct"]) for t in g["evidence"]["hw_tolerance"]] == [("HwTSR_0204", True),
                                                                                  ("HwTSR_0203", False)]
    s = hw_tolerance_summary(g["evidence"])
    assert (s["certain"], s["blocks"], s["largest"]) == (True, ["HwTSR_0204"], "0.255")


def _append(requirements, system=None, hw=None):
    tcs, out = [], []

    def build(**kw):
        return _build_tc_dict(test_env="HILS", derive_inputs=False, is_safety=False, **kw)
    stats = append_requirement_boundary_tcs(tcs, requirements, build, _make_tc_id, _classify_steps, system=system,
                                            review_out=out, hw=hw)
    return stats, tcs


def test_the_tc_the_sheet_and_the_report():
    hw = {"HwTSR_0204": BATTERY, "HwTSR_0203": NODE, "HwTSR_0101": CURRENT}
    reqs = [_req("- u16g_ApiIn_Vsup 가 8.50V 미만 시 저전압", "SySM_04, SyTSR_0116"),
            _req("- 차속이 3km/h 초과 시 정지", "SyTR_0101", rid="SwTR_0101")]
    stats, tcs = _append(reqs, hw=hw)
    battery = [tc for tc in tcs if tc["srs_id"] == "SwTSR_0104"][0]
    # the hub links the node monitor too: undecided; the largest known candidate bounds the points — a short line
    assert ("HW 측정 허용오차(감시 경로 미정 — 후보 HwTSR_0204 ±0.255V · HwTSR_0203 ±0.15V(척도 불명))가 경계 점 간격 "
            "0.01V 보다 큼 — HIL 에서는 경계 포함(이상/초과·이하/미만)을 가를 수 없다") in battery["precondition"]
    # outside the largest tolerance by one step, 8.50 ∓ (0.255 + 0.01), rounded outward to the written step (r2 I4)
    assert "가장 큰 후보의 허용오차 밖 8.23V · 8.77V 에서 방향만 확인" in battery["precondition"]
    # the node-scale candidate may be larger still: the points are outside the known ones only, and it says so
    assert "척도 불명 후보 HwTSR_0203 는 분압비 확인 전이라 이 밖 점에 넣지 않았다" in battery["precondition"]
    speed = [tc for tc in tcs if tc["srs_id"] == "SwTR_0101"][0]
    assert "HW 측정 허용오차" not in (speed.get("precondition") or "")
    assert (stats["hw_tolerance_groups"], stats["hw_tolerance_inside_step"], stats["hw_tolerance_path_undecided"],
            stats["hw_tolerance_scale_unknown"]) == (1, 1, 1, 0)
    wb = openpyxl.Workbook()
    write_requirement_evidence_sheet(wb, tcs)
    rows = list(wb[REQUIREMENT_EVIDENCE_SHEET].iter_rows(values_only=True))
    col = REQUIREMENT_EVIDENCE_HEADERS.index("HW Tolerance (Quoted)")
    assert rows[0][col] == "HW Tolerance (Quoted)"
    by_value = {r[6]: r[col] for r in rows[1:]}
    assert by_value["8.50"].startswith("감시 경로 미정 — HwTSR_0204 Battery Voltage Monitor: 허용 오차: ±3% → ±0.255V "
                                       "[한 눈금 0.01V 가 그 안] [공유 SySM_04, SyTSR_0116]")
    assert "HwTSR_0203 Hall Sensor Block 전원 Monitor: 허용 오차: ±0.15V → 척도 불명" in by_value["8.50"]
    assert by_value["3"] == "— (Related ID 로 이어진 HW 허용오차 없음)"
    # a spacing wider than the tolerance (``9V`` — one written step 1 V > ±0.27 V): linked and quoted, no HIL note
    wide_stats, wide = _append([_req("- u16g_ApiIn_Vsup 가 9V 미만 시 저전압", "SySM_04, SyTSR_0116", rid="SwTSR_0105")],
                               hw={"HwTSR_0204": BATTERY})
    assert (wide_stats["hw_tolerance_groups"], wide_stats["hw_tolerance_inside_step"]) == (1, 0)
    assert "HW 측정 허용오차" not in (wide[0].get("precondition") or "")
    assert wide[0]["requirement_evidence"][0]["hw_tolerance"][0]["inside"] is False


def test_one_path_one_number_and_an_unknown_scale_no_number():
    stats, tcs = _append([_req("- u16g_ApiIn_Vsup 가 8.50V 미만 시 저전압", "SyTSR_0116")], hw={"HwTSR_0204": BATTERY})
    note = tcs[0]["precondition"]
    assert "HW 측정 허용오차(HwTSR_0204 ±0.255V)가" in note and "허용오차 밖 8.23V · 8.77V" in note
    assert "가장 큰 후보" not in note and stats["hw_tolerance_path_undecided"] == 0
    # (review W1) only a node-scale tolerance: no points, the divider to confirm first
    stats, tcs = _append([_req("- u16g_ApiIn_HallSnsrLevel 이 4.85V 이하 시 고장", "SyTSR_0114")], hw={"HwTSR_0203": NODE})
    note = tcs[0]["precondition"]
    assert "HwTSR_0203 ±0.15V(척도 불명)" in note and "척도(분압비)를 HW 문서로 확인한 뒤" in note
    assert "V · " not in note.split("가를 수 없다")[1]                # no numeric outside points
    assert (stats["hw_tolerance_inside_step"], stats["hw_tolerance_scale_unknown"]) == (1, 1)
    # (review r2 W1) KJPDS02 Hall level 4.85 V: the node monitor writes ±0.3V (unknown scale) — more than the Battery's
    #   3% (0.1455 V) that has a number: points from the smaller would not be outside the larger — no number at all
    wide_node = {**NODE, "tolerances": [{**NODE["tolerances"][0], "text": "허용 오차: ±0.3V", "value": "0.3"}]}
    stats, tcs = _append([_req("- u16g_ApiIn_HallSnsrLevel 이 4.85V 이하 시 고장", "SySM_04")],
                         hw={"HwTSR_0203": wide_node, "HwTSR_0204": BATTERY})
    note = tcs[0]["precondition"]
    assert "감시 경로 미정" in note and "척도(분압비)를 HW 문서로 확인한 뒤" in note and "가장 큰 후보" not in note
    assert stats["hw_tolerance_scale_unknown"] == 1


def test_without_the_document_nothing_changes_and_the_report_says_so():
    reqs = [_req("- u16g_ApiIn_Vsup 가 8.50V 미만 시 저전압", "SySM_04, SyTSR_0116")]
    with_none, tcs = _append(reqs)
    assert "hw_tolerance_groups" not in with_none
    assert "HW 측정 허용오차" not in (tcs[0].get("precondition") or "")
    wb = openpyxl.Workbook()
    write_requirement_evidence_sheet(wb, tcs)
    col = REQUIREMENT_EVIDENCE_HEADERS.index("HW Tolerance (Quoted)")
    assert list(wb[REQUIREMENT_EVIDENCE_SHEET].iter_rows(values_only=True))[1][col] == \
        "— (HW 요구사항서 없음 — 미입력 또는 읽기 실패, 생성 공시 참조)"
    empty, _ = _append(reqs, hw={})
    assert (empty["hw_tolerance_groups"], empty["hw_tolerance_inside_step"]) == (0, 0)


def _disclosure(rb):
    from report_gen.generation_disclosures import build_disclosures
    items = {i["key"]: i for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": rb}})}
    return items.get("sts_hw_tolerance")


def test_the_disclosure_states():
    base = {"tcs": 1, "steps": 3, "facts": 1, "facts_used": 1}
    assert _disclosure(base) is None                       # an output from before R49: no item (absent ≠ not given)
    missing = _disclosure({**base, "hw_document": {"given": False}})    # not given: what giving it would add
    assert missing["value"] == "미입력" and "HW 요구사항서" in missing["note"] and "주면" in missing["note"]
    failed = _disclosure({**base, "hw_document": {"file": "hrs.docx", "error": "OSError: x"}})
    assert failed["value"] == "읽기 실패" and failed["tone"] == "warning"
    # (review W3) read, but no HW requirement table — another document registered
    empty = _disclosure({**base, "hw_document": {"file": "SyRS.docx", "blocks": 0, "blocks_with_tolerance": 0},
                         "hw_tolerance_groups": 0, "hw_tolerance_inside_step": 0})
    assert empty["value"] == "HW 요구 표 없음" and empty["tone"] == "warning"
    got = _disclosure({**base, "hw_document": {"file": "hrs.docx", "blocks": 47, "blocks_with_tolerance": 4},
                       "hw_tolerance_groups": 83, "hw_tolerance_inside_step": 76, "hw_tolerance_path_undecided": 20,
                       "hw_tolerance_scale_unknown": 5})
    assert got["value"] == "경계 사실 83 · 한 눈금이 허용오차 안 76" and got["tone"] == "warning"
    # (review I3) what to do comes first — the gate board keeps the head
    assert got["note"].startswith("한 눈금이 HW 측정 허용오차 안인 경계 76 개는 HIL 에서 경계 포함")
    assert "SW 변수 직접 주입" in got["note"][:230] and "감시 경로가 하나로 정해지지 않는다" in got["note"]
    assert "분압비" in got["note"] and "SySM_04" not in got["note"]           # (review I2) no project ID hard-coded
    none = _disclosure({**base, "hw_document": {"file": "hrs.docx", "blocks": 47, "blocks_with_tolerance": 0},
                        "hw_tolerance_groups": 0, "hw_tolerance_inside_step": 0})
    assert none["value"] == "0" and none["tone"] != "warning"


def test_the_gate_and_the_readiness_texts():
    from backend.services.docgen_requirements import IN_HWRS, INPUT_LABELS, requirements_for
    from backend.services.issue_explainer import rule_action
    assert INPUT_LABELS[IN_HWRS].startswith("HwRS")
    assert "허용오차" in requirements_for("sts")["optional"][IN_HWRS]
    assert "SW 변수" in rule_action("disclosure:sts_hw_tolerance")


def test_the_resolver_the_registry_and_the_middleware_carry_the_input(monkeypatch):
    from backend.middleware import _CLOUDIUM_PATH_KEYS
    from backend.schemas import ScmLinkedDocs
    from backend.services import resolver_helpers
    assert ScmLinkedDocs().hwrs == "" and "hwrs_path" in _CLOUDIUM_PATH_KEYS     # (review C1)
    monkeypatch.setattr(resolver_helpers, "resolve_builder_input", lambda p, label, reasons: p or None)
    got = resolver_helpers.resolve_system_requirement_docs("a.docx", "", "hw.docx")
    assert got == {"syrs_path": "a.docx", "syds_path": None, "hwrs_path": "hw.docx", "hwds_path": None,
                   "system_input_skips": []}
    assert resolver_helpers.resolve_system_requirement_docs("", "")["hwrs_path"] is None


@pytest.mark.parametrize("path", ["/api/jenkins/sts/generate-async", "/api/local/sts/generate",
                                  "/api/local/sts/generate-stream", "/api/local/sts/generate-async"])
def test_every_sts_handler_declares_the_hw_input(path):
    # FastAPI drops an undeclared form field silently — as the system inputs (R29)
    import inspect

    from backend.main import app
    route = next(r for r in app.routes if getattr(r, "path", "") == path and "POST" in getattr(r, "methods", set()))
    assert "hwrs_path" in inspect.signature(route.endpoint).parameters


def test_the_generator_records_what_it_was_given(tmp_path):
    """``hw_document``: not given · not localised (review W2) · unreadable · read."""
    from generators.sts import generate_sts
    text = ["SwTSR_0104: u16g_ApiIn_Vsup 가 8.50V 미만 시 저전압을 검출한다. Related ID: SySM_04"]
    none = generate_sts(text, {}, str(tmp_path / "a.xlsm"), project_config={"project_id": "T"})
    assert none["quality_report"]["generation_stats"]["requirement_boundary"]["hw_document"] == {"given": False}
    skipped = generate_sts(text, {}, str(tmp_path / "s.xlsm"), project_config={"project_id": "T"},
                           system_input_skips=["SyDS: 파일 없음 (a)", "HwRS: 파일 없음 (b)"])
    rb = skipped["quality_report"]["generation_stats"]["requirement_boundary"]
    assert rb["hw_document"] == {"file": "", "error": "HwRS: 파일 없음 (b)"}
    assert rb["system_input_skips"] == ["SyDS: 파일 없음 (a)"]           # the system documents' own reasons only
    bad = generate_sts(text, {}, str(tmp_path / "b.xlsm"), project_config={"project_id": "T"},
                       hwrs_path=str(tmp_path / "gone.docx"))
    assert "error" in bad["quality_report"]["generation_stats"]["requirement_boundary"]["hw_document"]
    hrs = _docx(tmp_path / "hrs.docx", [[("ID", "HwTSR_0204"), ("Name", "Battery Voltage Monitor"),
                                         ("Description", "Battery Voltage Range: 9 ~16V / 허용 오차: ±3%"),
                                         ("Related ID", "SySM_04")]])
    ok = generate_sts(text, {}, str(tmp_path / "c.xlsm"), project_config={"project_id": "T"}, hwrs_path=hrs)
    rb = ok["quality_report"]["generation_stats"]["requirement_boundary"]
    assert rb["hw_document"] == {"file": "hrs.docx", "blocks": 1, "blocks_with_tolerance": 1}
    assert (rb["hw_tolerance_groups"], rb["hw_tolerance_inside_step"]) == (1, 1)
