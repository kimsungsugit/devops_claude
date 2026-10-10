"""R75 — the SwUDS readers that open the document with python-docx release its graph before they return.

A python-docx `Document` is a reference cycle (the package and its parts point at each other): when the reader returns,
reference counting does not free it — it waits for the cyclic collector, whose full (second-generation) pass rarely runs in
a generation process holding millions of long-lived objects. Measured on KJPDS02_PV (SwUDS 48 MB): the design-ID bridge
(`report_gen.uds_related.docx_tables_text`) and the description reader (`generators.sts._load_uds_descriptions`) each left
1.2 GB behind (70 MB after a collection), and they piled up: the SUTS generation peaked at 4,079 MB RSS — the reason the
KJPDS02_PV measurements were killed for memory on a 16 GB PC, and the web backend runs the same generation.

The tests switch the automatic collector off so that only the readers' own collection can free the graph, and compare
with the documents alive when the test body starts, after a collection (`_baseline` — a failed earlier test keeps its
traceback, and any document on it, until pytest clears it: an absolute 0 failed on its behalf, review W2 · W4).
"""
from __future__ import annotations

import gc
import io

import docx
import pytest

from generators import sts
from report_gen import uds_related


def _docx_bytes() -> bytes:
    d = docx.Document()
    d.add_heading("u8g_Foo", level=1)
    d.add_paragraph("Reads the foo.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Description"
    t.cell(1, 0).text, t.cell(1, 1).text = "u8g_Bar", "Reads the bar."
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _live_documents() -> int:
    from docx.parts.document import DocumentPart
    return sum(isinstance(o, DocumentPart) for o in gc.get_objects())


def _baseline() -> int:
    """Documents alive at the start of the test body — collected first: an earlier failed test's traceback (and the
    document on it) is released by pytest only as this test's call begins (review round 2 W4)."""
    gc.collect()
    return _live_documents()


@pytest.fixture
def data():
    was = gc.isenabled()
    out = _docx_bytes()
    gc.collect()        # the document the fixture built is a cycle too
    gc.disable()
    try:
        yield out
    finally:
        if was:
            gc.enable()


def test_the_table_reader_leaves_no_document_graph(data):
    base = _baseline()
    tables = uds_related.docx_tables_text(data)
    assert tables == [[["Name", "Description"], ["u8g_Bar", "Reads the bar."]]]
    assert _live_documents() == base


def test_a_failed_python_docx_read_leaves_no_graph_either(data, monkeypatch):
    def half_read(raw):
        # opened, then failed while holding it (a damaged part): the traceback holds this frame — the collection must
        # come after the handler let it go (review W1: a dropped document would not tell)
        doc = docx.Document(io.BytesIO(raw))
        assert doc.paragraphs is not None
        raise ValueError("damaged part")
    monkeypatch.setattr(uds_related, "_python_docx_tables", half_read)
    base = _baseline()
    tables = uds_related.docx_tables_text(data)
    assert tables == [[["Name", "Description"], ["u8g_Bar", "Reads the bar."]]]   # the document.xml fallback
    assert _live_documents() == base


def test_the_description_reader_leaves_no_document_graph(data, tmp_path):
    path = tmp_path / "uds.docx"
    path.write_bytes(data)
    base = _baseline()
    descriptions = sts._load_uds_descriptions(str(path))
    assert descriptions == {"u8g_foo": "Reads the foo.", "u8g_bar": "Reads the bar."}
    assert _live_documents() == base


def test_a_description_read_that_raises_leaves_no_graph(data, tmp_path, monkeypatch):
    # (review I1) python-docx raises on a table whose merges disagree (``row.cells``): the SUTS keeps going with the
    # reason only — the document the read held must not stay behind on the traceback until the next full collection
    path = tmp_path / "uds.docx"
    path.write_bytes(data)

    def raising(p):
        doc = docx.Document(str(p))
        assert doc.paragraphs is not None
        raise ValueError("no tc element at grid_offset=3")
    monkeypatch.setattr(sts, "_uds_descriptions", raising)
    base = _baseline()
    try:
        sts._load_uds_descriptions(str(path))
    except ValueError as exc:
        assert "grid_offset" in str(exc)                           # the reason still reaches the caller
        tb = exc.__traceback__
        while tb.tb_next:
            tb = tb.tb_next
        assert tb.tb_frame.f_code.co_name == "raising"            # with the frames it was raised from (review W3)
    else:
        pytest.fail("the read error must reach the caller")
    assert _live_documents() == base
