"""R90 — a reference STS used as the template keeps none of its own test-spec sheet.

The generator writes `3.SW Test Spec` (the KJPDS02 reference's name). The HDPDM01 reference names it
`3.SW Integration Test Spec`: with it as the template, the generated workbook still carried the reference's 59 test cases
beside ours (found reading the generated workbook back, 2026-10-16) — the defect SITS fixed in R88."""
from __future__ import annotations

import openpyxl

from generators.sts import _drop_template_spec_sheets, generate_sts_xlsm


def _template(path, names, active=None):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    for name in names:
        wb.create_sheet(name)["B7"] = "SwTC_OLD_REFERENCE_01"
    wb.create_sheet("Notes")["A1"] = "keep me"
    if active is not None:
        wb.active = active
    wb.save(path)
    return path


def _cells(ws):
    return {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}


def _tc():
    return [{"id": "SwTC_SwTR_0001_01", "srs_id": "SwTR_0001", "title": "T", "safety_related": "X",
             "test_env": "SwTE_01", "test_method": "RBT", "gen_method": "AOR", "fs_req": "", "description": "",
             "precondition": "", "steps": [{"action": "a", "expected": "e"}], "related_ids": ["SwTR_0001"]}]


def test_the_hdpdm01_reference_spec_sheet_is_dropped(tmp_path):
    tpl = _template(tmp_path / "hd.xlsm", ["3.SW Integration Test Spec"])
    out = tmp_path / "out.xlsx"
    dropped: list = []
    generate_sts_xlsm(str(tpl), _tc(), {"coverage": {}}, str(out), {"project_id": "T"}, dropped_sheets=dropped)
    wb = openpyxl.load_workbook(str(out))
    assert "3.SW Integration Test Spec" not in wb.sheetnames and dropped == ["3.SW Integration Test Spec"]
    assert "3.SW Test Spec" in wb.sheetnames and "Notes" in wb.sheetnames and "Cover" in wb.sheetnames
    assert not any("SwTC_OLD_REFERENCE_01" in _cells(wb[s]) for s in wb.sheetnames), "the reference's old test stayed"


def test_our_own_sheet_name_is_rewritten_not_dropped(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.create_sheet("3.SW Test Spec")["B7"] = "SwTC_OLD_REFERENCE_01"
    assert _drop_template_spec_sheets(wb) == []
    tpl = tmp_path / "pv.xlsm"
    wb.save(tpl)
    out = tmp_path / "out.xlsx"
    generate_sts_xlsm(str(tpl), _tc(), {"coverage": {}}, str(out), {"project_id": "T"})
    assert "SwTC_OLD_REFERENCE_01" not in _cells(openpyxl.load_workbook(str(out))["3.SW Test Spec"])


def test_a_spelling_of_our_own_name_is_dropped_and_rewritten(tmp_path):
    # review I1: the writer replaces only the exact name — `3. SW Test Spec` kept the reference's tests, and
    #   `3.SW test spec` made openpyxl name ours `3.SW Test Spec1`
    for variant in ("3. SW Test Spec", "3.SW test spec"):
        tpl = _template(tmp_path / f"v{len(variant)}.xlsm", [variant])
        out = tmp_path / f"out{len(variant)}.xlsx"
        dropped: list = []
        generate_sts_xlsm(str(tpl), _tc(), {"coverage": {}}, str(out), {"project_id": "T"}, dropped_sheets=dropped)
        wb = openpyxl.load_workbook(str(out))
        assert dropped == [variant] and "3.SW Test Spec" in wb.sheetnames and "3.SW Test Spec1" not in wb.sheetnames
        assert not any("SwTC_OLD_REFERENCE_01" in _cells(wb[s]) for s in wb.sheetnames), variant


def test_only_spec_sheets_match_and_a_dropped_active_sheet_does_not_stay_active():
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    for name in ("1.Introduction", "2.Test Environment", "3.SW Integration Test Spec", "5. Traceability(SwRS)",
                 "Requirement Evidence", "Requirement Review"):
        wb.create_sheet(name)
    wb.active = 3
    assert _drop_template_spec_sheets(wb) == ["3.SW Integration Test Spec"]
    assert wb.active.title == "Cover" and "5. Traceability(SwRS)" in wb.sheetnames


def test_the_drop_reaches_the_disclosure():
    from report_gen.generation_disclosures import build_disclosures

    item = {i["key"]: i for i in build_disclosures(
        "sts", {"generation_stats": {"template_sheets_dropped": ["3.SW Integration Test Spec"]}})}.get(
        "sts_template_sheets_dropped")
    assert item is not None and item["value"] == "1개" and "3.SW Integration Test Spec" in item["note"]
