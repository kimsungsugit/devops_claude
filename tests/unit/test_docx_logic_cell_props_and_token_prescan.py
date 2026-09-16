# tests/unit/test_docx_logic_cell_props_and_token_prescan.py
"""(R53) 로직 이미지 칸은 **내용만** 비운다(N48) · 토큰 치환은 원문 노드 선검사 뒤에만 문단을 읽는다(N27-c).

## N48 — 왜 이 파일이 있나

`_insert_logic_image_in_table._clear_cell` 이 `tc` 의 자식을 전부 지워 `w:tcPr`(gridSpan·tcW)까지 없앴다. 그 칸은
`_merge_function_info_table` 이 `cols-2` 열로 병합한 값 칸이라 gridSpan 이 사라지면 그 행만 좁아지고(ragged), 정규화
(`_normalize_function_info_tables`)는 그 표를 비직사각으로 판정해 표마다 `table.cell()` 제곱 경로 + IndexError 경고를 냈다 —
라이브 실측 989/989 표(Initial commit 이래, 경고 1,978건/run). 이제 python-docx `_Cell.text` setter 와 같은 규약(`clear_content`)
으로 비우므로 행 폭이 유지되고 정규화는 조용한 no-op 이다.

## N27-c — 토큰 치환 선검사

`_replace_docx_text` 는 문단마다 `paragraph.text`(run 마다 xpath)를 읽어 정본(셀 17만 개)에서 21초였다. 키의 첫 글자가 문단의
어느 `w:t` 에도 없으면 치환할 게 없으므로 그 문단은 읽지 않는다 — 결과는 참조 구현(변경 전 코드)과 같고, 읽는 문단 수만 준다.
탭/개행/하이픈은 요소(`w:tab` 등)로도 텍스트에 들어오므로 그 글자로 시작하는 키가 있으면 선검사를 끈다.
"""
from __future__ import annotations

import logging

import pytest

docx = pytest.importorskip("docx", reason="python-docx 없음")

from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from report_gen.docx_builder import (  # noqa: E402
    FN_ROW_FULL,
    FN_ROW_PAIR,
    _add_blank_table,
    _fill_function_info_table,
    _insert_logic_image_in_table,
    _is_uniform_grid,
    _merge_function_info_table,
    _normalize_function_info_tables,
    _replace_docx_text,
)

_PNG_1x1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da6364f8cfc000000201"
    "0100c9fe92ef0000000049454e44ae426082"
)
_LAYOUT = [
    (FN_ROW_FULL, ["[ Function Information ]"]),
    (FN_ROW_PAIR, ["ID", "", "SwUFn_0101"]),
    (FN_ROW_PAIR, ["Logic Diagram", "", ""]),
    (FN_ROW_PAIR, ["Called Function", "", "N/A"]),
    (FN_ROW_PAIR, ["", "", ""]),           # R54 N50 이전 라이브 516/989 표에 남던 빈 꼬리 행 — 되읽기·삽입 경로는 여전히 견뎌야 한다
]


def _doc():
    return docx.Document()


def _function_table(doc):
    t = _add_blank_table(doc, len(_LAYOUT), 6, None, None, None)
    _merge_function_info_table(t, 6, _LAYOUT)
    _fill_function_info_table(t, _LAYOUT)
    return t


@pytest.fixture
def png(tmp_path):
    p = tmp_path / "logic.png"
    p.write_bytes(_PNG_1x1)
    return str(p)


# ── N48 ─────────────────────────────────────────────────────────────────────

class TestLogicImageKeepsCellProperties:
    def test_value_cell_keeps_tcpr_and_the_row_stays_full_width(self, png):
        d = _doc()
        t = _function_table(d)
        row = t._tbl.tr_lst[2]
        value_tc = row.tc_lst[1]
        tcpr_before = value_tc.tcPr.xml
        assert value_tc.grid_span == 4                         # 전제: 값 칸은 cols 2..5 로 병합돼 있다
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert value_tc.tcPr is not None and value_tc.tcPr.xml == tcpr_before, "값 칸의 tcPr(gridSpan·tcW)가 지워졌다"
        assert [tc.grid_span for tc in row.tc_lst] == [2, 4]
        assert _is_uniform_grid(t, 6) is True
        assert len(value_tc.p_lst) == 1 and len(value_tc.xpath(".//w:drawing")) == 1

    def test_normalization_is_a_quiet_no_op_after_insertion(self, png, caplog):
        d = _doc()
        t = _function_table(d)
        assert _insert_logic_image_in_table(t, 6, png) is True
        before = t._tbl.xml
        with caplog.at_level(logging.WARNING):
            _normalize_function_info_tables(d)
        assert t._tbl.xml == before
        noisy = [r.getMessage() for r in caplog.records if "행 폭이 tblGrid" in r.getMessage() or "병합 중단" in r.getMessage()]
        assert noisy == [], noisy

    def test_negative_control_stripping_tcpr_makes_the_table_ragged_and_noisy(self, png, caplog):
        """(자기 점검) 예전 동작을 손으로 재현하면 위 두 가드가 잡던 증상이 그대로 난다 — 가드가 빈 단언이 아니라는 뜻."""
        d = _doc()
        t = _function_table(d)
        assert _insert_logic_image_in_table(t, 6, png) is True
        value_tc = t._tbl.tr_lst[2].tc_lst[1]
        value_tc.remove(value_tc.tcPr)                          # 예전 `_clear_cell` 이 남기던 모양
        assert _is_uniform_grid(t, 6) is False
        with caplog.at_level(logging.WARNING):
            _normalize_function_info_tables(d)
        assert any("행 폭이 tblGrid" in r.getMessage() for r in caplog.records)


# ── N27-c 병합 칸 중복 읽기/쓰기 제거 ───────────────────────────────────────

class TestMergedCellsAreVisitedOnce:
    """`row.cells` 는 병합 칸을 gridSpan 만큼 반복한다 — 읽기(any/집합/첫 칸)와 멱등 쓰기는 칸당 한 번이면 결과가 같다."""

    @pytest.fixture
    def text_calls(self, monkeypatch):
        from docx.table import _Cell
        calls = {"get": 0, "set": 0}
        real_get, real_set = _Cell.text.fget, _Cell.text.fset

        def _get(self):
            calls["get"] += 1
            return real_get(self)

        def _set(self, value):
            calls["set"] += 1
            return real_set(self, value)
        monkeypatch.setattr(_Cell, "text", property(_get, _set))
        return calls

    def test_distinct_cells_folds_spans_and_keeps_first_seen_order(self):
        from report_gen.docx_builder import _distinct_cells
        d = _doc()
        t = _function_table(d)
        rows = list(t.rows)
        assert len(rows[0].cells) == 6 and len(_distinct_cells(rows[0])) == 1          # 머리글: 6칸 → 1
        assert [c._tc for c in _distinct_cells(rows[1])] == [rows[1].cells[0]._tc, rows[1].cells[2]._tc]   # 라벨|값
        u = d.add_table(rows=2, cols=3)
        u.cell(0, 1).merge(u.cell(1, 1))                                                # 세로 병합 — 아래 행은 위 칸의 tc
        assert [c._tc for c in _distinct_cells(list(u.rows)[1])] == [u.cell(1, 0)._tc, u.cell(0, 1)._tc, u.cell(1, 2)._tc]

    def test_logic_image_lookup_reads_each_cell_once(self, text_calls, png):
        d = _doc()
        t = _function_table(d)
        text_calls["get"] = 0
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert text_calls["get"] == 1 + 2 + 2, text_calls        # 머리글 1 · ID 행 2 · Logic Diagram 행 2 (찾으면 멈춘다)

    def test_layout_inference_reads_each_cell_once(self, text_calls):
        from report_gen.docx_builder import _infer_function_info_layout
        d = _doc()
        t = _function_table(d)
        text_calls["get"] = 0
        kinds = [k for k, _ in _infer_function_info_layout(t)]
        assert kinds == [k for k, _ in _LAYOUT]
        assert text_calls["get"] == 1 + 2 * 4, text_calls

    def test_normalization_writes_the_banner_once_per_merged_table(self, text_calls):
        d = _doc()
        _function_table(d)                                         # 빌더가 만든(이미 병합된) 표 — 머리글 칸은 하나
        text_calls["set"] = 0
        _normalize_function_info_tables(d)
        assert text_calls["set"] == 1, text_calls                  # 예전엔 같은 칸에 6회
        d2 = _doc()
        t2 = d2.add_table(rows=3, cols=6)
        t2.cell(0, 0).text = "Function Information"                # 아직 병합되지 않은 표 — 머리글 칸이 6개라 6회가 맞다(예전과 같음)
        text_calls["set"] = 0
        _normalize_function_info_tables(d2)
        assert text_calls["set"] == 6, text_calls
        # 머리글 쓰기 뒤 행 0 이 전폭 병합된다(내용 있는 칸의 병합은 python-docx 원판이 맡아 문단을 합친다 — 예전과 같은 결과)
        row0 = list(t2.rows)[0]
        assert len(row0.cells) == 6 and len(t2._tbl.tr_lst[0].tc_lst) == 1
        assert row0.cells[0].text.split("\n") == ["[ Function Information ]"] * 6

    def test_fill_writes_each_cell_once(self, text_calls):
        from report_gen.docx_builder import FN_ROW_GRID
        d = _doc()
        layout = [(FN_ROW_FULL, ["h"]), (FN_ROW_PAIR, ["L", "", "V"]), (FN_ROW_FULL, ["[ Input Parameters ]"]),
                  (FN_ROW_GRID, ["1", "p", "U8", "0~1", "0", "d"]), (FN_ROW_PAIR, ["L2", "", "V2"])]
        t = _add_blank_table(d, len(layout), 6, None, None, None)
        _merge_function_info_table(t, 6, layout)
        text_calls["set"] = 0
        _fill_function_info_table(t, layout)
        assert text_calls["set"] == 1 + 2 + 1 + 6 + 2, text_calls   # 예전: 12 + 8 + 12 + 12 + 8
        rows = list(t.rows)
        assert rows[1].cells[0].text == "L" and rows[1].cells[2].text == "V" and rows[3].cells[5].text == "d"


# ── N27-c 토큰 치환 선검사 ───────────────────────────────────────────────────

def _ref_replace_docx_text(doc, replacements):
    """변경 전 코드 그대로(참조 구현) — 선검사 없이 모든 문단을 읽는다."""
    def _replace_in_paragraph(paragraph):
        full = paragraph.text
        if not full:
            return
        changed = full
        for key, val in replacements.items():
            if key in changed:
                changed = changed.replace(key, val)
        if changed == full:
            return
        runs = paragraph.runs
        if not runs:
            paragraph.text = changed
            return
        joined = "".join(r.text for r in runs)
        if joined != full:
            paragraph.text = changed
            return
        for key, val in replacements.items():
            if key not in joined:
                continue
            cursor = 0
            for run in runs:
                rt = run.text
                start = cursor
                end = cursor + len(rt)
                cursor = end
                seg = joined[start:end]
                if key in seg:
                    run.text = seg.replace(key, val)
                    joined = "".join(r.text for r in runs)
                    break
            else:
                idx = joined.find(key)
                if idx < 0:
                    continue
                new_joined = joined[:idx] + val + joined[idx + len(key):]
                pos = 0
                for run in runs:
                    rlen = len(run.text)
                    run.text = new_joined[pos:pos + rlen] if pos + rlen <= len(new_joined) else new_joined[pos:]
                    pos += rlen
                if pos < len(new_joined):
                    runs[-1].text += new_joined[pos:]
                joined = "".join(r.text for r in runs)

    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph)
    for section in doc.sections:
        for hf in [section.header, section.footer]:
            if not hf or not hasattr(hf, "paragraphs"):
                continue
            for paragraph in hf.paragraphs:
                _replace_in_paragraph(paragraph)


def _doc_with_tokens():
    d = _doc()
    for i in range(40):
        d.add_paragraph(f"plain paragraph {i}")
    d.add_paragraph("Project: {{project_name}} v{{build_number}}")
    p = d.add_paragraph()
    p.add_run("split {{proj")
    p.add_run("ect_name}} end")                                  # 토큰이 run 경계에 걸침
    t = d.add_table(rows=2, cols=2)
    t.cell(1, 1).text = "cell {{MODULE_NAME}}"
    inner = t.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "nested {{project_name}}"           # 원판은 중첩 표를 보지 않는다 — 그대로 보지 않아야 한다
    d.sections[0].header.paragraphs[0].text = "hdr {{project_name}}"
    # (리뷰 W1) 요소에서 오는 세 글자를 **전부** 심는다 — 이 셋이 없으면 해제 규칙 케이스가 빈 단언이 된다.
    p2 = d.add_paragraph()
    p2.add_run().add_tab()
    p2.add_run("X")                                              # paragraph.text == "\\tX"   ← w:tab
    p3 = d.add_paragraph()
    p3.add_run().add_break()
    p3.add_run("X")                                              # paragraph.text == "\\nX"   ← w:br
    p4 = d.add_paragraph()
    r4 = p4.add_run()
    r4._r.append(OxmlElement("w:noBreakHyphen"))
    p4.add_run("X")                                              # paragraph.text == "-X"    ← w:noBreakHyphen
    return d


_ELEMENT_KEY_PARAGRAPH = {"\tX": -3, "\nX": -2, "-X": -1}         # 위 세 문단의 위치


_REP = {"{{project_name}}": "KJPDS02", "{{build_number}}": "76", "{{MODULE_NAME}}": "MOD"}


def _xml_pair(d):
    return d.element.xml, d.sections[0].header._element.xml


class TestTokenReplacementPrescan:
    def test_matches_the_reference_implementation(self):
        a, b = _doc_with_tokens(), _doc_with_tokens()
        _ref_replace_docx_text(a, _REP)
        _replace_docx_text(b, _REP)
        assert _xml_pair(a) == _xml_pair(b)
        assert b.paragraphs[40].text == "Project: KJPDS02 v76" and b.paragraphs[41].text == "split KJPDS02 end"
        assert b.tables[0].cell(1, 1).text == "cell MOD"
        assert b.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "nested {{project_name}}"
        assert b.sections[0].header.paragraphs[0].text == "hdr KJPDS02"

    @pytest.mark.parametrize("key", ["\tX", "\nX", "-X"])
    def test_keys_starting_with_element_borne_characters_disable_the_prescan(self, key):
        a, b = _doc_with_tokens(), _doc_with_tokens()
        idx = _ELEMENT_KEY_PARAGRAPH[key]
        assert b.paragraphs[idx].text == key                    # 전제: 그 글자는 `w:t` 가 아니라 요소에서 온다
        assert not any(key[0] in (t.text or "") for t in b.paragraphs[idx]._p.iter(qn("w:t")))
        rep = {key: "HIT"}
        _ref_replace_docx_text(a, rep)
        _replace_docx_text(b, rep)
        assert _xml_pair(a) == _xml_pair(b)
        assert b.paragraphs[idx].text == "HIT", f"{key!r} 문단의 키를 선검사가 걸러 버렸다(치환 안 됨)"

    def test_prescan_reads_only_paragraphs_that_can_contain_a_key(self, monkeypatch):
        calls = {"text": 0}
        real = Paragraph.text.fget

        def _text(self):
            calls["text"] += 1
            return real(self)
        monkeypatch.setattr(Paragraph, "text", property(_text, Paragraph.text.fset))
        a = _doc_with_tokens()
        _ref_replace_docx_text(a, _REP)
        ref_calls = calls["text"]
        calls["text"] = 0
        b = _doc_with_tokens()
        _replace_docx_text(b, _REP)
        # 원판은 40개의 평문 문단(+ 표·머리글의 빈 문단)까지 전부 읽는다 — 새 구현은 `{` 가 있는 문단만 읽는다
        assert ref_calls > 45, ref_calls
        assert calls["text"] <= 8, calls["text"]
        assert _xml_pair(a) == _xml_pair(b)

    def test_empty_replacements_is_a_no_op(self):
        a, b = _doc_with_tokens(), _doc_with_tokens()
        _ref_replace_docx_text(a, {})
        _replace_docx_text(b, {})
        assert _xml_pair(a) == _xml_pair(b)
