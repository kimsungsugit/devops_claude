"""R19 — inputs the source reads (`generators/suts.py`, extended SUTS only) and MC/DC fill rows.

The oracle refuses an expected value with ``initial_value_not_in_inputs:X`` when the function reads program object X
before writing it and the row gives X no value. Two causes: the unit's input list (design table, source analysis) left
X out, or an MC/DC design vector sets only what its decision reads. The extended profile (1) adds such objects as input
columns, typed from their declaration in the translation unit — the catalog *structure* still comes from the design
input list, so the reference rows keep their places and design-column values (R75), with the added inputs fixed there
and moved only by extended rows (OAT, boundary) — and (2) adds, next to each MC/DC vector with blank inputs, a filled
companion row that is no pair member. Nothing is added for locals, parameters, members, const/volatile/float objects
or unknown names. Every value below is what C computes on a 16-bit-int target (as the oracle models it)."""
from __future__ import annotations

import os

import pytest

from generators import boundary_rows as br
from generators import c_project_context as cpc
from generators import suts
from generators.suts import (
    MCDC_FILL_PREFIX,
    _source_object_decl,
    complete_source_read_inputs,
    extended_unit_sequences,
    generate_sequences,
    input_list_gaps,
    is_extended_strategy,
    resolve_seq_gen_method,
    scope_input_types,
    source_read_names,
    summarize_source_read_inputs,
)

ROOT = os.path.join(os.sep, "proj")
COMMON = """#ifndef COMMON_H
#define COMMON_H
typedef unsigned char U8;
typedef unsigned int U16;
typedef signed int S16;
typedef float F32;
typedef unsigned long long U64;
typedef enum { M_A = 0, M_B = 1, M_C = 5 } Mode_t;
#endif
"""
UNIT = """#include "common.h"
U8 g_a;
U8 g_b;
U8 g_gain_open;
U8 g_gain_close;
U8 g_o;
U8 g_buf[4];
U8 g_step;
const U8 g_k = 3U;
const U8 g_kc;
volatile U8 g_reg;
F32 g_f;
Mode_t g_m;
static U8 s_state;
static S16 s_level;
U64 g_big;
U16 g_raw;
void close_gain(void) { if (g_a > 10U) { g_o = g_gain_close; } else { g_o = 0U; } }
void chain(void) { if (g_a > 10U) { if (g_step > 2U) { g_o = s_state; } else { g_o = 1U; } } else { g_o = 0U; } }
void elem(void) { g_o = g_buf[2]; }
void mode_out(void) { if (g_m == M_C) { g_o = 1U; } else { g_o = 2U; } }
void inout(void) { if (g_a > 10U) { g_o = 5U; } }
void st(void) { if (g_a > 10U) { g_o = (U8)s_level; } else { g_o = 0U; } }
void big(void) { if (g_a > 10U) { g_o = (U8)g_big; } else { g_o = 0U; } }
void two(void) { if ((g_a > 10U) && (g_b > 3U)) { g_o = g_gain_close; } else { g_o = 0U; } }
void edge(void) { S16 t = s_level + 1; g_o = g_a; g_o += (U8)t; }
void dec(void) { if (g_a > 10U) { g_o = 1U; } else { g_o = 2U; } }
void raw(void) { if (g_a > 10U) { g_o = (U8)g_raw; } else { g_o = 0U; } }
"""


def _unit(name, input_vars, output_vars=("g_o",), **kw):
    files = {os.path.join(ROOT, "common.h"): COMMON, os.path.join(ROOT, "unit.c"): UNIT}
    context = cpc.build_project_context(files)
    path = os.path.join(ROOT, "unit.c")
    unit = {"name": name, "fid": "F1", "source_text": UNIT, "source_path": path, "source_text_complete": True,
            "project_scope": cpc.build_scopes(context, [path])[path], "input_vars": list(input_vars),
            "output_vars": list(output_vars),
            "param_types": {"g_a": "U8", "g_b": "U8", "g_o": "U8", "g_gain_open": "U8"}}
    unit.update(kw)
    return unit


def _gen(unit, cache=None, **kw):
    return generate_sequences(unit, None, type_cache={} if cache is None else cache, extended=True, **kw)


def _complete(unit, cache=None):
    seqs = _gen(unit, cache, boundary_rows=False)
    return complete_source_read_inputs(unit, seqs, lambda: _gen(unit, cache, boundary_rows=False))


def _stated(seqs, *names, hi=255):
    """Rows that give every named input a value inside its type (BV_*_INV rows leave it on purpose; MC/DC vectors
    state only what their decision reads)."""
    return [s for s in seqs if all(n in s["inputs"] and 0 <= s["inputs"][n] <= hi for n in names)]


def _slots(seqs, var="g_o"):
    return [s["expected_evidence"][var] for s in seqs if var in (s.get("expected_evidence") or {})]


# ── what the reasons say ────────────────────────────────────────────────────────────────────────────────────────

def test_names_come_from_the_reasons_including_nested_ones():
    seqs = [{"seq_num": 1, "expected_evidence": {
                "g_o": {"reason": "initial_value_not_in_inputs:g_x"},
                "g_p": {"reason": "path_dependent:branch_on_unknown:initial_value_not_in_inputs:g_buf[3]"}}},
            {"seq_num": 2, "expected_evidence": {"g_o": {"reason": "initial_value_not_in_inputs:g_x"},
                                                 "g_q": {"reason": "initial_value_not_in_inputs:s.member"},
                                                 "g_r": {"reason": "initial_value_not_in_inputs:m[1][2]"},
                                                 "g_s": {"reason": "undefined_behavior:signed_overflow"}}}]
    got = source_read_names(seqs)
    assert got == {"g_x": {"slots": 2, "sequences": [1, 2]}, "g_buf[3]": {"slots": 1, "sequences": [1]}}


def test_only_program_objects_of_the_unit_are_inputs():
    scope = _unit("close_gain", ["g_a"])["project_scope"]
    assert _source_object_decl(scope, "g_gain_close") == ("U8", "")
    assert _source_object_decl(scope, "s_state") == ("U8", "")          # a file-scope static is state the test sets
    assert _source_object_decl(scope, "g_buf[2]") == ("U8", "")
    assert _source_object_decl(scope, "g_buf[4]") == (None, "element_out_of_range")
    assert _source_object_decl(scope, "g_buf") == (None, "array_object")
    assert _source_object_decl(scope, "g_kc") == (None, "const_object")   # const without a visible value: not settable
    assert _source_object_decl(scope, "g_reg") == (None, "volatile_object")
    assert _source_object_decl(scope, "g_f") == (None, "float_object")
    assert _source_object_decl(scope, "t_local") == (None, "not_a_program_object")
    assert _source_object_decl(scope, "g_o[1]") == (None, "element_of_unmodeled_array")


def test_the_translation_unit_types_what_the_table_cannot_and_says_why_it_cannot():
    unit = _unit("st", ["g_a"])
    types, enums, why = scope_input_types(unit, ["s_state", "s_level", "g_buf[1]", "g_m", "g_big", "g_raw", "t_local",
                                                 "g_reg"])
    assert types == {"s_state": "uint8_t", "s_level": "int16_t", "g_buf[1]": "uint8_t", "g_raw": "uint16_t"}
    assert enums == {"g_m": {"values": [0, 1, 5], "source": "translation_unit_declaration"}}   # no globals map needed
    assert why == {"g_big": "no_boundary_table_for_type", "t_local": "type_not_resolved",
                   "g_reg": "not_a_settable_integer"}


# ── adding inputs ───────────────────────────────────────────────────────────────────────────────────────────────

def test_a_read_the_input_list_missed_becomes_a_typed_input_and_the_rows_derive():
    # the design table names the open gain; the source reads the close gain (HDPDM01 s_MoveStartClose_GainMeasure)
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    before = _gen(dict(unit))
    assert any("initial_value_not_in_inputs:g_gain_close" in str(e.get("reason")) for e in _slots(before))
    seqs = _complete(unit)
    assert unit["input_vars"] == ["g_a", "g_gain_open", "g_gain_close"]
    assert unit["design_input_vars"] == ["g_a", "g_gain_open"]
    rec = unit["source_read_inputs"]
    assert list(rec["added"]) == ["g_gain_close"] and rec["added"]["g_gain_close"]["type"] == "U8"
    assert rec["added"]["g_gain_close"]["value_type"] == "uint8_t"
    assert rec["rounds"] == 1 and rec["remaining"] == {}
    assert unit["var_types"]["g_gain_close"] == "uint8_t"
    rows = _stated(seqs, "g_a", "g_gain_close")
    assert rows
    for s in rows:
        ev = s["expected_evidence"]["g_o"]
        assert ev["status"] == "derived", (s["strategy"], ev)
        assert s["expected"]["g_o"] == (s["inputs"]["g_gain_close"] if s["inputs"]["g_a"] > 10 else 0)
    assert {0, 255} <= {s["inputs"].get("g_gain_close") for s in seqs}    # its extremes come from the OAT rows


def test_the_reference_rows_keep_their_places_and_design_values_after_inputs_are_added():
    # review R1 C1: added inputs used to add condition-combination rows inside the reference region and push the MC/DC
    # slots down — the extended document no longer began with the reference document
    ref_unit = _unit("close_gain", ["g_a"])
    ref = generate_sequences(ref_unit, None, type_cache={})
    unit = _unit("close_gain", ["g_a"])
    ext = _complete(unit)
    assert "g_gain_close" in unit["input_vars"]
    n = unit["base_strategy_count"]
    assert n == ref_unit["base_strategy_count"] == len(ref)
    assert [s["strategy"] for s in ext[:n]] == [s["strategy"] for s in ref]
    assert [{k: v for k, v in s["inputs"].items() if k != "g_gain_close"} for s in ext[:n]] == [s["inputs"] for s in ref]
    assert not any(s.get("tc_profile") for s in ext[:n]) and all(s.get("tc_profile") == "extended" for s in ext[n:])
    assert unit["mcdc_design"]["selected_inputs"] == ref_unit["mcdc_design"]["selected_inputs"]
    fixed = {s["inputs"]["g_gain_close"] for s in ext[:n] if "g_gain_close" in s["inputs"]}
    assert len(fixed) == 1 and all("고정값" in s["description"] for s in ext[:n] if "g_gain_close" in s["inputs"])
    oat = [s for s in ext if s["strategy"].startswith("OAT_") and s["inputs"].get("g_gain_close") in (0, 255)]
    assert len(oat) == 2 and all(is_extended_strategy(s["strategy"]) for s in oat)


def test_an_added_input_in_a_reference_row_does_not_turn_a_derived_slot_into_undefined_behaviour():
    # review R1 W3: BV_MAX pushed the added `s_level` to 32767 too; `s_level + 1` then overflows and the whole run —
    # including `g_a`'s part the row had derived before — became undefined behaviour
    unit = _unit("edge", ["g_a"])
    seqs = _complete(unit)
    assert "s_level" in unit["input_vars"]
    for name in ("BV_MIN", "BV_MID", "BV_MAX", "MIXED"):
        row = next(s for s in seqs if s["strategy"] == name)
        assert row["expected_evidence"]["g_o"]["status"] == "derived", (name, row["expected_evidence"]["g_o"])
    oat_max = next(s for s in seqs if s["strategy"].startswith("OAT_") and s["inputs"].get("s_level") == 32767)
    assert "undefined_behavior" in oat_max["expected_evidence"]["g_o"]["reason"]   # the extreme is still exercised


def test_a_read_behind_a_newly_opened_path_is_added_in_the_next_round():
    unit = _unit("chain", ["g_a"])
    seqs = _complete(unit)
    rec = unit["source_read_inputs"]
    # the oracle names one unknown per slot: `s_state` first, `g_step` once `s_state` has a value
    assert set(rec["added"]) == {"s_state", "g_step"} and rec["rounds"] == 2
    assert {a["round"] for a in rec["added"].values()} == {1, 2}
    rows = _stated(seqs, "g_a", "g_step", "s_state")
    assert rows and all(s["expected_evidence"]["g_o"]["status"] == "derived" for s in rows)


def test_an_array_element_is_added_by_its_element_name():
    unit = _unit("elem", ["g_a"])
    seqs = _complete(unit)
    assert "g_buf[2]" in unit["input_vars"]
    rows = _stated(seqs, "g_buf[2]")
    assert rows and all(s["expected"]["g_o"] == s["inputs"]["g_buf[2]"] for s in rows)


def test_an_enum_object_takes_the_enumerators_of_its_own_translation_unit():
    # review R1 I5: the globals map is keyed by name — a same-named static elsewhere must not lend its enumerators
    unit = _unit("mode_out", ["g_a"])
    seqs = _complete(unit)
    assert unit["source_read_inputs"]["added"]["g_m"]["value_type"] == "enum"
    assert unit["value_domains"]["g_m"]["values"] == [0, 1, 5]
    rows = [s for s in _stated(seqs, "g_m") if not s["strategy"].endswith("_INV")]   # INV rows step off the set
    assert {s["inputs"]["g_m"] for s in rows} <= {0, 1, 5} and {1, 2} == {s["expected"]["g_o"] for s in rows}


def test_an_output_whose_initial_value_is_read_becomes_an_input_too():
    # `g_o` keeps its initial value on the path that does not write it — the row must say what that value was
    unit = _unit("inout", ["g_a"])
    seqs = _complete(unit)
    assert "g_o" in unit["input_vars"]
    rows = _stated(seqs, "g_a", "g_o")
    assert rows and any(s["inputs"]["g_a"] <= 10 for s in rows)
    for s in rows:
        assert s["expected"]["g_o"] == (5 if s["inputs"]["g_a"] > 10 else s["inputs"]["g_o"])


def test_a_read_whose_type_has_no_boundary_table_is_not_added_and_remains():
    unit = _unit("big", ["g_a"])
    _complete(unit)
    assert "g_big" not in unit["input_vars"]
    rec = unit["source_read_inputs"]
    assert rec["not_added"] == {"g_big": "no_boundary_table_for_type"}
    assert set(rec["remaining"]) == {"g_big"}      # review R1 W2: counted whether or not anything else was added


@pytest.mark.parametrize("cache, values", [({"g_raw": "unsigned int"}, {0, 65535}),   # map type the table cannot read
                                           ({"g_raw": "U8"}, {0, 255})])               # a readable map type wins
def test_the_declared_type_wins_only_where_the_globals_map_type_is_unreadable(cache, values):
    # review R1 W1: a map type the table cannot read used to shadow the declaration — the added column stayed blank
    unit = _unit("raw", ["g_a"])
    seqs = _complete(unit, cache)
    assert "g_raw" in unit["input_vars"]
    assert values <= {s["inputs"].get("g_raw") for s in seqs}
    assert unit["source_read_inputs"]["remaining"] == {}


def test_the_reference_profile_is_not_changed():
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    seqs = generate_sequences(unit, 24, type_cache={})
    assert unit["input_vars"] == ["g_a", "g_gain_open"] and "source_read_inputs" not in unit
    assert not any(s.get("tc_profile") for s in seqs) and "_boundary_ctx" not in unit
    assert any("initial_value_not_in_inputs:g_gain_close" in str(e.get("reason")) for e in _slots(seqs))


def test_a_unit_without_inputs_or_outputs_is_left_alone():
    unit = _unit("close_gain", [], output_vars=())
    seqs = [{"seq_num": 1, "expected_evidence": {"g_o": {"reason": "initial_value_not_in_inputs:g_a"}}}]
    assert complete_source_read_inputs(unit, seqs, lambda: 1 / 0) is seqs
    assert unit["source_read_inputs"]["not_added"] == {"g_a": "unit_without_inputs_or_outputs"}


def test_a_full_input_list_is_not_widened_and_says_so(monkeypatch):
    monkeypatch.setattr(suts, "_INPUT_COL_END", suts._INPUT_COL_START + 1)   # two input columns
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    _complete(unit)
    assert unit["input_vars"] == ["g_a", "g_gain_open"]
    assert unit["source_read_inputs"]["not_added"] == {"g_gain_close": "input_columns_full"}
    monkeypatch.setattr(suts, "_INPUT_COL_END", suts._INPUT_COL_START + 2)   # exactly one free column: taken
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    _complete(unit)
    assert unit["input_vars"] == ["g_a", "g_gain_open", "g_gain_close"]


# ── the production path ─────────────────────────────────────────────────────────────────────────────────────────

def test_the_extended_path_searches_boundaries_once_after_the_inputs_are_complete(monkeypatch):
    # review R1 W6/W7: the first search was thrown away whenever inputs were added; no test ran the production path
    monkeypatch.setattr(suts, "_globals_type_cache", {"g_a": "U8", "g_o": "U8", "g_gain_open": "U8"})
    calls = []
    real = suts.find_boundaries

    def counting(unit, *a, **k):
        calls.append(list(unit["input_vars"]))
        return real(unit, *a, **k)
    monkeypatch.setattr(suts, "find_boundaries", counting)
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    seqs = extended_unit_sequences(unit)
    assert calls == [["g_a", "g_gain_open", "g_gain_close"]]
    assert "_boundary_ctx" not in unit
    rows = [s for s in seqs if s["strategy"].startswith(br.BOUNDARY_PREFIX)]
    assert rows and all(s["expected_evidence"]["g_o"]["status"] == "derived" for s in rows)
    assert [s["seq_num"] for s in seqs] == list(range(1, len(seqs) + 1))


def test_a_failing_completion_leaves_the_design_inputs_and_still_makes_the_document(monkeypatch):
    monkeypatch.setattr(suts, "_globals_type_cache", {"g_a": "U8", "g_o": "U8", "g_gain_open": "U8"})

    def boom(unit, *a, **k):
        raise ValueError("probe")
    monkeypatch.setattr(suts, "scope_input_types", boom)
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    seqs = extended_unit_sequences(unit)
    assert unit["input_vars"] == ["g_a", "g_gain_open"] and "design_input_vars" not in unit
    assert unit["source_read_inputs"]["error"] == "ValueError" and seqs


# ── MC/DC fill rows ─────────────────────────────────────────────────────────────────────────────────────────────

def test_an_mcdc_vector_that_leaves_a_read_input_blank_gets_a_filled_companion_row():
    unit = _unit("st", ["g_a", "s_level"], param_types={"g_a": "U8", "g_o": "U8", "s_level": "S16"})
    seqs = _gen(unit)
    vectors = [s for s in seqs if s["strategy"].startswith("MCDC_")]
    fills = [s for s in seqs if s["strategy"].startswith(MCDC_FILL_PREFIX)]
    assert vectors and all("s_level" not in s["inputs"] for s in vectors)           # the design vectors stay as designed
    assert fills
    for f in fills:
        src = next(s for s in vectors if s["seq_num"] == f["mcdc_fill"]["from_sequence"])
        assert f["inputs"]["g_a"] == src["inputs"]["g_a"] and set(f["mcdc_fill"]["filled"]) == {"s_level"}
        assert f["expected_evidence"]["g_o"]["status"] == "derived"
        assert f["expected"]["g_o"] == (f["inputs"]["s_level"] & 0xFF if f["inputs"]["g_a"] > 10 else 0)
        assert is_extended_strategy(f["strategy"]) and f["tc_profile"] == "extended"
        assert resolve_seq_gen_method(f["strategy"]) == resolve_seq_gen_method("MCDC_0")
        assert "구성원이 아니다" in f["description"]
    assert unit["mcdc_design"]["decisions"][0]["pairs"][0]["retained_status"] == "retained"   # pairs untouched


def test_a_fill_row_that_derives_nothing_new_is_dropped():
    # review R1 W4: `g_b` is not read — filling it gives the vector's own values; a third of the first run's fill rows
    unit = _unit("dec", ["g_a", "g_b"])
    seqs = _gen(unit)
    assert any(s["strategy"].startswith("MCDC_") and "g_b" not in s["inputs"] for s in seqs)
    assert not any(s["strategy"].startswith(MCDC_FILL_PREFIX) for s in seqs)
    assert unit["mcdc_fill"]["pruned"] == unit["mcdc_fill"]["rows_with_blanks"] > 0 and unit["mcdc_fill"]["rows"] == 0
    assert [s["seq_num"] for s in seqs] == list(range(1, len(seqs) + 1))


def test_fill_rows_skip_vectors_already_in_the_suite():
    seqs = [{"seq_num": 1, "strategy": "BV_MID", "inputs": {"a": 1, "b": 5}},
            {"seq_num": 2, "strategy": "MCDC_0", "inputs": {"a": 1}},
            {"seq_num": 3, "strategy": "MCDC_1", "inputs": {"a": 2}},
            {"seq_num": 4, "strategy": "MCDC_2", "inputs": {"a": 2}}]
    unit = {}
    suts._append_mcdc_fill_rows(unit, seqs, ["a", "b"], ["o"], {"a": "uint8_t", "b": "uint8_t"}, {"b": 5})
    assert [s["inputs"] for s in seqs[4:]] == [{"a": 2, "b": 5}]
    assert unit["mcdc_fill"] == {"mcdc_rows": 3, "rows_with_blanks": 3, "rows": 1, "not_fillable": 0, "duplicates": 2,
                                 "pruned": 0}


def test_the_reference_profile_has_no_fill_rows_and_a_guessed_type_is_not_filled():
    unit = _unit("st", ["g_a", "s_level"], param_types={"g_a": "U8", "g_o": "U8", "s_level": "S16"})
    assert not any(s["strategy"].startswith(MCDC_FILL_PREFIX) for s in generate_sequences(unit, 40, type_cache={}))
    # review R1 W5: `s_level` typed by nothing but its name pattern is a guess — MC/DC refuses it, so does the fill
    unit = _unit("st", ["g_a", "s_level"], param_types={"g_a": "U8", "g_o": "U8"})
    seqs = _gen(unit)
    assert not any(s["strategy"].startswith(MCDC_FILL_PREFIX) for s in seqs)
    assert unit["mcdc_fill"]["not_fillable"] == unit["mcdc_fill"]["rows_with_blanks"] > 0


def test_a_failing_fill_step_costs_only_the_fill_rows(monkeypatch):
    def boom(*a, **k):
        raise KeyError("probe")
    monkeypatch.setattr(suts, "_append_mcdc_fill_rows", boom)
    unit = _unit("st", ["g_a", "s_level"], param_types={"g_a": "U8", "g_o": "U8", "s_level": "S16"})
    seqs = _gen(unit)
    assert seqs and unit["mcdc_fill"] == {"error": "KeyError"}


# ── what the document says ──────────────────────────────────────────────────────────────────────────────────────

def test_the_reference_profile_reports_the_gap_as_a_source_finding():
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    seqs = generate_sequences(unit, 24, type_cache={})
    seqs.append({"seq_num": 99, "expected_evidence": {"g_o": {"reason": "initial_value_not_in_inputs:t_local"},
                                                      "g_p": {"reason": "initial_value_not_in_inputs:g_a"}}})
    gaps = input_list_gaps(unit, seqs)
    assert list(gaps) == ["g_gain_close"]          # a local is not a program object; g_a already is an input
    assert gaps["g_gain_close"]["slots"] >= 1 and gaps["g_gain_close"]["sequences"] and not gaps["g_gain_close"]["added"]


def test_the_extended_profile_reports_what_it_added_as_not_blocked_here():
    unit = _unit("chain", ["g_a"])
    seqs = _complete(unit)
    gaps = input_list_gaps(unit, seqs)
    assert set(gaps) == {"s_state", "g_step"}
    assert all(g == {"slots": 0, "sequences": [], "added": True} for g in gaps.values())   # review R1 W5


def test_the_findings_sheet_has_one_gap_row_per_function_and_the_ub_count_stays_apart():
    import openpyxl

    from generators import source_findings as sf
    from generators.suts import _summarize_source_findings
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    seqs = generate_sequences(unit, 24, type_cache={})
    seqs[0]["expected_evidence"]["g_x"] = {"reason": "undefined_behavior:division_by_zero"}
    summary = _summarize_source_findings([unit], {"F1": seqs})
    assert summary["findings"] == 1 and summary["by_kind"] == {"division_by_zero": 1}    # UB only
    assert summary["input_list_gaps"]["functions"] == 1 and summary["input_list_gaps"]["names"] == 1
    rows = sf.collect_input_gap_findings([unit], {"F1": input_list_gaps(unit, seqs)}, {"F1": "SwUTC_F1"})
    wb = openpyxl.Workbook()
    assert sf.write_source_findings_sheet(wb, rows) == 1
    ws = wb[sf.FINDINGS_SHEET]
    got = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]], strict=True))
    assert got["Kind"] == "read_not_in_unit_inputs" and got["Observables"] == "g_gain_close"
    assert got["Example Test Case ID"] == "SwUTC_F1" and got["Reproduction"] == sf.INPUT_GAP_REPRODUCTION
    assert "입력 목록" in got["Meaning"]


def test_the_summary_counts_units_names_reasons_and_what_remains():
    units = [{"source_read_inputs": {"added": {"a": {"slots": 3}, "b": {"slots": 1}}, "not_added": {"c": "float_object"},
                                     "rounds": 2, "remaining": {}}},
             {"source_read_inputs": {"added": {}, "not_added": {"d": "input_columns_full", "e": "float_object"},
                                     "rounds": 0, "remaining": {"f": 2}}},
             {"name": "no record"}]
    assert summarize_source_read_inputs(units) == {
        "units": 2, "units_with_added": 1, "names_added": 2, "slots_first_seen": 4,
        "not_added": {"float_object": 2, "input_columns_full": 1}, "units_input_columns_full": 1, "max_rounds": 2,
        "units_with_remaining": 1, "names_remaining": 1, "errors": 0}


def test_the_disclosures_say_what_was_added_what_was_not_and_what_remains():
    from report_gen.generation_disclosures import build_disclosures
    qr = {"source_read_inputs": {"units": 10, "units_with_added": 4, "names_added": 7, "slots_first_seen": 40,
                                 "not_added": {"float_object": 2, "input_columns_full": 1}, "units_input_columns_full": 1,
                                 "max_rounds": 2, "units_with_remaining": 1, "names_remaining": 3},
          "source_findings": {"findings": 0, "functions": 0, "sequences": 0, "by_kind": {},
                              "input_list_gaps": {"functions": 4, "names": 7, "slots": 40}}}
    items = {i["key"]: i for i in build_disclosures("suts", qr)}
    add = items["suts_source_read_inputs"]
    assert add["value"] == "unit 4/10 · 이름 7 · 남은 이름 3"
    for text in ("float_object", "input_columns_full", "입력 열 상한에 막힌 unit 1", "남은 이름 3(unit 1)", "설계서 입력 표와 다른 열",
                 "고정값", "MC/DC"):
        assert text in add["note"], text
    assert "enum 값 집합 > HSIS" not in add["note"]   # review R1 W5: added names get no HSIS range
    assert add["tone"] == "warning"
    gap = items["suts_input_list_gaps"]
    assert gap["value"] == "함수 4 · 객체 7" and gap["tone"] == "warning"
    assert "read_not_in_unit_inputs" in gap["note"]
    keys = {i["key"] for i in build_disclosures("suts", {})}
    assert "suts_source_read_inputs" not in keys and "suts_input_list_gaps" not in keys   # not recorded → no item
    older = {"source_findings": {"findings": 2, "functions": 1, "sequences": 2, "by_kind": {"division_by_zero": 2}}}
    assert "suts_input_list_gaps" not in {i["key"] for i in build_disclosures("suts", older)}


def test_the_fill_rows_are_disclosed_with_what_could_not_be_filled():
    from report_gen.generation_disclosures import build_disclosures
    qr = {"mcdc_fill": {"mcdc_rows": 50, "rows_with_blanks": 30, "rows": 20, "not_fillable": 3, "duplicates": 1,
                        "pruned": 6, "units_with_rows": 9}}
    item = {i["key"]: i for i in build_disclosures("suts", qr)}["suts_mcdc_fill_rows"]
    assert item["value"] == "20행 (공란이 있는 MC/DC 벡터 30/50)"
    for text in ("구성원이 아니며", "채우지 못한 벡터 3", "더하지 않은 것 1", "더 도출하는 칸이 없어 뺀 행 6", "열거자",
                 "선언 타입"):
        assert text in item["note"], text
    assert "suts_mcdc_fill_rows" not in {i["key"] for i in build_disclosures("suts", {})}


# ── review round 2 ──────────────────────────────────────────────────────────────────────────────────────────────

def test_a_global_row_keeps_moving_its_global_when_the_completion_added_it_as_an_input():
    # review R2 C1: the fixed-value step overwrote the GLOBAL row's target — two GLOBAL rows became identical and kept
    # a "minimum" label (KJPDS02_PV `s_MotorCompensation`)
    ref_unit = _unit("close_gain", ["g_a"], indirect_vars=["g_gain_close"])
    ref = generate_sequences(ref_unit, None, type_cache={})
    unit = _unit("close_gain", ["g_a"], indirect_vars=["g_gain_close"])
    ext = _complete(unit)
    assert "g_gain_close" in unit["input_vars"]
    n = unit["base_strategy_count"]
    assert [s["strategy"] for s in ext[:n]] == [s["strategy"] for s in ref]
    g_ref = next(s for s in ref if s["strategy"] == "GLOBAL_0")
    g_ext = next(s for s in ext if s["strategy"] == "GLOBAL_0")
    assert g_ext["inputs"] == g_ref["inputs"] and g_ext["inputs"]["g_gain_close"] == 0
    assert "고정값" not in g_ext["description"]


def test_a_global_the_completion_types_does_not_add_a_global_row_to_the_reference_region():
    # review R2 C1 (latent): a global the map cannot type gets a type from R19 — its GLOBAL row would slide into the
    # reference region and renumber the rows after it
    cache = {"g_a": "U8", "g_o": "U8", "g_raw": "unsigned int"}
    ref_unit = _unit("raw", ["g_a"], indirect_vars=["g_raw"])
    ref = generate_sequences(ref_unit, None, type_cache=cache)
    unit = _unit("raw", ["g_a"], indirect_vars=["g_raw"])
    ext = _complete(unit, cache)
    assert "g_raw" in unit["input_vars"]
    assert unit["base_strategy_count"] == ref_unit["base_strategy_count"]
    assert [s["strategy"] for s in ext[:unit["base_strategy_count"]]] == [s["strategy"] for s in ref]


def test_one_design_input_gets_no_one_at_a_time_rows_that_repeat_the_boundary_rows():
    # review R2 W2': with the added input fixed, OAT on the only design input equals BV_MIN / BV_MAX
    unit = _unit("close_gain", ["g_a"])
    seqs = _complete(unit)
    oat = [s["strategy"] for s in seqs if s["strategy"].startswith("OAT_")]
    assert oat == ["OAT_1_MIN", "OAT_1_MAX"]          # only the added input moves alone
    vectors = [tuple(sorted(s["inputs"].items())) for s in seqs if s["strategy"].startswith(("BV_", "OAT_"))]
    assert len(vectors) == len(set(vectors))


def test_a_decision_on_an_added_input_says_why_it_has_no_mcdc_design_and_the_reference_slots_stay():
    # review R2 W3' / W4' N3: the design runs on the design input list; the refusal names the added input as such
    ref_unit = _unit("two", ["g_a"])
    generate_sequences(ref_unit, None, type_cache={})
    unit = _unit("two", ["g_a"])
    _complete(unit)
    assert {"g_b", "g_gain_close"} <= set(unit["input_vars"])
    assert unit["mcdc_design"]["selected_inputs"] == ref_unit["mcdc_design"]["selected_inputs"]
    reasons = [d["reason"] for d in unit["mcdc_design"]["decisions"]]
    assert "decision_reads_source_read_input_not_designed:g_b" in reasons
    assert not any(r.startswith("decision_variable_not_in_unit_inputs:") for r in reasons)


def test_a_second_round_keeps_the_design_list_of_the_first():
    # review R2 W4' N6: overwriting `design_input_vars` in round 2 would bring round-1 inputs into the catalog
    ref_unit = _unit("chain", ["g_a"])
    ref = generate_sequences(ref_unit, None, type_cache={})
    unit = _unit("chain", ["g_a"])
    ext = _complete(unit)
    assert unit["source_read_inputs"]["rounds"] == 2 and unit["design_input_vars"] == ["g_a"]
    n = unit["base_strategy_count"]
    assert [s["strategy"] for s in ext[:n]] == [s["strategy"] for s in ref]


def test_an_added_column_no_row_can_fill_is_reported_as_remaining(monkeypatch):
    # review R2 W4' N7: a type key the boundary table cannot read leaves the column blank everywhere
    real = suts.scope_input_types

    def odd(unit, names):
        types, enums, why = real(unit, names)
        return {n: "struct Odd" for n in types}, enums, why
    monkeypatch.setattr(suts, "scope_input_types", odd)
    unit = _unit("close_gain", ["g_a"])
    _complete(unit)
    rec = unit["source_read_inputs"]
    assert rec["added"]["g_gain_close"]["no_value_in_any_row"] is True
    assert "g_gain_close" in rec["remaining"]


def test_a_completion_failing_in_a_later_round_restores_the_design_inputs(monkeypatch):
    # review R2 W4' N9: the failure happens after round 1 already added `s_state`
    monkeypatch.setattr(suts, "_globals_type_cache", {"g_a": "U8", "g_o": "U8"})
    real, calls = suts.scope_input_types, []

    def second_fails(unit, names):
        calls.append(names)
        if len(calls) > 1:
            raise ValueError("probe")
        return real(unit, names)
    monkeypatch.setattr(suts, "scope_input_types", second_fails)
    unit = _unit("chain", ["g_a"])
    seqs = extended_unit_sequences(unit)
    assert len(calls) == 2
    assert unit["input_vars"] == ["g_a"] and "design_input_vars" not in unit and "source_input_types" not in unit
    assert "s_state" not in (unit.get("value_domains") or {})
    assert unit["source_read_inputs"]["error"] == "ValueError"
    assert seqs and all(set(s["inputs"]) <= {"g_a"} for s in seqs)


def test_the_disclosures_say_when_a_unit_fell_back_after_an_error():
    from report_gen.generation_disclosures import build_disclosures
    qr = {"source_read_inputs": {"units": 3, "units_with_added": 1, "names_added": 1, "slots_first_seen": 2,
                                 "not_added": {}, "units_input_columns_full": 0, "max_rounds": 1,
                                 "units_with_remaining": 0, "names_remaining": 0, "errors": 1},
          "mcdc_fill": {"mcdc_rows": 4, "rows_with_blanks": 2, "rows": 1, "not_fillable": 0, "duplicates": 0,
                        "pruned": 1, "units_with_rows": 1, "errors": 1}}
    items = {i["key"]: i for i in build_disclosures("suts", qr)}
    assert "오류로 설계 입력 목록 그대로 둔 unit 1" in items["suts_source_read_inputs"]["note"]
    assert items["suts_source_read_inputs"]["tone"] == "warning"
    assert "채움 단계 오류로 채움 행 없이 둔 unit 1" in items["suts_mcdc_fill_rows"]["note"]
    assert items["suts_mcdc_fill_rows"]["tone"] == "warning"
    assert "더 도출하는 칸이 없어 뺀 행 1" in items["suts_mcdc_fill_rows"]["note"]


# ── review round 3 ──────────────────────────────────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("map_domain", [{"values": [0, 1, 5]}, None])
def test_an_enum_global_keeps_or_lacks_its_global_row_as_in_the_reference(map_domain):
    # review R3 W1: the completion replaces the map's value set with the unit's enumerators — structure is decided on the
    # value set from before (row kept when the map had one; no row added when it had none)
    cache = {"g_a": "U8", "g_o": "U8", "g_m": "Mode_t"}
    extra = {"indirect_vars": ["g_m"]}
    if map_domain:
        extra["value_domains"] = {"g_m": dict(map_domain)}
    ref_unit = _unit("mode_out", ["g_a"], **extra)
    ref = generate_sequences(ref_unit, None, type_cache=cache)
    unit = _unit("mode_out", ["g_a"], **{k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in extra.items()})
    ext = _complete(unit, cache)
    assert "g_m" in unit["input_vars"] and unit["value_domains"]["g_m"]["source"] == "translation_unit_declaration"
    assert unit["base_strategy_count"] == ref_unit["base_strategy_count"]
    assert [s["strategy"] for s in ext[:unit["base_strategy_count"]]] == [s["strategy"] for s in ref]
    assert any(s["strategy"] == "GLOBAL_0" for s in ref) is bool(map_domain)


def test_generate_suts_aggregates_fill_pruning_and_errors_into_the_quality_report(tmp_path, monkeypatch):
    # review R3 W2: the disclosure tests fed numbers in directly — nothing checked what the generator aggregates
    (tmp_path / "m.c").write_text("typedef unsigned char U8;\nU8 Fn(U8 a, U8 b)\n{\n    if (a > 1) { return b; }\n"
                                  "    return 0;\n}\n", encoding="utf-8")

    calls = []

    def fake(unit):
        calls.append(unit["fid"])
        seqs = generate_sequences(unit, None)
        unit["mcdc_fill"] = {"mcdc_rows": 2, "rows_with_blanks": 2, "rows": 0, "not_fillable": 0, "duplicates": 0,
                             "pruned": 2}
        unit["source_read_inputs"] = {"error": "ValueError", "added": {}, "not_added": {}, "rounds": 0, "remaining": {}}
        return seqs
    monkeypatch.setattr(suts, "extended_unit_sequences", fake)
    out = suts.generate_suts(source_root=str(tmp_path), output_path=str(tmp_path / "u.xlsm"), scope="source",
                             tc_profile="extended")
    q = out["quality_report"]
    assert calls and q["mcdc_fill"]["pruned"] == 2 * len(calls) and q["mcdc_fill"]["errors"] == 0
    assert q["source_read_inputs"]["errors"] == len(calls)
    monkeypatch.setattr(suts, "extended_unit_sequences",
                        lambda unit: (unit.update(mcdc_fill={"error": "KeyError"}), generate_sequences(unit, None))[1])
    q = suts.generate_suts(source_root=str(tmp_path), output_path=str(tmp_path / "v.xlsm"), scope="source",
                           tc_profile="extended")["quality_report"]
    assert q["mcdc_fill"]["errors"] == len(calls)


def test_remaining_is_counted_again_after_the_boundary_rows(monkeypatch):
    # review R3 W2 (R2 I4): a boundary row may read a program object no earlier row read
    monkeypatch.setattr(suts, "_globals_type_cache", {"g_a": "U8", "g_o": "U8", "g_gain_open": "U8"})
    real = suts.append_boundary_rows

    def plus_one(unit, seqs, ctx=None):
        seqs = real(unit, seqs, ctx)
        seqs.append({"seq_num": len(seqs) + 1, "strategy": "BND_x", "inputs": {"g_a": 1},
                     "expected_evidence": {"g_o": {"reason": "initial_value_not_in_inputs:g_step"}}})
        return seqs
    monkeypatch.setattr(suts, "append_boundary_rows", plus_one)
    unit = _unit("close_gain", ["g_a", "g_gain_open"])
    extended_unit_sequences(unit)
    assert unit["source_read_inputs"]["remaining"] == {"g_step": 1}
