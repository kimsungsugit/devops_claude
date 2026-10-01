"""R53 — STS 요구 문서 발견 사항을 RAG 지식베이스에 적재하고 채팅이 그것을 찾게 한다.

사용자 방향(2026-10-01) "중간중간에 RAG 저장할 것들 … 리스트 뽑아주고 진행해": 생성기가 요구 문서에서 찾은 것(검토 항목 ·
경계 포함 불일치 · 값 차이 후보 · 요구별 경계 근거와 HW 허용오차 경로)을 공유 KB(`reports/kb_store`)에 넣고, 채팅 어시스턴트의
질문이 요구·HW ID 를 적으면 그 ID 의 항목을 찾는다(질문 종류 · 검색 칸과 무관한 자기 블록).

지키는 것(리뷰 1 차 반영):
* 원문·사유는 시트에 적힌 그대로(요약·재서술 없음), 자기 분류 `sts_findings` — UDS 생성 AI 가 프롬프트에 넣지 않는다(W1),
* 채팅의 KB 는 세션 폴더라 공유 KB 를 따로 찾는다(C1),
* 키는 내용까지 — 바뀐 내용은 새 항목, 옛 항목은 지운다(W2), 같은 키 중복도 지운다(W6),
* 시트가 없는 파일 · 다른 프로젝트 파일 · sqlite 가 아닌 저장소는 지우기 전에 멈춘다(W3 · W4),
* 긴 근거는 임베딩 상한 안으로 나눈다(W7), 쓴 뒤 저장소를 다시 읽어 실제 건수를 보고한다(W8),
* 지우기는 메모리·엔트리 파일·DB 행 모두(다음 로드가 DB 에서 오므로 DB 가 본체).
"""
from __future__ import annotations

import openpyxl
import pytest

import config
import workflow.rag.embedder as emb
from workflow.rag import KnowledgeBase, _clear_kb_cache, get_kb


@pytest.fixture
def kb_env(monkeypatch, tmp_path):
    """sqlite 저장소 + 임베딩 네트워크 차단(`test_rag_kb_cache.kb_env` 와 같은 격리)."""
    calls = []

    def _stub_embedding(text, *, meta_out=None):
        calls.append(text)
        if meta_out is not None:
            meta_out.update({"embed_source": "stub", "embed_model": "stub", "embed_dim": 4,
                             "degraded": "degraded" in str(text)})
        return [1.0, 0.0, 0.0, float(len(str(text)) % 7)]

    monkeypatch.setattr(emb, "get_embedding", _stub_embedding)
    monkeypatch.setattr(config, "KB_GLOBAL_DIR", "")
    monkeypatch.setattr(config, "KB_SOURCES_DIR", "")
    monkeypatch.setattr(config, "KB_STORAGE", "sqlite")
    monkeypatch.setattr(config, "FORCE_PGVECTOR", False)
    monkeypatch.setattr(config, "KB_CACHE_ENABLED", False)
    monkeypatch.setattr(config, "KB_MAX_ENTRIES", 5000, raising=False)
    _clear_kb_cache()
    yield tmp_path, calls
    _clear_kb_cache()


def test_removed_documents_stay_removed_after_a_reload(kb_env):
    tmp, _calls = kb_env
    kb = KnowledgeBase(tmp / "kb_store")
    kb.add_document("a", "first", category="sts_findings", tags=["x"])
    kb.add_document("b", "second", category="sts_findings", tags=["y"])
    first = kb.data[0]["id"]
    assert kb.remove_documents([first, None, ""]) == 1
    assert [e["error_raw"] for e in kb.data] == ["b"]
    assert not (tmp / "kb_store" / f"{first}.json").exists()
    # the next load comes from the DB: the row is gone too (a stale entry must not come back)
    assert [e["error_raw"] for e in KnowledgeBase(tmp / "kb_store").data] == ["b"]
    assert kb.remove_documents([]) == 0


def _sts(path, review_rows, evidence_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Requirement Review"
    ws.append(["Kind", "SRS ID", "Source Document", "Fact", "Reason", "To Decide", "If Filled",
               "Evidence (Other Sentences)", "Source Line", "Line SHA-256"])
    for r in review_rows:
        ws.append(r)
    ev = wb.create_sheet("Requirement Evidence")
    ev.append(["Test Case ID", "SRS ID", "Kind", "Subject", "Reference Constant", "Operator", "Value", "Unit", "Step",
               "Step Basis", "Stimulus Points (holds)", "Other Conditions", "Source Line", "Line SHA-256",
               "Source Document", "HW Tolerance (Quoted)"])
    for r in evidence_rows:
        ev.append(r)
    wb.save(path)
    return str(path)


REVIEW = [["요구 문서 값 차이 후보", "SwTSR_0104", "SRS ↔ SyRS SyTSR_0116", "500ms ↔ 300ms", "같은 단위·같은 쪽", "같은 임계인지",
           "—", "—", "u16s_X 500ms 이상", "aa" * 32],
          ["검토 필요 조건", "SwTSR_0101, SwTSR_0102", "SyDS SySM_04 · Verification criteria", "4.85V 미만",
           "값이 무엇의 값인지(주어) 원문에 없음", "신호 정하기", "비교하는 신호 이름을 값 앞에", "—", "Port 를 4.85V 미만", "bb"]]
EVIDENCE = [["SwTC_SwTSR_0104_05", "SwTSR_0104", "threshold", "u16g_ApiIn_Vsup", None, "<", "8.5", "V", "0.01", "written",
             "8.49(T) · 8.5(F)", "—", "입력전원 8.50V 미만", "cc", "SRS",
             "감시 경로 미정 — HwTSR_0204 Battery Voltage Monitor: 허용 오차: ±3% → ±0.255V [공유 SySM_04]; HwTSR_0203 "
             "Hall Sensor Block 전원 Monitor: 허용 오차: ±0.15V → ±0.3V [공유 SySM_04] [척도 환산 — HwC_07 Sensor Power "
             "Switch: 모니터 전압 = Hall Sensor Power *0.5 · HW 설계서 분압식 · HSIS 행과 이름 일치]"],
            ["SwTC_SwTSR_0104_06", "SwTSR_0104", "threshold", "u16s_X", None, ">=", "500", "ms", "1", "written",
             "499 · 500", "—", "500ms 이상", "dd", "SRS", "— (Related ID 로 이어진 HW 허용오차 없음)"]]


def test_the_documents_quote_the_sheets(tmp_path):
    from scripts.sts_findings_to_kb import build_documents
    docs = build_documents(_sts(tmp_path / "s.xlsm", REVIEW, EVIDENCE), "HD")
    assert [d["group"] for d in docs] == ["value_difference", "review_held_back", "evidence", "hw_tolerance"]
    diff, held, ev, hw = docs
    assert diff["title"].startswith("[STS 요구 검토] HD SwTSR_0104 — 요구 문서 값 차이 후보: 500ms ↔ 300ms")
    assert "Source Line: u16s_X 500ms 이상" in diff["content"] and "If Filled" not in diff["content"]   # '—' left out
    assert "aaaa" not in diff["content"]                      # (review I2) the sentence hash is no reading matter
    assert held["tags"][:4] == ["HD", "sts_review", "review_held_back", "SwTSR_0101"]
    assert ev["title"] == "[STS 요구 경계 근거] HD SwTSR_0104" and "HW Tolerance (Quoted): 감시 경로 미정" in ev["content"]
    # (review W9) the HW blocks by their item heads — a bracketed scale source is no measured block — and undecided said
    assert hw["title"] == ("[STS HW 측정 허용오차] HD SwTSR_0104: HwTSR_0204 Battery Voltage Monitor · HwTSR_0203 Hall "
                           "Sensor Block 전원 Monitor (감시 경로 미정)")
    assert "u16g_ApiIn_Vsup < 8.5V (한 눈금 0.01V, SwTC_SwTSR_0104_05) → 감시 경로 미정" in hw["content"]
    assert "500ms" not in hw["content"] and len({d["key"] for d in docs}) == 4


def test_a_long_evidence_is_split_under_the_embedding_cut(tmp_path):
    """(review W7) the embedder keeps 6000 characters: a requirement's long evidence is parts, each whole."""
    from scripts.sts_findings_to_kb import MAX_DOC_CHARS, build_documents
    rows = [[f"SwTC_SwTSR_0104_{i:02d}", "SwTSR_0104", "threshold", "u16g_ApiIn_Vsup", None, "<", "8.5", "V", "0.01",
             "written", "x" * 300, "—", "입력전원 8.50V 미만", "cc", "SRS", "—"] for i in range(40)]
    docs = [d for d in build_documents(_sts(tmp_path / "s.xlsm", [], rows), "HD") if d["group"] == "evidence"]
    assert len(docs) > 1 and all(len(d["content"]) <= MAX_DOC_CHARS for d in docs)
    assert docs[0]["title"].endswith(f"(1/{len(docs)})")


def test_a_file_that_is_not_this_projects_generated_sts_is_refused(tmp_path):
    """(review W3) no findings sheet (a reference STS) or another project's file: refused before anything is deleted."""
    from scripts.sts_findings_to_kb import IngestRefused, build_documents, check_project, main
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "HDPDM01_STS_ref.xlsm")
    with pytest.raises(IngestRefused):
        build_documents(str(tmp_path / "HDPDM01_STS_ref.xlsm"), "HDPDM01")
    assert build_documents(str(tmp_path / "HDPDM01_STS_ref.xlsm"), "HDPDM01", allow_empty=True) == []
    with pytest.raises(IngestRefused):
        check_project(str(tmp_path / "KJPDS02_STS_sy.xlsm"), "HDPDM01")
    check_project(str(tmp_path / "KJPDS02_STS_sy.xlsm"), "KJPDS02_PV")
    assert main(["--sts", str(tmp_path / "HDPDM01_STS_ref.xlsm"), "--project", "hdpdm01", "--dry-run"]) == 2


def test_ingest_is_idempotent_follows_the_content_and_drops_what_is_gone(kb_env):
    from scripts.sts_findings_to_kb import CATEGORY, build_documents, ingest
    tmp, calls = kb_env
    report = tmp / "reports"
    docs = build_documents(_sts(tmp / "s.xlsm", REVIEW, EVIDENCE), "HD")
    out = ingest(docs, "HD", report)
    assert (out["added"], out["kept"], out["removed"], out["embedding_degraded"]) == (4, 0, 0, 0)
    assert (out["verified_in_store"], out["verified_matches"]) == (4, True)
    assert {e["category"] for e in get_kb(report).data} == {CATEGORY} == {"sts_findings"}
    assert all(not e.get("source_file") for e in get_kb(report).data)        # (review W5) no NTFS-stream name
    n_calls = len(calls)
    again = ingest(docs, "HD", report)
    assert (again["added"], again["kept"], again["removed"]) == (0, 4, 0) and len(calls) == n_calls   # no re-embedding
    # (review W2) the same rows, another content (the HW path found): a new entry, the old one goes
    changed = [r[:-1] + ["HwTSR_0204 Battery Voltage Monitor: 허용 오차: ±3% → ±0.255V"] for r in EVIDENCE]
    out = ingest(build_documents(_sts(tmp / "c.xlsm", REVIEW, changed), "HD"), "HD", report)
    assert (out["added"], out["removed"], out["verified_in_store"]) == (2, 2, 4)
    # the value difference fixed in the documents: its entry goes, the others stay
    out = ingest(build_documents(_sts(tmp / "t.xlsm", REVIEW[1:], changed), "HD"), "HD", report)
    assert (out["added"], out["kept"], out["removed"]) == (0, 3, 1)
    assert not any("값 차이" in e["error_raw"] for e in get_kb(report).data)
    # another project's entries are never touched
    other = ingest(build_documents(_sts(tmp / "u.xlsm", REVIEW, []), "PV"), "PV", report)
    assert other["added"] == 2 and len(get_kb(report).data) == 5
    assert ingest([], "HD", report)["removed"] == 3 and len(get_kb(report).data) == 2


def test_a_duplicated_key_is_cleaned(kb_env):
    """(review W6) two entries under one key (two runs at once): the extra one goes."""
    from scripts.sts_findings_to_kb import build_documents, ingest
    tmp, _calls = kb_env
    report = tmp / "reports"
    docs = build_documents(_sts(tmp / "s.xlsm", REVIEW[:1], []), "HD")
    kb = get_kb(report)
    for _ in range(2):
        kb.add_document(docs[0]["title"], docs[0]["content"], category="sts_findings",
                        tags=["HD", "sts_findings:HD", f"sts_key:{docs[0]['key']}"])
    out = ingest(docs, "HD", report)
    assert (out["kept"], out["removed"], out["verified_in_store"]) == (1, 1, 1)


def test_a_non_sqlite_store_is_refused(kb_env, monkeypatch):
    """(review W4) pgvector keeps no entry list in memory: no idempotence, no stale removal — refused."""
    from scripts.sts_findings_to_kb import IngestRefused, ingest
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    monkeypatch.setattr(kb, "storage", "pgvector")
    monkeypatch.setattr("workflow.rag.get_kb", lambda d: kb)
    with pytest.raises(IngestRefused):
        ingest([], "HD", tmp / "reports")


def test_a_dry_run_writes_nothing(kb_env):
    from scripts.sts_findings_to_kb import build_documents, ingest
    tmp, calls = kb_env
    out = ingest(build_documents(_sts(tmp / "s.xlsm", REVIEW, EVIDENCE), "HD"), "HD", tmp / "reports", dry_run=True)
    assert (out["new"], out["added"], out["dry_run"]) == (4, 0, True) and calls == []
    assert get_kb(tmp / "reports").data == []


def test_a_degraded_embedding_is_counted(kb_env):
    """키워드 검색만 되는 엔트리(무작위 폴백 벡터)는 그 수를 보고한다 — 'semantic' 근거로 오독하지 않게."""
    from scripts.sts_findings_to_kb import ingest
    tmp, _calls = kb_env
    doc = {"key": "k1", "title": "t", "content": "degraded text", "tags": ["HD", "sts_review", "x"], "group": "x"}
    assert ingest([doc], "HD", tmp / "reports")["embedding_degraded"] == 1


def _seed_findings(shared):
    kb = get_kb(shared)
    kb.add_document("[STS 요구 검토] HD SwTSR_0104 — 요구 문서 값 차이 후보: 500ms ↔ 300ms", "유지시간 500ms 300ms 문서",
                    category="sts_findings", tags=["HD", "sts_review", "SwTSR_0104"])
    kb.add_document("[STS 요구 검토] HD SwTR_0605 — 검토 필요 조건: 문서 9V", "요구 문서 9V 문서",
                    category="sts_findings", tags=["HD", "sts_review", "SwTR_0605"])
    kb.add_document("UDS 설명", "SwTSR_0104 유지시간 500ms 300ms 다른 문서", category="requirements", tags=["SwTSR_0104"])


def test_findings_are_asked_by_the_ids_the_question_names(kb_env, monkeypatch):
    """(review r2 W-A) only for a requirement / HW ID the question names, only entries carrying it — '문서' in any
    question matched every review title."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    _seed_findings(tmp / "reports")
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    hits = hybrid.requirement_findings_hits("SwTSR_0104 유지시간 500ms 300ms 차이는?")
    assert [h["label"][:30] for h in hits] == ["[STS 요구 검토] HD SwTSR_0104 — 요구"]
    assert hybrid.requirement_findings_hits("API 문서 어디 있어?") == []          # no ID named: nothing
    assert hybrid.requirement_findings_hits("SwXX_9999 는?") == []                # an ID with no finding


def test_the_chat_context_carries_them_first_whatever_the_question_type(kb_env, monkeypatch):
    """(review C1, r2 C1) through the chat's own context builder: a general question (the KB policy gate closed) and a
    docs question (the docs search fills the retrieval slots) both get the findings — first in the evidence list."""
    from backend.services import assistant_service as svc
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    _seed_findings(tmp / "reports")
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    # this venv has no langchain_core (the backend's has): the MCP tools go direct, as the adapter does without it
    monkeypatch.setattr(svc, "get_langchain_mcp_tool_map", lambda: {})
    for question in ("SwTSR_0104 유지시간 500ms 300ms 차이는?", "SwTSR_0104 요구사항 문서 설명해줘"):
        context, sources, citations = svc._build_context(mode="local", question=question,
                                                         report_dir=tmp / "reports" / "sessions" / "t1", ui_context=None)
        assert "requirement_findings" in context, question
        assert sources[0].startswith("requirement_findings:") and citations[0]["label"].startswith("[STS 요구 검토]")
        evidence = svc._build_evidence(citations=citations, sources=sources)
        assert evidence[0]["title"].startswith("[STS 요구 검토] HD SwTSR_0104")
    context, _s, _c = svc._build_context(mode="local", question="API 문서 어디 있어?", report_dir=None, ui_context=None)
    assert "requirement_findings" not in context


def test_a_mass_removal_needs_force_and_a_degraded_vector_is_redone(kb_env):
    """(review r2 W-B) a DV STS given as PV (one file-name token) would replace most of PV's entries — refused without
    --force; (r2 W-C) a random fallback vector is not kept for good: the next run re-embeds it."""
    from scripts.sts_findings_to_kb import IngestRefused, ingest
    tmp, _calls = kb_env
    report = tmp / "reports"
    docs = [{"key": f"k{i}", "title": f"t{i}", "content": f"c{i}", "tags": ["PV", "sts_review", "x"], "group": "x"}
            for i in range(12)]
    ingest(docs, "PV", report)
    with pytest.raises(IngestRefused):
        ingest(docs[:3], "PV", report)
    # (review r3 W3) a dry run shows the plan and the refusal instead of raising
    plan = ingest(docs[:3], "PV", report, dry_run=True)
    assert plan["removals"] == 9 and "0 개는 같은 제목에 내용만 바뀌었고" in plan["would_refuse"]
    assert ingest(docs[:3], "PV", report, allow_mass_removal=True)["removed"] == 9
    bad = {"key": "kd", "title": "td", "content": "degraded text", "tags": ["PV", "sts_review", "x"], "group": "x"}
    assert ingest([*docs[:3], bad], "PV", report)["embedding_degraded"] == 1
    again = ingest([*docs[:3], bad], "PV", report)
    assert (again["added"], again["removed"]) == (1, 1)      # re-embedded, the degraded one replaced


def test_the_cli_exit_codes(kb_env, tmp_path, monkeypatch):
    """(review r2 I-3 · W-C) a missing file is refused (2), a degraded or unverified write is no success (3)."""
    from scripts import sts_findings_to_kb as mod
    tmp, _calls = kb_env
    assert mod.main(["--sts", str(tmp_path / "HDPDM01_none.xlsm"), "--project", "HDPDM01",
                     "--report-dir", str(tmp / "reports")]) == 2
    path = _sts(tmp_path / "HDPDM01_STS.xlsm", [REVIEW[0][:3] + ["degraded 500ms"] + REVIEW[0][4:]], [])
    assert mod.main(["--sts", path, "--project", "HDPDM01", "--report-dir", str(tmp / "reports")]) == 3


def test_an_id_with_a_particle_or_in_lower_case_is_read():
    """(review r3 W1) Python's \\w holds Hangul: ``SwTSR_0104는`` failed \\b — the most natural Korean writing."""
    from workflow.retrieval.hybrid import REQUIREMENT_ID
    for q in ("SwTSR_0104는 값 차이가 왜 나?", "SwTSR_0104의 경계 근거", "swtsr_0104 값 차이", "(HwTSR_0204) 허용오차",
              "SyTSR_0116, 300ms"):
        assert REQUIREMENT_ID.findall(q), q
    assert not REQUIREMENT_ID.findall("xSwTSR_0104 · SwTSR_x")


def test_a_hw_or_system_id_finds_the_entries_that_write_it(kb_env, monkeypatch):
    """(review r3 W2) the tags carry SRS IDs only: an entry is also about the HW / system IDs its text writes."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    kb.add_document("[STS HW 측정 허용오차] HD SwTSR_0104: HwTSR_0204 Battery Voltage Monitor",
                    "u16g_ApiIn_Vsup < 8.5V → HwTSR_0204 Battery Voltage Monitor: 허용 오차: ±3%",
                    category="sts_findings", tags=["HD", "sts_hw_tolerance", "SwTSR_0104", "sts_findings:HD"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    assert [h["label"][:16] for h in hybrid.requirement_findings_hits("HwTSR_0204 허용오차는?")] == ["[STS HW 측정 허용오차]"]


def test_the_project_named_narrows_and_kinds_are_spread(kb_env, monkeypatch):
    """(review r3 W4) an ID's entries first, then ranked; the project the question names narrows them; one per project
    and kind before a second — the value difference asked about is not crowded out by two review items."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    for proj in ("HDPDM01", "KJPDS02_PV"):
        for i in range(3):
            kb.add_document(f"[STS 요구 검토] {proj} SwTSR_0104 — 검토 필요 조건: 값 {i}", f"검토 {i}",
                            category="sts_findings",
                            tags=[proj, "sts_review", "review_held_back", "SwTSR_0104", f"sts_findings:{proj}"])
        kb.add_document(f"[STS 요구 검토] {proj} SwTSR_0104 — 요구 문서 값 차이 후보: 500ms ↔ 300ms", "값 차이",
                        category="sts_findings",
                        tags=[proj, "sts_review", "value_difference", "SwTSR_0104", f"sts_findings:{proj}"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    hits = hybrid.requirement_findings_hits("HDPDM01 SwTSR_0104 값 차이는?")
    assert all("HDPDM01" in h["label"] for h in hits) and any("값 차이" in h["label"] for h in hits)
    both = hybrid.requirement_findings_hits("SwTSR_0104 값 차이는?")
    assert {h["label"].split()[3] for h in both} == {"HDPDM01", "KJPDS02_PV"}


def test_a_degraded_rerun_and_an_allowed_empty_are_not_refused(kb_env, tmp_path):
    """(review r3 W3) the re-embedding rc 3 promised is no mass removal; --allow-empty clears as its help says."""
    from scripts import sts_findings_to_kb as mod
    tmp, _calls = kb_env
    report = tmp / "reports"
    docs = [{"key": f"k{i}", "title": f"degraded t{i}", "content": f"degraded c{i}", "tags": ["PV", "sts_review", "x"],
             "group": "x"} for i in range(12)]
    assert mod.ingest(docs, "PV", report)["embedding_degraded"] == 12
    again = mod.ingest(docs, "PV", report)                  # all twelve redone under their own keys: no refusal
    assert (again["removals"], again["removed"], again["added"]) == (0, 12, 12)
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "PV_STS_empty.xlsm")
    assert mod.main(["--sts", str(tmp_path / "PV_STS_empty.xlsm"), "--project", "PV", "--allow-empty",
                     "--report-dir", str(report)]) == 0
    assert not [e for e in get_kb(report).data if "sts_findings:PV" in (e.get("tags") or [])]


def test_a_question_without_an_id_never_touches_the_kb(monkeypatch):
    """No ID named: no shared-KB load and no query embedding in the request path (review r3 I4 — latency)."""
    from workflow.retrieval import hybrid

    def no_kb(*_a, **_k):
        raise AssertionError("the KB was loaded for a question naming no ID")
    monkeypatch.setattr(hybrid, "get_kb", no_kb)
    notes: list = []
    assert hybrid.requirement_findings_hits("API 문서 어디 있어?", notes_out=notes) == [] and notes == []


def test_the_session_report_hits_skip_the_findings(kb_env, monkeypatch):
    """(review r3 I6) with one KB for all (KB_GLOBAL_DIR) the session's report hits would take the findings again —
    unfiltered by ID; they have their own block."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    kb.add_document("[STS 요구 검토] HD SwTSR_0104 — 값 차이", "유지시간 500ms 300ms", category="sts_findings", tags=["HD"])
    kb.add_document("빌드 실패 로그 요약", "유지시간 500ms 300ms 빌드", category="build", tags=["x"])
    hits = hybrid._report_hits("유지시간 500ms 300ms", tmp / "reports", top_k=5)
    labels = [h["label"] for h in hits]
    assert "sts_findings" not in labels and "build" in labels


def test_the_model_reads_every_finding_the_evidence_list_cites(kb_env, monkeypatch):
    """(review r4 W1) the block's trim kept two or three of four long findings while the evidence list cited four."""
    from backend.services import assistant_service as svc
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    for proj in ("HDPDM01", "KJPDS02_PV"):
        for kind in ("value_difference", "review_held_back"):
            kb.add_document(f"[STS 요구 검토] {proj} SwTSR_0104 — {kind}", f"{kind} 근거 " + "가" * 1100,
                            category="sts_findings",
                            tags=[proj, "sts_review", kind, "SwTSR_0104", f"sts_findings:{proj}"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    monkeypatch.setattr(svc, "get_langchain_mcp_tool_map", lambda: {})
    text, sources, _c = svc._requirement_findings_hints("SwTSR_0104 값 차이는?")
    assert len(sources) == 4 and len(text) <= svc.FINDINGS_BLOCK_CHARS
    assert all(f"REQ#{i} (" in text for i in range(1, 5))
    assert text.count("[STS 요구 검토]") == 4          # the title once (the label), not again from the chunk
    context, _s, _c = svc._build_context(mode="local", question="SwTSR_0104 값 차이는?",
                                         report_dir=tmp / "reports" / "sessions" / "t1", ui_context=None)
    assert "REQ#4 (" in context


def test_allow_empty_keeps_the_mass_removal_guard_on_a_workbook_with_findings(kb_env, tmp_path):
    """(review r4 W2) ``--allow-empty`` lifted the guard for any workbook — another project's STS with sheets then
    replaced this project's entries."""
    from scripts import sts_findings_to_kb as mod
    tmp, _calls = kb_env
    report = tmp / "reports"
    rows = [REVIEW[0][:3] + [f"{500 + i}ms ↔ 300ms"] + REVIEW[0][4:] for i in range(12)]
    full = _sts(tmp_path / "PV_STS_full.xlsm", rows, [])
    assert mod.main(["--sts", full, "--project", "PV", "--report-dir", str(report)]) == 0
    few = _sts(tmp_path / "PV_STS_few.xlsm", rows[:2], [])
    assert mod.main(["--sts", few, "--project", "PV", "--allow-empty", "--report-dir", str(report)]) == 2
    assert mod.main(["--sts", few, "--project", "PV", "--allow-mass-removal", "--report-dir", str(report)]) == 0


def test_equal_scores_break_on_the_kind_asked(kb_env, monkeypatch):
    """(review r4 I1) a keyword-only search scores every candidate 0 — the title's alphabet then put ``[STS HW…``
    before the value difference the question asked about."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    kb.add_document("[STS HW 측정 허용오차] HD SwTSR_0104: Battery Monitor", "허용 오차 ±3%", category="sts_findings",
                    tags=["HD", "sts_hw_tolerance", "SwTSR_0104", "sts_findings:HD"])
    kb.add_document("[STS 요구 검토] HD SwTSR_0104 — 검토 필요 조건: 9V", "주어 없음", category="sts_findings",
                    tags=["HD", "sts_review", "review_held_back", "SwTSR_0104", "sts_findings:HD"])
    kb.add_document("[STS 요구 검토] HD SwTSR_0104 — 요구 문서 값 차이 후보: 500ms ↔ 300ms", "같은 단위",
                    category="sts_findings",
                    tags=["HD", "sts_review", "value_difference", "SwTSR_0104", "sts_findings:HD"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")

    def zero(self, query, top_k=5, category=None, stats_out=None):
        return [{**e, "score": 0.0} for e in self.data if not category or e.get("category") == category]
    monkeypatch.setattr(KnowledgeBase, "search", zero)
    first = hybrid.requirement_findings_hits("SwTSR_0104는 값 차이가 왜 나?", top_k=1)
    assert "값 차이" in first[0]["label"]
    first = hybrid.requirement_findings_hits("SwTSR_0104 허용오차는?", top_k=1)
    assert first[0]["label"].startswith("[STS HW 측정 허용오차]")


def test_an_id_nothing_was_found_about_costs_no_search(kb_env, monkeypatch):
    """(review r4 I2) the candidates are taken before the search — an unknown ID paid a full search and an embedding."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    _seed_findings(tmp / "reports")
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")

    def no_search(*_a, **_k):
        raise AssertionError("searched for an ID with no finding")
    monkeypatch.setattr(KnowledgeBase, "search", no_search)
    notes: list = []
    assert hybrid.requirement_findings_hits("SwXX_9999 는?", notes_out=notes) == [] and notes == []


def test_a_store_without_an_entry_list_says_so(monkeypatch):
    """(review r4 I3) pgvector keeps no entry list in memory — the ingest refuses it, and the chat said nothing."""
    from workflow.retrieval import hybrid

    class _Pg:
        storage = "pgvector"
        data: list = []
    monkeypatch.setattr(hybrid, "get_kb", lambda _d: _Pg())
    notes: list = []
    assert hybrid.requirement_findings_hits("SwTSR_0104 는?", notes_out=notes) == []
    assert notes and "sqlite" in notes[0]


def test_the_report_hits_are_cut_after_the_findings_are_dropped(kb_env):
    """(review r4 I9) with one KB for all the top places were all findings — the reports got none."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    for i in range(20):         # (review r5 I-4) more findings than three times the places
        kb.add_document(f"[STS 요구 검토] HD SwTSR_01{i:02d} — 값 차이", "유지시간 500ms 300ms 유지시간", category="sts_findings",
                        tags=["HD"])
    kb.add_document("빌드 실패 로그 요약", "빌드 로그", category="build", tags=["x"])
    hits = hybrid._report_hits("유지시간 500ms 300ms", tmp / "reports", top_k=3)
    labels = [h["label"] for h in hits]
    assert "build" in labels and "sts_findings" not in labels


def _scored(monkeypatch, scores):
    def search(self, query, top_k=5, category=None, stats_out=None):
        return [{**e, "score": scores(e)} for e in self.data if not category or e.get("category") == category]
    monkeypatch.setattr(KnowledgeBase, "search", search)


def test_the_kind_asked_about_comes_first_whatever_its_score(kb_env, monkeypatch):
    """(review r5 W-A) by score alone '허용오차는?' lost its HW tolerance entry to four review items scoring higher."""
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    for kind in ("review_held_back", "review_read", "inclusion_conflict", "value_difference"):
        kb.add_document(f"[STS 요구 검토] PV SwTSR_0104 — {kind}", "검토", category="sts_findings",
                        tags=["PV", "sts_review", kind, "SwTSR_0104", "sts_findings:PV"])
    kb.add_document("[STS HW 측정 허용오차] PV SwTSR_0104: Battery Monitor", "허용 오차 ±3%", category="sts_findings",
                    tags=["PV", "sts_hw_tolerance", "SwTSR_0104", "sts_findings:PV"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    _scored(monkeypatch, lambda e: 1.0 if "sts_hw_tolerance" in e["tags"] else 2.2)
    labels = [h["label"] for h in hybrid.requirement_findings_hits("PV SwTSR_0104 허용오차는?")]
    assert len(labels) == 4 and labels[0].startswith("[STS HW 측정 허용오차]")
    # (r6 I-1) a word every threshold question writes is the second tier: '허용오차' still comes first, and the
    #   inclusion conflict '초과' asks about is kept though it scores below the other review items
    _scored(monkeypatch, lambda e: 1.0 if "sts_hw_tolerance" in e["tags"] else
            1.5 if "inclusion_conflict" in e["tags"] else 2.2)
    labels = [h["label"] for h in hybrid.requirement_findings_hits("PV SwTSR_0104 8.5V 초과 허용오차는?")]
    assert labels[0].startswith("[STS HW 측정 허용오차]") and "inclusion_conflict" in labels[1]


def test_the_asked_kinds_are_the_ingests_kinds(tmp_path):
    """(review r5 X5 · r6 W-2) the question-word table copies the ingest's kind tags — a renamed tag would lose its
    priority. Bound to the tags the ingest really writes (naming two of them here let their renaming pass)."""
    from scripts.sts_findings_to_kb import _KIND_TAG, build_documents
    from workflow.retrieval.hybrid import _ASKED_KIND, _ASKED_KIND_LOOSE
    docs = build_documents(_sts(tmp_path / "HD_STS.xlsm", REVIEW, EVIDENCE), "HD")
    written = {d["tags"][1] for d in docs} | {d["tags"][2] for d in docs if d["tags"][1] == "sts_review"}
    assert {"sts_evidence", "sts_hw_tolerance"} & written           # the fixture writes both group tags
    assert {k for _w, k in (*_ASKED_KIND, *_ASKED_KIND_LOOSE)} <= written | set(_KIND_TAG.values())


def test_two_findings_with_one_title_stay_two_in_the_evidence_list(kb_env, monkeypatch):
    """(review r5 W-B) the evidence list folds equal labels — three review items differing only in their source block
    were one in the list while the model read three."""
    from backend.services import assistant_service as svc
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    for sid in ("SyII_23", "SyII_24"):
        kb.add_document("[STS 요구 검토] HD SwEI_05 — 검토 필요 조건: 0 ~ 2400",
                        f"프로젝트 HD STS 생성 시 요구 문서에서 찾은 검토 항목.\nKind: 검토 필요 조건\nSource Document: SyDS {sid} · Range",
                        category="sts_findings", tags=["HD", "sts_review", "review_held_back", "SwEI_05", "sts_findings:HD"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    text, sources, citations = svc._requirement_findings_hints("SwEI_05 범위 검토 항목은?")
    evidence = svc._build_evidence(citations=citations, sources=sources)
    assert len(sources) == 2 and len({e["title"] for e in evidence}) == 2
    assert any("SyII_24" in e["title"] for e in evidence)
    assert "프로젝트 HD STS 생성 시" not in text          # (r5 I-2) the embedding preamble is no reading matter


def test_no_more_findings_than_the_budget_holds_readably(kb_env, monkeypatch):
    """(review r5 I-1) at 200 characters a finding is its title only — fewer findings, each readable."""
    from backend.services import assistant_service as svc
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    _seed_findings(tmp / "reports")
    kb = get_kb(tmp / "reports")
    kb.add_document("[STS 요구 검토] HD SwTSR_0104 — 검토 필요 조건: 9V", "주어 없음 " * 80, category="sts_findings",
                    tags=["HD", "sts_review", "review_held_back", "SwTSR_0104", "sts_findings:HD"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    text, sources, _c = svc._requirement_findings_hints("SwTSR_0104 은?", budget=400)
    assert len(sources) == 1 and len(text) <= 400


def test_a_failed_findings_search_is_said_to_the_model(monkeypatch):
    """(review r5 I-7) the failure reached the user's notes only — the model could not tell 'none' from 'could not
    look'."""
    from backend.services import assistant_service as svc
    from workflow.retrieval import hybrid

    def broken(_d):
        raise RuntimeError("db locked")
    monkeypatch.setattr(hybrid, "get_kb", broken)
    text, sources, citations = svc._requirement_findings_hints("SwTSR_0104 값 차이는?")
    assert "볼 수 없었다" in text and "RuntimeError" in text and sources == [] and citations == []
    assert svc._requirement_findings_hints("API 문서 어디 있어?") == ("", [], [])


def test_an_upload_cannot_take_the_findings_category():
    """(review r5 I-5) an uploaded file under 'sts_findings' would be cited first as 'STS 생성 결과'."""
    import asyncio

    from fastapi import HTTPException

    from backend.routers.local import local_rag_ingest_files
    with pytest.raises(HTTPException) as err:
        asyncio.run(local_rag_ingest_files(files=[], category=" STS_Findings ", tags="", report_dir="", chunk_size=None,
                                           chunk_overlap=None, max_chunks=None))
    assert err.value.status_code == 400


def test_two_findings_of_one_source_block_stay_two(kb_env, monkeypatch):
    """(review r6 W-1) one title and one source document (two source lines of one block) folded again — the source
    suffix was the same; the REQ number tells them apart. The citation's snippet is the body (r6 I-2)."""
    from backend.services import assistant_service as svc
    from workflow.retrieval import hybrid
    tmp, _calls = kb_env
    kb = get_kb(tmp / "reports")
    for line in ("0~5V 입력", "0~5V 출력"):
        kb.add_document("[STS 요구 검토] HD SwEI_04 — 검토 필요 조건: 0~5V",
                        f"프로젝트 HD STS 생성 시 요구 문서에서 찾은 검토 항목.\nKind: 검토 필요 조건\n"
                        f"Source Document: SyDS SyII_19 · Range\nSource Line: {line}",
                        category="sts_findings", tags=["HD", "sts_review", "review_held_back", "SwEI_04", "sts_findings:HD"])
    monkeypatch.setattr(hybrid, "_shared_report_dir", lambda: tmp / "reports")
    _text, sources, citations = svc._requirement_findings_hints("SwEI_04 주어가 빠진 곳은?")
    evidence = svc._build_evidence(citations=citations, sources=sources)
    assert len(sources) == 2 and len({e["title"] for e in evidence}) == 2 and len(evidence) == 2
    assert all(c["snippet"].startswith("Kind: 검토 필요 조건") for c in citations)
