# tests/unit/test_docx_logic_diagram_rows.py
"""(R54) Logic Diagram 행 배치가 정본과 다르다(N49) · 템플릿 행 수 하한이 빈 꼬리 행을 남긴다(N50).

## 실측 (2026-09-16, KJPDS02 정본 v3.03 · 라이브 run 2084 산출물 — 함수 정보 표 989개 전수, lxml 직독)

| | 정본 | 생성기(R53 까지) |
|---|---|---|
| Logic Diagram 행 | `[ Logic Diagram ]` 전폭 머리행(1칸·gridSpan 6) + 전폭 그림행(1칸·gridSpan 6·trHeight 2175) — **989/989** | `Logic Diagram` 라벨(2칸) \\| 그림(4칸) **한 행** — 989/989 |
| 표 끝 빈 행 | 0 | **516/989** 표에 빈 라벨\\|값 행 — `max(len(data_rows), 템플릿 행 수)` 하한 (무템플릿 경로는 18) |

같은 전제를 가진 소비자가 넷이라 한 세트로 움직인다(라벨 판정은 `function_analyzer` 한 곳):
· `function_analyzer._build_function_info_layout` — 행 배치
· `docx_builder._insert_logic_image_in_table` / `_infer_function_info_layout` — 그림을 어느 칸에 넣나 · 되짚기
· `validation.validate_uds_docx_structure` — "logic 행 ↔ 이미지" 가 `issues` → `ok` → 품질 게이트 입력(배치만 바꾸면 989/989 거짓 실패)
· `backend.routers.health._extract_docx_sheets` — 미리보기 Functions 시트의 그림 열(라벨 정확 일치라 정본 그림은 원래 못 봤다)
"""
from __future__ import annotations

import logging

import pytest

docx = pytest.importorskip("docx", reason="python-docx 없음")

from docx.oxml.ns import qn  # noqa: E402

from report_gen.docx_builder import (  # noqa: E402
    _add_blank_table,
    _distinct_cells,
    _fill_function_info_table,
    _infer_function_info_layout,
    _insert_logic_image_in_table,
    _is_uniform_grid,
    _merge_function_info_table,
    _normalize_function_info_tables,
    _PictureSink,
)
from report_gen.function_analyzer import (  # noqa: E402
    FN_ROW_FULL,
    FN_ROW_PAIR,
    LOGIC_DIAGRAM_HEADER,
    LOGIC_DIAGRAM_LABEL,
    _build_function_info_layout,
    _function_info_pairs,
    is_logic_diagram_header,
    is_logic_diagram_label,
)
from report_gen.validation import validate_uds_docx_structure  # noqa: E402

_PNG_1x1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da6364f8cfc000000201"
    "0100c9fe92ef0000000049454e44ae426082"
)

# 정본 989/989 실측 모양 — 행마다 (서로 다른 tc 수, [gridSpan…], 그림 유무)
_REF_LOGIC_SHAPE = [(1, [6], False), (1, [6], True)]

# 옛 배치(라벨|값 한 행) — 좁은 표(6열 미만) 폴백이 아직 쓴다. 대괄호 없는 라벨 + 두 칸.
_OLD = [
    (FN_ROW_FULL, ["[ Function Information ]"]),
    (FN_ROW_PAIR, ["ID", "", "SwUFn_0101"]),
    (FN_ROW_PAIR, ["Logic Diagram", "", ""]),
    (FN_ROW_PAIR, ["Called Function", "", "N/A"]),
]


def _info(**over):
    base = {
        "id": "SwUFn_0101", "name": "Motor_Init", "prototype": "void Motor_Init(void)",
        "description": "Init motor", "asil": "B", "related": "SwTR_001",
        "inputs": [], "outputs": [], "globals_global": [], "globals_static": [], "logic": "",
    }
    base.update(over)
    return base


def _layout(cols=6, **over):
    return [(FN_ROW_FULL, ["[ Function Information ]"])] + _build_function_info_layout(_info(**over), cols)


def _table(doc, layout, cols=6):
    """레이아웃을 실제 표로 — 생성 경로와 같은 순서(크기 → 병합 → 채움)."""
    t = _add_blank_table(doc, len(layout), cols, None, None, None)
    _merge_function_info_table(t, cols, layout)
    _fill_function_info_table(t, layout)
    return t


def _row_shape(t):
    return [(len(tr.tc_lst), [tc.grid_span for tc in tr.tc_lst], bool(tr.xpath(".//w:drawing"))) for tr in t._tbl.tr_lst]


def _row_text(tr) -> str:
    return "".join(x.text or "" for x in tr.iter(qn("w:t")))


@pytest.fixture
def png(tmp_path):
    p = tmp_path / "logic.png"
    p.write_bytes(_PNG_1x1)
    return str(p)


# ── 라벨 판정 한 곳 ─────────────────────────────────────────────────────────

class TestLabelJudgement:
    @pytest.mark.parametrize("text", ["Logic Diagram", "[ Logic Diagram ]", "[LogicDiagram]", "  logic   diagram ", "LOGIC DIAGRAM"])
    def test_variants_are_the_same_label(self, text):
        assert is_logic_diagram_label(text)

    @pytest.mark.parametrize("text", ["", None, "Logic Diagrams", "Logic", "[ Input Parameters ]", "Logic Diagram: present",
                                      "see the Logic Diagram below"])
    def test_other_text_is_not(self, text):
        assert not is_logic_diagram_label(text)

    def test_header_is_the_bracketed_form_only(self):
        assert is_logic_diagram_header(LOGIC_DIAGRAM_HEADER) and is_logic_diagram_header(" [LogicDiagram] ")
        assert not is_logic_diagram_header(LOGIC_DIAGRAM_LABEL) and not is_logic_diagram_header("[ Logic Diagram")

    def test_the_pairs_builder_uses_the_shared_label(self):
        assert _function_info_pairs(_info(logic="흐름"))[-1] == (LOGIC_DIAGRAM_LABEL, "흐름")


# ── N49 배치 ────────────────────────────────────────────────────────────────

class TestLayoutMatchesTheReference:
    def test_logic_diagram_is_a_full_width_header_plus_a_full_width_body_row(self):
        layout = _build_function_info_layout(_info(logic="분기 설명"), 6)
        assert layout[-2:] == [(FN_ROW_FULL, [LOGIC_DIAGRAM_HEADER]), (FN_ROW_FULL, ["분기 설명"])]
        assert not any(k == FN_ROW_PAIR and is_logic_diagram_label(c[0]) for k, c in layout)

    def test_empty_logic_text_still_makes_the_body_row(self):
        """그림이 들어갈 자리 — 글이 없어도 행은 있다(정본은 그림만 있다)."""
        assert _build_function_info_layout(_info(), 6)[-2:] == [(FN_ROW_FULL, [LOGIC_DIAGRAM_HEADER]), (FN_ROW_FULL, [""])]

    def test_seven_column_template_gets_the_same_two_rows(self):
        """실측 템플릿 415개 중 124개가 7열 — 폭이 달라도 같은 두 행."""
        assert [k for k, _ in _build_function_info_layout(_info(), 7)[-2:]] == [FN_ROW_FULL, FN_ROW_FULL]

    def test_narrow_table_keeps_the_packed_pair_fallback(self):
        """6열 미만은 그리드가 안 들어가므로 예전 그대로 라벨/값 쌍 — 대괄호 머리행을 만들지 않는다."""
        layout = _build_function_info_layout(_info(), 4)
        assert {k for k, _ in layout} == {FN_ROW_PAIR}
        assert any(LOGIC_DIAGRAM_LABEL in c for _, c in layout) and not any(LOGIC_DIAGRAM_HEADER in c for _, c in layout)

    def test_written_table_has_the_reference_row_shape(self, png):
        """정본 989/989 모양: 머리행 (1칸, [6]) · 그림행 (1칸, [6], 그림) — 그리고 그림행이 마지막 행이다."""
        t = _table(docx.Document(), _layout())
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert _row_shape(t)[-2:] == _REF_LOGIC_SHAPE
        assert _row_text(t._tbl.tr_lst[-2]) == LOGIC_DIAGRAM_HEADER and _row_text(t._tbl.tr_lst[-1]) == ""


class TestPictureGoesIntoTheBodyRow:
    def test_next_row_cell_gets_the_picture_and_keeps_its_tcpr(self, png):
        t = _table(docx.Document(), _layout())
        hdr, body = t._tbl.tr_lst[-2], t._tbl.tr_lst[-1]
        body_tc = body.tc_lst[0]
        tcpr_before = body_tc.tcPr.xml
        assert body_tc.grid_span == 6
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert len(body_tc.xpath(".//w:drawing")) == 1 and not hdr.xpath(".//w:drawing")
        assert body_tc.tcPr.xml == tcpr_before, "(N48) 그림 칸의 tcPr(gridSpan·tcW)가 지워졌다"
        assert len(body_tc.p_lst) == 1
        assert _is_uniform_grid(t, 6) is True
        assert _row_text(hdr) == LOGIC_DIAGRAM_HEADER, "머리행 글이 지워졌다"

    def test_body_text_is_replaced_by_the_picture(self, png):
        """옛 라벨|값 행과 같은 규칙 — 그림이 있으면 `logic` 글 대신 그림."""
        t = _table(docx.Document(), _layout(logic="글"))
        assert _row_text(t._tbl.tr_lst[-1]) == "글"
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert _row_text(t._tbl.tr_lst[-1]) == "" and t._tbl.tr_lst[-1].xpath(".//w:drawing")

    def test_goes_through_the_sink_when_given(self, png):
        d = docx.Document()
        t = _table(d, _layout())
        sink = _PictureSink(d)
        assert _insert_logic_image_in_table(t, 6, png, picture_sink=sink) is True
        assert sink.pictures == 1 and t._tbl.tr_lst[-1].xpath(".//w:drawing")

    def test_old_label_value_row_still_takes_the_value_cell(self, png):
        """좁은 표 폴백(라벨|값 한 행)은 값 칸 그대로 — 대괄호 없는 라벨 + 두 칸."""
        t = _table(docx.Document(), _OLD)
        assert _insert_logic_image_in_table(t, 6, png) is True
        row = t._tbl.tr_lst[2]
        assert [tc.grid_span for tc in row.tc_lst] == [2, 4]
        assert row.tc_lst[1].xpath(".//w:drawing") and not row.tc_lst[0].xpath(".//w:drawing")
        assert not t._tbl.tr_lst[3].xpath(".//w:drawing")

    def test_unmerged_bracketed_header_sends_the_picture_to_the_next_row(self, png):
        """대괄호 표기는 칸이 아직 6개여도 머리행이다 — 병합 전 표에서도 같은 자리."""
        t = docx.Document().add_table(rows=3, cols=6)
        t.cell(1, 0).text = LOGIC_DIAGRAM_HEADER
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert t.cell(2, 0)._tc.xpath(".//w:drawing") and not t.cell(1, 2)._tc.xpath(".//w:drawing")

    def test_merged_single_cell_header_without_brackets_also_uses_the_next_row(self, png):
        """전폭 한 칸이면 '값 칸(3열)' 이 곧 라벨 칸이라 거기 넣으면 라벨이 지워진다 → 다음 행."""
        t = _table(docx.Document(), [(FN_ROW_FULL, ["[ Function Information ]"]), (FN_ROW_FULL, ["Logic Diagram"]), (FN_ROW_FULL, [""])])
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert t._tbl.tr_lst[2].xpath(".//w:drawing") and _row_text(t._tbl.tr_lst[1]) == "Logic Diagram"

    def test_a_label_row_right_after_the_header_is_not_overwritten(self, png, caplog):
        """(리뷰 W3) 다음 행이 그림행 모양이 아니면 지우지 않고 False — 라벨 칸을 그림으로 덮지 않는다."""
        t = _table(docx.Document(), [(FN_ROW_FULL, ["[ Function Information ]"]), (FN_ROW_FULL, [LOGIC_DIAGRAM_HEADER]),
                                     (FN_ROW_PAIR, ["Called Function", "", "Foo_Bar"])])
        before = t._tbl.xml
        with caplog.at_level(logging.WARNING):
            assert _insert_logic_image_in_table(t, 6, png) is False
        assert t._tbl.xml == before and any("그림행 모양" in r.getMessage() for r in caplog.records)

    def test_ragged_next_row_is_addressed_by_row_not_by_grid_index(self, png):
        """(리뷰 W3) 6-4-6 처럼 행 폭이 다른 표 — ragged 행 **뒤의** 행은 그리드 평면 인덱스가 어긋난다(뮤턴트 M27 이 살아남아 표본을 고쳤다:
        머리행 자신이 ragged 여야 다음 행의 평면 인덱스가 두 칸 밀린다). 그림은 다음 행의 **첫** 칸이어야 한다."""
        t = docx.Document().add_table(rows=3, cols=6)
        t.cell(0, 0).text = "[ Function Information ]"
        t.cell(1, 0).text = LOGIC_DIAGRAM_HEADER
        tr1 = t._tbl.tr_lst[1]
        for tc in tr1.tc_lst[4:]:
            tr1.remove(tc)                                  # 머리행(둘째 행)을 4칸으로 — 평면 인덱스 (1+1)*6 은 셋째 행의 3번째 칸을 가리킨다
        tr2 = t._tbl.tr_lst[2]
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert tr2.tc_lst[0].xpath(".//w:drawing") and not any(tc.xpath(".//w:drawing") for tc in tr2.tc_lst[1:])

    def test_header_as_the_last_row_returns_false_and_leaves_the_table_alone(self, png, caplog):
        t = _table(docx.Document(), [(FN_ROW_FULL, ["[ Function Information ]"]), (FN_ROW_FULL, [LOGIC_DIAGRAM_HEADER])])
        before = t._tbl.xml
        with caplog.at_level(logging.WARNING):
            assert _insert_logic_image_in_table(t, 6, png) is False
        assert t._tbl.xml == before
        assert any("그림행이 없다" in r.getMessage() for r in caplog.records)


class TestInferenceAndNormalisation:
    def test_inferred_kinds_match_the_builder_layout(self, png):
        layout = _layout()
        t = _table(docx.Document(), layout)
        assert _insert_logic_image_in_table(t, 6, png) is True
        assert [k for k, _ in _infer_function_info_layout(t)] == [k for k, _ in layout]

    def test_old_row_is_still_a_pair(self):
        t = _table(docx.Document(), _OLD)
        assert [k for k, _ in _infer_function_info_layout(t)][2] == FN_ROW_PAIR

    def test_normalisation_is_a_quiet_byte_no_op(self, png, caplog):
        d = docx.Document()
        t = _table(d, _layout())
        assert _insert_logic_image_in_table(t, 6, png) is True
        before = t._tbl.xml
        with caplog.at_level(logging.WARNING):
            _normalize_function_info_tables(d)
        assert t._tbl.xml == before
        assert not [r for r in caplog.records if "행 폭이 tblGrid" in r.getMessage() or "병합 중단" in r.getMessage()]

    def test_a_label_row_right_after_the_header_is_left_as_a_pair(self):
        """(리뷰 W2) 머리행 다음이 라벨|값 행인 표 — 정규화가 그 행을 전폭으로 접으면 라벨/값 구조가 사라진다."""
        d = docx.Document()
        t = _table(d, [(FN_ROW_FULL, ["[ Function Information ]"]), (FN_ROW_FULL, [LOGIC_DIAGRAM_HEADER]),
                       (FN_ROW_PAIR, ["Called Function", "", "Foo_Bar"])])
        _normalize_function_info_tables(d)
        assert _row_shape(t)[-1] == (2, [2, 4], False)
        assert t.rows[-1].cells[0].text == "Called Function" and t.rows[-1].cells[2].text == "Foo_Bar"

    def test_normalisation_merges_an_unmerged_reference_layout_to_full_width(self):
        """머리행·그림행이 아직 6칸이면 정규화가 둘 다 전폭으로 — 옛 pair 판정이면 두 칸으로 접힌다(음성 대조군 [1, 2, 2])."""
        d = docx.Document()
        t = d.add_table(rows=3, cols=6)
        t.cell(0, 0).text = "Function Information"
        t.cell(1, 0).text = LOGIC_DIAGRAM_HEADER
        _normalize_function_info_tables(d)
        assert [len(tr.tc_lst) for tr in t._tbl.tr_lst] == [1, 1, 1]


class TestLogicTextRoundTrip:
    """(리뷰 W1) `logic` 본문은 docx 왕복에서 살아남아야 한다 — 새 배치는 머리행 다음 전폭 행이 본문이다."""

    @staticmethod
    def _extract(layout, png=None):
        from report_gen.requirements import _extract_function_info_from_docx
        d = docx.Document()
        t = _table(d, layout)
        if png:
            assert _insert_logic_image_in_table(t, 6, png) is True
        return _extract_function_info_from_docx(d)["SwUFn_0101"]

    def test_reference_layout_text_survives(self):
        assert self._extract(_layout(logic="STEP1 -> STEP2")).get("logic") == "STEP1 -> STEP2"

    def test_old_layout_text_survives_too(self):
        old = [(FN_ROW_FULL, ["[ Function Information ]"]), (FN_ROW_PAIR, ["ID", "", "SwUFn_0101"]),
               (FN_ROW_PAIR, ["Logic Diagram", "", "STEP1 -> STEP2"]), (FN_ROW_PAIR, ["Called Function", "", "N/A"])]
        assert self._extract(old).get("logic") == "STEP1 -> STEP2"

    def test_picture_row_reads_back_as_empty(self, png):
        block = self._extract(_layout(logic="글"), png)       # 그림이 글을 대체했다 — 되읽기는 빈 값, 다른 축은 그대로
        assert not str(block.get("logic") or "").strip()
        assert block.get("name") == "Motor_Init" and str(block.get("asil")).upper() == "B"

    def test_a_label_row_after_the_header_is_not_read_as_logic(self):
        block = self._extract([(FN_ROW_FULL, ["[ Function Information ]"]), (FN_ROW_PAIR, ["ID", "", "SwUFn_0101"]),
                               (FN_ROW_FULL, [LOGIC_DIAGRAM_HEADER]), (FN_ROW_PAIR, ["Called Function", "", "Foo_Bar"])])
        assert not str(block.get("logic") or "").strip() and "Foo_Bar" in str(block.get("called") or "")


# ── N50 빈 꼬리 행 ──────────────────────────────────────────────────────────

class TestNoTrailingBlankRows:
    """표의 행 수는 데이터가 정한다 — 두 빌더(템플릿/무템플릿) 모두. 실제 생성 경로로 문서를 만들어 표를 되읽는다."""

    @pytest.fixture(autouse=True)
    def _no_reference_suds(self, monkeypatch, tmp_path):
        """저장소 고정 참조 SUDS(40MB)·기본 템플릿(430 heading)을 읽지 않는다(`test_uds_docx_gen_stats` 와 같은 이유·같은 방법)."""
        import config
        monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(tmp_path / "no_such_reference.docx"), raising=False)
        monkeypatch.setattr(config, "resolve_uds_template_path", lambda: "", raising=False)

    @staticmethod
    def _payload(names):
        fd = {}
        for i, n in enumerate(names, start=1):
            fd[f"SwUFn_{i:04d}"] = {
                "id": f"SwUFn_{i:04d}", "name": n, "prototype": f"void {n}(void);", "description": "설명", "asil": "QM",
                "related": "", "inputs": [], "outputs": [], "precondition": "", "globals_global": [], "globals_static": [],
                "called": "", "logic": "",
            }
        return {"project_name": "T", "overview": "o", "requirements": "r", "interfaces": "i", "uds_frames": "u", "notes": "n",
                "function_details": fd}

    @staticmethod
    def _function_tables(path):
        d = docx.Document(str(path))
        return [t for t in d.tables if t.rows and "Function Information" in t.rows[0].cells[0].text]

    @staticmethod
    def _assert_tight(tables, expected_count):
        assert len(tables) == expected_count, f"함수 정보 표 {len(tables)}개 — 표본이 잘못됐다"
        for t in tables:
            rows = list(t.rows)
            texts = [[c.text.strip() for c in _distinct_cells(r)] for r in rows]
            hdr = [i for i, cells in enumerate(texts) if any(is_logic_diagram_header(c) for c in cells)]
            assert len(hdr) == 1, texts
            assert hdr[0] == len(rows) - 2, f"그림행이 마지막 행이 아니다 — 꼬리에 {len(rows) - 2 - hdr[0]}행이 남았다"
            empty = [i for i, cells in enumerate(texts) if all(c == "" for c in cells)]
            assert empty in ([], [len(rows) - 1]), f"빈 행 {empty} (행 {len(rows)}개) — 그림행만 글이 없을 수 있다"

    def test_no_template_branch(self, tmp_path):
        from report_gen.docx_builder import generate_uds_docx
        out = tmp_path / "u.docx"
        generate_uds_docx(None, self._payload(["alpha", "beta"]), str(out))
        self._assert_tight(self._function_tables(out), 2)

    def test_template_branch_with_an_oversized_template_table(self, tmp_path):
        """템플릿 표가 30행이어도 우리 표는 데이터 행 수 — 예전엔 30행 하한이라 빈 꼬리 행이 남았다(라이브 516/989)."""
        from report_gen.docx_builder import generate_uds_docx
        tpl = docx.Document()
        tpl.add_heading("Software Unit Design", level=1)
        for i, n in enumerate(("alpha", "beta"), start=1):
            tpl.add_heading(f"SwUFn_{i:04d}: {n}", level=4)
            t = tpl.add_table(rows=30, cols=6)
            t.cell(0, 0).text = "[ Function Information ]"
        tpl_path = tmp_path / "tpl.docx"
        tpl.save(str(tpl_path))
        out = tmp_path / "u.docx"
        generate_uds_docx(str(tpl_path), self._payload(["alpha", "beta"]), str(out))
        tables = self._function_tables(out)
        self._assert_tight(tables, 2)
        assert all(len(t.columns) == 6 and len(t.rows) < 30 for t in tables), "템플릿 표(6열 30행)를 채택했다는 전제가 서지 않는다(리뷰 I8)"


# ── 검증기(게이트 입력) ─────────────────────────────────────────────────────

class TestValidatorFindsThePictureInBothLayouts:
    @staticmethod
    def _doc_with(tmp_path, kinds):
        """kinds: 'ref'(정본 배치+그림) / 'old'(라벨|그림 한 행) / 'ref_nopic'(머리행만, 그림 없음) / 'mention'(설명 칸에 그 말만)."""
        d = docx.Document()
        png = tmp_path / "p.png"
        png.write_bytes(_PNG_1x1)
        for i, kind in enumerate(kinds, start=1):
            d.add_heading(f"SwUFn_{i:04d}: f{i}", level=3)
            if kind == "old":
                t = _table(d, _OLD)
            elif kind == "mention":
                t = _table(d, [(FN_ROW_FULL, ["[ Function Information ]"]), (FN_ROW_PAIR, ["ID", "", f"SwUFn_{i:04d}"]),
                               (FN_ROW_PAIR, ["Description", "", "see the Logic Diagram of the caller"])])
                continue
            else:
                t = _table(d, _layout())
            if kind != "ref_nopic":
                assert _insert_logic_image_in_table(t, 6, str(png)) is True
        p = tmp_path / "v.docx"
        d.save(str(p))
        return str(p)

    def test_reference_and_old_layouts_both_count_as_rows_with_image(self, tmp_path):
        r = validate_uds_docx_structure(self._doc_with(tmp_path, ["ref", "old", "ref"]))
        assert (r["logic_row_count"], r["logic_with_image_count"]) == (3, 3)
        assert not [i for i in r["issues"] if "Logic rows" in i]

    def test_a_header_without_a_picture_is_still_reported(self, tmp_path):
        """음성 대조군 — 배치를 받아들이느라 검사를 죽이지 않았다."""
        r = validate_uds_docx_structure(self._doc_with(tmp_path, ["ref", "ref_nopic"]))
        assert (r["logic_row_count"], r["logic_with_image_count"]) == (2, 1)
        assert any("Logic rows(2) != rows with image(1)" in i for i in r["issues"]) and r["ok"] is False

    def test_the_picture_row_and_a_description_mentioning_the_words_are_not_logic_rows(self, tmp_path):
        """그림행(글 없음)과 문장 속 언급은 logic 행이 아니다 — 예전 부분문자열 판정은 후자를 '이미지 없는 logic 행' 으로 셌다."""
        r = validate_uds_docx_structure(self._doc_with(tmp_path, ["ref", "mention"]))
        assert (r["logic_row_count"], r["logic_with_image_count"]) == (1, 1)


# ── 미리보기 리더 ───────────────────────────────────────────────────────────

class TestPreviewReaderFindsThePicture:
    @staticmethod
    def _functions_sheet(tmp_path, kind):
        from backend.routers.health import _extract_docx_sheets
        d = docx.Document()
        png = tmp_path / "p.png"
        png.write_bytes(_PNG_1x1)
        t = _table(d, _OLD if kind == "old" else _layout())
        if kind != "ref_nopic":
            assert _insert_logic_image_in_table(t, 6, str(png)) is True
        return next(s for s in _extract_docx_sheets(d) if s["name"].startswith("Functions"))

    def test_reference_layout_row_has_the_image_column(self, tmp_path):
        fn = self._functions_sheet(tmp_path, "ref")
        assert fn["headers"][-1] == "Logic Diagram" and fn["rows"][0][-1].startswith("__IMG__rId")

    def test_old_layout_still_has_it(self, tmp_path):
        fn = self._functions_sheet(tmp_path, "old")
        assert fn["headers"][-1] == "Logic Diagram" and fn["rows"][0][-1].startswith("__IMG__rId")

    def test_no_picture_means_no_image_column(self, tmp_path):
        fn = self._functions_sheet(tmp_path, "ref_nopic")
        assert "Logic Diagram" not in fn["headers"]


# ── 판정 한 곳 ──────────────────────────────────────────────────────────────

class TestOneJudgementEverywhere:
    """라벨 판정은 `function_analyzer` 한 곳 — 소비자 셋이 그걸 쓴다(복제본 셋 중 하나만 고침 방지)."""

    def test_consumers_import_the_shared_judgement(self):
        import ast

        from backend.routers import health
        from report_gen import docx_builder, validation
        from tests.unit._source_probe import source_of
        for mod in (docx_builder, validation, health):
            names = {a.name for n in ast.walk(ast.parse(source_of(mod)))
                     if isinstance(n, ast.ImportFrom) and n.module == "report_gen.function_analyzer" for a in n.names}
            assert {"is_logic_diagram_label", "is_logic_diagram_header"} <= names, (mod.__name__, names)   # (리뷰 I6) 주석이 아니라 import 문
        assert 'c.replace(" ", "") == "LogicDiagram"' not in source_of(docx_builder)
        assert 'label == "Logic Diagram"' not in source_of(health)
        assert '"Logic Diagram" in c' not in source_of(validation)
