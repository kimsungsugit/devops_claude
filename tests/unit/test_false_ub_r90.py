"""R90 — false "undefined behaviour" refusals of the source oracle (`generators/c_source_oracle.py`).

HDPDM01 SITS (r89c) had 3,612 expected cells refused as ``undefined_behavior:*``. Three causes were not undefined behaviour
of the program:

(1) ``sizeof`` was never evaluated — ``for (i = 0u; i < sizeof(p) / sizeof(UINT8); i++) … p[i]`` (`sf_Crc32Init`) was a
    branch on an unknown value every pass, and the pass past the array's length indexed past it: a proven out-of-bounds
    index on a path the program never takes. A character type is 1 and a character array is its length whatever the
    implementation pads (C11 6.5.3.4p4); the type is ``size_t``, whose rank the source does not say — only results exact
    in every such type are known.
(2) Undefined behaviour proven on one path refused every other path, even when a branch on an unknown value chose that
    path (`lin_process_uncd_frame`: ``for (i = 1; i < len + 1; i++) buf[i] = …`` with ``len`` read from a table the model
    does not hold). On such a path it is possible undefined behaviour: disclosed with the branch that chose it, the
    faulting result unknown, the path runs on (review C1: never dropped — dropping it took "no UB" as evidence about the
    selector). Refused when the path is the run's only one or every final path proved it. A source finding stays.
(3) Behind (1): ``crc = tbl[(crc ^ *p++) & 0xFFu] ^ (crc >> 8u)`` (`sf_Crc32Cal`) was refused as unsequenced — ``p++``
    writes the pointer object, not through it.

Genuine undefined behaviour still refuses: a proven one on the only path or on every path (``(U8)191 << 8`` on a 16-bit
int, ``~1 << 7``, a divisor of zero the inputs set, an index the inputs set past the array).

(4) review W4: an enumeration input may be any member (0xB2) as well as 0..127 — the type holds them all.
"""
from __future__ import annotations

import os

from generators import c_project_context as cpc
from generators.c_source_oracle import evaluate_outputs, observe_decisions
from generators.integration_oracle import CalleeProvider
from generators.source_findings import collect_findings, write_source_findings_sheet
from generators.test_evidence import apply_sequence_evidence, summarize_expected_evidence

ROOT = os.path.join(os.sep, "proj")
COMMON = """#ifndef COMMON_H
#define COMMON_H
typedef unsigned char U8;
typedef unsigned char UINT8;
typedef unsigned int U16;
typedef unsigned long U32;
typedef unsigned long UINT32;
typedef signed int S16;
typedef signed long S32;
#endif
"""
H = '#include "common.h"\n'


def _p(name):
    return os.path.join(ROOT, name)


def _unit(text, name, extra=None, provider=False):
    files = {_p("common.h"): COMMON, _p("unit.c"): text}
    files.update({_p(k): v for k, v in (extra or {}).items()})
    context = cpc.build_project_context(files)
    scopes = cpc.build_scopes(context, [p for p in files if p.endswith(".c")])
    unit = {"name": name, "source_text": text, "source_path": _p("unit.c"), "source_text_complete": True,
            "project_scope": scopes[_p("unit.c")]}
    if provider:
        unit["callee_provider"] = CalleeProvider(context, files, scopes, cpc.shared_parser())
    return unit


def _run(text, name, inputs, outputs, **kw):
    return evaluate_outputs(_unit(text, name, **kw), [inputs], [outputs])[0]


def _values(result):
    return {k: v.get("value", "?" + v.get("reason", "")) for k, v in result["outputs"].items()}


# ── (1) sizeof ──────────────────────────────────────────────────────────────────────────────────

CRC_INIT = H + """U32 g_poly; U32 g_n;
static const U8 g_terms[] = {0u, 1u, 2u, 4u, 5u, 7u, 8u, 10u, 11u, 12u, 16u, 22u, 23u, 26u};
void crc_init(void)
{
    UINT32 i;
    static const UINT8 p[] = {0u, 1u, 2u, 4u, 5u, 7u, 8u, 10u, 11u, 12u, 16u, 22u, 23u, 26u};
    g_n = 0uL;
    for (i = 0u; i < ((sizeof (p)) / (sizeof (UINT8))); i++)
    {
        g_n = g_n + 1uL;
    }
    g_poly = 0uL;
    for (i = 0u; i < ((sizeof (g_terms)) / (sizeof (UINT8))); i++)
    {
        g_poly |= 1uL << (31u - g_terms[i]);
    }
}
"""


def test_sizeof_of_a_character_array_bounds_the_loop_it_controls():
    # HDPDM01 sf_Crc32Init: before R90 the bound was unknown, the pass with i == 14 indexed past ``p`` and the whole
    # run was refused as array_index_out_of_bounds
    r = _run(CRC_INIT, "crc_init", {}, ["g_n", "g_poly"])
    assert r["status"] == "supported", r["reason"]
    assert _values(r) == {"g_n": 14, "g_poly": 0xEDB88320}
    assert not r["possible_undefined_behavior"]


def test_sizeof_values_c_fixes_and_the_ones_it_does_not():
    text = H + """U8 g_buf[300]; U16 g_w; U8 g_c; U16 g_a; U16 g_b; U16 g_d; U16 g_e; U16 g_f; U16 g_g; U16 g_h; S16 g_s;
U8 g_one;
void f(U8 prm[4], U8 c)
{
    U8 loc[6];
    g_w = sizeof(g_buf);                 /* 300 */
    g_c = (U8)sizeof(g_buf);             /* converts by its value: 44 */
    g_a = sizeof(U8) + sizeof(unsigned char) + sizeof(char) + sizeof c + sizeof(g_buf[3]);   /* 5 */
    g_b = sizeof(loc) / sizeof(loc[0]);  /* 6 */
    g_d = sizeof(U16);                   /* padding bits are the implementation's: unknown */
    g_e = sizeof(prm);                   /* a pointer */
    g_f = sizeof(U8 *);                  /* a pointer */
    g_s = (S16)(sizeof(loc) - 7);        /* wraps in size_t, whose width the source does not say */
    g_g = (sizeof(loc) > -1);            /* -1 converts by size_t's rank: unknown */
    g_h = (sizeof(loc) == 6U) && (sizeof(g_buf) >= 300);
    g_one = 0U;
    if (sizeof(loc)) { g_one = 1U; }
}
"""
    v = _values(_run(text, "f", {"c": 1}, ["g_w", "g_c", "g_a", "g_b", "g_d", "g_e", "g_f", "g_s", "g_g", "g_h",
                                           "g_one"]))
    assert (v["g_w"], v["g_c"], v["g_a"], v["g_b"], v["g_h"], v["g_one"]) == (300, 44, 5, 6, 1, 1)
    for name in ("g_d", "g_e", "g_f", "g_s", "g_g"):
        assert str(v[name]).startswith("?"), (name, v[name])


def test_a_sizeof_value_is_not_kept_by_an_operator_that_does_not_know_it():
    # ``-sizeof`` is size_t arithmetic (wraps by its width): the unary operator drops the value — never read back as 3
    text = H + "U8 g_b[3]; U16 g_x;\nvoid f(void) { g_x = (U16)(-sizeof(g_b) + 6U); }\n"
    assert str(_values(_run(text, "f", {}, ["g_x"]))["g_x"]).startswith("?")


# ── (2) undefined behaviour on a path a branch on an unknown value chose ─────────────────────────

UNCD = H + """U8 g_buf[10]; U8 g_len; U8 g_after; U8 g_src[10];
void make(void)
{
    U8 i;
    for (i = 1U; (i < (U8)(g_len + 1U)) && (i < 12U); i++)
    {
        g_buf[i] = g_src[i];
    }
    g_after = 5U;
}
"""


def test_index_past_the_array_on_a_path_an_unknown_bound_chose_is_possible_ub_and_the_path_runs_on():
    # HDPDM01 lin_process_uncd_frame: the passes where the unknown length is 10 or more write past g_buf
    r = _run(UNCD, "make", {}, ["g_after", "g_buf[3]"])
    assert r["status"] == "supported"
    v = _values(r)
    assert v["g_after"] == 5   # every path writes it, the faulting ones too
    assert str(v["g_buf[3]"]).startswith("?")   # the write past the array leaves the whole array unknown
    tag = "array_index_out_of_bounds_on_unknown_branch(selector=initial_value_not_in_inputs:g_len)"
    assert tag in r["possible_undefined_behavior"] and "possible_ub=" in r["reason"]
    assert r["forked_path_undefined_behavior"] == [
        {"kind": "array_index_out_of_bounds", "selector": "initial_value_not_in_inputs:g_len"}]
    assert any(a.startswith("possible_ub_on_forked_path:") for a in r["assumptions"])
    # the length known: the loop stays inside the array, nothing to disclose
    known = _run(UNCD, "make", {"g_len": 3, "g_src[1]": 7, "g_src[2]": 8, "g_src[3]": 9}, ["g_after", "g_buf[3]"])
    assert _values(known) == {"g_after": 5, "g_buf[3]": 9} and not known["possible_undefined_behavior"]
    assert "forked_path_undefined_behavior" not in known


PARADOX = H + """U8 g_mode; U8 g_tbl[4]; U8 g_out; volatile U8 g_reg; U8 get_mode(void);
void safe_g(void) { if (g_mode == 1U) { g_out = g_tbl[1]; } else { g_out = 2U; } }
void ub_g(void)   { if (g_mode == 1U) { g_out = g_tbl[5]; } else { g_out = 2U; } }
void safe_v(void) { if (g_reg == 1U) { g_out = g_tbl[1]; } else { g_out = 2U; } }
void ub_v(void)   { if (g_reg == 1U) { g_out = g_tbl[5]; } else { g_out = 2U; } }
void safe_p(U8 mode) { if (mode == 1U) { g_out = g_tbl[1]; } else { g_out = 2U; } }
void ub_p(U8 mode)   { if (mode == 1U) { g_out = g_tbl[5]; } else { g_out = 2U; } }
void safe_c(void) { if (get_mode() == 1U) { g_out = g_tbl[1]; } else { g_out = 2U; } }
void ub_c(void)   { if (get_mode() == 1U) { g_out = g_tbl[5]; } else { g_out = 2U; } }
"""


def test_undefined_behaviour_never_makes_a_value_more_certain_than_without_it():
    # review C1: dropping the faulting path took "the program has no UB" as evidence about the selector — ``g_out = 2``
    #   derived with g_tbl[5] where g_tbl[1] in its place leaves it unknown. Selector: a global the row leaves open, a
    #   volatile register, a parameter the row does not set, a callee's return value
    for sel in ("g", "v", "p", "c"):
        for inputs in ({}, {"g_tbl[1]": 2}, {"g_tbl[1]": 7}):
            safe = _run(PARADOX, "safe_" + sel, inputs, ["g_out"])
            ub = _run(PARADOX, "ub_" + sel, inputs, ["g_out"])
            assert ub["status"] == "supported", (sel, ub["reason"])
            got = ub["outputs"]["g_out"]
            assert "value" not in got, (sel, inputs, got)
            assert got["reason"] == "possible_undefined_behavior_on_unknown_branch:array_index_out_of_bounds"
            if "value" in safe["outputs"]["g_out"]:
                assert safe["outputs"]["g_out"]["value"] == 2 and inputs == {"g_tbl[1]": 2}
    r = _run(PARADOX, "ub_v", {}, ["g_out"])
    assert r["forked_path_undefined_behavior"] == [{"kind": "array_index_out_of_bounds",
                                                    "selector": "volatile_object:g_reg"}]


DIV = H + """U16 g_u; U16 g_zero; U16 g_x; U16 g_y; U16 g_z;
void f(void)
{
    if (g_u > 3U) { g_x = 10U / g_zero; } else { g_y = 1U; }
    g_z = 7U;
}
void every(void)
{
    if (g_u > 3U) { g_x = 10U / g_zero; } else { g_y = 20U / g_zero; }
    g_z = 7U;
}
void only(void) { g_x = 10U / g_zero; g_z = 7U; }
void cond(void)
{
    if (g_u > 3U) { if ((10U / g_zero) > 1U) { g_x = 1U; } }
    else { switch (20U / g_zero) { case 1U: g_x = 2U; break; default: break; } }
    if (g_u > 9U) { g_y = 3U; }
    g_z = 7U;
}
void upd(void)
{
    U16 i;
    for (i = 0U; (i < g_u) && (i < 3U); i++, g_y = 10U / g_zero) { g_x = i; }
    g_z = 7U;
}
"""
DIV_TAG = "division_by_zero_on_unknown_branch(selector=initial_value_not_in_inputs:g_u)"


def test_division_by_zero_on_one_arm_of_an_unknown_branch_is_disclosed_and_its_result_unknown():
    r = _run(DIV, "f", {"g_zero": 0}, ["g_z", "g_y", "g_x"])
    assert r["status"] == "supported"
    v = _values(r)
    assert v["g_z"] == 7   # both paths write it
    # the faulting path runs on: g_y keeps its unset initial value there — not "1" (that was the dropped-path claim)
    assert v["g_y"] == "?unwritten_output_initial_not_in_inputs:g_y"
    assert str(v["g_x"]).startswith("?possible_undefined_behavior_on_unknown_branch:division_by_zero")
    assert r["possible_undefined_behavior"] == [DIV_TAG]
    assert any("chose the path it was proven on" in a for a in r["assumptions"])


def test_undefined_behaviour_in_a_branch_or_switch_condition_is_judged_on_its_path():
    # g_u > 3: the inner ``if`` condition divides by zero; g_u <= 3: the ``switch`` controlling expression does — every
    #   final path proved it, so no run is without it
    assert _run(DIV, "cond", {"g_zero": 0}, ["g_z"])["reason"] == "undefined_behavior:division_by_zero"
    r = _run(DIV, "cond", {"g_zero": 1}, ["g_z", "g_x"])
    assert r["status"] == "supported" and _values(r)["g_z"] == 7 and not r["possible_undefined_behavior"]
    text = DIV.replace("switch (20U / g_zero)", "switch (g_u)")
    r = _run(text, "cond", {"g_zero": 0}, ["g_z", "g_x"])
    assert r["status"] == "supported" and _values(r)["g_z"] == 7 and str(_values(r)["g_x"]).startswith("?")
    assert r["possible_undefined_behavior"] == [DIV_TAG]
    text = DIV.replace("if ((10U / g_zero) > 1U)", "if (g_u > 5U)")
    r = _run(text, "cond", {"g_zero": 0}, ["g_z"])
    assert r["status"] == "supported" and _values(r) == {"g_z": 7}
    assert r["possible_undefined_behavior"] == [DIV_TAG]


def test_a_switch_on_an_unknown_value_chooses_the_path_like_an_if():
    # review W3 M9: a case arm a switch on an open input chose is a forked path — not the run's only one
    text = H + """U16 g_u; U16 g_zero; U16 g_x; U16 g_y; U16 g_z;
void sw(void) { switch (g_u) { case 1U: g_x = 10U / g_zero; break; default: g_y = 1U; break; } g_z = 7U; }
"""
    r = _run(text, "sw", {"g_zero": 0}, ["g_z"])
    assert r["status"] == "supported" and _values(r) == {"g_z": 7}
    assert r["possible_undefined_behavior"] == [DIV_TAG]
    assert _run(text, "sw", {"g_zero": 0, "g_u": 1}, ["g_z"])["reason"] == "undefined_behavior:division_by_zero"


def test_undefined_behaviour_in_a_loop_update_is_judged_on_the_paths_that_reach_it():
    r = _run(DIV, "upd", {"g_zero": 0}, ["g_z", "g_y"])   # the path that leaves before the first pass has none
    assert r["status"] == "supported" and _values(r)["g_z"] == 7 and str(_values(r)["g_y"]).startswith("?")
    assert r["possible_undefined_behavior"] == [DIV_TAG]
    r = _run(DIV, "upd", {"g_zero": 0, "g_u": 0}, ["g_z"])
    assert _values(r) == {"g_z": 7} and not r["possible_undefined_behavior"]
    assert _run(DIV, "upd", {"g_zero": 0, "g_u": 5}, ["g_z"])["reason"] == "undefined_behavior:division_by_zero"


def test_undefined_behaviour_inside_a_statement_macro_is_judged_by_the_path_that_ran_it():
    text = H + """#define DO_DIV() do { g_x = 10U / g_zero; } while (0)
U16 g_u; U16 g_zero; U16 g_x; U16 g_z;
void f(void) { if (g_u > 3U) { DO_DIV(); } g_z = 7U; }
void g(void) { DO_DIV(); g_z = 7U; }
"""
    r = _run(text, "f", {"g_zero": 0}, ["g_z"])
    assert r["status"] == "supported" and _values(r) == {"g_z": 7}
    assert r["possible_undefined_behavior"] == [DIV_TAG]
    assert _run(text, "g", {"g_zero": 0}, ["g_z"])["reason"] == "undefined_behavior:division_by_zero"


def test_undefined_behaviour_on_every_path_or_the_only_path_still_refuses():
    assert _run(DIV, "every", {"g_zero": 0}, ["g_z"])["reason"] == "undefined_behavior:division_by_zero"
    assert _run(DIV, "only", {"g_zero": 0}, ["g_z"])["reason"] == "undefined_behavior:division_by_zero"
    # a known condition selects the arm: proven on the run itself
    assert _run(DIV, "f", {"g_zero": 0, "g_u": 9}, ["g_z"])["reason"] == "undefined_behavior:division_by_zero"
    assert _values(_run(DIV, "f", {"g_zero": 0, "g_u": 1}, ["g_z", "g_y"])) == {"g_z": 7, "g_y": 1}


def test_genuine_shift_undefined_behaviour_on_a_16_bit_int_still_refuses():
    # HDPDM01 lin_update_word_status_lin21 ``pid << 8`` (pid an U8 >= 128) and lin_process_parity ``(~(…)) << 7``
    text = H + "U16 g_w;\nvoid f(U8 pid) { g_w = (U16)(pid << 8); }\n"
    assert _run(text, "f", {"pid": 191}, ["g_w"])["reason"] == "undefined_behavior:signed_left_shift_overflow"
    assert _values(_run(text, "f", {"pid": 100}, ["g_w"])) == {"g_w": 25600}
    text = H + "U8 g_k;\nvoid g(U8 b) { g_k = (U8)((~((S16)b & 1)) << 7); }\n"
    assert _run(text, "g", {"b": 0}, ["g_k"])["reason"] == "undefined_behavior:signed_left_shift_overflow"


def test_a_callee_path_with_undefined_behaviour_runs_on_and_a_callee_with_it_on_every_path_marks_the_callers():
    callee = H + """U16 g_zero; U16 g_u; U16 g_x; U16 g_y;
void maybe(void) { if (g_u > 3U) { g_x = 10U / g_zero; } else { g_y = 1U; } }
void always(void) { if (g_u > 3U) { g_x = 10U / g_zero; } else { g_x = 20U / g_zero; } }
"""
    text = H + """U16 g_x; U16 g_y; U16 g_z; U16 g_k;
void maybe(void);
void always(void);
void entry(void) { maybe(); g_z = 7U; }
void entry2(U16 k) { g_k = 1U; if (k > 2U) { always(); } g_z = 7U; }
void entry3(void) { always(); g_z = 7U; }
"""
    extra = {"callee.c": callee}
    r = _run(text, "entry", {"g_zero": 0}, ["g_z", "g_y", "g_x"], extra=extra, provider=True)
    assert r["status"] == "supported" and _values(r)["g_z"] == 7
    assert str(_values(r)["g_y"]).startswith("?") and str(_values(r)["g_x"]).startswith("?")
    assert DIV_TAG in r["possible_undefined_behavior"]
    assert r["interprocedural"]["inlined"] == {"maybe": 1}
    # every path of ``always`` divides by zero: the caller's path that called it proved it — the other stands
    r = _run(text, "entry2", {"g_zero": 0}, ["g_z", "g_k"], extra=extra, provider=True)
    assert r["status"] == "supported" and _values(r) == {"g_z": 7, "g_k": 1}
    # (review W1) every branch that chose the path: the caller's on ``k`` (the one that really chose — the callee faults on
    #   every path) and the callee's own on g_u
    assert r["possible_undefined_behavior"] == [
        "division_by_zero_on_unknown_branch(selector=parameter_not_in_inputs:k · initial_value_not_in_inputs:g_u)"]
    # ... and on the caller's only path it is the run's: refused
    assert _run(text, "entry2", {"g_zero": 0, "k": 5}, ["g_z"], extra=extra, provider=True)["reason"] == \
        "undefined_behavior:division_by_zero"
    assert _run(text, "entry3", {"g_zero": 0}, ["g_z"], extra=extra, provider=True)["reason"] == \
        "undefined_behavior:division_by_zero"


def test_undefined_behaviour_on_the_callers_only_path_after_a_forking_callee_refuses():
    # review W3 M1: the callee's branch on g_u leaves two paths; the caller's division after it runs on both — no run
    #   without it
    callee = H + """U16 g_zero; U16 g_u; U16 g_x; U16 g_y;
void maybe(void) { if (g_u > 3U) { g_x = 10U / g_zero; } else { g_y = 1U; } }
"""
    text = H + """U16 g_z; U16 g_zero;
void maybe(void);
void entry4(void) { maybe(); g_z = 10U / g_zero; }
"""
    r = _run(text, "entry4", {"g_zero": 0}, ["g_z"], extra={"callee.c": callee}, provider=True)
    assert (r["status"], r["reason"]) == ("unsupported", "undefined_behavior:division_by_zero")


def test_every_branch_that_chose_the_path_is_named_and_an_open_input_is_claimed_only_when_all_are():
    # review W1: a register chose the path, then an input the row leaves open — naming only the last called it "an input
    #   the test did not set", but setting g_x0 leaves the undefined behaviour where it is
    text = H + """volatile U8 g_reg; U16 g_x0; U16 g_zero; U16 g_out; U16 g_z;
void w1(void) { if (g_reg > 0U) { if (g_x0 > 3U) { g_out = 10U / g_zero; } } g_z = 7U; }
"""
    unit = _unit(text, "w1")
    seqs = [{"seq_num": 1, "inputs": {"g_zero": 0}, "expected": {"g_z": ""}}]
    apply_sequence_evidence({**unit, "output_vars": ["g_z"]}, seqs)
    assert seqs[0]["forked_path_undefined_behavior"] == [
        {"kind": "division_by_zero", "selector": "volatile_object:g_reg · initial_value_not_in_inputs:g_x0"}]
    findings = collect_findings([{"fid": "F1", "name": "w1"}], {"F1": seqs}, {"F1": "TC_1"})
    assert findings[0]["unknown_branch_selectors"] == ["volatile_object:g_reg", "initial_value_not_in_inputs:g_x0"]
    import openpyxl
    wb = openpyxl.Workbook()
    write_source_findings_sheet(wb, findings)
    meaning = wb["Source Findings"].cell(row=2, column=12).value
    assert "모델이 값을 정하지 못한 분기(volatile_object:g_reg, initial_value_not_in_inputs:g_x0" in meaning
    assert "시험 입력이 정하지 않은" not in meaning


def test_a_sequence_with_several_tags_of_one_kind_is_one_sequence():
    # review W3 M17: two paths, two selector lists, one kind — one sequence, each branch named once
    text = H + """U16 g_u; U16 g_w; U16 g_zero; U16 g_x; U16 g_y; U16 g_z;
void two(void) { if (g_u > 3U) { g_x = 10U / g_zero; } if (g_w > 3U) { g_y = 10U / g_zero; } g_z = 7U; }
"""
    unit = _unit(text, "two")
    seqs = [{"seq_num": 1, "inputs": {"g_zero": 0}, "expected": {"g_z": ""}}]
    apply_sequence_evidence({**unit, "output_vars": ["g_z"]}, seqs)
    assert len(seqs[0]["forked_path_undefined_behavior"]) == 2
    findings = collect_findings([{"fid": "F1", "name": "two"}], {"F1": seqs}, {"F1": "TC_1"})
    assert [(f["kind"], f["unknown_branch_sequences"], f["unknown_branch_selectors"]) for f in findings] == [
        ("division_by_zero", 1, ["initial_value_not_in_inputs:g_u", "initial_value_not_in_inputs:g_w"])]


def test_many_selectors_are_cut_with_a_count_never_read_as_an_open_input():
    from generators.c_source_oracle import FORKED_UB_SELECTOR_SEP, selector_text
    from generators.source_findings import FORKED_UB_SELECTOR_SEP as FINDINGS_SEP

    assert FINDINGS_SEP == FORKED_UB_SELECTOR_SEP
    names = [f"branch_on_unknown:initial_value_not_in_inputs:g_{i}" for i in range(6)]
    text = selector_text(names + names[:1])
    assert text.split(FORKED_UB_SELECTOR_SEP) == [f"initial_value_not_in_inputs:g_{i}" for i in range(4)] + ["… 2 more"]
    from generators.source_findings import _unknown_branch_note

    note = _unknown_branch_note({"unknown_branch_sequences": 1, "unknown_branch_selectors": text.split(FORKED_UB_SELECTOR_SEP)})
    assert "시험 입력이 정하지 않은" not in note, "hidden selectors may be registers — no open-input claim"


def test_undefined_behaviour_in_an_operand_that_runs_on_some_executions_is_judged_alike_on_any_path():
    # review W2: on a forked path it was tagged with the branch that chose the path (g_u), not the ``&&`` that runs it
    text = H + """volatile U8 g_reg; U16 g_u; U16 g_zero; U16 g_x; U16 g_y; U16 g_z;
void fork(void) { if (g_u > 3U) { g_y = 1U; } if ((g_reg > 0U) && ((10U / g_zero) > 1U)) { g_x = 1U; } g_z = 7U; }
void one(void) { if ((g_reg > 0U) && ((10U / g_zero) > 1U)) { g_x = 1U; } g_z = 7U; }
"""
    for fn in ("fork", "one"):
        r = _run(text, fn, {"g_zero": 0}, ["g_z"])
        assert r["status"] == "supported" and _values(r) == {"g_z": 7}, fn
        assert r["possible_undefined_behavior"] == ["operand_under_unknown_condition"], fn
        assert "forked_path_undefined_behavior" not in r, fn


def test_an_enumeration_object_holds_only_what_a_short_enum_would_too():
    # review C2: a {0, 300} enumeration on a 32-bit int target — the char types cannot hold it, so int alone decided,
    #   and 70000 stored was derived as 70000 (an AAPCS short enum holds 4464). With or without testimony of short's width
    for common in ("typedef unsigned char U8;\ntypedef unsigned short U16;\ntypedef unsigned int U32;\n",
                   "typedef unsigned char U8;\ntypedef unsigned int U32;\n"):
        files = {_p("c32.h"): common, _p("unit.c"): '#include "c32.h"\n' + """typedef enum { E_A = 0, E_B = 300 } WIDE_T;
WIDE_T g_e; U32 g_w;
void st(void) { g_e = (WIDE_T)70000U; g_w = (U32)g_e; }
void st_small(void) { g_e = (WIDE_T)200U; g_w = (U32)g_e; }
"""}
        ctx = cpc.build_project_context(files)
        scope = cpc.build_scopes(ctx, [_p("unit.c")])[_p("unit.c")]
        assert scope["target"]["widths"]["int"] == 32
        unit = {"name": "st", "source_text": files[_p("unit.c")], "source_path": _p("unit.c"), "source_text_complete": True,
                "project_scope": scope}
        r = evaluate_outputs(unit, [{}], [["g_w"]])[0]
        assert r["outputs"]["g_w"].get("reason") == "enum_object_type_implementation_defined", r["outputs"]
        r = evaluate_outputs({**unit, "name": "st_small"}, [{}], [["g_w"]])[0]
        assert r["outputs"]["g_w"].get("value") == 200


def test_a_downgraded_cell_drops_the_forked_path_disclosure_too():
    # review W3 M14: a cell the row's hidden input over the column cap downgrades says no basis of the derivation
    from generators import suts

    cap = suts._INPUT_COL_END - suts._INPUT_COL_START + 1
    unit = {"input_vars": [f"i{k}" for k in range(cap)], "output_vars": ["o"]}
    seq = {"inputs": {"hidden": 1}, "expected": {"o": 3},
           "expected_evidence": {"o": {"status": "derived", "possible_ub_on_forked_path": ["division_by_zero (selector g)"]}}}
    suts.render_row_io(unit, [seq])
    ev = seq["expected_evidence"]["o"]
    assert ev["status"] == "unknown" and "possible_ub_on_forked_path" not in ev


def test_a_downgraded_cell_drops_every_derivation_field():
    # (R93 integration) the SUTS downgrade and the SITS one read one list — a cell carrying every derivation field the
    #   evidence attaches (the register model's included) says no basis once downgraded
    from generators import suts
    from generators.test_evidence import DERIVATION_FIELDS

    cap = suts._INPUT_COL_END - suts._INPUT_COL_START + 1
    unit = {"input_vars": [f"i{k}" for k in range(cap)], "output_vars": ["o"]}
    full = {k: ["x"] for k in DERIVATION_FIELDS}
    seq = {"inputs": {"hidden": 1}, "expected": {"o": 3}, "expected_evidence": {"o": {"status": "derived", **full}}}
    suts.render_row_io(unit, [seq])
    ev = seq["expected_evidence"]["o"]
    assert ev["status"] == "unknown" and not set(DERIVATION_FIELDS) & set(ev), ev
    assert {"polling_loops_assumed", "storage_objects", "storage_registers"} <= set(DERIVATION_FIELDS)


def test_the_decision_search_reads_the_same_paths_as_the_evidence():
    # (R78) a search that ranks runs by their derived outputs reads them as the evidence will — and sees the run unclean
    unit = _unit(DIV, "f")
    run = observe_decisions(unit, [{"g_zero": 0}], [], outputs=["g_z", "g_y"])[0]
    assert run["status"] == "supported" and run["outputs"]["g_z"].get("value") == 7
    assert "value" not in run["outputs"]["g_y"]
    assert DIV_TAG in run["possible_undefined_behavior"]
    assert observe_decisions(unit, [{"g_zero": 0, "g_u": 9}], [], outputs=["g_z"])[0]["reason"] == \
        "undefined_behavior:division_by_zero"
    # review W3 M6: every final path proved it — the search refuses the run as the evidence does
    every = observe_decisions(_unit(DIV, "every"), [{"g_zero": 0}], [], outputs=["g_z"])[0]
    assert (every["status"], every["reason"]) == ("unsupported", "undefined_behavior:division_by_zero")


def test_a_forked_path_undefined_behaviour_stays_a_source_finding_naming_its_selector():
    unit = _unit(DIV, "f")
    seqs = [{"seq_num": 1, "inputs": {"g_zero": 0}, "expected": {"g_z": "", "g_x": ""}}]
    apply_sequence_evidence({**unit, "output_vars": ["g_z", "g_x"]}, seqs)
    assert seqs[0]["forked_path_undefined_behavior"] == [
        {"kind": "division_by_zero", "selector": "initial_value_not_in_inputs:g_u"}]
    ev = seqs[0]["expected_evidence"]["g_z"]
    assert ev["status"] == "derived" and ev["possible_ub_on_forked_path"] == [
        "division_by_zero (selector initial_value_not_in_inputs:g_u)"]
    assert any(a.startswith("possible_ub_on_forked_path:") for a in ev["assumptions"])
    assert summarize_expected_evidence(seqs)["derived_with_possible_ub_on_forked_path"] == 1
    findings = collect_findings([{"fid": "F1", "name": "f"}], {"F1": seqs}, {"F1": "TC_1"})
    assert [(f["kind"], f["occurrences"], f["sequences"], f.get("unknown_branch_sequences"),
             f.get("unknown_branch_selectors")) for f in findings] == [
        ("division_by_zero", 0, [1], 1, ["initial_value_not_in_inputs:g_u"])]
    import openpyxl
    wb = openpyxl.Workbook()
    write_source_findings_sheet(wb, findings)
    meaning = wb["Source Findings"].cell(row=2, column=12).value
    assert "시험 입력이 정하지 않은 값(initial_value_not_in_inputs:g_u)" in meaning
    # a selector the model does not know (a register): not claimed as an input left open (review X8)
    text = DIV.replace("U16 g_u;", "volatile U16 g_u;")
    unit = _unit(text, "f")
    seqs = [{"seq_num": 1, "inputs": {"g_zero": 0}, "expected": {"g_z": ""}}]
    apply_sequence_evidence({**unit, "output_vars": ["g_z"]}, seqs)
    findings = collect_findings([{"fid": "F1", "name": "f"}], {"F1": seqs}, {"F1": "TC_1"})
    wb = openpyxl.Workbook()
    write_source_findings_sheet(wb, findings)
    meaning = wb["Source Findings"].cell(row=2, column=12).value
    assert "모델이 값을 정하지 못한 분기(volatile_object:g_u" in meaning and "시험 입력이 정하지 않은" not in meaning
    # a run with the arm known refuses: the finding blocks the sequence's slots, as before R90
    unit = _unit(DIV, "f")
    seqs = [{"seq_num": 2, "inputs": {"g_zero": 0, "g_u": 9}, "expected": {"g_z": ""}}]
    apply_sequence_evidence({**unit, "output_vars": ["g_z"]}, seqs)
    assert "forked_path_undefined_behavior" not in seqs[0]
    findings = collect_findings([{"fid": "F1", "name": "f"}], {"F1": seqs}, {"F1": "TC_1"})
    assert [(f["kind"], f["occurrences"], f.get("unknown_branch_sequences")) for f in findings] == [
        ("division_by_zero", 1, None)]


def test_the_sits_disclosure_counts_values_derived_beside_forked_path_undefined_behaviour():
    from generators.integration_oracle import attach_integration_evidence
    files = {_p("common.h"): COMMON, _p("unit.c"): DIV}
    ctx = cpc.build_project_context(files)
    subs = [{"inputs": {"g_zero": 0}, "expected": {"g_z": "[검증 필요] x", "g_x": "[검증 필요] x"}},
            {"inputs": {"g_zero": 1}, "expected": {"g_z": "[검증 필요] x"}}]
    itcs = [{"tc_id": "T1", "entry_fn": "f", "sub_cases": subs}]
    stats = attach_integration_evidence(itcs, {"project_context": ctx, "source_files": files},
                                        {"f1": {"name": "f", "source_path": _p("unit.c")}})
    assert subs[0]["expected"]["g_z"] == 7 and subs[1]["expected"]["g_z"] == 7
    assert stats["derived"] == 2 and stats["derived_with_possible_ub_on_forked_path"] == 1


# ── (3) unsequenced: a pointer object written beside reads of other objects ─────────────────────

CRC_CAL = H + """U32 g_tbl[256]; U32 g_crc; U8 g_arr[4]; U8 g_v;
U32 crc_cal(const U8 *buf, U32 size)
{
    const U8 *p = buf;
    U32 crc;
    crc = ~0ul;
    while (size > 0u)
    {
        crc = g_tbl[(crc ^ *p++) & 0xFFu] ^ (crc >> 8u);
        size--;
    }
    return crc ^ ~0ul;
}
void both(const U8 *q) { const U8 *p = q; g_v = (U8)(*p++ + *p); }
void through(U8 *q) { g_arr[0] = (U8)((*q)++ + g_arr[0]); }
"""


def test_an_incremented_pointer_beside_a_table_read_is_sequenced():
    # HDPDM01 sf_Crc32Cal: before R90 ``p++`` counted as a write through ``p`` that could reach ``g_tbl``
    r = _run(CRC_CAL, "crc_cal", {"size": 0}, ["return"])
    assert r["status"] == "supported" and _values(r) == {"return": 0}
    r = _run(CRC_CAL, "crc_cal", {"size": 2}, ["return"])
    assert r["status"] == "supported", r["reason"]


def test_a_pointer_read_beside_its_own_increment_or_a_write_through_a_pointer_still_refuses():
    assert _run(CRC_CAL, "both", {}, ["g_v"])["reason"] == "unsequenced_side_effects"
    assert _run(CRC_CAL, "through", {}, ["g_arr[0]"])["reason"] == "unsequenced_side_effects"


# ── (4) enumeration inputs: 0..127 and the members (review W4) ────────────────────────────────────

ENUM = H + """typedef enum { SID_RDBI = 0xB2, SID_WDBI = 0xB7, SID_RESET = 0xB0 } SID_T;
typedef enum { E_A = 0, E_B = 1 } SMALL_T;
typedef enum { N_M = -3, N_P = 5 } NEG_T;
U8 g_r; SID_T g_sid; SMALL_T g_small; NEG_T g_neg;
void f(void) { if (g_sid == SID_RDBI) { g_r = 1U; } else { g_r = 2U; } }
void g(void) { if (g_small == E_B) { g_r = 1U; } else { g_r = 2U; } }
void h(void) { if (g_neg == N_M) { g_r = 1U; } else { g_r = 2U; } }
void k(void) { g_sid = SID_WDBI; if (g_sid == SID_WDBI) { g_r = 1U; } else { g_r = 2U; } }
"""


def test_an_enumeration_input_may_be_any_member_and_compares_as_one():
    # C11 6.7.2.2p4: the type the implementation picks holds every member — 0xB2 is a value of SID_T, and no signed char
    assert _values(_run(ENUM, "f", {"g_sid": 0xB2}, ["g_r"])) == {"g_r": 1}
    assert _values(_run(ENUM, "f", {"g_sid": 0xB0}, ["g_r"])) == {"g_r": 2}
    assert _values(_run(ENUM, "f", {"g_sid": 100}, ["g_r"])) == {"g_r": 2}
    # past the members and 0..127 — no type is sure to hold it
    assert str(_values(_run(ENUM, "f", {"g_sid": 0xB8}, ["g_r"]))["g_r"]).startswith("?")
    assert str(_values(_run(ENUM, "g", {"g_small": 128}, ["g_r"]))["g_r"]).startswith("?")
    assert _values(_run(ENUM, "g", {"g_small": 127}, ["g_r"])) == {"g_r": 2}
    assert _values(_run(ENUM, "h", {"g_neg": -3}, ["g_r"])) == {"g_r": 1}
    assert str(_values(_run(ENUM, "h", {"g_neg": -4}, ["g_r"]))["g_r"]).startswith("?")
    # stored in the enumeration object: every type that holds the members holds 0xB7 (a signed char would not, and is
    #   not one the implementation may pick for SID_T)
    assert _values(_run(ENUM, "k", {}, ["g_r", "g_sid"])) == {"g_r": 1, "g_sid": 0xB7}
