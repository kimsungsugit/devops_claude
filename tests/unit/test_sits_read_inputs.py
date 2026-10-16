"""(S0 · S2) SITS 행이 쓰는 입력 — `generators/sits_read_inputs.py` · `generators/sits.py`.

S0: GLOBAL 행이 입력으로 두는 간접 전역이 TC 의 열에 없으면 기대값이 문서에 없는 자극으로 도출된다(HDPDM01 SwITC_94/96/97
`u8g_SysUds_WriteData` · `u8g_SysUds_WdbiCmd`). 그 이름을 기존 열 뒤에 보이고, 열 상한에 막힌 입력을 쓰는 행의 확정 칸은 내린다.
목록 앞부분만 보던 자리(진입 매개변수 20 · 전역 15 · 출력 10 …)의 절단을 세고, 예산 절단의 이름을 남긴다.

S2: 입력 도메인 해상을 `_input_domain_resolver` 로 뽑아냈다 — 열 · 값은 바이트 단위로 그대로(골든). 입력 보완 계획 함수(순수)는
등급 · 순위 · 제외 사유 · 값이 입력 순서와 무관하게 정해진다.
"""
from __future__ import annotations

import json
import os
import random
import re
from pathlib import Path

import pytest

from generators import c_project_context as cpc
from generators import sits as S
from generators import sits_read_inputs as R
from generators.integration_oracle import attach_integration_evidence
from generators.test_evidence import VERIFY_PREFIX
from tests.unit import _sits_golden_project as G


def _f(name, file, calls=(), inputs=(), outputs=(), gg=()):
    return {"name": name, "file": file, "calls_list": list(calls), "inputs": list(inputs), "outputs": list(outputs),
            "globals_global": list(gg), "globals_static": [], "asil": "B", "logic_flow": [{"type": "if"}]}


def _hidden_fd():
    """두 모양 — Ap_Top: 경계를 세 홉 아래에서 넘어(거리 순 훑기 밖) GLOBAL 행의 `g_Hidden` 이 열에 없다(입력 열 여유 있음).
    Big_Top: 입력 82 · 기대 113 이 다른 전역으로 찬 뒤라 GLOBAL 행의 `g_Over` 를 붙일 자리가 없다."""
    return {
        "A1": _f("Ap_Top", "Ap.c", ["Ap_Mid"], inputs=["[IN] U8 u8_cmd"]),
        "A2": _f("Ap_Mid", "Ap.c", ["Ap_Low"]),
        "A3": _f("Ap_Low", "Ap.c", ["Drv_Set"]),
        "D1": _f("Drv_Set", "Drv.c", gg=[f"[IN] u8g_D{i}" for i in range(4)] + ["[INDIRECT] g_Hidden"]),
        "B1": _f("Big_Top", "Big.c", ["Fill", "Drv_Set2"], inputs=["[IN] U8 u8_sel"]),
        "F1": _f("Fill", "Fill.c", gg=[f"[IN] u8g_F{i:03d}" for i in range(120)]),
        "D2": _f("Drv_Set2", "Drv2.c", gg=[f"[IN] u8g_E{i}" for i in range(4)] + ["[INDIRECT] g_Over"]),
    }


def _render_like_generate_sits(itcs, stats):
    """`generate_sits` 와 같은 순서 — 렌더 → 앞부분 상한 축 맞춤 → (oracle 없음) → 요약."""
    for t in itcs:
        R.render_hidden_inputs(t, S._MAX_INPUT_PARAMS, S._MAX_EXP_PARAMS)
    stats.update(R.reconcile_candidate_cuts(itcs, S._TRUNC_INPUT_AXES, S._TRUNC_EXPECTED_AXES, S._TRUNC_AXES))
    R.downgrade_hidden_input_rows(itcs, None)
    stats.update(R.summarize_hidden_io(itcs))


def _hidden_itcs(fd=None, render=False, fi_flows=None):
    stats: dict = {}
    flows = S.collect_integration_flows(fd or _hidden_fd(), max_flows=None, stats_out=stats)
    itcs = S.generate_itc_list(flows, max_subcases=15, stats_out=stats, fi_flows=fi_flows)
    if render:
        _render_like_generate_sits(itcs, stats)
    return {t["tc_id"] if t.get("fi_design_id") else t["entry_fn"]: t for t in itcs}, stats


def _keys_outside_columns(itc):
    """행이 쓰는데 열에 없는 (행, 입력 이름) · (행, 기대 이름)."""
    cin, cexp = set(itc.get("input_vars") or []), set(itc.get("expected_vars") or [])
    return ([(i, n) for i, sc in enumerate(itc["sub_cases"]) for n in sc.get("inputs") or {} if n not in cin],
            [(i, n) for i, sc in enumerate(itc["sub_cases"]) for n in sc.get("expected") or {} if n not in cexp])


# ── S0: 숨은 입력을 열로 ─────────────────────────────────────────────────────────────────────────────────────────
class TestHiddenInputsBecomeColumns:
    def test_global_row_input_is_hidden_before_rendering(self):
        """대조군 — 고치기 전 모양이 실제로 있다(가드가 공허하지 않다)."""
        by, _ = _hidden_itcs()
        hid_in, hid_exp = _keys_outside_columns(by["Ap_Top"])
        assert {n for _r, n in hid_in} == {"g_Hidden"} and {n for _r, n in hid_exp} == {"g_Hidden"}
        g = next(sc for sc in by["Ap_Top"]["sub_cases"] if str(sc["case_label"]).startswith("GLOBAL_1"))
        assert "g_Hidden" in g["inputs"] and "g_Hidden" in g["expected"]

    def test_global_row_gv_becomes_an_appended_column(self):
        by, _ = _hidden_itcs()
        t = by["Ap_Top"]
        before_in, before_exp = list(t["input_vars"]), list(t["expected_vars"])
        rows = json.dumps([[sc["case_label"], sc["inputs"], sc["expected"]] for sc in t["sub_cases"]], sort_keys=True)
        rec = R.render_hidden_inputs(t)
        assert t["input_vars"] == before_in + ["g_Hidden"]           # 기존 열 그대로, 뒤에 붙인다(R75)
        assert t["expected_vars"] == before_exp + ["g_Hidden"]
        assert rec["inputs_shown"] == ["g_Hidden"] and rec["inputs_over_cap"] == []
        assert _keys_outside_columns(t) == ([], [])
        assert json.dumps([[sc["case_label"], sc["inputs"], sc["expected"]] for sc in t["sub_cases"]],
                          sort_keys=True) == rows                     # 행 · 라벨 · 값은 안 바뀐다
        assert not any(R.ROW_MARK in sc for sc in t["sub_cases"])

    def test_first_use_order(self):
        t = {"input_vars": ["a"], "expected_vars": [],
             "sub_cases": [{"inputs": {"a": 1, "z": 1}}, {"inputs": {"a": 1, "m": 2, "b": 3}}, {"inputs": {"b": 0}}]}
        R.render_hidden_inputs(t, 10, 10)
        assert t["input_vars"] == ["a", "z", "m", "b"]

    def test_over_cap_input_marks_the_row_and_is_not_shown(self):
        by, _ = _hidden_itcs()
        t = by["Big_Top"]
        assert len(t["input_vars"]) == S._MAX_INPUT_PARAMS and len(t["expected_vars"]) == S._MAX_EXP_PARAMS
        rec = R.render_hidden_inputs(t)
        assert rec["inputs_over_cap"] == ["g_Over"] and rec["expected_over_cap"] == ["g_Over"]
        assert len(t["input_vars"]) == S._MAX_INPUT_PARAMS                 # 시트 배치 상수 그대로
        marked = [sc for sc in t["sub_cases"] if sc.get(R.ROW_MARK)]
        assert [sc[R.ROW_MARK] for sc in marked] == [["g_Over"]] and str(marked[0]["case_label"]).startswith("GLOBAL")
        assert rec["rows_input_not_in_document"] == 1

    def test_inputs_keys_are_columns_on_the_synthetic_project(self):
        """불변식 — 렌더 후 행의 입력 이름(Inputs JSON 키) ⊆ TC 의 입력 열, **단 `inputs_not_in_document` 가 붙은 행은 예외**이고 그
        행의 열 밖 이름은 정확히 표시된 이름이다. 기대 이름 ⊆ 기대 열 ∪ `expected_over_cap`."""
        by, _ = _hidden_itcs()
        golden = G.run()["itcs"]
        for t in list(by.values()) + golden:
            R.render_hidden_inputs(t)
            hid_in, hid_exp = _keys_outside_columns(t)
            rec = t[R.RECORD_KEY]
            for r, sc in enumerate(t["sub_cases"]):
                outside = [n for rr, n in hid_in if rr == r]
                assert outside == list(sc.get(R.ROW_MARK) or []), (t["tc_id"], r, outside)
            assert {n for _r, n in hid_exp} <= set(rec["expected_over_cap"]), t["tc_id"]

    def test_render_twice_keeps_the_record(self):
        by, _ = _hidden_itcs()
        t = by["Ap_Top"]
        R.render_hidden_inputs(t)
        rec = R.render_hidden_inputs(t)
        assert rec["inputs_shown"] == ["g_Hidden"] and t["input_vars"].count("g_Hidden") == 1

    def test_rerender_with_room_clears_the_row_marks(self):
        """앞 회차의 표시는 이 회차의 것이 아니다 — 자리가 생겨 열로 보이면 행 표시를 지운다."""
        by, _ = _hidden_itcs()
        t = by["Ap_Top"]
        R.render_hidden_inputs(t, 1, 113)
        assert any(sc.get(R.ROW_MARK) for sc in t["sub_cases"])
        rec = R.render_hidden_inputs(t, 82, 113)
        assert not any(R.ROW_MARK in sc for sc in t["sub_cases"]) and rec["rows_input_not_in_document"] == 0
        assert rec["inputs_shown"] == ["g_Hidden"] and rec["inputs_over_cap"] == []


# ── S0: 상한 초과 → 확정 칸을 내린다(통합 oracle 실행) ───────────────────────────────────────────────────────────────
ROOT = os.path.join(os.sep, "virt_s0")
_H = "typedef unsigned char U8;\ntypedef unsigned int U16;\nextern U8 g_in;\nextern U8 g_out;\nU8 drv(U8 v);\n"
_A = '#include "m.h"\nU8 g_in;\nU8 g_out;\nvoid entry(U8 k) { if (g_in > 10U) { g_out = drv(k); } else { g_out = 1U; } }\n'
_B = '#include "m.h"\nU8 drv(U8 v) { if (v > 10U) { return 3U; } return 4U; }\n'


def _oracle_case(input_cap=None):
    """진입 `entry(k)` — GLOBAL 행만 `g_in` 을 준다. 그 행은 `g_out`(callee 가 쓴 값) · `g_in`(그대로 둔 값)을 도출한다."""
    files = {os.path.join(ROOT, "m.h"): _H, os.path.join(ROOT, "a.c"): _A, os.path.join(ROOT, "b.c"): _B}
    ctx = cpc.build_project_context(files)
    a = os.path.join(ROOT, "a.c")
    pend = f"{VERIFY_PREFIX} unsupported_output_binding"
    itc = {"tc_id": "SwITC_T", "entry_fn": "entry", "input_vars": ["k"], "expected_vars": ["g_out"],
           "sub_cases": [{"case_num": 1, "case_label": "1", "inputs": {"k": 1}, "expected": {"g_out": pend}},
                         {"case_num": 2, "case_label": "GLOBAL_1 [g_in=min]", "inputs": {"k": 1, "g_in": 20},
                          "expected": {"g_out": pend, "g_in": pend}}]}
    R.render_hidden_inputs(itc, input_cap if input_cap is not None else S._MAX_INPUT_PARAMS, S._MAX_EXP_PARAMS)
    stats = attach_integration_evidence([itc], {"project_context": ctx, "source_files": files},
                                        {"f": {"name": "entry", "source_path": a}})
    return itc, stats


class TestOverCapDowngrade:
    def test_with_room_the_row_keeps_its_derived_cells(self):
        itc, stats = _oracle_case()
        assert itc["input_vars"] == ["k", "g_in"] and itc["expected_vars"] == ["g_out", "g_in"]
        g = itc["sub_cases"][1]
        assert g["expected"] == {"g_out": 4, "g_in": 20}
        before = dict(stats)
        assert R.downgrade_hidden_input_rows([itc], stats)["cells"] == 0
        assert stats == before

    def test_over_cap_derived_cells_are_downgraded(self):
        itc, stats = _oracle_case(input_cap=1)
        g = itc["sub_cases"][1]
        assert g["expected"] == {"g_out": 4, "g_in": 20}                # oracle 은 문서에 없는 g_in 으로 도출했다
        assert g["expected_evidence"]["g_out"]["callees_interpreted"] == ["drv"]
        derived0, assigned0, unchanged0 = stats["derived"], stats["derived_assigned"], stats["derived_unchanged_input"]
        kinds0 = stats["unknown_reason_kinds"]
        out = R.downgrade_hidden_input_rows([itc], stats)
        reason = "input_not_in_document:g_in"
        assert g["expected"] == {"g_out": f"{VERIFY_PREFIX} {reason}", "g_in": f"{VERIFY_PREFIX} {reason}"}
        for ev in g["expected_evidence"].values():
            assert ev["status"] == "unknown" and ev["oracle_kind"] == "none" and ev["reason"] == reason
            assert not set(R.DERIVATION_FIELDS) & set(ev), ev                 # 내린 칸은 근거를 말하지 않는다
        assert "근거: 소스 계산" not in g.get("description", "")
        # 통계 — 같은 자리에서 맞춘다
        assert out == {"cells": 2, "derived_assigned": 1, "derived_unchanged_input": 1,
                       "derived_with_possible_ub_on_forked_path": 0, "tc_with_derived": 1}
        assert (stats["derived"], stats["derived_assigned"], stats["derived_unchanged_input"]) == (
            derived0 - 2, assigned0 - 1, unchanged0 - 1)
        assert stats["tc_with_derived"] == 0 and stats["unknown_reasons"]["input_not_in_document"] == 2
        assert stats["unknown_reason_kinds"] == kinds0 + 1
        assert stats["hidden_input_downgrade"]["cells"] == 2
        assert itc[R.RECORD_KEY]["downgraded_cells"] == 2
        # 요약 — 리포트 축과 같다
        summ = R.summarize_hidden_io([itc])
        assert summ["hidden_inputs_over_cap"] == 1 and summ["hidden_input_cells_downgraded"] == 2
        assert summ["hidden_input_rows_not_in_document"] == 1

    def test_report_counts_agree_with_the_document_after_the_post_pass(self):
        """품질 리포트의 기대값 집계(`summarize_expected_evidence`)와 oracle 블록이 같은 칸을 센다."""
        itc, stats = _oracle_case(input_cap=1)
        R.downgrade_hidden_input_rows([itc], stats)
        qr = S.generate_sits_quality_report([itc], 1, flow_stats={"integration_oracle": stats})
        assert qr["expected_evidence"]["derived"] == stats["derived"] == 0
        assert qr["hidden_io_columns"] == {"SwITC_T": {"inputs_over_cap": ["g_in"], "expected_shown": ["g_in"]}}

    def test_not_evaluated_stats_are_left_alone(self):
        itc = {"input_vars": [], "expected_vars": ["x"], "sub_cases": [
            {R.ROW_MARK: ["g"], "inputs": {"g": 1}, "expected": {"x": 1},
             "expected_evidence": {"x": {"status": "derived", "basis": "assigned"}}}]}
        stats = {"status": "error:RuntimeError"}
        assert R.downgrade_hidden_input_rows([itc], stats)["cells"] == 1
        assert stats == {"status": "error:RuntimeError"}
        assert itc["sub_cases"][0]["expected"]["x"].startswith(VERIFY_PREFIX)

    def test_adjusted_stats_equal_a_fresh_recount(self):
        """(리뷰 W2 ①) 맞춘 통계 == 마지막 칸을 oracle 과 같은 규칙으로 다시 센 값. T1 은 표시 없는 행에 확정 칸이 남고(TC 수
        유지), T2 는 전부 잃는다(TC 수 감소)."""
        a = ('#include "m.h"\nU8 g_in;\nU8 g_out;\nU8 g_cnt;\n'
             'void entry(U8 k) { g_cnt = k; if (g_in > 10U) { g_out = drv(k); } else { g_out = 1U; } }\n'
             'void entry2(U8 k) { if (g_in > 10U) { g_out = drv(k); } else { g_out = 1U; } }\n')
        files = {os.path.join(ROOT, "m.h"): _H.replace("extern U8 g_out;", "extern U8 g_out;\nextern U8 g_cnt;"),
                 os.path.join(ROOT, "a.c"): a, os.path.join(ROOT, "b.c"): _B}
        ctx = cpc.build_project_context(files)
        pend = f"{VERIFY_PREFIX} unsupported_output_binding"

        def itc(tc, entry, exp):
            return {"tc_id": tc, "entry_fn": entry, "input_vars": ["k"], "expected_vars": list(exp),
                    "sub_cases": [{"case_num": 1, "case_label": "1", "inputs": {"k": 1}, "expected": {e: pend for e in exp}},
                                  {"case_num": 2, "case_label": "GLOBAL_1", "inputs": {"k": 1, "g_in": 20},
                                   "expected": {**{e: pend for e in exp}, "g_in": pend}}]}
        itcs = [itc("T1", "entry", ["g_cnt", "g_out"]), itc("T2", "entry2", ["g_out"])]
        for t in itcs:
            R.render_hidden_inputs(t, 1, 113)
        a_path = os.path.join(ROOT, "a.c")
        stats = attach_integration_evidence(itcs, {"project_context": ctx, "source_files": files},
                                            {"f1": {"name": "entry", "source_path": a_path},
                                             "f2": {"name": "entry2", "source_path": a_path}})
        assert stats["tc_with_derived"] == 2
        R.downgrade_hidden_input_rows(itcs, stats)
        cnt: dict = {"cells": 0, "derived": 0, "derived_assigned": 0, "derived_unchanged_input": 0, "tc_with_derived": 0}
        reasons: dict = {}
        for t in itcs:
            here = 0
            for sub in t["sub_cases"]:
                for var, ev in sub["expected_evidence"].items():
                    if var not in sub["expected"]:
                        continue
                    cnt["cells"] += 1
                    if ev["status"] == "derived":
                        here += 1
                        cnt["derived"] += 1
                        cnt["derived_unchanged_input" if ev.get("basis") == "unchanged_input" else "derived_assigned"] += 1
                    else:
                        k = str(ev.get("reason") or "").split(":", 1)[0] or "unrecorded"
                        reasons[k] = reasons.get(k, 0) + 1
            cnt["tc_with_derived"] += 1 if here else 0
        assert {k: stats[k] for k in cnt} == cnt
        assert cnt["tc_with_derived"] == 1 and cnt["derived"] >= 1           # T1 은 g_cnt 가 남는다
        assert stats["unknown_reasons"] == dict(sorted(reasons.items(), key=lambda kv: -kv[1]))
        assert stats["unknown_reason_kinds"] == len(reasons)

    @staticmethod
    def _fake(evidence_rows, marks):
        subs = []
        for ev, mark in zip(evidence_rows, marks, strict=True):
            sc = {"inputs": {}, "expected": {k: 1 for k in ev}, "expected_evidence": ev}
            if mark:
                sc[R.ROW_MARK] = list(mark)
            subs.append(sc)
        return {"tc_id": "T", "input_vars": [], "expected_vars": [], "sub_cases": subs,
                R.RECORD_KEY: {"downgraded_cells": 0}}

    def test_basis_and_possible_ub_are_subtracted_per_kind(self):
        """(리뷰 W2 ②) 근거가 다른 칸을 섞는다 — 흐름이 쓴 값 2(그중 미정의 동작 가능 경로 1) · 입력을 그대로 둔 값 1."""
        itc = self._fake([
            {"a": {"status": "derived", "basis": "assigned"},
             "b": {"status": "derived", "basis": "assigned", "possible_ub_on_forked_path": ["x"]},
             "c": {"status": "derived", "basis": "unchanged_input"}},
            {"d": {"status": "derived", "basis": "assigned"}}], [["g"], None])
        stats = {"status": "evaluated", "derived": 10, "derived_assigned": 6, "derived_unchanged_input": 4,
                 "derived_with_possible_ub_on_forked_path": 2, "tc_with_derived": 1,
                 "unknown_reasons": {"x": 3}, "unknown_reason_kinds": 1}
        out = R.downgrade_hidden_input_rows([itc], stats)
        assert out == {"cells": 3, "derived_assigned": 2, "derived_unchanged_input": 1,
                       "derived_with_possible_ub_on_forked_path": 1, "tc_with_derived": 0}
        assert (stats["derived"], stats["derived_assigned"], stats["derived_unchanged_input"],
                stats["derived_with_possible_ub_on_forked_path"], stats["tc_with_derived"]) == (7, 4, 3, 1, 1)
        assert itc["sub_cases"][1]["expected_evidence"]["d"]["status"] == "derived"   # 표시 없는 행은 그대로

    def test_register_model_counters_are_subtracted_for_downgraded_cells(self):
        """(R93 통합) 저장소 모델 · 레지스터 메모리 · 폴링 종료 가정에 기댄 칸을 내리면 oracle 통계의 그 셋도 준다. 근거가 없는 칸은
        빼지 않고, 통계에 없던 키는 만들지 않는다(oracle 은 한 칸이라도 있을 때만 키를 만든다)."""
        itc = self._fake([
            {"a": {"status": "derived", "basis": "assigned", "storage_objects": ["REG_A"], "storage_registers": ["REG_A"]},
             "b": {"status": "derived", "basis": "assigned", "storage_objects": ["U"], "polling_loops_assumed": ["l1"]},
             "c": {"status": "derived", "basis": "assigned"}},
            {"d": {"status": "derived", "basis": "assigned", "storage_objects": ["REG_A"]}}], [["g"], None])
        stats = {"status": "evaluated", "derived": 10, "derived_assigned": 10, "derived_unchanged_input": 0,
                 "derived_with_possible_ub_on_forked_path": 0, "tc_with_derived": 1,
                 "unknown_reasons": {}, "unknown_reason_kinds": 0,
                 "derived_on_storage_model": 5, "derived_on_register_memory": 3, "derived_on_assumed_polling_exit": 2}
        out = R.downgrade_hidden_input_rows([itc], stats)
        assert out["cells"] == 3 and out["derived_on_storage_model"] == 2
        assert out["derived_on_register_memory"] == 1 and out["derived_on_assumed_polling_exit"] == 1
        assert (stats["derived_on_storage_model"], stats["derived_on_register_memory"],
                stats["derived_on_assumed_polling_exit"]) == (3, 2, 1)
        for ev in itc["sub_cases"][0]["expected_evidence"].values():
            assert not {"storage_objects", "storage_registers", "polling_loops_assumed"} & set(ev), ev
        assert itc["sub_cases"][1]["expected_evidence"]["d"]["storage_objects"] == ["REG_A"]   # 표시 없는 행은 그대로
        # 통계에 없던 키는 만들지 않는다
        st2 = {"status": "evaluated", "derived": 1, "derived_assigned": 1, "derived_unchanged_input": 0,
               "derived_with_possible_ub_on_forked_path": 0, "tc_with_derived": 1,
               "unknown_reasons": {}, "unknown_reason_kinds": 0}
        R.downgrade_hidden_input_rows([self._fake([{"x": {"status": "derived", "basis": "assigned",
                                                          "storage_objects": ["U"]}}], [["g"]])], st2)
        assert "derived_on_storage_model" not in st2

    def test_top12_is_reselected_and_kinds_counted_once(self):
        """(리뷰 W2 ③) 사유 13 종(상위 12 만 기록) — 새 사유가 들어가면 가장 작은 것이 빠지고 종류 수가 하나 는다. 이미 있던 사유면
        종류 수는 그대로다."""
        reasons = {f"r{i:02d}": 20 - i for i in range(12)}          # 20 … 9 (13 번째 종류는 기록 밖)
        stats = {"status": "evaluated", "derived": 5, "derived_assigned": 5, "derived_unchanged_input": 0,
                 "derived_with_possible_ub_on_forked_path": 0, "tc_with_derived": 1,
                 "unknown_reasons": dict(reasons), "unknown_reason_kinds": 13}
        itc = self._fake([{f"v{i}": {"status": "derived", "basis": "assigned"} for i in range(10)}], [["g"]])
        R.downgrade_hidden_input_rows([itc], stats)
        assert len(stats["unknown_reasons"]) == 12 and stats["unknown_reasons"]["input_not_in_document"] == 10
        assert "r11" not in stats["unknown_reasons"] and "r10" in stats["unknown_reasons"]   # 9 가 빠지고 10 은 남는다
        assert list(stats["unknown_reasons"].values()) == sorted(stats["unknown_reasons"].values(), reverse=True)
        assert stats["unknown_reason_kinds"] == 14
        st2 = dict(stats, unknown_reasons=dict(stats["unknown_reasons"]), derived=1, derived_assigned=1)
        R.downgrade_hidden_input_rows([self._fake([{"w": {"status": "derived", "basis": "assigned"}}], [["g"]])], st2)
        assert st2["unknown_reason_kinds"] == 14 and st2["unknown_reasons"]["input_not_in_document"] == 11

    def test_reason_names_three_and_counts_the_rest(self):
        """(리뷰 W2 ⑥) 이름이 넷 이상이면 앞 셋 + `+N`."""
        itc = self._fake([{"x": {"status": "derived", "basis": "assigned"}}], [["a", "b", "c", "d"]])
        R.downgrade_hidden_input_rows([itc], None)
        ev = itc["sub_cases"][0]["expected_evidence"]["x"]
        assert ev["reason"] == "input_not_in_document:a,b,c+1"
        assert itc["sub_cases"][0]["expected"]["x"] == f"{VERIFY_PREFIX} input_not_in_document:a,b,c+1"

    def test_downgraded_cells_accumulate_over_calls(self):
        itc, stats = _oracle_case(input_cap=1)
        R.downgrade_hidden_input_rows([itc], stats)
        assert R.downgrade_hidden_input_rows([itc], stats)["cells"] == 0      # 두 번째엔 내릴 칸이 없다
        assert itc[R.RECORD_KEY]["downgraded_cells"] == 2                      # 앞 회차의 기록은 남는다

    def test_downgrade_drops_every_derivation_field_the_oracle_sets(self):
        """`apply_sequence_evidence` 가 확정 칸에 다는 필드가 늘면 내릴 때 뗄 목록도 늘어야 한다(근거를 말하는 미상 칸 방지)."""
        src = (Path(__file__).resolve().parents[2] / "generators" / "test_evidence.py").read_text(encoding="utf-8")
        set_on_item = set(re.findall(r'item\["(\w+)"\]\s*=', src)) - {"project_scope", "reason"}
        assert set_on_item and set_on_item <= set(R.DERIVATION_FIELDS), sorted(set_on_item - set(R.DERIVATION_FIELDS))

    def test_evidence_sheet_inputs_are_document_columns(self, tmp_path):
        """불변식(문서) — 'Test Evidence' 시트의 Inputs JSON 이름은 그 TC 의 입력 열이다. 명세 시트도 붙인 열과 값을 싣는다."""
        import openpyxl

        from generators.integration_oracle import write_integration_evidence_sheet
        itc, _ = _oracle_case()
        wb = openpyxl.Workbook()
        assert write_integration_evidence_sheet(wb, [itc]) > 0
        rows = list(wb["Test Evidence"].iter_rows(values_only=True))
        col = rows[0].index("Inputs JSON")
        for r in rows[1:]:
            assert set(json.loads(r[col])) <= set(itc["input_vars"]), r
        out = S.generate_sits_xlsm(None, [dict(itc, call_chain="entry -> drv", related_ids=[], asil="")],
                                   str(tmp_path / "s.xlsm"))
        ws = openpyxl.load_workbook(out)[S._SPEC_SHEET_NAME]
        head = [ws.cell(S._DATA_START_ROW, S._INPUT_COL_START + i).value for i in range(3)]
        assert head == ["k", "g_in", None]
        assert ws.cell(S._DATA_START_ROW + 2, S._INPUT_COL_START + 1).value == 20     # GLOBAL 행의 g_in
        assert ws.cell(S._DATA_START_ROW + 1, S._INPUT_COL_START + 1).value in (None, "")


# ── S0: 생성 경로 배선 ───────────────────────────────────────────────────────────────────────────────────────────────
_PH = "typedef unsigned char U8;\ntypedef unsigned int U16;\nextern U8 g_hid;\nextern U8 g_out;\nvoid b_set(U8 v);\n"
_PA = ('#include "m.h"\nU8 g_out;\nstatic void a_low(U8 k) { b_set(k); }\nstatic void a_mid(U8 k) { a_low(k); }\n'
       'void a_top(U8 k) { a_mid(k); }\n')
_PB = ('#include "m.h"\nU8 g_hid;\nstatic void b_helper(U8 v) { if (g_hid > 5U) { g_out = v; } else { g_out = 1U; } }\n'
       'void b_set(U8 v) { b_helper(v); }\n')


def test_generate_sits_shows_hidden_inputs_in_the_document(tmp_path):
    """소스 파싱부터 — `a_top` 은 경계를 세 홉 아래(b.c)에서 넘어 GLOBAL 행의 `g_out` · `g_hid` 가 열에 없었다. 생성 문서의 Test
    Evidence Inputs JSON 이름은 전부 그 TC 의 입력 열이고, 품질 리포트 · 공시에 축이 실린다."""
    import openpyxl
    src = tmp_path / "src"
    src.mkdir()
    (src / "m.h").write_text(_PH)
    (src / "a.c").write_text(_PA)
    (src / "b.c").write_text(_PB)
    out = S.generate_sits(source_root=str(src), output_path=str(tmp_path / "o.xlsm"), max_subcases=15)
    assert out.get("output_path"), out.get("error")
    cov = out["quality_report"]["integration_flow_coverage"]
    assert cov["hidden_inputs_shown"] >= 2 and cov["hidden_inputs_over_cap"] == 0
    shown = out["quality_report"]["hidden_io_columns"]
    top = next(v for k, v in shown.items())
    assert set(top["inputs_shown"]) == {"g_out", "g_hid"}
    wb = openpyxl.load_workbook(out["output_path"], read_only=True)
    cols: dict = {}
    for row in wb[S._SPEC_SHEET_NAME].iter_rows(min_row=S._DATA_START_ROW, values_only=True):
        tc = row[S._TCID_COL - 1]
        if tc and str(tc).startswith("SwITC_"):
            cols[tc] = [v for v in row[S._INPUT_COL_START - 1:S._INPUT_COL_END] if v not in (None, "")]
    ev = list(wb["Test Evidence"].iter_rows(values_only=True))
    i_tc, i_in = ev[0].index("Test Case ID"), ev[0].index("Inputs JSON")
    for r in ev[1:]:
        assert set(json.loads(r[i_in])) <= set(cols[r[i_tc]]), (r[i_tc], r[i_in], cols[r[i_tc]])
    wb.close()
    from report_gen.generation_disclosures import build_disclosures
    item = {i["key"]: i for i in build_disclosures("sits", out["quality_report"])}["sits_hidden_inputs"]
    assert item["value"].startswith("입력 ") and item["tone"] == "info"


_PB_WIDE = ('#include "m.h"\nU8 g_hid;\nU8 g_d0;\nU8 g_d1;\nU8 g_d2;\nU8 g_d3;\n'
            'static void b_helper(U8 v) { if (g_hid > 5U) { g_out = v; } else { g_out = 1U; } }\n'
            'void b_set(U8 v) { g_d0 = v; g_d1 = v; g_d2 = v; g_d3 = v; b_helper(v); }\n')


def test_generate_sits_does_not_count_shown_names_as_lost(tmp_path):
    """(리뷰 W1) 소스 파싱부터 — `b_set` 의 전역은 직접 넷 뒤에 간접 `g_out` · `g_hid` 다. `a_top` 은 경계를 세 홉 아래에서 넘어 그
    둘이 기대 앞부분 상한(callee 전역 앞 4)에 잘리지만 `g_out` 은 간접 전역 경로로 GLOBAL 행에 들어가 열로 보인다. 문서 · 리포트의
    `callee_observables_truncated` 는 보이지 않은 `g_hid` 하나만 센다."""
    src = tmp_path / "src"
    src.mkdir()
    (src / "m.h").write_text(_PH)
    (src / "a.c").write_text(_PA)
    (src / "b.c").write_text(_PB_WIDE)
    out = S.generate_sits(source_root=str(src), output_path=str(tmp_path / "o.xlsm"), max_subcases=15)
    assert out.get("output_path"), out.get("error")
    qr = out["quality_report"]
    cuts, shown = qr["candidate_cut_names"], qr["hidden_io_columns"]
    top = next(tc for tc, v in cuts.items() if "callee_observables" in v)
    assert cuts[top]["callee_observables"] == ["g_hid"] and "g_out" in shown[top]["expected_shown"]
    assert qr["integration_flow_coverage"]["callee_observables_truncated"] == 1
    for tc, axes in cuts.items():
        rec = shown.get(tc) or {}
        for ax in S._TRUNC_INPUT_AXES:
            assert not set(axes.get(ax) or []) & set(rec.get("inputs_shown") or []), (tc, ax)
        for ax in S._TRUNC_EXPECTED_AXES:
            assert not set(axes.get(ax) or []) & set(rec.get("expected_shown") or []), (tc, ax)


def test_generate_sits_downgrades_rows_whose_input_cannot_be_shown(tmp_path, monkeypatch):
    """(리뷰 W2 ④) 입력 열 상한을 1 로 — GLOBAL 행의 간접 전역을 붙일 자리가 없다. 생성 경로가 그 행의 확정 칸을 내리고(후처리
    배선), 통합 oracle 블록 · 흐름 축 · 문서 칸이 같은 사실을 말한다."""
    import openpyxl
    monkeypatch.setattr(S, "_MAX_INPUT_PARAMS", 1)
    src = tmp_path / "src"
    src.mkdir()
    (src / "m.h").write_text(_PH)
    (src / "a.c").write_text(_PA)
    (src / "b.c").write_text(_PB)
    out = S.generate_sits(source_root=str(src), output_path=str(tmp_path / "o.xlsm"), max_subcases=15)
    assert out.get("output_path"), out.get("error")
    qr = out["quality_report"]
    cov, io = qr["integration_flow_coverage"], qr["integration_oracle"]
    assert cov["hidden_inputs_over_cap"] >= 1 and cov["hidden_input_cells_downgraded"] > 0
    assert io["hidden_input_downgrade"]["cells"] == cov["hidden_input_cells_downgraded"]
    assert io["unknown_reasons"]["input_not_in_document"] == cov["hidden_input_cells_downgraded"]
    assert qr["expected_evidence"]["derived"] == io["derived"]
    wb = openpyxl.load_workbook(out["output_path"], read_only=True)
    cells = [v for row in wb[S._SPEC_SHEET_NAME].iter_rows(min_row=S._DATA_START_ROW, values_only=True)
             for v in row[S._EXP_COL_START - 1:S._EXP_COL_END] if isinstance(v, str) and "input_not_in_document:" in v]
    wb.close()
    assert len(cells) == cov["hidden_input_cells_downgraded"] > 0       # 기대 열 여유가 있어 내린 칸이 전부 문서에 보인다


# ── S0: 앞부분 상한 절단 · 예산 절단의 이름 ───────────────────────────────────────────────────────────────────────
class TestTruncationCounters:
    def test_entry_lists_cut_by_the_head_limits_are_counted_with_names(self):
        got = G.run()
        big = next(f for f in got["flows"] if f["entry_fn"] == "Big_Entry")
        assert big["candidate_cuts"]["entry_params"] == ["p20", "p21", "p22", "p23"]
        assert big["candidate_cuts"]["entry_globals"] == ["u8g_B15", "u8g_B16", "u8g_B17"]
        assert big["candidate_cuts"]["entry_outputs"] == ["u8g_O10", "u8g_O11"]
        st = got["stats"]
        assert (st["entry_params_truncated"], st["entry_globals_truncated"], st["entry_outputs_truncated"]) == (4, 3, 2)
        assert st["callee_params_truncated"] == 0

    def test_a_cut_name_recovered_elsewhere_is_not_counted(self):
        """16 번째 전역이 다른 경로(거리 순 훑기)로 후보가 되면 잘린 게 아니다."""
        fd = {"E": _f("Ent", "E.c", ["Cal"], inputs=["[IN] U8 x"], gg=[f"[IN] u8g_G{i:02d}" for i in range(16)]),
              "C": _f("Cal", "C.c", gg=["[IN] u8g_G15"])}
        st: dict = {}
        flow = S.collect_integration_flows(fd, max_flows=None, stats_out=st)[0]
        assert st["entry_globals_truncated"] == 0 and flow["candidate_cuts"]["entry_globals"] == []
        assert "u8g_G15" in flow["input_vars"]
        fd["C"]["globals_global"] = []
        st = {}
        flow = S.collect_integration_flows(fd, max_flows=None, stats_out=st)[0]
        assert st["entry_globals_truncated"] == 1 and flow["candidate_cuts"]["entry_globals"] == ["u8g_G15"]

    def test_callee_params_boundaries(self):
        """(리뷰 W2 ⑤) 입력 없는 진입 — 다른 모듈 callee 다섯(앞 넷만 매개변수를 본다), 각 매개변수 일곱(앞 여섯만)."""
        fd = {"E": _f("Ent", "E.c", [f"C{k}" for k in range(5)]),
              **{f"C{k}": _f(f"C{k}", f"C{k}.c", inputs=[f"[IN] U8 c{k}_p{i}" for i in range(7)]) for k in range(5)}}
        st: dict = {}
        flow = next(f for f in S.collect_integration_flows(fd, max_flows=None, stats_out=st) if f["entry_fn"] == "Ent")
        want = [f"c{k}_p6" for k in range(4)] + [f"c4_p{i}" for i in range(7)]
        assert flow["candidate_cuts"]["callee_params"] == want
        assert st["callee_params_truncated"] == len(want) == 11

    def test_callee_observables_boundaries(self):
        """(리뷰 W2 ⑤) 경계를 세 홉 아래에서 넘는 진입(거리 순 훑기 밖) — 다른 모듈 callee 여섯(앞 다섯만 출력 · 전역을 본다),
        각 출력 여섯(앞 다섯) · 전역 여섯(앞 넷)."""
        fd = {"E": _f("Ent", "E.c", ["Mid"], inputs=["[IN] U8 x"]), "M": _f("Mid", "E.c", ["Low"]),
              "L": _f("Low", "E.c", [f"C{k}" for k in range(6)]),
              **{f"C{k}": _f(f"C{k}", f"C{k}.c", outputs=[f"[OUT] o{k}_{i}" for i in range(6)],
                             gg=[f"[IN] g{k}_{i}" for i in range(6)]) for k in range(6)}}
        st: dict = {}
        flow = next(f for f in S.collect_integration_flows(fd, max_flows=None, stats_out=st) if f["entry_fn"] == "Ent")
        cut = flow["candidate_cuts"]["callee_observables"]
        near = [n for k in range(5) for n in (f"o{k}_5", f"g{k}_4", f"g{k}_5")]
        far = [f"o5_{i}" for i in range(6)] + [f"g5_{i}" for i in range(6)]
        assert sorted(cut) == sorted(near + far) and len(cut) == 27
        assert flow["candidate_cuts"]["entry_outputs"] == []

    def test_function_names_are_not_counted(self):
        """잘린 진입 전역의 이름이 프로젝트 함수 이름이면(ISR 부산물) 관측 후보가 아니다 — 세지 않는다."""
        fd = {"E": _f("Ent", "E.c", ["Cal"], inputs=["[IN] U8 x"],
                      gg=[f"[IN] u8g_G{i:02d}" for i in range(15)] + ["[IN] Cal"]),
              "C": _f("Cal", "C.c")}
        st: dict = {}
        S.collect_integration_flows(fd, max_flows=None, stats_out=st)
        assert st["entry_globals_truncated"] == 0

    def test_expected_axes_are_checked_against_expected_candidates(self):
        """잘린 진입 출력이 **입력** 후보이기만 하면 기대 축에선 잃은 것이다(관측 대상이 아니다)."""
        fd = {"E": _f("Ent", "E.c", ["Cal"], inputs=["[IN] U8 x"], gg=["[IN] u8g_Shared"],
                      outputs=[f"[OUT] u8g_O{i:02d}" for i in range(10)] + ["[OUT] u8g_Shared"]),
              "C": _f("Cal", "C.c")}
        st: dict = {}
        flow = S.collect_integration_flows(fd, max_flows=None, stats_out=st)[0]
        assert "u8g_Shared" in flow["input_vars"]
        assert flow["candidate_cuts"]["entry_outputs"] == ["u8g_Shared"] and st["entry_outputs_truncated"] == 1

    def test_indirect_state_candidates_cut_by_the_list_caps(self):
        """(리뷰 I6) 간접 전역 목록의 세 상한(callee 4 · callee 전역 앞 5 · 목록 5) — 잘린 이름을 센다."""
        fd = {"E": _f("Ent", "E.c", [f"C{k}" for k in range(5)], inputs=["[IN] U8 x"]),
              **{f"C{k}": _f(f"C{k}", f"C{k}.c", gg=[f"[INDIRECT] s{k}"]) for k in range(5)}}
        fd["C0"]["globals_global"] = [f"[INDIRECT] s0_{i}" for i in range(7)]
        st: dict = {}
        flow = S.collect_integration_flows(fd, max_flows=None, stats_out=st)[0]
        assert flow["indirect_vars"] == [f"s0_{i}" for i in range(5)]
        assert flow["candidate_cuts"]["indirect_vars"] == ["s0_5", "s0_6", "s1", "s2", "s3", "s4"]
        assert st["indirect_vars_truncated"] == 6

    def test_collect_alone_counts_before_rendering(self):
        """렌더 전 값은 상한이다 — Ap_Top 의 `g_Hidden`(Drv_Set 의 다섯째 전역)은 기대 앞부분 상한에 잘렸지만 간접 전역 경로로
        GLOBAL 행에 들어간다. 흐름만 모은 값은 1, 생성 경로(렌더 → 맞춤)는 0."""
        _by, st = _hidden_itcs()
        assert st["callee_observables_truncated"] == 1
        by, st = _hidden_itcs(render=True)
        assert st["callee_observables_truncated"] == 0 and by["Ap_Top"]["candidate_cuts"]["callee_observables"] == []
        assert by["Ap_Top"][R.RECORD_KEY]["expected_shown"] == ["g_Hidden"]

    @pytest.mark.parametrize("label, fd, axis", [
        ("entry_global_16th_indirect",
         {"E": _f("Ent", "E.c", ["Cal"], inputs=["[IN] U8 x"],
                  gg=[f"[IN] u8g_G{i:02d}" for i in range(15)] + ["[INDIRECT] g_Ind", "[IN] u8g_Lost"]),
          "C": _f("Cal", "C.c", gg=["[IN] u8g_Other"])}, "entry_globals"),
        ("callee_global_5th_indirect", None, "callee_observables"),
    ])
    def test_no_name_is_both_lost_and_shown(self, label, fd, axis):
        """(리뷰 W1) 잘렸다고 센 이름과 S0 가 열로 보인 이름이 겹치지 않는다 — 리뷰 프로브의 두 모양."""
        _pre, st_pre = _hidden_itcs(fd)
        by, st = _hidden_itcs(fd, render=True)
        for t in by.values():
            rec = t[R.RECORD_KEY]
            shown_in, shown_exp = set(rec["inputs_shown"]), set(rec["expected_shown"])
            for ax in S._TRUNC_INPUT_AXES:
                assert not set(t["candidate_cuts"].get(ax) or []) & shown_in, (label, t["entry_fn"], ax)
            for ax in S._TRUNC_EXPECTED_AXES:
                assert not set(t["candidate_cuts"].get(ax) or []) & shown_exp, (label, t["entry_fn"], ax)
        assert st[f"{axis}_truncated"] == st_pre[f"{axis}_truncated"] - 1     # 정확히 보인 이름 하나만 빠진다
        if label == "entry_global_16th_indirect":
            assert by["Ent"]["candidate_cuts"]["entry_globals"] == ["u8g_Lost"]   # 보이지 않은 이름은 남는다
            assert by["Ent"][R.RECORD_KEY]["inputs_shown"] == ["g_Ind"]

    def test_reconcile_pairs_input_axes_with_inputs_and_expected_axes_with_expected(self):
        """입력 축은 입력 열로 보인 이름과, 기대 축은 기대 열로 보인 이름과 맞춘다 — 한쪽에만 보인 이름은 다른 축에서 남는다."""
        itc = {"candidate_cuts": {"entry_globals": ["a", "b"], "entry_outputs": ["a", "c"], "budget_input": ["a"]},
               R.RECORD_KEY: {"inputs_shown": ["a"], "expected_shown": ["c"]}}
        totals = R.reconcile_candidate_cuts([itc], S._TRUNC_INPUT_AXES, S._TRUNC_EXPECTED_AXES, S._TRUNC_AXES)
        assert itc["candidate_cuts"] == {"entry_globals": ["b"], "entry_outputs": ["a"], "budget_input": ["a"]}
        assert totals["entry_globals_truncated"] == 1 and totals["entry_outputs_truncated"] == 1
        assert set(totals) == {f"{a}_truncated" for a in S._TRUNC_AXES}

    def test_name_shown_only_as_input_clears_the_input_axis(self):
        """실제 흐름 — 진입의 16 번째 전역 `g_Ind`(간접)가 GLOBAL 행으로 입력 열에 보이지만 기대 열엔 callee 출력으로 이미 있다."""
        fd = {"E": _f("Ent", "E.c", ["Cal"], inputs=["[IN] U8 x"],
                      gg=[f"[IN] u8g_G{i:02d}" for i in range(15)] + ["[INDIRECT] g_Ind"]),
              "C": _f("Cal", "C.c", outputs=["[OUT] g_Ind"])}
        by, st = _hidden_itcs(fd, render=True)
        rec = by["Ent"][R.RECORD_KEY]
        assert rec["inputs_shown"] == ["g_Ind"] and rec["expected_shown"] == []
        assert by["Ent"]["candidate_cuts"]["entry_globals"] == [] and st["entry_globals_truncated"] == 0

    def test_fi_tc_is_not_counted_twice(self):
        """같은 흐름의 전용 FI TC 는 흐름 단위 합계에 넣지 않는다."""
        flows = G.run()["flows"]
        big = next(f for f in flows if f["entry_fn"] == "Big_Entry")
        st: dict = {}
        S.collect_integration_flows(G.function_details(), max_flows=None, stats_out=st,
                                    globals_info_map=dict(G.GLOBALS_INFO), struct_member_types=dict(G.STRUCT_MEMBER_TYPES),
                                    uds_unit_io=G.uds_unit_io())
        pre = {k: st[k] for k in st if k.endswith("_truncated")}
        itcs = S.generate_itc_list(flows, max_subcases=15, fi_flows=[dict(big, fi_design_id="SwFn_01")])
        assert sum(1 for t in itcs if t.get("fi_design_id")) == 1
        _render_like_generate_sits(itcs, st)
        assert {k: st[k] for k in pre} == pre == {**{k: 0 for k in pre}, "entry_params_truncated": 4,
                                                  "entry_globals_truncated": 3, "entry_outputs_truncated": 2}

    def test_fi_tc_without_io_gets_the_scenario_columns_not_hidden_ones(self):
        """(리뷰 I3) 입출력 없는 흐름의 FI TC 도 일반 TC 와 같이 `Scenario` · `Result` 를 열로 낸다 — S0 가 숨은 이름으로 세지 않는다."""
        fd = {"E": _f("Ent", "E.c", ["Cal"]), "C": _f("Cal", "C.c")}
        flows = S.collect_integration_flows(fd, max_flows=None)
        itcs = S.generate_itc_list(flows, max_subcases=7, fi_flows=[dict(flows[0], fi_design_id="SwFn_07")])
        fi = next(t for t in itcs if t.get("fi_design_id"))
        assert (fi["input_vars"], fi["expected_vars"]) == (["Scenario"], ["Result"]) == (
            itcs[0]["input_vars"], itcs[0]["expected_vars"])
        st: dict = {}
        _render_like_generate_sits(itcs, st)
        assert st["hidden_inputs_shown"] == st["hidden_expected_shown"] == 0

    def test_budget_cut_names_are_recorded(self):
        by, st = _hidden_itcs(render=True)
        cuts = by["Big_Top"]["candidate_cuts"]
        assert len(cuts["budget_input"]) == st["var_budget_cut_input"] and cuts["budget_input"][-1] == "g_Over"
        assert cuts["budget_input"][0] == "u8g_F081"                        # 만난 순서 그대로
        assert len(cuts["budget_expected"]) == st["var_budget_cut_expected"]
        qr = S.generate_sits_quality_report(list(by.values()), 1, flow_stats=st)
        names = qr["candidate_cut_names"][by["Big_Top"]["tc_id"]]["budget_input"]
        # (리뷰 I5) TC · 축마다 앞 20 개 + `외 N`
        assert names == cuts["budget_input"][:20] + [f"외 {len(cuts['budget_input']) - 20}"]
        assert len(cuts["budget_input"]) > 20
        assert by["Ap_Top"]["tc_id"] not in qr["candidate_cut_names"]       # 보인 이름은 잃은 게 아니다(W1)
        assert by["Ap_Low"]["tc_id"] not in qr["candidate_cut_names"]       # 잘린 것 없음 → 싣지 않는다
        assert S._head_names(["a"] * 20) == ["a"] * 20

    def test_keys_reach_report_loss_axes_and_disclosure(self):
        from report_gen.generation_disclosures import build_disclosures
        by, st = _hidden_itcs(render=True)
        cov = S.generate_sits_quality_report(list(by.values()), 1, flow_stats=st)["integration_flow_coverage"]
        new = [k for k in st if k.endswith("_truncated") or k.startswith("hidden_")]
        assert new and all(k in cov for k in new), sorted(set(new) - set(cov))
        assert {"hidden_inputs_over_cap", "hidden_input_cells_downgraded", "entry_params_truncated",
                "indirect_vars_truncated"} <= set(S._FLOW_LOSS_KEYS)
        got = {i["key"]: i for i in build_disclosures("sits", {"integration_flow_coverage": cov})}
        assert "sits_candidates_truncated" not in got       # 맞춘 뒤 이 프로젝트엔 잃은 이름이 없다(W1 — 예전엔 1)
        assert got["sits_hidden_inputs"]["value"] == (
            "입력 1 · 기대 1 열 추가 · 상한 초과 입력 1 · 기대 1 · 내린 확정 칸 0 (TC 2)")
        assert got["sits_hidden_inputs"]["tone"] == "warning"
        gst: dict = {}
        _render_like_generate_sits(S.generate_itc_list(G.run()["flows"], max_subcases=15), gst)
        gcov = {k: v for k, v in gst.items() if k.endswith("_truncated")}
        item = {i["key"]: i for i in build_disclosures("sits", {"integration_flow_coverage": gcov})}[
            "sits_candidates_truncated"]
        assert item["value"] == "진입 매개변수 4 · 진입 전역 3 · 진입 출력 2 · callee 매개변수 0 · callee 출력·전역 0"
        ind = {i["key"]: i for i in build_disclosures("sits", {"integration_flow_coverage": {
            **gcov, "indirect_vars_truncated": 6}})}["sits_candidates_truncated"]
        assert ind["value"].endswith(" · 전역 상태 행 후보 6") and "앞 2개만" in ind["note"]
        # 키가 없는 구판 · 전부 0 은 감춘다
        quiet = build_disclosures("sits", {"integration_flow_coverage": {k: 0 for k in new}})
        assert not {"sits_candidates_truncated", "sits_hidden_inputs"} & {i["key"] for i in quiet}


# ── S2: 리팩터 골든 ─────────────────────────────────────────────────────────────────────────────────────────────────
class TestResolverGolden:
    def test_columns_and_values_are_byte_identical_to_before_the_extraction(self):
        fixture = G.FIXTURE.read_text(encoding="utf-8")
        old = json.loads(fixture)
        snap = G.snapshot(G.run())
        assert set(old["stats"]) <= set(snap["stats"]), sorted(set(old["stats"]) - set(snap["stats"]))
        snap["stats"] = {k: snap["stats"][k] for k in old["stats"]}
        assert G.dumps(snap) + "\n" == fixture

    def test_the_fixture_exercises_the_resolution_paths(self):
        """골든이 공허하지 않다 — 출처 · 사유가 고루 들어 있다."""
        st = json.loads(G.FIXTURE.read_text(encoding="utf-8"))["stats"]
        assert all(v > 0 for v in st["input_domain_source"].values())
        assert set(st["input_documented_range_unused"]) == {
            "documented_outside_type", "documented_range_unread:digit_count_differs", "documented_ranges_disagree",
            "pointer_address_range", "float_not_ranged"}
        assert set(st["input_type_resolution"]["unresolved"]) == {
            "no_declaration", "pointer_object", "struct_object", "member_type_unresolved", "type_unresolved"}
        # (리뷰 I9) 간접 전역 — 포인터는 주소 범위를 받지 않고(None), 구조체 통째는 설계서 범위를 받는다(뽑아내기 전과 같은 동작)
        odd = next(f for f in json.loads(G.FIXTURE.read_text(encoding="utf-8"))["flows"] if f["entry_fn"] == "Odd_Entry")
        kinds = {n: (d or {}).get("kind") for n, d in zip(odd["indirect_vars"], odd["indirect_types"], strict=True)}
        assert kinds == {"g_ptr2": None, "REG_ODD": "documented", "f32g_Gain2": "table"}

    def test_itc_carries_the_chain_functions(self):
        got = G.run()
        by = {f["entry_fn"]: f for f in got["flows"]}
        for t in got["itcs"]:
            assert t["chain_functions"] == by[t["entry_fn"]]["functions"] and t["chain_functions"][0] == t["entry_fn"]

    def test_resolver_is_the_flow_judgment(self):
        resolve = _resolver()
        r = resolve(["u8g_Speed", "g_ptr", "REG_FSTAT"], ["", "", ""], ["Ap_Run", "Drv_Write"])
        assert r[0].source == "documented_range" and (r[0].domain["vlo"], r[0].domain["vhi"]) == (10, 20)
        assert r[1].type_reason == "pointer_object" and r[1].domain is None
        assert r[2].type_reason == "struct_object"
        assert resolve(["u8g_Speed"], [""], ["Ap_Run"])[0].source == "declared_type"   # 체인 밖 표는 안 본다


def _resolver():
    import copy
    return S._input_domain_resolver(G.function_details(), copy.deepcopy(G.GLOBALS_INFO),
                                    copy.deepcopy(G.STRUCT_MEMBER_TYPES), G.uds_unit_io())


# ── S2: 계획 함수(순수) ────────────────────────────────────────────────────────────────────────────────────────────
def _ev(reason, status="unknown"):
    return {"status": status, "reason": reason}


def _plan_itc():
    return {"input_vars": ["k", "REG_FSTAT"], "expected_vars": ["o1", "o2", "o3"],
            "chain_functions": ["Ap_Run", "Drv_Write", "Io_Set"],
            "sub_cases": [
                {"expected_evidence": {
                    "o1": _ev("initial_value_not_in_inputs:u8g_Speed"),
                    "o2": _ev("unwritten_output_initial_not_in_inputs:o2b"),
                    "o3": _ev("path_dependent:branch_on_unknown:initial_value_not_in_inputs:g_DoorState")}},
                {"expected_evidence": {
                    "o1": _ev("initial_value_not_in_inputs:u8g_Speed"),
                    "o2": _ev("parameter_not_in_inputs:k"),
                    "o3": _ev("initial_value_not_in_inputs:REG_FSTAT.Bits.CCIF"),
                    "o4": _ev("initial_value_not_in_inputs:tbl[1][2]")}},
                {"expected_evidence": {
                    "o1": _ev("whatever", status="derived"),
                    "o2": _ev("initial_value_not_in_inputs:s16g_Pos;also"),
                    "o3": _ev("unwritten_output_initial_not_in_inputs:arr[2].m"),
                    "o4": _ev("input_not_in_document:g_x")}}]}


def _shuffled(itc, seed):
    rnd = random.Random(seed)
    out = json.loads(json.dumps(itc))
    for sc in out["sub_cases"]:
        items = list(sc["expected_evidence"].items())
        rnd.shuffle(items)
        sc["expected_evidence"] = dict(items)
    return out


class TestNeededInputs:
    def test_tiers_cells_rows_first(self):
        need = R.needed_inputs(_plan_itc())
        assert need["u8g_Speed"] == {"tier": "value_read", "cells": 2, "rows": [0, 1], "first": (0, 0)}
        assert need["o2b"] == {"tier": "observed_initial", "cells": 1, "rows": [0], "first": (0, 1)}
        assert need["g_DoorState"]["tier"] == "decision"
        assert need["k"]["tier"] == "value_read"
        assert need["tbl[1][2]"]["first"] == (1, 3)                       # 열에 없는 기대값은 열 수 뒤로
        assert "g_x" not in need and "whatever" not in need               # 내린 칸 · 확정 칸은 보지 않는다
        assert need["s16g_Pos"]["cells"] == 1

    def test_reads_the_oracle_reasons(self):
        """실제 통합 oracle 사유 — 첫 행은 `g_in` 을 주지 않아 막혔다. S0 가 그 이름을 이미 열로 보였으므로 보완 대상이 아니다."""
        itc, _ = _oracle_case()
        reason = itc["sub_cases"][0]["expected_evidence"]["g_out"]["reason"]
        need = R.needed_inputs(itc)
        assert "g_in" in need and need["g_in"]["rows"] == [0], (reason, need)
        _kept, excluded = R.exclude_needed(need, itc)
        assert excluded["g_in"] == "already_a_column"

    def test_best_tier_wins(self):
        itc = {"expected_vars": ["a", "b"], "sub_cases": [{"expected_evidence": {
            "a": _ev("path_dependent:branch_on_unknown:initial_value_not_in_inputs:g"),
            "b": _ev("unwritten_output_initial_not_in_inputs:g")}}]}
        assert R.needed_inputs(itc)["g"]["tier"] == "observed_initial" and R.needed_inputs(itc)["g"]["cells"] == 2

    def test_order_insensitive(self):
        base = R.needed_inputs(_plan_itc())
        for seed in range(5):
            assert R.needed_inputs(_shuffled(_plan_itc(), seed)) == base


class TestExclusion:
    def test_reasons(self):
        itc = _plan_itc()
        need = R.needed_inputs(itc)
        need["F() return"] = {"tier": "value_read", "cells": 3, "rows": [0], "first": (0, 0)}
        need["Cal() p[0]"] = {"tier": "value_read", "cells": 1, "rows": [0], "first": (0, 0)}
        need["u8g_Zero"] = {"tier": "value_read", "cells": 0, "rows": [], "first": (0, 0)}
        need["g_ptr"] = {"tier": "value_read", "cells": 1, "rows": [0], "first": (0, 0)}
        need["REG_X"] = {"tier": "value_read", "cells": 1, "rows": [0], "first": (0, 0)}
        need["g_flt"] = {"tier": "value_read", "cells": 1, "rows": [0], "first": (0, 0)}
        import copy
        gim = dict(copy.deepcopy(G.GLOBALS_INFO), REG_X={"type": "FSTATSTR", "array": ""},
                   g_flt={"type": "float", "base_type": "float", "array": ""})
        resolve = S._input_domain_resolver(G.function_details(), gim, copy.deepcopy(G.STRUCT_MEMBER_TYPES),
                                           G.uds_unit_io())
        kept, excluded = R.exclude_needed(need, itc, resolve)
        assert excluded == {
            "F() return": "stub_input", "Cal() p[0]": "stub_input",
            "tbl[1][2]": "subscript_shape", "arr[2].m": "subscript_shape",
            "k": "already_a_column", "REG_FSTAT.Bits.CCIF": "storage_root_already_set",
            "u8g_Zero": "blocks_no_cell", "g_ptr": "pointer_object", "REG_X": "struct_object",
            "o2b": "unknown_type:no_declaration", "g_flt": "float_not_ranged"}
        assert sorted(kept) == ["g_DoorState", "s16g_Pos", "u8g_Speed"]

    def test_structural_filter_without_a_resolver(self):
        need = {"F() return": {"tier": "value_read", "cells": 1, "rows": [0], "first": (0, 0)},
                "g": {"tier": "value_read", "cells": 1, "rows": [0], "first": (0, 0)}}
        kept, excluded = R.exclude_needed(need, {"input_vars": []})
        assert list(kept) == ["g"] and excluded == {"F() return": "stub_input"}

    def test_order_insensitive(self):
        resolve = _resolver()
        base = R.exclude_needed(R.needed_inputs(_plan_itc()), _plan_itc(), resolve)
        for seed in range(5):
            need = R.needed_inputs(_shuffled(_plan_itc(), seed))
            items = list(need.items())
            random.Random(seed).shuffle(items)
            assert R.exclude_needed(dict(items), _plan_itc(), resolve) == base


def _rec(tier, cells, first):
    return {"tier": tier, "cells": cells, "rows": [first[0]], "first": first}


class TestRanking:
    _NEED = {
        "a_decision": _rec("decision", 9, (0, 0)),
        "b_value": _rec("value_read", 1, (3, 0)),
        "c_value_more": _rec("value_read", 4, (5, 2)),
        "d_value_more_earlier": _rec("value_read", 4, (2, 7)),
        "e_value_more_same_row": _rec("value_read", 4, (2, 1)),
        "z_observed": _rec("observed_initial", 1, (9, 9)),
        "arr[10]": _rec("value_read", 1, (3, 0)),
        "arr[2]": _rec("value_read", 1, (3, 0)),
    }

    def test_tier_then_cells_then_first_then_root_numeric_subscript(self):
        assert R.rank_needed(self._NEED) == ["z_observed", "e_value_more_same_row", "d_value_more_earlier",
                                             "c_value_more", "arr[2]", "arr[10]", "b_value", "a_decision"]

    def test_shuffle_insensitive(self):
        base = R.rank_needed(self._NEED)
        for seed in range(10):
            items = list(self._NEED.items())
            random.Random(seed).shuffle(items)
            assert R.rank_needed(dict(items)) == base


class TestValueSelection:
    _FUNCS = ["Ap_Run", "Drv_Write", "Io_Set"]

    def test_values_are_the_mid_value_moved_by_rank_inside_the_domain(self):
        resolve = _resolver()
        ranked = ["u8g_Speed", "g_DoorState", "REG_FSTAT.Bits.CCIF", "s16g_Pos", "g_ptr"]
        values, excluded = R.select_values(ranked, resolve, self._FUNCS)
        # u8g_Speed: 설계서 10~20 의 가운데 15(넓은 축 첫째 → 그대로) · g_DoorState: 설계서가 좁힌 열거자 0 · 2 · 4 의 가운데 2
        # CCIF: 1 비트 0~1 의 0(좁은 범위는 옮기지 않는다) · s16g_Pos: 2253~2672 의 가운데 2462 + 넓은 축 셋째(2) = 2464
        assert values == {"u8g_Speed": 15, "g_DoorState": 2, "REG_FSTAT.Bits.CCIF": 0, "s16g_Pos": 2464}
        assert excluded == {"g_ptr": "pointer_object"}

    def test_fixed_once_and_deterministic(self):
        resolve = _resolver()
        need = {n: _rec("value_read", 1, (0, i)) for i, n in enumerate(["s16g_Pos", "u8g_Speed", "g_DoorState"])}
        first = R.select_values(R.rank_needed(need), resolve, self._FUNCS)
        for seed in range(5):
            items = list(need.items())
            random.Random(seed).shuffle(items)
            assert R.select_values(R.rank_needed(dict(items)), resolve, self._FUNCS) == first

    def test_rank_order_decides_the_offsets(self):
        """값은 순위 순서로 정한다 — 넓은 범위 둘의 순서를 바꾸면 옮김이 바뀐다(순위가 값에 닿는다는 근거)."""
        resolve = _resolver()
        a = R.select_values(["u8g_Speed", "s16g_Pos"], resolve, self._FUNCS)[0]
        b = R.select_values(["s16g_Pos", "u8g_Speed"], resolve, self._FUNCS)[0]
        assert a["s16g_Pos"] != b["s16g_Pos"] and a["u8g_Speed"] != b["u8g_Speed"]

    def test_empty(self):
        assert R.select_values([], _resolver()) == ({}, {})


@pytest.mark.parametrize("name, shape", [("g", True), ("g.a.b", True), ("a[3]", True), ("a[1][2]", False),
                                         ("a[1].m", False), ("F() return", False)])
def test_shape_rule_is_the_suts_read_rule(name, shape):
    assert R._shape_ok(name) is shape
