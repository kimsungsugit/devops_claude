"""R57 (N54) — 같은 이름의 정의가 둘(APP/FBL)이면 각자의 정의로, 정본에만 있는 heading 은 지어내지 않는다.

실측(run 2088 ↔ KJPDS02 정본 v3.03): 정본은 main·WriteBlock·EEPROM_SetByte … 9개 함수를 APP 절과 Bootloader 절(SwCom_35)에
따로 싣고 내용이 다른데, 생성본은 두 표 모두 APP 내용이었다 — 분석이 두 번째 정의를 파일째 버리고(first-wins) 문서는 이름 키
하나로 골랐기 때문. 그리고 payload 에 없는 정본 heading 36개엔 이름 낱말로 지어낸 동작 문장이 실려 있었다.
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
import pytest

from report_gen.docx_builder import (
    UNMATCHED_HEADING_NOTE,
    _candidates_by_name,
    _pick_function_candidate,
    _resolve_reference_target,
)

_APP_CALLS = ["s_System_InitSequence", "s_System_MainLoop"]
_FBL_CALLS = ["MainInit", "MainServiceLoop"]


def _fn(fid: str, name: str, calls, proto: str = "", file: str = "") -> dict:
    return {"id": fid, "name": name, "prototype": proto or f"void {name}(void)", "description": "설명", "asil": "QM",
            "related": "", "inputs": [], "outputs": [], "precondition": "", "globals_global": [], "globals_static": [],
            "logic": "", "calls_list": list(calls), "called": "\n".join(calls), "calling": "", "file": file}


# ────────────────────────────── 1. 분석: 교차 파일 쌍둥이는 둘 다 남는다 ──────────────────────────────
@pytest.fixture(scope="module")
def twin_sections(tmp_path_factory):
    root = tmp_path_factory.mktemp("twin")
    app = root / "APP" / "Sources"
    fbl = root / "FBL" / "Sources"
    app.mkdir(parents=True)
    fbl.mkdir(parents=True)
    (app / "SysOs_Main.c").write_text(
        "void s_System_InitSequence(void){}\nvoid s_System_MainLoop(void){}\n"
        "void main(void){ s_System_InitSequence(); s_System_MainLoop(); }\n"
        "void Twin_Helper(void){ s_System_InitSequence(); }\n"          # 이름 override 가 없는 쌍둥이(module_map 검사용)
        # 같은 파일 안의 같은 이름(#ifdef 변형이 preprocess=False 로 둘 다 파싱되는 경우의 모델) — 종전대로 하나만
        "void twin_same(void){ s_System_InitSequence(); }\nvoid twin_same(void){ s_System_MainLoop(); }\n",
        encoding="utf-8")
    (fbl / "main.c").write_text(
        "void MainInit(void){}\nvoid MainServiceLoop(void){}\nvoid main(void){ MainInit(); MainServiceLoop(); }\n"
        "void Twin_Helper(void){ MainInit(); }\n",
        encoding="utf-8")
    from report_gen.uds_generator import generate_uds_source_sections

    return generate_uds_source_sections(str(root), preprocess=False, max_files=20, max_items=100)


class TestSourceSectionsKeepCrossFileTwins:
    def test_two_definitions_two_ids_each_with_its_own_callees(self, twin_sections):
        fd = twin_sections["function_details"]
        mains = [v for v in fd.values() if str(v.get("name")).lower() == "main"]
        assert len(mains) == 2, [v.get("file") for v in fd.values()]
        assert len({v["id"] for v in mains}) == 2
        by_file = {Path(str(v["file"])).parts[-3]: sorted(c.lower() for c in v["calls_list"]) for v in mains}
        assert by_file["APP"] == sorted(c.lower() for c in _APP_CALLS)
        assert by_file["FBL"] == sorted(c.lower() for c in _FBL_CALLS)

    def test_same_file_duplicate_still_collapses_to_one(self, twin_sections):
        fd = twin_sections["function_details"]
        assert sum(1 for v in fd.values() if str(v.get("name")).lower() == "twin_same") == 1

    def test_by_name_keeps_the_first_definition(self, twin_sections):
        """이름 하나만 보는 소비자(영향분석 등)의 기본값은 예전(두 번째 정의 폐기)과 같아야 한다."""
        info = twin_sections["function_details_by_name"]["main"]
        assert "APP" in Path(str(info["file"])).parts

    def test_call_map_is_the_union_and_collisions_list_both_files(self, twin_sections):
        cm = {k.lower(): v for k, v in twin_sections["call_map"].items()}
        assert {c.lower() for c in cm["main"]} == {c.lower() for c in _APP_CALLS + _FBL_CALLS}
        files = twin_sections["function_collisions"]["main"]["files"]
        assert len(files) == 2 and any("APP" in Path(f).parts for f in files) and any("FBL" in Path(f).parts for f in files)


# ────────────────────────────── 2. 후보 고르기(단일 출처) ──────────────────────────────
class TestPickFunctionCandidate:
    def _cands(self):
        return [_fn("SwUFn_0101", "main", _APP_CALLS, file="APP/SysOs_Main.c"),
                _fn("SwUFn_3563", "main", _FBL_CALLS, file="FBL/main.c")]

    def test_reference_callee_overlap_decides_regardless_of_order(self):
        cands = self._cands()
        ref_fbl = {"name": "main", "called": "MainInit\nMainServiceLoop\nMainStartReprogSession"}
        assert _pick_function_candidate(cands, ref_fbl) is cands[1]
        assert _pick_function_candidate(list(reversed(cands)), ref_fbl) is cands[1]
        ref_app = {"name": "main", "called": "s_System_InitSequence"}
        assert _pick_function_candidate(list(reversed(cands)), ref_app) is cands[0]

    def test_prototype_breaks_a_callee_tie(self):
        a = _fn("A", "BackupSector", [], proto="static void BackupSector(void)")
        b = _fn("B", "BackupSector", [], proto="void BackupSector(byte sector)")
        assert _pick_function_candidate([a, b], {"name": "BackupSector", "called": "N/A", "prototype": "void BackupSector( byte sector )"}) is b

    def test_tie_goes_to_the_unconsumed_candidate_in_order(self):
        a = _fn("A", "BackupSector", [])
        b = _fn("B", "BackupSector", [])
        consumed: set = set()
        assert _pick_function_candidate([a, b], None, consumed) is a
        assert _pick_function_candidate([a, b], None, consumed) is b
        assert _pick_function_candidate([a, b], None, consumed) is a     # 셋째 heading — 다시 순서대로
        assert consumed == {id(a), id(b)}

    def test_single_candidate_and_empty(self):
        a = _fn("A", "x", [])
        assert _pick_function_candidate([a], {"name": "x", "called": "zzz"}) is a
        with pytest.raises(ValueError):
            _pick_function_candidate([], None)

    def test_candidates_by_name_groups_in_payload_order(self):
        fd = {"SwUFn_0101": _fn("SwUFn_0101", "main", _APP_CALLS), "SwUFn_0002": _fn("SwUFn_0002", "Other", []),
              "SwUFn_3563": _fn("SwUFn_3563", "main", _FBL_CALLS)}
        c = _candidates_by_name(fd)
        assert [x["id"] for x in c["main"]] == ["SwUFn_0101", "SwUFn_3563"] and len(c["other"]) == 1
        assert _candidates_by_name(None) == {}


# ────────────────────────────── 3. 정본 병합: 쌍둥이 블록은 각자의 정의로, 축 충돌 아님 ──────────────────────────────
def _stats():
    return {"by_name": 0, "by_name_and_id": 0, "id_collision_blocks": 0, "unmatched_blocks": 0, "unnamed_blocks": 0,
            "blocked_axes": {}, "ambiguous_names": 0, "ambiguous_sample": [], "twin_definition_blocks": 0}


class TestResolveReferenceTargetTwins:
    def test_two_blocks_two_targets_no_blocked_axes(self):
        app = _fn("SwUFn_0101", "main", _APP_CALLS)
        fbl = _fn("SwUFn_3563", "main", _FBL_CALLS)
        fd = {"SwUFn_0101": app, "SwUFn_3563": fbl}
        by_name = {"main": app}
        ref_map = {"SwUFn_0101": {"name": "main", "called": "\n".join(_APP_CALLS), "asil": "A"},
                   "SwUFn_3563": {"name": "main", "called": "\n".join(_FBL_CALLS), "asil": "B"}}
        by_ref = {"main": ["SwUFn_0101", "SwUFn_3563"]}
        st, seen, consumed, pairing = _stats(), set(), set(), {}
        cands = _candidates_by_name(fd)
        # 정본 블록 순서를 뒤집어 넣어도 각자의 정의를 찾는다
        t_fbl, bl_fbl = _resolve_reference_target("SwUFn_3563", ref_map["SwUFn_3563"], fd, by_name, by_ref, ref_map, st, seen,
                                                  candidates_by_name=cands, consumed=consumed, twin_pairing=pairing)
        t_app, bl_app = _resolve_reference_target("SwUFn_0101", ref_map["SwUFn_0101"], fd, by_name, by_ref, ref_map, st, seen,
                                                  candidates_by_name=cands, consumed=consumed, twin_pairing=pairing)
        assert t_fbl is fbl and t_app is app
        assert bl_fbl == set() and bl_app == set()          # ASIL A vs B 는 두 함수의 값이지 한 함수의 두 의견이 아니다
        assert st["twin_definition_blocks"] == 2 and st["ambiguous_names"] == 0
        assert pairing == {"main": {"SwUFn_0101": app, "SwUFn_3563": fbl}}     # heading 채움이 같은 표를 본다

    def test_without_a_pairing_table_the_axis_conflict_is_kept_conservatively(self):
        """짝 표 없이 부르면 distinct 를 알 수 없다 — 갈린 ASIL 은 종전대로 막는다(덮어쓰기보다 안전한 쪽)."""
        app = _fn("SwUFn_0101", "main", _APP_CALLS)
        fbl = _fn("SwUFn_3563", "main", _FBL_CALLS)
        fd = {"SwUFn_0101": app, "SwUFn_3563": fbl}
        ref_map = {"SwUFn_0101": {"name": "main", "called": "\n".join(_APP_CALLS), "asil": "A"},
                   "SwUFn_3563": {"name": "main", "called": "\n".join(_FBL_CALLS), "asil": "B"}}
        st = _stats()
        t, blocked = _resolve_reference_target("SwUFn_3563", ref_map["SwUFn_3563"], fd, {"main": app},
                                               {"main": ["SwUFn_0101", "SwUFn_3563"]}, ref_map, st, set(),
                                               candidates_by_name=_candidates_by_name(fd), consumed=set())
        assert t is fbl and "asil" in blocked

    def test_three_definitions_two_blocks_on_the_same_definition_is_not_a_split(self):
        """(리뷰 W3) 후보 3 · 블록 2 인데 두 블록이 같은 정의를 고르면 그건 한 함수의 두 의견 — 축 충돌을 본다."""
        from report_gen.docx_builder import _pair_twin_blocks, _twin_pairing_is_split

        a = _fn("A", "f", ["x", "y"])
        b = _fn("B", "f", ["p"])
        c = _fn("C", "f", ["q"])
        ref_map = {"R1": {"name": "f", "called": "x\ny"}, "R2": {"name": "f", "called": "x"}}
        pm = _pair_twin_blocks(["R1", "R2"], ref_map, [a, b, c])
        assert pm["R1"] is a and pm["R2"] is a
        assert _twin_pairing_is_split(pm) is False
        assert _twin_pairing_is_split({"R1": a, "R2": b}) is True
        assert _twin_pairing_is_split({"R1": a}) is False and _twin_pairing_is_split(None) is False

    def test_single_definition_with_two_disagreeing_blocks_still_blocks_the_axis(self):
        """대조군 — 정본이 한 함수를 두 절에 실은 경우(정의는 하나)는 종전대로 갈린 축을 막는다."""
        one = _fn("SwUFn_3369", "s_UDS_WDBI_ProcessDynamicPWMTable", ["ld_send_message"])
        fd = {"SwUFn_3369": one}
        ref_map = {"SwUFn_3369": {"name": one["name"], "called": "ld_send_message", "asil": "A"},
                   "SwUFn_3371": {"name": one["name"], "called": "ld_send_message", "asil": "B"}}
        by_ref = {one["name"].lower(): ["SwUFn_3369", "SwUFn_3371"]}
        st = _stats()
        t, blocked = _resolve_reference_target("SwUFn_3369", ref_map["SwUFn_3369"], fd, {one["name"].lower(): one}, by_ref, ref_map,
                                               st, set(), candidates_by_name=_candidates_by_name(fd), consumed=set())
        assert t is one and "asil" in blocked and st["twin_definition_blocks"] == 0

    def test_default_arguments_keep_the_old_signature_working(self):
        one = _fn("SwUFn_0001", "f", [])
        st = _stats()
        t, blocked = _resolve_reference_target("SwUFn_0001", {"name": "f"}, {"SwUFn_0001": one}, {"f": one}, {"f": ["SwUFn_0001"]},
                                               {}, st, set())
        assert t is one and blocked == set() and st["by_name_and_id"] == 1


# ────────────────────────────── 4. 문서 끝까지: 쌍둥이 heading 두 표의 내용이 다르다 ──────────────────────────────
class TestGeneratedTwinTables:
    @pytest.fixture(autouse=True)
    def _no_template_resolution(self, monkeypatch, tmp_path):
        import config

        monkeypatch.setattr(config, "resolve_uds_template_path", lambda: "", raising=False)
        monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(tmp_path / "no_such_reference.docx"), raising=False)

    @staticmethod
    def _payload():
        fd = {
            "SwUFn_0101": _fn("SwUFn_0101", "main", _APP_CALLS, file="APP/SysOs_Main.c"),
            "SwUFn_0102": _fn("SwUFn_0102", "s_System_InitSequence", [], file="APP/SysOs_Main.c"),
            "SwUFn_3563": _fn("SwUFn_3563", "main", _FBL_CALLS, file="FBL/main.c"),
            "SwUFn_3501": _fn("SwUFn_3501", "MainInit", [], file="FBL/main.c"),
        }
        return {"project_name": "T", "overview": "o", "requirements": "r", "interfaces": "i", "uds_frames": "u", "notes": "n",
                "function_details": fd, "function_details_by_name": {"main": fd["SwUFn_0101"], "s_system_initsequence": fd["SwUFn_0102"],
                                                                     "maininit": fd["SwUFn_3501"]},
                "call_map": {"main": _APP_CALLS + _FBL_CALLS}}

    @staticmethod
    def _template(tmp_path, order):
        tpl = docx.Document()
        tpl.add_heading("Software Unit Design", level=1)
        for swcom, items in order:
            tpl.add_heading(swcom, level=2)
            for fid, n in items:
                tpl.add_heading(f"{fid}: {n}", level=4)
                t = tpl.add_table(rows=3, cols=6)
                t.cell(0, 0).text = "[ Function Information ]"
        p = tmp_path / "tpl.docx"
        tpl.save(str(p))
        return str(p)

    @staticmethod
    def _main_tables(out: Path):
        from report_gen.validation import read_back_function_info

        return {str(v.get("id") or ""): v for v in read_back_function_info(str(out)).values() if str(v.get("name")).lower() == "main"}

    def _generate(self, tmp_path, order, name="u"):
        from report_gen.docx_builder import generate_uds_docx

        out = tmp_path / f"{name}.docx"
        generate_uds_docx(self._template(tmp_path, order), self._payload(), str(out))
        return out

    def test_without_reference_the_two_headings_get_the_two_definitions_in_order(self, tmp_path):
        order = [("SwCom_01 (System OS)", [("SwUFn_0101", "main"), ("SwUFn_0102", "s_System_InitSequence")]),
                 ("SwCom_35 (Bootloader)", [("SwUFn_3501", "MainInit"), ("SwUFn_3563", "main")])]
        mains = self._main_tables(self._generate(tmp_path, order))
        assert set(mains) == {"SwUFn_0101", "SwUFn_3563"}, mains.keys()
        assert mains["SwUFn_0101"]["called"] == "\n".join(_APP_CALLS)
        assert mains["SwUFn_3563"]["called"] == "\n".join(_FBL_CALLS)

    def test_reference_block_decides_even_against_heading_order(self, tmp_path, monkeypatch):
        """정본이 있으면 heading 순서가 아니라 **그 heading 의 정본 블록** 피호출자로 고른다 — Bootloader 절이 앞에 와도."""
        import config
        import report_gen.docx_builder as db

        tpl_order = [("SwCom_35 (Bootloader)", [("SwUFn_3563", "main"), ("SwUFn_3501", "MainInit")]),
                     ("SwCom_01 (System OS)", [("SwUFn_0101", "main"), ("SwUFn_0102", "s_System_InitSequence")])]
        ref_map = {"SwUFn_3563": {"name": "main", "called": "\n".join(_FBL_CALLS), "prototype": "void main(void)"},
                   "SwUFn_0101": {"name": "main", "called": "\n".join(_APP_CALLS), "prototype": "void main(void)"}}
        tpl = self._template(tmp_path, tpl_order)
        monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", tpl, raising=False)       # 실재 파일이면 정본 경로가 열린다
        monkeypatch.setattr(db, "_extract_function_info_from_docx", lambda _doc: dict(ref_map))
        out = tmp_path / "ref.docx"
        db.generate_uds_docx(tpl, self._payload(), str(out))
        mains = self._main_tables(out)
        assert mains["SwUFn_3563"]["called"] == "\n".join(_FBL_CALLS)
        assert mains["SwUFn_0101"]["called"] == "\n".join(_APP_CALLS)

    def test_a_reference_block_of_another_function_does_not_steer_the_pick(self, tmp_path, monkeypatch):
        """(리뷰 W1) 정본 ID 만 같은 **남의 함수** 블록(이름이 다름)은 쌍둥이 선택을 조종하지 못한다 — 그 heading 은 짝 표가
        다른 heading 에 준 정의를 피해 나머지 정의를 받는다."""
        import config
        import report_gen.docx_builder as db

        tpl_order = [("SwCom_35 (Bootloader)", [("SwUFn_3563", "main"), ("SwUFn_3501", "MainInit")]),
                     ("SwCom_01 (System OS)", [("SwUFn_0101", "main"), ("SwUFn_0102", "s_System_InitSequence")])]
        # 3563 블록은 이름이 'Other' 인데 피호출자가 APP 쪽 — 가드가 없으면 Bootloader heading 이 APP 정의를 받고 두 표가 같아진다
        ref_map = {"SwUFn_3563": {"name": "Other", "called": "\n".join(_APP_CALLS)},
                   "SwUFn_0101": {"name": "main", "called": "\n".join(_APP_CALLS)}}
        tpl = self._template(tmp_path, tpl_order)
        monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", tpl, raising=False)
        monkeypatch.setattr(db, "_extract_function_info_from_docx", lambda _doc: dict(ref_map))
        out = tmp_path / "guard.docx"
        db.generate_uds_docx(tpl, self._payload(), str(out))
        mains = self._main_tables(out)
        assert mains["SwUFn_0101"]["called"] == "\n".join(_APP_CALLS)
        assert mains["SwUFn_3563"]["called"] == "\n".join(_FBL_CALLS)

    def test_twin_with_no_callees_is_not_recovered_from_the_shared_call_map(self, tmp_path):
        """(리뷰 I1) 이름 키 `call_map` 은 쌍둥이의 합집합 — 피호출자가 없는 쪽 정의에 상대 정의의 피호출자를 복구해 싣지 않는다."""
        from report_gen.docx_builder import generate_uds_docx

        payload = self._payload()
        payload["function_details"]["SwUFn_3563"]["calls_list"] = []
        payload["function_details"]["SwUFn_3563"]["called"] = ""
        order = [("SwCom_01 (System OS)", [("SwUFn_0101", "main"), ("SwUFn_0102", "s_System_InitSequence")]),
                 ("SwCom_35 (Bootloader)", [("SwUFn_3501", "MainInit"), ("SwUFn_3563", "main")])]
        out = tmp_path / "norecover.docx"
        generate_uds_docx(self._template(tmp_path, order), payload, str(out))
        mains = self._main_tables(out)
        assert mains["SwUFn_0101"]["called"] == "\n".join(_APP_CALLS)
        assert mains["SwUFn_3563"]["called"] == "N/A"


class TestNameKeyedMapsKeepTheFirstDefinition:
    def test_module_map_is_first_wins_like_by_name(self, twin_sections):
        """(리뷰 W4) `module_map` 도 이름 키 — 쌍둥이의 뒤 정의가 앞 정의의 모듈을 덮지 않는다."""
        mm = twin_sections["module_map"]
        # `main` 은 이름 override(docs/uds_function_swcom_override.json)로 라벨이 고정되므로 override 없는 쌍둥이로 본다
        assert mm["Twin_Helper"] == mm["s_System_InitSequence"]        # APP 형제와 같은 모듈
        assert mm["Twin_Helper"] != mm["MainInit"]                      # FBL 형제와 다르다


class TestTwinCounterIsDisclosed:
    _REF = {"configured": True, "document": "X_v3.03.docx", "identity": {"same_project": True}, "safety_fields_overridden": {}}

    def test_twin_definition_blocks_become_a_risk_item(self):
        """(리뷰 I2) 쌍둥이 9개는 이제 `ambiguous_names` 에서 빠진다 — 대체 공시 없이 사라지면 안 된다."""
        from backend.helpers import uds as U
        from report_gen.gen_issues import IssueCollector

        c = IssueCollector()
        U._note_docx_outcome(c, {"reference_suds": {**self._REF, "matching": {
            "ambiguous_names": 0, "id_collision_blocks": 0, "twin_definition_blocks": 18}}})
        items = {i["code"]: i for i in c.as_list()}
        assert list(items) == ["reference_twin_definitions"]
        it = items["reference_twin_definitions"]
        assert (it["severity"], it["kind"]) == ("risk", "actual")
        assert "18개" in it["message"] and it["facts"] == {"twin_definition_blocks": 18}

    @pytest.mark.parametrize("value", [0, None, "18", True])
    def test_zero_missing_or_non_int_makes_no_item(self, value):
        from backend.helpers import uds as U
        from report_gen.gen_issues import IssueCollector

        c = IssueCollector()
        m = {"ambiguous_names": 0, "id_collision_blocks": 0}
        if value is not None:
            m["twin_definition_blocks"] = value
        U._note_docx_outcome(c, {"reference_suds": {**self._REF, "matching": m}})
        assert c.as_list() == []

    def test_explainer_has_an_action_for_the_code(self):
        from backend.services import issue_explainer as IE

        assert IE._RULE_ACTION["reference_twin_definitions"]


# ────────────────────────────── 5. 정본에만 있는 heading — 지어내지 않는다 ──────────────────────────────
class TestUnmatchedHeadingIsNotFabricated:
    @pytest.fixture(autouse=True)
    def _isolate(self, monkeypatch, tmp_path):
        import config

        monkeypatch.setattr(config, "resolve_uds_template_path", lambda: "", raising=False)
        monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(tmp_path / "no_such_reference.docx"), raising=False)

    def test_shell_table_states_the_fact_and_invents_no_prototype(self, tmp_path):
        from report_gen.docx_builder import generate_uds_docx
        from report_gen.validation import read_back_function_info

        tpl = docx.Document()
        tpl.add_heading("Software Unit Design", level=1)
        for fid, n in (("SwUFn_0001", "Top_Init"), ("SwUFn_0002", "s_Ap_Diagnostic"), ("SwUFn_0003", "main")):
            tpl.add_heading(f"{fid}: {n}", level=4)
            t = tpl.add_table(rows=3, cols=6)
            t.cell(0, 0).text = "[ Function Information ]"
        tpl_path = tmp_path / "tpl.docx"
        tpl.save(str(tpl_path))
        fd = {"SwUFn_0001": _fn("SwUFn_0001", "Top_Init", [])}
        payload = {"project_name": "T", "overview": "o", "requirements": "r", "interfaces": "i", "uds_frames": "u", "notes": "n",
                   "function_details": fd, "function_details_by_name": {"top_init": fd["SwUFn_0001"]}, "call_map": {}}
        out = tmp_path / "u.docx"
        generate_uds_docx(str(tpl_path), payload, str(out))
        info = {str(v.get("name")).lower(): v for v in read_back_function_info(str(out)).values()}
        shell = info["s_ap_diagnostic"]
        assert shell["description"] == UNMATCHED_HEADING_NOTE
        assert "모듈에서" not in shell["description"] and "진단" not in shell["description"]
        assert shell["prototype"] == "" and shell["called"] == "N/A"
        assert info["main"]["prototype"] == ""                   # `void main( void )` 리터럴 없음
        assert info["top_init"]["description"].rstrip(".") == "설명"     # 분석된 함수는 그대로(마침표는 라이터 규약)

    def test_writer_pairs_do_not_invent_main_prototype(self):
        from report_gen.function_analyzer import _function_info_pairs

        pairs = dict(_function_info_pairs({"name": "main", "prototype": "", "description": "d"}))
        assert pairs.get("Prototype", "") in ("", "N/A"), pairs.get("Prototype")

    def test_note_is_not_a_generic_phrase_the_builder_would_replace(self):
        from report_gen.function_analyzer import _is_generic_description

        assert not _is_generic_description(UNMATCHED_HEADING_NOTE)


# ────────────────────────────── 6. 계수 키가 증거 요약까지 간다 ──────────────────────────────
class TestTwinCounterReachesEvidence:
    def test_matching_keys_include_twin_definition_blocks(self):
        from report_gen.evidence import _MATCHING_KEYS, _matching_summary

        assert "twin_definition_blocks" in _MATCHING_KEYS
        assert _matching_summary({"twin_definition_blocks": 2})["twin_definition_blocks"] == 2

    def test_builder_initialises_the_counter(self):
        src = Path("report_gen/docx_builder.py").read_text(encoding="utf-8")
        assert re.search(r'"twin_definition_blocks":\s*0', src)
