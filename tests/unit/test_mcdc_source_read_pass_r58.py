"""R58 — the extended SUTS profile designs MC/DC again over the inputs the source reads (`generators/suts.py`).

The base MC/DC design runs on the design input list (R19): its vectors fill the reference-size slots, so the extended
document contains the reference one (R75). A decision reading an input the source reads but the design table omits
(`g_b` below) got no pair there. The extended profile now designs the decisions the base design left *without any pair*
once more, over every row input (`_mcdc_source_read_pass`), and appends the new vectors after the base ones. Base
slots, base decisions and their labels are unchanged; each design is revalidated on its own unit and its own rows.
Measured: plan record R58 (`docs/plans/2026-09-21_test_generation_superiority.md`)."""
from __future__ import annotations

import json
import os
import sys

import openpyxl
from openpyxl.styles import Border, Font, PatternFill

from generators import suts
from generators.mcdc_design import build_mcdc_design
from generators.suts import generate_sequences, summarize_mcdc_design

sys.path.insert(0, os.path.dirname(__file__))
from test_suts_source_read_inputs import COMMON, ROOT, UNIT, _complete, _unit  # noqa: E402

from generators import c_project_context as cpc  # noqa: E402

_LABEL = "소스가 읽어 더한 입력까지 쓴 설계(확장)"
# (review R58 W3) one function with both designs: D1 reads the design input only (base design), D2 reads the added
#   g_b (second design); `never` reads g_b in a contradiction (no pair on any input list)
EXTRA = """
void both(void) { if (g_a > 10U) { g_o = 1U; } else { g_o = 3U; } if ((g_a > 5U) && (g_b > 3U)) { g_o = g_gain_close; } else { g_o = 0U; } }
void never(void) { if ((g_b > 3U) && (g_b < 2U)) { g_o = 1U; } else { g_o = 0U; } }
void cmp(void) { if (g_b < g_a) { g_o = 1U; } else { g_o = 0U; } }
void huge(void) {
 if (g_a > 1U) { g_o = 1U; } if (g_a > 20U) { g_o = 2U; } if (g_a > 40U) { g_o = 3U; } if (g_a > 60U) { g_o = 4U; }
 if (g_a > 80U) { g_o = 5U; } if (g_a > 100U) { g_o = 6U; } if (g_a > 120U) { g_o = 8U; } if (g_a > 140U) { g_o = 9U; }
 if (g_b > g_a) { g_o = 7U; } if ((g_a == 0U) && (g_b == 4U)) { g_o = 11U; } }
"""


def _unit2(name, inputs):
    files = {os.path.join(ROOT, "common.h"): COMMON, os.path.join(ROOT, "unit.c"): UNIT + EXTRA}
    context = cpc.build_project_context(files)
    path = os.path.join(ROOT, "unit.c")
    return {"name": name, "fid": "F1", "source_text": UNIT + EXTRA, "source_path": path, "source_text_complete": True,
            "project_scope": cpc.build_scopes(context, [path])[path], "input_vars": list(inputs),
            "output_vars": ["g_o"], "param_types": {"g_a": "U8", "g_b": "U8", "g_o": "U8", "g_gain_open": "U8"}}


def _first_line(s):
    return s["description"].split("\n")[0]


def _two():
    ref = _unit("two", ["g_a"])
    ref_rows = generate_sequences(ref, None, type_cache={})
    unit = _unit("two", ["g_a"])
    rows = _complete(unit)
    return ref, ref_rows, unit, rows


def test_a_decision_on_an_added_input_is_designed_over_every_input():
    _ref, _ref_rows, unit, _rows = _two()
    (d,) = unit["mcdc_design"]["decisions"]
    assert d["status"] == "designed" and d["reason"] == "unique_cause_pairs_found"
    assert d["input_scope"] == "with_source_read_inputs"
    # what the base design said stays visible
    assert d["design_input_reason"] == "decision_variable_not_in_unit_inputs:g_b"
    assert [p["retained_status"] for p in d["pairs"]] == ["retained", "retained"]
    rec = unit["mcdc_design"]["source_read_pass"]
    assert {k: rec[k] for k in ("decisions_searched", "decisions_adopted", "decisions_designed", "vectors",
                                "designed_refused_for_added")} == {"decisions_searched": 1, "decisions_adopted": 1,
                                                                   "decisions_designed": 1, "vectors": 3,
                                                                   "designed_refused_for_added": 1}
    assert d["design_refused_for_added_input"] is True
    assert d["source_read_inputs_moved"] == ["g_b"]
    assert unit["mcdc_design"]["summary"]["retained_pairs"] == 2


def test_the_reference_slots_stay_and_the_new_vectors_come_after_them():
    ref, ref_rows, unit, rows = _two()
    n = unit["base_strategy_count"]
    assert n == ref["base_strategy_count"]
    # design-column values as in the reference document (added inputs are fixed values there — R19)
    assert [(s["strategy"], {k: s["inputs"].get(k) for k in r["inputs"]}) for s, r in zip(rows[:n], ref_rows,
                                                                                       strict=True)] == \
        [(r["strategy"], r["inputs"]) for r in ref_rows]
    mcdc = [s for s in rows if s["strategy"].startswith("MCDC_")]
    # (review R58 W1) the reference slot names MCDC_0..6 stay the base design's — the new vectors start at MCDC_7
    assert [s["strategy"] for s in mcdc] == ["MCDC_7", "MCDC_8", "MCDC_9"]
    assert all(s.get("tc_profile") == suts.TC_PROFILE_EXTENDED and suts.is_extended_strategy(s["strategy"])
               for s in mcdc)
    assert all(rows.index(s) >= n for s in mcdc)
    # the base design had no vector: the reference document has no MC/DC row for this decision
    assert ref["mcdc_design"]["selected_inputs"] == []
    assert unit["mcdc_design"]["selected_inputs"] == [s["inputs"] for s in mcdc]


def test_every_new_row_is_a_pair_member_and_says_where_its_design_came_from():
    _ref, _ref_rows, unit, rows = _two()
    for s in rows:
        if s["strategy"].startswith("MCDC_"):
            assert s["mcdc_design"], s["strategy"]
            assert _LABEL in s["description"].split("\n")[0]
            assert set(s["inputs"]) == {"g_a", "g_b"}
    pairs = unit["mcdc_design"]["decisions"][0]["pairs"]
    by_seq = {s["seq_num"]: s for s in rows}
    for p in pairs:
        assert by_seq[p["seq_a"]]["inputs"] == p["inputs_a"] and by_seq[p["seq_b"]]["inputs"] == p["inputs_b"]


def test_base_rows_keep_their_label_when_the_function_has_both_designs():
    # `close_gain` reads only g_a in its decision: designed by the base design, nothing for the second design
    ref = _unit("close_gain", ["g_a"])
    ref_rows = generate_sequences(ref, None, type_cache={})
    unit = _unit("close_gain", ["g_a"])
    rows = _complete(unit)
    base = [s for s in rows if s["strategy"].startswith("MCDC_")]
    assert base and not any(_LABEL in s["description"] for s in base)
    ref_mcdc = [s for s in ref_rows if s["strategy"].startswith("MCDC_")]
    # (the extended label also names the blank added inputs — R81 — before which it is the reference label)
    assert [s["description"].split("\n")[0].split(" · 설계 밖 입력")[0] for s in base[:len(ref_mcdc)]] == \
        [s["description"].split("\n")[0] for s in ref_mcdc]
    assert [s["inputs"] for s in base[:len(ref_mcdc)]] == [s["inputs"] for s in ref_mcdc]
    assert unit["mcdc_design"]["source_read_pass"]["decisions_searched"] == 0
    assert all(d.get("input_scope") is None for d in unit["mcdc_design"]["decisions"])


def test_a_function_with_both_designs_keeps_the_base_rows_as_the_reference_has_them():
    # review R58 W3 (mutants M1 · M9: a finalize call given the other design's rows empties their pair records)
    ref = _unit2("both", ["g_a"])
    ref_rows = generate_sequences(ref, None, type_cache={})
    unit = _unit2("both", ["g_a"])
    rows = _complete(unit)
    d1, d2 = unit["mcdc_design"]["decisions"]
    assert d1.get("input_scope") is None and d1["status"] == "designed"
    assert d2["input_scope"] == "with_source_read_inputs" and d2["status"] == "designed"
    ref_mcdc = [s for s in ref_rows if s["strategy"].startswith("MCDC_")]
    base = [s for s in rows if s["strategy"].startswith("MCDC_") and not suts.is_extended_strategy(s["strategy"])]
    assert [s["strategy"] for s in base] == [s["strategy"] for s in ref_mcdc]
    assert [s["inputs"] for s in base] == [s["inputs"] for s in ref_mcdc]
    assert [s["mcdc_design"] for s in base] == [s["mcdc_design"] for s in ref_mcdc]
    assert [_first_line(s).split(" · 설계 밖 입력")[0] for s in base] == [_first_line(s) for s in ref_mcdc]
    new = [s for s in rows if s["strategy"].startswith("MCDC_") and suts.is_extended_strategy(s["strategy"])]
    assert new and all(s["mcdc_design"] and _LABEL in _first_line(s) for s in new)
    assert {r["decision_id"] for s in new for r in s["mcdc_design"]} == {d2["decision_id"]}
    assert {r["decision_id"] for s in base for r in s["mcdc_design"]} == {d1["decision_id"]}
    assert [p["retained_status"] for d in (d1, d2) for p in d["pairs"]] == ["retained"] * (len(d1["pairs"]) + len(d2["pairs"]))


def test_more_than_seven_base_vectors_keep_their_pairs_next_to_a_second_design():
    # review R58 W-B (mutant MA: splitting at the slot count instead of the base vector count made the base pairs whose
    #   vectors sit in extended slots `truncated` and labelled their rows as second-design rows)
    unit = _unit2("huge", ["g_a"])
    rows = _complete(unit)
    base_n = len([d for d in unit["mcdc_design"]["decisions"] if d.get("input_scope") is None])
    second = [d for d in unit["mcdc_design"]["decisions"] if d.get("input_scope")]
    assert second and base_n >= 8
    pairs = [p for d in unit["mcdc_design"]["decisions"] for p in d["pairs"]]
    assert pairs and all(p["retained_status"] == "retained" for p in pairs)
    mcdc = [s for s in rows if s["strategy"].startswith("MCDC_")]
    n_base_vectors = len(unit["mcdc_design"]["selected_inputs"]) - sum(
        len({json.dumps(p[f"inputs_{side}"], sort_keys=True) for p in d["pairs"] for side in "ab"}) for d in second)
    assert n_base_vectors > suts._BASE_MCDC_SLOTS
    for s in mcdc:
        is_second = any(r["decision_id"] in {d["decision_id"] for d in second} for r in s["mcdc_design"])
        assert (_LABEL in _first_line(s)) == is_second, s["strategy"]


def test_a_pair_on_a_design_input_alone_is_still_counted_as_refused_for_the_added_one():
    # review R58 W-A: `g_b < g_a` — the pair moves g_a only, but exists because g_b has a value in the row
    unit = _unit2("cmp", ["g_a"])
    _complete(unit)
    (d,) = unit["mcdc_design"]["decisions"]
    assert d["status"] == "designed" and d["source_read_inputs_moved"] == []
    assert d["design_refused_for_added_input"] is True
    assert unit["mcdc_design"]["source_read_pass"]["designed_refused_for_added"] == 1


def test_a_decision_refused_for_an_added_input_takes_the_second_design_s_reason():
    # review R58 W3 (mutant M8): the contradiction has no pair on any input list — the refusal "reads an added input"
    #   is no longer true once the second design read it; its own reason is the one to show
    unit = _unit2("never", ["g_a"])
    _complete(unit)
    (d,) = unit["mcdc_design"]["decisions"]
    assert d["input_scope"] == "with_source_read_inputs" and not d["pairs"]
    assert d["design_input_reason"] == "decision_variable_not_in_unit_inputs:g_b"
    assert not str(d["reason"]).startswith(("decision_variable_not_in_unit_inputs", "decision_reads_source_read_input"))
    rec = unit["mcdc_design"]["source_read_pass"]
    # (review R58 round 3 W-1) adopted but not designed: every count says so, and the sentence names it
    assert (rec["decisions_adopted"], rec["decisions_designed"], rec["designed_refused_for_added"]) == (1, 0, 0)
    from report_gen.generation_disclosures import build_disclosures
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": summarize_mcdc_design([unit])})
               if i["key"] == "suts_mcdc_design"]
    assert "설계까지 못 간 채택 1" in item["note"] and "거절했던 결정 0" not in item["note"]
    assert "가르지 않는다" not in item["note"]


def test_the_reference_profile_has_no_second_design():
    ref = _unit("two", ["g_a"])
    generate_sequences(ref, None, type_cache={})
    assert "source_read_pass" not in ref["mcdc_design"]
    assert ref["mcdc_design"]["decisions"][0]["status"] == "unsupported"


def test_a_decision_with_a_pair_is_not_designed_again(monkeypatch):
    # a partial decision keeps its pair: designing it again would change the label of a reference-region row
    seen = {}
    real = suts.build_mcdc_design

    def spy(unit, **kw):
        seen["only"] = kw.get("only_occurrences")
        return real(unit, **kw)
    monkeypatch.setattr(suts, "build_mcdc_design", spy)
    base = {"decisions": [
        {"occurrence_id": "x:1", "status": "partial", "pairs": [{"pair_id": "D1:C1:P1"}], "reason": ""},
        {"occurrence_id": "x:2", "status": "no_pair_found", "pairs": [], "reason": "no_pair_in_candidate_domain"},
        {"occurrence_id": "x:3", "status": "designed", "pairs": [{}], "reason": "unique_cause_pairs_found"},
        {"occurrence_id": "x:4", "status": "unenumerated", "pairs": [], "reason": "source_unavailable"}]}
    suts._mcdc_source_read_pass(_unit("two", ["g_a", "g_b"]), base, lambda: {}, ["g_b"])
    # only the pairless, enumerated decisions are searched again (review R58 I8: the ids here are synthetic, so this
    #   is the whole claim — the real adoption is in the `both` / `never` tests)
    assert seen["only"] == {"x:2"}
    # on a real function: the base-designed decision is not searched again
    seen.clear()
    unit = _unit2("both", ["g_a"])
    _complete(unit)
    d1, d2 = unit["mcdc_design"]["decisions"]
    assert seen["only"] == {d2["occurrence_id"]}


def test_only_occurrences_lists_the_other_decisions_without_a_search():
    unit = _unit("chain", ["g_a", "g_step"])
    full = build_mcdc_design(dict(unit))
    first = full["decisions"][0]["occurrence_id"]
    part = build_mcdc_design(dict(unit), only_occurrences={first})
    assert [d["status"] for d in part["decisions"][1:]] == ["not_requested"] * (len(part["decisions"]) - 1)
    assert part["decisions"][0]["status"] == full["decisions"][0]["status"]
    assert part["decisions"][0]["pairs"] == full["decisions"][0]["pairs"]


def test_the_summary_counts_the_second_design_and_its_failures(monkeypatch):
    _ref, _ref_rows, unit, _rows = _two()
    s = summarize_mcdc_design([unit])
    assert (s["source_read_pass_decisions"], s["source_read_pass_adopted"], s["source_read_pass_designed"],
            s["source_read_pass_errors"]) == (1, 1, 1, 0)
    assert s["designed"] == 1

    def fail(*_a, **_k):
        raise RuntimeError("down")
    monkeypatch.setattr(suts, "_mcdc_source_read_pass", fail)
    broken = _unit("two", ["g_a"])
    _complete(broken)
    s = summarize_mcdc_design([broken])
    assert s["source_read_pass_errors"] == 1 and s["designed"] == 0 and s["unsupported"] == 1
    # (review R58 W2) a failure alone is disclosed — not only a warning tone
    from report_gen.generation_disclosures import build_disclosures
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": s}) if i["key"] == "suts_mcdc_design"]
    assert "2차 설계 실패로 기본 설계를 유지한 함수 1" in item["note"] and item["tone"] == "warning"


def test_a_failure_inside_the_second_design_keeps_what_it_searched(monkeypatch):
    # review R58 W2: the record says which decisions were to be designed again, plus the error
    real = suts.build_mcdc_design

    def fail_second(unit, **kw):
        if kw.get("only_occurrences") is not None:
            raise RuntimeError("down")
        return real(unit, **kw)
    monkeypatch.setattr(suts, "build_mcdc_design", fail_second)
    unit = _unit("two", ["g_a"])
    _complete(unit)
    rec = unit["mcdc_design"]["source_read_pass"]
    assert rec["decisions_searched"] == 1 and rec["error"] == "RuntimeError" and rec["decisions_adopted"] == 0
    assert unit["mcdc_design"]["decisions"][0]["status"] == "unsupported"
    s = summarize_mcdc_design([unit])
    assert (s["source_read_pass_decisions"], s["source_read_pass_errors"]) == (1, 1)


def test_the_disclosure_says_what_the_second_design_did():
    from report_gen.generation_disclosures import build_disclosures
    _ref, _ref_rows, unit, _rows = _two()
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": summarize_mcdc_design([unit])})
               if i["key"] == "suts_mcdc_design"]
    for text in ("쌍을 하나도 못 만든 결정 1", "결과로 바꾼 결정 1", "그중 설계 1", "'더한 입력을 읽는다' 로 거절했던 결정 1",
                 "Input Scope", "부분 설계", "입력 열이 더해질 수 있고"):
        assert text in item["note"], text
    # nothing designed for another reason — that clause is absent (round 3 W-1)
    assert "가르지 않는다" not in item["note"] and "설계까지 못 간" not in item["note"]
    ref = _unit("two", ["g_a"])
    generate_sequences(ref, None, type_cache={})
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": summarize_mcdc_design([ref])})
               if i["key"] == "suts_mcdc_design"]
    assert "소스가 읽어 더한 입력까지" not in item["note"]


def test_the_mcdc_design_sheet_names_the_input_scope():
    _ref, _ref_rows, unit, rows = _two()
    base = _unit("close_gain", ["g_a"])
    base["fid"] = "F2"
    base_rows = _complete(base)
    wb = openpyxl.Workbook()
    suts._write_mcdc_design_sheet(wb, [unit, base], {"F1": rows, "F2": base_rows}, {"F1": "TC_1", "F2": "TC_2"},
                                  Border(), PatternFill(), Font(), Font())
    ws = wb["MCDC Design"]
    values = list(ws.iter_rows(values_only=True))
    header = list(values[0])
    assert "Input Scope" in header and header == suts._MCDC_HEADERS
    col = header.index("Input Scope")
    scope = {(r[0], r[2], r[4]): r[col] for r in values[1:]}
    assert {v for (tc, _d, _p), v in scope.items() if tc == "TC_1"} == {"with_source_read_inputs"}
    assert {v for (tc, _d, _p), v in scope.items() if tc == "TC_2"} == {"design_inputs"}
