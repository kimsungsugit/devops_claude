"""R12 — change impact on a generated SUTS (`scripts/change_impact.py`): every derived row is re-derived on the current
tree; a changed expectation is the regression impact, and no changed behaviour may go unflagged — a header change
included (review round 1 C1)."""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

import change_impact  # noqa: E402

# U16/S16 as ``unsigned int``/``signed int`` testify the 16-bit target int the oracle needs (as in the R3 tests)
COMMON = "typedef unsigned char U8;\ntypedef unsigned int U16;\ntypedef signed int S16;\n"
UNIT = """#include "common.h"
U8 g_i;
U8 g_o;
U8 g_p;
void f(void) { g_o = (U8)(g_i + 1U); }
void k(void) { g_p = g_i; }
void gone(void) { g_p = 7U; }
"""
OTHER = '#include "common.h"\nU8 g_q;\nvoid h(void) { g_q = 3U; }\n'
HEADER = ["Function ID", "Function", "Sequence", "Observable", "Expected", "Status", "Oracle", "Execution",
          "Source SHA256", "Source path", "Reason", "Test Case ID", "Source hash scope", "Inputs JSON", "Basis",
          "Assumptions"]


def _tree(tmp_path):
    root = tmp_path / "src"
    root.mkdir()
    (root / "common.h").write_text(COMMON, encoding="utf-8", newline="\n")
    (root / "unit.c").write_text(UNIT, encoding="utf-8", newline="\n")
    (root / "other.c").write_text(OTHER, encoding="utf-8", newline="\n")
    return root


def _suts(tmp_path, root):
    rows = [("f", "g_o", 2, {"g_i": 1}, "unit.c"), ("f", "g_o", 6, {"g_i": 5}, "unit.c"),
            ("k", "g_p", 9, {"g_i": 9}, "unit.c"), ("gone", "g_p", 7, {}, "unit.c"), ("h", "g_q", 3, {}, "other.c")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Evidence"
    ws.append(HEADER)
    for i, (fn, obs, val, inputs, file) in enumerate(rows):
        path = root / file
        sha = hashlib.sha256(path.read_bytes().decode("utf-8").encode()).hexdigest()
        ws.append([f"SwUFn_{fn}", fn, i + 1, obs, val, "derived", "source", "not_run", sha, str(path.resolve()), "",
                   f"SwUTC_{fn}", "captured_decoded_utf8_text", json.dumps(inputs), "assigned", ""])
    ws.append(["SwUFn_x", "x", 9, "g_o", "[검증 필요] call_return_value:y", "unknown", "none", "not_run", "",
               str((root / "unit.c").resolve()), "call_return_value:y", "SwUTC_x", "", "{}", "", ""])
    out = tmp_path / "old.xlsm"
    wb.save(out)
    return str(out)


def test_an_unchanged_tree_has_no_impact(tmp_path):
    root = _tree(tmp_path)
    report = change_impact.impact(change_impact.read_derived_rows(_suts(tmp_path, root)), [root])
    assert report["status"] == {"same_value": 5} and report["impacted_count"] == 0 and report["changed_files"] == []


def test_a_function_without_derived_rows_is_listed_not_forgotten(tmp_path):
    root = _tree(tmp_path)
    _rows, without = change_impact.read_rows(_suts(tmp_path, root))
    assert without == ["x (unit.c)"]


def test_only_the_changed_behaviour_is_flagged_and_nothing_is_missed(tmp_path):
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    (root / "unit.c").write_text(UNIT.replace("g_i + 1U", "g_i + 2U").replace(
        "void gone(void) { g_p = 7U; }\n", ""), encoding="utf-8", newline="\n")
    report = change_impact.impact(change_impact.read_derived_rows(old), [root])
    assert report["status"] == {"function_removed": 1, "value_changed": 2, "same_value": 2}
    assert sorted(report["impacted_test_cases"]) == ["SwUTC_f", "SwUTC_gone"]
    changed = report["impacted_test_cases"]["SwUTC_f"]
    assert [(c["expected"], c["now"]) for c in changed] == [(2, 3), (6, 7)]
    assert report["changed_files"] and report["changed_files"][0].endswith("unit.c")


def test_a_missing_file_is_reported_and_counted_as_impact(tmp_path):
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    (root / "other.c").unlink()
    report = change_impact.impact(change_impact.read_derived_rows(old), [root])
    assert report["status"]["source_missing"] == 1 and "SwUTC_h" in report["impacted_test_cases"]


def test_a_changed_header_is_caught_though_the_c_file_is_untouched(tmp_path):
    # (review round 1 C1) the value depends on the header's macro: re-derived, not skipped by the file hash
    root = _tree(tmp_path)
    (root / "common.h").write_text(COMMON + "#define K_ADD 1U\n", encoding="utf-8", newline="\n")
    (root / "unit.c").write_text(UNIT.replace("g_i + 1U", "g_i + K_ADD"), encoding="utf-8", newline="\n")
    old = _suts(tmp_path, root)
    (root / "common.h").write_text(COMMON + "#define K_ADD 5U\n", encoding="utf-8", newline="\n")
    report = change_impact.impact(change_impact.read_derived_rows(old), [root])
    assert report["status"]["value_changed"] == 2 and "SwUTC_f" in report["impacted_test_cases"]
    assert report["changed_files"] == []   # the .c text did not change — and that did not matter


def test_crlf_text_is_the_same_text(tmp_path):
    # (review W1) the generator hashes text read with universal newlines
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    (root / "unit.c").write_bytes(UNIT.replace("\n", "\r\n").encode("utf-8"))
    report = change_impact.impact(change_impact.read_derived_rows(old), [root])
    assert report["changed_files"] == [] and report["impacted_count"] == 0


def test_a_moved_function_is_followed_and_an_unreadable_file_is_not_called_missing(tmp_path):
    # (review W2)
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    (root / "unit.c").write_text(UNIT.replace("void k(void) { g_p = g_i; }\n", ""), encoding="utf-8", newline="\n")
    (root / "moved.c").write_text('#include "common.h"\nextern U8 g_i;\nextern U8 g_p;\nvoid k(void) { g_p = g_i; }\n',
                                  encoding="utf-8", newline="\n")
    (root / "other.c").write_bytes(b"/* \xb0\xa1 */\n" + OTHER.encode("utf-8"))
    report = change_impact.impact(change_impact.read_derived_rows(old), [root])
    rows = [r for rs in report["impacted_test_cases"].values() for r in rs]
    assert report["status"]["source_unreadable"] == 1 and report["unreadable_source_files"] == 1
    assert "SwUTC_k" not in report["impacted_test_cases"]   # moved, same value
    assert all(r["impact"] != "source_missing" for r in rows)



# ── review round 2 ──────────────────────────────────────────────

def test_a_changed_file_with_only_unverified_rows_still_names_its_tests(tmp_path):
    """(W1) ``x`` has only a ``[검증 필요]`` row in unit.c; a file whose every row is unverified still changes."""
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    wb = openpyxl.load_workbook(old)
    ws = wb["Test Evidence"]
    sha = hashlib.sha256(OTHER.encode()).hexdigest()
    ws.append(["SwUFn_z", "z", 1, "g_q", "[검증 필요] x", "unknown", "none", "not_run", sha,
               str((root / "only_unknown.c").resolve()), "x", "SwUTC_z", "", "{}", "", ""])
    wb.save(old)
    (root / "only_unknown.c").write_text(OTHER.replace("3U", "4U"), encoding="utf-8", newline="\n")
    rows, _without, files = change_impact.read_workbook(old)
    report = change_impact.impact(rows, [root], files)
    assert [Path(p).name for p in report["changed_files_without_derived_rows"]] == ["only_unknown.c"]
    assert report["tests_on_changed_files_not_rederived"] == ["SwUTC_z"]
    assert any(p.endswith("only_unknown.c") for p in report["changed_files"])


def test_a_workbook_without_the_needed_columns_is_refused(tmp_path):
    """(W2) without ``Inputs JSON`` every row would be re-derived with no inputs — a false impact."""
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    wb = openpyxl.load_workbook(old)
    wb["Test Evidence"].delete_cols(HEADER.index("Inputs JSON") + 1)
    wb.save(old)
    import pytest
    with pytest.raises(ValueError, match="Inputs JSON"):
        change_impact.read_workbook(old)


def test_a_path_that_no_longer_resolves_is_not_called_unchanged(tmp_path):
    """(I4) the tree moved: nothing is compared — said so, not "no changed files"."""
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    moved = tmp_path / "moved"
    root.rename(moved)
    rows, _without, files = change_impact.read_workbook(old)
    report = change_impact.impact(rows, [moved], files)
    assert len(report["source_paths_not_compared"]) == 2 and report["changed_files"] == []


# ── review round 3 ──────────────────────────────────────────────

def test_a_recorded_path_without_a_hash_is_not_called_unchanged(tmp_path):
    """(W-a) rows that record a path but no hash: its text cannot be compared — listed, never silently unchanged."""
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    (root / "nohash.c").write_text(OTHER, encoding="utf-8", newline="\n")
    wb = openpyxl.load_workbook(old)
    wb["Test Evidence"].append(["SwUFn_n", "n", 1, "g_q", "[검증 필요] x", "unknown", "none", "not_run", "",
                                str((root / "nohash.c").resolve()), "x", "SwUTC_n", "", "{}", "", ""])
    wb.save(old)
    rows, _without, files = change_impact.read_workbook(old)
    report = change_impact.impact(rows, [root], files)
    assert [Path(p).name for p in report["source_paths_not_compared"]] == ["nohash.c"]
    assert not any(p.endswith("nohash.c") for p in report["changed_files"])


def test_an_unreadable_file_is_not_moved_to_a_same_named_function_elsewhere(tmp_path):
    """(W-c) the I3 order matters only when the name is defined elsewhere: ``h`` in ``dup.c`` is not the unreadable
    ``other.c``'s ``h``."""
    root = _tree(tmp_path)
    old = _suts(tmp_path, root)
    (root / "other.c").write_bytes(b"/* \xb0\xa1 */\n" + OTHER.encode("utf-8"))
    (root / "dup.c").write_text('#include "common.h"\nstatic U8 g_r;\nstatic void h(void) { g_r = 5U; }\n',
                                encoding="utf-8", newline="\n")
    rows = change_impact.read_derived_rows(old)
    change_impact.impact(rows, [root])
    h = [r for r in rows if r["function"] == "h"]
    assert [r["impact"] for r in h] == ["source_unreadable"] and "moved_to" not in h[0]


def test_a_hex_expectation_is_the_same_number():
    assert change_impact._same(6, "0x6") and change_impact._same(6, " 6 ") and not change_impact._same(6, "0x7")
    assert not change_impact._same(6, "[검증 필요] x")
