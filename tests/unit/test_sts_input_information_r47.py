"""R47 — a SyDS element's ``Input Information`` is stepped, not only quoted.

R46 read the cell as evidence for review items only (``BAT 전압(9V 이상 16.0V 이하), 5V, …`` — the only other field of
the HDPDM01 / KJPDS02 SyRS / SyDS that states comparisons). It writes where an element's input is — a normal range
(``9V 이상 16.0V 이하``) or a trigger (``5000RPM 초과``), the cell does not say which; an SRS requirement that cites the block (``SwTSR_0104`` → ``SySM_04``) now gets that range's boundary steps, in TCs of
their own. A step says only the input's place — inside or outside what the cell writes (a normal range and a trigger
read alike — review W1); the element's reaction (a fault, a reset) is the requirement's other sentences' to decide, and
the cross scorer holds no mutant against it (review W2).

Two extractor gaps the cell showed: ``RPM`` was no unit (``0~5000RPM`` lost it, ``5000RPM 초과`` was not read), and
``Pulse 주기 5ms 이하`` was read as a hold time (the step set "Pulse 주기 지속 시간") — a name that is itself a period
compares the period.
"""
from __future__ import annotations

import re
from collections import Counter

import docx

from generators.requirement_oracle import extract
from generators.sts import _build_tc_dict, _classify_steps, _make_tc_id
from generators.sts_requirement_tc import (
    append_requirement_boundary_tcs,
    boundary_steps,
    parse_system_requirement_docx,
    review_evidence_corpus,
)

SYDS = {"SySM_04": {"doc": "SyDS", "name": "Power Monitoring",
                    "fields": {"Input Information": "BAT 전압(9V 이상 16.0V 이하), 5V, Buzzer chk, Sensor chk"}},
        "SySM_06": {"doc": "SyDS", "name": "Watchdog", "fields": {"Input Information": "Watchdog Input(Pulse 주기 5ms 이하)"}}}
# the cross scorer reads a step's verdict with this pattern (scripts/requirement_oracle_eval.py `_VERDICT`)
_SCORER_VERDICT = re.compile(r"^조건 \[.*?\] (성립|불성립) →", re.S)


def _req(related, desc="- 전원 감시", rid="SwTSR_0104"):
    return {"id": rid, "name": "n", "description": desc, "verification": "", "asil": "B", "related_id": related}


def _facts(text):
    return [f for b in extract(f"ID\tSwTR_0001\n{text}\n") for line in b["lines"] for f in line["facts"]]


def test_parsing_reads_input_information_as_a_text_field(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=4, cols=2)
    for row, (k, v) in zip(t.rows, [("ID", "SySM_04"), ("Name", "Power Monitoring"), ("Description", "전원 감시"),
                                    ("Input Information", "BAT 전압(9V 이상 16.0V 이하), 5V")], strict=True):
        row.cells[0].text, row.cells[1].text = k, v
    path = tmp_path / "syds.docx"
    d.save(path)
    blocks, _ = parse_system_requirement_docx(str(path), "SyDS")
    assert blocks["SySM_04"] == {"doc": "SyDS", "name": "Power Monitoring",
                                 "fields": {"Description": "전원 감시",
                                            "Input Information": "BAT 전압(9V 이상 16.0V 이하), 5V"}}


def test_a_cited_input_range_steps_both_bounds_inside_and_outside():
    """HDPDM01 / KJPDS02 ``SwTSR_0104`` cites ``SySM_04``: 9 V (written precision 1 V) and 16.0 V (0.1 V)."""
    stats, review = Counter(), []
    groups = boundary_steps(_req("SySM_04"), stats, SYDS, review)
    assert [(g["evidence"]["signal"], g["evidence"]["op"], g["evidence"]["value"], g["evidence"]["source"]["field"])
            for g in groups] == [("BAT 전압", ">=", "9", "Input Information"),
                                 ("BAT 전압", "<=", "16.0", "Input Information")]
    assert [[(p["point"], p["holds"]) for p in g["evidence"]["points"]] for g in groups] == [
        [("8", False), ("9", True), ("10", True)], [("15.9", True), ("16.0", True), ("16.1", False)]]
    low, at, _ = groups[0]["steps"]
    assert low["action"].startswith("입력 설정 (요구 경계): BAT 전압 = 8V — 근거: 9V 이상 [SyDS SySM_04 · Input Information")
    assert low["expected"].startswith("조건 [BAT 전압 9V 이상 그리고 16.0V 이하] 불성립 → 시스템 요구 SySM_04 의 입력 칸이 적은 "
                                      "범위 밖 — 요소의 반응(이상 판정·리셋 등)은 같은 요구의 다른 문장 판정을 따른다")
    assert at["expected"].startswith("조건 [BAT 전압 9V 이상 그리고 16.0V 이하] 성립 → 시스템 요구 SySM_04 의 입력 칸이 적은 "
                                     "범위 안 — 요소의 반응(이상 판정·리셋 등)은 같은 요구의 다른 문장 판정을 따른다")
    # (review W1) no step claims a reaction — a power monitor reports *no* fault inside its normal range
    assert not any("동작 수행" in s["expected"] or "설계 기능" in s["expected"] for g in groups for s in g["steps"])
    # the cross scorer still reads each step's verdict
    assert [_SCORER_VERDICT.match(s["expected"]).group(1) for g in groups for s in g["steps"]] == \
        ["불성립", "성립", "성립", "성립", "성립", "불성립"]
    # no connective between the bounds: the range is a reading, listed to confirm (R45)
    assert [(r["kind"], r["reason"], r["fact"]) for r in review] == [
        ("read", "read_as_range", "9V 이상"), ("read", "read_as_range", "16.0V 이하")]
    used = (stats["traced:facts_used"], stats["traced:input_range_facts_used"], stats["traced:input_range_steps"])
    assert used == (2, 2, 6)


def test_input_ranges_come_after_every_other_traced_fact():
    """A TC numbered before R47 keeps its number: the new input-range facts follow all other traced facts, also those
    of a block cited after (``SwTSR_0302`` cites ``SySM_04`` then ``SySM_06``)."""
    system = {"SySM_04": {"doc": "SyDS", "name": "n", "fields": {
                  "Description": "- 저전압 8.5V 미만 시 정지", "Input Information": SYDS["SySM_04"]["fields"]["Input Information"]}},
              "SySM_06": {"doc": "SyDS", "name": "n", "fields": {"Description": "- 차속 3km/h 초과 시 정지"}}}
    groups = boundary_steps(_req("SySM_04, SySM_06", rid="SwTSR_0302"), Counter(), system, [])
    assert [(g["evidence"]["source"]["id"], g["evidence"]["source"]["field"], g["evidence"]["value"]) for g in groups] == [
        ("SySM_04", "Description", "8.5"), ("SySM_06", "Description", "3"),
        ("SySM_04", "Input Information", "9"), ("SySM_04", "Input Information", "16.0")]


def test_an_uncited_block_adds_nothing():
    """Only the Related ID links a block — a value is never matched by guess (R29)."""
    assert boundary_steps(_req("SyTR_0101"), Counter(), SYDS, []) == []


def test_other_fields_keep_their_expected_text():
    system = {"SyTR_0101": {"doc": "SyRS", "name": "n", "fields": {"Description": "도어 속도가 1.3m/s 초과 시 정지"}}}
    (group,) = boundary_steps(_req("SyTR_0101"), Counter(), system, [])
    assert group["steps"][2]["expected"].startswith("조건 [도어 속도 1.3m/s 초과] 성립 → 시스템 요구 SyTR_0101 의 이 문장이 "
                                                    "기술한 동작 수행 확인")


def test_a_period_is_the_compared_quantity_not_a_hold_time():
    """``Watchdog Input(Pulse 주기 5ms 이하)`` — before R47 the step set "Pulse 주기 지속 시간"."""
    (fact,) = _facts("Watchdog Input(Pulse 주기 5ms 이하)")
    assert (fact["kind"], fact["signal"], fact["op"], fact["unit"]) == ("threshold", "Pulse 주기", "<=", "ms")
    (group,) = boundary_steps(_req("SySM_06", rid="SwTSR_0210"), Counter(), SYDS, [])
    assert [s["action"].split(" — ")[0] for s in group["steps"]] == [
        "입력 설정 (요구 경계): Pulse 주기 = 4ms", "입력 설정 (요구 경계): Pulse 주기 = 5ms", "입력 설정 (요구 경계): Pulse 주기 = 6ms"]
    assert [p["holds"] for p in group["evidence"]["points"]] == [True, True, False]


def test_a_named_hold_time_is_still_a_hold_time():
    """The period rule is the name's last word only: ``LIN 통신이 15초이상 끊길 경우`` lasts (review W5 of R6)."""
    (fact,) = _facts("LIN 통신이 15초이상 끊길 경우")
    assert (fact["kind"], fact["signal"], fact.get("predicate")) == ("duration", "LIN 통신", "끊길")


def test_rpm_is_a_unit():
    rng, over = _facts("Motor Rotation : 모터의 물리적 회전(Close Direction), 0~5000RPM, 5000RPM 초과")
    assert (rng["kind"], rng["value"], rng["unit"]) == ("range", [0, 5000], "RPM")
    assert (over["kind"], over["op"], over["value"], over["unit"]) == ("threshold", ">", 5000, "RPM")
    (bare,) = _facts("Motor Current: 0~65535")
    assert (bare["kind"], bare["unit"], bare["signal"]) == ("range", "", "Motor Current")


def test_the_evidence_corpus_still_reads_input_information():
    corpus = review_evidence_corpus([], SYDS)
    (cand,) = corpus["ranges"][("V", 9, 16)]
    assert cand["where"] == "SyDS SySM_04 · Input Information" and cand["condition"] == "BAT 전압 9V 이상 그리고 16.0V 이하"


def _tcs(requirements, system):
    tcs = []

    def build(**kw):
        return _build_tc_dict(test_env="HILS", derive_inputs=False, is_safety=False, **kw)
    stats = append_requirement_boundary_tcs(tcs, requirements, build, _make_tc_id, _classify_steps, system=system)
    return stats, tcs


def test_an_input_range_never_joins_a_behaviour_tc():
    """(review W3) ``Pulse 주기 5ms 이하`` has no condition note and the same document as the Description fact before it
    — it joined that TC (6 → 9 steps). A TC made before R47 keeps its steps; the input range gets its own TC."""
    system = {"SySM_06": {"doc": "SyDS", "name": "Watchdog", "fields": {
        "Description": "- Watchdog 입력이 10ms 초과 시 리셋", "Input Information": "Watchdog Input(Pulse 주기 5ms 이하)"}}}
    before = {"SySM_06": {"doc": "SyDS", "name": "Watchdog", "fields": {"Description": "- Watchdog 입력이 10ms 초과 시 리셋"}}}
    req = _req("SySM_06", rid="SwTSR_0210")
    _s, old = _tcs([req], before)
    stats, new = _tcs([req], system)
    assert [tc["id"] for tc in new] == ["SwTC_SwTSR_0210_01", "SwTC_SwTSR_0210_02"]
    assert new[0]["steps"] == old[0]["steps"] and len(new[1]["steps"]) == 3
    assert new[1]["requirement_evidence"][0]["source"]["field"] == "Input Information"
    assert (stats["traced:tcs"], stats["traced:input_range_tcs"]) == (2, 1)


def test_a_split_quantity_in_an_input_cell_claims_no_action():
    """(review W4) the R45 split branch came first and said "이 조건에 기술한 동작 수행 확인" for an input cell."""
    system = {"SyFN_02": {"doc": "SyDS", "name": "n", "fields": {
        "Input Information": "Manual Assist조건(0.8m/s 미만) 또는 Tip-To-Run 조건(0.8m/s 이상 1.3m/s 이하)"}}}
    groups = boundary_steps(_req("SyFN_02", rid="SwTR_0104"), Counter(), system, [])
    texts = [s["expected"] for g in groups for s in g["steps"]]
    assert texts and all("입력 칸이 적은 범위" in t and "동작 수행" not in t for t in texts)
    assert any("같은 양을 나눈 조건" in t for t in texts)


def test_the_period_rule_is_the_word_itself():
    """(review I5) ``주기적`` is no period — a hold time stays one; lower-case ``rpm`` is a unit too."""
    (fact,) = _facts("LIN 통신이 주기적 15초이상 끊길 경우")
    assert fact["kind"] == "duration"
    (rpm,) = _facts("모터 5000rpm 초과 시 정지")
    assert (rpm["unit"], rpm["value"]) == ("rpm", 5000)


def test_the_cross_scorer_holds_no_mutant_against_an_input_cell(tmp_path):
    """(review W2) PV ``SwTSR_0104``: the input range's 성립 at 9 V read as the requirement's verdict against the reference
    ``8.9V 이하`` → ``<= 9.0`` mutant (``separated_against_verdict``). An input cell's point has no verdict."""
    import sys
    from pathlib import Path

    import openpyxl
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
    import requirement_oracle_eval as ev
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3.SW Test Spec"
    for r in (["", "Test Case"], ["", "Test Case ID"], ["", "ID"]):
        ws.append(r)
    rows = [("SwTC_1", "입력 설정 (요구 경계): BAT 전압 = 9V — 근거: 9V 이상 [SyDS SySM_04 · Input Information — SwTSR_0104 "
                       "Related ID]", "조건 [BAT 전압 9V 이상 그리고 16.0V 이하] 성립 → 시스템 요구 SySM_04 의 입력 칸이 적은 범위 안"),
            ("", "입력 설정 (요구 경계): 전원 = 8.5V — 근거: 8.5V 이하 [SyRS SyTR_0703 · Description — SwTSR_0104 Related ID]",
             "조건 [전원 8.5V 이하] 성립 → 시스템 요구 SyTR_0703 의 이 문장이 기술한 동작 수행 확인")]
    for tc, action, expected in rows:
        ws.append(["", tc, "", "", "", "", "", "", "d", "", action, expected, "SwTSR_0104"])
    wb.save(tmp_path / "g.xlsx")
    points = ev.read_sts(str(tmp_path / "g.xlsx"))["SwTSR_0104"]["points"]
    assert [(p["value"], p["traced"], p["holds"]) for p in points] == [(9.0, "SyDS SySM_04", None),
                                                                      (8.5, "SyRS SyTR_0703", True)]
