"""R29 (G4(b)) — STS requirement boundary TCs from the system requirements an SRS block cites.

An SRS block's Related ID names the SyRS / SyDS blocks it realises. When those documents are given, the text fields of
each **directly cited** block are read with the same rules as the requirement's own text, and their facts get steps
under the citing SRS requirement, naming the block they came from. A block the requirement does not cite contributes
nothing, whatever values it writes.
"""
from __future__ import annotations

import inspect
import sys
from collections import Counter
from pathlib import Path

import openpyxl
import pytest
from docx import Document

from generators.sts import _build_tc_dict, _classify_steps, _make_tc_id
from generators.sts_requirement_tc import (
    ACTION_PREFIX,
    BASIS_MARK,
    REQUIREMENT_EVIDENCE_HEADERS,
    append_requirement_boundary_tcs,
    boundary_steps,
    cited_system_ids,
    load_system_requirements,
    parse_system_requirement_docx,
    write_requirement_evidence_sheet,
)
from report_gen.generation_disclosures import build_disclosures

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
import requirement_oracle_eval as ev  # noqa: E402

SYSTEM = {
    "SyII_06": {"doc": "SyDS", "fields": {"Description": "Power 이상 발생 시 Diagnostic block 에 전송한다.",
                                          "Range": "저전압 : 8.5V 이하 500ms 초과"}},
    "SyTR_0602": {"doc": "SyRS", "fields": {"Description": "도어속도 1.3m/s 초과 시 정지한다."}},
    "SyTR_0999": {"doc": "SyRS", "fields": {"Description": "도어속도 2.5m/s 초과 시 경보한다."}},   # cited by nobody
}


def _req(desc="- u16g_ApiIn_Vsup: 1604 초과", rid="SwTR_0601", related="SyII_06, SyEIF_77"):
    return {"id": rid, "name": "n", "description": desc, "verification": "", "asil": "B", "related_id": related}


def _build(**kw):
    return _build_tc_dict(test_env="SwTE_01", derive_inputs=False, is_safety=False, **kw)


def _docx(path: Path, tables) -> str:
    doc = Document()
    for rows in tables:
        t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    doc.save(str(path))
    return str(path)


# ── system document parsing ───────────────────────────────────────────────────────────────────────────────────────

def test_system_tables_are_read_by_id_with_merged_key_cells_and_duplicates_counted(tmp_path):
    path = _docx(tmp_path / "sy.docx", [
        [("ID", "SyII_06"), ("Name", "Power Diag"), ("Range", "저전압 : 8.5V 이하"), ("Related ID", "SyEIF_06")],
        [("ID", "ID", "SyOS_01"), ("Description", "Description", "Latch Open 수신 시 정지")],   # key merged over 2 columns
        [("ID", "SyII_06"), ("Range", "다른 판")],                                            # a later copy
        [("Attribute", "Contents"), ("System Name", "PDSM")],                                 # not a requirement table
        [("ID", "SwTR_0101"), ("Description", "SW 요구")],                                   # not a system ID
        [("ID", "SyTR_0101"), ("Verification Criteria\n(optional)", "5V 이상")],               # field name variant
    ])
    blocks, duplicates = parse_system_requirement_docx(path, "SyDS")
    assert set(blocks) == {"SyII_06", "SyOS_01", "SyTR_0101"} and duplicates == 1
    assert blocks["SyII_06"] == {"doc": "SyDS", "fields": {"Range": "저전압 : 8.5V 이하"}}   # first table wins
    assert blocks["SyOS_01"]["fields"] == {"Description": "Latch Open 수신 시 정지"}
    assert blocks["SyTR_0101"]["fields"] == {"Verification criteria": "5V 이상"}


def test_documents_merge_first_wins_and_an_unreadable_one_is_recorded_not_raised(tmp_path):
    a = _docx(tmp_path / "a.docx", [[("ID", "SyTR_0101"), ("Description", "A")]])
    b = _docx(tmp_path / "b.docx", [[("ID", "SyTR_0101"), ("Description", "B")], [("ID", "SyDB_0102"), ("Range", "C")]])
    system, records = load_system_requirements([("SyRS", a), ("SyDS", b), ("SyDS", str(tmp_path / "missing.docx")),
                                                ("SyRS", None)])
    assert system["SyTR_0101"]["fields"]["Description"] == "A" and system["SyDB_0102"]["doc"] == "SyDS"
    assert records[0] == {"doc": "SyRS", "file": "a.docx", "blocks": 1, "ids_in_earlier_document": 0,
                          "duplicate_tables": 0}
    assert records[1]["ids_in_earlier_document"] == 1 and records[1]["blocks"] == 2
    assert records[2]["file"] == "missing.docx" and records[2]["error"] and "blocks" not in records[2]
    assert len(records) == 3                                      # a document not given is not a record


def test_cited_ids_come_from_the_related_id_in_order_once():
    assert cited_system_ids({"related_id": "SyTR_0602, SyII_06,SyTR_0602 / SwCom_01"}) == ["SyTR_0602", "SyII_06"]
    assert cited_system_ids({"related_id": ""}) == [] and cited_system_ids({}) == []


# ── traced facts ───────────────────────────────────────────────────────────────────────────────────────────────────

def test_a_cited_blocks_facts_get_steps_that_name_the_block_and_the_citing_requirement():
    stats = Counter()
    groups = boundary_steps(_req(), stats, SYSTEM)
    traced = [g for g in groups if g["evidence"]["source"]]
    own = [g for g in groups if not g["evidence"]["source"]]
    assert len(own) == 1 and own[0]["evidence"]["value"] == "1604"
    assert {(g["evidence"]["signal"], g["evidence"]["value"]) for g in traced} >= {("저전압", "8.5")}
    g = next(g for g in traced if g["evidence"]["value"] == "8.5")
    assert g["evidence"]["source"] == {"doc": "SyDS", "id": "SyII_06", "field": "Range"}
    act = g["steps"][0]["action"]
    assert act.startswith(ACTION_PREFIX + "저전압 = 8.4V")
    assert act.endswith("[SyDS SyII_06 · Range — SwTR_0601 Related ID]")
    assert act.split(BASIS_MARK)[1].startswith("8.5V 이하")          # the system sentence is the basis, not a stimulus
    assert "시스템 요구 SyII_06 의 이 문장이 기술한 동작" in g["steps"][0]["expected"]
    # the SRS's own counters are unchanged by the trace; the trace has its own, including the ID no document has
    assert stats["facts_used"] == 1 and stats["traced:cited"] == 2 and stats["traced:cited_not_in_documents"] == 1
    assert stats["traced:facts_used"] == len(traced) and "tcs" not in stats


def test_a_block_the_requirement_does_not_cite_contributes_nothing_whatever_it_writes():
    groups = boundary_steps(_req(related="SyTR_0602"), None, SYSTEM)
    values = {g["evidence"]["value"] for g in groups if g["evidence"]["source"]}
    assert values == {"1.3"}                                          # SyTR_0999's 2.5 m/s is never read


def test_no_system_documents_leaves_the_requirement_boundary_exactly_as_before():
    with_none, without = Counter(), Counter()
    a = boundary_steps(_req(), with_none, None)
    b = boundary_steps(_req(), without)
    assert a == b and with_none == without and not any(k.startswith("traced:") for k in with_none)


def test_a_fact_the_srs_already_steps_is_not_repeated_from_the_system_block():
    stats = Counter()
    req = _req(desc="- 저전압 : 8.5V 이하 인 경우 고장", related="SyII_06")
    groups = boundary_steps(req, stats, SYSTEM)
    assert [g["evidence"]["value"] for g in groups if not g["evidence"]["source"]] == ["8.5"]
    assert "8.5" not in {g["evidence"]["value"] for g in groups if g["evidence"]["source"]}
    assert stats["traced:skipped:already_stepped"] == 1


def test_traced_facts_never_share_a_tc_with_the_srs_facts_and_are_counted_apart():
    tcs: list = []
    stats = append_requirement_boundary_tcs(tcs, [_req(related="SyTR_0602")], _build, _make_tc_id, _classify_steps,
                                            max_steps=12, system=SYSTEM)
    sources = [{bool(e["source"]) for e in tc["requirement_evidence"]} for tc in tcs]
    assert sources == [{False}, {True}]                               # would fit in one TC by steps, split by document
    assert (stats["tcs"], stats["traced:tcs"]) == (1, 1)
    assert [tc["id"] for tc in tcs] == ["SwTC_SwTR_0601_01", "SwTC_SwTR_0601_02"]


def test_the_evidence_sheet_names_the_source_document_of_every_row():
    tcs: list = []
    append_requirement_boundary_tcs(tcs, [_req(related="SyTR_0602")], _build, _make_tc_id, _classify_steps,
                                    system=SYSTEM)
    wb = openpyxl.Workbook()
    assert write_requirement_evidence_sheet(wb, tcs) == 2
    rows = list(wb["Requirement Evidence"].iter_rows(values_only=True))
    assert list(rows[0]) == REQUIREMENT_EVIDENCE_HEADERS and rows[0][-1] == "Source Document"
    assert [r[-1] for r in rows[1:]] == ["SRS", "SyRS SyTR_0602 · Description"]


# ── review round 1: what a system block states as an outcome is never a stimulus ──────────────────────────────────

def _traced(text, field="Verification criteria", doc="SyDS", sid="SyFN_13", stats=None):
    system = {sid: {"doc": doc, "fields": {field: text}}}
    groups = boundary_steps(_req(desc="", related=sid), stats if stats is not None else Counter(), system)
    return [g for g in groups if g["evidence"]["source"]]


def test_a_verification_cases_output_line_is_not_stepped_and_its_input_line_is():
    # (C2) HDPDM01 SyFN_13: ``Output : 암전류가 10mA 이상으로 상승하는지 확인`` was stepped as ``암전류 = 9mA``
    stats = Counter()
    groups = _traced("Case 1\nInput : 입력전원 8.9V 미만\nOutput : 암전류가 10mA 이상으로 상승하는지 확인\n"
                     "Case 2\nInput : 입력전원 16.1V 이상\nOutput : DTC 송신 확인", stats=stats)
    assert [(g["evidence"]["value"], g["evidence"]["unit"]) for g in groups] == [("8.9", "V"), ("16.1", "V")]
    assert stats["traced:skipped:outcome_section"] == 1


def test_action_and_system_behavior_fields_are_outcomes():
    stats = Counter()
    assert _traced("모터 출력 10V 이상 유지", field="Action", stats=stats) == []
    assert _traced("속도 1.3m/s 초과로 구동", field="System Behavior", stats=stats) == []
    assert stats["traced:skipped:outcome_section"] == 2


def test_an_obligation_is_an_output_requirement_but_a_held_condition_is_not():
    stats = Counter()
    assert _traced("시스템의 암전류는 0.3mA 이하여야 한다.", field="Description", stats=stats) == []
    assert _traced("이때 최소 전류는 0.3mA 이하로 유지한다.", field="Description", stats=stats) == []
    assert stats["traced:skipped:output_requirement"] == 2
    held = _traced("모터 전류 5A 이상 유지시 정지한다", field="Description")
    assert [g["evidence"]["value"] for g in held] == ["5"]


def test_a_subject_read_from_a_separator_is_not_a_subject():
    stats = Counter()
    groups = _traced("Pull-up Port를 4.85V 미만/5.15V초과로 Setting", field="Description", stats=stats)
    assert "5.15" not in [g["evidence"]["value"] for g in groups]
    assert stats["traced:skipped:subject_unclear"] == 1


@pytest.mark.parametrize("text", ["Motor Driver Block 와의 통신이 300ms 이상 단절 시 정지",
                                  "Watchdog Input 을 10ms 이상의 신호를 주입"])
def test_a_hold_time_named_by_a_particle_or_a_verb_phrase_is_not_stepped(text):
    # (KJPDS02_PV SySM_03 / SySM_06) ``와 통신 지속 시간`` · ``신호를 주입 지속 시간`` are misreads, not quantities
    stats = Counter()
    assert _traced(text, field="Description", stats=stats) == []
    assert stats["traced:skipped:subject_unclear"] == 1


def test_a_condition_that_goes_on_after_its_must_is_a_condition_not_an_obligation():
    # (review r2 W-4) ``이어야 모터가 동작한다`` / ``유지되어야 고장으로 판정한다`` are conditions of what follows
    stats = Counter()
    assert [g["evidence"]["value"] for g in _traced("전원이 9V 이상이어야 모터가 동작한다", field="Description",
                                                    stats=stats)] == ["9"]
    assert [g["evidence"]["kind"] for g in _traced("고장 조건이 500ms 이상 유지되어야 고장으로 판정한다",
                                                   field="Description", stats=stats)] == ["duration"]
    assert "traced:skipped:output_requirement" not in stats


@pytest.mark.parametrize("text", ["도어를 여는 속도 : 1.3m/s 초과", "전압 이상 : 9V 이상"])
def test_a_label_subject_with_an_object_or_a_trailing_comparison_word_is_kept(text):
    # (review r2 W-4) ``도어를 여는 속도`` names a quantity; ``전압 이상`` is "voltage abnormality"
    assert len(_traced(text, field="Description")) == 1


def test_two_unnamed_hold_times_of_two_conditions_are_both_stepped():
    # (review r2 W-3) SySM_04: the over-voltage hold was dropped as "already stepped" after the under-voltage hold
    stats = Counter()
    groups = _traced("Input : 입력전원 8.5V 이하 500ms 초과 유지\nInput : 입력전원 16.5V 이상 500ms 초과 유지", stats=stats)
    holds = [g for g in groups if g["evidence"]["kind"] == "duration"]
    assert len(holds) == 2 and "traced:skipped:already_stepped" not in stats


def test_a_condition_without_a_subject_is_not_held_true_by_an_and_the_words_do_not_say():
    # (review r2 W-1) ``… 500ms 미만 유지 후 9V 이상으로 복귀``: the hold time's precondition must not assert ``9V 이상``
    (hold,) = [g for g in boundary_steps(_req(desc="- 전원 8.5V 이하 500ms 미만 유지 후 9V 이상으로 복귀", related=""))
               if g["evidence"]["kind"] == "duration"]
    note = hold["evidence"]["combination_note"]
    assert "[9V 이상]: 결합이 원문에 명시되지 않음" in note and "[9V 이상]: 성립" not in note


def test_a_time_after_a_condition_on_another_quantity_is_that_conditions_hold_time():
    # (C2 / W5) ``저전압 : 8.5V 이하 500ms 초과`` — the time is how long low voltage lasts, never ``저전압 = 499ms``
    groups = _traced("저전압 : 8.5V 이하 500ms 초과", field="Range", sid="SyII_06")
    by_value = {g["evidence"]["value"]: g for g in groups}
    assert by_value["500"]["evidence"]["kind"] == "duration"
    assert by_value["500"]["steps"][0]["action"].startswith(ACTION_PREFIX + "저전압 지속 시간 = 499ms")
    assert by_value["8.5"]["steps"][0]["action"].startswith(ACTION_PREFIX + "저전압 = 8.4V")
    assert "[저전압 지속 시간 500ms 초과]: 성립 상태로 둔다" in by_value["8.5"]["evidence"]["combination_note"]


def test_the_not_holding_step_names_the_system_requirement_too():
    (g,) = _traced("도어속도 1.3m/s 초과 시 정지한다.", field="Description", sid="SyTR_0602")
    below, _at, above = g["steps"]
    assert "시스템 요구 SyTR_0602 의 이 문장의 동작 대상 아님" in below["expected"]
    assert "시스템 요구 SyTR_0602 의 이 문장이 기술한 동작 수행 확인" in above["expected"]


def test_traced_points_outside_the_subject_type_are_counted_apart_from_the_srs():
    stats = Counter()
    _traced("u8g_Level: 0 초과", field="Description", stats=stats)
    assert stats["traced:points_outside_subject_type"] == 1 and "points_outside_subject_type" not in stats


def test_facts_of_the_syrs_and_the_syds_never_share_a_tc():
    # (W2) HDPDM01 SwTC_SwTSR_0101_09 held a SyRS fact and a SyDS fact
    system = {"SyTR_0602": {"doc": "SyRS", "fields": {"Description": "도어속도 1.3m/s 초과 시 정지한다."}},
              "SyII_06": {"doc": "SyDS", "fields": {"Range": "배터리 전압 8.5V 이하"}}}
    tcs: list = []
    append_requirement_boundary_tcs(tcs, [_req(desc="", related="SyTR_0602, SyII_06")], _build, _make_tc_id,
                                    _classify_steps, max_steps=12, system=system)
    assert [{e["source"]["doc"] for e in tc["requirement_evidence"]} for tc in tcs] == [{"SyRS"}, {"SyDS"}]


def test_the_blocks_carrying_the_traced_facts_are_counted():
    system = {"SyTR_0602": {"doc": "SyRS", "fields": {"Description": "도어속도 1.3m/s 초과 시 정지한다."}}}
    reqs = [_req(desc="", rid=f"SwTR_060{i}", related="SyTR_0602") for i in (1, 2)]
    stats = append_requirement_boundary_tcs([], reqs, _build, _make_tc_id, _classify_steps, system=system)
    assert stats["traced_facts_by_block"] == {"SyTR_0602": 2} and stats["traced_blocks_used"] == 1


def test_a_key_written_on_two_rows_keeps_both_and_a_key_repeated_as_its_value_is_empty(tmp_path):
    path = _docx(tmp_path / "sy.docx", [[("ID", "SySM_04"), ("Time Constraint", "FRTI : 100 ms"),
                                          ("Time Constraint", "FDTI : 300 ms"), ("Description", "Description")]])
    blocks, _ = parse_system_requirement_docx(path, "SyDS")
    assert blocks["SySM_04"]["fields"] == {"Time Constraint": "FRTI : 100 ms\nFDTI : 300 ms"}
    # (review r3 I1) the same value on two rows is kept once; a value that only contains another is its own row
    path = _docx(tmp_path / "sy2.docx", [[("ID", "SySM_05"), ("Range", "8.5V 이하\n500ms 초과"),
                                           ("Range", "8.5V 이하\n500ms 초과"), ("Range", "8.5V 이하")]])
    blocks, _ = parse_system_requirement_docx(path, "SyDS")
    assert blocks["SySM_05"]["fields"] == {"Range": "8.5V 이하\n500ms 초과\n8.5V 이하"}


# ── generate_sts wiring and the disclosure ─────────────────────────────────────────────────────────────────────────

_REQ_TEXT = ["SwTR_0601: 저전압 고장을 검출한다. Related ID: SyII_06"]


def test_generate_sts_reads_the_given_system_documents_and_discloses_them(tmp_path):
    from generators.sts import generate_sts
    syds = _docx(tmp_path / "(P_SyDS) System Design.docx",
                 [[("ID", "SyII_06"), ("Range", "저전압 : 8.5V 이하 500ms 초과")]])
    res = generate_sts(_REQ_TEXT, {}, str(tmp_path / "o.xlsm"), project_config={"project_id": "T"}, syds_path=syds)
    rb = res["quality_report"]["generation_stats"]["requirement_boundary"]
    assert rb["system_documents"] == [{"doc": "SyDS", "file": "(P_SyDS) System Design.docx", "blocks": 1,
                                       "ids_in_earlier_document": 0, "duplicate_tables": 0}]
    assert rb["traced:tcs"] >= 1 and rb["traced:facts_used"] >= 1
    item = {i["key"]: i for i in build_disclosures("sts", res["quality_report"])}["sts_traced_system_boundary"]
    assert item["tone"] == "info" and item["value"].startswith(f"TC {rb['traced:tcs']} ")
    assert "SyDS (P_SyDS) System Design.docx 블록 1" in item["note"]
    wb = openpyxl.load_workbook(res["output_path"], read_only=True)
    try:
        sources = [r[-1] for r in wb["Requirement Evidence"].iter_rows(min_row=2, values_only=True)]
    finally:
        wb.close()
    assert "SyDS SyII_06 · Range" in sources


def test_no_system_documents_is_disclosed_as_no_input_and_an_unreadable_one_as_a_warning(tmp_path):
    from generators.sts import generate_sts
    none = generate_sts(_REQ_TEXT, {}, str(tmp_path / "a.xlsm"), project_config={"project_id": "T"})
    assert none["quality_report"]["generation_stats"]["requirement_boundary"]["system_documents"] == []
    item = {i["key"]: i for i in build_disclosures("sts", none["quality_report"])}["sts_traced_system_boundary"]
    assert (item["value"], item["tone"]) == ("입력 없음", "info")
    bad = generate_sts(_REQ_TEXT, {}, str(tmp_path / "b.xlsm"), project_config={"project_id": "T"},
                       syrs_path=str(tmp_path / "gone.docx"), system_input_skips=["SyDS: 접근 거부 — prefix"])
    rb = bad["quality_report"]["generation_stats"]["requirement_boundary"]
    assert rb["system_documents"][0]["error"] and rb["system_input_skips"] == ["SyDS: 접근 거부 — prefix"]
    item = {i["key"]: i for i in build_disclosures("sts", bad["quality_report"])}["sts_traced_system_boundary"]
    assert (item["value"], item["tone"]) == ("읽지 못함", "warning")
    assert "gone.docx 읽기 실패" in item["note"] and "로컬화 실패 — SyDS: 접근 거부" in item["note"]
    assert not any(k.startswith("traced:") for k in rb)                # nothing was traced


def _item_of(rb):
    qr = {"generation_stats": {"requirement_boundary": {"tcs": 0, "facts": 0, **rb}}}
    return {i["key"]: i for i in build_disclosures("sts", qr)}["sts_traced_system_boundary"]


def test_a_read_document_and_a_failed_one_warn_and_name_both():
    # (review M15) one read, one unreadable: counts shown, the failure keeps the warning tone
    item = _item_of({"traced:tcs": 2, "system_documents": [
        {"doc": "SyRS", "file": "a.docx", "blocks": 3, "ids_in_earlier_document": 0, "duplicate_tables": 0},
        {"doc": "SyDS", "file": "b.docx", "error": "PackageNotFoundError: x"}]})
    assert item["tone"] == "warning" and item["value"].startswith("TC 2 ")
    assert "SyRS a.docx 블록 3" in item["note"] and "SyDS b.docx 읽기 실패" in item["note"]


def test_a_document_with_no_system_table_warns_and_the_one_not_given_is_named():
    # (review W3) an SRS registered as the SyRS reads fine and holds no system block — not "no input", not info
    item = _item_of({"system_documents": [{"doc": "SyRS", "file": "wrong.docx", "blocks": 0,
                                           "ids_in_earlier_document": 0, "duplicate_tables": 0}]})
    assert item["tone"] == "warning" and "하나도 찾지 못했다" in item["note"]
    assert "받지 않은 문서: SyDS" in item["note"]


def test_the_blocks_carrying_most_facts_are_disclosed():
    item = _item_of({"traced:tcs": 1, "traced_facts_by_block": {"SySM_04": 35, "SyII_06": 3}, "traced_blocks_used": 2,
                     "system_documents": [{"doc": "SyDS", "file": "s.docx", "blocks": 9,
                                           "ids_in_earlier_document": 0, "duplicate_tables": 0}]})
    assert "SySM_04 35, SyII_06 3 (블록 2 개" in item["note"] and item["tone"] == "info"


def test_an_artifact_from_before_the_trace_gets_no_item():
    qr = {"generation_stats": {"requirement_boundary": {"tcs": 1, "steps": 3, "facts": 1, "facts_used": 1}}}
    keys = {i["key"] for i in build_disclosures("sts", qr)}
    assert "sts_requirement_boundary" in keys and "sts_traced_system_boundary" not in keys


# ── evaluator: a kill only a traced point makes is split out ──────────────────────────────────────────────────────

def _sts(path: Path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3.SW Integration Test Spec"
    for _ in range(3):
        ws.append([""])
    for tc, action, srs in rows:
        ws.append(["", tc, "", "", "", "", "", "", "d", "", action, "", srs])
    wb.save(path)
    return str(path)


def test_the_evaluator_credits_a_kill_to_the_trace_only_when_the_requirements_own_points_miss(tmp_path):
    srs = "ID\tSwTR_0601\nName\tT\n- 저전압 고장 검출\n"
    ref = _sts(tmp_path / "ref.xlsx", [("SwTC_1", "저전압 500ms 초과 유지", "SwTR_0601"),
                                       ("SwTC_2", "도어 1.3m/s 초과 속도", "SwTR_0602")])
    trace = " — 근거: 500ms 초과 [SyDS SyII_06 · Range — SwTR_0601 Related ID]"
    gen = _sts(tmp_path / "gen.xlsx", [
        ("SwTC_9", f"입력 설정 (요구 경계): 저전압 지속 시간 = 501ms{trace}", "SwTR_0601"),
        ("SwTC_8", "입력 설정 (요구 경계): 도어 속도 = 1.4m/s — 근거: 1.3m/s 초과", "SwTR_0602"),
        ("", f"입력 설정 (요구 경계): 도어 속도 = 1.4m/s{trace.replace('SwTR_0601', 'SwTR_0602')}", ""),
    ])
    points = ev.read_sts(gen)
    assert [p["traced"] for p in points["SwTR_0601"]["points"]] == ["SyDS SyII_06"]
    assert [p["traced"] for p in points["SwTR_0602"]["points"]] == [None, "SyDS SyII_06"]
    g = ev.evaluate(srs, ref, gen)["cross_source"]["generated"]
    by = {(m["req_id"], m["mutant"], m["m_op"], m["m_value"]): m for m in g["rows"]}
    assert by[("SwTR_0601", "value_shift", ">", 501.0)]["killed_by_trace"] == "SyDS SyII_06"   # 501 > 500, not > 501
    assert by[("SwTR_0601", "boundary_inclusion", ">=", 500.0)]["killed_by_point"] is None       # needs a point at 500
    # both suites' points kill SwTR_0602's inclusion mutant: the requirement's own point gets the credit
    assert by[("SwTR_0602", "value_shift", ">", 1.4)]["killed_by_trace"] is None
    assert g["killed_via_traced_system"] == sum(1 for m in g["rows"] if m["killed_by_trace"])
    assert g["killed_via_traced_system"] >= 1 and g["killed"] > g["killed_via_traced_system"]


def test_the_requirements_own_point_gets_the_credit_even_when_a_traced_point_comes_first(tmp_path):
    # (review M5) the document order puts the traced point first; the requirement's own point still gets the kill, and
    #   (review I1) that kill is classified among the requirement's own points — a single subject
    srs = "ID\tSwTR_0602\nName\tT\n- 도어 속도\n"
    ref = _sts(tmp_path / "ref.xlsx", [("SwTC_1", "도어 1.3m/s 초과 속도", "SwTR_0602")])
    trace = " — 근거: 1.3m/s 초과 [SyRS SyTR_0602 · Description — SwTR_0602 Related ID]"
    gen = _sts(tmp_path / "gen.xlsx", [
        ("SwTC_9", f"입력 설정 (요구 경계): 동작 속도 = 1.4m/s{trace}", "SwTR_0602"),
        ("", "입력 설정 (요구 경계): 도어 속도 = 1.4m/s — 근거: 1.3m/s 초과", ""),
    ])
    g = ev.evaluate(srs, ref, gen)["cross_source"]["generated"]
    killed = [m for m in g["rows"] if m["killed_by_point"] is not None]
    assert killed and all(m["killed_by_trace"] is None and m["killed_by_signal"] == "도어 속도" for m in killed)
    assert g["killed_via_traced_system"] == 0
    assert (g["killed_single_subject"], g["killed_multi_subject"]) == (len(killed), 0)


def test_the_generators_traced_step_reads_back_as_traced_in_the_evaluator(tmp_path):
    # (review I5) one format, two modules: the evaluator reads the generator's own action text, not a hand-written copy
    (g,) = _traced("도어속도 1.3m/s 초과 시 정지한다.", field="Description", doc="SyRS", sid="SyTR_0602")
    gen = _sts(tmp_path / "gen.xlsx", [("SwTC_9" if i == 0 else "", s["action"], "SwTR_0601" if i == 0 else "")
                                       for i, s in enumerate(g["steps"])])
    points = ev.read_sts(gen)["SwTR_0601"]["points"]
    assert [p["traced"] for p in points] == ["SyRS SyTR_0602"] * 3 and [p["value"] for p in points] == [1.2, 1.3, 1.4]


# ── backend: one resolver, four handlers, the web form ────────────────────────────────────────────────────────────

def test_the_resolver_leaves_out_a_missing_document_with_its_reason(tmp_path, monkeypatch):
    import backend.services.resolver_helpers as rh
    monkeypatch.setattr(rh, "_needs_resolver_read", lambda: False)
    ok = _docx(tmp_path / "sy.docx", [[("ID", "SyTR_0101")]])
    got = rh.resolve_system_requirement_docs(ok, str(tmp_path / "gone.docx"))
    assert got["syrs_path"] == str(Path(ok).resolve()) and got["syds_path"] is None
    assert len(got["system_input_skips"]) == 1 and got["system_input_skips"][0].startswith("SyDS: 파일 없음")
    assert rh.resolve_system_requirement_docs("", "") == {"syrs_path": None, "syds_path": None,
                                                          "system_input_skips": []}


@pytest.mark.parametrize("path", ["/api/jenkins/sts/generate-async", "/api/local/sts/generate",
                                  "/api/local/sts/generate-stream", "/api/local/sts/generate-async"])
def test_every_sts_handler_declares_the_system_inputs(path):
    # FastAPI drops an undeclared form field silently — a handler without them would generate without the trace
    from backend.main import app
    route = next(r for r in app.routes if getattr(r, "path", "") == path and "POST" in getattr(r, "methods", set()))
    params = inspect.signature(route.endpoint).parameters
    assert {"syrs_path", "syds_path"} <= set(params)


@pytest.mark.parametrize("path", ["/api/jenkins/sts/generate-async", "/api/local/sts/generate",
                                  "/api/local/sts/generate-stream", "/api/local/sts/generate-async"])
def test_every_sts_handler_passes_the_system_inputs_to_the_generator(path, tmp_path, monkeypatch):
    # (review W4) declared but not passed on is the same silent loss — run each handler to the generator call
    import threading

    from fastapi.testclient import TestClient

    import sts_generator
    from backend.main import app

    seen: dict = {}
    done = threading.Event()

    def fake(**kw):
        seen.update(kw)
        done.set()
        raise RuntimeError("stop after the call")   # the handlers report a generator failure; only the call matters

    monkeypatch.setattr(sts_generator, "generate_sts", fake)
    import backend.helpers.session as session_helpers
    import backend.routers.local as local_router
    monkeypatch.setattr(session_helpers, "_resolve_base_dir", lambda _x: tmp_path)   # outputs stay in tmp
    monkeypatch.setattr(local_router, "_discover_sds_docx", lambda *a, **k: None)   # no repo docs/ discovery
    src = tmp_path / "src"
    src.mkdir()
    (src / "a.c").write_text("int f(void){return 0;}", encoding="utf-8")
    srs = tmp_path / "(P_SRS) Software Requirements Specification.docx"
    srs.write_text("x", encoding="utf-8")
    syrs = _docx(tmp_path / "sy.docx", [[("ID", "SyTR_0101")]])
    body = {"source_root": str(src), "srs_path": str(srs), "syrs_path": syrs,
            "syds_path": str(tmp_path / "gone_SyDS.docx")}
    if "jenkins" in path:
        (tmp_path / "cache").mkdir()
        body.update(job_url="http://ci/job/x/", cache_root=str(tmp_path / "cache"))
    with TestClient(app, raise_server_exceptions=False) as client:
        r = client.post(path, data=body, headers={"X-User": "tester"})
        _ = r.text                                           # a stream is read to its end
    assert done.wait(30), f"{path}: generate_sts was not called ({r.status_code} {r.text[:300]})"
    assert seen["syrs_path"] == str(Path(syrs).resolve()) and seen["syds_path"] is None
    assert [s.split(":")[0] for s in seen["system_input_skips"]] == ["SyDS"]


def test_the_web_form_sends_the_system_inputs_for_sts_only():
    src = (Path(__file__).resolve().parents[2] / "frontend-v2/src/components/sections/DocGenSection.jsx").read_text(
        encoding="utf-8")
    block = src[src.index("if (docType === 'sts') {"):]
    block = block[:block.index("}\n") + 1]
    assert "formData.append('syrs_path', syrsPath)" in block and "formData.append('syds_path', sydsPath)" in block
    assert "docPaths.syrs || linkedDocs.syrs" in block and "docPaths.syds || linkedDocs.syds" in block


def test_the_readiness_gate_lists_the_system_inputs_as_optional_for_sts():
    from backend.routers.docgen_preflight import _DOC_KEY_TO_INPUT
    from backend.services.docgen_requirements import DOC_REQUIREMENTS, IN_SYDS, IN_SYRS
    opt = DOC_REQUIREMENTS["sts"]["optional"]
    assert opt[IN_SYRS] and opt[IN_SYDS] and IN_SYRS not in DOC_REQUIREMENTS["sts"]["required"]
    assert _DOC_KEY_TO_INPUT["syrs"] == IN_SYRS and _DOC_KEY_TO_INPUT["syds"] == IN_SYDS
