"""R56(N52) 가드 — Called/Calling 행의 방향·내용을 정본 관례로.

착수 실측(run 2086 산출물 ↔ 정본 KJPDS02 SwUDS v3.03, 분석과 대조되는 951 함수):
  정본 "Called Function" 행 = **호출자**, "Calling Function" 행 = **피호출자** — 이 방향 921 · 반대 0.
  생성 문서는 989개 표 **전부** 반대였고, accuracy 리포트는 문서를 자기 관례로 재서 99.8% 라 적었다.
  잔여 불일치 3건은 방향이 아니라 **여러 줄 프로토타입**이었다 — 소스가 `(` 를 다음 줄에 두면(LIN 드라이버
  50건) 칸에 실린 프로토타입의 이름 줄에 `(` 가 없어 줄 단위 이름 파서가 이름을 잃었다. 정본은 칸에 이름만 적고
  (951 중 914), 프로토타입은 한 줄이며 697개가 `static` 으로 시작한다(생성본 0 — tree-sitter 가 `type` 필드에서만
  `static` 을 찾았다).

가드:
  1. 라벨 ↔ 내부 키 대응은 `validation_labels.CALL_ROW_LABEL_TO_KEY` 하나 — 라이터(pairs·블록 렌더)와 리더(docx·텍스트)가
     같은 표를 쓰고 왕복이 항등이다.
  2. 정본 관례 문서를 읽으면 `called`=피호출자·`calling`=호출자 — accuracy 100%; 옛 관례 문서(대조군)는 아니다.
  3. 소스 분석: 프로토타입 한 줄(주석 없음) · static 보존 · `called` 텍스트는 이름만.
  4. 이름 파서는 여러 줄 프로토타입·주석에서도 이름을 잃지 않는다(run 2086 의 실제 세 문자열).
  5. 함수명 리터럴 특례(wake_up_setting → l_ifc_init · main → _Startup)가 없다. 캐시 스키마가 올라갔다.
  6. 프로덕션 코드에 라벨 리터럴이 단일 출처 밖에 없다.
"""
from __future__ import annotations

import ast
import re
from pathlib import Path

import docx
import pytest

from report_gen.function_analyzer import _function_info_pairs
from report_gen.requirements import _extract_function_blocks, _extract_function_info_from_docx
from report_gen.uds_generator import _format_function_block_lines
from report_gen.utils import _extract_call_names, normalize_prototype_text
from report_gen.validation import generate_called_calling_accuracy_report
from report_gen.validation_labels import (
    CALL_KEY_TO_ROW_LABEL,
    CALL_ROW_LABEL_TO_KEY,
    LABEL_CALLED_FUNCTION,
    LABEL_CALLING_FUNCTION,
)

REPO = Path(__file__).resolve().parents[2]


# ────────────────────────────────────────── 표본 ──────────────────────────────────────────
def _table_doc(path: Path, rows_by_fn):
    """빌더 표의 최소 모양(라벨|값 2열). rows_by_fn: [(fn_id, name, [(label, value), ...])]."""
    d = docx.Document()
    for fid, name, extra in rows_by_fn:
        d.add_heading(f"{fid}: {name}", level=3)
        t = d.add_table(rows=0, cols=2)
        for label, value in [("[ Function Information ]", ""), ("ID", fid), ("Name", name),
                             ("Prototype", f"void {name}(void)"), ("ASIL", "B"), ("Related ID", "SwFn_01")] + list(extra):
            c = t.add_row().cells
            c[0].text, c[1].text = label, value
    d.save(str(path))
    return path


def _sections():
    return {
        "function_details_by_name": {
            "f_one": {"name": "f_one", "calls_list": ["f_two"]},
            "f_two": {"name": "f_two", "calls_list": []},
        },
        "function_table_rows": [["SwCom_01", "", "", "f_one"], ["SwCom_01", "", "", "f_two"]],
    }


def _accuracy(doc_path: Path, out: Path):
    generate_called_calling_accuracy_report(str(doc_path), "unused", str(out), source_sections=_sections())
    text = out.read_text(encoding="utf-8")
    m1 = re.search(r"Called exact match: `(\d+)` / `(\d+)`", text)
    m2 = re.search(r"Calling exact match: `(\d+)` / `(\d+)`", text)
    return (int(m1.group(1)), int(m1.group(2))), (int(m2.group(1)), int(m2.group(2)))


# ────────────────────────────── 1. 대응표 단일 출처 · 왕복 ──────────────────────────────
class TestLabelKeyMapping:
    def test_reference_convention(self):
        """정본 실측 921:0 — Called 행은 호출자(`calling`), Calling 행은 피호출자(`called`)."""
        assert CALL_ROW_LABEL_TO_KEY == {LABEL_CALLED_FUNCTION: "calling", LABEL_CALLING_FUNCTION: "called"}
        assert CALL_KEY_TO_ROW_LABEL == {"calling": LABEL_CALLED_FUNCTION, "called": LABEL_CALLING_FUNCTION}

    def test_writer_pairs_put_callers_under_called_row(self):
        pairs = dict(_function_info_pairs({"name": "f", "prototype": "void f(void)", "called": "callee_a\ncallee_b", "calling": "caller_x"}))
        assert pairs[LABEL_CALLED_FUNCTION] == "caller_x"
        assert pairs[LABEL_CALLING_FUNCTION] == "callee_a\ncallee_b"

    def test_writer_pairs_empty_is_na_for_both(self):
        pairs = dict(_function_info_pairs({"name": "f", "prototype": "void f(void)"}))
        assert pairs[LABEL_CALLED_FUNCTION] == "N/A" and pairs[LABEL_CALLING_FUNCTION] == "N/A"

    def test_docx_reader_maps_rows_back_to_the_same_keys(self, tmp_path):
        p = _table_doc(tmp_path / "ref.docx", [
            ("SwUFn_0001", "f_one", [(LABEL_CALLED_FUNCTION, "caller_x"), (LABEL_CALLING_FUNCTION, "callee_a\ncallee_b")]),
        ])
        info = _extract_function_info_from_docx(docx.Document(str(p)))["SwUFn_0001"]
        assert info["calling"] == "caller_x"
        assert info["called"] == "callee_a\ncallee_b"

    def test_docx_reader_continuation_rows_follow_the_mapping(self, tmp_path):
        """라벨 없는 이어짐 행(값만 있는 행)도 직전 라벨의 키로 붙는다."""
        p = _table_doc(tmp_path / "cont.docx", [
            ("SwUFn_0001", "f_one", [(LABEL_CALLED_FUNCTION, "caller_x"), ("", "caller_y"),
                                     (LABEL_CALLING_FUNCTION, "callee_a"), ("", "callee_b")]),
        ])
        info = _extract_function_info_from_docx(docx.Document(str(p)))["SwUFn_0001"]
        assert info["calling"].split("\n") == ["caller_x", "caller_y"]
        assert info["called"].split("\n") == ["callee_a", "callee_b"]

    def test_writer_to_reader_round_trip_is_identity(self, tmp_path):
        src = {"name": "f_one", "prototype": "void f_one(void)", "called": "callee_a\ncallee_b", "calling": "caller_x"}
        pairs = [(lbl, val) for lbl, val in _function_info_pairs(dict(src)) if lbl in CALL_ROW_LABEL_TO_KEY]
        p = _table_doc(tmp_path / "rt.docx", [("SwUFn_0001", "f_one", pairs)])
        info = _extract_function_info_from_docx(docx.Document(str(p)))["SwUFn_0001"]
        assert (info["called"], info["calling"]) == (src["called"], src["calling"])

    def test_text_block_parser_and_renderer_round_trip(self):
        blocks = _extract_function_blocks("SwUFn_0201: Bar\nCalled Function caller_x\nCalling Function callee_a\n")
        b = blocks[0]
        assert b["calling"] == "caller_x" and b["called"] == "callee_a"      # 라벨 절반이 값에 남지 않는다
        lines = _format_function_block_lines(b)
        assert f"{LABEL_CALLED_FUNCTION}\tcaller_x" in lines
        assert f"{LABEL_CALLING_FUNCTION}\tcallee_a" in lines
        # 렌더된 "ID\tSwUFn_…" 줄이 파서의 `\bSwUFn_\d+` 블록 경계에 걸려 블록이 둘로 갈린다(파서의 옛 한계, 이 라운드 밖) —
        # Called/Calling 값을 가진 블록을 집는다.
        again = [x for x in _extract_function_blocks("\n".join(lines)) if "calling" in x or "called" in x][-1]
        assert (again["called"], again["calling"]) == (b["called"], b["calling"])

    def test_text_block_parser_accepts_colon_and_tab_separators(self):
        b = _extract_function_blocks("SwUFn_0201: Bar\nCalled Function: caller_x\nCalling Function\tcallee_a\n")[0]
        assert (b["calling"], b["called"]) == ("caller_x", "callee_a")


# ────────────────────────────── 2. 정본 관례 문서의 accuracy ──────────────────────────────
class TestReferenceConventionDocumentIsReadCorrectly:
    def test_reference_convention_doc_matches_the_analysis(self, tmp_path):
        p = _table_doc(tmp_path / "ref.docx", [
            ("SwUFn_0001", "f_one", [(CALL_KEY_TO_ROW_LABEL["called"], "f_two"), (CALL_KEY_TO_ROW_LABEL["calling"], "N/A")]),
            ("SwUFn_0002", "f_two", [(CALL_KEY_TO_ROW_LABEL["called"], "N/A"), (CALL_KEY_TO_ROW_LABEL["calling"], "f_one")]),
        ])
        assert _accuracy(p, tmp_path / "a.md") == ((2, 2), (2, 2))

    def test_old_convention_doc_is_a_mismatch(self, tmp_path):
        """대조군 — 옛 배치(Called 행에 피호출자)는 이제 불일치로 잡힌다. 이게 안 잡히면 방향 회귀가 침묵한다."""
        p = _table_doc(tmp_path / "old.docx", [
            ("SwUFn_0001", "f_one", [(LABEL_CALLED_FUNCTION, "f_two"), (LABEL_CALLING_FUNCTION, "N/A")]),
            ("SwUFn_0002", "f_two", [(LABEL_CALLED_FUNCTION, "N/A"), (LABEL_CALLING_FUNCTION, "f_one")]),
        ])
        (c_ok, c_n), (g_ok, g_n) = _accuracy(p, tmp_path / "b.md")
        assert (c_n, g_n) == (2, 2) and c_ok < 2 and g_ok < 2


# ────────────────────────────── 3. 소스 분석: 한 줄 프로토타입 · static · 이름만 ──────────────────────────────
_LIN_STYLE_C = """
static void helper(void)
{
}

/* LIN 드라이버 스타일 — `(` 가 다음 줄, 인자 앞에 주석 */
l_bool l_ifc_init
(
/* [IN] interface name */
l_ifc_handle iii
)
{
    helper();
    return 0;
}

STATIC void aliased(void)
{
}

void top(void)
{
    l_ifc_init(0);
    aliased();
}
"""


@pytest.fixture
def lin_tree(tmp_path):
    (tmp_path / "lin.c").write_text(_LIN_STYLE_C, encoding="utf-8")
    return tmp_path


class TestPrototypeText:
    @pytest.mark.parametrize("raw, expected", [
        ("l_bool l_ifc_init\n(\n/* [IN] interface name */\nl_ifc_handle iii\n)", "l_bool l_ifc_init ( l_ifc_handle iii )"),
        ("U8 s_Lib_SafeWriteQueue_Enqueue(ST_SAFE_WRITE_QUEUE* pst_Queue, \r\n      U16 u16t_Addr1, \r\n  U8 u8t_Data)",
         "U8 s_Lib_SafeWriteQueue_Enqueue(ST_SAFE_WRITE_QUEUE* pst_Queue, U16 u16t_Addr1, U8 u8t_Data)"),
        ("void f(void) // trailing", "void f(void)"),
        ("", ""),
        (None, ""),
    ])
    def test_one_line_no_comments(self, raw, expected):
        assert normalize_prototype_text(raw) == expected

    @pytest.mark.parametrize("text, expected", [
        ("static", True), ("static void", True), ("STATIC U8", True),          # 한 노드에 두 토큰(리뷰 W5)
        ("FAST_STATIC", True), ("U8 static_helper", False), ("statically", False), ("", False), (None, False),
    ])
    def test_static_token_judgement_is_token_based(self, text, expected):
        from workflow.code_parser.c_parser import _has_static_token

        assert _has_static_token(text) is expected

    def test_regex_fallback_uses_the_same_static_judgement(self, lin_tree):
        """tree-sitter 없이(regex 폴백)도 같은 답 — 예전엔 한쪽만 소문자화해 정책이 갈렸다."""
        from workflow.code_parser.c_parser import _extract_function_defs_regex_fallback

        text = (lin_tree / "lin.c").read_text(encoding="utf-8")
        fns = {f.name: f for f in _extract_function_defs_regex_fallback(text, "lin.c", set())}
        assert fns["helper"].is_static is True and fns["aliased"].is_static is True and fns["top"].is_static is False
        for f in fns.values():
            assert "\n" not in f.signature and "/*" not in f.signature

    def test_analysis_cache_key_carries_the_schema_version(self):
        """(리뷰 W2) cloudium/원격 루트는 소스 시그니처가 None 이라 `_sig_ok` 가 항상 True — 버전은 **키에** 있어야 인메모리 캐시를 깬다."""
        src = (REPO / "backend" / "helpers" / "uds.py").read_text(encoding="utf-8")
        m = re.search(r'key = f"\{source_root\}\\x00pp=.*?"', src)
        assert m and "_SOURCE_SECTIONS_SCHEMA_VERSION" in m.group(0), m.group(0) if m else "cache key f-string not found"

    def test_single_source_reexport(self):
        from workflow.code_parser.c_parser import normalize_prototype_text as parser_side

        assert normalize_prototype_text is parser_side

    def test_parser_signatures_are_one_line_and_static_is_seen(self, lin_tree):
        from workflow.code_parser.c_parser import parse_c_project

        fns = {f["name"]: f for f in parse_c_project(str(lin_tree), preprocess=False)["functions"]}
        assert {"helper", "l_ifc_init", "aliased", "top"} <= set(fns)
        for f in fns.values():
            assert "\n" not in f["signature"] and "/*" not in f["signature"], f
        assert fns["l_ifc_init"]["signature"] == "l_bool l_ifc_init ( l_ifc_handle iii )"
        assert fns["helper"]["is_static"] is True
        assert fns["aliased"]["is_static"] is True          # 매크로 별칭(STATIC) — 변수 static 별칭과 같은 출처
        assert fns["top"]["is_static"] is False
        assert fns["l_ifc_init"]["is_static"] is False

    def test_analysis_prototype_keeps_static_and_called_is_names_only(self, lin_tree):
        from report_gen.uds_generator import generate_uds_source_sections

        sec = generate_uds_source_sections(str(lin_tree), preprocess=False, max_files=10, max_items=50)
        d = sec["function_details_by_name"]
        assert d["helper"]["prototype"].startswith("static ")
        assert "\n" not in d["l_ifc_init"]["prototype"] and "/*" not in d["l_ifc_init"]["prototype"]
        # called 텍스트 = 이름만(프로토타입 아님) — 여러 줄 원문이 칸에 들어갈 길을 뿌리에서 막는다
        assert sorted(d["top"]["called"].split("\n")) == ["aliased", "l_ifc_init"]     # 분석은 호출 목록을 정렬해 준다
        assert d["l_ifc_init"]["called"] == "helper"
        assert d["top"]["called"].split("\n") == list(d["top"]["calls_list"])          # 텍스트 = 이름 목록 그대로
        # (리뷰 W4) static 복원의 2차 효과 — 접두사(s_/g_) 없는 함수의 Type 은 static 이면 Internal(내부 링키지), 아니면 I/F
        types = {str(r[3]).strip(): next((str(v).strip() for v in r if str(v).strip() in ("Internal", "I/F")), "")
                 for r in sec["function_table_rows"] if isinstance(r, list) and len(r) > 3}
        assert types["helper"] == "Internal" and types["aliased"] == "Internal" and types["top"] == "I/F"


# ────────────────────────────── 3b. 실제 생성 경로 — 문서 칸의 방향·내용 ──────────────────────────────
class TestGeneratedDocumentCells:
    """`generate_uds_docx` 로 문서를 만들어 되읽는다 — Called 행 = 호출자 · Calling 행 = 피호출자 · 칸엔 이름만(괄호 0)."""

    @pytest.fixture(autouse=True)
    def _no_reference_suds(self, monkeypatch, tmp_path):
        import config

        monkeypatch.setattr(config, "UDS_REF_SUDS_PATH", str(tmp_path / "no_such_reference.docx"), raising=False)
        monkeypatch.setattr(config, "resolve_uds_template_path", lambda: "", raising=False)

    @staticmethod
    def _payload():
        base = {"description": "설명", "asil": "QM", "related": "", "inputs": [], "outputs": [], "precondition": "",
                "globals_global": [], "globals_static": [], "logic": ""}
        # 호출자 이름은 대소문자를 섞는다 — 빌더의 callers_map 은 소문자 정규화라 원래 표기로 되돌리는지까지 본다.
        fd = {
            "SwUFn_0001": dict(base, id="SwUFn_0001", name="Top_Init", prototype="void Top_Init(void)",
                               calls_list=["l_ifc_init"], called="l_ifc_init", calling=""),
            "SwUFn_0002": dict(base, id="SwUFn_0002", name="l_ifc_init", prototype="l_bool l_ifc_init ( l_ifc_handle iii )",
                               calls_list=[], called="", calling=""),
        }
        return {"project_name": "T", "overview": "o", "requirements": "r", "interfaces": "i", "uds_frames": "u", "notes": "n",
                "function_details": fd,
                "function_details_by_name": {v["name"].lower(): v for v in fd.values()},
                "call_map": {"Top_Init": ["l_ifc_init"]}}

    @staticmethod
    def _template(tmp_path):
        """정본 heading 을 가진 템플릿 — 라이브(KJPDS02)가 타는 `_build_function_info_table` 분기. 무템플릿 분기는 다른 라이터라
        한쪽만 검사하면 다른 쪽 회귀가 침묵한다(뮤테이션 1차: 템플릿 분기 변이 4개가 전부 생존했다)."""
        tpl = docx.Document()
        tpl.add_heading("Software Unit Design", level=1)
        for i, n in enumerate(("Top_Init", "l_ifc_init"), start=1):
            tpl.add_heading(f"SwUFn_{i:04d}: {n}", level=4)
            t = tpl.add_table(rows=3, cols=6)
            t.cell(0, 0).text = "[ Function Information ]"
        p = tmp_path / "tpl.docx"
        tpl.save(str(p))
        return str(p)

    def _generate(self, tmp_path, branch: str):
        from report_gen.docx_builder import generate_uds_docx

        out = tmp_path / f"u_{branch}.docx"
        generate_uds_docx(self._template(tmp_path) if branch == "template" else None, self._payload(), str(out))
        return out

    @pytest.mark.parametrize("branch", ["template", "no_template"])
    def test_rows_follow_the_reference_convention_and_hold_names_only(self, tmp_path, branch):
        from report_gen.validation import read_back_function_info

        out = self._generate(tmp_path, branch)
        info = {v["name"]: v for v in read_back_function_info(str(out)).values()}
        top, lif = info["Top_Init"], info["l_ifc_init"]
        # 되읽기 키는 내부 의미(called=피호출자 · calling=호출자); 행 라벨은 대응표가 정한다
        assert top["called"] == "l_ifc_init" and top["calling"] == "N/A"
        assert lif["calling"] == "Top_Init" and lif["called"] == "N/A"       # 원래 표기(소문자 'top_init' 이 아니다)
        for v in (top, lif):
            assert "(" not in v["called"] and "(" not in v["calling"], v      # 프로토타입이 칸에 실리지 않는다
        assert lif["prototype"] == "l_bool l_ifc_init ( l_ifc_handle iii )"

    @pytest.mark.parametrize("branch", ["template", "no_template"])
    def test_raw_rows_put_callers_under_called_function(self, tmp_path, branch):
        """되읽기 매핑을 거치지 않고 표의 원문 라벨|값을 직접 본다 — 리더가 뒤집혀도 이 검사는 안 속는다."""
        from report_gen.docx_text import cell_text

        out = self._generate(tmp_path, branch)
        rows = {}
        for t in docx.Document(str(out)).tables:
            cells0 = [cell_text(c).strip() for c in t.rows[1].cells] if len(t.rows) > 1 else []
            if not any("SwUFn_0002" in c for c in cells0):
                continue
            for r in t.rows:
                labels = [cell_text(c).strip() for c in r.cells]
                if labels and labels[0] in CALL_ROW_LABEL_TO_KEY:
                    rows[labels[0]] = labels[-1]
        assert rows[LABEL_CALLED_FUNCTION] == "Top_Init"     # l_ifc_init 을 부르는 쪽(원래 표기)
        assert rows[LABEL_CALLING_FUNCTION] == "N/A"         # l_ifc_init 이 부르는 쪽 없음


# ────────────────────────────── 4. 이름 파서 ──────────────────────────────
class TestExtractCallNamesMultiLine:
    @pytest.mark.parametrize("cell, expected", [
        # run 2086 accuracy 잔여 3건의 실제 칸 텍스트
        ("l_bool l_ifc_init\n(\n/* [IN] interface name */\nl_ifc_handle iii\n)\nl_bool l_sys_init ()\nvoid ld_init(void)",
         ["l_ifc_init", "l_sys_init", "ld_init"]),
        ("l_bool l_ifc_init\n(\n/* [IN] interface name */\nl_ifc_handle iii\n)\nl_bool l_sys_init ()\nvoid ld_init(void)\nvoid lin_lld_tx_wake_up ()",
         ["l_ifc_init", "l_sys_init", "ld_init", "lin_lld_tx_wake_up"]),
        ("l_u8 lin_lld_init\n(\n)\nvoid LinInit (void)\nvoid LinInt_StopInterruptMode(void)\nvoid Comms_Init(  )",
         ["lin_lld_init", "LinInit", "LinInt_StopInterruptMode", "Comms_Init"]),
        # 이름만 · 한 줄 프로토타입 · 주석 줄 · 빈 줄 — 예전과 같은 답
        ("a\nb", ["a", "b"]),
        ("void f(void)\n// note\n\nU8 g(U8 x)", ["f", "g"]),
        ("/* only a comment */", []),
        ("N/A", []),
        # 주석 안의 호출 모양 토큰은 이름이 아니다 — 주석을 먼저 지우지 않으면 `bar`·`helper` 가 이름으로 잡힌다
        ("foo /* was bar() */", ["foo"]),
        ("a\n/* see\n helper() */\nb", ["a", "b"]),
        ("x // calls y()", ["x"]),
    ])
    def test_names(self, cell, expected):
        assert _extract_call_names(cell) == expected


# ────────────────────────────── 5. 리터럴 특례 · 캐시 스키마 ──────────────────────────────
class TestNoNameLiteralSpecialCases:
    def test_docx_builder_has_no_function_name_literals(self):
        src = (REPO / "report_gen" / "docx_builder.py").read_text(encoding="utf-8")
        code = "\n".join(ln for ln in src.splitlines() if not ln.lstrip().startswith("#"))
        assert 'fn_key == "wake_up_setting"' not in code
        assert "_Startup" not in code
        assert '"l_ifc_init"' not in code

    def test_schema_version_advanced_past_v16(self):
        from backend.helpers.uds import _SOURCE_SECTIONS_SCHEMA_VERSION as v

        assert re.fullmatch(r"v\d+", v) and int(v[1:]) >= 17


# ────────────────────────────── 5b. 기대측 정의가 **라이브 기록 관문**을 지난다 ──────────────────────────────
class TestExpectedSideReachesTheLiveRecord:
    """R55 W2 는 recorder 가 `quality_eval["accuracy"]` 에서 인양하게 했지만 라이브 경로는 quick gate dict(accuracy 없음)를
    기록해 한 번도 발화하지 않았다(run 2086·2087 meta 에 키 없음). 이제 단일 관문 `_record_uds_run` 이 리포트에서 직접 싣는다."""

    def _capture(self, monkeypatch):
        import workflow.quality.recorder as R

        seen = {}

        def fake(quality_eval, **kwargs):
            seen["kwargs"] = kwargs
            return 1

        monkeypatch.setattr(R, "record_uds_run", fake)
        return seen

    def _artifacts(self, tmp_path, side_line: str):
        out = tmp_path / "u.docx"
        docx.Document().save(str(out))
        acc = tmp_path / "u.accuracy.md"
        acc.write_text("# Called/Calling Accuracy Report\n\n" + side_line + "\n- Called exact match: `1` / `1` (100.0%)\n", encoding="utf-8")
        return out, acc

    def test_quick_gate_record_carries_expected_side_from_the_report(self, tmp_path, monkeypatch):
        from backend.helpers.uds import _record_uds_run

        seen = self._capture(monkeypatch)
        out, acc = self._artifacts(tmp_path, "- Expected side: pipeline source_sections (the analysis that built the document)")
        _record_uds_run({"quick_gate": {}}, source_root="", out_path=out, ai_used=False, extra_meta={"entry": "t"}, accuracy_report=acc)
        assert seen["kwargs"]["meta"]["accuracy_expected_side"] == "pipeline"
        assert seen["kwargs"]["meta"]["entry"] == "t"

    def test_re_analysis_side_is_recorded_as_such(self, tmp_path, monkeypatch):
        from backend.helpers.uds import _record_uds_run

        seen = self._capture(monkeypatch)
        out, acc = self._artifacts(tmp_path, "- Expected side: re-analysis of source_root `x` (may differ)")
        _record_uds_run({}, source_root="", out_path=out, accuracy_report=acc)
        assert seen["kwargs"]["meta"]["accuracy_expected_side"] == "re-analysis"

    @pytest.mark.parametrize("case", ["no_report", "missing_file", "no_side_line"])
    def test_no_definition_means_no_key(self, tmp_path, monkeypatch, case):
        from backend.helpers.uds import _record_uds_run

        seen = self._capture(monkeypatch)
        out, acc = self._artifacts(tmp_path, "- Relation mode: `code`")
        report = {"no_report": None, "missing_file": tmp_path / "nope.accuracy.md", "no_side_line": acc}[case]
        _record_uds_run({}, source_root="", out_path=out, accuracy_report=report)
        assert "accuracy_expected_side" not in seen["kwargs"]["meta"]

    def test_live_async_caller_passes_the_report(self):
        src = (REPO / "backend" / "helpers" / "uds.py").read_text(encoding="utf-8")
        i = src.index('"entry": "jenkins_generate_async"')
        call = src[src.rfind("_record_uds_run(", 0, i):i]
        assert re.search(r"accuracy_report\s*=\s*accuracy_path", call), "라이브 비동기 기록 호출부가 accuracy 리포트를 넘기지 않는다"


# ────────────────────────────── 6. 라벨 리터럴 단일 출처 ──────────────────────────────
def _string_constants_outside_docstrings(path: Path):
    tree = ast.parse(path.read_text(encoding="utf-8-sig"))   # `workflow/pipeline.py` 는 BOM 으로 시작한다
    doc_ids = set()
    for node in ast.walk(tree):
        if isinstance(node, (ast.Module, ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef)):
            body = getattr(node, "body", [])
            if body and isinstance(body[0], ast.Expr) and isinstance(getattr(body[0], "value", None), ast.Constant):
                doc_ids.add(id(body[0].value))
    for node in ast.walk(tree):
        if isinstance(node, ast.Constant) and isinstance(node.value, str) and id(node) not in doc_ids:
            yield node.lineno, node.value


class TestLabelLiteralSingleSource:
    def test_production_code_uses_the_constants(self):
        offenders = []
        wanted = {LABEL_CALLED_FUNCTION.casefold(), LABEL_CALLING_FUNCTION.casefold()}
        for sub in ("report_gen", "backend", "workflow", "generators", "tools"):     # (리뷰 I2) tools 포함 · 대소문자 무시
            for py in (REPO / sub).rglob("*.py"):
                if "venv" in py.parts or "site-packages" in py.parts or py.name == "validation_labels.py":
                    continue
                for lineno, val in _string_constants_outside_docstrings(py):
                    if val.strip().casefold() in wanted:
                        offenders.append(f"{py.relative_to(REPO)}:{lineno}")
        assert offenders == [], offenders
