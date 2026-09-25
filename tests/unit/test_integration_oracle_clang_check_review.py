"""(R16b review round 1) clang check of SITS integration values — each reproduced false verdict, pinned.

Each case is a small project; claims come from the interprocedural oracle (or are forced where the case is a wrong
claim the harness must not agree with). clang (``--target=msp430``) must be on PATH for the checks that compile.
"""
from __future__ import annotations

import json
import os
import shutil

import pytest

from generators import c_project_context as cpc
from generators.c_source_oracle import evaluate_outputs
from generators.integration_oracle import TEST_EVIDENCE_HEADERS, CalleeProvider
from scripts import integration_oracle_clang_check as ioc

ROOT = os.path.join(os.sep, "virt_r16b_review")
needs_clang = pytest.mark.skipif(shutil.which("clang") is None, reason="clang (msp430 target) not installed")
H = """#ifndef C_H
#define C_H
typedef unsigned char U8;
typedef unsigned int U16;
typedef signed int S16;
%s
#endif
"""


def _project(files):
    files = {os.path.join(ROOT, k): v for k, v in files.items()}
    ctx = cpc.build_project_context(files)
    scopes = cpc.build_scopes(ctx, [p for p in files if p.endswith(".c")])
    parser = cpc.shared_parser()
    return files, scopes, CalleeProvider(ctx, files, scopes, parser), parser


def _run(proj, name, cases, tmp_path, unit="a.c", interpreted: str | None = "model"):
    """Claims (the model's derived values, plus forced ones) and the harness verdicts. ``interpreted``: the callees
    the model ran, as the Test Evidence sheet records them (``"model"``), or None — run every body (a forced claim
    about what C does, whatever the model could interpret)."""
    files, scopes, provider, parser = proj
    path = os.path.join(ROOT, unit)
    u = {"name": name, "source_text": files[path], "source_path": path, "source_text_complete": True,
         "project_scope": scopes[path], "callee_provider": provider}
    claims, derived = [], []
    for inputs, outputs, forced in cases:
        rec = evaluate_outputs(u, [inputs], [list(outputs)])[0]
        d = {k: v["value"] for k, v in rec["outputs"].items() if "value" in v}
        derived.append(d)
        ip = rec.get("interprocedural") or {}
        ran = set(ip.get("inlined") or ()) if interpreted == "model" else None
        effects = set(ip.get("effects_only") or ()) if interpreted == "model" else set()
        claims.append({"inputs": inputs, "outputs": {**d, **forced},
                       "possible_ub": list(rec.get("possible_undefined_behavior") or []), "interpreted": ran,
                       "effects_only": effects})
    final, notes = ioc.check_group(name, path, claims, files, scopes, provider, parser, work_dir=str(tmp_path / name))
    return derived, {k: v for k, (v, _d) in (final or {}).items()}, notes


# ── the model: a declaration after an earlier case label is in scope for the later cases ────────────────────────

def test_a_local_declared_after_a_case_label_is_that_local_in_a_later_case():
    # review R1 C2 (model side): jumping to ``case 1`` passed ``U8 cnt`` — the later ``cnt = 7U`` writes the local,
    # not the global of the same name; the global stays at its input (C11 6.2.1p4, 6.8.4.2p7)
    for decl in ("U8 cnt = 9U;", "static U8 cnt;"):
        text = ('#include "c.h"\n'
            "U8 cnt;\nU8 g_out;\n"
            "void f(U8 k) {\n    switch (k) {\n    case 0U:\n        k = 1U;\n        " + decl + "\n        cnt = 3U;\n"
            "        break;\n    case 1U:\n        cnt = 7U;\n        g_out = cnt;\n        break;\n"
            "    default:\n        break;\n    }\n}\n")
        files = {os.path.join(ROOT, "c.h"): H % "extern U8 cnt;\nextern U8 g_out;", os.path.join(ROOT, "s.c"): text}
        ctx = cpc.build_project_context(files)
        scope = cpc.build_scopes(ctx, [os.path.join(ROOT, "s.c")])[os.path.join(ROOT, "s.c")]
        unit = {"name": "f", "source_text": text, "source_path": os.path.join(ROOT, "s.c"), "source_text_complete": True,
                "project_scope": scope}
        rec = evaluate_outputs(unit, [{"k": 1, "cnt": 2}], [["cnt", "g_out"]])[0]
        assert rec["outputs"]["cnt"].get("value") == 2, (decl, rec["outputs"])      # the global is untouched
        assert rec["outputs"]["g_out"].get("value") == 7, (decl, rec["outputs"])    # the local was written


# ── the harness ───────────────────────────────────────────────────────────────────────────────────────────────

@needs_clang
def test_a_static_local_after_a_case_label_is_renamed_to_the_end_of_the_switch(tmp_path):
    # review R1 C2: the rename stopped at the ``case`` — ``cnt = 7U`` in ``case 1`` wrote the global
    proj = _project({"c.h": H % "extern U8 cnt;\nextern U8 g_out;\nU8 b_f(U8 k);",
                     "a.c": '#include "c.h"\nU8 cnt;\nU8 g_out;\nvoid e(U8 k) { g_out = b_f(k); }\n',
                     "b.c": '#include "c.h"\nU8 b_f(U8 k)\n{\n    switch (k) {\n    case 0U:\n        k = 1U;\n'
                            '        static U8 cnt;\n        cnt = 3U;\n        break;\n    case 1U:\n        cnt = 7U;\n'
                            '        break;\n    default:\n        break;\n    }\n    return cnt;\n}\n'})
    derived, verdicts, notes = _run(proj, "e", [({"k": 1, "cnt": 2}, ["g_out", "cnt"], {}),
                                                ({"k": 1, "cnt": 2}, ["cnt"], {"cnt": 7})], tmp_path)
    assert derived[0] == {"g_out": 2, "cnt": 2}   # after the switch, ``cnt`` is the global again
    assert verdicts == {(0, "g_out"): "agree", (0, "cnt"): "agree", (1, "cnt"): "mismatch"}, notes


@needs_clang
def test_a_const_table_keeps_its_values_whichever_declaration_comes_first(tmp_path):
    # review R1 C3: the entry's ``extern const U8 tab[4];`` won over the definition's ``{1,2,3,4}`` → false mismatch
    proj = _project({"c.h": H % "extern U8 g_out;\nextern U8 g_first;\nextern const U8 tab[4];\nU8 b_get(U8 i);",
                     "a.c": '#include "c.h"\nU8 g_out;\nU8 g_first;\nvoid e(void) { g_first = tab[0]; g_out = b_get(1U); }\n',
                     "b.c": '#include "c.h"\nconst U8 tab[4] = {1U, 2U, 3U, 4U};\nU8 b_get(U8 i) { return tab[i]; }\n'})
    derived, verdicts, notes = _run(proj, "e", [({}, ["g_out"], {}), ({}, ["g_out"], {"g_out": 3})], tmp_path)
    assert derived[0] == {"g_out": 2}
    assert verdicts == {(0, "g_out"): "agree", (1, "g_out"): "mismatch"}, notes


@needs_clang
def test_a_callee_name_in_a_call_note_is_not_an_evaluation_error(tmp_path):
    # review R1 W1: ``in call to 'this->do_shift(1)'`` matched "shift" — a goto (harness limit) became a contradiction
    proj = _project({"c.h": H % "extern U8 g_in;\nextern U8 g_out;\nextern U16 g_sum;\nU8 do_shift(U8 v);",
                     "a.c": '#include "c.h"\nU8 g_in;\nU8 g_out;\nU16 g_sum;\n'
                            'void e(void) { (void)do_shift(g_in); g_out = 7U; }\n',
                     "b.c": '#include "c.h"\nU8 do_shift(U8 v) { if (v) { goto done; } g_sum = 1U; done: return v; }\n'})
    # every body run (the model reads do_shift's goto as effects — with its record the stub would stand in)
    derived, verdicts, _notes = _run(proj, "e", [({"g_in": 1}, ["g_out"], {})], tmp_path, interpreted=None)
    assert derived[0] == {"g_out": 7} and verdicts == {(0, "g_out"): "unchecked:constexpr_limit"}


@needs_clang
def test_an_observable_is_the_object_not_the_entry_units_constant_of_that_name(tmp_path):
    # review R1 W2: ``LIMIT`` is an object in b.c and a ``static const`` in a.c — the observer read a.c's constant
    proj = _project({"c.h": H % "extern U8 g_in;\nextern U8 g_out;\nvoid b_set(void);",
                     "a.c": '#include "c.h"\nstatic const U8 LIMIT = 5U;\nU8 g_in;\nU8 g_out;\n'
                            'void e(void) { g_out = LIMIT; b_set(); }\n',
                     "b.c": '#include "c.h"\nU8 LIMIT;\nvoid b_set(void) { LIMIT = (U8)(g_in + 1U); }\n'})
    derived, verdicts, notes = _run(proj, "e", [({"g_in": 8}, ["LIMIT", "g_out"], {}),
                                                ({"g_in": 8}, ["LIMIT"], {"LIMIT": 5})], tmp_path)
    assert derived[0] == {"LIMIT": 9, "g_out": 5}
    assert verdicts == {(0, "LIMIT"): "agree", (0, "g_out"): "agree", (1, "LIMIT"): "mismatch"}, notes


@needs_clang
def test_an_input_named_by_a_callee_units_enumerator_is_read_there(tmp_path):
    # review R1 W3: ``"g_mode": "K_RUN"`` — the enumerator is visible only in b.c; the entry's constants left the fill
    proj = _project({"c.h": H % "extern U8 g_mode;\nextern U8 g_out;\nvoid b_step(void);",
                     "bh.h": "#ifndef BH_H\n#define BH_H\nenum { K_RUN = 7 };\n#endif\n",
                     "a.c": '#include "c.h"\nU8 g_mode;\nU8 g_out;\nvoid e(void) { b_step(); }\n',
                     "b.c": '#include "c.h"\n#include "bh.h"\n'
                            'void b_step(void) { if (g_mode == K_RUN) { g_out = 1U; } else { g_out = 2U; } }\n'})
    derived, verdicts, _notes = _run(proj, "e", [({"g_mode": "K_RUN"}, ["g_out"], {})], tmp_path)
    assert derived[0] == {"g_out": 1} and verdicts == {(0, "g_out"): "agree"}


@needs_clang
def test_a_sequence_stub_never_replaces_the_function_under_test(tmp_path):
    # review R1 W4: ``e() return`` made the entry return the stub value at once (the model ignores it)
    proj = _project({"c.h": H % "extern U8 g_out;", "a.c": '#include "c.h"\nU8 g_out;\n'
                     'U8 e(void) { g_out = 5U; return 6U; }\n'})
    derived, verdicts, _notes = _run(proj, "e", [({"e() return": 9}, ["g_out", "return"], {})], tmp_path)
    assert derived[0] == {"g_out": 5, "return": 6} and verdicts == {(0, "g_out"): "agree", (0, "return"): "agree"}


@needs_clang
def test_a_run_that_reaches_a_function_clang_finds_unsequenced_is_set_aside(tmp_path):
    # review R1 W5: the unsequenced diagnostic cut the callee (weak agree) or failed the entry's group
    proj = _project({"c.h": H % "extern U8 g_in;\nextern U8 g_out;\nU8 b_unseq(U8 v);\nU8 b_ok(U8 v);",
                     "a.c": '#include "c.h"\nU8 g_in;\nU8 g_out;\n'
                            'void e(void) { g_out = 4U; if (g_in > 5U) { (void)b_unseq(g_in); } else { g_out = b_ok(1U); } }\n',
                     "b.c": '#include "c.h"\nU8 b_unseq(U8 v) { U8 i = v; U8 j = (U8)(i++ + i++); return j; }\n'
                            'U8 b_ok(U8 v) { return (U8)(v + 1U); }\n'})
    claims = [({"g_in": 9}, ["g_out"], {"g_out": 4}), ({"g_in": 1}, ["g_out"], {})]
    derived, verdicts, notes = _run(proj, "e", claims, tmp_path, interpreted=None)
    assert derived[1] == {"g_out": 2}
    assert verdicts == {(0, "g_out"): "unchecked:unsequenced_function_reached", (1, "g_out"): "agree"}, notes
    assert notes["unsequenced"] == ["b_unseq"]


@needs_clang
def test_a_disclosure_must_explain_every_failing_fill(tmp_path):
    # review R1 W6: only fill 0's division by zero was disclosed; fills 90/201 index out of bounds
    proj = _project({"c.h": H % "extern U8 g_d;\nextern U8 g_out;\nextern U8 g_buf[4];",
                     "a.c": '#include "c.h"\nU8 g_d;\nU8 g_out;\nU8 g_buf[4];\n'
                            'void e(void) { g_out = (U8)(100U / g_d); g_out = g_buf[g_d]; g_out = 5U; }\n'})
    files, scopes, provider, parser = proj
    path = os.path.join(ROOT, "a.c")
    for disclosed, expected in ((["division_by_unknown"], "eval_error"),
                                (["division_by_unknown", "index_unknown"], "unchecked:possible_ub_disclosed")):
        claims = [{"inputs": {}, "outputs": {"g_out": 5}, "possible_ub": disclosed}]
        final, _notes = ioc.check_group("e", path, claims, files, scopes, provider, parser,
                                        work_dir=str(tmp_path / expected.replace(":", "_")))
        assert final[(0, "g_out")][0] == expected, (disclosed, final)


@needs_clang
def test_a_static_array_sized_by_a_macro_and_same_named_statics_of_two_units(tmp_path):
    # review R1 W11 (the size macro was lost with the declaration) and W12 (two ``static U8 s_state`` failed the group)
    proj = _project({"c.h": H % "#define BUF_LEN 4U\nextern U8 g_out;\nU8 b_buf(void);\nvoid b_tick(void);",
                     "a.c": '#include "c.h"\nstatic U8 s_state;\nU8 g_out;\n'
                            'void e(void) { s_state = 1U; b_tick(); g_out = (U8)(b_buf() + 1U); }\n',
                     "b.c": '#include "c.h"\nstatic U8 s_state;\n'
                            'U8 b_buf(void) { static U8 buf[BUF_LEN]; buf[1] = 6U; return buf[1]; }\n'
                            'void b_tick(void) { s_state = 2U; }\n'})
    derived, verdicts, notes = _run(proj, "e", [({}, ["g_out"], {"g_out": 7}), ({}, ["s_state"], {"s_state": 1}),
                                                ({}, ["s_state"], {"s_state": 2})], tmp_path, interpreted=None)
    assert derived[0] == {}   # the model refuses both callees (``linkage_collision:s_state``) — the claims are forced
    # the entry's ``s_state`` is a.c's: b_tick writes b.c's — a claim of 2 is the other object's value
    assert verdicts == {(0, "g_out"): "agree", (1, "s_state"): "agree", (2, "s_state"): "mismatch"}, notes


def _sheet(tmp_path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Evidence"
    ws.append(TEST_EVIDENCE_HEADERS)
    for r in rows:
        row = dict.fromkeys(TEST_EVIDENCE_HEADERS, "")
        row.update({"Status": "derived", "Case": 1, **r})
        ws.append([row[h] for h in TEST_EVIDENCE_HEADERS])
    path = tmp_path / "s.xlsx"
    wb.save(path)
    return str(path)


@needs_clang
def test_groups_never_share_a_file_and_the_exit_code_says_when_nothing_was_checked(tmp_path):
    # review R1 C1: ``Init`` (a.c) and ``init`` (b.c) compiled in one folder on a case-insensitive disk — with two
    # workers one read the other's result and agreed with a wrong claim. W7: a run that checked nothing exited 0.
    root = tmp_path / "src"
    root.mkdir()
    (root / "c.h").write_text(H % "extern U8 g_a;\nextern U8 g_b;", encoding="utf-8", newline="\n")
    (root / "a.c").write_text('#include "c.h"\nU8 g_a;\nvoid Init(void) { g_a = 1U; }\n', encoding="utf-8", newline="\n")
    (root / "b.c").write_text('#include "c.h"\nU8 g_b;\nvoid init(void) { g_b = 2U; }\n', encoding="utf-8", newline="\n")
    a, b = str((root / "a.c").resolve()), str((root / "b.c").resolve())
    xlsm = _sheet(tmp_path, [{"Test Case ID": "T1", "Entry function": "Init", "Observable": "g_a", "Expected": 77,
                              "Source path": a, "Inputs JSON": "{}"},
                             {"Test Case ID": "T2", "Entry function": "init", "Observable": "g_b", "Expected": 2,
                              "Source path": b, "Inputs JSON": "{}"}])
    for workers in ("1", "2"):
        out = tmp_path / f"r{workers}.json"
        rc = ioc.main(["--xlsm", xlsm, "--source-root", str(root), "--out", str(out), "--workers", workers])
        report = json.loads(out.read_text(encoding="utf-8"))
        assert rc == 1 and report["verdicts"] == {"mismatch": 1, "agree": 1}, (workers, report)
        assert report["mismatches"][0]["output"] == "g_a"
    out = tmp_path / "bad.json"
    assert ioc.main(["--xlsm", xlsm, "--source-root", str(root), "--out", str(out), "--target", "no-such-target"]) == 2
    empty = _sheet(tmp_path, [{"Test Case ID": "T3", "Entry function": "Init", "Observable": "g_a", "Expected": 1,
                               "Source path": str(tmp_path / "elsewhere.c"), "Inputs JSON": "{}"}])
    assert ioc.main(["--xlsm", empty, "--source-root", str(root), "--out", str(out)]) == 2


# ── review round 2 ──────────────────────────────────────────────────────────────────────────────────────────────

_BIG = "".join(f"    if ((g_x & {1 << i}U) != 0U) {{ a = (U8)(a + {i + 1}U); }}\n" for i in range(7))


@needs_clang
def test_a_callee_the_model_did_not_run_is_never_run_for_its_undefined_behaviour(tmp_path):
    # review R2 CR2-1: over its path budget the model reads ``b_big`` as effects — it never saw ``g_tab[g_idx]``. The
    # harness ran the body, hit ``g_tab[90]`` and reported a contradiction for a value (7) the model rightly derived.
    proj = _project({"c.h": H % "extern U8 g_x;\nextern U8 g_idx;\nextern U8 g_tab[4];\nextern U8 g_w;\n"
                                "extern U8 g_res;\nvoid b_big(void);",
                     "a.c": '#include "c.h"\nU8 g_x;\nU8 g_idx;\nU8 g_tab[4];\nU8 g_w;\nU8 g_res;\n'
                            'void e(void) { b_big(); g_res = 7U; }\n',
                     "b.c": '#include "c.h"\nvoid b_big(void)\n{\n    U8 a = 0U;\n' + _BIG +
                            '    g_w = (U8)(a + g_tab[g_idx]);\n}\n'})
    derived, verdicts, notes = _run(proj, "e", [({}, ["g_res"], {})], tmp_path)
    assert derived[0] == {"g_res": 7}
    assert verdicts == {(0, "g_res"): "agree_with_stubbed_callees"}, notes
    # the same claim with every body run (no record of what the model ran) is the old, false contradiction
    _d, verdicts, _n = _run(proj, "e", [({}, ["g_res"], {})], tmp_path / "all", interpreted=None)
    assert verdicts == {(0, "g_res"): "eval_error"}


@needs_clang
def test_a_sequence_stub_whose_value_cannot_be_read_still_replaces_the_callee(tmp_path):
    # review R2 CR2-1 (s6): ``"F() return": "RET_UNKNOWN_NAME"`` — the model stubs F all the same; the harness ran F
    # (a division by an unset global) and reported a contradiction
    proj = _project({"c.h": H % "extern U8 g_div;\nextern U8 g_out;\nextern U8 g_res;\nU8 F(void);",
                     "a.c": '#include "c.h"\nU8 g_div;\nU8 g_out;\nU8 g_res;\nvoid e(void) { (void)F(); g_res = 7U; }\n',
                     "b.c": '#include "c.h"\nU8 F(void) { g_out = (U8)(100U / g_div); return g_out; }\n'})
    derived, verdicts, notes = _run(proj, "e", [({"F() return": "RET_UNKNOWN_NAME"}, ["g_res"], {})], tmp_path)
    assert derived[0] == {"g_res": 7}
    assert verdicts == {(0, "g_res"): "agree_with_stubbed_callees"}, notes


@needs_clang
def test_a_name_means_the_object_the_entry_unit_declares(tmp_path):
    # review R2 WR2-1 (s1, s5): the entry never touches ``s_state``/``x`` itself — the name must still mean the entry
    # unit's object (its helper writes 3; b.c's own ``s_state`` gets 2; c.c's external ``x`` gets 5)
    proj = _project({"c.h": H % "extern U8 g_out;\nvoid b_tick(void);\nvoid c_set(void);",
                     "a.c": '#include "c.h"\nstatic U8 s_state;\nstatic U8 x;\nU8 g_out;\n'
                            'static void helper(void) { s_state = 3U; }\n'
                            'void e(void) { helper(); b_tick(); c_set(); g_out = 4U; }\n',
                     "b.c": '#include "c.h"\nstatic U8 s_state;\nvoid b_tick(void) { s_state = 2U; }\n',
                     "c.c": '#include "c.h"\nU8 x;\nvoid c_set(void) { x = 5U; }\n'})
    cases = [({"s_state": 5, "x": 1}, ["s_state"], {"s_state": 3}), ({"s_state": 5, "x": 1}, ["s_state"], {"s_state": 5}),
             ({"s_state": 5, "x": 1}, ["x"], {"x": 1}), ({"s_state": 5, "x": 1}, ["x"], {"x": 5})]
    _derived, verdicts, notes = _run(proj, "e", cases, tmp_path, interpreted=None)
    assert verdicts == {(0, "s_state"): "agree", (1, "s_state"): "mismatch",
                        (2, "x"): "agree", (3, "x"): "mismatch"}, notes


def test_the_model_refuses_a_case_label_nested_in_a_block():
    # review R2 WR2-3: ``else { case 3U: … }`` — only direct labels were modelled; k=3 matched none and the model
    # derived the value of "no case taken"
    text = ('#include "c.h"\nU8 cnt;\nU8 g_out;\n'
            "void f(U8 k) {\n    switch (k) {\n    case 0U:\n        if (k == 0U) {\n            cnt = 1U;\n"
            "        } else {\n    case 3U:\n            cnt = 9U;\n        }\n        break;\n    default:\n"
            "        break;\n    }\n    g_out = cnt;\n}\n")
    files = {os.path.join(ROOT, "c.h"): H % "extern U8 cnt;\nextern U8 g_out;", os.path.join(ROOT, "n.c"): text}
    ctx = cpc.build_project_context(files)
    scope = cpc.build_scopes(ctx, [os.path.join(ROOT, "n.c")])[os.path.join(ROOT, "n.c")]
    unit = {"name": "f", "source_text": text, "source_path": os.path.join(ROOT, "n.c"), "source_text_complete": True,
            "project_scope": scope}
    rec = evaluate_outputs(unit, [{"k": 3, "cnt": 2}], [["cnt", "g_out"]])[0]
    assert "value" not in rec["outputs"]["cnt"] and "case_label_in_nested_block" in str(rec.get("reason")), rec


@needs_clang
def test_a_contradiction_proven_for_one_enum_type_survives_a_later_type_failing(tmp_path, monkeypatch):
    # review R2 IR2-1 (s12): the second underlying type timed out and the whole group — with its proven mismatch —
    # became unchecked
    proj = _project({"c.h": H % "typedef enum { E_A = 0, E_B = 1 } E_T;\nextern E_T g_e;\nextern U8 g_out;",
                     "a.c": '#include "c.h"\nE_T g_e;\nU8 g_out;\nvoid e(void) { g_e = E_B; g_out = 2U; }\n'})
    real = ioc._compile

    def flaky(source, layout, path, *a, **k):
        if os.path.basename(path).startswith("b1_"):
            return {"failure": "clang_failed:TimeoutExpired"}
        return real(source, layout, path, *a, **k)
    monkeypatch.setattr(ioc, "_compile", flaky)
    _derived, verdicts, notes = _run(proj, "e", [({}, ["g_out"], {"g_out": 3}), ({}, ["g_out"], {})], tmp_path)
    assert verdicts == {(0, "g_out"): "mismatch", (1, "g_out"): "unchecked:enum_base_failed"}, notes
    assert notes["failure_partial"].startswith("enum_base_unsigned int:")


def test_a_project_typedef_of_bool_or_a_stdint_name_is_what_the_project_says():
    # R16b (found through IR2-4): tree-sitter reads ``bool``/``int8_t`` as primitive types, so the scanner dropped
    # ``typedef unsigned char bool;`` — ``bool`` stayed ``_Bool`` (``(bool)2`` → 1) and ``uint8_t`` rank-unknown
    # (unguarded here: without build-configuration evidence an ``#ifndef uint8_t`` guard stays undecided — R17)
    files = {os.path.join(ROOT, "t.h"): H % "typedef unsigned char bool;\ntypedef unsigned char uint8_t;\n"
                                            "extern bool g_flag;\nextern uint8_t g_u;",
             os.path.join(ROOT, "t.c"): '#include "t.h"\nbool g_flag;\nuint8_t g_u;\n'
                                        "void f(void) { g_flag = (bool)2U; g_u = (uint8_t)300U; }\n"}
    ctx = cpc.build_project_context(files)
    scope = cpc.build_scopes(ctx, [os.path.join(ROOT, "t.c")])[os.path.join(ROOT, "t.c")]
    assert scope["types"]["bool"]["kind"] == "char" and scope["types"]["uint8_t"]["kind"] == "char"
    unit = {"name": "f", "source_text": files[os.path.join(ROOT, "t.c")], "source_path": os.path.join(ROOT, "t.c"),
            "source_text_complete": True, "project_scope": scope}
    rec = evaluate_outputs(unit, [{}], [["g_flag", "g_u"]])[0]
    assert {k: v.get("value") for k, v in rec["outputs"].items()} == {"g_flag": 2, "g_u": 44}
    # without the project's typedef, ``bool`` is still ``_Bool`` and ``uint8_t`` still the target's choice
    files2 = {os.path.join(ROOT, "u.c"): "_Bool g_b;\nvoid f(void) { g_b = (_Bool)2U; }\n"}
    ctx2 = cpc.build_project_context(files2)
    scope2 = cpc.build_scopes(ctx2, list(files2))[os.path.join(ROOT, "u.c")]
    assert scope2["globals"]["g_b"]["type"]["kind"] == "_Bool"


@needs_clang
def test_a_typedef_named_like_a_cxx_keyword_is_renamed_in_its_blocks(tmp_path):
    # review R2 IR2-4: ``typedef unsigned char bool;`` (KJPDS02 PE_Types.h) failed every group that used it
    proj = _project({"c.h": H % "typedef unsigned char bool;\nextern bool g_flag;\nextern U8 g_out;",
                     "a.c": '#include "c.h"\nbool g_flag;\nU8 g_out;\n'
                            'void e(void) { bool b = (bool)2U; g_flag = b; g_out = (U8)(g_flag + 1U); }\n'})
    derived, verdicts, notes = _run(proj, "e", [({}, ["g_flag", "g_out"], {})], tmp_path)
    assert derived[0] == {"g_flag": 2, "g_out": 3}   # an unsigned char, not C++'s bool (which would give 1)
    assert verdicts == {(0, "g_flag"): "agree", (0, "g_out"): "agree"}, notes


@needs_clang
def test_a_run_where_every_claim_ends_unchecked_exits_2(tmp_path):
    # review R2 WR2-2 (s14): the only claim hit a goto (a harness limit) — nothing was checked, yet the exit said 0
    root = tmp_path / "src"
    root.mkdir()
    (root / "c.h").write_text(H % "extern U8 g_in;\nextern U8 g_out;\nU8 b_goto(U8 v);", encoding="utf-8", newline="\n")
    (root / "a.c").write_text('#include "c.h"\nU8 g_in;\nU8 g_out;\nvoid e(void) { (void)b_goto(g_in); g_out = 7U; }\n',
                              encoding="utf-8", newline="\n")
    (root / "b.c").write_text('#include "c.h"\nU8 b_goto(U8 v) { if (v) { goto done; } v = 1U; done: return v; }\n',
                              encoding="utf-8", newline="\n")
    xlsm = _sheet(tmp_path, [{"Test Case ID": "T1", "Entry function": "e", "Observable": "g_out", "Expected": 7,
                              "Source path": str((root / "a.c").resolve()), "Inputs JSON": '{"g_in": 1}',
                              "Callees interpreted": "b_goto"}])
    out = tmp_path / "r.json"
    assert ioc.main(["--xlsm", xlsm, "--source-root", str(root), "--out", str(out)]) == 2
    assert json.loads(out.read_text(encoding="utf-8"))["verdicts"] == {"unchecked": 1}


# ── review round 3 ──────────────────────────────────────────────────────────────────────────────────────────────

_HEAVY = ("{\n    U8 a = 0U;\n" + _BIG + "    g_w = (U8)(a + (U8)(100U / g_x));\n}\n")


@needs_clang
def test_a_callee_run_at_one_site_and_read_as_effects_at_another_does_not_contradict(tmp_path):
    # review R3 W3-1 (s18): b_big runs with g_x = 128 (known) and again with g_x = g_in, over the path budget — read as
    # effects there. The harness ran both, hit ``100U / 0`` at the second, and called the model's g_res = 7 a
    # contradiction; the sheet now records which callees the model read as effects anywhere
    proj = _project({"c.h": H % "extern U8 g_x;\nextern U8 g_in;\nextern U8 g_w;\nextern U8 g_res;\nvoid b_big(void);",
                     "a.c": '#include "c.h"\nU8 g_x;\nU8 g_in;\nU8 g_w;\nU8 g_res;\n'
                            'void e(void) { g_x = 128U; b_big(); g_x = g_in; b_big(); g_res = 7U; }\n',
                     "b.c": '#include "c.h"\nvoid b_big(void)\n' + _HEAVY})
    derived, verdicts, notes = _run(proj, "e", [({}, ["g_res"], {})], tmp_path)
    assert derived[0] == {"g_res": 7}
    assert verdicts == {(0, "g_res"): "unchecked:callee_partly_interpreted"}, notes


@needs_clang
def test_same_named_statics_one_run_one_read_as_effects(tmp_path):
    # review R3 W3-1 (s21): a.c's light ``static h`` is interpreted, b.c's heavy ``static h`` (reached through b_tick)
    # is not — the C names are equal, so both members count as partly interpreted
    proj = _project({"c.h": H % "extern U8 g_x;\nextern U8 g_in;\nextern U8 g_w;\nextern U8 g_a;\nextern U8 g_res;\n"
                                "void b_tick(void);",
                     "a.c": '#include "c.h"\nU8 g_x;\nU8 g_in;\nU8 g_w;\nU8 g_a;\nU8 g_res;\n'
                            'static void h(void) { g_a = 1U; }\n'
                            'void e(void) { h(); g_x = g_in; b_tick(); g_res = 7U; }\n',
                     "b.c": '#include "c.h"\nstatic void h(void)\n' + _HEAVY + 'void b_tick(void) { h(); }\n'})
    derived, verdicts, notes = _run(proj, "e", [({}, ["g_res"], {})], tmp_path)
    assert derived[0] == {"g_res": 7}
    assert verdicts[(0, "g_res")] in {"unchecked:callee_partly_interpreted", "agree_with_stubbed_callees"}, notes
    assert verdicts[(0, "g_res")] != "eval_error"


@needs_clang
def test_a_call_kept_as_effects_for_sequencing_is_recorded(tmp_path):
    # review R3 W3-1 (s23): ``g_acc += F()`` is read as effects on purpose (no failure recorded) while F runs elsewhere
    proj = _project({"c.h": H % "extern U8 g_div;\nextern U8 g_in;\nextern U8 g_acc;\nextern U8 g_res;\nU8 F(void);",
                     "a.c": '#include "c.h"\nU8 g_div;\nU8 g_in;\nU8 g_acc;\nU8 g_res;\n'
                            'void e(void) { g_div = 5U; (void)F(); g_div = g_in; g_acc += F(); g_res = 7U; }\n',
                     "b.c": '#include "c.h"\nU8 F(void) { g_acc = 1U; return (U8)(100U / g_div); }\n'})
    files, scopes, provider, _parser = proj
    path = os.path.join(ROOT, "a.c")
    unit = {"name": "e", "source_text": files[path], "source_path": path, "source_text_complete": True,
            "project_scope": scopes[path], "callee_provider": provider}
    rec = evaluate_outputs(unit, [{}], [["g_res"]])[0]
    ip = rec["interprocedural"]
    assert "F" in ip["inlined"] and "F" in ip["effects_only"], ip
    derived, verdicts, notes = _run(proj, "e", [({}, ["g_res"], {})], tmp_path)
    assert derived[0] == {"g_res": 7} and verdicts == {(0, "g_res"): "unchecked:callee_partly_interpreted"}, notes


def test_the_sheet_carries_the_callees_read_as_effects(tmp_path):
    import openpyxl

    from generators.integration_oracle import write_integration_evidence_sheet
    itcs = [{"tc_id": "T1", "entry_fn": "e", "sub_cases": [{"case_num": 1, "inputs": {}, "expected": {"g": 1},
                                                            "expected_evidence": {"g": {
                                                                "status": "derived", "oracle_kind": "source",
                                                                "callees_interpreted": ["F"],
                                                                "callees_effects_only": ["F", "G"]}}}]}]
    wb = openpyxl.Workbook()
    assert write_integration_evidence_sheet(wb, itcs) == 1
    rows = list(wb["Test Evidence"].iter_rows(values_only=True))
    row = dict(zip(rows[0], rows[1], strict=True))
    assert row["Callees interpreted"] == "F" and row["Callees effects-only"] == "F, G"


_BOOL_H = H % "typedef unsigned char bool;\nextern U8 g_x;\nextern U8 g_out;\nextern U8 g_out3;\nU8 b_id(const bool v);"


@needs_clang
def test_a_qualified_project_bool_is_the_projects_type_too(tmp_path):
    # review R3 W3-2 (s20): ``const bool k = g_x;`` and a ``const bool`` parameter went through ``base_kind`` → _Bool,
    # so the model derived 1 where C gives 2 (the harness caught it: mismatch 2 == 1)
    proj = _project({"c.h": _BOOL_H,
                     "a.c": '#include "c.h"\nU8 g_x;\nU8 g_out;\nU8 g_out3;\n'
                            "void e(void) { const bool k = g_x; g_out = (U8)k; g_out3 = b_id(g_x); }\n",
                     "b.c": '#include "c.h"\nU8 b_id(const bool v) { return (U8)v; }\n'})
    derived, verdicts, notes = _run(proj, "e", [({"g_x": 2}, ["g_out", "g_out3"], {})], tmp_path)
    assert derived[0] == {"g_out": 2, "g_out3": 2}
    assert verdicts == {(0, "g_out"): "agree", (0, "g_out3"): "agree"}, notes
    from generators.mcdc_design import _scope_type
    scope = proj[1][os.path.join(ROOT, "a.c")]
    assert _scope_type(scope, "const bool")["kind"] == "char"
    with pytest.raises(cpc.Unresolved):   # unchanged: a qualified ordinary typedef stays unresolved (conservative)
        _scope_type(scope, "const U8")


@needs_clang
def test_the_unit_check_renames_a_bool_typedef_as_well(tmp_path):
    # review R3 W3-3 (s24): the unit-level harness wrote ``typedef unsigned char bool;`` → every KJPDS02 function that
    # reaches ``bool`` failed to compile (unchecked)
    from scripts.source_oracle_clang_check import check_claims
    proj = _project({"c.h": H % "typedef unsigned char bool;",
                     "a.c": '#include "c.h"\nU8 f(bool b) { U8 r = 0U; if (b != 2U) { r = 3U; } return r; }\n'})
    files, scopes, _provider, _parser = proj
    path = os.path.join(ROOT, "a.c")
    unit = {"name": "f", "source_text": files[path], "source_path": path, "source_text_complete": True,
            "project_scope": scopes[path]}
    rec = evaluate_outputs(unit, [{"b": 2}], [["return"]])[0]
    assert rec["outputs"]["return"].get("value") == 0   # an unsigned char 2, not _Bool 1
    report = check_claims([{"unit": unit, "inputs": {"b": 2}, "outputs": {"return": 0}, "possible_ub": []},
                           {"unit": unit, "inputs": {"b": 2}, "outputs": {"return": 3}, "possible_ub": []}],
                          work_dir=str(tmp_path / "unit"))
    assert (report["agree"], report["mismatch"], report["unchecked"]) == (1, 1, 0), report


# ── review round 4 ──────────────────────────────────────────────────────────────────────────────────────────────

@needs_clang
def test_an_effects_only_frame_deep_in_the_call_stack_is_seen(tmp_path):
    # review R4 W4-1 (s26): e → c1 … c11, c5 read as effects at one site (``g_acc += c5()``), a division by zero in
    # c11 — clang's default backtrace (10 frames) skipped c5 and the claim became a false contradiction
    chain = "".join(f"U8 c{i}(void) {{ return c{i + 1}(); }}\n" for i in range(1, 11) if i not in (4,))
    chain = chain.replace("U8 c3(void) { return c4(); }\n",
                          "U8 c3(void) { return c4(); }\nU8 c4(void) { (void)c5(); g_acc += c5(); return 1U; }\n")
    decls = "\n".join(f"U8 c{i}(void);" for i in range(1, 12))
    proj = _project({"c.h": H % ("extern U8 g_div;\nextern U8 g_acc;\nextern U8 g_res;\n" + decls),
                     "a.c": '#include "c.h"\nU8 g_div;\nU8 g_acc;\nU8 g_res;\n'
                            "void e(void) { (void)c1(); g_res = 7U; }\n" + chain +
                            "U8 c11(void) { g_acc = 1U; return (U8)(100U / g_div); }\n"})
    files, scopes, provider, parser = proj
    path = os.path.join(ROOT, "a.c")
    claims = [{"inputs": {}, "outputs": {"g_res": 7}, "possible_ub": [], "interpreted": {f"c{i}" for i in range(1, 12)},
               "effects_only": {"c5"}}]
    final, notes = ioc.check_group("e", path, claims, files, scopes, provider, parser, work_dir=str(tmp_path / "deep"))
    assert final[(0, "g_res")][0] == "unchecked:callee_partly_interpreted", (final, notes)


@needs_clang
def test_the_effects_only_column_reaches_the_verdict_through_the_sheet(tmp_path):
    # review R4 W4-3: the reader is the only production path for R3 W3-1 — a column dropped there would bring the false
    # contradictions back silently
    root = tmp_path / "src"
    root.mkdir()
    (root / "c.h").write_text(H % "extern U8 g_x;\nextern U8 g_in;\nextern U8 g_w;\nextern U8 g_res;\nvoid b_big(void);",
                              encoding="utf-8", newline="\n")
    (root / "a.c").write_text('#include "c.h"\nU8 g_x;\nU8 g_in;\nU8 g_w;\nU8 g_res;\n'
                              "void e(void) { g_x = 128U; b_big(); g_x = g_in; b_big(); g_res = 7U; }\n",
                              encoding="utf-8", newline="\n")
    (root / "b.c").write_text('#include "c.h"\nvoid b_big(void)\n' + _HEAVY, encoding="utf-8", newline="\n")
    row = {"Test Case ID": "T1", "Entry function": "e", "Observable": "g_res", "Expected": 7,
           "Source path": str((root / "a.c").resolve()), "Inputs JSON": "{}", "Callees interpreted": "b_big"}
    out = tmp_path / "r.json"
    xlsm = _sheet(tmp_path, [{**row, "Callees effects-only": "b_big"}])
    grouped, _unreadable = ioc._read_claims(xlsm)
    assert next(iter(grouped.values()))["effects_only"] == {"b_big"}
    assert ioc.main(["--xlsm", xlsm, "--source-root", str(root), "--out", str(out)]) == 2   # nothing agreed
    report = json.loads(out.read_text(encoding="utf-8"))
    assert report["unchecked_reasons"] == {"callee_partly_interpreted": 1}, report
    # the model recorded no effects-only callee: the division by zero is a contradiction
    (tmp_path / "none").mkdir()
    xlsm = _sheet(tmp_path / "none", [{**row, "Callees effects-only": ""}])
    assert ioc.main(["--xlsm", xlsm, "--source-root", str(root), "--out", str(out)]) == 1
    # a sheet from before the column: not recorded (None), read conservatively
    import openpyxl
    wb = openpyxl.load_workbook(xlsm)
    ws = wb["Test Evidence"]
    col = [c.value for c in ws[1]].index("Callees effects-only") + 1
    ws.delete_cols(col)
    old = tmp_path / "old.xlsx"
    wb.save(old)
    grouped, _unreadable = ioc._read_claims(str(old))
    assert next(iter(grouped.values()))["effects_only"] is None

