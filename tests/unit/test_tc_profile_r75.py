"""R75 — 시험 물량 프로파일(정본 규모 / 확장) + STS 라이터의 O(n²) 병합.

사용자 결정(2026-09-19): 정본과 비슷한 기본은 그대로 두고, 품질 좋은 TC 를 더 많이 내는 **옵션**을 둔다.
두 프로파일은 **포함 관계**다 — 확장 문서는 기본 문서의 시험을 그대로(같은 ID·순서·값) 담고 그 뒤에 덧붙인다.

실측(KJPDS02): STS TC 295 → 2,210(시험 가진 함수 94 → 1,037/1,037) · SUTS 시퀀스 8,741 → 17,565 · SITS TC 120 → 360.
상한만 올리는 길은 막혀 있었다 — 요구 하나에 함수가 600개씩 매핑돼 (요구, 함수) 쌍이 9,434 이고, 라이터의
`ws.merge_cells` 가 O(n²) 라 요구당 상한 1000 의 라이브 생성은 27분이 지나도 끝나지 않았다.
"""
from __future__ import annotations

import pytest

from generators import sits as sits_mod
from generators._xlsx_merge import merge_fresh
from generators.sts import _classify_steps, _generate_steps_from_flow, generate_sts_xlsm, generate_test_cases
from generators.suts import generate_sequences, is_extended_strategy
from generators.tc_profile import TC_PROFILE_EXTENDED, TC_PROFILE_REFERENCE, is_extended, normalize_tc_profile


class TestProfileValue:
    @pytest.mark.parametrize("raw, expected", [
        ("", (TC_PROFILE_REFERENCE, "")), (None, (TC_PROFILE_REFERENCE, "")), ("reference", (TC_PROFILE_REFERENCE, "")),
        ("extended", (TC_PROFILE_EXTENDED, "")), (" EXTENDED ", (TC_PROFILE_EXTENDED, "")),
        ("extnded", (TC_PROFILE_REFERENCE, "extnded")), ("all", (TC_PROFILE_REFERENCE, "all")), (1, (TC_PROFILE_REFERENCE, "1")),
    ])
    def test_unknown_values_fall_back_to_reference_and_are_reported(self, raw, expected):
        """오타 하나가 문서를 열 배로 키우면 안 된다 — 모르는 값은 기본으로 가고, 원래 값은 돌려줘서 공시한다."""
        assert normalize_tc_profile(raw) == expected
        assert is_extended(raw) is (expected[0] == TC_PROFILE_EXTENDED)


# ── 병합 ─────────────────────────────────────────────────────────────────────

class TestMergeFresh:
    def test_same_result_as_merge_cells(self):
        from openpyxl import Workbook
        from openpyxl.cell.cell import MergedCell

        a, b = Workbook().active, Workbook().active
        for ws in (a, b):
            for r in range(1, 7):
                for c in range(1, 4):
                    ws.cell(row=r, column=c, value=f"{r}.{c}")
        for r0, r1 in ((1, 3), (4, 6)):
            a.merge_cells(start_row=r0, start_column=2, end_row=r1, end_column=2)
            merge_fresh(b, r0, 2, r1, 2)
        assert sorted(str(r) for r in a.merged_cells.ranges) == sorted(str(r) for r in b.merged_cells.ranges) == ["B1:B3", "B4:B6"]
        for r in range(1, 7):
            assert isinstance(a.cell(row=r, column=2), MergedCell) == isinstance(b.cell(row=r, column=2), MergedCell)
        assert b.cell(row=1, column=2).value == "1.2" and b.cell(row=1, column=1).value == "1.1"

    def test_reversed_range_is_an_error_not_a_silent_skip(self):
        from openpyxl import Workbook

        with pytest.raises(ValueError):
            merge_fresh(Workbook().active, 5, 2, 3, 2)

    def test_written_file_opens_with_the_merges(self, tmp_path):
        from openpyxl import Workbook, load_workbook

        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="x")
        merge_fresh(ws, 1, 1, 4, 1)
        wb.save(tmp_path / "m.xlsx")
        assert [str(r) for r in load_workbook(tmp_path / "m.xlsx").active.merged_cells.ranges] == ["A1:A4"]

    def test_sts_writer_does_not_scan_existing_merges_per_tc(self, tmp_path, monkeypatch):
        """TC 블록 병합이 `MultiCellRange.add`(기존 병합 전부를 훑는다)를 거치면 TC 수에 비례해 호출이 는다 — 문서 전체로 O(n²)."""
        from openpyxl.worksheet.cell_range import MultiCellRange

        calls = {"n": 0}
        real = MultiCellRange.add

        def _counting(self, coord):
            calls["n"] += 1
            return real(self, coord)

        monkeypatch.setattr(MultiCellRange, "add", _counting)

        def _tcs(n):
            return [{"id": f"SwTC_{i:03d}", "title": "t", "safety_related": "X", "test_environment": "SwTE_01", "test_method": "RBT",
                     "gen_method": "ECA", "fs_req": "", "description": "d", "precondition": "p", "srs_id": "SwTR_0001",
                     "steps": [{"action": "a1", "expected": "e1"}, {"action": "a2", "expected": "e2"}]} for i in range(n)]

        per_n = {}
        for n in (1, 6):
            calls["n"] = 0
            generate_sts_xlsm(None, _tcs(n), {"matrix": [], "coverage": {}}, str(tmp_path / f"s{n}.xlsm"), {})
            per_n[n] = calls["n"]
        assert per_n[1] == per_n[6], per_n

        from openpyxl import load_workbook

        ws = load_workbook(tmp_path / "s6.xlsm", keep_vba=True)["3.SW Test Spec"]
        tc_merges = [r for r in ws.merged_cells.ranges if r.min_row >= 7 and r.max_row == r.min_row + 1]
        assert len(tc_merges) >= 6 and len(tc_merges) % 6 == 0, "TC 6개 × 병합 열 수"


# ── STS ──────────────────────────────────────────────────────────────────────

_IF = [{"type": "if", "condition": "n > 10", "true_body": [{"type": "call", "name": "A"}],
        "false_body": [{"type": "call", "name": "B"}]}]


def _fn(name, inputs=("[IN] U8 n",), flow=None):
    return {"id": name, "name": name, "prototype": f"void {name}(U8 n)", "inputs": list(inputs), "outputs": [],
            "calls_list": [], "logic_flow": _IF if flow is None else flow}


def _sts(profile, max_tc=5, reqs=None, mapping=None, n_fns=6):
    fd = {f"F{i}": _fn(f"S_Fn_{i}") for i in range(1, n_fns + 1)}
    reqs = reqs or [{"id": "SwTR_0001", "req_type": "TR"}, {"id": "SwTR_0002", "req_type": "TR"}]
    mapping = mapping or {"SwTR_0001": list(fd), "SwTR_0002": ["F5", "F6"]}
    cfg = {"max_tc_per_req": max_tc}
    if profile is not None:
        cfg["tc_profile"] = profile
    stats: dict = {}
    import copy

    tcs = generate_test_cases(copy.deepcopy(reqs), fd, mapping, cfg, stats_out=stats)
    return tcs, stats


def _key(tc):
    return (tc["id"], tc["srs_id"], tc["title"], tc["gen_method"], [(s["action"], s["expected"]) for s in tc["steps"]])


class TestStsProfile:
    def test_reference_is_the_default_and_unchanged(self):
        base, st0 = _sts(None)
        for profile in ("", "reference", "extnded"):
            tcs, st = _sts(profile)
            assert [_key(t) for t in tcs] == [_key(t) for t in base], profile
            assert st["extended_tcs"] == 0 and st["extended_functions"] == 0 and st["tc_profile"] == TC_PROFILE_REFERENCE
        assert _sts("extnded")[1]["tc_profile_unknown_value"] == "extnded" and st0["tc_profile_unknown_value"] == ""
        assert not any("tc_profile" in t for t in base)

    def test_extended_contains_the_reference_document_unchanged(self):
        base, _ = _sts(None)
        ext, st = _sts("extended")
        assert [_key(t) for t in ext if t.get("tc_profile") != TC_PROFILE_EXTENDED] == [_key(t) for t in base]
        assert len(ext) == len(base) + st["extended_tcs"] and st["extended_tcs"] > 0
        assert len({t["id"] for t in ext}) == len(ext), "확장 TC 의 ID 는 기본 TC 뒤를 잇는다 — 겹치면 안 된다"

    def test_every_mapped_function_gets_a_tc_exactly_once_in_the_extension(self):
        base, st0 = _sts(None)
        ext, st = _sts("extended")
        assert st0["functions_with_tc"] < st0["mapped_functions"], "전제: 기본은 상한에 걸려 일부 함수만 시험한다"
        assert st["functions_with_tc"] == st["mapped_functions"] == 6 and st["functions_without_tc"] == 0
        added = [t for t in ext if t.get("tc_profile") == TC_PROFILE_EXTENDED]
        fns = {t["title"].rsplit(" - ", 1)[-1] for t in added}
        base_fns = {t["title"].rsplit(" - ", 1)[-1] for t in base}
        assert not (fns & base_fns), "이미 시험이 있는 함수는 되풀이하지 않는다"
        assert fns | base_fns == {f"S_Fn_{i}" for i in range(1, 7)}
        for fn in fns:
            assert len({t["srs_id"] for t in added if t["title"].endswith(f" - {fn}")}) == 1, "함수마다 한 요구 밑에만"

    def test_function_goes_under_its_narrowest_requirement_and_ties_follow_document_order(self):
        # 기본 구간이 F5·F6 을 쓰지 못하게 상한을 1 로 — F5 는 두 요구에 다 매핑되고 좁은 쪽은 SwTR_0002(2개) 다.
        ext, _ = _sts("extended", max_tc=1, mapping={"SwTR_0001": ["F1", "F2", "F5"], "SwTR_0002": ["F6", "F5"]})
        home = {t["title"].rsplit(" - ", 1)[-1]: t["srs_id"] for t in ext if t.get("tc_profile") == TC_PROFILE_EXTENDED}
        assert home["S_Fn_5"] == "SwTR_0002" and home["S_Fn_2"] == "SwTR_0001"
        tie, _ = _sts("extended", max_tc=1, mapping={"SwTR_0001": ["F1", "F5"], "SwTR_0002": ["F6", "F5"]})
        home = {t["title"].rsplit(" - ", 1)[-1]: t["srs_id"] for t in tie if t.get("tc_profile") == TC_PROFILE_EXTENDED}
        assert home["S_Fn_5"] == "SwTR_0001", "동률이면 문서에서 먼저 나온 요구"

    def test_document_stays_grouped_by_requirement_in_document_order(self):
        ext, _ = _sts("extended", max_tc=1)
        order = [t["srs_id"] for t in ext]
        assert order == sorted(order, key=["SwTR_0001", "SwTR_0002"].index)
        ids_0001 = [t["id"] for t in ext if t["srs_id"] == "SwTR_0001"]
        assert ids_0001 == sorted(ids_0001), "요구 안에서 번호는 이어진다"

    def test_boundary_tc_survives_the_function_cap_only_in_the_extension(self):
        assert not any(_classify_steps(t)[1] == "BAA" for t in _generate_steps_from_flow(_IF, _fn("X"), max_steps=15, max_tc=2))
        kept = _generate_steps_from_flow(_IF, _fn("X"), max_steps=15, max_tc=2, keep_boundary=True)
        assert len(kept) == 3 and _classify_steps(kept[-1])[1] == "BAA" and _classify_steps(kept[0])[1] != "BAA"
        ext, st = _sts("extended", max_tc=2)
        assert st["extended_boundary_tcs"] == st["extended_functions"] > 0
        assert sum(1 for t in ext if t.get("tc_profile") == TC_PROFILE_EXTENDED and t["gen_method"] == "BAA") == st["extended_boundary_tcs"]

    def test_unknown_typed_inputs_still_get_no_invented_boundary(self):
        """확장은 상한을 풀 뿐 값을 지어내지 않는다 — 모르는 타입뿐인 함수는 확장에서도 경계값 TC 가 없다."""
        fd = {"F1": _fn("A"), "F2": _fn("B", inputs=("[IN] void *p",))}
        stats: dict = {}
        tcs = generate_test_cases([{"id": "SwTR_0001", "req_type": "TR"}], fd, {"SwTR_0001": ["F1", "F2"]},
                                  {"max_tc_per_req": 1, "tc_profile": "extended"}, stats_out=stats)
        b_tcs = [t for t in tcs if t["title"].endswith(" - B")]
        assert b_tcs and not any(t["gen_method"] == "BAA" for t in b_tcs) and stats["extended_boundary_tcs"] == 0

    def test_safety_requirement_wins_over_a_narrower_non_safety_one(self):
        """(리뷰 C2) 폭만 보면 안전 요구(넓음)·비안전 요구(좁음)에 함께 매핑된 함수가 비안전 요구 밑으로 가 `X` 로 찍힌다 —
        "시험이 없다" 를 "비안전이다" 로 바꾸는 under-classification."""
        reqs = [{"id": "SwTSR_0001", "req_type": "TSR", "asil": "B"}, {"id": "SwTR_0002", "req_type": "TR", "asil": "QM"}]
        ext, st = _sts("extended", max_tc=1, reqs=reqs, mapping={"SwTSR_0001": ["F1", "F2", "F3", "F5"], "SwTR_0002": ["F6", "F5"]})
        f5 = [t for t in ext if t["title"].endswith(" - S_Fn_5")]
        assert f5 and {t["srs_id"] for t in f5} == {"SwTSR_0001"} and {t["safety_related"] for t in f5} == {"O"}
        assert st["extended_functions_placed_by_safety"] == 1, "좁은 요구가 아닌 곳에 놓은 수를 센다"
        both_qm = [{**r, "asil": "QM"} for r in reqs]
        ext2, st2 = _sts("extended", max_tc=1, reqs=both_qm, mapping={"SwTSR_0001": ["F1", "F2", "F3", "F5"], "SwTR_0002": ["F6", "F5"]})
        assert {t["srs_id"] for t in ext2 if t["title"].endswith(" - S_Fn_5")} == {"SwTR_0002"} and st2["extended_functions_placed_by_safety"] == 0

    def test_duplicate_requirement_id_does_not_reuse_a_number(self):
        """(리뷰 I1) 같은 요구 ID 가 두 번 들어오면 기본 구간이 번호를 두 번 쓴다 — 확장은 그 **끝** 뒤를 이어야 한다."""
        reqs = [{"id": "SwTR_0001", "req_type": "TR"}, {"id": "SwTR_0001", "req_type": "TR"}]
        ext, _ = _sts("extended", max_tc=3, reqs=reqs, mapping={"SwTR_0001": ["F1", "F2", "F3"]})
        added = [t["id"] for t in ext if t.get("tc_profile") == TC_PROFILE_EXTENDED]
        base_ids = {t["id"] for t in ext if t.get("tc_profile") != TC_PROFILE_EXTENDED}
        assert added and not (set(added) & base_ids) and len(set(added)) == len(added)

    def test_duplicated_fid_does_not_widen_a_requirement(self):
        """(리뷰 I2) 폭은 서로 다른 함수 수다 — 상류의 fid 중복이 집 요구를 바꾸면 안 된다."""
        ext, _ = _sts("extended", max_tc=1, mapping={"SwTR_0001": ["F1", "F5", "F2"], "SwTR_0002": ["F6", "F5", "F5", "F5"]})
        assert {t["srs_id"] for t in ext if t["title"].endswith(" - S_Fn_5")} == {"SwTR_0002"}

    def test_what_the_extension_could_not_do_is_counted(self):
        """(리뷰 W4·W5) 확장의 침묵 축 — 상세 없는 함수 · 경계값을 못 붙인 함수(모르는 타입). "값을 지어내지 않는다" 의 수치."""
        fd = {"F1": _fn("A"), "F2": _fn("B", inputs=("[IN] void *p",)), "F3": {}}
        stats: dict = {}
        generate_test_cases([{"id": "SwTR_0001", "req_type": "TR"}], fd, {"SwTR_0001": ["F1", "F2", "F3"]},
                            {"max_tc_per_req": 1, "tc_profile": "extended"}, stats_out=stats)
        assert stats["extended_functions_without_detail"] == 1 and stats["functions_without_tc"] == 1
        assert stats["extended_boundary_tc_unavailable"] == 1 and stats["extended_boundary_tc_candidates"] == 0

    def test_leftover_warning_does_not_blame_the_requirement_cap_under_the_extension(self):
        from generators.sts import generate_quality_report, generate_traceability_matrix

        fd = {"F1": _fn("A"), "F3": {}}
        reqs = [{"id": "SwTR_0001", "req_type": "TR"}]
        for profile, blames_cap in (("", True), ("extended", False)):
            stats: dict = {}
            tcs = generate_test_cases([dict(r) for r in reqs], fd, {"SwTR_0001": ["F1", "F3"]},
                                      {"max_tc_per_req": 1, "tc_profile": profile}, stats_out=stats)
            q = generate_quality_report(tcs, generate_traceability_matrix(tcs, reqs), generation_stats=stats)
            text = " ".join(q["coverage_warnings"])
            assert ("요구당 TC 상한(max_tc_per_req=" in text) is blames_cap, (profile, text)
            if not blames_cap:
                assert "확장 프로파일인데도" in text and "함수 상세 없음 1개" in text

    def test_stats_are_json_serializable(self):
        import json

        json.dumps(_sts("extended")[1])


# ── SUTS ─────────────────────────────────────────────────────────────────────

def _unit(**kw):
    base = {"fid": "F", "name": "Fn", "prototype": "void Fn(void)", "input_vars": ["a", "b", "c", "d", "e"], "output_vars": ["o"],
            "indirect_vars": [], "logic_flow": [], "calls_list": [], "value_domains": {}, "uds_param_info": {},
            "param_types": {"a": "U8", "b": "U8", "c": "U16", "d": "U8", "e": "U8", "o": "U8"}}
    base.update(kw)
    return base


def _seq_key(s):
    return (s["seq_num"], s["strategy"], s["inputs"], s["expected"], s["description"])


_SWITCH8 = [{"type": "switch", "variable": "a", "cases": [{"value": v, "label": f"C{v}"} for v in range(8)]}]


class TestSutsProfile:
    def test_default_call_is_unchanged_and_has_no_extension(self):
        seqs = generate_sequences(_unit(), 24, type_cache={})
        assert not any(is_extended_strategy(s["strategy"]) for s in seqs)
        assert [s["strategy"] for s in seqs[:6]] == ["BV_MIN_INV", "BV_MIN", "BV_MID", "BV_MAX", "BV_MAX_INV", "MIXED"]

    def test_extended_starts_with_the_reference_sequences_unchanged(self):
        for unit_kw in ({}, {"logic_flow": _SWITCH8}, {"indirect_vars": ["g1", "g2", "g3", "g4", "g5"],
                                                       "param_types": {**_unit()["param_types"], **{f"g{i}": "U8" for i in range(1, 6)}}}):
            ref = generate_sequences(_unit(**unit_kw), 24, type_cache={})
            ext = generate_sequences(_unit(**unit_kw), None, type_cache={}, extended=True)
            assert [_seq_key(s) for s in ext[:len(ref)]] == [_seq_key(s) for s in ref], unit_kw
            assert len(ext) > len(ref) and all(is_extended_strategy(s["strategy"]) for s in ext[len(ref):]), unit_kw

    def test_extension_is_marked_where_it_is_made_and_the_name_rule_agrees(self):
        """(리뷰 W2·X5) 표시는 만든 자리에서 단다. 이름 규칙(`is_extended_strategy`)은 같은 상수를 보므로 둘이 늘 같아야 한다."""
        from generators import suts as suts_mod

        kw = {"logic_flow": _SWITCH8, "indirect_vars": ["g1", "g2", "g3", "g4"],
              "param_types": {**_unit()["param_types"], **{f"g{i}": "U8" for i in range(1, 5)}}}
        ref = generate_sequences(_unit(**kw), 24, type_cache={})
        u = _unit(**kw)
        ext = generate_sequences(u, None, type_cache={}, extended=True)
        assert not any("tc_profile" in s for s in ref)
        assert [s.get("tc_profile") == TC_PROFILE_EXTENDED for s in ext] == [is_extended_strategy(s["strategy"]) for s in ext]
        assert u["base_strategy_count"] == len(ref) == sum(1 for s in ext if "tc_profile" not in s)
        assert (suts_mod._BASE_SWITCH_SLOTS, suts_mod._BASE_GLOBAL_SLOTS, suts_mod._BASE_MCDC_SLOTS) == (6, 3, 6), "정본 규모의 정의"

    def test_one_at_a_time_covers_every_input_side_not_already_made(self):
        ext = {s["strategy"]: s for s in generate_sequences(_unit(), None, type_cache={}, extended=True)}
        oat = sorted(k for k in ext if k.startswith("OAT_"))
        # 조건 조합이 만든 것: a=min, b=max, c=min, d=max. 남은 것: a=max, b=min, c=max, d=min, e=min, e=max.
        assert oat == ["OAT_0_MAX", "OAT_1_MIN", "OAT_2_MAX", "OAT_3_MIN", "OAT_4_MAX", "OAT_4_MIN"]
        assert ext["OAT_2_MAX"]["inputs"] == {"a": 127, "b": 127, "c": 65535, "d": 127, "e": 127}
        assert ext["OAT_4_MIN"]["inputs"] == {"a": 127, "b": 127, "c": 32767, "d": 127, "e": 0}
        assert "c=최댓값" in ext["OAT_2_MAX"]["description"] and "e=최솟값" in ext["OAT_4_MIN"]["description"]

    def test_single_input_gets_no_one_at_a_time(self):
        """입력이 하나면 BV_MIN/BV_MAX 가 이미 단독 경계다 — 같은 행을 한 번 더 만들지 않는다."""
        ext = generate_sequences(_unit(input_vars=["a"]), None, type_cache={}, extended=True)
        assert not any(s["strategy"].startswith("OAT_") for s in ext)

    def test_unknown_typed_input_is_never_the_varied_one_and_stays_blank(self):
        u = _unit(input_vars=["a", "pt", "b"], param_types={"a": "U8", "pt": "ST_THING *", "b": "U8", "o": "U8"})
        ext = generate_sequences(u, None, type_cache={}, extended=True)
        oat = [s for s in ext if s["strategy"].startswith("OAT_")]
        assert oat and all("pt" not in s["inputs"] for s in ext), "모르는 타입은 확장에서도 값을 비운다"
        assert all("pt=" not in s["description"].splitlines()[0] for s in oat)
        # 값을 만들 수 있는 입력은 a·b 둘 — 조건 조합이 a=min·b=max 를 이미 만들었으니 남은 것은 정확히 a=max·b=min 이다.
        assert sorted(s["strategy"] for s in oat) == ["OAT_0_MAX", "OAT_1_MIN"]
        mid = next(s["inputs"] for s in ext if s["strategy"] == "BV_MID")
        assert all(s["inputs"] != mid for s in oat), "아무것도 바꾸지 않은 행이 단독 경계 라벨로 서면 안 된다"

    def test_all_switch_cases_appear_only_in_the_extension(self):
        ref = [s["strategy"] for s in generate_sequences(_unit(logic_flow=_SWITCH8), 24, type_cache={})]
        ext = {s["strategy"]: s for s in generate_sequences(_unit(logic_flow=_SWITCH8), None, type_cache={}, extended=True)}
        assert "SWITCH_5" in ref and "SWITCH_6" not in ref
        assert ext["SWITCH_7"]["inputs"]["a"] in (7, "0x07") and "C7" in ext["SWITCH_7"]["description"]

    def test_void_side_effect_row_is_the_same_in_both_profiles(self):
        """확장에서 전역 목록이 길어져도 기본 시퀀스의 기대값 칸은 그대로다(포함 관계) — 4번째 이후 전역은 뒤에 붙는 GLOBAL_3+ 의 몫."""
        kw = {"output_vars": [], "indirect_vars": ["g1", "g2", "g3", "g4", "g5"],
              "param_types": {**_unit()["param_types"], **{f"g{i}": "U8" for i in range(1, 6)}}}
        ref = {s["strategy"]: s for s in generate_sequences(_unit(**kw), 24, type_cache={})}
        ext = {s["strategy"]: s for s in generate_sequences(_unit(**kw), None, type_cache={}, extended=True)}
        assert "VOID_SIDE_EFFECT" in ref, "전제: 출력 없는 함수 + 전역 → 부작용 시퀀스가 선다"
        assert sorted(ref["VOID_SIDE_EFFECT"]["expected"]) == ["g1", "g2", "g3"]
        assert ext["VOID_SIDE_EFFECT"]["expected"] == ref["VOID_SIDE_EFFECT"]["expected"]
        assert "GLOBAL_4" in ext and "GLOBAL_3" not in ref and "g5" in ext["GLOBAL_4"]["inputs"]

    def test_no_cap_means_no_truncation(self):
        kw = {"logic_flow": _SWITCH8, "indirect_vars": ["g1", "g2", "g3", "g4", "g5"],
              "param_types": {**_unit()["param_types"], **{f"g{i}": "U8" for i in range(1, 6)}}}
        assert len(generate_sequences(_unit(**kw), 3, type_cache={}, extended=True)) == 3
        # BV 6 + 조건 조합 4 + switch 6 + 전역 3 = 19(기본) · 확장: switch 2 + 전역 2 + 단독 경계 6 = 10 → 29 (> 기본 상한 24)
        full = generate_sequences(_unit(**kw), None, type_cache={}, extended=True)
        assert len(full) == 29 and [s["seq_num"] for s in full] == list(range(1, 30))

    def test_condition_combination_label_says_which_side(self):
        """값은 짝수 번째=최솟값·홀수 번째=최댓값인데 라벨은 늘 "최솟값" 이었다 — 표의 값과 설명이 어긋났다."""
        by = {s["strategy"]: s for s in generate_sequences(_unit(), 24, type_cache={})}
        assert by["COND_COMB_1"]["inputs"]["b"] in (255, "0xFF") and "b=최댓값" in by["COND_COMB_1"]["description"]
        assert by["COND_COMB_0"]["inputs"]["a"] in (0, "0x00") and "a=최솟값" in by["COND_COMB_0"]["description"]

    @pytest.mark.parametrize("name, expected", [
        ("OAT_0_MIN", True), ("SWITCH_6", True), ("SWITCH_5", False), ("GLOBAL_3", True), ("GLOBAL_2", False),
        ("MCDC_6", True), ("MCDC_5", False), ("MCDC_BASE", False), ("BV_MAX", False), ("COND_COMB_3", False), ("", False), (None, False),
    ])
    def test_extended_strategy_names(self, name, expected):
        assert is_extended_strategy(name) is expected


# ── SITS ─────────────────────────────────────────────────────────────────────

class TestSitsProfile:
    def test_reference_keeps_the_requested_caps(self):
        assert sits_mod.resolve_profile_caps("", 7, 120) == (TC_PROFILE_REFERENCE, "", 7, 120)
        assert sits_mod.resolve_profile_caps("extnded", 7, 120) == (TC_PROFILE_REFERENCE, "extnded", 7, 120)

    def test_extended_lifts_both_caps_and_never_lowers_a_larger_request(self):
        cat = sits_mod._SUBCASE_CATALOG_MAX
        assert sits_mod.resolve_profile_caps("extended", 7, 120) == (TC_PROFILE_EXTENDED, "", cat, None)
        assert sits_mod.resolve_profile_caps("extended", cat + 5, 120)[2] == cat + 5

    def test_no_flow_cap_keeps_every_flow_and_a_cap_really_cuts(self):
        flows = [{"fn_name": f"f{i}", "asil": ""} for i in range(7)]
        assert len(sits_mod._select_flows_within_cap(list(flows), 3, {})) == 3, "전제: 상한은 실제로 자른다"
        assert len(sits_mod._select_flows_within_cap(list(flows), None, {})) == 7

    def test_fallback_tc_id_does_not_shift_with_the_selection(self):
        """(리뷰 C3) 설계 ID 를 못 찾은 흐름의 ID 는 **후보 전체 기준** 순번이다 — 선별된 목록 안의 위치로 매기면 상한(프로파일)에
        따라 같은 흐름의 ID 가 밀려 포함 관계가 깨진다."""
        def _flow(name, cand):
            return {"entry_fn": name, "call_chain": f"{name} -> Drv_Set", "module_name": "Ap", "candidate_index": cand,
                    "input_vars": ["u8g_x"], "input_raws": ["[IN] U8 u8g_x"], "expected_vars": ["u8g_out"],
                    "expected_raws": ["[OUT] U8 u8g_out"], "indirect_vars": [], "asil": "A", "logic_flow": [], "related_ids": ["SwCom_01"]}

        all_flows = [_flow(f"f{i}", i) for i in range(1, 6)]
        picked = [all_flows[1], all_flows[3]]   # 상한이 2번·4번 후보만 남긴 기본 문서
        ref_ids = [t["tc_id"] for t in sits_mod.generate_itc_list(picked, max_subcases=1)]
        ext_ids = [t["tc_id"] for t in sits_mod.generate_itc_list(all_flows, max_subcases=1)]
        assert ref_ids == ["SwITC_02", "SwITC_04"], "선별 목록 안의 위치(01·02)가 아니라 후보 순번이다"
        assert set(ref_ids) <= set(ext_ids) and len(set(ext_ids)) == 5
        no_index = [{k: v for k, v in f.items() if k != "candidate_index"} for f in picked]
        assert [t["tc_id"] for t in sits_mod.generate_itc_list(no_index, max_subcases=1)] == ["SwITC_01", "SwITC_02"], "순번이 없으면 목록 위치"


# ── 준비 게이트 ──────────────────────────────────────────────────────────────

class TestGateRow:
    """게이트는 타이핑이 아니라 **만들어질 문서**를 잰다 — 확장을 골랐는데 상한 행이 기본값 기준으로 손실을 말하면 반대말이다."""

    @staticmethod
    def _steps(doc_type, tmp_path, caps=None):
        from fastapi.testclient import TestClient

        from backend.main import app

        body = {"doc_type": doc_type, "source_root": str(tmp_path)}
        if caps is not None:
            body["caps"] = caps
        r = TestClient(app).post("/api/docgen/preflight", json=body, headers={"X-User": "tester"})
        assert r.status_code == 200, r.text
        return {s["id"]: s for s in r.json()["steps"]}

    @pytest.mark.parametrize("doc_type", ["sts", "suts", "sits"])
    def test_row_shows_the_default_and_the_pick(self, doc_type, tmp_path):
        row = self._steps(doc_type, tmp_path)["tc_profile"]
        assert row["measured"]["value"] == TC_PROFILE_REFERENCE and row["measured"]["choice"] == "tc_profile"
        assert [o["value"] for o in row["measured"]["options"]] == ["", TC_PROFILE_EXTENDED]
        assert "정본 규모" in row["reason"] and "현재 **확장**" not in row["reason"]
        picked = self._steps(doc_type, tmp_path, {"tc_profile": "extended"})["tc_profile"]
        assert picked["measured"]["value"] == TC_PROFILE_EXTENDED and picked["measured"]["picked"] == "extended"
        assert "현재 **확장**" in picked["reason"] and picked["state"] == "ok"

    def test_unknown_stored_value_is_shown_not_silently_dropped(self, tmp_path):
        row = self._steps("sts", tmp_path, {"tc_profile": "extnded"})["tc_profile"]
        assert row["measured"]["value"] == TC_PROFILE_REFERENCE and row["measured"]["stored"] == "extnded"
        assert row["state"] == "degraded" and "extnded" in row["reason"]

    def test_uds_has_no_row(self, tmp_path):
        assert "tc_profile" not in self._steps("uds", tmp_path, {"tc_profile": "extended"})

    @pytest.mark.parametrize("doc_type, caps", [("sts", ["max_tc_per_req"]), ("suts", ["max_sequences"]),
                                                ("sits", ["max_subcases", "max_flows"])])
    def test_lifted_caps_stop_reporting_a_loss_under_the_extension(self, doc_type, caps, tmp_path):
        ext = self._steps(doc_type, tmp_path, {"tc_profile": "extended"})
        ref = self._steps(doc_type, tmp_path)
        for cap in caps:
            row = ext[f"cap_{cap}"]
            assert row["state"] == "ok" and row["measured"]["lifted_by_profile"] == TC_PROFILE_EXTENDED, row
            assert "확장 프로파일" in row["reason"] and "suggested" not in row["measured"]
            assert "lifted_by_profile" not in ref[f"cap_{cap}"]["measured"] and "확장 프로파일" not in ref[f"cap_{cap}"]["reason"]

    def test_measured_rows_do_not_contradict_the_extension(self, tmp_path, monkeypatch):
        """(리뷰 C1) 결정 행만 고치면 같은 패널의 **측정 행**이 동시에 반대말을 한다 — "N개가 시험되지 않습니다"/"빠집니다".

        ⚠ 측정 캐시가 **차 있어야** 이 행들이 생긴다. 빈 디렉터리로만 재면 손실을 말하는 행이 애초에 시야 밖이다.
        """
        from tests.unit.test_docgen_preflight import _fake_cat_cache

        _fake_cat_cache(monkeypatch, {
            "ok": True, "sits": {"flows_total": 145, "cap": 120, "headroom": -25, "at_cap_boundary": True},
            "sts_mapping": {"measured": True, "requirements": 2, "mapped": 2, "causes": {}, "cause_samples": {}, "sds_reason": "",
                            "bridge": {"on": True}, "cap": 5, "mapped_functions": 9, "functions_beyond_cap": 1,
                            "requirements_over_cap": 1, "max_functions_per_req": 6, "req_fid_lists": [[0, 1, 2, 3, 4, 5], [6, 7, 8]]}})

        ref = self._steps("sits", tmp_path)["sits_flows"]
        assert ref["state"] == "degraded" and "빠집니다" in ref["reason"], "전제: 기본에서는 이 행이 손실을 말한다"
        ext = self._steps("sits", tmp_path, {"tc_profile": "extended"})["sits_flows"]
        assert ext["state"] == "ok" and "빠집니다" not in ext["reason"] and "145" in ext["reason"]
        assert ext["measured"].get("lifted_by_profile") == TC_PROFILE_EXTENDED and "headroom" not in ext["measured"]

        ref = self._steps("sts", tmp_path)["sts_tc_cap"]
        assert ref["state"] == "degraded" and "시험되지" in ref["reason"] and ref["measured"]["beyond_cap"] == 1, "전제"
        ext = self._steps("sts", tmp_path, {"tc_profile": "extended"})["sts_tc_cap"]
        assert ext["state"] == "ok" and "시험되지" not in ext["reason"] and ext["measured"]["beyond_cap"] == 0
        assert ext["measured"]["beyond_cap_in_reference"] == 1 and "확장 프로파일" in ext["reason"]

    def test_step_cap_is_judged_as_before_under_the_extension(self, tmp_path):
        ext = self._steps("sts", tmp_path, {"tc_profile": "extended"})["cap_max_steps_per_tc"]
        ref = self._steps("sts", tmp_path)["cap_max_steps_per_tc"]
        assert ext == ref


# ── 배선 ─────────────────────────────────────────────────────────────────────

class TestWiring:
    @pytest.mark.parametrize("doc_type, module, func", [
        ("sts", "backend.routers.jenkins", "jenkins_sts_generate_async"),
        ("suts", "backend.routers.jenkins", "jenkins_suts_generate_async"),
        ("sits", "backend.routers.local", "local_sits_generate_async"),
    ])
    def test_handler_declares_the_form_field(self, doc_type, module, func):
        """선언이 없으면 FastAPI 가 보낸 값을 조용히 버린다(R70 N88 · UDS `reference_doc_path` 와 같은 결함)."""
        import importlib
        import inspect

        mod = importlib.import_module(module)
        fn = getattr(mod, func, None)
        if fn is None:  # 이름이 다르면 경로로 찾는다
            from backend.services import docgen_requirements as req

            path = req.requirements_for(doc_type)["handler"].split(" ", 1)[1]
            fn = next(r.endpoint for r in mod.router.routes if getattr(r, "path", "") == path)
        assert "tc_profile" in inspect.signature(fn).parameters, doc_type

    @pytest.mark.parametrize("profile, want_flows, want_extended", [("", 120, False), ("extended", None, True)])
    def test_generators_act_on_the_profile(self, profile, want_flows, want_extended, tmp_path, monkeypatch):
        """핸들러가 넘겨도 생성기 진입점이 안 쓰면 아무 일도 안 일어난다 — 상한·전략 스위치가 실제 호출에 실리는지."""
        from generators import suts as suts_mod

        (tmp_path / "m.c").write_text("typedef unsigned char U8;\nU8 Fn(U8 a, U8 b)\n{\n    if (a > 1) { return b; }\n    return 0;\n}\n",
                                      encoding="utf-8")
        seen: dict = {}

        def _flows(*_a, **kw):
            seen["max_flows"] = kw.get("max_flows")
            return []

        monkeypatch.setattr(sits_mod, "collect_integration_flows", _flows)
        sits_mod.generate_sits(source_root=str(tmp_path), output_path=str(tmp_path / "o.xlsm"), max_subcases=7, max_flows=120,
                               tc_profile=profile)
        assert seen["max_flows"] == want_flows

        real = suts_mod.generate_sequences

        def _seqs(unit, max_seq=24, type_cache=None, extended=False):
            seen["suts"] = (max_seq, extended)
            return real(unit, max_seq, type_cache, extended)

        monkeypatch.setattr(suts_mod, "generate_sequences", _seqs)
        out = suts_mod.generate_suts(source_root=str(tmp_path), output_path=str(tmp_path / "u.xlsm"), scope="source", tc_profile=profile)
        assert seen["suts"] == ((None, True) if want_extended else (24, False))
        assert out["quality_report"]["tc_profile"] == (TC_PROFILE_EXTENDED if want_extended else TC_PROFILE_REFERENCE)
        assert (out["quality_report"]["extended_sequences"] > 0) is want_extended

    def test_handlers_pass_the_value_to_the_generators(self):
        from pathlib import Path

        root = Path(__file__).resolve().parents[2]
        jenkins = (root / "backend/routers/jenkins.py").read_text(encoding="utf-8")
        local = (root / "backend/routers/local.py").read_text(encoding="utf-8")
        assert '"tc_profile": tc_profile,' in jenkins and jenkins.count("tc_profile=tc_profile,") == 1
        assert local.count("tc_profile=tc_profile,") == 1

    @pytest.mark.parametrize("doc_type", ["sts", "suts", "sits"])
    def test_choice_is_declared_with_the_generator_spelling(self, doc_type):
        from backend.services import docgen_requirements as req

        ch = req.requirements_for(doc_type)["choices"]["tc_profile"]
        assert ch["param"] == "tc_profile" and ch["api"] == "" and ch["adjustable"] is True
        # 같은 뜻의 값 둘(""·"reference")을 다 내면 화면에 "정본 규모" 가 두 줄 선다(리뷰 I6) — 미설정 하나만 낸다.
        assert [o["value"] for o in ch["options"]] == ["", TC_PROFILE_EXTENDED]
        assert ch["effect"]

    def test_uds_has_no_such_choice(self):
        from backend.services import docgen_requirements as req

        assert "tc_profile" not in req.requirements_for("uds")["choices"]

    def test_gate_names_only_caps_the_generators_really_lift(self):
        """게이트가 "확장이 이 상한을 풉니다" 라고 말하는 상한 = 생성기가 실제로 푸는 상한. 어긋나면 게이트가 거짓말을 한다."""
        from backend.routers.docgen_preflight import _CAPS_LIFTED_BY_EXTENDED
        from backend.services import docgen_requirements as req

        declared = {c for d in ("sts", "suts", "sits") for c in req.requirements_for(d)["caps"]}
        assert set(_CAPS_LIFTED_BY_EXTENDED) == {"max_tc_per_req", "max_sequences", "max_subcases", "max_flows"}
        assert set(_CAPS_LIFTED_BY_EXTENDED) <= declared
        assert "max_steps_per_tc" not in _CAPS_LIFTED_BY_EXTENDED, "TC 당 스텝 상한은 확장에서도 그대로 자른다"
