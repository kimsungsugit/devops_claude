"""R52 — AI 제안: 검토 항목의 빠진 주어를 LLM 이 후보 이름 중에서 고르고 원문 인용으로 근거를 댄다.

사용자 방향(2026-10-01) "AI 활용해서(LLM) 진행해야할 것들은 … 진행해". 정직성 규약이 본체다:
* 후보 밖 이름 · 원문에 없는 인용 · 이름과 값을 함께 적지 않은 인용은 버린다(HDPDM01 실측: '인용·이름 실재' 만으로는
  5.1 V LDO 입력에 배터리 변수, 0~16 V 출력에 PWM duty 변수가 통과했다),
* 통과해도 스텝·판정에 쓰지 않는다(검토 시트 'AI 제안' 열 · 공시만),
* 캐시는 판정이 아니라 **원 답**을 남겨 오늘의 검증으로 다시 판정한다, AI 설정이 없으면 캐시만 읽는다.
"""
from __future__ import annotations

import json

from generators.sts_ai_review import candidate_names, proposal_text, propose_subjects

REQS = [{"id": "SwTSR_0101", "description": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면 센서 전원 고장",
         "verification": "Position Sensor 입력 전압을 4.85V 미만으로 설정"}]
SYSTEM = {"SyII_03": {"doc": "SyDS", "fields": {"Name": "Power Monitoring", "Range": "9~ 16V(Operation Range)"}}}
ROWS = [{"id": "HSI_23", "vars": ["u16g_ApiIn_HallSnsrLevel"], "ids": [], "nets": ["VCC_HALL_MON"], "name": "VCC_HALL_MON"}]


def _items():
    return [{"kind": "held_back", "reason": "no_subject_for_value", "source": "SRS", "srs_ids": ["SwTSR_0101"],
             "line": "Position Sensor 입력 전압을 4.85V 미만으로 설정", "line_sha256": "a", "fact": "4.85V 미만",
             "block_name": ""},
            {"kind": "held_back", "reason": "no_subject_in_value_field", "source": "SyDS SyII_03 · Range",
             "srs_ids": ["SwTSR_0101"], "line": "9~ 16V(Operation Range)", "line_sha256": "b", "fact": "9~ 16V",
             "block_name": "Power Monitoring"},
            {"kind": "held_back", "reason": "kind_range", "source": "SRS", "srs_ids": ["SwTSR_0101"], "line": "x",
             "line_sha256": "c", "fact": "1~2V", "block_name": ""},          # not a subject question: not asked
            {"kind": "read", "reason": "read_as_range", "source": "SRS", "srs_ids": [], "line": "y", "line_sha256": "d",
             "fact": "3V", "block_name": ""}]


def _llm(answers, calls=None):
    def call(cfg, messages, stage=None, meta_out=None):
        if calls is not None:
            calls.append(json.loads(messages[1]["content"]))
        return "```json\n" + json.dumps({"answers": answers}) + "\n```"
    return call


CFG = {"model": "m-test"}


def test_only_subject_questions_are_asked_and_candidates_are_real_names():
    items = _items()
    cands = candidate_names(items[0], {r["id"]: r for r in REQS}, SYSTEM, ROWS)
    assert cands == ["u16g_ApiIn_HallSnsrLevel", "VCC_HALL_MON"]
    assert "Power Monitoring" in candidate_names(items[1], {r["id"]: r for r in REQS}, SYSTEM, ROWS)
    calls = []
    stats = propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, llm=_llm([], calls))
    assert stats["eligible"] == 2 and len(calls) == 1 and len(calls[0]["items"]) == 2
    assert "ai_proposal" not in items[2] and "ai_proposal" not in items[3]


def test_a_proposal_must_quote_a_text_that_ties_the_name_to_the_value():
    """The LLM's inference (the Hall level is 'the Position Sensor input') is not in the text: rejected, shown."""
    items = _items()
    stats = propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "Position Sensor 입력 전압을 4.85V 미만으로", "reason": "r"},
        {"id": 1, "subject": "Power Monitoring", "quote": "9~ 16V(Operation Range)", "reason": "r"}]))
    assert items[0]["ai_proposal"]["status"] == "rejected" and items[0]["ai_proposal"]["reason"] == "quote_does_not_link"
    assert "검증 실패로 버림" in proposal_text(items[0]) and "이름과 이 값을 함께 적지 않음" in proposal_text(items[0])
    # a value-only field: its block's own name is what it measures — passes, said how
    assert (items[1]["ai_proposal"]["status"], items[1]["ai_proposal"]["by"]) == ("verified", "block_name")
    assert "값만 적는 칸 '9~ 16V(Operation Range)' 의 블록 이름" in proposal_text(items[1]) and "스텝·판정에 쓰지 않음" in proposal_text(items[1])
    assert (stats["verified"], stats["verified_by_quote"], stats["verified_by_block_name"], stats["rejected"]) == \
        (1, 0, 1, 1)


def test_a_quote_with_the_name_and_the_value_passes():
    items = _items()
    propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면", "reason": ""}]))
    p = items[0]["ai_proposal"]
    assert (p["status"], p["by"], p["model"], p["cached"]) == ("verified", "quote", "m-test", False)


def test_invented_names_and_quotes_are_rejected():
    items = _items()
    propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_Invented", "quote": "4.85V 미만", "reason": ""},
        {"id": 1, "subject": "Power Monitoring", "quote": "9~16V 정격 전원(지어낸 문장)", "reason": ""}]))
    assert items[0]["ai_proposal"]["reason"] == "name_not_in_candidates"
    assert items[1]["ai_proposal"]["reason"] == "quote_not_in_texts"


def test_no_ai_config_reads_the_cache_only_and_re_verifies(tmp_path):
    cache = tmp_path / "ai.json"
    items = _items()
    propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, cache_path=cache, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면"}]))
    saved = json.loads(cache.read_text(encoding="utf-8"))
    assert all(set(v) == {"model", "answer"} for v in saved.values())          # the raw answer, not the verdict
    # (review W3) the second item got no answer: a failure, never cached as "could not decide"
    assert items[1]["ai_proposal"]["status"] == "call_failed" and len(saved) == 1
    again = _items()

    def no_call(*a, **k):
        raise AssertionError("no network without ai_config")
    stats = propose_subjects(again, REQS, SYSTEM, ROWS, cache_path=cache, llm=no_call)
    assert (again[0]["ai_proposal"]["status"], again[0]["ai_proposal"]["cached"]) == ("verified", True)
    assert (stats["cached"], stats["not_asked"], stats["called"]) == (1, 1, False)
    # a document change (another sentence sha) is another key: not asked, not reused
    moved = _items()
    moved[0]["line_sha256"] = "z"
    assert propose_subjects(moved, REQS, SYSTEM, ROWS, cache_path=cache)["not_asked"] == 2    # + the uncached one


def test_a_failed_call_is_counted_not_silent():
    items = _items()

    def boom(*a, **k):
        raise RuntimeError("down")
    stats = propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, llm=boom)
    assert stats["call_failed"] == 2 and "error" not in stats and proposal_text(items[0]).startswith("[AI 호출 실패]")


# ── wiring: generate_sts, the review sheet, the disclosure ─────────────────────────────────────────────────────────


def test_the_review_sheet_carries_the_checked_proposal():
    import openpyxl

    from generators.sts_requirement_tc import REQUIREMENT_REVIEW_HEADERS, write_requirement_review_sheet
    items = _items()[:2]
    propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면"},
        {"id": 1, "subject": "Power Monitoring", "quote": "9~ 16V(Operation Range)"}]))
    for it in items:
        it.update(srs_ids=it["srs_ids"], line_span=[0, 1], fact_kind="threshold", if_filled="x", fill_check="verified")
    wb = openpyxl.Workbook()
    write_requirement_review_sheet(wb, items)
    rows = list(wb["Requirement Review"].iter_rows(values_only=True))
    assert rows[0][-1] == REQUIREMENT_REVIEW_HEADERS[-1] == "AI Proposal (Checked)"
    assert rows[1][-1].startswith("AI 제안: 주어 = u16g_ApiIn_HallSnsrLevel") and "스텝·판정에 쓰지 않음" in rows[1][-1]


def test_the_generator_reads_the_cache_without_ai_and_discloses(tmp_path):
    """Without an AI config the generation never calls out — the counts say so; a cache made earlier is used."""
    from generators.sts import generate_sts
    from report_gen.generation_disclosures import build_disclosures
    text = ["SwTSR_0101: 센서 전원 이상 시 고장을 검출한다. Verification: Pull-up Port를 4.85V 미만으로 Setting"]
    out = generate_sts(text, {}, str(tmp_path / "a.xlsm"),
                       project_config={"project_id": "T", "ai_review_cache": str(tmp_path / "ai.json")})
    rb = out["quality_report"]["generation_stats"]["requirement_boundary"]
    ai = rb.get("ai_review")
    assert isinstance(ai, dict) and ai["called"] is False and ai["verified"] == 0
    items = {i["key"]: i for i in build_disclosures("sts", out["quality_report"])}
    if ai["eligible"]:
        assert "AI 설정 없음 — 캐시만 읽음" in items["sts_ai_review"]["note"]
    else:
        assert "sts_ai_review" not in items


def test_the_disclosure_counts_apart_and_a_failure_is_a_warning():
    from report_gen.generation_disclosures import build_disclosures

    def item(ai):
        return {i["key"]: i for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
            "tcs": 1, "steps": 1, "facts": 1, "facts_used": 1, "ai_review": ai}}})}.get("sts_ai_review")
    got = item({"eligible": 45, "verified": 3, "verified_by_quote": 3, "verified_by_block_name": 0, "rejected": 14,
                "unknown": 28, "not_asked": 0, "call_failed": 0, "cached": 0, "called": True, "model": "m"})
    assert got["value"] == "인용 확인 3 · 블록 이름 0 / 45" and "모델 m" in got["note"] and got["tone"] == "info"
    assert item({"error": "RuntimeError: x"})["tone"] == "warning"
    assert item({"eligible": 0, "called": False}) is None


def test_the_kb_ingest_leaves_the_ai_column_out(tmp_path):
    """The KB holds what the generator read, never a model's answer — and the column's '묻지 않음' would change every
    key with the AI setting (R53's mass-removal guard on every switch)."""
    import openpyxl

    from generators.sts_requirement_tc import write_requirement_review_sheet
    from scripts.sts_findings_to_kb import build_documents
    items = _items()[:1]
    for it in items:
        it.update(line_span=[0, 1], fact_kind="threshold", if_filled="x", fill_check="verified")
    keys = []
    for answer in ([{"id": 0, "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면"}],
                   None):
        for it in items:
            it.pop("ai_proposal", None)
        if answer is not None:
            propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, llm=_llm(answer))
        wb = openpyxl.Workbook()
        write_requirement_review_sheet(wb, items)
        wb.save(tmp_path / "HD_STS.xlsm")
        docs = build_documents(str(tmp_path / "HD_STS.xlsm"), "HD")
        assert docs and not any("AI 제안" in d["content"] or "AI Proposal" in d["content"] for d in docs)
        keys.append([d["key"] for d in docs])
    assert keys[0] == keys[1]


def test_a_quote_that_fits_several_names_points_to_none():
    """KJPDS02_PV: 'Motor' is in nearly every name — it tied a door speed (m/s) to ``u16g_DrvIn_MOTOR_C_FB``."""
    item = {"kind": "held_back", "reason": "no_subject_for_value", "source": "SyDS SyFN_03 · Description",
            "srs_ids": [], "line": "일정 속도 이상인 경우(2.1m/s 초과) Motor Short Brake 로 정지", "line_sha256": "e",
            "fact": "2.1m/s 초과", "block_name": ""}
    rows = [{"id": "H1", "vars": ["u16g_DrvIn_MOTOR_C_FB"], "ids": [], "nets": [], "name": ""},
            {"id": "H2", "vars": ["u16g_DrvIn_MOTOR_A_FB"], "ids": [], "nets": [], "name": ""}]
    propose_subjects([item], [], None, rows, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_DrvIn_MOTOR_C_FB", "quote": "일정 속도 이상인 경우(2.1m/s 초과) Motor Short Brake"}]))
    # (review r2 W4) one word of the name ('motor') is no tie: every word of it must be in the quote
    assert item["ai_proposal"]["status"] == "rejected" and item["ai_proposal"]["reason"] == "quote_does_not_link"
    # all of the name's words, but another candidate's too: it points to neither
    speed = {**item, "line": "Motor Speed 2.1m/s 초과 시 정지", "line_sha256": "f"}
    rows2 = [{"id": "H1", "vars": ["u16g_DrvIn_MotorSpeed"], "ids": [], "nets": [], "name": ""},
             {"id": "H2", "vars": ["u16g_ApiIn_MotorSpeed"], "ids": [], "nets": [], "name": ""}]
    propose_subjects([speed], [], None, rows2, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_DrvIn_MotorSpeed", "quote": "Motor Speed 2.1m/s 초과 시 정지"}]))
    assert speed["ai_proposal"]["reason"] == "quote_names_several" and "다른 후보도 가리킴" in proposal_text(speed)


def test_a_requirement_tag_is_no_candidate():
    """KJPDS02_PV: ``[FS_REQ_MD_20]`` — a requirement's tag — was proposed as the subject of '9V 이상'."""
    item = {"kind": "held_back", "reason": "no_subject_for_value", "source": "SRS", "srs_ids": [],
            "line": "[FS_REQ_MD_20] 배터리 저전압(B+ < 8.6V)이 9V 이상으로 복귀", "line_sha256": "f", "fact": "9V 이상",
            "block_name": ""}
    cands = candidate_names(item, {}, None, None)
    assert "FS_REQ_MD_20" not in cands and not any("REQ" in c for c in cands)


# ── review round 1 ─────────────────────────────────────────────────────────────────────────────────────────────────


def _one(fact, line, reason="no_subject_for_value", block_name="", source="SRS", srs_ids=()):
    return {"kind": "held_back", "reason": reason, "source": source, "srs_ids": list(srs_ids), "line": line,
            "line_sha256": line[:8], "fact": fact, "block_name": block_name}


def test_the_block_name_and_the_hsis_names_come_before_the_cap_and_a_cut_is_counted():
    """(review C1) HDPDM01 had 54~72 names: the cap of 40 took the HSIS names and the block name off every item."""
    from generators.sts_ai_review import MAX_CANDIDATES, _candidates
    reqs = {"SwX_01": {"description": " ".join(f"u16s_MOTOR_ERR_{i:03d}" for i in range(MAX_CANDIDATES + 10))}}
    item = _one("9~ 16V", "9~ 16V", "no_subject_in_value_field", "Power Monitoring", "SyDS SyII_03 · Range", ["SwX_01"])
    names, cut = _candidates(item, reqs, SYSTEM, ROWS)
    assert names[0] == "Power Monitoring" and "u16g_ApiIn_HallSnsrLevel" in names and "VCC_HALL_MON" in names
    assert len(names) == MAX_CANDIDATES and cut == 3 + MAX_CANDIDATES + 10 - MAX_CANDIDATES   # block · HSIS 2 · 90
    stats = propose_subjects([item], [{"id": "SwX_01", **reqs["SwX_01"]}], SYSTEM, ROWS)
    assert stats["candidates_cut"] == 1


def test_every_number_of_the_value_is_written_as_a_number():
    """(review W1) '5V 미만' passed on '15V'; '0~5V' passed on any '0'."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_Vsup"], "ids": [], "nets": [], "name": ""}]
    for fact, quote in (("5V 미만", "u16g_ApiIn_Vsup 이 15V 초과"), ("0~5V", "u16g_ApiIn_Vsup 0~ 3V 로 제한")):
        item = _one(fact, quote)
        propose_subjects([item], [], None, rows, ai_config=CFG, llm=_llm([
            {"id": 0, "subject": "u16g_ApiIn_Vsup", "quote": quote}]))
        assert item["ai_proposal"]["reason"] == "quote_does_not_link", fact
    item = _one("0~5V", "u16g_ApiIn_Vsup 0~5V 범위")
    propose_subjects([item], [], None, rows, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_Vsup", "quote": "u16g_ApiIn_Vsup 0~5V 범위"}]))
    assert item["ai_proposal"]["status"] == "verified"


def test_a_quote_writing_two_candidates_points_to_neither_but_one_rows_names_are_one_signal():
    """(review W2) two variables and one value passed both; (I1) a row's net and its variable are one signal."""
    rows = [{"id": "H1", "vars": ["u16g_ApiIn_A"], "ids": [], "nets": ["NET_A"], "name": "NET_A"},
            {"id": "H2", "vars": ["u16g_ApiIn_B"], "ids": [], "nets": [], "name": ""}]
    two = _one("5V 미만", "u16g_ApiIn_A 와 u16g_ApiIn_B 가 5V 미만")
    propose_subjects([two], [], None, rows, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_A", "quote": "u16g_ApiIn_A 와 u16g_ApiIn_B 가 5V 미만"}]))
    assert two["ai_proposal"]["reason"] == "quote_names_several"
    synonym = _one("5V 미만", "NET_A(u16g_ApiIn_A) 5V 미만")
    propose_subjects([synonym], [], None, rows, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_A", "quote": "NET_A(u16g_ApiIn_A) 5V 미만"}]))
    assert synonym["ai_proposal"]["status"] == "verified"


def test_an_unread_answer_is_a_failure_never_cached_and_a_string_id_is_read(tmp_path):
    """(review W3) a truncated or refused reply made every item 'could not decide' — and cached it for good."""
    cache = tmp_path / "ai.json"
    items = _items()[:2]

    def refuse(cfg, messages, stage=None, meta_out=None):
        return "I cannot help with that."
    stats = propose_subjects(items, REQS, SYSTEM, ROWS, ai_config=CFG, cache_path=cache, llm=refuse)
    assert stats["call_failed"] == 2 and stats["unknown"] == 0 and not cache.exists()
    assert items[0]["ai_proposal"]["reason"] == "응답을 해석하지 못함"

    def truncated(cfg, messages, stage=None, meta_out=None):
        meta_out["truncated"] = True
        return json.dumps({"answers": [{"id": 0, "subject": None}]})
    assert propose_subjects(_items()[:2], REQS, SYSTEM, ROWS, ai_config=CFG, llm=truncated)["call_failed"] == 2
    again = _items()[:1]
    propose_subjects(again, REQS, SYSTEM, ROWS, ai_config=CFG, llm=_llm([
        {"id": "0", "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면"}]))
    assert again[0]["ai_proposal"]["status"] == "verified"


def test_the_cache_key_holds_the_texts_the_model_saw(tmp_path):
    """(review W4) a changed description kept the old answer; two blocks with one sentence shared one."""
    cache = tmp_path / "ai.json"
    propose_subjects(_items()[:1], REQS, SYSTEM, ROWS, ai_config=CFG, cache_path=cache, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면"}]))
    changed = [{**REQS[0], "description": REQS[0]["description"] + " (개정)"}]
    assert propose_subjects(_items()[:1], changed, SYSTEM, ROWS, cache_path=cache)["not_asked"] == 1


def test_a_cache_that_is_no_json_is_kept_aside_and_a_failed_save_is_said(tmp_path, monkeypatch):
    """(review W5) a corrupted cache was overwritten by the next save; a replace over an open file (Windows) lost the
    run's answers and reported the whole step as failed."""
    import os

    from generators import sts_ai_review as mod
    cache = tmp_path / "ai.json"
    cache.write_text("{not json", encoding="utf-8")
    propose_subjects(_items()[:1], REQS, SYSTEM, ROWS, ai_config=CFG, cache_path=cache, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_HallSnsrLevel", "quote": "u16g_ApiIn_HallSnsrLevel 이 4.85V 미만이면"}]))
    assert [p.read_text(encoding="utf-8") for p in tmp_path.glob("ai.json.corrupt-*")] == ["{not json"]

    def locked(src, dst):
        raise PermissionError(5, "Access is denied")
    monkeypatch.setattr(mod.os, "replace", locked)
    monkeypatch.setattr(mod.time, "sleep", lambda _s: None)
    stats = propose_subjects(_items()[:1], [{**REQS[0], "description": "다른 설명"}], SYSTEM, ROWS, ai_config=CFG,
                             cache_path=cache, llm=_llm([{"id": 0, "subject": None}]))
    assert stats["cache_error"].startswith("캐시 저장 실패") and stats["unknown"] == 1
    assert not [p for p in os.listdir(tmp_path) if p.endswith(".tmp")]


def test_a_failed_call_stops_the_asking():
    """(review W6) every batch was tried after the first failed — up to half an hour each on a stalled service."""
    calls = []

    def down(cfg, messages, stage=None, meta_out=None):
        calls.append(cfg)
        raise ConnectionError("refused")
    items = [_one("5V 미만", f"전압 {i} 5V 미만") for i in range(10)]
    rows = [{"id": "H", "vars": ["u16g_ApiIn_Vsup"], "ids": [], "nets": [], "name": ""}]
    progress: list = []
    stats = propose_subjects(items, [], None, rows, ai_config={**CFG, "retries": 10, "read_timeout": 300}, llm=down,
                             on_progress=progress.append)
    assert len(calls) == 1 and stats["call_failed"] == 10 and stats["asked"] == 8
    assert (calls[0]["retries"], calls[0]["read_timeout"], calls[0]["temperature"]) == (2, 120, 0)
    assert progress == ["AI 제안 8/10", "AI 제안 10/10"]


def test_the_disclosure_warns_on_failures_and_says_when_no_ai_was_called():
    """(review W7) all calls failed was an info line nobody saw; a path that calls no AI was 'AI 설정 없음'."""
    from report_gen.generation_disclosures import build_disclosures

    def item(ai):
        return {i["key"]: i for i in build_disclosures("sts", {"generation_stats": {"requirement_boundary": {
            "tcs": 1, "steps": 1, "facts": 1, "facts_used": 1, "ai_review": ai}}})}.get("sts_ai_review")
    base = {"eligible": 9, "verified": 0, "verified_by_quote": 0, "verified_by_block_name": 0, "rejected": 0, "unknown": 0,
            "not_asked": 0, "cached": 0, "model": "m"}
    assert item({**base, "called": True, "call_failed": 9, "asked": 8})["tone"] == "warning"
    assert item({**base, "called": True, "cache_error": "캐시 저장 실패 — PermissionError"})["tone"] == "warning"
    assert "이번 생성은 AI 를 부르지 않음" in item({**base, "called": False, "not_asked": 9})["note"]


def test_the_hsis_names_are_candidates_without_a_hwrs(tmp_path, monkeypatch):
    """(review W8) the HSIS rows were read only with a HwRS — without one the AI saw no HSIS name."""
    from generators import sts as sts_mod
    from generators import sts_ai_review as mod
    seen = []
    real = mod.propose_subjects

    def spy(items, requirements, system, hsis_rows=None, **kw):
        seen.append(hsis_rows)
        return real(items, requirements, system, hsis_rows, **kw)
    monkeypatch.setattr(mod, "propose_subjects", spy)
    monkeypatch.setattr(sts_mod, "_load_hsis_signals", lambda _p: {
        "sw_var_names": ["u16g_ApiIn_HallSnsrLevel"], "pat": sts_mod._HW_SIGNAL_PAT, "signals": [
            {"id": "HSI_1", "signal_name": "VCC_HALL_MON", "sw_var_name": "u16g_ApiIn_HallSnsrLevel", "related_id": ""}]})
    text = ["SwTSR_0101: 센서 전원 이상 시 고장을 검출한다. Verification: Pull-up Port를 4.85V 미만으로 Setting"]
    sts_mod.generate_sts(text, {}, str(tmp_path / "a.xlsm"), hsis_path=str(tmp_path / "hsis.xlsx"),
                         project_config={"project_id": "T", "ai_review_cache": str(tmp_path / "ai.json")})
    assert seen and seen[0] and seen[0][0]["vars"] == ["u16g_ApiIn_HallSnsrLevel"]


def test_the_subject_reasons_are_the_generators():
    """(review X5) a renamed reason would leave the AI step with no item, silently."""
    from generators.sts_ai_review import SUBJECT_REASONS
    from generators.sts_requirement_tc import REVIEW_TEXT
    assert SUBJECT_REASONS <= set(REVIEW_TEXT)



def test_a_name_listed_with_others_in_one_parenthesis_is_not_pinned():
    """HDPDM01 SyII_19: 'Control Signal(EN1, EN2, INA, INB, SE_EN) : 0~5V' — the value is all five's; SE_EN was 'checked'."""
    item = _one("0~5V", "Control Signal(EN1, EN2, INA, INB, SE_EN) : 0~5V")
    propose_subjects([item], [], None, None, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "SE_EN", "quote": "Control Signal(EN1, EN2, INA, INB, SE_EN) : 0~5V"}]))
    assert item["ai_proposal"]["reason"] == "quote_names_several"


def test_a_block_name_proposal_shows_the_field_not_the_models_quote():
    """A value-only field's block name is no reading of the model's quote — a description without the value was shown
    as its ground (KJPDS02_PV SyII_28)."""
    item = _one("0~5V", "0~5V", "no_subject_in_value_field", "Motor Driver Current", "SyDS SyII_28 · Range")
    system = {"SyII_28": {"doc": "SyDS", "fields": {"Range": "0~5V", "Description": "Motor 의 Current 를 Feedback"}}}
    propose_subjects([item], [], system, None, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "Motor Driver Current", "quote": "Motor 의 Current 를 Feedback"}]))
    assert item["ai_proposal"]["by"] == "block_name" and item["ai_proposal"]["quote"] == "0~5V"
    assert "값만 적는 칸 '0~5V' 의 블록 이름" in proposal_text(item) and "Feedback" not in proposal_text(item)


def test_a_cached_answer_is_judged_by_todays_rules(tmp_path):
    """The cache holds the model's answer, never a verdict — a rule tightened since (a requirement tag is no
    candidate) applies to answers cached before it."""
    from generators.sts_ai_review import _candidates, _key, _texts
    item = _items()[0]
    reqs = {r["id"]: r for r in REQS}
    key = _key(item, _candidates(item, reqs, SYSTEM, ROWS)[0], _texts(item, reqs, SYSTEM))
    cache = tmp_path / "ai.json"
    cache.write_text(json.dumps({key: {"model": "old", "answer": {
        "subject": "FS_REQ_MD_20", "quote": "Position Sensor 입력 전압을 4.85V 미만으로 설정", "reason": ""}}}),
        encoding="utf-8")
    propose_subjects([item], REQS, SYSTEM, ROWS, cache_path=cache)
    assert (item["ai_proposal"]["status"], item["ai_proposal"]["cached"]) == ("rejected", True)


# ── review round 2 ─────────────────────────────────────────────────────────────────────────────────────────────────


def _ask(item, subject, quote, rows=None, system=None):
    propose_subjects([item], [], system, rows, ai_config=CFG, llm=_llm([{"id": 0, "subject": subject, "quote": quote}]))
    return item["ai_proposal"]


def test_a_block_name_the_quote_ties_to_the_value_meets_the_competitors():
    """(review r2 W1) the block-name path skipped the competitor check even when the quote, not the field, tied it."""
    item = _one("4.85V 미만", "u16g_ApiIn_MagnetLevel 및 Battery Monitor 4.85V 미만", "no_subject_in_value_field",
                "Battery Monitor", "SyDS SyII_99 · Range")
    rows = [{"id": "H", "vars": ["u16g_ApiIn_MagnetLevel"], "ids": [], "nets": [], "name": ""}]
    got = _ask(item, "Battery Monitor", "u16g_ApiIn_MagnetLevel 및 Battery Monitor 4.85V 미만", rows)
    assert got["reason"] == "quote_names_several"


def test_a_value_is_read_with_its_sign_and_a_hyphen_range_is_no_sign():
    """(review r2 W2) '-0.5V 이하' passed on '0.5V 이하', '-40 ℃' on '40 ℃'."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_Vbat"], "ids": [], "nets": [], "name": ""}]
    for fact, quote in (("-0.5V 이하", "u16g_ApiIn_Vbat 0.5V 이하"), ("-40 ℃ 이상", "u16g_ApiIn_Vbat 40 ℃ 이상")):
        assert _ask(_one(fact, quote), "u16g_ApiIn_Vbat", quote, rows)["reason"] == "quote_does_not_link", fact
    for fact, quote in (("-40 ℃ 이상", "u16g_ApiIn_Vbat 가 -40 ℃ 이상"), ("0~5V", "u16g_ApiIn_Vbat 0-5V 범위")):
        assert _ask(_one(fact, quote), "u16g_ApiIn_Vbat", quote, rows)["status"] == "verified", fact


def test_a_parenthesis_with_one_name_and_its_unit_or_values_is_no_list():
    """(review r2 W3) '(u8g_ApiIn_DoorSpeed [m/s])' and '(u16g_X, 0x0~0xFF, offset : 4V)' were taken for name lists."""
    rows = [{"id": "H1", "vars": ["u8g_ApiIn_DoorSpeed"], "ids": [], "nets": [], "name": ""},
            {"id": "H2", "vars": ["u16g_ApiIn_MotorCur"], "ids": [], "nets": [], "name": ""}]
    q1 = "도어 속도(u8g_ApiIn_DoorSpeed [m/s]) 2.1 초과"
    assert _ask(_one("2.1 초과", q1), "u8g_ApiIn_DoorSpeed", q1, rows)["status"] == "verified"
    q2 = "출력(u16g_ApiIn_MotorCur, 0x0~0xFF, offset : 4V) 4V 이상"
    assert _ask(_one("4V 이상", q2), "u16g_ApiIn_MotorCur", q2, rows)["status"] == "verified"


def test_the_names_words_without_the_name_are_counted_apart():
    """(review r2 W4) a quote writing every word of the name but not the name itself is weaker evidence: 'words'."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_MagnetLevel"], "ids": [], "nets": [], "name": ""}]
    item = _one("4.85V 미만", "Magnet Level 이 4.85V 미만")
    stats = propose_subjects([item], [], None, rows, ai_config=CFG, llm=_llm([
        {"id": 0, "subject": "u16g_ApiIn_MagnetLevel", "quote": "Magnet Level 이 4.85V 미만"}]))
    assert item["ai_proposal"]["by"] == "words" and (stats["verified_by_words"], stats["verified_by_quote"]) == (1, 0)
    assert "이름 그대로는 아님" in proposal_text(item)


def test_a_net_name_with_a_hyphen_is_not_its_prefix():
    """(review r2 I1) ``V_BAT`` was 'written' in ``V_BAT-MON`` — the right name rejected for a candidate it contains."""
    rows = [{"id": "H1", "vars": ["u16g_ApiIn_Vbat"], "ids": [], "nets": ["V_BAT-MON"], "name": "V_BAT-MON"},
            {"id": "H2", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": ["V_BAT"], "name": "V_BAT"}]
    assert _ask(_one("5V 미만", "V_BAT-MON 5V 미만"), "V_BAT-MON", "V_BAT-MON 5V 미만", rows)["status"] == "verified"


def test_a_name_written_with_its_own_value_does_not_compete():
    """KJPDS02_PV SySM_04: the time constant written with its 300 ms rejected the voltage's own variable."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_HallSnsrLevel"], "ids": [], "nets": [], "name": ""}]
    quote = "Main 전원(u16g_ApiIn_HallSnsrLevel)이 4.85V 이하인 상태로 특정 시간(u16s_HALLPOWER_ERR_TM : 300ms) 이상"
    assert _ask(_one("4.85V 미만", quote), "u16g_ApiIn_HallSnsrLevel", quote, rows)["status"] == "verified"
    rows2 = [{"id": "H1", "vars": ["u16g_ApiIn_A"], "ids": [], "nets": [], "name": ""},
             {"id": "H2", "vars": ["u16g_ApiIn_B"], "ids": [], "nets": [], "name": ""}]
    shared = "u16g_ApiIn_A 와 u16g_ApiIn_B 가 5V 미만"           # followed by the item's value: it competes
    assert _ask(_one("5V 미만", shared), "u16g_ApiIn_B", shared, rows2)["reason"] == "quote_names_several"


def test_a_cache_of_no_answer_map_is_kept_aside_and_a_failed_dump_leaves_no_file(tmp_path, monkeypatch):
    """(review r2 I3 · I2) JSON that is no map was overwritten by the next save; a dump failing mid-write left its
    temporary file."""
    from generators import sts_ai_review as mod
    cache = tmp_path / "ai.json"
    cache.write_text("[1, 2, 3]", encoding="utf-8")
    assert mod._load_cache(cache) == ({}, "") and [p.read_text() for p in tmp_path.glob("ai.json.corrupt-*")] == ["[1, 2, 3]"]

    def full(*_a, **_k):
        raise OSError(28, "No space left on device")
    monkeypatch.setattr(mod.json, "dump", full)
    assert mod._save_cache(cache, {"k": {"model": "m", "answer": {}}}).startswith("캐시 저장 실패")
    assert not list(tmp_path.glob("*.tmp"))


def test_zero_retries_is_one_not_the_default():
    """(review r2 I4) `llm_call` reads 0 as its default (10)."""
    calls = []

    def once(cfg, messages, stage=None, meta_out=None):
        calls.append(cfg)
        return json.dumps({"answers": []})
    propose_subjects(_items()[:1], REQS, SYSTEM, ROWS, ai_config={**CFG, "retries": 0, "read_timeout": 0}, llm=once)
    assert (calls[0]["retries"], calls[0]["read_timeout"]) == (1, 10)



def test_a_hyphen_after_a_unit_is_a_range_not_a_sign():
    """'4.5V-5.1V' is a range: the '-' after a letter is no sign (the sign rule's other half)."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_Vbat"], "ids": [], "nets": [], "name": ""}]
    quote = "u16g_ApiIn_Vbat 4.5V-5.1V 범위"
    assert _ask(_one("4.5V ~ 5.1V", quote), "u16g_ApiIn_Vbat", quote, rows)["status"] == "verified"


def test_a_name_written_with_the_items_own_value_competes():
    """A candidate written right before the item's value is bound to it too — it competes (only another value excuses)."""
    rows = [{"id": "H1", "vars": ["u16g_ApiIn_A"], "ids": [], "nets": [], "name": ""},
            {"id": "H2", "vars": ["u16g_ApiIn_B"], "ids": [], "nets": [], "name": ""}]
    quote = "u16g_ApiIn_A = 5V 또는 u16g_ApiIn_B = 5V 미만"
    assert _ask(_one("5V 미만", quote), "u16g_ApiIn_B", quote, rows)["reason"] == "quote_names_several"


# ── review round 3 ─────────────────────────────────────────────────────────────────────────────────────────────────


def test_a_digit_inside_a_name_is_no_value():
    """(review r3 W1) '16V 초과' passed on 'u16g_ApiIn_BatVolt 감시' — the 16 of ``u16g``."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_BatVolt"], "ids": [], "nets": ["V_BAT_12V"], "name": "V_BAT_12V"}]
    for fact, quote in (("16V 초과", "u16g_ApiIn_BatVolt 감시"), ("12V 이상", "V_BAT_12V 감시 u16g_ApiIn_BatVolt")):
        assert _ask(_one(fact, quote), "u16g_ApiIn_BatVolt", quote, rows)["reason"] == "quote_does_not_link", fact


def test_the_subject_written_with_a_value_of_its_own_is_rejected():
    """(review r3 W2) the rule that excuses a competitor bound to its own value rejects such a subject."""
    rows = [{"id": "H", "vars": ["u16s_HALLPOWER_ERR_TM"], "ids": [], "nets": [], "name": ""}]
    quote = "4.85V 미만이 u16s_HALLPOWER_ERR_TM : 300ms 이상"
    assert _ask(_one("4.85V 미만", quote), "u16s_HALLPOWER_ERR_TM", quote, rows)["reason"] == "subject_bound_elsewhere"


def test_a_parenthesised_range_does_not_excuse_a_competitor():
    """(review r3 W3) 'Y(0~5V)' describes Y's range — it is no value of Y's that frees the quote's 4.85 V from it."""
    rows = [{"id": "H1", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": [], "name": ""},
            {"id": "H2", "vars": ["u16g_ApiIn_Y"], "ids": [], "nets": [], "name": ""}]
    quote = "u16g_ApiIn_X 와 u16g_ApiIn_Y(0~5V) 가 4.85V 미만"
    assert _ask(_one("4.85V 미만", quote), "u16g_ApiIn_X", quote, rows)["reason"] == "quote_names_several"


def test_numbers_read_alike_whatever_their_spelling():
    """(review r3 I) a minus sign (U+2212), thousands, trailing zeros, hex — one reading on both sides."""
    from generators.sts_ai_review import _numbers
    assert _numbers("−0.5V") == ["-0.5"] and _numbers("최소-5V") == ["-5"] and _numbers("4.5V-5.1V") == ["4.5", "5.1"]
    assert _numbers("(4.5)-5.1V") == ["4.5", "5.1"] and _numbers("-40 ℃") == ["-40"]
    assert _numbers("1,604") == ["1604"] and _numbers("5.0V") == ["5"] and _numbers("0x0~0xFF") == ["0x0", "0xff"]
    rows = [{"id": "H", "vars": ["u16g_ApiIn_Vsup"], "ids": [], "nets": [], "name": ""}]
    quote = "u16g_ApiIn_Vsup 1604 이상"
    assert _ask(_one("1,604 이상", quote), "u16g_ApiIn_Vsup", quote, rows)["status"] == "verified"
    unit = "속도(u16g_ApiIn_Vsup, km/h) 30 초과"
    assert _ask(_one("30 초과", unit), "u16g_ApiIn_Vsup", unit, rows)["status"] == "verified"


# ── review round 4 ─────────────────────────────────────────────────────────────────────────────────────────────────


def test_the_unit_written_with_the_number_must_agree():
    """(review r4 W1) '5V 미만' passed on 'u16g_A_X 5ms 경과' — the number alone matched."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": [], "name": ""}]
    quote = "u16g_ApiIn_X 5ms 경과 후 판정"
    assert _ask(_one("5V 미만", quote), "u16g_ApiIn_X", quote, rows)["reason"] == "quote_does_not_link"
    door = "Door Speed (u16g_ApiIn_X) 2.1[m/s]초과 입력"            # a bracketed unit is read
    assert _ask(_one("2.1m/s 초과", door), "u16g_ApiIn_X", door, rows)["status"] == "verified"


def test_a_subject_given_its_own_value_by_a_comparison_or_a_particle_is_rejected():
    """(review r4 W2) '==' · '>=' · '는' bound the subject to 300 too — only ':' and '=' were seen."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": [], "name": ""}]
    for quote in ("u16g_ApiIn_X == 300 이면 4.85V 미만", "u16g_ApiIn_X >= 300ms 이고 4.85V 미만",
                  "u16g_ApiIn_X 는 300ms, 4.85V 미만"):
        assert _ask(_one("4.85V 미만", quote), "u16g_ApiIn_X", quote, rows)["reason"] == "subject_bound_elsewhere", quote


def test_thousands_keep_their_decimals_and_plus_minus_stays():
    """(review r4 W3 · I) '1,604.5' was read as 1604; '±0.5V' as 0.5."""
    from generators.sts_ai_review import _numbers
    assert _numbers("1,604.5V") == ["1604.5"] and _numbers("±0.5V") == ["±0.5"] and _numbers("－0.5V") == ["-0.5"]
    rows = [{"id": "H", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": [], "name": ""}]
    quote = "u16g_ApiIn_X 1604 이상"
    assert _ask(_one("1,604.5V 이상", quote), "u16g_ApiIn_X", quote, rows)["reason"] == "quote_does_not_link"


def test_a_list_joined_by_words_is_a_list():
    """(review r4 I) '(EN1 및 SE_EN)' passed while '(EN1, SE_EN)' was rejected."""
    quote = "Control Signal(EN1 및 SE_EN) : 0~5V"
    assert _ask(_one("0~5V", quote), "SE_EN", quote)["reason"] == "quote_names_several"


# ── review round 5 ─────────────────────────────────────────────────────────────────────────────────────────────────


def test_a_count_after_a_competitor_does_not_excuse_it():
    """(review r5) 'X 또는 Y가 3회 연속 4.85V 미만' verified X — Y's '3회' was taken for a value of Y's own."""
    rows = [{"id": "H1", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": [], "name": ""},
            {"id": "H2", "vars": ["u16g_ApiIn_Y"], "ids": [], "nets": [], "name": ""}]
    for quote in ("u16g_ApiIn_X 또는 u16g_ApiIn_Y가 3회 연속 4.85V 미만", "u16g_ApiIn_X 또는 u16g_ApiIn_Y가 10ms 동안 4.85V 미만"):
        assert _ask(_one("4.85V 미만", quote), "u16g_ApiIn_X", quote, rows)["reason"] == "quote_names_several", quote


def test_a_count_after_the_subject_is_no_value_but_a_measure_is():
    """(review r5) 'X가 3회 연속 4.85V 미만' was rejected (the 3 read as X's value); '10ms 동안' still binds."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": [], "name": ""}]
    count = "u16g_ApiIn_X가 3회 연속 4.85V 미만"
    assert _ask(_one("4.85V 미만", count), "u16g_ApiIn_X", count, rows)["status"] == "verified"
    measure = "u16g_ApiIn_X이 10ms 동안 4.85V 미만"
    assert _ask(_one("4.85V 미만", measure), "u16g_ApiIn_X", measure, rows)["reason"] == "subject_bound_elsewhere"
    slash = "u16g_ApiIn_X 4.85V/5.14V 미만"                     # 'V/' is V
    assert _ask(_one("4.85V 미만", slash), "u16g_ApiIn_X", slash, rows)["status"] == "verified"


def test_a_parenthesis_after_the_subject_describes_it():
    """(review r3 W3 · r5) 'X(10ms 주기)' describes X — it gives X no value of its own, on the subject's side too."""
    rows = [{"id": "H", "vars": ["u16g_ApiIn_X"], "ids": [], "nets": [], "name": ""}]
    quote = "u16g_ApiIn_X(10ms 주기) 가 4.85V 미만"
    assert _ask(_one("4.85V 미만", quote), "u16g_ApiIn_X", quote, rows)["status"] == "verified"
