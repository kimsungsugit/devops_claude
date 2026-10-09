"""R71 — what stopped the MC/DC search for a decision it did not design (input loss audit #4 · #5 · #47).

The path search spends one run · step budget per function in decision order, cuts every input's value list to
`_PATH_SAMPLES`, and refuses a decision of more than ``max_conditions`` conditions; the expression engine looks at sample
values under a candidate cap; the Observable search stops at 48 candidates per pair. A decision those stopped read
``path_search_no_pair`` · ``path_evaluation:unreached`` · ``masked`` exactly like one searched to the end — nothing in the
reason, the summary or the sheet said 'budget'. Each limit is now on the decision (`mcdc_design.search_limits`), counted in
the summary, disclosed, and written in the 'Search Limit' column; no row, pair or expected value changes.
"""
from __future__ import annotations

import inspect
import os

import openpyxl
import pytest
from openpyxl.styles import Border, Font, PatternFill

from generators import c_project_context as cpc
from generators import mcdc_design as md
from generators import observable_mcdc as om
from generators import suts
from generators.mcdc_design import build_mcdc_design, search_limits
from generators.suts import summarize_mcdc_design
from report_gen import generation_disclosures as gd
from report_gen.generation_disclosures import build_disclosures
from tests.unit import test_mcdc_stub_return_inputs_r36 as r36
from tests.unit import test_mcdc_stubbed_call_in_condition_r37 as r37
from tests.unit import test_observable_mcdc_r60 as r60
from tests.unit.test_mcdc_path_design import H, _unit

INPUTS = ["a", "b", "c", "d", "e"]
# D1 ``(a + b) == 300U``: the expression engine (both operands are inputs) — no sample sums to 300 (b would be 45).
# D2: on two locals under D1 — never reached; its own search ends after the climb finds no way in.
# D3: R56's XOR decision — beyond 1,536 runs.
TWO = H + ("U8 g_o;\nvoid f(U8 a, U8 b, U8 c, U8 d, U8 e) { U8 t = (U8)(a ^ b); U8 u = (U8)(c ^ d); "
           "if ((a + b) == 300U) { if ((t == 1U) && (e == 2U)) { g_o = 3U; } } "
           "if ((t == 77U) && (u == 13U) && (e == 5U)) { g_o = 1U; } }\n")
# an enumeration input: its three values are the whole domain — nothing sampled, nothing cut
SMALL = H + ("typedef enum { M0 = 0, M1 = 1, M2 = 2 } MODE;\nU8 g_o;\n"
             "void g(MODE a) { U8 k = 3U; if ((k == 3U) && (a == M1)) { g_o = 1U; } }\n")
EXPR = H + ("U8 g_o;\nvoid h(U8 a, U8 b, U8 c) { if ((a + b) == 300U) { g_o = 1U; } "
            "if ((a == 1U) && (b == 2U)) { g_o = 2U; } "
            "if (((a == 1U) && (b == 2U)) || ((a == 1U) && (c == 3U))) { g_o = 3U; } }\n")
WIDE = H + ("U8 g_o;\nvoid w(" + ", ".join(f"U8 a{i}" for i in range(13)) + ") { if ("
            + " && ".join(f"(a{i} == {i}U)" for i in range(13)) + ") { g_o = 1U; } }\n")
WIDE_LOCAL = H + ("U8 g_o;\nvoid wl(U8 a) { U8 t = (U8)(a + 1U); if ("
                  + " && ".join(f"(t == {i}U)" for i in range(13)) + ") { g_o = 1U; } }\n")


def _kinds(decision):
    return [kind for kind, _detail in search_limits(decision)]


@pytest.fixture(scope="module")
def two():
    return build_mcdc_design(_unit(TWO, "f", INPUTS), max_path_runs=1536, max_path_steps=30_000)


def test_the_decision_whose_own_search_ran_out_of_budget_says_so(two):
    d1, d2, d3 = two["decisions"]
    assert (d3["status"], d3["reason"]) == ("no_pair_found", "path_search_no_pair")
    assert d3["path_search"]["budget_exhausted"] is True and "path_budget" in _kinds(d3)
    # D2's own search ended on its own (no vector reaches it) before D3 spent the rest: the function ran out, D2 did not
    assert d2["reason"] == "path_evaluation:unreached"
    assert d2["path_search"]["function_budget_exhausted"] is True
    assert d2["path_search"]["budget_exhausted"] is False and "path_budget" not in _kinds(d2)


def test_a_product_the_budget_cut_short_is_a_budget_limit_though_the_runs_did_not_fill_it():
    # behind ``(a == 0U) && (b == 0U)`` the frontier covers (0, 0, c) and its one-step neighbours; the influence product
    #   over (a, b, c) starts with those (already run) and goes on to vectors that never ran — the run budget's slice
    #   ended inside the part that had run, so the cache stays below the budget while the search was cut
    text = H + ("U8 g_o;\nvoid q(U8 a, U8 b, U8 c) { U8 t = c; if ((a == 0U) && (b == 0U)) "
                "{ if ((t == 77U) && (t == 78U)) { g_o = 1U; } } }\n")
    cut = build_mcdc_design(_unit(text, "q", ["a", "b", "c"]), max_path_runs=300, max_path_steps=10_000_000)
    d = cut["decisions"][1]
    assert d["status"] == "no_pair_found" and d["path_search"]["runs"] < 300
    assert (d["path_search"]["budget_exhausted"], d["path_search"]["function_budget_exhausted"]) == (True, False)
    assert _kinds(d) == ["path_budget", "sampled_values"]
    # the whole product (8 values each — 512 vectors) ran: the search ended on its own (over U8 samples)
    full = build_mcdc_design(_unit(text, "q", ["a", "b", "c"]))["decisions"][1]
    assert full["path_search"]["runs"] == 512 and _kinds(full) == ["sampled_values"]
    # (review W1) the step budget checked between chunks: the last chunk may overshoot it after every vector ran — no cut
    steps = full["path_search"]["steps"]
    over = build_mcdc_design(_unit(text, "q", ["a", "b", "c"]), max_path_steps=steps - 1)["decisions"][1]
    assert over["path_search"]["runs"] == 512 and over["path_search"]["function_budget_exhausted"] is True
    assert over["path_search"]["budget_exhausted"] is False and _kinds(over) == ["sampled_values"]
    # one chunk fewer is a cut
    short = build_mcdc_design(_unit(text, "q", ["a", "b", "c"]), max_path_steps=steps // 2)["decisions"][1]
    assert short["path_search"]["runs"] < 512 and "path_budget" in _kinds(short)


@pytest.mark.parametrize("text, name, inputs, budgets", [
    # the climb towards ``t`` behind ``a == 5U`` · ``b == 6U``: a budget that ends right after a climb step (every vector
    #   of it ran) still stopped the climb
    (H + "U8 g_o;\nvoid cl(U8 a, U8 b, U8 c) { U8 t = c; if (a == 5U) { if (b == 6U) "
         "{ if ((t == 1U) && (c == 2U)) { g_o = 1U; } } } }\n", "cl", ["a", "b", "c"], range(2, 60)),
    # the frontier around a decision every vector reaches and none can pair (``k`` is a constant local): a budget that
    #   ends right after a frontier expansion still left frontier vectors unexpanded (no product follows: no input
    #   changes the observation)
    (H + "U8 g_o;\nvoid fr(U8 a, U8 b) { U8 k = 3U; if ((k == 77U) && (k == 78U)) { g_o = 1U; } }\n", "fr", ["a", "b"],
     range(2, 60)),
])
def test_a_budget_that_ends_between_search_steps_is_a_cut(text, name, inputs, budgets):
    # (review W1) every budget below the natural end is a cut, wherever it falls
    natural = build_mcdc_design(_unit(text, name, inputs))["decisions"][-1]["path_search"]["runs"]
    assert natural > max(budgets)       # (review N6) otherwise the loop below would check nothing
    for runs in [b for b in budgets if b < natural]:
        d = build_mcdc_design(_unit(text, name, inputs), max_path_runs=runs)["decisions"][-1]
        assert d["path_search"]["budget_exhausted"] is True and "path_budget" in _kinds(d), runs


def test_the_first_search_count_wins_where_both_searches_cut_a_list():
    d = {"status": "no_pair_found", "evaluation": "source_path", "path_search": {"samples_capped": {"a": 14}},
         "stub_search_limit": {"samples_capped": {"a": 16, "get() return": 20}}}
    assert search_limits(d) == [("path_samples_capped", "a(14),get() return(20)")]


def test_no_budget_at_all_is_a_budget_limit():
    # not even the base vector ran: every path decision was left unsearched
    (d,) = build_mcdc_design(_unit(SMALL, "g", ["a"]), max_path_runs=0)["decisions"]
    assert d["reason"] == "path_evaluation:no_run" and _kinds(d) == ["path_budget"]


def test_the_value_lists_the_search_cut_are_named(two):
    _d1, d2, d3 = two["decisions"]
    # every input had 14 candidates (0 · 1 · 255 · each constant ±1 inside U8) — the search tried 12
    assert d3["path_search"]["samples_capped"] == dict.fromkeys(INPUTS, 14)
    assert dict(search_limits(d2))["path_samples_capped"] == "a(14),b(14),c(14),d(14),e(14)"


def test_a_search_that_ended_on_its_own_has_no_limit():
    (d,) = build_mcdc_design(_unit(SMALL, "g", ["a"]))["decisions"]
    assert (d["status"], d["reason"]) == ("partial", "path_search_incomplete")   # ``k == 3U`` never flips
    ps = d["path_search"]
    assert (ps["budget_exhausted"], ps["function_budget_exhausted"], ps["samples_capped"], ps["values_sampled"]) == (
        False, False, {}, False)
    assert search_limits(d) == []


def test_the_path_search_says_it_looked_at_samples():
    # (review W5) the same decision through a local: the path engine also looks at samples of a wide domain — a = 11
    #   (``a * 3U == 33U``) is no sample of U8 here, as for the expression engine
    expr = build_mcdc_design(_unit(H + "U8 g_o;\nvoid x(U8 a, U8 b) { if (((a * 3U) == 33U) && (b == 1U)) { g_o = 1U; } }\n",
                                   "x", ["a", "b"]))["decisions"][0]
    path = build_mcdc_design(_unit(H + "U8 g_o;\nvoid y(U8 a, U8 b) { U8 t = (U8)(a * 3U); "
                                   "if ((t == 33U) && (b == 1U)) { g_o = 1U; } }\n", "y", ["a", "b"]))["decisions"][0]
    assert expr["status"] != "designed" and _kinds(expr) == ["sampled_values"]
    assert path["evaluation"] == "source_path" and path["status"] != "designed"
    assert path["path_search"]["samples_capped"] == {} and _kinds(path) == ["sampled_values"]


def test_a_refusal_the_values_steer_still_has_a_value_limit():
    # (review N1) every sample indexes past ``g_arr[4]`` (a / 20 - 3 for a in 0 · 1 · 255 · 29..31 ...): all runs are
    #   undefined behaviour — yet a = 60 runs and pairs ``b == 1U``. The samples are why: the limit is written
    text = H + ("U8 g_o;\nU8 g_arr[4];\nvoid u(U8 a, U8 b) { U8 t = a; g_arr[(U8)(a / 20U - 3U)] = 1U; "
                "if ((t > 30U) && (b == 1U)) { g_o = 1U; } }\n")
    (d,) = build_mcdc_design(_unit(text, "u", ["a", "b"]))["decisions"]
    assert set(d["path_search"]["observations"]) == {"unsupported:undefined_behavior"}
    assert d["status"] != "designed" and set(_kinds(d)) & {"sampled_values", "path_samples_capped"}


def test_no_value_limit_where_every_run_is_the_same_model_refusal():
    # (review W4) a ``goto`` the model does not run: every vector gives ``unsupported:goto_unmodeled`` — more or other
    #   values would not change that, so no value limit is written (the value lists were cut: 22 candidates each)
    text = H + ("U8 g_o;\nvoid gt(U8 a, U8 b) { U8 t = (U8)(a + b); goto L1; L1: "
                "if ((t == 3U) && (b == 4U) && (a != 5U) && (t != 6U) && (b != 7U)) { g_o = 1U; } }\n")
    (d,) = build_mcdc_design(_unit(text, "gt", ["a", "b"]))["decisions"]
    assert set(d["path_search"]["observations"]) == {"unsupported:goto_unmodeled"}
    assert d["path_search"]["samples_capped"] == {} and d["path_search"]["values_sampled"] is False
    assert search_limits(d) == []
    # (review N5) a call in a condition outside the call group: the same static ``effectful_condition`` on every run
    text = H + "U8 rd(void);\nvoid f(U8 a) { U8 t = a; if ((t > 1U) && (rd() == 2U)) { } }\n"
    (d,) = build_mcdc_design(_unit(text, "f", ["a"]))["decisions"]
    assert set(d["path_search"]["observations"]) == {"effectful_condition"} and "search_group" not in d
    assert d["path_search"]["values_sampled"] is False and search_limits(d) == []


def test_no_value_limit_where_every_run_ends_on_the_same_unknown_value():
    # (review P1) a volatile read first: every run ends ``undetermined:volatile_object`` whatever the inputs — no value
    #   limit (it would count among the limited decisions of the disclosure head)
    text = H + ("U8 g_o;\nvolatile U8 g_reg;\nvoid vc(U8 a, U8 b) { U8 r = g_reg; if ((r == 3U) && (a == 1U) && (b == 2U)) "
                "{ g_o = 10U; } else { g_o = (U8)(20U + 30U + 40U + 50U); } }\n")
    (d,) = build_mcdc_design(_unit(text, "vc", ["a", "b"]))["decisions"]
    assert set(d["path_search"]["observations"]) == {"undetermined:volatile_object"}
    assert d["path_search"]["samples_capped"] == {} and search_limits(d) == []
    # the volatile read second: ``a != 1U`` short-circuits it — the values decide what the run sees, the limit stays
    text = text.replace("(r == 3U) && (a == 1U)", "(a == 1U) && (r == 3U)")
    (d,) = build_mcdc_design(_unit(text, "vc", ["a", "b"]))["decisions"]
    assert len(d["path_search"]["observations"]) > 1 and search_limits(d)
    # two unknown values the input steers between (``a == 5U`` reads the volatile, otherwise a call): no run got past the
    #   model, but the values decided which unknown it met — the limit stays
    text = H + ("U8 g_o;\nvolatile U8 g_reg;\nU8 rd(void);\nvoid mx(U8 a, U8 b) { U8 r; if (a == 5U) { r = g_reg; } "
                "else { r = rd(); } if ((r == 3U) && (b == 2U)) { g_o = 1U; } }\n")
    d = build_mcdc_design(_unit(text, "mx", ["a", "b"]))["decisions"][-1]
    assert set(d["path_search"]["observations"]) == {"undetermined:call_return_value", "undetermined:volatile_object"}
    assert _kinds(d) == ["sampled_values"]


def test_a_designed_decision_has_no_limit():
    from tests.unit.test_mcdc_path_budget_r56 import XOR
    d = build_mcdc_design(_unit(XOR, "f", INPUTS))["decisions"][-1]     # R56: designed under the full budget
    assert d["status"] == "designed" and d["path_search"]["samples_capped"]    # its value lists were cut, yet it paired
    assert search_limits(d) == []


def test_the_expression_engine_says_its_values_were_samples_or_its_candidates_ran_out():
    d1, d2, d3 = build_mcdc_design(_unit(EXPR, "h", ["a", "b", "c"]))["decisions"]
    assert (d1["reason"], d1["values_sampled"]) == ("no_pair_in_candidate_domain", True)
    assert _kinds(d1) == ["sampled_values"]          # a + b == 300 at a = 255 · b = 45 — 45 is no sample
    assert d2["status"] == "designed" and _kinds(d2) == []
    # ``A && B || A && C``: the twin A is proven unpairable on the compiled atoms — no sample caveat
    assert d3["reason"] == "unique_cause_infeasible:coupled_condition" and _kinds(d3) == []
    capped = build_mcdc_design(_unit(EXPR, "h", ["a", "b", "c"]), max_candidates=2)["decisions"]
    assert capped[0]["reason"] == "candidate_budget_exhausted" and _kinds(capped[0]) == ["candidate_budget"]


def test_an_enumeration_tried_whole_is_no_sample():
    # (review W3) ``domain_exhaustive`` compares with ``max - min + 1`` — false for E0 · E5 although both were tried
    text = H + ("typedef enum { E0 = 0, E5 = 5 } MODE;\nU8 g_o;\n"
                "void e(MODE m, U8 k) { if ((m == E0) && (m == E5)) { g_o = 1U; } }\n")
    (d,) = build_mcdc_design(_unit(text, "e", ["m", "k"]))["decisions"]
    assert d["status"] == "no_pair_found" and d["search_complete"] is True and d["domain_exhaustive"] is False
    assert d["values_sampled"] is False and search_limits(d) == []


def test_a_decision_over_the_caps_was_not_searched_and_says_which_cap():
    (d,) = build_mcdc_design(_unit(WIDE, "w", [f"a{i}" for i in range(13)]))["decisions"]
    assert d["reason"] == "decision_or_condition_budget" and search_limits(d) == [("condition_cap", "13>12")]
    (d,) = build_mcdc_design(_unit(WIDE_LOCAL, "wl", ["a"]))["decisions"]
    assert d["reason"] == "path_refused:decision_or_condition_budget"
    assert search_limits(d) == [("condition_cap", "13>12")]
    first, second, _third = build_mcdc_design(_unit(EXPR, "h", ["a", "b", "c"]), max_decisions=1)["decisions"]
    assert "decision_cap" not in _kinds(first)
    assert second["reason"] == "decision_or_condition_budget" and search_limits(second) == [("decision_cap", "#2>1")]
    # (review W2) one past the cap the path search would have taken (a local): its reason is the static refusal, the
    #   cap is why it was not searched
    text = H + ("U8 g_o;\nvoid c(U8 a, U8 b) { U8 t = (U8)(a + 1U); if ((a == 1U) && (b == 2U)) { g_o = 1U; } "
                "if ((t == 3U) && (b == 4U)) { g_o = 2U; } }\n")
    _d1, d2 = build_mcdc_design(_unit(text, "c", ["a", "b"]), max_decisions=1)["decisions"]
    assert d2["reason"].startswith("local_variable_not_input:") and d2.get("evaluation") != "source_path"
    assert search_limits(d2) == [("decision_cap", "#2>1")]


def test_the_stub_value_search_records_its_own_budget():
    # ``(r == 3U) && (r == 4U)`` on a returned value: the first search (no stub input) ends on its own; the second
    #   (stub values) is cut at 4 runs. Under the default budget this toy is not limited and R36's "left as it was"
    #   holds; in real code the stub search's value lists are often cut (review I4) and then recorded
    (d,) = build_mcdc_design(r36._unit("none", ["get() return"]), max_path_runs=4, max_path_steps=100_000)["decisions"]
    assert d["path_search"]["budget_exhausted"] is False
    assert d["stub_search_limit"] == {"budget_exhausted": True} and _kinds(d) == ["stub_path_budget"]
    (d,) = build_mcdc_design(r36._unit("none", ["get() return"]))["decisions"]
    assert "stub_search_limit" not in d and search_limits(d) == []


def test_a_stub_climb_cut_inside_one_chunk_is_a_cut():
    # (review N7 — mutant M24) the strict stub search: the vector that would reach the decision is refused for the step
    #   budget inside the climb's one chunk, the climb then breaks (no progress) and nothing else looks at the budget —
    #   the refusal itself is the record
    text = ('#include "common.h"\nU8 g_a;\nU8 g_x;\nU8 g_o;\nU8 get(void) { return g_x; }\n'
            "void f(void) { U8 r = get(); if (g_a == 7U) { if ((r == 3U) && (g_a == 7U)) { g_o = 1U; } } }\n")
    path = os.path.join(r36.ROOT, "m24.c")
    context = cpc.build_project_context({os.path.join(r36.ROOT, "common.h"): r36.COMMON, path: text})
    unit = {"name": "f", "source_text": text, "source_path": path, "source_text_complete": True,
            "project_scope": cpc.build_scopes(context, [path])[path], "input_vars": ["g_a", "get() return"]}
    d2 = build_mcdc_design(unit, max_path_steps=20)["decisions"][-1]
    assert d2.get("stub_search_limit", {}).get("budget_exhausted") is True and "stub_path_budget" in _kinds(d2)


def test_a_decision_the_run_cannot_observe_tried_no_values(monkeypatch):
    # a stub writing through its pointer argument: the call-condition group does not search the decision (R37) — so no
    #   value list was cut for it, even where every input's list is
    monkeypatch.setattr(md, "_PATH_SAMPLES", 1)
    (d,) = build_mcdc_design(r37._unit("pointer", ["g_en", "getp() return"]))["decisions"]
    assert d["reason"] == "path_evaluation:effectful_condition"
    assert d["path_search"]["samples_capped"] == {} and d["path_search"]["values_sampled"] is False
    assert search_limits(d) == []


def test_summary_disclosure_and_sheet(two):
    units = [{"fid": "F1", "name": "f", "mcdc_design": two},
             {"fid": "F2", "name": "w", "mcdc_design": build_mcdc_design(_unit(WIDE, "w", [f"a{i}" for i in range(13)]))}]
    s = summarize_mcdc_design(units)
    # D1 sampled_values only · D2 path_samples_capped · D3 path_budget + path_samples_capped · w condition_cap
    assert (s["search_limited_decisions"], s["sampled_only_decisions"]) == (3, 1)
    assert s["search_limited_by_status"] == {"partial": 0, "no_pair_found": 1, "unsupported": 2}
    assert {k: v for k, v in s["search_limits"].items() if v} == {
        "path_budget": 1, "path_samples_capped": 2, "sampled_values": 1, "condition_cap": 1}
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": s}) if i["key"] == "suts_mcdc_design"]
    for text in ("탐색이 상한에 닿은 결정 3(쌍 못 찾음 1 · 미지원 2)", "함수 실행 모델 탐색이 예산에 잘림 1",
                 "입력 값 목록을 12 개로 자른 탐색 2",
                 "그 밖에 상한에는 닿지 않았지만 도메인보다 좁은 표본 값(0 · ±1 · 상수 ±1 · 양 끝)만 본 결정 1",
                 "조건 수 상한(12 초과 — 탐색하지 않음) 1", "끝까지 찾아본 결과가 아니다", "상한이 원인이라는 뜻은 아니다",
                 "'Search Limit' 열", "path_evaluation:unsupported:path_budget", "'Search Complete = yes'"):
        assert text in item["note"], text
    assert "stub 값 탐색이 예산에 잘림" not in item["note"]          # a kind with no decision is not written
    assert "911" not in item["note"]                             # (review I6) no project's numbers in every document
    # (review N2) the sampled-only kind is not among the limits in the head count
    assert "표본 값(0 · ±1 · 상수 ±1 · 양 끝)만 본 탐색 1" not in item["note"]
    # an older summary (no key) and a summary with none: no sentence
    for summary in ({k: v for k, v in s.items() if not k.startswith(("search_limit", "sampled_only"))},
                    {**s, "search_limited_decisions": 0, "sampled_only_decisions": 0}):
        (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": summary})
                   if i["key"] == "suts_mcdc_design"]
        assert "상한에 닿은 결정" not in item["note"] and "표본 값" not in item["note"]
    # only sampled decisions: no limit reached, the sampled count still said
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": {**s, "search_limited_decisions": 0}})
               if i["key"] == "suts_mcdc_design"]
    assert "상한에 닿은 결정은 없다. 다만 그중 1 개는 도메인보다 좁은 표본 값" in item["note"]
    assert "입력이 고르는 분기 안에만 있는데" in item["note"]       # (review Q3) the rule's limit is said
    # (review P2) no limit reached: no word on raising limits or on caps that did not apply
    assert "상한을 올리면" not in item["note"] and "조건 수 · 결정 수 상한" not in item["note"]
    wb = openpyxl.Workbook()
    suts._write_mcdc_design_sheet(wb, units, {}, {"F1": "TC_1", "F2": "TC_2"}, Border(), PatternFill(), Font(), Font())
    rows = list(wb["MCDC Design"].iter_rows(values_only=True))
    header = list(rows[0])
    assert header == suts._MCDC_HEADERS and header.index("Search Limit") == header.index("Search Complete") + 1
    cells = {(r[1], r[2]): r[header.index("Search Limit")] for r in rows[1:] if not r[3]}
    assert cells[("f", "D1")] == "sampled_values"
    assert cells[("f", "D3")] == "path_budget · path_samples_capped:a(14),b(14),c(14),d(14),e(14)"
    assert cells[("w", "D1")] == "condition_cap:13>12"


def test_the_observable_search_says_where_it_stopped_at_its_candidate_cap():
    unit = r60._unit()
    seqs = r60._ext(unit)
    by_num = {s["seq_num"]: s for s in seqs}
    pair = next(p for p in r60._outer(unit)["pairs"] if p["condition_id"] == "C1")
    a_inputs = {k: int(v) for k, v in by_num[pair["seq_a"]]["inputs"].items()}
    # inside (0, 3) g_inner never opens the inner write: every candidate is masked
    bases = [{"seq_num": 997, "inputs": {**a_inputs, "g_inner": 2}}, {"seq_num": 996, "inputs": {**a_inputs, "g_inner": 1}}]

    def mark(bases, cap):
        found = om.find_observable_pairs(unit, seqs, ["g_out"], {"g_inner": (0, 3)}, bases, max_candidates=cap)
        return {id(p): m for p, m in found["marks"]}[id(pair)], found["report"]

    capped, report = mark(bases, 1)
    assert (capped["status"], capped["candidates_tried"], capped.get("candidates_capped")) == ("masked", 1, True)
    flagged = sum(1 for m in om.find_observable_pairs(unit, seqs, ["g_out"], {"g_inner": (0, 3)}, bases,
                                                     max_candidates=1)["marks"] if m[1].get("candidates_capped"))
    assert report["candidates_capped"] == flagged >= 1
    whole, _ = mark(bases, 2)
    assert whole["status"] == "masked" and "candidates_capped" not in whole
    # what is left over gives no new overlay (the same g_inner again): nothing was left untried
    same, _ = mark([bases[0], {"seq_num": 995, "inputs": {**a_inputs, "g_inner": 2}}], 1)
    assert "candidates_capped" not in same
    assert suts._observable_cells(capped)[0] == f"masked [own:{capped['own_rows']}] [candidates_capped:1]"
    pair["observable"] = capped
    s = summarize_mcdc_design([unit])
    assert s["observable_candidates_capped"] == 1
    (item,) = [i for i in build_disclosures("suts", {"mcdc_design_summary": s}) if i["key"] == "suts_mcdc_design"]
    assert "쌍당 후보 48 에서 멈춰 남은 후보를 보지 않은 쌍 1" in item["note"] and "`candidates_capped`" in item["note"]
    # (R71 measurement) the clause names the states it counts within — not read as a part of the item before it; only
    #   states that are in the list (here masked alone), and '출력을 못 찾음', not the head's '쌍 못 찾음' (review R1 · R2)
    assert "앞의 '출력을 못 찾음' 중 쌍당 후보 48" in item["note"] and "그중 쌍당 후보" not in item["note"]
    assert "'복사와 같은 값'" not in item["note"]
    # a capped pair is one of those four states: the count never exceeds them
    assert s["observable_candidates_capped"] <= sum(s.get(k, 0) for k in (
        "observable_copy_only", "observable_masked", "observable_underived", "observable_recheck_refused"))


def test_the_observable_cap_clause_names_the_nonzero_states_in_list_order():
    text = gd._observable_text({"observable_pairs": 3, "observable_copy_only": 2, "observable_masked": 0,
                                "observable_underived": 4, "observable_recheck_refused": 0,
                                "observable_candidates_capped": 5})
    assert "앞의 '복사와 같은 값' · '도출된 출력 없음' 중 쌍당 후보 48 에서 멈춰 남은 후보를 보지 않은 쌍 5" in text
    assert "'출력을 못 찾음'" not in text and "'판정을 확인하지 못함'" not in text
    assert "쌍당 후보 48 에서 멈춰" not in gd._observable_text({"observable_pairs": 3, "observable_masked": 2})


def test_the_disclosed_caps_are_the_generator_caps():
    defaults = inspect.signature(build_mcdc_design).parameters
    assert (gd.MAX_MCDC_PATH_SAMPLES, gd.MAX_MCDC_CONDITIONS, gd.MAX_MCDC_DECISIONS) == (
        md._PATH_SAMPLES, defaults["max_conditions"].default, defaults["max_decisions"].default)
    assert [k for k, _label in gd._SEARCH_LIMIT_LABELS] == list(md.SEARCH_LIMITS)
