"""Requirement boundary test cases for STS (R6, P3/G4) — built from the requirement text alone.

`generators/requirement_oracle.py` reads a requirement's own sentences into facts (``u16g_ApiIn_Vsup: 1604 초과``,
``Battery 전압이 8.5V미만``, ``특정시간(300ms)초과하여 유지``). For each parsed threshold or hold time this module writes
one test step per stimulus point around it:

* threshold ``subject op v``: ``v − δ``, ``v``, ``v + δ`` — the points where ``<`` and ``<=`` (``>`` and ``>=``) differ
  and where ``v`` moved by one step would change the verdict;
* hold time ``≥ t``: the condition lasting ``t − δ``, ``t``, ``t + δ``.

``δ`` is **one unit of the precision the value is written with** (``8.5`` → ``0.1``, ``4.00`` → ``0.01``,
``1604``/``0x1FF0`` → ``1``): the requirement states no tolerance or resolution, so none is invented — the step is
disclosed as ``written_precision`` per fact.

**What is set**: the subject the requirement compares. For ``Param_X( 3도 ) 초과한 열림각`` that is the opening angle
(``열림각 = 2도``, with ``Param_X = 3도`` as the reference) — never the constant itself; a parameter threshold whose
compared quantity the sentence does not name is not stepped. A hold time is the condition's duration (``LIN 통신 지속
시간 = 14초 (끊길 상태)``); a named time constant is its reference.

**The verdict** of each step is the one **the sentence** gives at that point: every condition of the line on the same
subject, joined by the words between them (``4.00V 이하 또는 4.74V 이상`` at 5 V: the second holds → the behaviour
the sentence describes happens). A point where the sentence's condition does not hold is *not in that sentence's
scope* — it is not claimed that the requirement does nothing there (another sentence of the same requirement may cover
it: ``3km/h 이하`` / ``3km/h 초과`` — review round 3 C-1). Conditions on other subjects are held false when joined
by *or*, true when joined by *and* (a hold time is always *and*); a join the words do not state — also with the next or
previous line — is left to the tester ("결합 조건 확인 필요"). A fact with such conditions gets a TC of its own, so its
precondition never contradicts another fact's steps. Not turned into steps, each with its reason: a fact without a
subject, a non-numeric comparison, a range, an operator that is not an order comparison, a negated clause (``… 아닌
경우``), a response constraint (``100ms 이내``), an outcome (``<완료조건>`` / ``<Output>`` sections), a label that names
a reference (``기준 전압: …``), a same-subject join the line does not state, a raw value outside the width the subject's
name declares (``u8g_X 45000 이상``), and a sentence already written.

**Traced system requirements (R29, G4(b))**: an SRS block's ``Related ID`` names the system requirements it realises
(``SyTR_0602``, ``SyII_06`` …). When the SyRS / SyDS are given, the same reading is applied to the text fields of
**each directly cited** system block (one hop — the system block's own Related IDs are not followed), and its facts get
steps under the citing SRS requirement, the step quoting the system sentence and naming the block and field it came
from. Nothing is matched by value or by guess: a system block contributes only because the SRS cites it. A fact the SRS
requirement already stepped is not repeated (only the same subject wording counts as the same — an alias is not
guessed); facts of different documents (the SRS, the SyRS, the SyDS) never share a TC, so every TC comes from one
document. What a system block states as an outcome is not stepped: a verification case's ``Output :`` line, the
``Action`` / ``System Behavior`` fields, and an obligation (``암전류는 0.3mA 이하여야 한다``).

**What the text leaves unwritten (R45)**: a lower bound below an upper bound of one subject side by side with no
connective has one meaningful reading, a range (`requirement_oracle.joint`; a lower bound above the upper one is a
hysteresis, not read), and two names that split one unit at one value (``Manual Assist조건(0.8m/s 미만) 또는 Tip-To-Run
조건(0.8m/s 이상 …)``) are one quantity in two regions — each step writes the other region's verdict at its point. What
was read so is listed to confirm; a fact still held back for a reason a reader can settle (a missing subject, a range's
inclusion, an unwritten join …) is listed with what to decide — the 'Requirement Review' sheet and the generation
disclosure — never filled by a guess.

**Evidence and what filling would improve (R46)**: for each held-back fact, the other sentences of the given documents
(every SRS requirement, every SyRS / SyDS block — also a SyDS element's ``Input Information``) that state **the same
value in the same unit in a form this module steps** are quoted as candidates (``BAT 전압(9V 이상 16.0V 이하)`` for a
``Range`` cell's ``9 ~ 16V``) — never used: the same number may be another quantity. Each item also says what writing
the missing part into the requirement would produce, **tried in that very sentence** (`fill_guidance`: the fill is
written in, the sentence read again; the value must step under the written name and nothing stepped before may stop —
else the text says what still holds it back); what only a rewrite can settle (a negation, a deadline, a label) gets an
example that keeps the sentence's meaning.

**Input cells (R47)**: a SyDS element's ``Input Information`` (``BAT 전압(9V 이상 16.0V 이하)`` · ``Watchdog Input(Pulse
주기 5ms 이하)`` · ``Motor Rotation : …, 5000RPM 초과``) is read as the other text fields are — its bounds step under the
SRS requirement that cites the block, in TCs of their own after every other traced fact. A cell writes a normal range
and a trigger alike, so a step says only whether the input is inside or outside what the cell writes; the element's
reaction is the requirement's other sentences' to decide.

**Value differences (R48)**: an SRS condition value that no cited system block states, where a cited block states a
condition on the same side in the same unit that is likely the same threshold — within 5%, the same subject, one written
step, or the only value each side writes; a hold time, the hold time of the same condition — ``5.14V 이상`` ↔ ``5.15V
초과`` · ``500ms`` ↔ ``300ms`` of ``8.5V 미만`` — is listed with R33's inclusion conflicts, never decided, never used.

**HW measurement tolerance (R49)**: given the HW requirements specification, a boundary group whose requirement cites
a system block that a HW block cites too gets that HW block's monitor accuracy (``허용 오차: ±3%`` — a percentage only
for the units the block measures; an absolute one only on the scale the block writes) quoted next to it; where the
step spacing lies inside it, the TC's precondition says a HIL run cannot tell the inclusion — decide by injecting the
SW variable, or check the direction outside the largest candidate (the monitor path is undecided when a hub block
links several). Quoted, never used to move a point.

**Monitor path and scale (R50)**: the HSIS row of a fact's SW signal (the same name, or the same name under another
layer — ``u16g_ApiIn_Vsup`` / ``u16g_DrvIn_Vsup``) narrows a hub to the one HW block that alone cites one of the row's
system IDs, or whose text names the row's net (``V-BAT``); the HW design specification's ``Sensor Power Monitor =
Sensor Power *0.5`` (a block citing the HW requirement) puts a monitor node's ``±0.3V`` on the 5 V scale (``±0.6V``) —
only where the value is off the node's scale and on the converted one.
"""
from __future__ import annotations

import hashlib
import re
from collections import Counter
from decimal import Decimal
from typing import Any

from generators.requirement_oracle import (
    condition_part_join,
    condition_units,
    exact_value,
    extract,
    is_outcome_section,
    join_is_inferred,
    joint,
    line_holds,
    same_subject,
    subject_of,
    written_step,
)

_ORDER = {"<", "<=", ">", ">="}
_OP_TEXT = {"<": "미만", "<=": "이하", ">": "초과", ">=": "이상", "==": "와 같음", "!=": "이 아님"}
ACTION_PREFIX = "입력 설정 (요구 경계): "   # the STS step classifier reads this as boundary value analysis (BAA)
BASIS_MARK = " — 근거: "                   # the requirement text quoted after it is not a stimulus
DURATION_SUFFIX = " 지속 시간"
_VERDICT_TEXT = {True: "성립", False: "불성립", None: "판정 미상(결합 미기재)"}

# (R45) held-back facts a reader can settle — 2026-09-30 user direction: what is wrong or undecided is shown in the
#   document and on the web even when there is no basis to fill it. The rest (an outcome, a response constraint, a
#   symbolic or equality comparison, a repeat) is no condition to step at all. Each: why it is held back, what to decide.
REVIEW_TEXT = {
    "no_subject_for_value": ("값이 무엇의 값인지(주어) 원문에 없음", "이 값과 비교하는 신호·변수를 정하면 경계 TC 가 된다"),
    "subject_unclear": ("주어로 읽힌 말이 신호 이름이 아님", "비교하는 신호·변수 확인"),
    "kind_range": ("범위(~)의 경계 포함 여부가 적혀 있지 않음", "하한·상한의 포함(이상/초과, 이하/미만) 확인"),
    "parenthesis_labels_may_share_a_quantity": ("괄호 앞 이름 둘이 같은 단위 — 같은 양인지 원문에 없음",
                                                "두 조건이 한 신호의 구간인지, 다른 신호인지 확인"),
    "same_subject_combination_unstated": ("같은 주어의 조건 결합(그리고/또는)이 적혀 있지 않음", "결합 확인"),
    "negated_condition": ("부정 절 — 부정이 조건 전체에 걸리는지 서술어에 걸리는지 원문으로 가려야 함", "부정의 범위 확인"),
    "monitored_quantity_unknown": ("기준 상수가 비교하는 감시량이 적혀 있지 않음", "감시하는 신호 확인"),
    # (R45 review W1) a value-only field (a block's ``Range``) names its quantity in the block name, not in the cell —
    #   the document did not leave it out; the generator does not read the name as a subject
    "no_subject_in_value_field": ("값 전용 칸(Range)이라 주어가 이 칸에 없음 — 블록 이름이 이 값의 주어일 수 있음",
                                  # (R46 review I7) the generator never reads the Name: no promise here — the 'If
                                  #   Filled' column says what writing the name into the cell produces
                                  "블록 이름(Name)이 이 값과 비교하는 신호인지 확인"),
}
    # (R45 review r2 W-R2-2) ``20ms 이내에 … 되지 않으면`` / ``500ms 이내에 9V 이상으로 복귀할 경우``: the extractor reads
    #   every ``이내`` as a deadline, but in a conditional clause it may be the condition's time window
REVIEW_TEXT["deadline_or_window"] = ("'이내' 가 조건 절 안에 있음 — 응답 기한인지 조건의 시간 창인지 원문에 가려야 함",
                                     "응답 기한이면 측정 대상(스텝 아님), 조건의 시간 창이면 경계 TC 대상 — 어느 쪽인지 확인")
# (review r2 W-R2-1) a range never becomes a boundary step here, even with a subject: no promise of one
RANGE_DECIDE = {
    "no_subject_for_value": "이 범위와 비교하는 신호, 그리고 하한·상한의 포함(이상/초과, 이하/미만) 확인",
    "no_subject_in_value_field": "블록 이름(Name)이 이 범위의 주어인지, 그리고 하한·상한의 포함(이상/초과, 이하/미만) 확인",
}
REVIEW_REASONS = frozenset(REVIEW_TEXT) - {"no_subject_in_value_field", "deadline_or_window"}   # skip → review
_VALUE_FIELDS = frozenset({"Range"})
_WINDOW_CLAUSE = re.compile(r"않으면|하면|되면|이면|경우|(?<=[가-힣])면(?=\s|$|,)|\s시(?=\s|$|,)|때")
# (R45 review I3) what the generator **read** where the words are silent — stepped, and shown so a reader can confirm it
READ_TEXT = {
    "read_as_range": ("연결어 없이 붙은 하한·상한을 범위(그리고)로 읽어 스텝함",
                      "범위가 맞는지 확인(다른 뜻이면 이 경계 TC 의 판정이 바뀐다)"),
    "read_as_one_quantity": ("같은 단위를 한 값에서 나누는 두 이름을 한 양의 두 구간으로 읽어 스텝함",
                             "두 이름이 한 신호의 구간인지 확인(다른 신호면 상대 구간의 판정은 근거가 없다)"),
}
MAX_REVIEW_ITEMS = 30     # items kept for the quality report (all counted; the document lists all)


def _fmt(value: Decimal, fact: dict[str, Any]) -> str:
    text = str(fact.get("value_text") or "")
    bare = text.lower().lstrip("-")
    if bare.startswith("0x"):
        digits = len(bare) - 2
        return ("-" if value < 0 else "") + "0x" + format(abs(int(value)), "X").zfill(digits)   # keep base and width
    decimals = len(text.split(".", 1)[1]) if "." in text else 0
    return f"{value:.{decimals}f}"


# (R51) a verification line that only states a limit — ``- Load(Average)는 70% 이하.`` · ``- DDM, ADM : 88% 이하.``
#   (HDPDM01 / KJPDS02 ``SwNTR_0301``, read since the parser takes ``Verification Criteria`` by any capitals) — is a pass
#   criterion: what a test measures and judges, not the range a test sets its input from (R33's reading of verification
#   lines — right for ``Input : 입력전원 8.5V 이하``). CPU load cannot be set: stepping 69.99 % / 70 % / 70.01 % claimed
#   discrimination no test can deliver. The line's last comparison ends the sentence and no test-input label precedes it.
ACCEPT_PREFIX = "측정 (요구 판정 기준): "
_ACCEPT_TAIL = re.compile(r"\s*[.。]?\s*$")
_SETTABLE_NAME = re.compile(r"(?<![A-Za-z0-9_])[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9_]+(?![A-Za-z0-9_])")
_TEST_INPUT_LABEL = re.compile(r"input|precondition|입력|사전\s*조건|설정|setting|인가", re.IGNORECASE)


# (R51 review W2) a pass criterion whose subject the line leaves to the line above (``2) Slack Time`` / ``- Task 할당
#   시간의 10% 이상.``) is still a criterion — measured, so its subject is no stimulus to name
_CRITERION_MAY_SKIP = frozenset({"", "no_subject_for_value", "subject_unclear", "same_subject_combination_unstated",
                                 "parenthesis_labels_may_share_a_quantity", "monitored_quantity_unknown"})


def _acceptance_line(line: dict, req: dict[str, Any]) -> bool:
    """(R51) A verification line of limits only (`ACCEPT_PREFIX`) — read on the line, so ``그 외 : Proto 60% 이하,
    Master 80% 이하.`` is one criterion line, not a stimulus and a criterion. Only of a non-functional requirement
    (``SwNTR`` · ``SwNTSR`` — resources, timing) and only where no SW variable is named: ``- u16s_C: 30 초과`` is a test
    input a SIL run sets, ``- 차속 3km/h 이하`` of a functional requirement is one a HIL bench sets."""
    if not str(req.get("id") or "").startswith("SwNT"):
        return False
    facts = [f for f in line.get("facts") or [] if f.get("line_span")]
    if not facts or _SETTABLE_NAME.search(line.get("raw") or line["text"]):
        return False            # a variable or a net (``u16s_C`` · ``VCC_BAT``): something a test sets
    raw = line.get("raw") or line["text"]
    last = max(facts, key=lambda f: f["line_span"][1])
    first = min(facts, key=lambda f: f["line_span"][0])
    return bool(_ACCEPT_TAIL.match(raw[last["line_span"][1]:])) and \
        not _TEST_INPUT_LABEL.search(raw[:first["line_span"][0]])


# a sub-line for the heading search — ``-5% 이상.`` too (bullet or minus sign, it names no subject: the heading does)
_BULLET = re.compile(r"^\s*[-•·*ㆍ–]")
# (R51 review r3 W-1) what is cut off a quoted criterion: a hyphen only with a space after it — ``- -5% 이상.`` keeps
#   its sign, ``-5% 이상.`` is left whole
_LEAD_BULLET = re.compile(r"^\s*(?:[•·*ㆍ–]\s*|-\s+)")
_NUMBERING = re.compile(r"^\s*(?:\d+[.)]|[①-⑳]|[a-zA-Z][.)])\s*")


def _extract_head(anchor: str, field: str | None = None) -> str:
    """What `extract` reads before a field's text — ``ID<TAB>anchor`` and, for an outcome field, ``<Output>``. A line's
    ``span`` (absolute in that text) minus its length is the line's place in the field (R51 review r4 W-A)."""
    return f"ID\t{anchor}\n" + ("<Output>\n" if field in SYSTEM_OUTCOME_FIELDS else "")


def _criterion_heading(text: str, at: int) -> str:
    """(R51 review W-B) What a criterion line measures when the line itself does not say it: the nearest line above it
    that is no bullet (``2) Slack Time`` · ``3) 메모리 점유율``), numbering removed — ``""`` when the line is no bullet
    (it names its own subject) or ``at`` is outside ``text``. (review r4 W-A) ``at`` is the line's place in the field
    (`_extract_head`) — matching the sentence's text took the first of two ``- 70% 이하.`` under ``RAM 사용률`` and
    ``ROM 사용률`` (r3 W-2), and with another spacing or a longer row holding it, still the wrong one."""
    text = str(text or "")
    if not 0 <= at <= len(text):
        return ""
    start = text.rfind("\n", 0, at) + 1
    if not _BULLET.match(text[start:].split("\n", 1)[0]):
        return ""
    for above in reversed(text[:start].split("\n")):
        if above.strip() and not _BULLET.match(above):
            return _NUMBERING.sub("", above).strip()
    return ""


def _requirement_facts_by_field(req: dict[str, Any]) -> list[tuple[dict, dict, str]]:
    """``(line, fact, field)`` — `requirement_facts` with the field each fact is in (R51: a sentence repeated in both
    fields is two lines here, told apart by identity, never by its text)."""
    out = []
    for key in ("description", "verification"):
        blocks = extract(f"{_extract_head(str(req.get('id', '')))}{req.get(key) or ''}\n")
        out += [(line, fact, key) for block in blocks for line in block["lines"] for fact in line["facts"]]
    return out


def requirement_facts(req: dict[str, Any]) -> list[tuple[dict, dict]]:
    """(line, fact) pairs of a requirement's description and verification criteria — read as two texts, so an
    ``<Output>`` heading at the end of the description never makes the criteria an outcome (review r4 W-4)."""
    return [(line, fact) for line, fact, _key in _requirement_facts_by_field(req)]


# ── (R29) system requirements the SRS block cites ────────────────────────────────────────────────────────────────
SYSTEM_ID = re.compile(r"\bSy[A-Za-z]{1,6}_\d+(?:_\d+)?\b")
# the fields of a SyRS / SyDS block that state behaviour or conditions (measured over HDPDM01 and KJPDS02: the rest —
#   Name, Type, ASIL, Priority, Rationale, pin allocation, element names, failure rates — carry no stimulus).
#   (R47) a SyDS element's ``Input Information`` (``BAT 전압(9V 이상 16.0V 이하), 5V, …`` · ``Watchdog Input(Pulse 주기
#   5ms 이하)`` · ``5000RPM 초과``) states where the element's input is — R46 read it as evidence only
SYSTEM_TEXT_FIELDS = ("Description", "Functional Behavior", "System Behavior", "Condition", "Action", "Range",
                      "Time Constraint", "Verification criteria", "Input Information")
# (R47) what a step on an input cell says: the input is inside or outside what the cell writes (a normal range or a
#   trigger — the cell does not say which), never the element's reaction. ⚠ a field stating what the element puts out
#   (an "Output Information") belongs to `SYSTEM_OUTCOME_FIELDS`, not here (review X5)
SYSTEM_INPUT_FIELDS = frozenset({"Input Information"})
# (R29 review C2) what the system does — a state table's ``Action`` and a block's ``System Behavior`` — is read as an
#   outcome: its values are counted (``outcome_section``) and never stepped
SYSTEM_OUTCOME_FIELDS = frozenset({"Action", "System Behavior"})


def _field_key(name: str) -> str:
    """``Verification Criteria\\n(optional)`` and ``Verification criteria`` are one field."""
    return re.sub(r"\s+", " ", re.sub(r"\(optional\)", "", name, flags=re.IGNORECASE)).strip().casefold()


_FIELD_BY_KEY = {_field_key(f): f for f in SYSTEM_TEXT_FIELDS}


def parse_system_requirement_docx(path: str, label: str) -> tuple[dict[str, dict[str, Any]], int]:
    """``({system ID: {"doc": label, "fields": {field: text}}}, duplicate tables)`` from the attribute tables of a
    SyRS / SyDS docx (the SRS layout: ``ID | SyTR_0602``, ``Description | …``). A key cell merged over two columns
    (``ID | ID | SyOS_01``) reads the last distinct cell; a key written on two rows keeps both, in order (review I4:
    ``Time Constraint`` twice). The first table of an ID wins; later copies are counted. A nested table (a min/typ/max
    grid) is not part of its cell's text — it states no comparison sentence. Raises if the file cannot be read: the
    caller discloses it."""
    from docx import Document
    doc = Document(path)
    out: dict[str, dict[str, Any]] = {}
    duplicates = 0
    for table in doc.tables:
        parts: dict[str, list[str]] = {}
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 2 or not cells[0]:
                continue
            values = [c for c in cells[1:] if c and c != cells[0]]
            if values:
                bucket = parts.setdefault(_field_key(cells[0]), [])
                if values[-1] not in bucket:        # the same (also multi-line) value twice is kept once (r2 I4, r3 I1)
                    bucket.append(values[-1])
        cells_map = {k: "\n".join(v) for k, v in parts.items()}
        rid = cells_map.get("id", "")
        if not SYSTEM_ID.fullmatch(rid):
            continue
        if rid in out:
            duplicates += 1
            continue
        out[rid] = {"doc": label, "fields": {_FIELD_BY_KEY[k]: v for k, v in cells_map.items() if k in _FIELD_BY_KEY},
                    # (R45 review W1) not read for facts — a ``Range`` value's subject is often the block's name
                    "name": cells_map.get("name", "")}
    return out, duplicates


def load_system_requirements(documents: list[tuple[str, str | None]]) -> tuple[dict[str, dict[str, Any]], list]:
    """Parse the given ``(label, path)`` system documents (``("SyRS", …), ("SyDS", …)``) into one ID map — the first
    document naming an ID wins. Returns the map and one record per document for the quality report: its blocks, the
    IDs another document already had, duplicate tables — or the error that kept it out (never raised: the STS is still
    generated, and the report says the trace was not read)."""
    system: dict[str, dict[str, Any]] = {}
    records: list[dict[str, Any]] = []
    for label, path in documents:
        if not path:
            continue
        name = re.split(r"[\\/]", str(path))[-1]
        try:
            blocks, duplicates = parse_system_requirement_docx(str(path), label)
        except Exception as exc:  # noqa: BLE001 — an unreadable optional input is disclosed, not fatal
            records.append({"doc": label, "file": name, "error": f"{type(exc).__name__}: {exc}"[:200]})
            continue
        shadowed = [k for k in blocks if k in system]
        for k, v in blocks.items():
            system.setdefault(k, v)
        records.append({"doc": label, "file": name, "blocks": len(blocks), "ids_in_earlier_document": len(shadowed),
                        "duplicate_tables": duplicates})
    return system, records


def cited_system_ids(req: dict[str, Any]) -> list[str]:
    """The system IDs the SRS block's Related ID names, in order, once each."""
    return list(dict.fromkeys(SYSTEM_ID.findall(str(req.get("related_id") or ""))))


def traced_system_facts(req: dict[str, Any], system: dict[str, dict[str, Any]],
                        stats: Counter | None = None) -> list[tuple[dict, dict, dict]]:
    """(line, fact, source) of every text field of each system block the requirement cites. Each field is read as
    its own text (a heading in one never makes the next an outcome), under the SRS requirement's ID. (R47) The input
    ranges come after every other traced fact, so a TC numbered before R47 keeps its number."""
    stats = stats if stats is not None else Counter()
    out, inputs = [], []
    for sid in cited_system_ids(req):
        stats["traced:cited"] += 1
        block = system.get(sid)
        if block is None:
            stats["traced:cited_not_in_documents"] += 1
            continue
        for field in SYSTEM_TEXT_FIELDS:
            text = block["fields"].get(field)
            if not text:
                continue
            source = {"doc": block["doc"], "id": sid, "field": field}
            (inputs if field in SYSTEM_INPUT_FIELDS else out).extend(
                (line, fact, source) for line in _field_lines(str(req.get("id", "")), field, text)
                for fact in line["facts"])
    return out + inputs


def _field_lines(anchor: str, field: str, text: str) -> list[dict[str, Any]]:
    """The fact lines of one system block field, read under the SRS ID ``anchor`` (`extract` opens a block at an SRS
    ID line — the ID names nothing else here). An outcome field is read under an ``<Output>`` heading."""
    return [line for b in extract(f"{_extract_head(anchor, field)}{text}\n") for line in b["lines"]]


_TYPED = re.compile(r"^([us])(8|16|32)[a-z]*_", re.IGNORECASE)


def subject_type_bounds(fact: dict[str, Any]) -> tuple[int, int] | None:
    """(review r4 I-2) ``u8g_…`` / ``s16s_…``: the declared width in the name bounds a unit-less raw value; a value with
    a unit (``8.5V``) is a physical quantity whose scaling the name does not give — no bound."""
    name = subject_of(fact) or ""
    m = _TYPED.match(name)
    if not m or fact.get("unit") or fact["kind"] != "threshold":
        return None
    bits = int(m.group(2))
    return (0, 2 ** bits - 1) if m.group(1).lower() == "u" else (-(2 ** (bits - 1)), 2 ** (bits - 1) - 1)


def _clip(text: str, n: int) -> str:
    return text if len(text) <= n else text[:n - 1] + "…"   # a cut quote says so (review r4 I-3)


# (review r2 W-4) an obligation ends its sentence: ``0.3mA 이하여야 한다`` / ``…이하로 유지한다`` — never a condition that
#   goes on (``9V 이상이어야 모터가 동작한다`` / ``500ms 이상 유지되어야 고장으로 판정한다`` / ``…유지시``)
_MUST = r"\s*(?:한다|함|합니다)"
_OBLIGATION = re.compile(r"\s*\)?\s*(?:로|으로|를|을|가|이)?\s*(?:유지(?:한다|합니다|해야" + _MUST + r"|하여야" + _MUST
                         + r"|되어야" + _MUST + r")|(?:이)?어야" + _MUST + r"|여야" + _MUST + r"|되어야" + _MUST
                         + r"|하여야" + _MUST + r")\s*(?:[.。]|$|\()")
# a subject that ends in a separator or a postposition is a misread, not a quantity (review W5: ``미만/``, ``위치에
#   따라``, ``내``), and so is one starting with a particle (``와 통신`` from ``Block 와의 통신``). (review r2 W-4) not a
#   comparison word at the end — ``전압 이상 :`` is "voltage abnormality"
_UNCLEAR_SUBJECT = re.compile(r"[/·,:;~\-]$|(?:^|\s)(?:따라|의해|대해|위해|내|중|후|시|때)$|^(?:와|과|의|을|를|에|로)\s")
# an object and its verb read *after* the value (``10ms 이상의 신호를 주입``, KJPDS02_PV SySM_06) — only there: a label
#   before the value may hold one (``도어를 여는 속도 :``, review r2 W-4)
_VERB_PHRASE = re.compile(r"\S(?:을|를)\s+\S")


# (R45) the comparison that holds exactly where the other does not, at the same value
_COMPLEMENT = {">=": "<", "<": ">=", ">": "<=", "<=": ">"}


def _condition_fact(f: dict) -> bool:
    return f["kind"] == "threshold" and f.get("op") in _COMPLEMENT and not f.get("negated") \
        and isinstance(f.get("value"), (int, float)) and not isinstance(f.get("value"), bool)


def _partition_partners(line: dict, fact: dict) -> list[dict]:
    """(R45) ``Manual Assist조건(0.8m/s 미만) 또는 Tip-To-Run 조건(0.8m/s 이상 1.3m/s 이하)``: two names, one unit, joined
    by *or*, that **split the unit at one value** — one of them ``< v`` (``<= v``), the other ``>= v`` (``> v``), one name
    a parenthesis owner — are one quantity in two regions, not two quantities to hold apart. R32 review C1 held both back
    because a step on one told the tester to hold the other false where it holds; with the split written that way, each
    point's verdict of the other name is known, and the step writes it. The other names' facts of that unit on the line
    are returned (empty: not such a split)."""
    unit = fact.get("unit") or ""
    name = subject_of(fact)
    if not unit or not name or not _condition_fact(fact):
        return []
    in_unit = [f for f in line["facts"] if _condition_fact(f) and (f.get("unit") or "") == unit and subject_of(f)]
    mine = [f for f in in_unit if subject_of(f) == name]
    others = [f for f in in_unit if subject_of(f) != name]
    split = {subject_of(g) for f in mine for g in others
             if g["op"] == _COMPLEMENT[f["op"]] and exact_value(g) == exact_value(f)
             and "paren_owner" in (f.get("signal_kind"), g.get("signal_kind")) and joint(line, f, g) == "or"}
    partners = [g for g in others if subject_of(g) in split]
    heads: dict[str, dict] = {}
    for g in partners:
        heads.setdefault(str(subject_of(g)), g)
    # (R45 review W4) a region whose own conditions do not join as written (``B조건(0.8m/s 이상, 2.0m/s 이상)``) has no
    #   verdict at a point: no split — back to the R32 C1 hold
    if any(line_holds(exact_value(g), g, line) is None for g in heads.values()):
        return []
    return partners


def _label_shares_a_quantity(line: dict, fact: dict, partners: list[dict] | None = None) -> bool:
    """(R32 review C1 · r2 W2) ``Manual Assist조건(0.8m/s 미만) 또는 Tip-To-Run 조건(0.8m/s 이상 …)`` / ``도어 속도가
    0.8m/s 미만이거나 Tip-To-Run 조건(0.8m/s 이상)``: a label before a parenthesis names a condition, not a quantity —
    with another condition (a threshold or a range, not negated, not an obligation) in the same non-empty unit under
    another name, joined to it by a **stated** *or* / *and*, the two may be one quantity (the door speed). Then the
    stated join would make the step hold the other "false" / "true" at a point where it is not the step's to say (a step
    said 불성립 at 0.8 m/s and told the tester to hold Tip-To-Run false there). An unstated join already tells the tester
    to check it and stays. (R45) A partner of a split at one value (`_partition_partners`) is no longer a doubt."""
    unit = fact.get("unit") or ""
    if not unit:
        return False
    raw = line.get("raw") or line["text"]
    for other in line["facts"]:
        if other is fact or any(other is p for p in partners or ()) \
                or other["kind"] not in {"threshold", "range"} or other.get("status") != "parsed" \
                or other.get("negated") or (other.get("unit") or "") != unit \
                or "paren_owner" not in (fact.get("signal_kind"), other.get("signal_kind")) \
                or (subject_of(other) if other["kind"] == "threshold" else other.get("signal")) == subject_of(fact) \
                or _OBLIGATION.match(raw[other["line_span"][1]:]):
            continue
        if joint(line, fact, other) in {"or", "and"}:
            return True
    return False


def _skip_reason(line: dict, fact: dict) -> str:
    if fact["status"] != "parsed":
        return str(fact.get("reason") or "not_parsed")
    if fact["kind"] not in {"threshold", "duration"}:
        return "kind_" + fact["kind"]
    if fact.get("op") not in _ORDER or not isinstance(fact.get("value"), (int, float)) \
            or isinstance(fact.get("value"), bool):
        return "not_an_order_comparison"
    if fact.get("role") == "response_constraint":
        return "response_constraint"
    if fact.get("negated"):
        return "negated_condition"
    if is_outcome_section(line.get("section") or ""):
        return "outcome_section"            # ``<완료조건>`` / ``<Output>``: produced, not stimulated (review r3 W6)
    if fact.get("signal_kind") == "reference_label":
        return "reference_label"            # ``기준 전압: 8.50V 이하`` names the threshold, not what to set (r3 W1)
    if fact["kind"] == "threshold" and fact.get("signal_kind") == "parameter" and not fact.get("monitored"):
        return "monitored_quantity_unknown"
    if fact["kind"] == "threshold" and _label_shares_a_quantity(line, fact, _partition_partners(line, fact)):
        return "parenthesis_labels_may_share_a_quantity"
    if _OBLIGATION.match((line.get("raw") or line["text"])[fact["line_span"][1]:]):
        # (R29 review C2) ``암전류는 0.3mA 이하여야 한다`` / ``최소 전류는 0.3mA 이하로 유지한다``: what the system must
        #   produce, not a condition to set — ``300ms 이상 유지시`` (a condition) is not matched
        return "output_requirement"
    named = subject_of(fact) if fact["kind"] == "threshold" else \
        (fact.get("signal") if fact.get("signal_kind") not in {"parameter", "identifier"} else None)
    if _UNCLEAR_SUBJECT.search(str(named or "").strip()) or \
            (fact.get("signal_kind") == "phrase_after" and _VERB_PHRASE.search(str(named or ""))):
        return "subject_unclear"            # ``4.85V 미만/5.15V 초과`` read ``미만/`` as the subject (review W5)
    if line_holds(exact_value(fact), fact, line) is None:
        return "same_subject_combination_unstated"
    bounds = subject_type_bounds(fact)
    if bounds and not bounds[0] <= exact_value(fact) <= bounds[1]:
        return "value_outside_subject_type"   # ``u8g_X 45000 이상``: no stimulus can reach it (review r4 I-2)
    return ""


def _not_a_condition(line: dict, fact: dict) -> bool:
    """A deadline (``100ms 이내``), an outcome line (``<Output>`` / ``Output :`` / a system Action) or an obligation
    (``…이하여야 한다``) — what the requirement produces or bounds, never a stimulus (`_conflict_role`)."""
    return fact.get("role") == "response_constraint" or _conflict_role(line, fact, False) == "outcome"


def _review_reason(line: dict, fact: dict, why: str, source: dict | None) -> str:
    """(R45) Why a held-back fact is a review item — ``""`` when it is not one. (review W2) A response constraint, an
    outcome line or an obligation is no condition to decide (a subject would not make it a step) — except (r2 W-R2-2)
    an ``이내`` inside a conditional clause (``… 이내에 … 되지 않으면``), which may be the condition's time window."""
    raw = line.get("raw") or line["text"]
    if fact.get("role") == "response_constraint" and _conflict_role(line, fact, False) != "outcome" \
            and _WINDOW_CLAUSE.search(raw[fact["line_span"][1]:fact["line_span"][1] + 60]):
        return "deadline_or_window"
    if why not in REVIEW_REASONS or _not_a_condition(line, fact):
        return ""
    if why == "no_subject_for_value" and source and source["field"] in _VALUE_FIELDS:
        return "no_subject_in_value_field"
    return why


def _review_item(req: dict, line: dict, fact: dict, source: dict | None, reason: str, kind: str,
                 block_name: str = "", criterion: bool = False) -> dict[str, Any]:
    item = {"srs_id": req.get("id", ""), "source": source_label(source), "reason": reason, "kind": kind,
            "line": line["text"], "line_sha256": line["sha256"], "fact": str(fact.get("raw") or ""),
            "fact_kind": fact.get("kind", ""), "line_span": list(fact.get("line_span") or []),
            "block_name": block_name, **_numeric_identity(fact)}
    if kind == "held_back" and criterion:
        # (R51 review r2 W-F) a criterion line (measured, not set): written out, it is a criterion again — never a step
        item["if_filled"] = ("판정 기준(측정해 확인하는 상·하한) 줄이다 — 원문대로 고쳐 적으면 '측정' 스텝의 판정 기준이 되고, "
                             "입력으로 설정하는 경계 스텝은 만들지 않는다")
        item["fill_check"] = "not_applicable"
    elif kind == "held_back":
        # (R46 review W1) what filling would do, tried in this very sentence
        try:
            item["if_filled"], item["fill_check"] = fill_guidance(line, fact, reason)
        except Exception as exc:  # noqa: BLE001 — guidance never costs the review item or the boundary TCs
            item["if_filled"], item["fill_check"] = f"— (채우기 안내를 만들지 못함: {type(exc).__name__})", "unchecked"
    return item


def _numeric_identity(fact: dict) -> dict[str, Any]:
    """(R46) What a review item's value is, exactly and JSON-safe — ``{"unit", "op", "value"}`` for a comparison,
    ``{"unit", "lo", "hi"}`` for a range (``0x1000 ~ 0x2000`` → ``4096`` / ``8192``); ``{}`` for anything else. The
    evidence search matches other sentences on it."""
    unit = fact.get("unit") or ""
    value = fact.get("value")
    if fact.get("kind") == "range" and isinstance(value, list) and len(value) == 2:
        texts = fact.get("value_text") or [None, None]
        try:
            lo, hi = (exact_value({"value_text": t, "value": v}) for t, v in zip(texts, value, strict=True))
        except (ArithmeticError, ValueError, TypeError):
            return {}
        return {"unit": unit, "lo": str(lo), "hi": str(hi)}
    if fact.get("kind") in {"threshold", "duration"} and isinstance(value, (int, float)) and not isinstance(value, bool):
        return {"unit": unit, "op": fact.get("op"), "value": str(exact_value(fact))}
    return {}


def _subject_text(fact: dict) -> str:
    if fact["kind"] == "duration":
        named = fact.get("signal") if fact.get("signal_kind") not in {"parameter", "identifier"} else None
        return (named or "조건") + DURATION_SUFFIX
    return subject_of(fact) or ""


def _condition_text(line: dict, fact: dict) -> str:
    def one(f):
        return f"{f.get('value_text') or f['value']}{f.get('unit') or ''} {_OP_TEXT[f['op']]}"
    units, join = condition_units(line, fact)
    if any(len(u) > 1 for u in units):
        # (R45 review W3) a range read from meaning is one unit, as `line_holds` judges it: ``4V 이하 또는 (5V 이상
        #   그리고 6V 이하)`` — in written order
        joiner = " 또는 " if join == "or" else " 그리고 "
        return _subject_text(fact) + " " + joiner.join(
            ("(" if len(units) > 1 and len(u) > 1 else "") + " 그리고 ".join(one(f) for f in u)
            + (")" if len(units) > 1 and len(u) > 1 else "") for u in units)
    parts = [fact] + [f for f in line["facts"] if same_subject(fact, f)]
    joins = {joint(line, fact, f) for f in parts[1:]}
    joiner = " 또는 " if joins == {"or"} else " 그리고 "
    return _subject_text(fact) + " " + joiner.join(one(f) for f in parts)


def _fact_text(f: dict) -> str:
    subject = _subject_text(f) if f["kind"] in {"threshold", "duration"} else (f.get("signal") or "")
    if f["kind"] in {"threshold", "duration"} and f.get("op") in _OP_TEXT:
        ref = f" (기준 {f['signal']})" if f.get("signal_kind") == "parameter" else ""
        return f"{subject} {f.get('value_text') or f['value']}{f.get('unit') or ''} {_OP_TEXT[f['op']]}{ref}".strip()
    return str(f.get("raw", "")).strip()


def _other_conditions(line: dict, fact: dict, joined_prev: str = "") -> str:
    """How to hold the line's conditions on other subjects while this one is stepped — one entry per subject (its
    conditions joined as the line joins them, review r3 W2), plus a join with a neighbouring line (r3 W6)."""
    notes, done = [], []
    partners = _partition_partners(line, fact)
    siblings = [g for g in line["facts"] if same_subject(fact, g)]
    if any(join_is_inferred(line, a, b) for a in [fact] + siblings for b in [fact] + siblings if a is not b):
        # (R45) the words give no connective between the subject's bounds; the one meaningful reading was taken
        notes.append(f"같은 주어의 결합 [{_condition_text(line, fact)}]: 원문에 연결어 없음 — 항상 참·항상 거짓이 아닌 "
                     "유일한 읽기로 판정")
    for f in line["facts"]:
        if f is fact or same_subject(fact, f) or f["kind"] not in {"threshold", "duration", "symbolic", "range"}:
            continue
        if any(same_subject(f, g) for g in done):
            continue
        done.append(f)
        group = [f] + [g for g in line["facts"] if g is not fact and same_subject(f, g)]
        inner = {joint(line, f, g) for g in group[1:]}
        text = (" 또는 " if inner == {"or"} else " 그리고 ").join(_fact_text(g) for g in group)
        if any(f is p for p in partners):
            # (R45) one quantity in two regions: the other name is neither held true nor false — its verdict at each
            #   point is written in the step (R32 review C1 said "불성립 상태로 둔다" where it held)
            notes.append(f"같은 양을 나눈 조건 [{_condition_text(line, f)}]: 원문의 두 이름이 같은 단위를 한 값에서 "
                         "나눈다 — 한 양의 두 구간으로 읽어 각 점에서의 판정을 스텝 기대 결과에 적는다")
            continue
        how = joint(line, fact, f)
        if f.get("status") != "parsed":
            # (R29 review r2 W-1) a condition without a subject cannot be held "true" or "false" by the words between:
            #   ``전원 8.5V 이하 500ms 미만 유지 후 9V 이상으로 복귀`` — `joint` joins a hold time with anything by *and*,
            #   which asserted both ``8.5V 이하`` and ``9V 이상`` hold. Say the join is for the tester to settle
            how = "unstated"
        state = {"or": "불성립 상태로 둔다", "and": "성립 상태로 둔다"}.get(how, "결합이 원문에 명시되지 않음 — 결합 조건 확인 필요")
        notes.append(f"다른 조건 [{text}]: {state}")
    hold = {"or": "불성립 상태로 둔다", "and": "성립 상태로 둔다"}
    if not notes and len([f for f in line["facts"] if f["kind"] in {"threshold", "symbolic", "range", "duration"}]) == 1:
        # (review r4 W-1) a non-numeric condition in the same clause (``… 이상이고 MotorDirection 이 Close``) is not a
        #   fact, but the words say how it joins: hold it accordingly — or say the join is unclear (review r5 W-B)
        how = condition_part_join(line, fact)
        if how:
            notes.append("같은 줄의 다른 조건(원문): "
                         + hold.get(how, "결합이 원문에 명시되지 않음 — 결합 조건 확인 필요"))
    if line.get("continues"):
        quoted = f" [{_clip(line['next_text'], 80)}]" if line.get("next_text") else ""
        notes.append(f"다음 줄 조건{quoted}과 {line['continues'].upper()} 결합(원문): {hold[line['continues']]}")
    if joined_prev:
        # every line of the run joined up to here (``A AND`` / ``B AND`` / ``C``), not only the last (review r5 W-A)
        quoted = f" [{_clip(line.get('previous_text') or '', 300)}]" if line.get("previous_text") else ""
        notes.append(f"앞 줄 조건{quoted}과 {joined_prev.upper()} 결합(원문): {hold[joined_prev]}")
    return "; ".join(notes)


def _dedupe_text(text: str) -> str:
    """``2) X``, ``- X``, ``① X`` are the same sentence (review r3 I1)."""
    return re.sub(r"^[\s\-·•*]*(?:\(?\d+[.)]|[①-⑳])?\s*", "", text).strip()


def _fact_key(fact: dict) -> tuple:
    return (fact["kind"], subject_of(fact), fact.get("op"), str(fact.get("value_text")), fact.get("unit"))


def boundary_steps(req: dict[str, Any], stats: Counter | None = None,
                   system: dict[str, dict[str, Any]] | None = None,
                   review: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """One group of steps per fact used: ``[{"steps": [...], "evidence": {...}}]``. With ``system`` (parsed SyRS /
    SyDS blocks), the facts of the system blocks the requirement cites follow its own (``evidence["source"]`` names
    the block; counters prefixed ``traced:``). (R45) ``review`` receives every fact held back for a reason a reader can
    settle (`REVIEW_REASONS` — a missing subject, an unwritten join, a range's inclusion …): the requirement ID, the
    source, the sentence, the fact and the reason — shown in the document and on the web, never filled by a guess."""
    stats = stats if stats is not None else Counter()
    groups: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    stepped: set[tuple] = set()
    own = _requirement_facts_by_field(req)
    items = [(line, fact, None) for line, fact, _key in own]
    measuring = measurement_spec(req)                           # (R51 review W1)
    verification_lines = {id(line) for line, _fact, key in own if key == "verification"}     # (R51)
    if system is not None:
        items += traced_system_facts(req, system, stats)
    accept: dict[tuple, list[tuple]] = {}
    for line, fact, source in items:
        prefix = "traced:" if source else ""
        stats[prefix + "facts"] += 1
        why = _skip_reason(line, fact)
        key = (_dedupe_text(line["text"]), *_fact_key(fact))
        if not why and key in seen:
            why = "duplicate_fact"   # the description and the verification criteria repeat the sentence (review W6)
        # (review r2 W-3) only a fact that names its subject is "the same" as another — two unnamed hold times of two
        #   conditions are not one quantity (`same_subject`, review r4 C-2): SySM_04's over-voltage hold was dropped
        if not why and source and subject_of(fact) is not None and _fact_key(fact) in stepped:
            why = "already_stepped"  # the SRS (or another cited block) already steps this subject at this value
        conversion = _CONVERSION.search(line.get("raw") or line["text"]) if not source and measuring else None
        if conversion and conversion.group(0).strip() == measuring["lines"][2] and fact.get("line_span") and \
                conversion.start() <= fact["line_span"][0] < conversion.end():
            # (review r2 W-E) a fact of that sentence outside the conversion is read as any other
            # (R51 review W1) ``(0x0~0xFF, offset : 4V, …, 4V ~ 16.75V)`` is the measurement TC's conversion: its ranges
            #   are no subjectless review items (writing a subject in, as they advised, broke the measurement spec)
            why = "measurement_conversion"
        in_verification = source["field"] == "Verification criteria" if source else id(line) in verification_lines
        if why in _CRITERION_MAY_SKIP and in_verification and _acceptance_line(line, req):
            # (R51) a pass criterion: a '측정' step in the requirement's criteria TC (`acceptance_groups`) — no points
            field_text = req.get("verification") if not source else \
                (((system or {}).get(source["id"]) or {}).get("fields") or {}).get(source["field"])
            # (review r3 W-2 · r4 W-A) one step per line object, headed by the heading above its own place
            at = line["span"][0] - len(_extract_head(str(req.get("id", "")), source["field"] if source else None))
            accept.setdefault((source_label(source), id(line)), []).append(
                (line, fact, source, _criterion_heading(str(field_text or ""), at)))
            why = "acceptance_criterion"
        if why:
            stats[prefix + "skipped:" + why] += 1
            reason = _review_reason(line, fact, why, source) if review is not None else ""
            if reason:
                block_name = str(((system or {}).get(source["id"]) or {}).get("name") or "") \
                    if reason == "no_subject_in_value_field" and source else ""
                review.append(_review_item(req, line, fact, source, reason, "held_back", block_name,
                                           criterion=in_verification and _acceptance_line(line, req)))
            continue
        seen.add(key)
        stepped.add(_fact_key(fact))
        value, delta = exact_value(fact), written_step(fact)
        unit = fact.get("unit") or ""
        reference = ""
        if fact.get("signal_kind") == "parameter":
            reference = f" (기준 {fact['signal']} = {fact.get('value_text') or fact['value']}{unit})"
        predicate = f" ({fact['predicate']} 상태)" if fact["kind"] == "duration" and fact.get("predicate") else ""
        condition = _condition_text(line, fact)
        others = _other_conditions(line, fact, line.get("joins_previous") or "")
        # (R29) a traced sentence names where it comes from — after the basis mark, so it is never read as a stimulus
        origin = f" [{source_label(source)} — {req.get('id', '')} Related ID]" if source else ""
        whose = f"시스템 요구 {source['id']} 의 " if source else ""
        steps, used = [], []
        bounds = subject_type_bounds(fact)
        input_cell = bool(source) and source["field"] in SYSTEM_INPUT_FIELDS   # (R47)
        # (R45) the other region(s) of one quantity split at a value: one fact per name — its verdict is the line's
        heads: dict[str, dict] = {}
        for g in _partition_partners(line, fact):
            heads.setdefault(str(subject_of(g)), g)  # the name's first fact: its conditions read in written order
        partner_heads = list(heads.values())
        if review is not None:
            # (R45 review I3) what was read, not written — stepped, and listed for a reader to confirm
            if partner_heads:
                review.append(_review_item(req, line, fact, source, "read_as_one_quantity", "read"))
            if any(join_is_inferred(line, fact, g) for g in line["facts"] if same_subject(fact, g)):
                review.append(_review_item(req, line, fact, source, "read_as_range", "read"))
        for p in (value - delta, value, value + delta):
            if bounds and not bounds[0] <= p <= bounds[1]:
                stats[prefix + "points_outside_subject_type"] += 1   # ``u16g_X 0 초과`` has no −1 (review r4 I-2)
                continue
            verdict = line_holds(p, fact, line)
            action = (f"{ACTION_PREFIX}{_subject_text(fact)} = {_fmt(p, fact)}{unit}{predicate}{reference}"
                      f"{BASIS_MARK}{fact['raw']}{origin}")
            # (review r3 C-1) the verdict is the sentence's, not the whole requirement's: outside its condition the
            #   sentence claims nothing (another sentence may describe what happens there)
            partner_verdicts = [(g, line_holds(p, g, line)) for g in partner_heads]
            other = " · ".join(f"[{_condition_text(line, g)}] {_VERDICT_TEXT[v]}" for g, v in partner_verdicts)
            if input_cell:
                # (R47 review W1 · W4) an input cell states where the input is, not what the element does: a normal range
                #   (SySM_04 ``BAT 전압(9V 이상 16.0V 이하)`` — inside it a power monitor reports *no* fault) and a
                #   trigger (SyFN_10 ``5000RPM 초과``) read alike, so neither side claims a reaction — the requirement's
                #   other sentences decide it. First, also over a split of one quantity (which would claim an action)
                expected = (f"조건 [{condition}] {'성립' if verdict else '불성립'} → {whose}입력 칸이 적은 범위 "
                            f"{'안' if verdict else '밖'}" + (f"(같은 양을 나눈 조건 {other})" if partner_heads else "")
                            + f" — 요소의 반응(이상 판정·리셋 등)은 같은 요구의 다른 문장 판정을 따른다: "
                            f"{_clip(line['text'], 160)}")
            elif partner_heads:
                # (R45) the stepped condition's verdict stays first (``조건 […] 성립 →`` is what a reader — and the
                #   cross scorer — takes as the step's verdict); the other region's follows, at the same point
                if verdict:
                    expected = (f"조건 [{condition}] 성립 → {whose}이 문장이 이 조건에 기술한 동작 수행 확인"
                                f"(같은 양을 나눈 조건 {other}): {_clip(line['text'], 160)}")
                elif any(v for _g, v in partner_verdicts):
                    expected = (f"조건 [{condition}] 불성립 → {whose}이 조건의 구간 아님(같은 양을 나눈 조건 {other} — "
                                f"그 조건의 판정을 따른다): {_clip(line['text'], 160)}")
                else:
                    # outside every region the line names (``1.4m/s`` past Tip-To-Run's 1.3 and Manual Assist's 0.8):
                    #   the sentence claims nothing there — or, when a region's verdict is unknown, says so
                    expected = (f"조건 [{condition}] 불성립 → {whose}이 문장의 동작 대상 아님(같은 양을 나눈 조건 {other} — "
                                f"같은 요구의 다른 문장 판정을 따른다): {_clip(line['text'], 160)}")
            else:
                expected = (f"조건 [{condition}] 성립 → {whose}이 문장이 기술한 동작 수행 확인: {_clip(line['text'], 160)}"
                            if verdict else
                            f"조건 [{condition}] 불성립 → {whose}이 문장의 동작 대상 아님(같은 요구의 다른 문장 판정을 "
                            f"따른다): {_clip(line['text'], 160)}")
            steps.append({"action": action, "expected": expected})
            used.append({"point": _fmt(p, fact), "holds": verdict})
        stats[prefix + "facts_used"] += 1
        stats[prefix + "steps"] += len(steps)
        if input_cell:
            # (R47 review I4) the disclosure tells the input ranges apart — the traced counts grew with them
            stats["traced:input_range_facts_used"] += 1
            stats["traced:input_range_steps"] += len(steps)
        groups.append({"steps": steps, "evidence": {
            "srs_id": req.get("id", ""), "line": line["text"], "line_sha256": line["sha256"], "span": fact["span"],
            "kind": fact["kind"], "signal": _subject_text(fact),
            "reference_constant": fact["signal"] if reference else None,
            "signal_kind": fact.get("signal_kind", ""), "op": fact["op"],
            "value": fact.get("value_text") or fact["value"], "unit": unit, "step": format(delta, "f"),
            "step_basis": "written_precision", "combination_note": others, "points": used,
            "source": dict(source) if source else None}})
    groups += acceptance_groups(req, accept, stats)
    return groups


def acceptance_groups(req: dict[str, Any], accept: dict[tuple, list], stats: Counter) -> list[dict[str, Any]]:
    """(R51) One group per pass-criterion line of the requirement (`_acceptance_line`, review r4 I-3): a '측정' step —
    measure what the line names and judge it against the limits as written (inclusion kept: ``이하`` includes the
    limit). No boundary points: a measured quantity is not set, so no discrimination is claimed."""
    out = []
    for facts in accept.values():
        line, fact, source, heading = facts[0]
        text = line["text"]
        # (R51 review W5) the clause as written — the bare comparisons dropped what they apply to (``SCM : 90%
        #   이하`` for ``BCM, SMK, IBU, SJB, PSM, SCM``); (r2 W-B) under the heading it sits under — what is measured
        clause = _clip(_LEAD_BULLET.sub("", text.strip()).strip(), 200)
        said = f"{heading} — {clause}" if heading else clause
        stats["acceptance_criteria"] += len(facts)
        stats["acceptance_steps"] += 1
        out.append({"steps": [{"action": f"{ACCEPT_PREFIX}{_clip(said, 200)}",
                               "expected": f"측정값이 원문 판정 기준을 만족 — '{_clip(said, 220)}'(대상·경계 포함은 원문대로, "
                                           "입력으로 설정하지 않고 측정해 확인)"}],
                    "evidence": {
            "srs_id": req.get("id", ""), "line": line["text"], "line_sha256": line["sha256"], "span": fact["span"],
            "kind": "acceptance", "signal": heading or "—", "reference_constant": None, "signal_kind": "acceptance",
            "op": "judged", "value": clause, "unit": "", "step": "—", "step_basis": "pass criterion (measured)",
            "combination_note": "", "points": [], "source": dict(source) if source else None, "hw_tolerance": []}})
    return out


# (R33) ``이상`` ↔ ``초과`` and ``이하`` ↔ ``미만``: the same value with the other boundary inclusion
_INCLUSION_FLIP = {">=": ">", ">": ">=", "<=": "<", "<": "<="}
MAX_INCLUSION_CONFLICTS = 30     # items kept for the report (all are counted)


# (R33 review W3) a hold time whose words state no inclusion (``500ms 동안 유지``) got ``>=`` from the extractor: it
#   wrote none, so it disagrees with nothing
_WRITTEN_INCLUSION = re.compile(r"이상|초과|이하|미만|이내|[<>]")


def _comparable_fact(fact: dict) -> bool:
    """(R48 review I6) An order comparison of a number, not negated, not a response time — what two documents can state
    two ways (R33 inclusion conflicts, R48 value differences): one predicate, so the two never drift apart."""
    return fact["kind"] in {"threshold", "duration"} and fact.get("op") in _INCLUSION_FLIP \
        and isinstance(fact.get("value"), (int, float)) and not isinstance(fact.get("value"), bool) \
        and not fact.get("negated") and fact.get("role") != "response_constraint"


def _conflict_role(line: dict, fact: dict, verification: bool) -> str:
    """``condition`` — or ``outcome`` for what the requirement produces (an ``<Output>``/``Output :`` section, a system
    block's Action / System Behavior, an obligation ``…이하여야 한다``) — the split `_skip_reason` makes for stepping
    (R33 review W2: an input condition and an output value of the same number are not one statement) — or
    ``stimulus`` for anything else of a **verification criteria** field (the SRS ``verification`` text, a block's
    "Verification criteria"): it describes tests — ``Input : 입력전원 8.5V 이하`` / ``Precondition : …`` and the SRS's
    unlabelled ``2. 500ms 초과 저전압 상황 동작 확인`` alike — the range a test chooses its input from, not the
    requirement's condition (R33 review r2 W-A, r3 W-2)."""
    raw = line.get("raw") or line["text"]
    if is_outcome_section(line.get("section") or "") or _OBLIGATION.match(raw[fact["line_span"][1]:]):
        return "outcome"
    return "stimulus" if verification else "condition"


def _joins(ra: str, a: dict, rb: str, b: dict) -> bool:
    """Do two facts of these roles make a candidate? A condition with a condition, an outcome with an outcome — and a
    test input with a condition only when the **input** includes the value the condition excludes (``Input : 8.5V 이하``
    against ``8.5V 미만``: the test expects a reaction where the requirement gives none). A stricter input (``Input :
    4.85V 미만`` against ``4.85V 이하``) only leaves the value untested — no disagreement."""
    if ra == rb:
        return ra != "stimulus"
    if {ra, rb} == {"stimulus", "condition"}:
        stimulus = a if ra == "stimulus" else b
        return stimulus["op"] in {"<=", ">="}
    return False


def inclusion_conflicts(req: dict[str, Any], system: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    """(R33) Where the requirement's own text and a system block it cites (or two such blocks) state the **same value in
    the same unit with the other boundary inclusion**, in roles that can disagree (`_joins`: two conditions, two
    outcomes, or a test input that includes a value a condition excludes) — HDPDM01
    ``SwTSR_0104``: SRS ``u16g_ApiIn_Vsup < 8.5V`` against SyDS ``SySM_04`` ``입력전원 8.5V 이하``: 8.5 V is inside one
    document's condition and outside the other's. The link is explicit (the SRS Related ID); the pairing is by value
    and unit only, so ``same_subject`` is true only when both facts name the same subject (`subject_of`) — otherwise the
    reviewer first checks they are one quantity. Which document is right is not decided here; the boundary TCs judge
    each sentence as written. Not candidates: a negated condition, a response time (``500ms 이내``), a hold time that
    writes no inclusion (``500ms 동안 유지``). Pairs are formed within (unit, value) buckets.

    (R43) Two different lines of **one source** (the SRS text, or the fields of one block) are candidates too, marked
    ``within_source`` — HDPDM01 ``SwTSR_0104``'s own text: ``저 전압 고장 검출 기준 전압: 8.50V 이하`` against ``8.50V
    미만인 전압으로 … 유지하는 경우``. Within one source two sentences may also be two conditions or two actions on
    purpose (a warning at ``이하`` and a DTC at ``미만``; items 1·2·3 of one block with their own hold times), so the
    mark tells the reviewer to read both first."""
    items = []
    for key in ("description", "verification"):          # `requirement_facts`, keeping which field a fact is in
        for b in extract(f"ID\t{req.get('id', '')}\n{req.get(key) or ''}\n"):
            items += [(line, fact, None, key == "verification") for line in b["lines"] for fact in line["facts"]]
    items += [(line, fact, source, source.get("field") == "Verification criteria")
              for line, fact, source in traced_system_facts(req, system, Counter())]
    buckets: dict[tuple, list] = {}
    for line, fact, source, verification in items:
        if not _comparable_fact(fact) or not _WRITTEN_INCLUSION.search(str(fact.get("raw") or "")):
            continue            # a response time (``500ms 이내``) is measured, not a condition — not the same statement
        buckets.setdefault((fact.get("unit") or "", exact_value(fact)), []).append(
            (line, fact, source, _conflict_role(line, fact, verification)))
    out: list[dict[str, Any]] = []
    seen: dict[tuple, int] = {}   # key → index in ``out``
    side = _statement

    for (unit, value), group in buckets.items():
        for i, (la, a, sa, ra) in enumerate(group):
            for lb, b, sb, rb in group[i + 1:]:
                within = ((sa or {}).get("doc"), (sa or {}).get("id")) == ((sb or {}).get("doc"), (sb or {}).get("id"))
                if within and _dedupe_text(la["text"]) == _dedupe_text(lb["text"]):
                    # (R43) one line says one thing: a pair within one source is two of its lines — and ``1) X`` in
                    #   the description is the same sentence as ``- X`` in the verification (review I3)
                    continue
                if _INCLUSION_FLIP[a["op"]] != b["op"] or not _joins(ra, a, rb, b):
                    continue
                role = "stimulus" if "stimulus" in (ra, rb) else ra
                key = (unit, value, role, frozenset({((sa or {}).get("doc"), (sa or {}).get("id"), a["op"]),
                                                   ((sb or {}).get("doc"), (sb or {}).get("id"), b["op"])}))
                named = subject_of(a) is not None and subject_of(b) is not None
                same = bool(named and subject_of(a) == subject_of(b))
                if key in seen:
                    # (R43 review I2) one finding per key — but a pair naming one subject says more than the first seen
                    k = seen[key]
                    if same and not out[k]["same_subject"]:
                        out[k].update(same_subject=True, subject_missing=False, a=side(la, a, sa), b=side(lb, b, sb))
                    continue
                seen[key] = len(out)
                out.append({"srs_id": req.get("id", ""), "value": str(exact_value(a)), "unit": unit, "role": role,
                            "same_subject": same, "subject_missing": not named, "within_source": within,
                            "a": side(la, a, sa), "b": side(lb, b, sb)})
    return out


def _statement(line: dict, fact: dict, source: dict | None) -> dict[str, Any]:
    """One side of a finding (R33 conflict · R48 difference): where, the subject (a hold time's reference constant
    apart), the comparison as written and the line."""
    subject = _subject_text(fact) if fact.get("signal_kind") == "parameter" else subject_of(fact)
    return {"source": source_label(source), "subject": subject or "주어 없음", "op": fact["op"],
            "reference": fact["signal"] if fact.get("signal_kind") == "parameter" else None,
            "text": f"{fact.get('value_text') or fact['value']}{fact.get('unit') or ''} {_OP_TEXT[fact['op']]}",
            "line": _clip(line["text"], 160)}


# ── (R48) requirement-document value differences ─────────────────────────────────────────────────────────────────
# 2026-09-30 user direction: what the documents state two ways is shown, even without a basis to decide it. R33 pairs
#   the same value written with the other inclusion; this pairs a value one document states with **another value** on
#   the same side that the other document states — only where the two are likely one threshold (`_DIFF_TIERS`).
_SIDE = {">": "lower", ">=": "lower", "<": "upper", "<=": "upper"}
# the strongest first. (review W1 · W2) every tier but a hold time's is within `_DIFF_MAX_RATIO` of the SRS value:
#   ``배터리 전압 9V 이상 구동`` ↔ ``16V 초과 정지`` (one quantity, two thresholds) and ``100ms`` ↔ ``1초`` (one step of an
#   integer second) were paired without it
_DIFF_TIERS = ("same_subject", "same_condition", "within_step", "only_pair")
_DIFF_MAX_RATIO = Decimal("0.05")
MAX_VALUE_DIFFERENCES = 30      # items kept for the report (all are counted; the document lists all)
MAX_DIFFERENCE_COUNTERPARTS = 3


def _difference_fact(fact: dict) -> bool:
    """A comparison with a unit that can differ in value (`_comparable_fact` — the R33 filter's own — with a unit)."""
    return _comparable_fact(fact) and bool(fact.get("unit"))


def _near(a: Decimal, b: Decimal) -> bool:
    return abs(a - b) <= _DIFF_MAX_RATIO * max(abs(a), abs(b))


# a line that ends its sentence (or its clause) — the next line starts anew; (review r2 I2) so does a list item
_SENTENCE_END = re.compile(r"(?:[.。]|다|시|때|경우|면)\s*$")
_LIST_MARK = re.compile(r"^\s*(?:\(?\d+[.)]|[-·•*]|[①-⑳])")


def _held_conditions(line: dict, above: dict | None = None) -> list[tuple[str, str, Decimal, str | None]]:
    """The (unit, side, value, subject) of a line's conditions that are not times — what its hold times hold. A line
    with none of its own continues the line right above it when that one does not end its sentence and it opens no list
    item (SyRS ``SyTSR_0116``: ``… BAT 범위 8.5V미만의 BAT`` / ``전압으로 특정시간(300ms) 이상 유지하는 경우.``)."""
    def conditions(ln: dict) -> list[tuple[str, str, Decimal, str | None]]:
        return [(k[0], _SIDE[c["op"]], k[1], subject_of(c)) for c in ln["facts"]
                if c["kind"] == "threshold" and _difference_fact(c)
                and (k := _match_key(c["unit"], exact_value(c)))[0] != "s"]
    own = conditions(line)
    if own or above is None or above["span"][1] + 1 != line["span"][0] or _SENTENCE_END.search(above["text"]) \
            or _LIST_MARK.match(line["text"]):
        return own
    return conditions(above)


def _lines_above(rows: list[tuple]) -> dict[int, dict]:
    """``{id(line): the fact line before it in the same text}`` for (line, fact[, source]) rows in written order — one
    position space per text (the SRS description; each block field). `_held_conditions` checks it is the line right
    above (its span)."""
    by_text: dict[tuple, list[dict]] = {}
    for row in rows:
        source = (row[2] if len(row) > 2 else None) or {}
        lines = by_text.setdefault((source.get("doc"), source.get("id"), source.get("field")), [])
        if not lines or lines[-1] is not row[0]:
            lines.append(row[0])
    return {id(b): a for lines in by_text.values() for a, b in zip(lines, lines[1:], strict=False)}


def value_differences(req: dict[str, Any], system: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    """(R48) Where the requirement's own text states a condition value that **no** condition of the system blocks it
    cites states (in any comparison — a value both state is one statement or an R33 inclusion conflict), and a cited
    block states a condition on the **same side** (both lower or both upper bounds) in the same unit that is likely the
    same threshold written differently. Likely means, within `_DIFF_MAX_RATIO` of each other, one of:

    * ``same_subject`` — both name the same subject: KJPDS02 ``SwTR_0605`` SRS ``저전압 8.6V 이하`` ↔ SyRS ``SyTR_1305``
      ``저전압 8.5V 이하``;
    * ``within_step`` — within one written step of the coarser writing: HDPDM01 ``SwTSR_0102`` SRS
      ``u16g_ApiIn_HallSnsrLevel 5.14V 이상`` ↔ SyRS ``SyTSR_0114`` ``5.15V 초과``, ``SwTR_0605`` ``16.04V`` ↔ ``16.1V``;
    * ``only_pair`` — the requirement's text and its cited blocks each state exactly one value in that unit and side;

    and, for hold times at any distance, ``same_condition`` — the two lines hold conditions on the same side at near
    values (HDPDM01 ``SwTSR_0104`` SRS ``8.50V 미만 … 500ms 이상 유지`` ↔ SyRS ``SyTSR_0116`` ``BAT 8.5V미만 … 300ms 이상
    유지``: the SW decides later than the system requirement). The opposite side (an entry ``미만`` and its recovery
    ``이상``) is never paired; two thresholds of one quantity on the same side (a warning and a fault) may be — the
    guidance says to check it is one threshold first. Conditions only on both sides — the SRS description (its
    verification text describes tests), a cited block's condition text (not its Verification criteria, an outcome, an
    obligation or an input cell, the element's design domain — R47). Which value is right is not decided here; each
    sentence's boundary TCs judge it as written."""
    own_all = [(line, f) for b in extract(f"ID\t{req.get('id', '')}\n{req.get('description') or ''}\n")
               for line in b["lines"] for f in line["facts"]]
    cited_all = traced_system_facts(req, system, Counter())
    above = {**_lines_above(own_all), **_lines_above(cited_all)}
    own = [(line, f) for line, f in own_all if _difference_fact(f) and _conflict_role(line, f, False) == "condition"]
    cited = [(line, f, s) for line, f, s in cited_all
             if s["field"] not in SYSTEM_INPUT_FIELDS and _difference_fact(f)
             and _conflict_role(line, f, s["field"] == "Verification criteria") == "condition"]
    if not own or not cited:
        return []

    def key(f: dict) -> tuple[str, Decimal]:
        return _match_key(f["unit"], exact_value(f))

    def step(f: dict) -> Decimal:
        return _match_key(f["unit"], written_step(f))[1]

    stated = {key(f) for _line, f, _s in cited}
    own_values: dict[tuple, set] = {}
    cited_values: dict[tuple, set] = {}
    for _line, f in own:
        own_values.setdefault((key(f)[0], _SIDE[f["op"]]), set()).add(key(f)[1])
    for _line, f, _s in cited:
        cited_values.setdefault((key(f)[0], _SIDE[f["op"]]), set()).add(key(f)[1])
    out: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    for line, f in own:
        (unit, value), side = key(f), _SIDE[f["op"]]
        if (unit, value) in stated or (unit, value, f["op"], subject_of(f)) in seen:
            continue
        seen.add((unit, value, f["op"], subject_of(f)))
        hold = f["kind"] == "duration"
        only = len(own_values[(unit, side)]) == 1 and len(cited_values.get((unit, side), ())) == 1
        held = _held_conditions(line, above.get(id(line))) if hold else []
        others: dict[tuple, dict[str, Any]] = {}
        for cline, g, source in cited:
            g_unit, g_value = key(g)
            if g_unit != unit or _SIDE[g["op"]] != side:
                continue
            near = _near(value, g_value)
            held_pair = next(((hs, cs) for u, sd, v, hs in held
                              for cu, csd, cv, cs in _held_conditions(cline, above.get(id(cline)))
                              if u == cu and sd == csd and _near(v, cv)), None) \
                if hold and g["kind"] == "duration" else None
            same_condition = held_pair is not None
            tiers = [t for t, ok in (("same_subject", near and subject_of(f) is not None and subject_of(f) == subject_of(g)),
                                     ("same_condition", same_condition),
                                     ("within_step", near and abs(g_value - value) <= max(step(f), step(g))),
                                     ("only_pair", near and only)) if ok]
            okey = (source_label(source), g["op"], g_value, subject_of(g))
            if tiers and okey not in others:   # a block repeating the sentence is one counterpart
                # (a hold time) the SRS larger: it decides later than the system requirement allows; smaller: it
                #   may be the system time shared out (a decomposition) — the reviewer reads which
                others[okey] = {**_statement(cline, g, source), "tiers": tiers, "srs_is": "larger" if value > g_value
                                else "smaller", "_gap": abs(g_value - value), "_named": subject_of(g) is not None,
                                # (review r2 I1) the subjects of the two held conditions — different names are shown
                                "held_subjects": list(held_pair) if held_pair else None}
        if not others:
            continue
        ranked = sorted(others.values(), key=lambda o: (min(_DIFF_TIERS.index(t) for t in o["tiers"]), o["_gap"]))
        # (review I2) the flags describe the first counterpart — the one the row leads with
        named = subject_of(f) is not None and ranked[0]["_named"]
        for o in ranked:
            del o["_gap"], o["_named"]
        out.append({"srs_id": req.get("id", ""), "unit": f["unit"], "value": str(exact_value(f)), "side": side,
                    "time": hold,     # (review W4) a hold time — a period (R47) is a threshold like any other
                    "tiers": [t for t in _DIFF_TIERS if any(t in o["tiers"] for o in ranked)],
                    "subject_missing": not named, "srs": _statement(line, f, None),
                    "others": ranked[:MAX_DIFFERENCE_COUNTERPARTS], "other_count": len(ranked)})
    return out


# ── (R49) HW measurement tolerance — quoted next to the boundary steps it concerns ───────────────────────────────────
# A boundary step's spacing is the written precision (``8.50`` → 0.01 V); a HW requirement writes how precisely the
#   monitor path measures the quantity the SW compares (HDPDM01 ``HwTSR_0204`` Battery Voltage Monitor ``허용 오차:
#   ±3%``). Where the spacing is inside that tolerance, a HIL run cannot tell the boundary's inclusion: the points are
#   decided by the injection path (the SW variable directly), or checked for direction only outside the tolerance.
#   Linked only through the Related IDs — a HW block citing a system ID the SRS requirement cites; never by value.
HW_ID = re.compile(r"\bHw[A-Za-z]{1,6}_\d+(?:_\d+)?\b")
# ``허용 오차: ±3%`` / ``허용 오차: ± 1A`` / ``허용 오차: ±0.15V`` — the accuracy of a monitor path. A value's own spec
#   (``5V(±0.15)`` — what a HW output must produce) and a test's measuring error (``측정오차±0.1``) are not it
_TOLERANCE = re.compile(r"허용\s*오차\s*[:：]?\s*±\s*(?P<num>\d+(?:\.\d+)?)\s*(?P<unit>%|mV|V|mA|A|ms|s)?(?![A-Za-z])")
_HW_FIELDS = ("Description", "Range")
# what a block measures, as its text writes it: a range (``Battery Voltage Range: 9 ~16V`` · ``전류(0~20A)``) — its
#   unit is what a ``%`` applies to — and a nominal node value (``Monitor전압: 2.5V(±0.15V)`` · ``2.5(±0.15)V``)
_HW_RANGE = re.compile(r"(?P<lo>\d+(?:\.\d+)?)\s*~\s*(?P<hi>\d+(?:\.\d+)?)\s*(?P<unit>mV|V|mA|A|ms|s)(?![A-Za-z])")
_HW_NOMINAL = re.compile(r"(?P<v>\d+(?:\.\d+)?)\s*(?:(?P<u1>mV|V|mA|A)\s*\(\s*±|\(\s*±\s*\d+(?:\.\d+)?\s*\)\s*(?P<u2>mV|V|mA|A))")
_HW_SCALE_MARGIN = Decimal("0.2")     # a value within 20% of what the block writes is on the block's scale
# (R50) ``Sensor Power Monitor = Sensor Power *0.5`` · ``Battery Voltage = Monitor Voltage * 9`` — one side a monitor node.
#   (review I4) an offset after the factor (``* 0.2 + 0.1``) is no ratio: not read
_HW_FORMULA = re.compile(r"(?P<lhs>[A-Za-z가-힣][A-Za-z가-힣 ()]{1,60}?)\s*=\s*(?P<rhs>[A-Za-z가-힣][A-Za-z가-힣 ()]{1,60}?)"
                         r"\s*\*\s*(?P<k>\d+(?:\.\d+)?)(?![\d.])(?!\s*[A-Za-z]*\s*[+\-]\s*\d)(?!\s*/)(?![eE]\d)")
_MONITOR_WORD = re.compile(r"monitor|모니터", re.IGNORECASE)
# (review W3a / I4) a formula between two quantities (``전류(A) = 모니터 전압((V)*6``) or of a time (``Monitoring Period =
#   … * 10``) is no divider ratio
_FORMULA_UNIT = re.compile(r"\((m?[VA])\)")
_CURRENT_WORD = re.compile(r"전류|current", re.IGNORECASE)
_VOLTAGE_WORD = re.compile(r"전압|voltage", re.IGNORECASE)
_NOT_A_NODE = re.compile(r"period|주기|time|시간|freq|주파수", re.IGNORECASE)
# (review W3b) a block that writes a monitor node: a nominal with its ± next to a monitor word (``Monitor전압: 2.5V(±0.15V)``)
_MONITOR_NODE = re.compile(r"(?:monitor|모니터)[^\n]{0,30}?\d+(?:\.\d+)?\s*(?:mV|V|mA|A)?\s*\(\s*±", re.IGNORECASE)
# a HSIS row's SW variables and net names (``VCC_BAT`` · ``VCC_HALL_MON``); a layer prefix (``u16g_ApiIn_`` / ``u16g_DrvIn_``)
#   — (review I2) the direction stays: ``ApiIn_X`` is not ``ApiOut_X``
_SW_VAR = re.compile(r"\b[us]\d+[gs]?_[A-Za-z]+_[A-Za-z0-9_]+\b")
_SW_LAYER = re.compile(r"^[us]\d+[gs]?_(?:[A-Z][a-z]+)?(?=(?:In|Out)_)")
_NET = re.compile(r"\b[A-Z][A-Z0-9]*(?:[_-][A-Z0-9]+)+\b")
# (review W2) words that name no particular signal — what is left names what a block, a formula or a HSIS row is about
_GENERIC_WORDS = frozenset({
    "block", "블록", "전원", "monitor", "모니터", "monitoring", "power", "voltage", "전압", "signal", "신호", "output", "출력",
    "input", "입력", "interface", "level", "mon", "vcc", "api", "drv", "srv", "in", "out", "adc", "analog", "switch",
    "supply", "control", "controller", "drive", "range"})
_NAME_TOKEN = re.compile(r"[A-Z]?[a-z]+|[A-Z]+(?![a-z])|[가-힣]+")
MAX_HW_TOLERANCES = 3
# (R50 review W4) where a ratio came from, and (W2) what checked it against the fact's signal
_SCALE_SOURCE_TEXT = {"design": "HW 설계서 분압식", "own": "HW 요구 블록의 식"}
_SCALE_CHECK_TEXT = {"confirmed": "HSIS 행과 이름 일치", "unchecked": "이름 대조 불가",
                     "no_row": "HSIS 에 이 신호 행 없음 — 네트 미확인", "no_hsis": "HSIS 없음(미입력 또는 읽기 실패 — 생성 공시 참조) — 네트 미확인",
                     "not_sw": "주어가 SW 변수가 아님 — 네트 미확인",
                     "row_ambiguous": "HSIS 에 이 신호 행이 둘 이상 — 네트 미확인"}


def parse_hw_requirement_docx(path: str) -> dict[str, dict[str, Any]]:
    """``{HW ID: {"name", "related": [system IDs], "tolerances": [...], "ranges": [...], "nominals": [...]}}`` from the
    attribute tables of a HW requirements specification (the SyRS layout: ``ID | HwTSR_0204``, ``Description | …``,
    ``Related ID | …``). Each tolerance: the words (``허용 오차: ±3%``), the number, its unit (``%`` — relative) and the
    units the block measures (a ``%`` applies only to them). ``ranges`` / ``nominals``: what the block writes it
    measures — an absolute tolerance is on that scale (``2.5V(±0.15V)`` is a monitor node, not the 5 V the SW compares).
    The first table of an ID wins. Raises if the file cannot be read: the caller discloses it."""
    out: dict[str, dict[str, Any]] = {}
    for cells_map in _attribute_maps(path):
        hid = cells_map.get("id", "")
        if not HW_ID.fullmatch(hid) or hid in out:
            continue
        text = "\n".join(cells_map.get(_field_key(f), "") for f in _HW_FIELDS).strip()
        ranges = [{"lo": m.group("lo"), "hi": m.group("hi"), "unit": m.group("unit")} for m in _HW_RANGE.finditer(text)]
        measured = sorted({r["unit"] for r in ranges})
        out[hid] = {"name": cells_map.get("name", ""),
                    "related": list(dict.fromkeys(SYSTEM_ID.findall(cells_map.get("related id", "")))),
                    "tolerances": [{"text": m.group(0).strip(), "value": m.group("num"), "unit": m.group("unit") or "",
                                    "measured_units": measured} for m in _TOLERANCE.finditer(text)],
                    "ranges": ranges,
                    "nominals": [{"value": m.group("v"), "unit": m.group("u1") or m.group("u2")}
                                 for m in _HW_NOMINAL.finditer(text)],
                    # (R50) every field — a HSIS net name is matched here (``V-BAT`` in a verification criteria)
                    "full_text": "\n".join(cells_map.values()),
                    # (review W3b) only a block that writes a monitor node takes a divider ratio
                    "monitor_node": bool(_MONITOR_NODE.search(text)),
                    "scales": [dict(s, via=f"{hid}: {s['via']}", source="own") for s in _formula_scales(text)]}
    return out


def _attribute_maps(path: str):
    """(R50 review r2 I8) The attribute tables of a HW specification, each as ``{field key: text}`` — a field written
    over several rows joined (each distinct value once, the last value cell of a row). One reader for the HW requirement
    and the HW design parsers: two copies drifted on empty values. Raises if the file cannot be read."""
    from docx import Document
    for table in Document(path).tables:
        parts: dict[str, list[str]] = {}
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 2 or not cells[0]:
                continue
            values = [c for c in cells[1:] if c and c != cells[0]]
            bucket = parts.setdefault(_field_key(cells[0]), [])
            if values and values[-1] not in bucket:
                bucket.append(values[-1])
        yield {k: "\n".join(v) for k, v in parts.items()}


def _formula_scales(text: str) -> list[dict[str, str]]:
    """(R50) The monitor-node ratios a text writes: ``k`` and the side the monitor is on — ``to_source`` (display) is what
    a node value is multiplied by on the quantity's own scale (``Sensor Power Monitor = Sensor Power *0.5`` → 2;
    ``Battery Voltage = Monitor Voltage * 9`` → 9). A formula whose both or neither side is a monitor says nothing about
    which is the node; (review W3a / I4) one between a current and a voltage, of a time, or with an offset is no ratio."""
    out = []
    for m in _HW_FORMULA.finditer(text):
        k = _decimal(m.group("k"))
        lhs, rhs = m.group("lhs"), m.group("rhs")
        lhs_monitor, rhs_monitor = bool(_MONITOR_WORD.search(lhs)), bool(_MONITOR_WORD.search(rhs))
        if not k or lhs_monitor == rhs_monitor or _NOT_A_NODE.search(lhs) or _NOT_A_NODE.search(rhs):
            continue
        units = {u.upper().lstrip("M") for side in (lhs, rhs) for u in _FORMULA_UNIT.findall(side)}
        mixed = (_CURRENT_WORD.search(lhs) and _VOLTAGE_WORD.search(rhs)) or (_VOLTAGE_WORD.search(lhs)
                                                                               and _CURRENT_WORD.search(rhs))
        if len(units) > 1 or mixed:
            continue
        side = "lhs" if lhs_monitor else "rhs"
        out.append({"k": _plain(k), "side": side, "to_source": _plain(_tidy(_to_source(Decimal(1), k, side))),
                    "via": m.group(0).strip(), "formula": m.group(0).strip()})
    return out


def _to_source(x: Decimal, k: Decimal, side: str) -> Decimal:
    """A node value ``x`` on the quantity's own scale: ``monitor = source * k`` → ``x / k``; ``source = monitor * k`` →
    ``x * k`` (exact division — ``0.15 / 0.3`` is ``0.5``, never ``0.4999…``)."""
    return x / k if side == "lhs" else x * k


def _to_node(x: Decimal, k: Decimal, side: str) -> Decimal:
    return x * k if side == "lhs" else x / k


def _tidy(d: Decimal) -> Decimal:
    """(review I4) A non-terminating quotient (``1 / 0.3``) rounded up to six decimals — a tolerance only grows."""
    exponent = d.as_tuple().exponent
    return d.quantize(Decimal("0.000001"), rounding="ROUND_CEILING") if isinstance(exponent, int) and exponent < -6 else d


def parse_hw_design_docx(path: str) -> dict[str, dict[str, Any]]:
    """(R50) ``{HW design ID: {"name", "related": [HW requirement IDs], "scales": [...]}}`` from a HW architecture design
    specification (``ID | HwC_07``, ``Description | … Sensor Power Monitor = Sensor Power *0.5``, ``Related ID |
    HwTR_0701, HwTSR_0203``). A field written over several rows is read whole, as the HW requirements (review I4).
    Raises if the file cannot be read: the caller discloses it."""
    out: dict[str, dict[str, Any]] = {}
    for cells_map in _attribute_maps(path):
        did = cells_map.get("id", "")
        if not HW_ID.fullmatch(did) or did in out:
            continue
        out[did] = {"name": cells_map.get("name", ""),
                    "related": list(dict.fromkeys(HW_ID.findall(cells_map.get("related id", "")))),
                    "scales": [dict(s, via=f"{did} {cells_map.get('name', '')}: {s['via']}".strip(), source="design")
                               for s in _formula_scales("\n".join(cells_map.get(_field_key(f), "")
                                                                  for f in ("Description", "Range")))]}
    return out


def load_hw_design(path: str | None) -> tuple[dict[str, dict[str, Any]] | None, dict[str, Any] | None]:
    """(R50) As `load_hw_requirements`, for the HW design specification."""
    if not path:
        return None, None
    name = re.split(r"[\\/]", str(path))[-1]
    try:
        blocks = parse_hw_design_docx(str(path))
    except Exception as exc:  # noqa: BLE001 — an unreadable optional input is disclosed, not fatal
        return None, {"file": name, "error": f"{type(exc).__name__}: {exc}"[:200]}
    return blocks, {"file": name, "blocks": len(blocks),
                    "blocks_with_formula": sum(1 for b in blocks.values() if b["scales"])}


def apply_hw_design(hw: dict[str, dict[str, Any]], design: dict[str, dict[str, Any]] | None) -> dict[str, int]:
    """(R50) Give each HW requirement block that writes a monitor node (review W3b — ``Monitor전압: 2.5V(±0.15V)``; a
    power switch's 5 V output is the source, not a node) the one ratio its own text or the design blocks citing it write
    (``scale`` — with its ``source``, own or design: review W4); two different ratios are a conflict, not a choice
    (``scale_conflict`` — shown, review W6). Returns the blocks given one, by source."""
    given = {"design": 0, "own": 0}
    for hid, block in hw.items():
        if not block.get("monitor_node"):
            continue
        scales = list(block.get("scales") or [])
        scales += [s for d in (design or {}).values() if hid in d["related"] for s in d["scales"]]
        ratios = {Decimal(s["to_source"]) for s in scales}
        if len(ratios) == 1:
            block["scale"] = scales[0]
            block["to_source"], block["scale_via"] = scales[0]["to_source"], scales[0]["via"]
            given[scales[0]["source"]] += 1
        elif len(ratios) > 1:
            block["scale_conflict"] = [s["via"] for s in scales]
    return given


def hsis_rows_from_signals(signals: list[dict[str, Any]] | None) -> list[dict[str, Any]]:
    """(R50 review I3) The HSIS signals the STS generator already reads (`generators.sts._load_hsis_signals` — header-based,
    cached, both layouts) as rows: ``{"id", "vars", "ids" (the Related ID column's system IDs), "nets" (the signal name
    when it is a net, ``VCC_HALL_MON``), "name"}``. A signal with no SW variable is no row."""
    rows = []
    for s in signals or []:
        sw_vars = list(dict.fromkeys(_SW_VAR.findall(str(s.get("sw_var_name") or ""))))
        if not sw_vars:
            continue
        name = str(s.get("signal_name") or "").strip()
        rows.append({"id": str(s.get("id") or ""), "vars": sw_vars,
                     "ids": sorted(set(SYSTEM_ID.findall(str(s.get("related_id") or "")))),
                     "nets": [name] if _NET.fullmatch(name) else [], "name": name})
    return rows


def _bare(var: str) -> str:
    return _SW_LAYER.sub("", var).lower()


def _hsis_row(subject: str, rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    """(R50) The one HSIS row of a SW variable — the same name, else the same name under another layer prefix (the
    direction kept — review I2). The row is returned with the variable it matched (``matched`` — review I1)."""
    same = _hsis_matches(subject, rows)
    if len({id(r) for r, _v in same}) != 1:
        return None
    row, var = same[0]
    return {**row, "matched": var}


def _hsis_matches(subject: str, rows: list[dict[str, Any]]) -> list[tuple[dict[str, Any], str]]:
    """Every ``(row, variable)`` a SW variable matches — the same name, else under another layer prefix."""
    if not subject or not _SW_VAR.fullmatch(subject):
        return []
    same = [(r, subject) for r in rows if subject in r["vars"]]
    return same or [(r, v) for r in rows for v in r["vars"] if _bare(v) == _bare(subject)]


def _name_words(*texts: str) -> set[str]:
    """(review W2) The words of a name that say which signal it is — ``Hall Sensor Block 전원 Monitor`` → {hall, sensor},
    ``u16g_ApiIn_MagnetLevel`` · ``V_MAGNET_MON`` → {magnet}; generic words (power, monitor, 전원 …) dropped."""
    words = set()
    for text in texts:
        for part in re.split(r"[^A-Za-z0-9가-힣]+", str(text or "")):     # ``u16g`` · ``ApiIn`` · ``MagnetLevel``
            if not re.fullmatch(r"[us]\d+[gs]?", part):
                words |= {w.lower() for w in _NAME_TOKEN.findall(part)}
    return {w for w in words if len(w) > 1 and w not in _GENERIC_WORDS}


def _names_agree(a: set[str], b: set[str]) -> bool | None:
    """``True`` when two names share a word (``bat`` · ``battery`` — one the other's start, 3 letters at least),
    ``False`` when both name something and nothing is shared, ``None`` when either names nothing particular."""
    if not a or not b:
        return None
    if all(re.fullmatch(r"[가-힣]+", w) for w in a) != all(re.fullmatch(r"[가-힣]+", w) for w in b) and \
            (all(re.fullmatch(r"[가-힣]+", w) for w in a) or all(re.fullmatch(r"[가-힣]+", w) for w in b)):
        return None         # (review r2 I1) ``배터리 전압 모니터`` vs ``Battery Power``: not comparable, not "different"
    return any(x == y or (min(len(x), len(y)) >= 3 and (x.startswith(y) or y.startswith(x))) for x in a for y in b)


def _row_words(row: dict[str, Any]) -> set[str]:
    """(review r2 I1) The row's name, nets and the variable the fact matched — not every variable of the row (HD's LIN
    row lists 32: its words agreed with any block name)."""
    return _name_words(row.get("name") or "", *row.get("nets", []), row.get("matched") or "")


def load_hw_requirements(path: str | None) -> tuple[dict[str, dict[str, Any]] | None, dict[str, Any] | None]:
    """The HW requirement blocks of ``path`` and a record for the quality report — ``(None, None)`` when not given,
    ``(None, {"file", "error"})`` when unreadable (never raised: the STS is still generated, the report says why)."""
    if not path:
        return None, None
    name = re.split(r"[\\/]", str(path))[-1]
    try:
        blocks = parse_hw_requirement_docx(str(path))
    except Exception as exc:  # noqa: BLE001 — an unreadable optional input is disclosed, not fatal
        return None, {"file": name, "error": f"{type(exc).__name__}: {exc}"[:200]}
    return blocks, {"file": name, "blocks": len(blocks),
                    "blocks_with_tolerance": sum(1 for b in blocks.values() if b["tolerances"])}


def _decimal(text: Any) -> Decimal | None:
    try:
        return Decimal(str(text))
    except (ArithmeticError, ValueError):
        return None


def _on_scale(block: dict[str, Any], base: str, value: Decimal) -> bool:
    """(R49 review W1) Does the block write a range or a nominal value on the scale of ``value`` (base unit)? A tolerance
    the block writes for a 2.5 V monitor node is not one of a 4.85 V threshold (the divider is not written)."""
    for r in block.get("ranges") or []:
        u_base, u_scale = _UNIT_SCALE.get(r["unit"], (r["unit"], Decimal(1)))
        lo, hi = Decimal(r["lo"]) * u_scale, Decimal(r["hi"]) * u_scale
        margin = _HW_SCALE_MARGIN * max(hi - lo, abs(hi))
        if u_base == base and lo - margin <= value <= hi + margin:
            return True
    for n in block.get("nominals") or []:
        u_base, u_scale = _UNIT_SCALE.get(n["unit"], (n["unit"], Decimal(1)))
        nominal = Decimal(n["value"]) * u_scale
        if u_base == base and abs(value - nominal) <= _HW_SCALE_MARGIN * abs(nominal):
            return True
    return False


def hw_tolerances_for(req: dict[str, Any], evidence: dict[str, Any], hw: dict[str, dict[str, Any]],
                      hsis_rows: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """(R49) The HW tolerances that may bound how a boundary group's value is measured: blocks citing a system ID the
    requirement cites, a tolerance in the value's unit (``±0.15V`` for ``4.85V``, ``mV`` too) or a relative one
    (``±3%``) of a block that measures that unit. Each: the block, the words, the size at the value (in the value's unit;
    ``None`` when an absolute tolerance is written on another scale — ``scale_known`` false), the shared IDs, whether
    the block cites the system block the group's fact comes from (``direct``) and whether the step spacing lies inside.
    Direct first, then the HSIS row's block, then the most shared IDs — a hub block (cited by every power requirement)
    links several; the caller says the monitor path is undecided when more than one block remains
    (`hw_tolerance_summary`). (R50) An absolute tolerance off the value's scale is put on it by the block's one
    monitor-node ratio (`apply_hw_design`) when the value then is on the node's scale (``scaled_by``) — unless the HSIS
    row of the fact's SW signal names another node than the formula (review W2: ``scale_net_mismatch``, not scaled);
    the HSIS row marks the one block it points to (``hsis`` — `_hsis_pick`), or the block it points to whose name is
    another signal's (``hsis_mismatch`` — no path). ``hsis_rows`` ``None``: HSIS not given."""
    unit = str(evidence.get("unit") or "")
    cited = set(cited_system_ids(req))
    value, step = _decimal(evidence.get("value")), _decimal(evidence.get("step"))
    if not unit or not cited or value is None or step is None:
        return []
    base, scale = _UNIT_SCALE.get(unit, (unit, Decimal(1)))
    source_id = ((evidence.get("source") or {}).get("id"))
    row = _hsis_row(str(evidence.get("signal") or ""), hsis_rows or [])
    out = []
    for hid, block in hw.items():
        shared = [s for s in block["related"] if s in cited]
        if not shared:
            continue
        for tol in block["tolerances"]:
            scaled = mismatch = None
            if tol["unit"] == "%":
                if base not in {_UNIT_SCALE.get(u, (u, Decimal(1)))[0] for u in tol["measured_units"]}:
                    continue            # the Battery monitor's ±3% is not a hold time's
                size = written = abs(value) * Decimal(tol["value"]) / 100   # relative: the same on any node's scale
            else:
                t_base, t_scale = _UNIT_SCALE.get(tol["unit"], (tol["unit"], Decimal(1)))
                if not tol["unit"] or t_base != base:
                    continue
                written = Decimal(tol["value"]) * t_scale / scale
                size = written if _on_scale(block, base, value * scale) else None
                ratio = block.get("scale")
                if size is None and ratio and _on_scale(block, base, _to_node(value * scale, Decimal(ratio["k"]),
                                                                                ratio["side"])):
                    agree = _names_agree(_name_words(ratio["formula"]), _row_words(row)) if row else None
                    if row and agree is False:
                        mismatch = {"formula": ratio["via"], "row": row["id"] or row["matched"],
                                    "name": row.get("name") or row["matched"]}
                    else:
                        # the tolerance is the node's; on the value's scale through the ratio
                        size = _tidy(_to_source(written, Decimal(ratio["k"]), ratio["side"]))
                        signal = str(evidence.get("signal") or "")
                        scaled = {"via": ratio["via"], "source": ratio["source"],
                                  "check": ("confirmed" if agree else "unchecked" if row else
                                            "no_hsis" if hsis_rows is None else
                                            "not_sw" if not _SW_VAR.fullmatch(signal) else
                                            "row_ambiguous" if _hsis_matches(signal, hsis_rows) else "no_row")}
            out.append({"hw_id": hid, "name": block["name"], "text": tol["text"], "shared": shared,
                        "scaled_by": scaled["via"] if scaled else None, "scale": scaled,
                        "scale_net_mismatch": mismatch,
                        # (review W6) two ratios for the block: shown, never chosen
                        "scale_conflict": (block.get("scale_conflict") if size is None and tol["unit"] != "%"
                                           else None),
                        "direct": bool(source_id) and source_id in block["related"],
                        # (R54 review W2) a relative tolerance is 'known' on any scale — no check that the block reads
                        #   this value
                        "relative": tol["unit"] == "%",
                        "scale_known": size is not None, "size": _plain(size) if size is not None else None,
                        # on an unknown scale, the number as written (in the value's unit) — never a point
                        "written_size": _plain(written) if size is None else None,
                        "unit": unit, "inside": step < size if size is not None else None})
    out.sort(key=lambda t: (not t["direct"], -len(t["shared"]), t["hw_id"]))
    # (review r2 I3) the fact's own system block cited decides first — HSIS picks nothing then (its mark would sit on a
    #   candidate the summary does not use)
    pick, refused = _hsis_pick(row, out, hw) if not any(t["direct"] for t in out) else (None, None)
    for t in out:
        t["hsis"] = pick if pick and t["hw_id"] == pick["hw_id"] else None
        t["hsis_mismatch"] = refused if refused and t["hw_id"] == refused["hw_id"] else None
    # (review C1) the HSIS block ahead of the other indirect ones: the evidence keeps the first `MAX_HW_TOLERANCES`, and
    #   the note cites from what is kept
    out.sort(key=lambda t: (not t["direct"], not t["hsis"], -len(t["shared"]), t["hw_id"]))
    return out


def _net_in(net: str, text: str) -> bool:
    """(review W5) A net as a whole token — ``VCC_BAT`` is not in ``VCC_BAT_MON``."""
    return bool(re.search(rf"(?<![A-Za-z0-9_-]){re.escape(net)}(?![A-Za-z0-9_-])", text))


def _hsis_pick(row: dict[str, Any] | None, candidates: list[dict[str, Any]],
               hw: dict[str, dict[str, Any]]) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    """(R50) The one candidate block the HSIS row of the fact's SW signal points to: the row names a system ID only that
    block cites among the candidates (a hub ID every candidate cites decides nothing), else a net the row names appears
    whole in that block's text alone — among the blocks the IDs left when they left several (review W5). ``(pick, None)``;
    ``(None, refused)`` when the pointed block's name names another signal than the row (review W2: ``Hall Sensor Block
    전원 Monitor`` for ``V_MAGNET_MON`` — the documents disagree; no path is decided); ``(None, None)`` when no row, no one
    block, or one candidate only. The pick carries the row's variable that matched (review I1)."""
    blocks = list(dict.fromkeys(t["hw_id"] for t in candidates))
    if not row or len(blocks) < 2:
        return None, None
    related = {b: set(hw[b]["related"]) for b in blocks}
    own = {b: {i for i in related[b] if sum(i in related[o] for o in blocks) == 1} for b in blocks}
    by_id = [b for b in blocks if own[b] & set(row["ids"])]
    if len(by_id) == 1:
        anchors = sorted(own[by_id[0]] & set(row["ids"]))
        # (review r2 W2) HD ``SyTSR_0114``: HwTSR_0203 and HwTSR_0202 'Position Sensor 전원 Monitor' (no tolerance, so no
        #   candidate) both cite it — unique among the candidates is not unique: the monitor block may be outside them.
        #   (r3 Info 1) any anchor no other block cites decides; refused only when every anchor is shared outside
        outside = {a: [o for o in hw if o not in related and a in hw[o]["related"]] for a in anchors}
        clean = [a for a in anchors if not outside[a]]
        if not clean:
            anchor = anchors[0]
            return None, {"hw_id": by_id[0], "row": row["id"], "var": row["matched"], "by": "system_id", "anchor": anchor,
                          "kind": "outside", "outside": [f"{o} {hw[o]['name']}".strip() for o in outside[anchor][:3]],
                          # (r3 W4) "no tolerance" only when none of them writes one — else "not a candidate"
                          "outside_without_tolerance": not any(hw[o]["tolerances"] for o in outside[anchor]),
                          "block_name": hw[by_id[0]]["name"], "row_name": row.get("name") or row["matched"]}
        found = {"hw_id": by_id[0], "row": row["id"], "var": row["matched"], "by": "system_id", "anchor": clean[0]}
    else:
        pool = by_id or blocks
        by_net = [b for b in pool if any(_net_in(n, hw[b].get("full_text") or "") for n in row["nets"])]
        if len(by_net) != 1:
            return None, None
        found = {"hw_id": by_net[0], "row": row["id"], "var": row["matched"], "by": "net",
                 "anchor": next(n for n in row["nets"] if _net_in(n, hw[by_net[0]].get("full_text") or ""))}
    agree = _names_agree(_name_words(hw[found["hw_id"]]["name"]), _row_words(row))
    if agree is False:
        return None, dict(found, kind="name", block_name=hw[found["hw_id"]]["name"],
                          row_name=row.get("name") or row["matched"])
    return dict(found, names_checked=bool(agree)), None


def hw_tolerance_summary(evidence: dict[str, Any], tolerances: list[dict[str, Any]] | None = None) -> dict[str, Any] | None:
    """(R49 review C2) One reading of a group's candidates (all of them — ``tolerances`` or the evidence's, before the
    sheet's cut, review r2 I1): the relevant ones (those citing the fact's own system block, else all), whether one
    monitor path remains (``certain``), the largest size on a known scale — ``None`` when there is none, or when a
    candidate on an unknown scale writes as much or more (review r2 W1: the points would not be outside it) — and whether
    the step spacing is, or may be, inside any candidate (an unknown scale compared with its written number: a divider
    to a larger input scale only widens it — r2 I2). (R50 review) What the HSIS and the ratios could not settle:
    ``path_refused`` (the row points to a block named for another signal), ``scale_unconfirmed`` (scaled without the
    row's net to check the formula against), ``net_mismatch`` (the formula is another node's — not scaled),
    ``scale_conflict`` (two ratios)."""
    tolerances = tolerances if tolerances is not None else (evidence.get("hw_tolerance") or [])
    if not tolerances:
        return None
    # the fact's own system block cited, else (R50) the HSIS row's block, else every candidate
    relevant = [t for t in tolerances if t["direct"]] or [t for t in tolerances if t.get("hsis")] or tolerances
    known = [Decimal(t["size"]) for t in relevant if t["scale_known"]]
    written = [Decimal(t["written_size"]) for t in relevant if not t["scale_known"] and t.get("written_size")]
    step = Decimal(str(evidence["step"]))
    largest = max(known) if known else None
    inside = (largest is not None and step < largest) or any(step < w for w in written)
    bounded = largest is not None and not any(w >= largest for w in written)
    certain = len({t["hw_id"] for t in relevant}) == 1
    # (R54 review W2) per block, from every candidate (before the sheet's cut): 'known' — an absolute tolerance on the
    #   value's scale; 'relative' — only a relative one (on any scale, so no check); 'unknown' — off its scale
    #   — (R54 review r2 I4) 'scaled_unconfirmed' when it is known only through a monitor ratio the HSIS did not confirm
    order = ("known", "scaled_unconfirmed", "relative", "unknown")
    scale_by_block: dict[str, str] = {}
    for t in relevant:
        if t["scale_known"] and not t.get("relative"):
            state = "scaled_unconfirmed" if (t.get("scale") or {}).get("check") not in (None, "confirmed") else "known"
        else:
            state = "relative" if t.get("relative") else "unknown"
        prev = scale_by_block.get(t["hw_id"])
        scale_by_block[t["hw_id"]] = state if prev is None else min(prev, state, key=order.index)
    return {"blocks": list(dict.fromkeys(t["hw_id"] for t in relevant)), "certain": certain,
            "scale_by_block": scale_by_block,
            "by_hsis": certain and not any(t["direct"] for t in relevant) and any(t.get("hsis") for t in relevant),
            "scaled": any(t.get("scaled_by") for t in relevant),
            "scale_unconfirmed": any((t.get("scale") or {}).get("check") not in (None, "confirmed") for t in relevant),
            "net_mismatch": any(t.get("scale_net_mismatch") for t in relevant),
            "scale_conflict": any(t.get("scale_conflict") for t in relevant),
            "path_refused": next((t["hsis_mismatch"] for t in tolerances if t.get("hsis_mismatch")), None),
            "scale": next((t["scale"] for t in relevant if t.get("scale")), None),
            "largest": _plain(largest) if bounded else None, "inside": inside}


def _plain(d: Decimal) -> str:
    text = format(d.normalize(), "f")
    return text if text != "-0" else "0"


def _tolerance_note(evidence: dict[str, Any]) -> str:
    """(R49) The TC precondition line — short; the candidates are in the 'Requirement Evidence' HW Tolerance column
    (review I5: a long text in a merged cell is cut when printed)."""
    summary = evidence.get("hw_tolerance_summary") or hw_tolerance_summary(evidence)
    if not summary or not summary["inside"]:
        return ""
    unit = str(evidence.get("unit") or "")
    tolerances = [t for t in evidence["hw_tolerance"] if t["hw_id"] in summary["blocks"]]
    if summary["certain"] and summary["largest"] is not None:
        # (review r2 I3) one block with two tolerances: cite the one the points use
        tolerances = sorted(tolerances, key=lambda t: t["size"] != summary["largest"])

    def said(t: dict) -> str:
        return f"{t['hw_id']} ±{t['size']}{unit}" if t["scale_known"] else f"{t['hw_id']} {t['text'].split(':')[-1].strip()}(척도 불명)"
    who = said(tolerances[0]) if summary["certain"] else \
        "감시 경로 미정 — 후보 " + " · ".join(dict.fromkeys(said(t) for t in tolerances))
    if summary.get("by_hsis"):
        h = tolerances[0]["hsis"]
        who += f", HSIS {h['row'] or h['var']} 로 경로 " + ("확정" if h.get("names_checked") else "정함(이름 대조 불가)")
    refused = summary.get("path_refused")
    if refused:
        who += ", " + _refusal_text(refused)
    scaled = summary.get("scale")
    if scaled:
        who += f", {_SCALE_SOURCE_TEXT[scaled['source']]}으로 척도 환산({_SCALE_CHECK_TEXT[scaled['check']]})"
    if summary.get("net_mismatch"):
        who += ", 분압식이 HSIS 의 다른 노드라 척도 환산 안 함(문서 확인)"
    if summary.get("scale_conflict"):
        who += ", 분압비 충돌(문서마다 다른 비율) — 척도 불명"
    head = (f"HW 측정 허용오차({who})가 경계 점 간격 {evidence['step']}{unit} 보다 큼 — HIL 에서는 경계 포함(이상/초과·이하/미만)을 "
            "가를 수 없다: SW 변수를 직접 주입해 판정하거나")
    if summary["largest"] is None:
        return (head + " 척도(분압비)를 HW 문서로 확인한 뒤 허용오차 밖에서 방향만 확인 ('Requirement Evidence' HW Tolerance 열)")
    value, step, largest = Decimal(str(evidence["value"])), Decimal(str(evidence["step"])), Decimal(summary["largest"])
    where = "" if summary["certain"] else "가장 큰 후보의 "
    # (review r2 I4) on the written step, outward — a point a HIL bench can set
    grid = Decimal(1).scaleb(step.as_tuple().exponent) if step.as_tuple().exponent < 0 else Decimal(1)
    low = (value - largest - step).quantize(grid, rounding="ROUND_FLOOR")
    high = (value + largest + step).quantize(grid, rounding="ROUND_CEILING")
    # a candidate on an unknown scale may be larger still: the points are outside the known ones only — said so
    unknown = list(dict.fromkeys(t["hw_id"] for t in tolerances if not t["scale_known"]))
    caveat = f" — 척도 불명 후보 {', '.join(unknown)} 는 분압비 확인 전이라 이 밖 점에 넣지 않았다" if unknown else ""
    return (head + f" {where}허용오차 밖 {_plain(low)}{unit} · {_plain(high)}{unit} 에서 "
            f"방향만 확인{caveat} ('Requirement Evidence' HW Tolerance 열)")


def _refusal_text(r: dict[str, Any]) -> str:
    """(R50 review W2, r2 W2) Why the HSIS row did not settle the path — the ID it names is cited by a HW block without a
    tolerance too (the monitor block may be outside the candidates), or the block it points to is named for another
    signal (the monitor block is outside the candidates, or the documents disagree). Shown, never decided."""
    row = r["row"] or r["var"]
    if r.get("kind") == "outside":
        which = "허용오차 없는 HW 블록" if r.get("outside_without_tolerance", True) else "후보가 아닌 HW 블록"
        return (f"HSIS {row} 의 시스템 ID {r['anchor']} 를 {which} {' · '.join(r['outside'])} 도 인용 — "
                "감시 블록이 후보 밖일 수 있어 경로 미정(확인)")
    return (f"HSIS {row}({r['row_name']}) 가 가리키는 {r['hw_id']}({r['block_name']}) 는 이름이 다른 신호 — 감시 블록이 "
            "후보 밖이거나 문서가 서로 다름, 경로로 쓰지 않음(확인)")


def _tolerance_text(evidence: dict[str, Any]) -> str:
    """(R49) The 'Requirement Evidence' cell: the linked tolerances — or why there are none (review I1: not given,
    none linked)."""
    if evidence.get("kind") == "acceptance":
        return "— (판정 기준 측정 — 해당 없음)"                       # (R51 review r2 I-4)
    if "hw_tolerance" not in evidence:
        return "— (HW 요구사항서 없음 — 미입력 또는 읽기 실패, 생성 공시 참조)"   # (review r2 I5)
    summary = evidence.get("hw_tolerance_summary") or hw_tolerance_summary(evidence)
    if not summary:
        return "— (Related ID 로 이어진 HW 허용오차 없음)"
    unit = str(evidence.get("unit") or "")

    def one(t: dict) -> str:
        if t["scale_known"]:
            size = f"±{t['size']}{unit}" + (f" [한 눈금 {evidence['step']}{unit} 가 그 안]" if t["inside"]
                                            and evidence.get("kind") != "measurement" else "")
        elif t.get("scale_conflict"):
            size = "척도 불명 — 분압비 충돌: " + " ↔ ".join(t["scale_conflict"])         # (R50 review W6)
        else:
            size = "척도 불명(블록이 적은 값의 척도와 다름 — 분압비 확인)"
        hsis, refused, mismatch, scaled = t.get("hsis"), t.get("hsis_mismatch"), t.get("scale_net_mismatch"), t.get("scale")
        how = {"system_id": "시스템 ID ", "net": "네트 "}
        return (f"{t['hw_id']} {t['name']}: {t['text']} → {size} [공유 {', '.join(t['shared'])}]"
                + (" [이 문장의 시스템 블록 인용]" if t["direct"] else "")
                # (R50) how the path and the scale were decided — and (review W2) where the documents disagree
                + (f" [HSIS {hsis['row'] or '—'} {hsis['var']} → {how[hsis['by']]}{hsis['anchor']}"
                   + ("" if hsis.get("names_checked") else " · 이름 대조 불가") + "]" if hsis else "")
                + (f" [{_refusal_text(refused)}]" if refused else "")
                + (f" [척도 환산 — {scaled['via']} · {_SCALE_SOURCE_TEXT[scaled['source']]} · "
                   f"{_SCALE_CHECK_TEXT[scaled['check']]}]" if scaled else "")
                + (f" [분압식 {mismatch['formula']} 은 HSIS {mismatch['row']}({mismatch['name']}) 와 다른 노드 — 환산 "
                   "안 함]" if mismatch else ""))
    first = (f"첫 입력값 {str(evidence.get('value')).split(',')[0].strip()}{unit} 기준 — "
             if evidence.get("kind") == "measurement" else "")    # (R51 review r2 I-5) each step's band is in its text
    return first + ("" if summary["certain"] else "감시 경로 미정 — ") + "; ".join(one(t) for t in evidence["hw_tolerance"])


def source_label(source: dict | None) -> str:
    """``SyDS SyII_06 · Range`` — or ``SRS`` for the requirement's own text."""
    return f"{source['doc']} {source['id']} · {source['field']}" if source else "SRS"


def _tc_class(group: dict) -> str:
    """(R51) what a TC of this group says: a boundary, a measurement, or pass criteria — one kind per TC."""
    kind = group["evidence"]["kind"]
    return kind if kind in ("measurement", "acceptance") else "boundary"


def _doc_of(group: dict) -> str:
    """The document a group's fact comes from — ``SRS`` for the requirement's own text, else ``SyRS`` / ``SyDS``."""
    source = group["evidence"]["source"]
    return source["doc"] if source else "SRS"


def _is_input(group: dict) -> bool:
    """(R47) A group stepped from a SyDS element's input range (`SYSTEM_INPUT_FIELDS`)."""
    source = group["evidence"]["source"]
    return bool(source) and source["field"] in SYSTEM_INPUT_FIELDS


# ── (R46) evidence for review items, and what filling the gap would improve ──────────────────────────────────────────
# 2026-09-30 user direction: what is missing is first looked for in the other documents; what is still missing is shown
#   with what filling it would improve. A candidate is quoted, never used — the same number may be another quantity.
MAX_EVIDENCE = 3          # candidates quoted per item (all are counted)
_EVIDENCE_ANCHOR = "SwEvidence_0"   # `extract` opens a block at an SRS ID line; a system block is read under this one
_DOC_ORDER = {"SRS": 0, "SyRS": 1, "SyDS": 2}


def _steppable(line: dict, fact: dict) -> bool:
    """A fact this module would turn into boundary steps (`_skip_reason` empty) — the form a candidate must have."""
    return fact["kind"] in {"threshold", "duration"} and not _skip_reason(line, fact)


# (R46 review I1) one quantity written in two units is one value: ``3초`` = ``3s`` = ``3000ms``
_UNIT_SCALE = {"ms": ("s", Decimal("0.001")), "s": ("s", Decimal(1)), "초": ("s", Decimal(1)), "분": ("s", Decimal(60)),
               "mV": ("V", Decimal("0.001")), "mA": ("A", Decimal("0.001")), "℃": ("°C", Decimal(1)),
               "KPH": ("km/h", Decimal(1)), "rpm": ("RPM", Decimal(1))}   # (R47 review I2)


def _match_key(unit: str, value: Decimal) -> tuple[str, Decimal]:
    base, scale = _UNIT_SCALE.get(unit, (unit, Decimal(1)))
    return base, value * scale


def review_evidence_corpus(requirements: list[dict], system: dict[str, dict[str, Any]] | None) -> dict[str, Any]:
    """Every steppable fact of the given documents, by value: all SRS requirements (description and verification
    criteria) and all SyRS / SyDS blocks — not only the cited ones — every text field (a SyDS element's ``Input
    Information`` too). ``singles[(unit, value)]`` — one comparison; ``ranges[(unit, lo, hi)]`` — a line's lower and upper bound of one
    subject (``BAT 전압(9V 이상 16.0V 이하)``); units normalised (`_match_key`). Each candidate: where (``SRS
    SwTR_0605`` / ``SyDS SyEL_05 · Input Information``), the condition as the steps would write it, its operator,
    whether it names its subject (an unnamed hold time does not) and how weakly (a noun before a parenthesis), the line
    and its hash, and the SRS requirement or system block it belongs to (a related sentence or not)."""
    singles: dict[tuple, list] = {}
    ranges: dict[tuple, list] = {}
    documents: list[str] = ["SRS"]

    def add(line: dict, where: str, owner: str) -> None:
        ok = [f for f in line["facts"] if _steppable(line, f)]
        for f in ok:
            singles.setdefault(_match_key(f.get("unit") or "", exact_value(f)), []).append(
                {"where": where, "owner": owner, "condition": _condition_text(line, f), "op": f["op"],
                 # an unnamed hold time (``특정시간(300ms) 이상 유지``) names no subject to borrow
                 "named": _subject_text(f) != "조건" + DURATION_SUFFIX,
                 "weak": f.get("signal_kind") == "paren_owner",     # (review I12, as R32 I4)
                 "line": _clip(line["text"], 160), "line_sha256": line["sha256"]})
        for lo in ok:
            for hi in ok:
                if lo is hi or lo["op"] not in {">", ">="} or hi["op"] not in {"<", "<="} or not same_subject(lo, hi) \
                        or not exact_value(lo) < exact_value(hi):
                    continue
                unit = lo.get("unit") or ""
                ranges.setdefault(_match_key(unit, exact_value(lo)) + _match_key(unit, exact_value(hi))[1:], []).append(
                    {"where": where, "owner": owner, "condition": _condition_text(line, lo), "op": "range",
                     "named": True, "weak": lo.get("signal_kind") == "paren_owner",
                     "line": _clip(line["text"], 160), "line_sha256": line["sha256"]})

    for req in requirements:
        seen: set[int] = set()
        for line, _ in requirement_facts(req):
            if id(line) not in seen:
                seen.add(id(line))
                add(line, f"SRS {req.get('id', '')}", str(req.get("id", "")))
    for sid, block in (system or {}).items():
        if block["doc"] not in documents:
            documents.append(block["doc"])
        for field, text in ((f, block["fields"].get(f)) for f in SYSTEM_TEXT_FIELDS):
            if text:
                for line in _field_lines(_EVIDENCE_ANCHOR, field, text):
                    add(line, source_label({"doc": block["doc"], "id": sid, "field": field}), sid)
    return {"singles": singles, "ranges": ranges, "documents": documents}


def _relation(item_op: str | None, cand_op: str) -> str:
    """(R46 review W3) how a candidate's comparison stands to the item's: ``same``; ``flip`` — the other boundary
    inclusion (``이상`` ↔ ``초과``, the R33 "경계 포함 불일치" of the same sheet); ``opposite`` — the other side
    (``이하`` ↔ ``초과``: the complement or another condition); ``range`` — a range item has no written inclusion."""
    if cand_op == "range":
        return "range"
    if cand_op == item_op:
        return "same"
    return "flip" if _INCLUSION_FLIP.get(str(item_op)) == cand_op else "opposite"


_RELATION_ORDER = {"range": 0, "same": 0, "flip": 1, "opposite": 2}


def _evidence_for(item: dict[str, Any], corpus: dict[str, Any], cited: dict[str, list[str]]) -> list[dict[str, Any]] | None:
    """The candidates for one held-back item — ``None`` when the item has no unit (a bare ``3`` — or ``0 ~ 65535``, the
    domain of every 16-bit variable — matches everywhere: review I2). The item's own fact is never a candidate (it is
    not steppable); another comparison of its sentence can be (review W4: ``u16g_A 가 4.85V 미만 시 경고, 4.85V 미만
    유지 시 정지``). An unnamed hold time counts only in a related requirement (it names no subject, and the value alone
    is no link). Related sentences (the item's requirements and the blocks they cite) first, then the same comparison,
    the other inclusion, the other side; then SRS · SyRS · SyDS."""
    unit = item.get("unit") or ""
    if not unit:
        return None
    if item.get("lo") is not None:
        pool = corpus["ranges"].get(_match_key(unit, Decimal(item["lo"])) + _match_key(unit, Decimal(item["hi"]))[1:], [])
    elif item.get("value") is not None:
        pool = corpus["singles"].get(_match_key(unit, Decimal(item["value"])), [])
    else:
        return None
    ids = [str(x) for x in item.get("srs_ids") or [item.get("srs_id")] if x]
    related = set(ids) | {sid for rid in ids for sid in cited.get(rid, [])}
    out: list[dict[str, Any]] = []
    keys: set[tuple] = set()
    for c in pool:
        key = (c["where"], c["condition"])
        if key in keys or not (c["named"] or c["owner"] in related):
            continue
        keys.add(key)
        out.append({"where": c["where"], "condition": c["condition"], "line": c["line"],
                    "line_sha256": c["line_sha256"], "related": c["owner"] in related, "weak_subject": c["weak"],
                    "inclusion": _relation(item.get("op"), c["op"])})
    out.sort(key=lambda c: (not c["related"], _RELATION_ORDER[c["inclusion"]],
                            _DOC_ORDER.get(c["where"].split(" ", 1)[0], len(_DOC_ORDER))))
    return out


# ── what filling the gap would do — tried in the item's own sentence (review W1) ─────────────────────────────────────
# The fill is written into the sentence, the sentence read again alone, and the item's value must then step under the
# written name while no value of the sentence that stepped before stops (``Pull-Up 전원 : 2.25V 이상 ~ Pull-Up 전원
# 2.75V 이하`` — the name written twice makes both unstated, and 2.25 V lost its step). What cannot be written in
# (a negation, a deadline, a label, a join whose meaning the shape does not tell) is a rewrite: its example keeps the
# sentence's meaning, and its form is exercised by tests.
# (review r2 I-b) placeholders a document does not use; (r2 W7) the Korean one is two words — a one-word name is read
#   together with an unparticled word before it (``따라 대상신호``), as a reader's own two-word name (``도어 각도``) is not
_FILL_NAME = "채움 신호"
_FILL_IDENT = "u16g_FillSignal"    # a C identifier — a noun after the value can take a Korean name's place
_FILL_MONITORED = "채움량"
_SEPARATOR_GAP = re.compile(r"[\s,~∼/]*")
_FILL_STEPS = "이 값의 경계 스텝(값−분해능·값·값+분해능)이 생긴다"
_VERIFIED = " — 이 문장에 넣어 확인"
_REWRITE = " — 예시 형태는 시험으로 확인, 원문 뜻대로 다시 쓸 것"
_SUBJECT_REASONS = frozenset({"no_subject_for_value", "subject_unclear", "no_subject_in_value_field"})


def _alone(raw: str) -> dict[str, Any] | None:
    """The line read again on its own (no heading, no neighbour) — both sides of a fill are read this way. (review r2)
    Whether a fact of a review item's line steps depends on the line alone: its section only says "outcome", and an
    outcome line makes no review item."""
    got = [ln for b in extract(f"ID\t{_EVIDENCE_ANCHOR}\n{raw}\n") for ln in b["lines"]]
    return got[0] if len(got) == 1 else None


def _value_key(f: dict) -> tuple:
    return (f.get("unit") or "", exact_value(f), f.get("op"))


def _stepped_by_key(line: dict | None) -> dict[tuple, list[dict]]:
    out: dict[tuple, list[dict]] = {}
    if line is None:
        return out
    for f in line["facts"]:
        if isinstance(f.get("value"), (int, float)) and not isinstance(f.get("value"), bool) and _steppable(line, f):
            out.setdefault(_value_key(f), []).append(f)
    return out


def _baseline(raw: str) -> Counter:
    """How many facts of each (unit, value, operator) step in the line as written (review r2 I-c: counted, so one of two
    equal comparisons lost is seen)."""
    return Counter({k: len(v) for k, v in _stepped_by_key(_alone(raw)).items()})


# (review r3 I-h) the other reasons a filled sentence may still not step, in words
_SKIP_TEXT = {"value_outside_subject_type": "값이 신호 이름이 선언한 폭 밖", "output_requirement": "의무문(출력)으로 읽힘",
              "outcome_section": "결과 절로 읽힘", "reference_label": "기준 이름표로 읽힘",
              "response_constraint": "응답 기한('이내')으로 읽힘", "not_an_order_comparison": "크기 비교가 아님",
              "kind_range": REVIEW_TEXT["kind_range"][0], "kind_symbolic": "값이 아니라 이름과 비교"}


def _try_fill(raw: str, targets: list[tuple[tuple, str | None]], baseline: Counter) -> str:
    """``""`` when every target value steps under its name (``None``: under any name) and every value that stepped
    before still steps as often; otherwise why not, in the review's own words."""
    line = _alone(raw)
    if line is None:
        return "다시 읽지 못함"
    stepped = _stepped_by_key(line)
    for key, name in targets:
        hits = stepped.get(key, [])
        if not hits:
            held = [f for f in line["facts"] if isinstance(f.get("value"), (int, float))
                    and not isinstance(f.get("value"), bool) and _value_key(f) == key]
            why = _skip_reason(line, held[0]) if held else ""
            if not why:
                return "그 값을 다시 읽지 못함"
            return REVIEW_TEXT[why][0] if why in REVIEW_TEXT else _SKIP_TEXT.get(why, f"스텝 조건이 아님({why})")
        if name is not None and not any(name in {subject_of(f), f.get("signal")} for f in hits):
            return f"적은 이름이 주어로 읽히지 않음(주어로 '{_subject_text(hits[0])}' 를 읽음)"
    lost = sorted(k for k, n in baseline.items() if len(stepped.get(k, [])) < n)
    if lost:
        return "같은 줄의 " + ", ".join(f"{k[1]}{k[0]} {_OP_TEXT.get(k[2], k[2])}" for k in lost) + " 스텝이 사라짐"
    return ""


def _placeholders(text: str) -> str:
    return text.replace(_FILL_IDENT, "u16g_<신호>").replace(_FILL_NAME, "<신호>").replace(_FILL_MONITORED, "<감시량>")


def _excerpt(raw: str, start: int, end: int) -> str:
    """The changed part of a filled sentence with a little context, placeholders shown as such."""
    a, b = max(0, start - 30), min(len(raw), end + 12)
    return _placeholders(("…" if a else "") + raw[a:b].strip() + ("…" if b < len(raw) else ""))


def _bounds(a: dict, b: dict) -> str:
    """How two comparisons of one quantity stand: ``range`` — a lower bound below an upper one (only *and* means
    anything, R45); ``apart`` — the lower above the upper: outside a band (*or* — ``4.85V 미만/5.15V 초과``) or an entry
    and a return (a hysteresis — ``8.5V 이하, 9.0V 이상 복귀``), which the shape does not tell (review r2 W8, as R45 C1);
    ``one_side`` — two bounds on one side (either join means something); ``empty`` — both sides of one value or an
    empty range on the written grid (``9V 미만 9V 이상`` · ``10 초과 11 미만``: *and* never holds, *or* always does —
    review r3 I-g, as R45 I1)."""
    lower = [f for f in (a, b) if f["op"] in {">", ">="}]
    upper = [f for f in (a, b) if f["op"] in {"<", "<="}]
    if len(lower) != 1 or len(upper) != 1:
        return "one_side"
    lo_f, hi_f = lower[0], upper[0]
    grid = min(written_step(lo_f), written_step(hi_f))
    lo = exact_value(lo_f) + (grid if lo_f["op"] == ">" else 0)
    hi = exact_value(hi_f) - (grid if hi_f["op"] == "<" else 0)
    if lo <= hi:
        return "range"
    return "apart" if exact_value(lo_f) > exact_value(hi_f) else "empty"


_APART = ("대역 밖(이거나)인지, 진입·복귀 같은 두 조건(히스테리시스 — 두 문장으로 나눠 적기)인지 원문 뜻으로 정해야 한다 — 모양으로는 "
          "가를 수 없다")
_EMPTY = ("같은 값의 양쪽이거나 빈 범위라 '이고' 면 늘 거짓, '이거나' 면 늘 참이다 — 원문 뜻을 정해 다시 써야 한다(결합어만으로는 뜻이 "
          "생기지 않는다)")


def _previous_in_unit(line: dict, fact: dict) -> dict | None:
    """The comparison right before ``fact`` in its unit, joined to it by nothing but a separator (``/`` ``~`` ``,``)."""
    raw = line.get("raw") or line["text"]
    s = fact["line_span"][0]
    before = [f for f in line["facts"] if f is not fact and f["kind"] in {"threshold", "duration"}
              and f.get("op") in _ORDER and (f.get("unit") or "") == (fact.get("unit") or "")
              and f["line_span"][1] <= s]
    if not before:
        return None
    prev = max(before, key=lambda f: f["line_span"][1])
    return prev if _SEPARATOR_GAP.fullmatch(raw[prev["line_span"][1]:s]) else None


def _joined(line: dict, prev: dict, fact: dict, conn: str, name: str | None) -> tuple[str, int, int]:
    """``prev`` and ``fact`` joined by ``conn`` (the separator between them replaced), with ``name`` written once before
    ``prev`` — the filled line and the span of the change."""
    raw = line.get("raw") or line["text"]
    s, e = fact["line_span"]
    p0, p1 = prev["line_span"]
    head = (name + " ") if name else ""
    filled = raw[:p0] + head + raw[p0:p1] + conn + " " + raw[s:]
    return filled, p0, p1 + len(head) + len(conn) + 1 + (e - s)


def _fill_subject(line: dict, fact: dict, baseline: Counter) -> tuple[str, str]:
    raw = line.get("raw") or line["text"]
    s, e = fact["line_span"]
    key = _value_key(fact)
    prev = _previous_in_unit(line, fact)
    if prev is not None:
        prev_name = subject_of(prev)
        name = None if prev_name else _FILL_NAME
        shape = _bounds(prev, fact)
        if shape == "range":
            # one quantity's lower and upper bound: join them (``2.25V 이상이고 2.75V 이하``) — naming the value again
            #   makes both an unstated join (review W1)
            filled, a, b = _joined(line, prev, fact, "이고", name)
            if not _try_fill(filled, [(key, prev_name or _FILL_NAME), (_value_key(prev), prev_name or _FILL_NAME)],
                             baseline):
                how = "결합어로 이어" if prev_name else "신호 이름을 한 번 적고 결합어로 이어"
                return (f"앞 값과 같은 신호면 {how} 적으면(예: '{_excerpt(filled, a, b)}') {_FILL_STEPS}{_VERIFIED}"
                        + ("(신호 이름만 다시 적으면 결합 미기재로 두 값이 모두 보류된다)" if prev_name else ""), "verified")
        elif shape == "apart":
            filled, a, b = _joined(line, prev, fact, "이거나", name)
            example = "" if _try_fill(filled, [(key, prev_name or _FILL_NAME)], baseline) else \
                f" 대역 밖이면 예: '{_excerpt(filled, a, b)}'(이 문장에 넣어 스텝 확인)."
            return f"앞 값과 한 신호면 {_APART}.{example}", "rewrite"
        elif shape == "empty":
            return f"앞 값과 한 신호면 {_EMPTY}. 다른 신호면 비교하는 신호 이름을 값 앞에 적는다", "rewrite"
    whys = []
    attempts = ((raw[:s] + _FILL_NAME + " " + raw[s:], _FILL_NAME, "비교하는 신호 이름을 값 앞에"),
                (raw[:s] + _FILL_IDENT + " " + raw[s:], _FILL_IDENT, "비교하는 변수 이름(C 식별자)을 값 앞에"))
    for filled, name, how in attempts:
        why = _try_fill(filled, [(key, name)], baseline)
        if not why:
            start = filled.index(name)
            note = f" — 한국어 이름은 이 문장에서 주어로 읽히지 않는다: {whys[0]}" if name == _FILL_IDENT else ""
            return (f"{how} 적으면(예: '{_excerpt(filled, start, start + len(name) + 1 + e - s)}') {_FILL_STEPS}"
                    f"{_VERIFIED}{note}", "verified")
        whys.append(why)
    return f"이 문장에서는 신호 이름을 적어도 스텝되지 않는다 — {whys[-1]}. 그것까지 원문에서 풀어야 한다", "needs_more"


def _fill_range(line: dict, fact: dict, baseline: Counter) -> tuple[str, str]:
    raw = line.get("raw") or line["text"]
    s, e = fact["line_span"]
    lo_t, hi_t = (str(t) for t in fact["value_text"])
    unit = fact.get("unit") or ""
    named = fact.get("signal") and not _UNCLEAR_SUBJECT.search(str(fact.get("signal")).strip())
    insert = "" if named else _FILL_NAME + " "
    written = raw[s:e]
    # (review r2 I-a) the range's match ends in the space before what follows (``0 ~ 5000 RPM``): keep it
    body = f"{lo_t}{unit} 이상이고 {hi_t}{unit} 이하" + written[len(written.rstrip()):]
    filled = raw[:s] + insert + body + raw[e:]
    lo = {"value_text": lo_t, "value": fact["value"][0]}
    hi = {"value_text": hi_t, "value": fact["value"][1]}
    name = None if named else _FILL_NAME
    why = _try_fill(filled, [((unit, exact_value(lo), ">="), name), ((unit, exact_value(hi), "<="), name)], baseline)
    if why:
        return f"이 문장에서는 두 비교로 적어도 스텝되지 않는다 — {why}. 그것까지 원문에서 풀어야 한다", "needs_more"
    return (f"하한·상한의 포함과 주어를 정해 두 비교로 적으면(예: '{_excerpt(filled, s, s + len(insert) + len(body))}' — "
            f"포함은 정한 대로) 하한·상한의 경계 스텝이 생긴다{_VERIFIED}(`~` 범위로는 주어가 있어도 스텝하지 않는다)", "verified")


def _fill_join(line: dict, fact: dict, baseline: Counter) -> tuple[str, str]:
    raw = line.get("raw") or line["text"]
    siblings = [g for g in line["facts"] if same_subject(fact, g) and joint(line, fact, g) == "unstated"]
    if not siblings:
        return "같은 주어의 두 조건 사이에 '이고'(그리고) 또는 '이거나'(또는)를 원문 뜻대로 적는다", "rewrite"
    other = min(siblings, key=lambda g: abs(g["line_span"][0] - fact["line_span"][0]))
    first, second = sorted((fact, other), key=lambda f: f["line_span"][0])
    shape = _bounds(first, second)
    if shape == "apart":
        return f"두 조건이 {_APART}", "rewrite"          # (review r2 W8) ``8.5V 이하, 9.0V 이상 복귀``
    if shape == "empty":
        return f"두 조건이 {_EMPTY}", "rewrite"          # (review r3 I-g) ``9V 미만 9V 이상``
    if not _SEPARATOR_GAP.fullmatch(raw[first["line_span"][1]:second["line_span"][0]]):
        return ("같은 주어의 두 조건 사이에 '이고'(그리고) 또는 '이거나'(또는)를 원문 뜻대로 적는다(사이의 말까지 다시 써야 해 이 "
                "문장에 넣어 보지는 못했다)", "rewrite")
    works = []
    for conn in ("이고",) if shape == "range" else ("이고", "이거나"):
        filled, a, b = _joined(line, first, second, conn, None)
        if not _try_fill(filled, [(_value_key(fact), subject_of(fact))], baseline):
            works.append((filled, a, b))
    if not works:
        return "이 문장에서는 결합어를 적어도 스텝되지 않는다 — 원문을 다시 써야 한다", "needs_more"
    filled, a, b = works[0]
    which = "이고·이거나 어느 쪽이든 스텝된다" if len(works) == 2 else "하한·상한이라 뜻 있는 결합은 이고뿐"
    return f"두 조건 사이에 원문 뜻대로 결합어를 적으면(예: '{_excerpt(filled, a, b)}' — {which}) {_FILL_STEPS}{_VERIFIED}", \
        "verified"


def _fill_monitored(line: dict, fact: dict, baseline: Counter) -> tuple[str, str]:
    raw = line.get("raw") or line["text"]
    e = fact["line_span"][1]
    filled = raw[:e] + "인 " + _FILL_MONITORED + raw[e:]
    why = _try_fill(filled, [(_value_key(fact), _FILL_MONITORED)], baseline)
    if why:
        return f"이 문장에서는 감시량을 적어도 스텝되지 않는다 — {why}", "needs_more"
    return (f"기준 상수 뒤에 비교하는 감시량을 적으면(예: '{_excerpt(filled, fact['line_span'][0], e + 2 + len(_FILL_MONITORED))}'"
            f") {_FILL_STEPS}{_VERIFIED}", "verified")


# (review r2 W5) a negation on the comparison itself: ``… 미만이 아닌/아니면/아닐/아님``, ``… 이상을 벗어나면``, ``… 이하 제외``,
#   ``… 이하가 되지 않으면`` — its example is the complement
#   (review r3 W9) and a hold time's own verb negated: ``300ms 이상 유지하지 않으면`` / ``… 지속되지 않으면`` — held less long
_NEGATED_COMPARISON = re.compile(r"\s*(?:이|가|을|를|은|는)?\s*(?:아(?:닌|니|닐|님)|벗어나|제외|(?:(?:유지|지속)\s*)?(?:하|되)지\s*않)")
# (review r3 W9 · r4 W10) a negated clause — every ending of ``않`` / ``못`` / ``아니`` / ``없`` (``되지 않은 경우``, ``받지
#   못하는``, ``아닐 경우``), but not the adverb ``없이`` (``이상 없이 복귀할 경우`` is a positive clause)
_NEGATIVE_WINDOW = re.compile(r"않|못|아(?:니|닌|닐|님)|없(?!이)")


def _fill_rewrite(line: dict, fact: dict, reason: str) -> str:
    """(review W2 · r2 W5 · W6) An example that keeps the sentence's meaning: a negated comparison is the complement
    (``8.5V 미만이 아닌`` → ``8.5V 이상``); a hold time with a negated predicate names the negation (``300ms 이상 수신되지
    않는`` → ``미수신 300ms 이상``); any other negation gets no example with a comparison word in it. ``이내`` is read as
    ``이하``: in a negative clause the non-event outlasts it (``20ms 이내에 … 되지 않으면`` → ``20ms 초과 유지되면``), in a
    positive one the condition lasts at most that long (``500ms 이내에 9V 이상으로 복귀`` → ``500ms 이하 유지 후``)."""
    raw = line.get("raw") or line["text"]
    e = fact["line_span"][1]
    value = f"{fact.get('value_text') or fact.get('value')}{fact.get('unit') or ''}"
    if reason == "negated_condition":
        if _NEGATED_COMPARISON.match(raw[e:]) and fact.get("op") in _COMPLEMENT:
            return (f"비교에 걸린 부정은 반대 비교로 적으면(예: '<신호> {value} {_OP_TEXT[_COMPLEMENT[fact['op']]]}') "
                    f"{_FILL_STEPS}{_REWRITE}")
        if (fact.get("unit") or "") in _TIME_UNIT_NAMES:
            return (f"부정을 조건 이름에 담아 부정 없이 적으면(예: '<신호> 미수신 {fact.get('raw')}' — '수신되지 않는' → '미수신' "
                    f"처럼) {_FILL_STEPS}{_REWRITE}")
        return ("부정이 비교에 걸리는지(반대 비교로) 서술어에 걸리는지(부정을 조건 이름에) 정해 부정 없이 적으면 이 값의 경계 스텝이 "
                "생긴다 — 원문 뜻대로 다시 쓸 것")
    if reason == "deadline_or_window":
        if _NEGATIVE_WINDOW.search(raw[e:e + 60]):
            example = f"'<조건> 이 {value} 초과 유지되면'"
        else:
            example = f"'<조건> 이 {value} 이하 유지 후'"
        return (f"조건의 시간 창이면 {example} 처럼 유지 시간으로 적으면 경계 스텝이 생기고(원문 '이내' 를 이하로 읽음 — 문서의 "
                f"다른 곳이 '미만'이면 경계 포함 불일치), 응답 기한이면 지금대로(측정 항목) 둔다{_REWRITE}")
    if reason == "parenthesis_labels_may_share_a_quantity":
        return ("괄호 앞 이름 대신 비교하는 신호 이름을 적는다 — 두 조건이 한 신호면 같은 이름과 결합어로, 다른 신호면 각자의 이름으로"
                f"(예: 'u16s_Speed 가 0.8m/s 미만 또는 u16s_Speed 가 2.0m/s 이상'){_REWRITE}")
    return "—"


_TIME_UNIT_NAMES = frozenset({"ms", "s", "초", "분"})


def fill_guidance(line: dict, fact: dict, reason: str) -> tuple[str, str]:
    """(R46) What writing the gap into the requirement would produce, and how that was checked: ``verified`` — written
    into this sentence and read again, the value steps and nothing stepped before stops; ``needs_more`` — tried here
    and not enough (the text says what else holds it back); ``rewrite`` — the sentence has to be rewritten or its
    meaning decided first (a negation, a deadline, a label, a lower bound above an upper one), the example's form is
    exercised by tests."""
    raw = line.get("raw") or line["text"]
    baseline = _baseline(raw)
    if fact["kind"] == "range" and isinstance(fact.get("value_text"), list) and len(fact["value_text"]) == 2:
        return _fill_range(line, fact, baseline)
    if fact["kind"] not in {"threshold", "duration"} or not isinstance(fact.get("value"), (int, float)) \
            or isinstance(fact.get("value"), bool):
        return "—", "rewrite"
    if reason in _SUBJECT_REASONS:
        return _fill_subject(line, fact, baseline)
    if reason == "same_subject_combination_unstated":
        return _fill_join(line, fact, baseline)
    if reason == "monitored_quantity_unknown":
        return _fill_monitored(line, fact, baseline)
    return _fill_rewrite(line, fact, reason), "rewrite"


_FILL_READ = "확인만 — 이미 스텝했다. 원문에 연결어나 한 신호 이름을 적으면 이 확인 항목이 사라진다"
_FILL_CONFLICT = "한쪽 포함으로 통일하면 이 후보가 사라진다(경계 스텝은 지금도 각 문장대로 판정한다)"
_FILL_DIFFERENCE = ("같은 임계(한 조건)면 맞는 값으로 한쪽을 고친다 — 그러면 이 후보가 사라지고 두 문서의 경계 TC 가 같은 점을 "
                    "시험한다(지금은 각 문장의 값대로 스텝한다). 같은 양의 다른 임계(경고·고장 두 단계, 동작 하한·차단 등)면 "
                    "고치지 않고 결함 아님으로 검토 기록에 닫는다")


def if_filled(item: dict[str, Any]) -> str:
    """(R46) The item's 'If Filled' text: a held-back item carries its own (`fill_guidance`, made where the sentence
    is at hand); a reading is to confirm; a conflict disappears when one inclusion is chosen."""
    if item.get("if_filled"):
        return str(item["if_filled"])
    if item.get("kind") == "inclusion_conflict":
        return _FILL_CONFLICT
    if item.get("kind") == "value_difference":
        return _FILL_DIFFERENCE
    if item.get("kind") == "read":
        return _FILL_READ
    return "—"


def attach_review_evidence(items: list[dict[str, Any]], requirements: list[dict],
                           system: dict[str, dict[str, Any]] | None) -> dict[str, Any]:
    """(R46) Give every held-back item with a value its ``evidence`` (the first `MAX_EVIDENCE` candidates) and
    ``evidence_total`` — ``None`` when not searched. Returns the counts for the quality report: items searched, items
    with a candidate (and with a related one), the documents searched, and how each 'If Filled' was checked."""
    held = [it for it in items if it.get("kind") == "held_back"]
    fill_checks = dict(sorted(Counter(str(it.get("fill_check") or "unchecked") for it in held).items()))
    if not held:
        return {"review_evidence_searched": 0, "review_evidence_found": 0, "review_evidence_related": 0,
                "review_evidence_documents": [], "review_fill_checks": fill_checks}
    corpus = review_evidence_corpus(requirements, system)
    cited = {str(r.get("id", "")): cited_system_ids(r) for r in requirements}
    # everything is computed before any item changes — a failure leaves every item as it was (the caller discloses it)
    updates: list[tuple[dict[str, Any], dict[str, Any]]] = []
    searched = found = related = 0
    for it in held:
        # (R51 review r4 W-B) a measured criterion is no gap another sentence fills — not searched, so the board's
        #   count and the sheet's cell ('탐색 대상 아님') say the same
        cands = None if it.get("fill_check") == "not_applicable" else _evidence_for(it, corpus, cited)
        if cands is None:
            updates.append((it, {"evidence": None}))    # not searched: no unit — the same number is everywhere
            continue
        searched += 1
        found += bool(cands)
        related += any(c["related"] for c in cands)
        updates.append((it, {"evidence": cands[:MAX_EVIDENCE], "evidence_total": len(cands)}))
    for it, update in updates:
        it.update(update)
    return {"review_evidence_searched": searched, "review_evidence_found": found,
            "review_evidence_related": related, "review_evidence_documents": corpus["documents"],
            "review_fill_checks": fill_checks}


_INCLUSION_TAG = {"flip": " [포함 다름]", "opposite": " [방향 반대]"}


def _evidence_text(item: dict[str, Any]) -> str:
    """The sheet's evidence cell: the candidates, or that none was found in the documents searched, why none was
    searched, or that the search failed — ``—`` for an item that is no gap (a reading, a conflict) or an output from
    before R46."""
    if item.get("evidence_error"):
        return f"근거 탐색 실패 — {item['evidence_error']}"
    if item.get("fill_check") == "not_applicable":
        return "— (측정 요구 · 판정 기준 항목 — 다른 문장의 같은 값 탐색 대상 아님)"
    if "evidence" not in item:
        return "—"
    cands = item["evidence"]
    if cands is None:
        return "탐색 안 함 — 단위 없는 값(같은 수가 흔해 짝짓지 않는다)"
    if not cands:
        return "준 문서의 다른 문장에 같은 값·단위를 스텝할 수 있게 적은 곳 없음 — 적을 말부터 요구 문서에서 정해야 한다"
    parts = [f"{c['where']}: {c['condition']}" + (" [괄호 앞 명사]" if c.get("weak_subject") else "")
             + (" [같은 요구·인용]" if c.get("related") else "") + _INCLUSION_TAG.get(str(c.get("inclusion")), "")
             + f" — {c['line']}" for c in cands]
    more = int(item.get("evidence_total") or len(cands)) - len(cands)
    return ("; ".join(parts) + (f" 외 {more}건" if more > 0 else "")
            + " (같은 값을 스텝할 수 있게 적은 다른 문장 — 같은 양인지 확인 후 옮길 것)")


# ── (R51) measurement requirements: listed inputs, the output that reports them, its conversion, a judgment tolerance ──
# HDPDM01 / KJPDS02 ``SwEI_01`` verification criteria: "입력 전원 인가 (9.0, 12.0, 16.0)" · "Canoe로 Output_Battery
#   Voltage(9.0, 12.0, 16.0)값 확인" · "(0x0~0xFF, offset : 4V, Resolution : 0x05V, 4V ~ 16.75V)". The output reports the
#   input: at each listed value the expected code is (v − offset) / resolution, judged within the HW measurement accuracy
#   of the path that measures the input (R49/R50 — `hw_tolerances_for` / `hw_tolerance_summary`, reused, not redone)
#   plus one step of the resolution (the rounding is not written). Nothing past the conversion range is stepped: the
#   document does not say what the output is there.
MEASURE_PREFIX = "입력 설정 (요구 측정): "     # not a boundary (BAA): the requirement's own listed values — AOR
_MEASURE_VALUES = r"-?\d+(?:\.\d+)?\s*(?:mV|mA|V|A)?(?:\s*,\s*-?\d+(?:\.\d+)?\s*(?:mV|mA|V|A)?)+"
_MEASURE_INPUT = re.compile(r"(?P<name>[가-힣A-Za-z][가-힣A-Za-z0-9_ ]{0,30}?)\s*인가\s*\((?P<vals>" + _MEASURE_VALUES
                            + r")\)")
_MEASURE_OUTPUT = re.compile(r"(?P<name>[A-Za-z][A-Za-z0-9_ ]{0,40}?)\s*\((?P<vals>" + _MEASURE_VALUES
                             + r")\)\s*값\s*확인")
_CONVERSION = re.compile(
    r"\(\s*(?P<rlo>0[xX][0-9A-Fa-f]+)\s*~\s*(?P<rhi>0[xX][0-9A-Fa-f]+)\s*,\s*(?i:offset)\s*:\s*(?P<off>-?\d+(?:\.\d+)?)\s*"
    r"(?P<unit>mV|mA|V|A)\s*,\s*(?i:resolution)\s*:\s*(?P<res>[0-9A-Fa-fx.]+?)\s*(?P<runit>mV|mA|V|A)\s*,\s*"
    r"(?P<lo>-?\d+(?:\.\d+)?)\s*(?P<lunit>mV|mA|V|A)?\s*~\s*(?P<hi>-?\d+(?:\.\d+)?)\s*(?P<hunit>mV|mA|V|A)\s*\)")
_MEASURE_NUMBER = re.compile(r"-?\d+(?:\.\d+)?")
# what the measurement review items say: why, what to decide (the 'Requirement Review' sheet — kind held_back)
MEASURE_TEXT = {
    "measurement_resolution_written_otherwise": (
        "분해능 표기가 범위·코드와 맞지 않음 — 범위(하한~상한)를 코드 폭으로 나눈 값으로 판정함",
        "분해능 값 확인(범위·코드로 계산한 값이 맞는지)"),
    "measurement_tolerance_unlinked": (
        "측정 판정 범위에 넣을 HW 측정 정확도가 이어지지 않음(Related ID 가 감시 HW 블록과 시스템 ID 를 공유하지 않음) — "
        "분해능만으로 판정함",
        "이 입력을 감시하는 HW 블록(허용 오차)과 이 요구가 같은 시스템 ID 를 인용하는지 확인"),
    "measurement_path_undecided": (
        "측정 경로 HW 블록이 하나로 정해지지 않음 — 후보 중 가장 큰 허용오차로 판정함",
        "이 입력을 감시하는 HW 블록 확인(HSIS 신호 행 · Related ID)"),
    "measurement_hw_not_given": (
        "HW 요구사항서를 주지 않아 판정 범위에 HW 측정 정확도가 없음 — 분해능만으로 판정함",
        "HW 요구사항서(HwRS·HRS) 등록"),
    "measurement_scale_unknown": (
        "이어진 HW 블록의 허용오차가 다른 척도(감시 노드) — 분압비를 몰라 분해능만으로 판정함",
        "그 블록의 감시 노드 분압비 확인"),
}


def measurement_spec(req: dict[str, Any]) -> dict[str, Any] | None:
    """(R51) A requirement whose text lists input values, an output checked at those same values and the output's
    linear conversion — ``None`` otherwise (also when the output is checked at other values, or the conversion is not
    read whole). The resolution is what the range and the codes give, ``(hi − lo) / (code span)``; a written one that
    reads otherwise (``0x05V`` for 0.05 V) is kept (``resolution_written``). The offset must be the range's low end, every
    unit one and the listed values inside the range — else the conversion is not clear: no spec."""
    text = f"{req.get('verification') or ''}\n{req.get('description') or ''}"
    inp, out, conv = _MEASURE_INPUT.search(text), _MEASURE_OUTPUT.search(text), _CONVERSION.search(text)
    if not (inp and out and conv):
        return None
    values = [Decimal(v) for v in _MEASURE_NUMBER.findall(inp.group("vals"))]
    if [Decimal(v) for v in _MEASURE_NUMBER.findall(out.group("vals"))] != values:
        return None
    unit = conv.group("unit")
    units = {unit, conv.group("runit"), conv.group("hunit"), conv.group("lunit") or unit}
    units |= set(re.findall(r"mV|mA|V|A", inp.group("vals") + out.group("vals")))
    if len(units) != 1:
        return None
    rlo, rhi = int(conv.group("rlo"), 16), int(conv.group("rhi"), 16)
    lo, hi, off = Decimal(conv.group("lo")), Decimal(conv.group("hi")), Decimal(conv.group("off"))
    if rhi <= rlo or hi <= lo or off != lo or not all(lo <= v <= hi for v in values):
        return None
    resolution = (hi - lo) / (rhi - rlo)       # (R51 review I1) exact for the codes; rounded for display only
    written = conv.group("res")
    written_value = None if written.lower().startswith("0x") else _decimal(written)
    if written_value is not None and resolution and abs(written_value - resolution) <= resolution / 100:
        written_value = resolution            # ``0.004888V`` for 5/1023: written within 1 % — no review item
    return {"input": inp.group("name").strip(), "output": out.group("name").strip(), "unit": unit,
            "values": [_plain(v) for v in values], "code_low": rlo, "code_high": rhi, "low": _plain(lo),
            "high": _plain(hi), "resolution": _plain(_tidy(resolution)), "resolution_exact": str(resolution),
            "resolution_written": None if written_value == resolution else f"{written}{conv.group('runit')}",
            "lines": [inp.group(0).strip(), out.group(0).strip(), conv.group(0).strip()]}


def _measure_code(value: Decimal, spec: dict[str, Any], rounding: str) -> int:
    """The code of ``value`` rounded ``rounding`` and clamped to the code range — ``ROUND_CEILING`` for a band's low
    end, ``ROUND_FLOOR`` for its high end: a code is inside the band only when its value is."""
    code = (value - Decimal(spec["low"])) / Decimal(spec["resolution_exact"]) + spec["code_low"]
    return int(min(max(code.to_integral_value(rounding=rounding), spec["code_low"]), spec["code_high"]))


def _outward(d: Decimal, up: bool) -> str:
    """(R51 review I1) A band end shown to six decimals at most, rounded outward — the exact quotient is in the codes."""
    exponent = d.as_tuple().exponent
    if isinstance(exponent, int) and exponent < -6:
        d = d.quantize(Decimal("0.000001"), rounding="ROUND_CEILING" if up else "ROUND_FLOOR")
    return _plain(d)


def _review_base(req: dict[str, Any], line: str) -> dict[str, Any]:
    return {"srs_id": req.get("id", ""), "source": "SRS", "kind": "held_back", "line": line,
            "line_sha256": hashlib.sha256(line.encode("utf-8")).hexdigest(), "fact_kind": "measurement",
            "line_span": [], "block_name": "", "fill_check": "not_applicable"}


def measurement_groups(req: dict[str, Any], stats: Counter, hw: dict[str, dict[str, Any]] | None = None,
                       hsis_rows: list[dict[str, Any]] | None = None,
                       review: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """(R51) One group per measurement requirement: a step per listed input value — the output's expected value and
    code, and the judgment band (HW measurement accuracy of the input's path + one resolution step) in value and code.
    The accuracy is R49/R50's reading of the HW blocks the requirement's system IDs link (the SW variable its text names
    marks the HSIS row); none linked or several left — the review sheet says so and the band says what it holds."""
    spec = measurement_spec(req)
    if not spec:
        return []
    unit = spec["unit"]
    step = Decimal(spec["resolution_exact"])
    signal = next(iter(_SW_VAR.findall(f"{req.get('description') or ''}\n{req.get('verification') or ''}")), "")
    basis = f"{BASIS_MARK}{' · '.join(spec['lines'])}"
    steps, points, summaries, first = [], [], [], []
    for v in (Decimal(x) for x in spec["values"]):
        probe = {"unit": unit, "value": _plain(v), "step": spec["resolution"], "source": None, "signal": signal}
        cands = hw_tolerances_for(req, probe, hw, hsis_rows) if hw is not None else []
        summary = hw_tolerance_summary(probe, cands)
        accuracy = Decimal(summary["largest"]) if summary and summary["largest"] is not None else None
        band = (accuracy or Decimal(0)) + step
        lo, hi = v - band, v + band
        if hw is None:
            source = "HW 측정 정확도 없음(HW 요구사항서 미입력)"            # (R51 review W4)
        elif summary and accuracy is None:
            # (R51 review W6) a block is linked, its tolerance on another scale (a divider unknown): no number
            source = f"HW 측정 정확도 척도 불명({' · '.join(summary['blocks'])} — 분압비 확인)"
        elif accuracy is None:
            source = "HW 측정 정확도 없음(이어진 HW 블록 없음)"
        else:
            blocks = " · ".join(summary["blocks"])
            source = (f"HW 측정 ±{_plain(accuracy)}{unit}({blocks}"
                      + ("" if summary["certain"] else " — 감시 경로 미정, 가장 큰 값") + ")")
        steps.append({"action": f"{MEASURE_PREFIX}{spec['input']} = {_plain(v)}{unit}{basis}",
                      "expected": (f"{spec['output']} = {_plain(v)}{unit}(코드 "
                                   f"0x{_measure_code(v, spec, 'ROUND_HALF_UP'):02X}) — 판정 범위 {_outward(lo, False)}{unit} ~ "
                                   f"{_outward(hi, True)}{unit}(코드 0x{_measure_code(lo, spec, 'ROUND_CEILING'):02X} ~ "
                                   f"0x{_measure_code(hi, spec, 'ROUND_FLOOR'):02X}): {source} + 분해능 "
                                   f"{spec['resolution']}{unit}(반올림 방식 미기재)")})
        points.append({"point": _plain(v), "holds": None})
        summaries.append(summary)
        if not first:
            first = cands[:MAX_HW_TOLERANCES]
    linked = [s for s in summaries if s and s["largest"] is not None]
    scale_unknown = not linked and any(s and s["largest"] is None for s in summaries)
    stats["measurement_tcs"] += 1
    stats["measurement_steps"] += len(steps)
    stats["measurement_hw_linked"] += bool(linked)
    stats["measurement_path_undecided"] += any(not s["certain"] for s in linked)
    stats["measurement_hw_not_given"] += hw is None                   # (R51 review W4)
    stats["measurement_scale_unknown"] += scale_unknown               # (R51 review W6)
    if review is not None:
        base = _review_base(req, spec["lines"][2])
        if spec["resolution_written"]:
            review.append({**base, "reason": "measurement_resolution_written_otherwise",
                           "fact": f"Resolution : {spec['resolution_written']} → {spec['resolution']}{unit}",
                           "if_filled": f"분해능을 '{spec['resolution']}{unit}' 로 적으면 표기가 범위·코드와 맞는다(지금도 "
                                        "그 값으로 판정 — 문서 표기만 바뀐다)"})
        if hw is None:
            review.append({**base, "reason": "measurement_hw_not_given", "fact": spec["output"],
                           "if_filled": "HW 요구사항서(HwRS·HRS)를 주면 이 입력을 감시하는 HW 블록의 허용 오차가 판정 범위에 "
                                        "들어간다(지금은 분해능만 — 실측이 벗어날 수 있다)"})
        elif scale_unknown:
            review.append({**base, "reason": "measurement_scale_unknown", "fact": spec["output"],
                           "if_filled": "그 HW 블록의 감시 노드 분압비를 HW 설계서(HwDS)나 HW 요구 블록에 적으면(노드 = 입력 × "
                                        "k) 허용오차가 입력 척도로 환산돼 판정 범위에 들어간다"})
        elif not linked:
            review.append({**base, "reason": "measurement_tolerance_unlinked", "fact": spec["output"],
                           # (review r2 W-A) the HSIS only picks among blocks the Related IDs already link — no
                           #   promise of it here (tried: a HSIS row with the monitor's ID and net linked nothing)
                           "if_filled": "이 입력을 감시하는 HW 요구 블록이 인용하는 시스템 ID 를 이 요구의 Related ID 에 "
                                        "적으면(또는 그 HW 블록이 이 요구의 시스템 ID 를 인용하면) 판정 범위에 HW 측정 정확도가 "
                                        "들어간다(지금은 분해능만 — 실측이 벗어날 수 있다)"})
        elif any(not s["certain"] for s in linked):
            review.append({**base, "reason": "measurement_path_undecided", "fact": spec["output"],
                           "if_filled": "이 입력의 HSIS 신호 행이 감시 HW 블록 하나를 가리키면(시스템 ID 또는 네트) 판정 "
                                        "범위가 그 블록의 허용오차로 좁혀진다"})
    line = spec["lines"][2]
    return [{"steps": steps, "evidence": {
        "srs_id": req.get("id", ""), "line": line, "line_sha256": hashlib.sha256(line.encode("utf-8")).hexdigest(),
        "span": [], "kind": "measurement", "signal": signal or spec["input"], "reference_constant": None,
        "signal_kind": "measurement", "op": "reports", "value": ", ".join(spec["values"]), "unit": unit,
        "step": spec["resolution"],
        "step_basis": f"변환 0x{spec['code_low']:02X}~0x{spec['code_high']:02X} = {spec['low']}~{spec['high']}{unit}",
        "combination_note": "", "points": points, "source": None,
        # (R51 review W4) no HW requirements given: no key — the cell says 'not given', not 'none linked'
        **({"hw_tolerance": first} if hw is not None else {}),
        "measurement": {k: v for k, v in spec.items() if k != "lines"}}}]


def append_requirement_boundary_tcs(test_cases: list[dict], requirements: list[dict], build_tc, make_tc_id,
                                    classify, max_steps: int = 12,
                                    system: dict[str, dict[str, Any]] | None = None,
                                    review_out: list[dict[str, Any]] | None = None,
                                    hw: dict[str, dict[str, Any]] | None = None,
                                    hsis_rows: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    """Append requirement-boundary TCs after each requirement's existing TCs, numbering on from them. A fact's points
    never split across TCs; a TC holds as many whole facts as fit in ``max_steps`` (at least one), all from the same
    document (the SRS, or the system requirements it cites — R29). Returns the counts for the quality report.
    (R45) ``review_out`` receives **every** review item for the document's 'Requirement Review' sheet — the held-back
    facts a reader can settle (one row per sentence and fact, however many requirements cite that block) and the
    requirement-document inclusion conflicts; the report keeps the first `MAX_REVIEW_ITEMS` and the counts.
    (R49) ``hw`` — the parsed HW requirement blocks (`load_hw_requirements`): each group gets its linked tolerances
    (`hw_tolerances_for`), and a TC whose step spacing lies inside one says so in its precondition. (R50) ``hsis_rows``
    — `hsis_rows_from_signals`; ``None`` when HSIS was not given (``[]``: given, no SW-variable row)."""
    stats: Counter = Counter()
    fanout: Counter = Counter()
    review: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    conflict_errors: list[str] = []
    differences: list[dict[str, Any]] = []
    difference_errors: list[str] = []
    compared_blocks = 0
    last: dict[str, int] = {}
    for tc in test_cases:
        rid = str(tc.get("srs_id") or "")
        tail = str(tc.get("id") or "").rsplit("_", 1)[-1]
        if rid and tail.isdigit():
            last[rid] = max(last.get(rid, 0), int(tail))
    added: list[dict] = []
    for req in requirements:
        groups = boundary_steps(req, stats, system, review)
        if hw is not None:
            for g in groups:
                if g["evidence"]["kind"] == "acceptance":
                    continue                     # (R51) a pass criterion has no point a HW tolerance could hide
                tolerances = hw_tolerances_for(req, g["evidence"], hw, hsis_rows)
                summary = hw_tolerance_summary(g["evidence"], tolerances)       # (review r2 I1) before the cut
                g["evidence"]["hw_tolerance"] = tolerances[:MAX_HW_TOLERANCES]
                if summary:
                    g["evidence"]["hw_tolerance_summary"] = summary
                    stats["hw_tolerance_groups"] += 1
                    stats["hw_tolerance_inside_step"] += summary["inside"]
                    stats["hw_tolerance_path_undecided"] += not summary["certain"]      # (review C2)
                    stats["hw_tolerance_scale_unknown"] += summary["largest"] is None  # (review W1)
                    stats["hw_tolerance_path_by_hsis"] += bool(summary["by_hsis"])     # (R50)
                    stats["hw_tolerance_scaled"] += bool(summary["scaled"])
                    # (R50 review) what the HSIS and the ratios could not settle — shown, not decided
                    refused = summary["path_refused"] or {}
                    stats["hw_tolerance_path_refused"] += refused.get("kind") == "name"
                    stats["hw_tolerance_path_outside"] += refused.get("kind") == "outside"     # (review r2 W2)
                    # (review r3 W1) a group with several of these is one boundary to settle, never two
                    stats["hw_tolerance_to_settle"] += bool(refused or summary["net_mismatch"]
                                                            or summary["scale_conflict"])
                    if not summary["certain"] and any(t["direct"] for t in tolerances):
                        # (review r3 W3) the fact's own system block is cited by several HW blocks: HSIS is not applied
                        stats["hw_tolerance_undecided_direct_hub"] += 1
                    stats["hw_tolerance_scale_unconfirmed"] += bool(summary["scale_unconfirmed"])
                    stats["hw_tolerance_scale_net_mismatch"] += bool(summary["net_mismatch"])
                    stats["hw_tolerance_scale_conflict"] += bool(summary["scale_conflict"])
                    if hsis_rows is not None and not summary["certain"] and not summary["path_refused"] and \
                            not any(t["direct"] for t in tolerances):       # (review r2 I3) HSIS does not decide there
                        # (review I8) why HSIS did not narrow: the subject is no SW variable, or its row is not there
                        signal = str(g["evidence"].get("signal") or "")
                        if not _SW_VAR.fullmatch(signal):
                            stats["hw_tolerance_undecided_not_sw"] += 1
                        elif len({id(r) for r, _v in _hsis_matches(signal, hsis_rows)}) > 1:
                            stats["hw_tolerance_undecided_row_ambiguous"] += 1           # (review r2 I7)
                        elif _hsis_row(signal, hsis_rows) is None:
                            stats["hw_tolerance_undecided_no_row"] += 1
                        else:   # the row names no ID or net only one candidate has
                            stats["hw_tolerance_undecided_row_silent"] += 1
        # (R51) a measurement requirement's group — after the loop above: it reads the HW accuracy itself (its value is a
        #   list of inputs, not one point)
        groups += measurement_groups(req, stats, hw, hsis_rows, review)
        if system is not None:
            compared_blocks += sum(1 for sid in cited_system_ids(req) if sid in system)
            try:
                conflicts += inclusion_conflicts(req, system)
            except Exception as exc:  # noqa: BLE001 — an optional finding never costs the boundary TCs (review I8)
                conflict_errors.append(f"{req.get('id', '')}: {type(exc).__name__}: {exc}"[:200])
            try:
                differences += value_differences(req, system)      # (R48)
            except Exception as exc:  # noqa: BLE001 — as the conflicts
                difference_errors.append(f"{req.get('id', '')}: {type(exc).__name__}: {exc}"[:200])
        chunks: list[list[dict]] = []
        for g in groups:
            # a fact with conditions to hold is its own TC: its precondition must not contradict another fact's steps
            #   (review r3 W2); facts without such notes share TCs up to ``max_steps``
            alone = bool(g["evidence"]["combination_note"]) or g["evidence"]["kind"] == "measurement"
            # (R47 review W3) an input range never joins a behaviour sentence's TC: a TC numbered before R47 keeps
            #   its steps (it did only by the data's luck — a note or a full last TC), and one TC says one kind of thing
            if not alone and chunks and not chunks[-1][0]["evidence"]["combination_note"] and \
                    _tc_class(chunks[-1][0]) == _tc_class(g) and chunks[-1][0]["evidence"]["kind"] != "measurement" and \
                    _doc_of(chunks[-1][0]) == _doc_of(g) and _is_input(chunks[-1][0]) == _is_input(g) and \
                    sum(len(x["steps"]) for x in chunks[-1]) + len(g["steps"]) <= max(max_steps, 1):
                chunks[-1].append(g)
            else:
                chunks.append([g])
        rid = req["id"]
        for chunk in chunks:
            steps = [s for g in chunk for s in g["steps"]]
            last[rid] = last.get(rid, 0) + 1
            method, gen, review_only = classify(steps)
            tc = build_tc(tc_id=make_tc_id(rid, last[rid]), req=req, steps=steps, test_method=method, gen_method=gen,
                          review_only=review_only)
            tc["requirement_evidence"] = [g["evidence"] for g in chunk]
            notes = list(dict.fromkeys(g["evidence"]["combination_note"] for g in chunk
                                       if g["evidence"]["combination_note"]))
            # (R49) — (R51) a measurement group's band is in its expected results; its value is a list, not one point
            notes += list(dict.fromkeys(n for g in chunk if g["evidence"]["kind"] not in ("measurement", "acceptance")
                                        and (n := _tolerance_note(g["evidence"]))))
            if notes:
                tc["precondition"] = "\n".join([p for p in [tc.get("precondition") or ""] if p] + notes)
            tc["requirement_boundary"] = True
            added.append(tc)
            if _tc_class(chunk[0]) == "acceptance":
                stats["acceptance_tcs"] += 1       # (review r2 W-D) not a boundary TC — counted apart
            elif _tc_class(chunk[0]) == "boundary":
                stats["traced:tcs" if chunk[0]["evidence"]["source"] else "tcs"] += 1
            if _is_input(chunk[0]):
                stats["traced:input_range_tcs"] += 1
            for g in chunk:
                if g["evidence"]["source"] and _tc_class(g) == "boundary":
                    fanout[g["evidence"]["source"]["id"]] += 1
    if added:   # nothing added: the TC order is left exactly as it was (review r4 I-4)
        order = {r["id"]: i for i, r in enumerate(requirements)}
        test_cases.extend(added)
        test_cases.sort(key=lambda tc: order.get(str(tc.get("srs_id") or ""), len(order)))  # stable
    if hw is not None:     # (R49) 0 is "looked, none linked" — absent is "not given"
        for k in ("hw_tolerance_groups", "hw_tolerance_inside_step", "hw_tolerance_path_undecided",
                  "hw_tolerance_scale_unknown", "hw_tolerance_path_by_hsis", "hw_tolerance_scaled",
                  "hw_tolerance_path_refused", "hw_tolerance_path_outside", "hw_tolerance_scale_unconfirmed",
                  "hw_tolerance_scale_net_mismatch", "hw_tolerance_scale_conflict", "hw_tolerance_to_settle",
                  "hw_tolerance_undecided_direct_hub"):
            stats[k] += 0
        if hsis_rows is not None:
            for k in ("hw_tolerance_undecided_not_sw", "hw_tolerance_undecided_no_row",
                      "hw_tolerance_undecided_row_silent", "hw_tolerance_undecided_row_ambiguous"):
                stats[k] += 0
    out: dict[str, Any] = {k: v for k, v in sorted(stats.items())}
    # (R45) one row per held-back sentence and fact: a block cited by several requirements names them all
    merged: dict[tuple, dict[str, Any]] = {}
    for item in review:
        # a held-back fact is its own row; a reading is one row per sentence (its facts listed)
        span = tuple(item["line_span"]) if item["kind"] == "held_back" else None
        key = (item["source"], item["line_sha256"], span, item["reason"])
        row = merged.setdefault(key, {**item, "srs_ids": [], "facts": []})
        if item["srs_id"] not in row["srs_ids"]:
            row["srs_ids"].append(item["srs_id"])
        if item["fact"] not in row["facts"]:
            row["facts"].append(item["fact"])
            row["fact"] = " · ".join(row["facts"])
    for row in merged.values():
        del row["facts"]
    review_items = [{k: v for k, v in r.items() if k != "srs_id"} for r in merged.values()]
    review_items.sort(key=lambda r: (not str(r["srs_ids"][0]).startswith("SwTSR"), str(r["srs_ids"][0]), r["source"]))
    # held back first (what the documents leave undecided), then what was read (to confirm)
    review_items.sort(key=lambda r: r["kind"] != "held_back")
    out["review_item_count"] = len(review_items)
    out["review_read_count"] = sum(1 for r in review_items if r["kind"] == "read")
    out["review_by_reason"] = dict(sorted(Counter(r["reason"] for r in review_items).items()))
    try:
        # (R46) the other sentences of the given documents that state an item's value steppably, and what filling
        #   the gap would improve
        out.update(attach_review_evidence(review_items, requirements, system))
    except Exception as exc:  # noqa: BLE001 — an optional finding never costs the boundary TCs (as review I8)
        out["review_evidence_error"] = f"{type(exc).__name__}: {exc}"[:200]
        for r in review_items:
            if r["kind"] == "held_back":
                r["evidence_error"] = out["review_evidence_error"]   # (review I4) the sheet says so, not "—"
    out["review_items"] = review_items[:MAX_REVIEW_ITEMS]
    if review_out is not None:
        review_out.extend(review_items)
    if fanout:
        # (review W5) one system block cited by many SRS requirements is stepped under each: say which blocks carry the
        #   traced facts (the most first, ties by ID) — HDPDM01 ``SySM_04`` alone carried 35 of 74
        out["traced_facts_by_block"] = dict(sorted(fanout.items(), key=lambda kv: (-kv[1], kv[0]))[:10])
        out["traced_blocks_used"] = len(fanout)
    if system is not None:
        # (R33) requirement-document inconsistencies found on the way — counted in full, the first ones itemised: safety
        #   requirements first (review I7: SwTSR_0104 fell behind "외 5건"), then by requirement ID
        # (R43 review W2) within a requirement, pairs across documents first — the disclosure shows five
        conflicts.sort(key=lambda c: (not str(c["srs_id"]).startswith("SwTSR"), str(c["srs_id"]),
                                      bool(c.get("within_source"))))
        # (R43 review W3) one block's own two lines are one finding, however many requirements cite that block (the hub
        #   block SySM_04 carries 35 of HDPDM01's 74 traced facts) — kept under the first requirement in the SORTED list
        #   (round 2 W3), so a safety requirement citing the block keeps it in front
        block_pairs: set = set()
        kept: list[dict[str, Any]] = []
        for c in conflicts:
            if c.get("within_source") and c["a"]["source"] != "SRS":
                bkey = (c["a"]["source"].split(" · ")[0], c["unit"], Decimal(c["value"]), c["role"],
                        frozenset((c["a"]["op"], c["b"]["op"])))
                if bkey in block_pairs:
                    continue
                block_pairs.add(bkey)
            kept.append(c)
        conflicts = kept
        if review_out is not None:
            review_out.extend({**c, "kind": "inclusion_conflict"} for c in conflicts)   # the document lists them all
        out["inclusion_conflicts"] = len(conflicts)
        # (R43) of them, two lines of one source — read both first: they may be two conditions or actions on purpose
        out["inclusion_conflicts_within_source"] = sum(1 for c in conflicts if c.get("within_source"))
        out["inclusion_conflict_items"] = conflicts[:MAX_INCLUSION_CONFLICTS]
        out["inclusion_conflict_requirements"] = len({c["srs_id"] for c in conflicts})
        out["inclusion_conflict_values"] = len({(c["srs_id"], Decimal(c["value"]), c["unit"]) for c in conflicts})
        # (review W4) how many cited blocks were there to compare with — 0 means nothing was compared, not "none found"
        out["inclusion_blocks_compared"] = compared_blocks
        if conflict_errors:
            out["inclusion_conflict_errors"] = conflict_errors[:5]
            out["inclusion_conflict_error_count"] = len(conflict_errors)
        # (R48) value differences — safety requirements first, as the conflicts; the document lists them all
        differences.sort(key=lambda d: (not str(d["srs_id"]).startswith("SwTSR"), str(d["srs_id"])))
        if review_out is not None:
            review_out.extend({**d, "kind": "value_difference"} for d in differences)
        out["value_differences"] = len(differences)
        out["value_difference_items"] = differences[:MAX_VALUE_DIFFERENCES]
        out["value_difference_requirements"] = len({d["srs_id"] for d in differences})
        out["value_difference_by_tier"] = {t: n for t in _DIFF_TIERS
                                           if (n := sum(1 for d in differences if t in d["tiers"]))}
        if difference_errors:
            out["value_difference_errors"] = difference_errors[:5]
            out["value_difference_error_count"] = len(difference_errors)
    return out


REQUIREMENT_EVIDENCE_HEADERS = ["Test Case ID", "SRS ID", "Kind", "Subject", "Reference Constant", "Operator", "Value",
                                "Unit", "Step", "Step Basis", "Stimulus Points (holds)", "Other Conditions",
                                "Source Line", "Line SHA-256", "Source Document",
                                # (R49) the HW monitor accuracy linked through the Related IDs — quoted, never used
                                "HW Tolerance (Quoted)",
                                # (R54) where the path is undecided: an LLM's block, with a quote of that block naming
                                #   the subject — checked, never used
                                "AI Path Proposal (Checked)"]
REQUIREMENT_EVIDENCE_SHEET = "Requirement Evidence"


def write_requirement_evidence_sheet(wb, test_cases: list[dict]) -> int:
    """'Requirement Evidence' sheet: which requirement sentence each boundary step comes from. Returns rows written.
    A sheet of that name already in the workbook (a template made from an earlier output) is removed first — also when
    this document has no boundary TC, so an old sheet never describes a document it does not belong to."""
    if REQUIREMENT_EVIDENCE_SHEET in wb.sheetnames:
        del wb[REQUIREMENT_EVIDENCE_SHEET]
    rows = [(tc["id"], e) for tc in test_cases for e in tc.get("requirement_evidence") or []]
    if not rows:
        return 0
    ws = wb.create_sheet(REQUIREMENT_EVIDENCE_SHEET)
    ws.append(REQUIREMENT_EVIDENCE_HEADERS)
    for tc_id, e in rows:
        # (R51) a point with no verdict (a measurement's listed value) is no 'F'
        points = ", ".join(f"{p['point']}({'T' if p['holds'] else '—' if p['holds'] is None else 'F'})"
                           for p in e["points"])
        subject = e.get("signal") or "—"
        if e.get("signal_kind") == "paren_owner":
            subject += " (괄호 앞 명사)"   # (R32 review I4) a weaker subject than a signal name: the reviewer sees it
        ws.append([tc_id, e["srs_id"], e["kind"], subject, e.get("reference_constant") or "—", e["op"],
                   e["value"], e["unit"] or "—", e["step"], e["step_basis"], points, e["combination_note"] or "—",
                   _clip(e["line"], 300), e["line_sha256"], source_label(e.get("source")), _tolerance_text(e),
                   _path_proposal_text(e)])
    return len(rows)


REQUIREMENT_REVIEW_HEADERS = ["Kind", "SRS ID", "Source Document", "Fact", "Reason", "To Decide",
                              # (R46) what writing the gap into the requirement produces · other sentences of the given
                              #   documents that state the value steppably (quoted, never used)
                              "If Filled", "Evidence (Other Sentences)", "Source Line", "Line SHA-256",
                              # (R52) an LLM's subject among real names, with a quote writing the name and the value —
                              #   checked, never used
                              "AI Proposal (Checked)"]
REQUIREMENT_REVIEW_SHEET = "Requirement Review"
_CONFLICT_ROLE_TEXT = {"condition": "조건끼리", "outcome": "결과 기준끼리", "stimulus": "시험 입력(검증 기준) ↔ 요구 조건"}
_DIFF_TIER_TEXT = {"same_subject": "같은 주어", "same_condition": "같은 조건의 유지시간", "within_step": "한 눈금 안",
                   "only_pair": "그 단위·쪽의 유일한 짝"}


def _path_proposal_text(evidence: dict[str, Any]) -> str:
    """(R54) The 'AI Path Proposal (Checked)' cell — ``—`` when the path was decided or not asked about."""
    if not evidence.get("ai_path"):
        return "—"
    from generators.sts_ai_review import path_proposal_text
    return _clip(path_proposal_text(evidence), 600)


def _ai_proposal_text(item: dict[str, Any]) -> str:
    """(R52) The 'AI Proposal (Checked)' cell — ``—`` when the item was not a subject question."""
    if not item.get("ai_proposal"):
        return "—"
    from generators.sts_ai_review import proposal_text
    return _clip(proposal_text(item), 600)


def write_requirement_review_sheet(wb, items: list[dict[str, Any]] | None) -> int:
    """(R45) 'Requirement Review' sheet: what the requirement documents leave undecided or state two ways — the facts
    held back from boundary steps for a reason a reader can settle (`REVIEW_TEXT`), what the generator read where the
    words are silent (`READ_TEXT`, stepped — to confirm), the inclusion conflicts (R33/R43) and the value differences
    (R48).
    Nothing here is filled by the generator: each row says what to decide. A sheet of that name from a template made of
    an earlier output is removed first; no item — no sheet (the quality report's counts say whether it was looked at).
    Returns rows written."""
    if REQUIREMENT_REVIEW_SHEET in wb.sheetnames:
        del wb[REQUIREMENT_REVIEW_SHEET]
    if not items:
        return 0
    ws = wb.create_sheet(REQUIREMENT_REVIEW_SHEET)
    ws.append(REQUIREMENT_REVIEW_HEADERS)
    for it in items:
        if it.get("kind") == "value_difference":
            srs = it.get("srs") or {}
            others = [o for o in (it.get("others") or []) if isinstance(o, dict)]

            def said(x):
                ref = f" (기준 {x.get('reference')})" if x.get("reference") else ""
                return f"{x.get('subject')} {x.get('text')}{ref}"
            more = int(it.get("other_count") or len(others)) - len(others)
            first = others[0] if others else {}
            tiers = " · ".join(_DIFF_TIER_TEXT[t] for t in (first.get("tiers") or []) if t in _DIFF_TIER_TEXT)
            # (review I2) the first counterpart's: the one the row leads with
            flag = (" · 한쪽 주어 없음" if it.get("subject_missing") else
                    "" if "same_subject" in (first.get("tiers") or []) else " · 주어 이름이 다름(같은 신호인지 먼저 확인)")
            held = first.get("held_subjects") or [None, None]
            if "same_condition" in (first.get("tiers") or []) and all(held) and held[0] != held[1]:
                flag += f" · 유지하는 조건의 주어 이름이 다름({held[0]} ↔ {held[1]} — 같은 조건인지 먼저 확인)"
            if others and it.get("time"):
                # (R48) a hold time: the direction says which reading to check first
                flag += (" · 유지시간: SRS 가 더 김 — 시스템 요구 시간 안에 판정하지 못할 수 있음"
                         if others[0].get("srs_is") == "larger" else
                         " · 유지시간: SRS 가 더 짧음 — 시스템 시간을 나눠 가진 것(분해)일 수 있음")
            ws.append(["요구 문서 값 차이 후보", it.get("srs_id", ""),
                       "SRS ↔ " + ", ".join(dict.fromkeys(str(o.get("source")) for o in others)),
                       f"{said(srs)} ↔ " + "; ".join(said(o) for o in others) + (f" 외 {more}" if more > 0 else ""),
                       f"같은 단위·같은 쪽({'하한' if it.get('side') == 'lower' else '상한'})의 조건을 다른 값으로 적음 — "
                       f"짝지은 근거: {tiers}" + flag,
                       "같은 임계(한 조건)인지 먼저 확인 — 같으면 어느 값이 맞는지 문서 검토로 정하고, 같은 양의 다른 임계(두 "
                       "단계 등)면 결함 아님으로 닫는다(경계 TC 는 각 문장대로 판정)",
                       if_filled(it), "—",
                       _clip(" ↔ ".join([str(srs.get("line") or "")] + [str(o.get("line") or "") for o in others]), 300),
                       "—", "—"])
            continue
        if it.get("kind") == "inclusion_conflict":
            a, b = it.get("a") or {}, it.get("b") or {}
            flag = (" · 한 출처 안의 두 줄(다른 조건·동작일 수 있음)" if it.get("within_source") else "") + (
                "" if it.get("same_subject") else " · 한쪽 주어 없음" if it.get("subject_missing") else " · 주어 이름이 다름")
            ws.append(["요구 문서 경계 포함 불일치", it.get("srs_id", ""), f"{a.get('source')} ↔ {b.get('source')}",
                       f"{a.get('subject')} {a.get('text')} ↔ {b.get('subject')} {b.get('text')}",
                       f"같은 값·단위를 경계 포함만 달리 적음({_CONFLICT_ROLE_TEXT.get(str(it.get('role')), it.get('role'))})"
                       + flag,
                       "어느 쪽이 맞는지 문서 검토로 정한다(경계 TC 는 각 문장대로 판정)",
                       if_filled(it), "—",
                       _clip(f"{a.get('line')} ↔ {b.get('line')}", 300), "—", "—"])
            continue
        reason = str(it.get("reason"))
        why, decide = {**REVIEW_TEXT, **READ_TEXT, **MEASURE_TEXT}.get(reason, (reason, "원문 확인"))
        if it.get("fact_kind") == "range" and reason in RANGE_DECIDE:
            decide = RANGE_DECIDE[reason]      # (review r2 W-R2-1) a range makes no boundary step, subject or not
        if it.get("block_name"):
            why += f" (블록 이름 '{_clip(str(it['block_name']), 80)}')"
        ws.append(["읽은 결합 — 확인" if it.get("kind") == "read" else "검토 필요 조건",
                   ", ".join(it.get("srs_ids") or []), it.get("source", ""), it.get("fact") or "—",
                   why, decide, if_filled(it), _clip(_evidence_text(it), 1000),
                   _clip(str(it.get("line") or ""), 300), it.get("line_sha256") or "—", _ai_proposal_text(it)])
    return len(items)
