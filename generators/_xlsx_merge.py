"""새로 쓴 영역의 셀 병합 — `ws.merge_cells` 의 O(기존 병합 수) 검사를 건너뛴다.

`Worksheet.merge_cells` 는 병합을 하나 넣을 때마다 `MultiCellRange.add` 가 **기존 병합 전부**를 훑어
포함 관계를 본다. 시험 규격서는 TC 마다 열 수만큼 병합하므로 문서 전체로는 O(n²)이다.

실측(KJPDS02 STS, 같은 PC): TC 297 → 5.8초 · 1,044 → 36.8초 · 2,687 → **264.5초**(병합 11,489회에 범위
비교 6,600만 회). 요구당 상한을 크게 올린 라이브 생성은 27분이 지나도 끝나지 않았다 — 준비 게이트가
"전량을 담으려면 N" 을 제안하는데 그 값을 넣으면 생성이 사실상 멈춘다는 뜻이다.
"""
from __future__ import annotations

from typing import Any


def merge_fresh(ws: Any, start_row: int, start_column: int, end_row: int, end_column: int) -> None:
    """`ws.merge_cells(start_row=…)` 과 같은 결과를 O(1) 로. 뒤집힌 범위는 openpyxl 이 ValueError 로 거부한다(가드가 본다).

    ⚠ 호출 조건: 이 범위가 **기존 병합과 겹치지 않음을 호출자가 보장**한다 — 새로 만든 시트에 위에서
      아래로 한 번씩 쓰는 데이터 영역이 그렇다. 템플릿의 병합이 남아 있을 수 있는 머리 영역에는 쓰지 말 것
      (겹친 병합은 Excel 이 "복구" 를 묻는 손상 파일이 된다).
    """
    from openpyxl.worksheet.cell_range import CellRange
    from openpyxl.worksheet.merge import MergedCellRange

    coord = CellRange(min_col=start_column, min_row=start_row, max_col=end_column, max_row=end_row).coord
    mcr = MergedCellRange(ws, coord)
    ws.merged_cells.ranges.add(mcr)     # `MultiCellRange.add` 의 포함 검사(전체 순회)를 건너뛴다
    ws._clean_merge_range(mcr)          # 좌상단 외 셀을 MergedCell 로 바꾸고 테두리를 되살린다(merge_cells 와 같은 후처리)
