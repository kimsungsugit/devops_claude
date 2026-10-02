"""(R53) STS 생성 결과의 요구 문서 발견 사항을 RAG 지식베이스(`reports/kb_store`)에 적재한다.

무엇을 — 생성된 STS xlsm 의 두 시트:
  * 'Requirement Review' — 요구 문서가 정하지 않았거나 두 가지로 적은 것(검토 필요 조건 — 측정 요구 · 판정 기준의 검토도
    이 종류 · 읽은 결합 · 경계 포함 불일치 · 값 차이 후보). 원문 문장 · 사유 · 정할 것 · 채우면 생기는 것 그대로.
  * 'Requirement Evidence' — 요구 ID 별 경계 근거(주어 · 비교 · 값 · 눈금 · 출처 문서 · HW 허용오차 경로), 그리고 요구 ID 별
    HW 측정 허용오차 경로(그 HW 블록 이름으로 찾히게 따로).

어디로 — 분류 `sts_findings` 하나(리뷰 W1: `requirements` 는 로컬 UDS 생성 AI 가 프로젝트 구분 없이 프롬프트에 넣는다).
  채팅 어시스턴트는 질문이 요구·HW ID 를 적으면 그 ID 의 항목을 공유 KB 에서 찾아 자기 컨텍스트 블록에 넣는다
  (`workflow/retrieval/hybrid.requirement_findings_hits` · `backend/services/assistant_service._build_context`).

어떻게 — 원문·사유는 시트에 적힌 그대로(LLM 재서술 없음). 키는 프로젝트 · 종류 · 제목 · **내용**의 sha256(리뷰 W2: 내용이
  바뀌면 새 키 — 옛 항목은 지워진다). 다시 돌리면 같은 키는 건너뛰고(임베딩 비용 0), 새 키만 더하고, 이 프로젝트의 이전 적재
  중 이번에 없는 키(같은 키의 중복 포함 — W6)는 지운다. 시트가 둘 다 없거나 파일 이름이 프로젝트와 다르면 지우기 전에 멈춘다
  (W3 — 정본 STS 를 잘못 주면 프로젝트 항목이 전부 지워졌다). sqlite 저장소만(W4 — pgvector 에서는 적재 목록을 메모리에서
  볼 수 없다). 긴 근거는 임베딩 상한(6000 자) 안으로 나눈다(W7). 쓴 뒤 저장소를 다시 읽어 실제 건수를 보고한다(W8).

사용:
  .venv/Scripts/python.exe scripts/sts_findings_to_kb.py --sts <STS.xlsm> --project HDPDM01 [--report-dir reports] [--dry-run]
    --force-project       파일 이름에 프로젝트 첫 토큰이 없어도 적재(같은 이름을 쓰는 DV/PV 는 구분하지 못한다)
    --allow-mass-removal  이 프로젝트 항목의 50% 넘게 지워지는 적재를 허용(생성기 형식이 바뀐 재적재 — dry-run 의
                          `would_refuse` 로 먼저 확인)
    --allow-empty         두 시트가 없거나 비어 발견 사항이 0 건인 산출물을 받아 이 프로젝트 항목을 비운다 — 발견 사항이 한 건이라도
                          있으면 대량 삭제 가드는 그대로
  종료 코드: 0 정상 · 2 거부(IngestRefused) · 3 저장소 재확인 불일치 또는 임베딩 열화
  ⚠ `scripts/reindex_kb.py` 와 동시에 돌리지 말 것 — 재색인이 먼저 읽은 항목을 다시 써서 방금 지운 항목이 되살아난다(다음
    적재가 다시 지우지만 그 사이 채팅이 인용한다).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

REPO = Path(__file__).resolve().parents[1]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

REVIEW_SHEET = "Requirement Review"
_NOT_INGESTED = frozenset({"Line SHA-256", "AI Proposal (Checked)"})
EVIDENCE_SHEET = "Requirement Evidence"
# (review W1) its own — read by the chat's requirement-findings domain only (one spelling: `config`)
CATEGORY = str(getattr(__import__("config"), "RAG_STS_FINDINGS_CATEGORY", "sts_findings"))
MAX_DOC_CHARS = 4500               # (review W7) under the embedder's 6000-character cut, with the title
_HW_COLUMN = "HW Tolerance (Quoted)"
# (review W9) an item head of the HW Tolerance cell: ``HwTSR_0204 Battery Voltage Monitor: 허용 오차 …`` — at the start
#   of the cell, after ``감시 경로 미정 — `` or after ``; `` (a bracketed ``[척도 환산 — HwC_07 …: …]`` is not one)
_HW_HEAD = re.compile(r"(?:^|; |감시 경로 미정 — )(?P<head>Hw[A-Za-z]{1,6}_\d+(?:_\d+)? [^:;\[\]]{2,60}?):")
_KIND_TAG = {"검토 필요 조건": "review_held_back", "읽은 결합 — 확인": "review_read",
             "요구 문서 경계 포함 불일치": "inclusion_conflict", "요구 문서 값 차이 후보": "value_difference"}


class IngestRefused(ValueError):
    """The input would delete what it should not (no sheet, another project's file) or the store cannot be listed."""


def _key(*parts: Any) -> str:
    return hashlib.sha256("\x1f".join(str(p) for p in parts).encode("utf-8")).hexdigest()[:20]


def _rows(wb, name: str) -> tuple[list[str], list[tuple]] | None:
    """The header and the non-empty rows of a sheet — ``None`` when the sheet is not there."""
    if name not in wb.sheetnames:
        return None
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return [], []
    return [str(h or "") for h in rows[0]], [r for r in rows[1:] if any(c not in (None, "") for c in r)]


def _chunks(lines: list[str], limit: int) -> list[list[str]]:
    out: list[list[str]] = [[]]
    size = 0
    for line in lines:
        if out[-1] and size + len(line) + 1 > limit:
            out.append([])
            size = 0
        out[-1].append(line[:limit])
        size += len(line) + 1
    return out


def _doc(project: str, group: str, title: str, content: str, tags: list[str]) -> dict[str, Any]:
    return {"key": _key(project, group, title, content), "title": title[:200], "content": content, "group": group,
            "tags": tags}


def build_documents(sts_path: str, project: str, *, allow_empty: bool = False) -> list[dict[str, Any]]:
    """The KB documents of one STS: ``{"key", "title", "content", "tags", "group"}`` — read only. Raises
    `IngestRefused` when neither sheet is there (an output without findings writes no sheet — say so with
    ``allow_empty``)."""
    from openpyxl import load_workbook
    wb = load_workbook(sts_path, read_only=True, data_only=True)
    try:
        review, evidence = _rows(wb, REVIEW_SHEET), _rows(wb, EVIDENCE_SHEET)
    finally:
        wb.close()
    if review is None and evidence is None and not allow_empty:
        raise IngestRefused(f"'{REVIEW_SHEET}' · '{EVIDENCE_SHEET}' 시트가 없다 — 생성된 STS 가 아니거나 발견 사항이 없는 산출물"
                            "(지우려면 --allow-empty)")
    docs: list[dict[str, Any]] = []
    head, rows = review or ([], [])
    for r in rows:
        cell = dict(zip(head, ("" if c is None else str(c) for c in r), strict=False))
        kind = cell.get("Kind", "")
        srs = [s.strip() for s in cell.get("SRS ID", "").split(",") if s.strip()]
        # (review I2) the sentence hash is no reading matter — kept out of the embedded text. (R52) Nor the AI column:
        #   the KB holds what the generator read from the documents, never a model's answer (cited back as evidence it
        #   would be AI on AI), and its '묻지 않음' would change every key with the AI setting
        content = "\n".join(f"{h}: {cell.get(h, '')}" for h in head
                            if h not in _NOT_INGESTED and cell.get(h, "") not in ("", "—"))
        group = _KIND_TAG.get(kind, "review_other")
        docs.append(_doc(project, group,
                         f"[STS 요구 검토] {project} {', '.join(srs[:3])} — {kind}: {cell.get('Fact', '')}",
                         f"프로젝트 {project} STS 생성 시 요구 문서에서 찾은 검토 항목.\n{content}"[:MAX_DOC_CHARS],
                         [project, "sts_review", group, *srs[:10]]))
    head, rows = evidence or ([], [])
    by_req: dict[str, list[dict[str, str]]] = {}
    for r in rows:
        cell = dict(zip(head, ("" if c is None else str(c) for c in r), strict=False))
        by_req.setdefault(cell.get("SRS ID", ""), []).append(cell)
    for srs, cells in sorted(by_req.items()):
        if not srs:
            continue
        lines = [" | ".join(f"{h}: {c[h]}" for h in head
                            if h not in ("SRS ID", "Line SHA-256") and c.get(h) not in ("", "—", None)) for c in cells]
        parts = _chunks(lines, MAX_DOC_CHARS - 200)
        for i, part in enumerate(parts, start=1):
            suffix = f" ({i}/{len(parts)})" if len(parts) > 1 else ""
            docs.append(_doc(project, "evidence", f"[STS 요구 경계 근거] {project} {srs}{suffix}",
                             f"프로젝트 {project} 요구 {srs} 의 STS 경계 근거(원문 인용 · 경계 점 · HW 허용오차).\n"
                             + "\n".join(part), [project, "sts_evidence", srs]))
        # the HW tolerance path of each boundary on its own: a question about a HW block finds it by the block's name
        hw = [c for c in cells if (c.get(_HW_COLUMN) or "—").strip()[:1] not in ("—", "")]
        if hw:
            blocks = list(dict.fromkeys(m.group("head").strip() for c in hw for m in _HW_HEAD.finditer(c[_HW_COLUMN])))
            undecided = any(c[_HW_COLUMN].startswith("감시 경로 미정") for c in hw)
            body = [f"{c.get('Subject', '')} {c.get('Operator', '')} {c.get('Value', '')}{c.get('Unit', '')} "
                    f"(한 눈금 {c.get('Step', '')}{c.get('Unit', '')}, {c.get('Test Case ID', '')}) → {c[_HW_COLUMN]}"
                    for c in hw]
            parts = _chunks(body, MAX_DOC_CHARS - 200)
            for i, part in enumerate(parts, start=1):
                suffix = f" ({i}/{len(parts)})" if len(parts) > 1 else ""
                docs.append(_doc(project, "hw_tolerance",
                                 f"[STS HW 측정 허용오차] {project} {srs}: {' · '.join(blocks[:3])}"
                                 + (" (감시 경로 미정)" if undecided else "") + suffix,
                                 f"프로젝트 {project} 요구 {srs} 의 경계 값을 재는 HW 감시 경로와 측정 허용오차(HW 요구사항서 "
                                 f"원문 인용).\n" + "\n".join(part), [project, "sts_hw_tolerance", srs]))
    # (review W6) two identical documents are one
    return list({d["key"]: d for d in docs}.values())


def check_project(sts_path: str, project: str) -> None:
    """(review W3) The file is this project's — its name carries the project's first token (``KJPDS02`` of
    ``KJPDS02_PV``); a PV STS given as HDPDM01 would otherwise replace HDPDM01's findings."""
    token = project.split("_")[0].upper()
    if token not in Path(sts_path).name.upper():
        raise IngestRefused(f"파일 이름 '{Path(sts_path).name}' 에 프로젝트 '{token}' 가 없다 — 다른 프로젝트의 STS 일 수 "
                            "있다(맞으면 --force-project)")


MASS_REMOVAL_SHARE = 0.5   # (review r2 W-B) removing more than this share of a project's entries needs a flag


def ingest(docs: list[dict[str, Any]], project: str, report_dir: Path, *, dry_run: bool = False,
           allow_mass_removal: bool = False) -> dict[str, Any]:
    """Add the new keys, keep the present ones, remove this project's stale (and duplicate) ones; re-read the store to
    report what is there. Returns the counts. Refuses (`IngestRefused`) to remove more than `MASS_REMOVAL_SHARE` of the
    project's entries without ``allow_mass_removal`` — ``KJPDS02`` (DV) and ``KJPDS02_PV`` share a file-name token. Only
    real removals count (review r3 W3): a degraded vector re-embedded under its own key and a duplicate are no loss of
    content; a dry run reports the refusal (``would_refuse``) instead of raising, so the plan can be seen first."""
    from workflow.rag import get_kb
    kb = get_kb(report_dir)
    if kb.storage != "sqlite":
        raise IngestRefused(f"KB 저장소가 {kb.storage} — 적재 목록을 읽을 수 없어 멱등·삭제가 동작하지 않는다(sqlite 만)")
    scope = f"sts_findings:{project}"
    wanted = {d["key"] for d in docs}
    existing: dict[str, list[tuple[str, str]]] = {}
    stale: list[str] = []
    removed_titles: list[str] = []
    held = 0
    for e in list(kb.data):
        tags = e.get("tags") or []
        if scope not in tags:
            continue
        held += 1
        key = next((t.split(":", 1)[1] for t in tags if str(t).startswith("sts_key:")), "")
        entry = (str(e.get("id")), str(e.get("error_raw") or ""))
        if ((e.get("metadata") or {}).get("embed") or {}).get("degraded"):
            # (review r2 W-C) a random fallback vector is re-embedded on the next run, not kept for good
            stale.append(entry[0])
            if key not in wanted:
                removed_titles.append(entry[1])
            continue
        existing.setdefault(key, []).append(entry)
    for key, entries in existing.items():
        if key not in wanted:
            stale += [i for i, _t in entries]
            removed_titles += [t for _i, t in entries]
        else:
            stale += [i for i, _t in entries[1:]]        # (review W6) a key twice: the extra one goes
    new = [d for d in docs if d["key"] not in existing]
    removals = len(removed_titles)
    refuse = ""
    if held >= 10 and removals > held * MASS_REMOVAL_SHARE:
        titles = {d["title"] for d in docs}
        same = sum(1 for t in removed_titles if t in titles)
        refuse = (f"이 프로젝트 항목 {held} 개 중 {removals} 개를 지우게 된다 — 그중 {same} 개는 같은 제목에 내용만 바뀌었고"
                  f"(생성 결과의 형식·근거 변경), {removals - same} 개는 이번 STS 에 없는 제목이다(다른 프로젝트·변형의 "
                  "STS 이거나 요구 문서가 크게 바뀜). 맞으면 --allow-mass-removal")
        if not allow_mass_removal and not dry_run:
            raise IngestRefused(refuse)
    out: dict[str, Any] = {"project": project, "documents": len(docs), "added": 0, "kept": len(docs) - len(new),
                           "removed": 0, "stale": len(stale), "removals": removals, "new": len(new), "dry_run": dry_run,
                           "kb_dir": str(kb.base_dir), "category": CATEGORY,
                           "by_group": dict(Counter(d["group"] for d in docs)),
                           "would_refuse": "" if allow_mass_removal else refuse}
    if dry_run:
        return out
    degraded = 0
    for d in new:
        before = len(kb.data)
        kb.add_document(d["title"], d["content"], category=CATEGORY,
                        tags=[*d["tags"], scope, f"sts_key:{d['key']}"], source_file=None)
        # (review W5) no source_file: `_persist_entry` (reindex) writes `{id}.json`; a ``sts:HD`` name was an NTFS stream
        added = kb.data[before:] if len(kb.data) > before else []
        degraded += sum(bool(((e.get("metadata") or {}).get("embed") or {}).get("degraded")) for e in added)
        out["added"] += bool(added)
    out["removed"] = kb.remove_documents(stale) if stale else 0
    out["embedding_degraded"] = degraded
    # (review W8) what the store holds now — a locked DB swallowed writes while the counts above said otherwise. Read
    #   the sqlite file itself (a new KnowledgeBase would also load the configured external sources)
    out["verified_in_store"] = _count_in_store(kb.db_path, scope)
    out["verified_matches"] = out["verified_in_store"] == len(wanted)
    return out


def _count_in_store(db_path: Path, scope: str) -> int | None:
    """The entries of ``scope`` the sqlite store holds — ``None`` when it cannot be read (said, never 0)."""
    import sqlite3
    try:
        conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True, timeout=10)
        try:
            rows = conn.execute("SELECT tags FROM kb_entries").fetchall()
        finally:
            conn.close()
    except sqlite3.Error:
        return None
    n = 0
    for (tags,) in rows:
        try:
            n += scope in (json.loads(tags) if isinstance(tags, str) else (tags or []))
        except (ValueError, TypeError):
            continue
    return n


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=(__doc__ or "").split("\n")[0])
    ap.add_argument("--sts", required=True)
    ap.add_argument("--project", required=True)
    ap.add_argument("--report-dir", default=str(REPO / "reports"))
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--allow-empty", action="store_true",
                    help="시트가 없는 STS 로 이 프로젝트 항목을 모두 지운다(대량 삭제 허용 포함)")
    # (review r3 W3) two protections, two flags — one flag turned both off
    ap.add_argument("--force-project", action="store_true", help="파일 이름의 프로젝트 대조를 건너뛴다")
    ap.add_argument("--allow-mass-removal", action="store_true",
                    help=f"이 프로젝트 항목의 {int(MASS_REMOVAL_SHARE * 100)}%% 넘게 지우는 적재를 허용한다")
    a = ap.parse_args(argv)
    project = a.project.strip().upper()          # (review I6) one scope per project, whatever the capitals
    try:
        if not project:
            raise IngestRefused("--project 가 비었다")
        if not a.force_project:
            check_project(a.sts, project)
        if not Path(a.sts).is_file():
            raise IngestRefused(f"STS 파일이 없다: {a.sts}")              # (review r2 I-3)
        docs = build_documents(a.sts, project, allow_empty=a.allow_empty)
        out = ingest(docs, project, Path(a.report_dir), dry_run=a.dry_run,
                     # (review r4 W2) --allow-empty lifts the mass-removal guard only for an output that has no
                     #   findings — on a workbook with sheets it would have let another project's file replace this one
                     allow_mass_removal=a.allow_mass_removal or (a.allow_empty and not docs))
        print(json.dumps(out, ensure_ascii=False, indent=2))
    except IngestRefused as exc:
        print(f"적재 중단: {exc}", file=sys.stderr)
        return 2
    # (review r2 W-C) what the store holds differs from what was meant, or a vector fell back: not a success
    if not out.get("dry_run") and (out.get("verified_matches") is False or out.get("embedding_degraded")):
        print("적재 확인 실패 — 저장소 건수 불일치 또는 임베딩 열화(다음 실행이 다시 적재한다)", file=sys.stderr)
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
