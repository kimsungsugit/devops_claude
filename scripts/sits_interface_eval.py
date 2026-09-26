"""SITS interface-fault discrimination (R8, G5) — the same annotation-based scoring for a reference and a generated SITS.

Neither suite is executed here (the reference SITS runs on the target / VectorCAST); both are judged on **what they
stimulate and observe**, against interface faults read from the current source by `generators/interface_contract.py`:

* ``return_not_propagated`` — at a cross-module call whose return value the caller uses (checked / assigned /
  returned), the callee's result is not acted on. A test tells it apart only if it **injects** that return (an input
  column ``callee() return`` / ``callee.return`` / ``callee return``) with at least two different values across its
  cases, in a test that exercises the caller and observes something;
* ``global_write_lost`` — a global written in one module and read in another along a flow; the producer's write is
  lost. A test tells it apart if it exercises the producer and **observes** that global (expected column);
* ``arguments_swapped`` — two arguments of the same declared type at a cross-module call are swapped. A test tells it
  apart if it exercises the caller, sets both argument sources (their base names are input columns) to different
  values in one case, and observes something.

**A test observes only what it writes an expectation for** (review round 1 C1): a row counts only where the observed
variable has a concrete expected value — ``[검증 필요]``, ``—``, ``N/A`` and blanks tell nothing apart.

A test *exercises* the functions on the call closure of the function it runs. Which one that is, is read from the
chain text the same way for both suites (review C2 — both write DFS call lists, not a single path): the primary rule is
the closure of the **first** function; ``rate_by_rule`` also reports the closure of the last function plus the chain, and the
chain alone, so the verdict's dependence on the rule is visible. Chain names not found in the source are counted.
These are necessary conditions for discrimination, not proof of it — reported as such.

Usage:
    .venv/Scripts/python.exe scripts/sits_interface_eval.py --source-root ROOT[,ROOT2] --reference REF.xlsm
        [--generated GEN.xlsm] --out r8.json
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

_SHEETS = ("4.SW Integration Test Spec", "3.SW Integration Test Spec")
_TC_ID = re.compile(r"^Sw[A-Za-z]*TC_\w+$")
_RETURN_COL = re.compile(r"^\s*([A-Za-z_]\w*)\s*(?:\(\s*\))?\s*(?:\.|\s)\s*return\s*$")
_IDENT = re.compile(r"[A-Za-z_]\w*")


def read_sits(path: str) -> list[dict[str, Any]]:
    """Tests of a SITS workbook: {tc_id, chain, inputs: [names], expected: [names], rows: [{name: value}]}."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        name = next((s for s in _SHEETS if s in wb.sheetnames), None)
        if name is None:
            raise ValueError(f"no SITS spec sheet in {Path(path).name}: {wb.sheetnames}")
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    finally:
        wb.close()
    head = next(i for i, r in enumerate(rows) if any(str(c or "").strip() == "Input" for c in r))
    in_col = next(j for j, c in enumerate(rows[head]) if str(c or "").strip() == "Input")
    exp_col = next(j for j, c in enumerate(rows[head]) if str(c or "").strip().startswith("Expected"))
    tests: list[dict[str, Any]] = []
    for r in rows[head + 2:]:
        r = r + [None] * (exp_col + 80 - len(r))
        first = str(r[1] or "").strip()
        if _TC_ID.match(first):
            inputs = [(j, str(r[j]).strip()) for j in range(in_col, exp_col) if r[j] not in (None, "")]
            expected = [(j, str(r[j]).strip()) for j in range(exp_col, len(r)) if r[j] not in (None, "")
                        and not _TC_ID.match(str(r[j]).strip())]
            tests.append({"tc_id": first, "chain": "", "title": str(r[2] or "").strip(), "inputs": inputs,
                          "expected": expected, "rows": [], "expected_rows": []})
            continue
        if not tests or all(c in (None, "") for c in r[in_col:]):
            continue
        t = tests[-1]
        chain_cell = str(r[3] or "").strip()
        if not t["chain"] and chain_cell and not re.fullmatch(r"[\d.\s]+", chain_cell):
            t["chain"] = chain_cell   # the first case row carries the full call chain (``main -> MainInit -> …``)
        t["rows"].append({n: r[j] for j, n in t["inputs"] if r[j] not in (None, "")})
        t["expected_rows"].append({n: r[j] for j, n in t["expected"] if r[j] not in (None, "")})
    for t in tests:
        if not t["chain"] and ("->" in t["title"] or "→" in t["title"]):
            t["chain"] = t["title"]   # a generated TC row's title (``Verify integration: entry → a → b``)
        t["inputs"] = [n for _j, n in t["inputs"]]
        t["expected"] = [n for _j, n in t["expected"]]
    return tests


_NO_EXPECTATION = re.compile(r"^\s*(?:\[\s*검증\s*필요\s*\].*|[—–-]|N/?A|TBD|\?+|don'?t\s*care|)\s*$",
                             re.IGNORECASE | re.DOTALL)


def concrete(value) -> bool:
    """A written expected value (``0``, ``0x10``, ``ERROR_OK (0)``) — not a placeholder (``[검증 필요]`` with its reason
    on the next line, a dash of any width, N/A, TBD, ``?``, don't care — round 2 Info 1)."""
    return value is not None and not _NO_EXPECTATION.match(str(value))


def _chain_functions(text: str) -> list[str]:
    # A leading label is not a function: ``Verify integration: a → b`` (generated) and ``Interface : main -> …``
    # (KJPDS02_PV reference, 47 of 54 tests — R17: reading it as part of the first name dropped the real entry function,
    # so the first-closure rule started one call deeper).
    text = re.sub(r"^\s*(?:Verify integration|Interface)\s*:\s*", "", text or "", flags=re.I)
    return [p.strip().rstrip("()").strip() for p in re.split(r"->|→|\n", text) if p.strip().rstrip("()").strip()]


def _closure(start: list[str], calls: dict[str, set[str]]) -> set[str]:
    seen, stack = set(), list(start)
    while stack:
        f = stack.pop()
        if f in seen:
            continue
        seen.add(f)
        stack.extend(calls.get(f, ()))
    return seen


def call_graph(index) -> dict[str, set[str]]:
    from generators.interface_contract import _callee_name, _walk
    graph: dict[str, set[str]] = {}
    # a cast ``(U16)(f())`` has no identifier callee: ``_callee_name`` gives "" and the inner call is its own node
    for name, (fn, raw, _path) in index.defs.items():
        body = fn.child_by_field_name("body")
        graph[name] = {c for n in (_walk(body) if body is not None else ()) if n.type == "call_expression"
                       and (c := _callee_name(n, raw)) and c in index.defs}
    return graph


def faults(index, graph, globals_: set[str]) -> list[dict[str, Any]]:
    """Every interface fault of the source: one contract over all functions (a flow is any caller → callee edge)."""
    from generators.interface_contract import flow_contract
    contract = flow_contract(sorted(index.defs), index, globals_)
    out = []
    for p in contract["calls"]:
        use = (p.get("return_use") or {}).get("use")
        if use in {"checked", "assigned", "returned"}:
            out.append({"kind": "return_not_propagated", "at": p["caller"], "callee": p["callee"], "line": p["line"],
                        "use": use})
        for a, b in p.get("swappable_pairs") or []:
            args = {x["param"]: x["arg"] for x in p.get("bindings") or []}
            out.append({"kind": "arguments_swapped", "at": p["caller"], "callee": p["callee"], "line": p["line"],
                        "params": [a, b], "args": [args.get(a, ""), args.get(b, "")]})
    consumers: dict[tuple[str, str], list[str]] = {}
    for g in contract["global_flows"]:
        consumers.setdefault((g["producer"], g["global"]), []).append(g["consumer"])
    for (producer, name), readers in sorted(consumers.items()):
        # one fault per write that another module reads — observing the global tells it apart whoever reads it
        out.append({"kind": "global_write_lost", "at": producer, "global": name, "consumers": sorted(set(readers))})
    # one fault per (kind, at, callee/global[, params]) — the same edge written twice is one interface
    seen, unique = set(), []
    for f in out:
        key = (f["kind"], f["at"], f.get("callee") or f.get("global"), tuple(f.get("params") or ()))
        if key not in seen:
            seen.add(key)
            unique.append(f)
    return unique


def _base(name: str) -> str:
    m = _IDENT.match(name.strip())
    return m.group(0) if m else ""


def discriminates(fault: dict, test: dict, exercised: set[str]) -> bool:
    """Necessary conditions: the test exercises the faulty function, stimulates the fault and writes an expectation
    that a changed behaviour would contradict (review C1)."""
    if fault["at"] not in exercised:
        return False
    judged = [i for i, exp in enumerate(test["expected_rows"]) if any(concrete(v) for v in exp.values())]
    if not judged:
        return False
    rows = [test["rows"][i] for i in judged if i < len(test["rows"])]
    if fault["kind"] == "return_not_propagated":
        cols = [n for n in test["inputs"] if (m := _RETURN_COL.match(n)) and m.group(1) == fault["callee"]]
        values = {str(row.get(c)) for c in cols for row in rows if concrete(row.get(c))}
        return len(values) >= 2
    if fault["kind"] == "global_write_lost":
        names = [n for n in test["expected"] if _base(n) == fault["global"]]
        return any(concrete(test["expected_rows"][i].get(n)) for i in judged for n in names)
    if fault["kind"] == "arguments_swapped":
        bases = [_base(a) for a in fault["args"]]
        if not all(bases) or bases[0] == bases[1]:
            return False
        cols = {b: [n for n in test["inputs"] if _base(n) == b] for b in bases}
        if not all(cols.values()):
            return False
        return any(any(concrete(row.get(c0)) and concrete(row.get(c1)) and str(row.get(c0)) != str(row.get(c1))
                       for c0 in cols[bases[0]] for c1 in cols[bases[1]]) for row in rows)
    return False


_RULES = ("first_closure", "last_closure_and_chain", "chain_only")


def _exercised(chain: list[str], graph, rule: str) -> set[str]:
    if rule == "first_closure":
        return _closure(chain[:1], graph)
    if rule == "last_closure_and_chain":
        return _closure(chain[-1:], graph) | set(chain)
    return set(chain)


def score(tests: list[dict], fault_list: list[dict], graph, rule: str = "first_closure") -> dict[str, Any]:
    known = set(graph)
    chains = [_chain_functions(t["chain"]) for t in tests]

    def run(rule_name):
        exercised = [_exercised([f for f in c if f in known], graph, rule_name) for c in chains]
        killed, by_kind = [], Counter()
        for f in fault_list:
            hit = next((t["tc_id"] for t, ex in zip(tests, exercised, strict=True) if discriminates(f, t, ex)), None)
            if hit:
                by_kind[f["kind"]] += 1
                killed.append({**f, "by": hit})
        return exercised, killed, by_kind

    exercised, killed, by_kind = run(rule)
    total = Counter(f["kind"] for f in fault_list)
    n = len(fault_list)
    sizes = sorted(len(e) for e in exercised)
    sensitivity = {}
    for other in _RULES:
        _e, k, _b = (exercised, killed, by_kind) if other == rule else run(other)
        sensitivity[other] = round(len(k) / n, 4) if n else None
    return {"tests": len(tests), "faults": n, "rule": rule, "discriminated": len(killed),
            "rate": round(len(killed) / n, 4) if n else None,
            "by_kind": {k: {"faults": total[k], "discriminated": by_kind[k]} for k in total},
            "rate_by_rule": sensitivity,
            # how much each test is credited with: the call-closure size behind the rule (``main`` exercises all)
            "exercised_functions": {"median": sizes[len(sizes) // 2] if sizes else None,
                                    "max": sizes[-1] if sizes else None,
                                    "tests_with_no_known_function": sum(1 for e in exercised if not e)},
            "chain_names_not_in_source": sum(1 for c in chains for f in c if f not in known),
            "tests_with_concrete_expectation": sum(1 for t in tests if any(concrete(v) for exp in t["expected_rows"]
                                                                          for v in exp.values())),
            "tests_with_return_injection": sum(any(_RETURN_COL.match(c) for c in t["inputs"]) for t in tests),
            "discriminated_faults": killed}


def evaluate(roots: list[Path], reference: str, generated: str | None = None) -> dict[str, Any]:
    from generators import c_project_context as cpc
    from generators.interface_contract import SourceIndex, file_scope_globals
    texts, unread = {}, []
    for root in roots:
        for p in sorted(root.rglob("*")):
            if p.suffix.lower() in (".c", ".h") and p.is_file():
                try:
                    texts[str(p)] = p.read_bytes().decode("utf-8")
                except UnicodeDecodeError:
                    unread.append(str(p))   # (round 2 W5) never read with a guessed encoding — counted and listed
    index = SourceIndex(texts, cpc.shared_parser())
    graph = call_graph(index)
    globals_ = file_scope_globals(texts, cpc.shared_parser())   # the contract's own definition (round 2 X5)
    fault_list = faults(index, graph, globals_)
    report = {"source_roots": [str(r) for r in roots], "reference_path": reference, "generated_path": generated,
              "functions": len(index.defs), "unreadable_source_files": unread,
              "duplicate_function_names": sorted(index.duplicates), "faults": len(fault_list),
              "faults_by_kind": dict(Counter(f["kind"] for f in fault_list)),
              "kind": "static annotation scoring — necessary conditions for discrimination, not an execution",
              "reference": score(read_sits(reference), fault_list, graph)}
    if generated:
        report["generated"] = score(read_sits(generated), fault_list, graph)
    return report


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--source-root", required=True)
    ap.add_argument("--reference", required=True)
    ap.add_argument("--generated", default="")
    ap.add_argument("--out", required=True)
    args = ap.parse_args(argv)
    roots = [Path(r.strip()) for r in re.split(r"[,;]", args.source_root) if r.strip()]
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from reference_alignment import _guard_output
    _guard_output(Path(args.out), args.reference, roots, inputs=(args.generated,))
    report = evaluate(roots, args.reference, args.generated or None)
    Path(args.out).write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({"faults": report["faults_by_kind"],
                      **{k: {"rate": report[k]["rate"], "rate_by_rule": report[k]["rate_by_rule"],
                             "by_kind": report[k]["by_kind"],
                             "tests_with_concrete_expectation": report[k]["tests_with_concrete_expectation"]}
                         for k in ("reference", "generated") if k in report}}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
