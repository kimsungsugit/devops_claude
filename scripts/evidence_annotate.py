"""Annotate a generated SUTS with an independent check of every derived expected value (R11).

The plan's R11 asked for a *host execution* of the derived vectors. On this target that would not prove anything: the
S12Z's ``int`` is 16 bits and a host's is 32, so a host run computes different values for the same C text. The
independent check that keeps the target's integer model is clang **constant evaluation on a 16-bit-int target**
(``--target=msp430``), which `scripts/source_oracle_clang_check.py` already runs. This script writes a **copy** of the
workbook whose "Test Evidence" sheet gains one column per row:

* ``clang_constexpr_agree`` — clang computes the same value (every fill of the state the vector leaves open);
* ``clang_constexpr_agree_stubbed_callees`` — the same, in a function that calls others (the harness stubs them:
  their effects are not what clang confirmed);
* ``clang_mismatch`` / ``clang_eval_error`` — a contradiction (the oracle model is wrong there);
* ``unchecked:<reason>`` — the harness could not evaluate it (a limit, not a verdict); ``unchecked:source_changed``
  — the file's text is not the one the row was derived from (its SHA-256 differs): the current text says nothing about
  that row;
* ``—`` — the row is not a derived value (nothing to check).

A "Check Legend" sheet in the copy says what each label means.

This is not an execution on the ECU and says so in the column header; ``Execution`` stays ``not_run``.

Usage:
    .venv/Scripts/python.exe scripts/evidence_annotate.py --xlsm SUTS.xlsm --source-root ROOT --out SUTS_checked.xlsm
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

COLUMN = "Independent Check (clang constexpr, 16-bit int target — not an execution)"
_LABEL = {"agree": "clang_constexpr_agree", "agree_with_stubbed_callees": "clang_constexpr_agree_stubbed_callees",
          "mismatch": "clang_mismatch", "eval_error": "clang_eval_error"}


def verdicts(xlsm: str, source_root: str, clang: str = "clang") -> tuple[dict[tuple, str], dict]:
    """(source path, function, test case, sequence, observable) → label, and the check's report."""
    from source_oracle_clang_check import _claims_from_xlsm, check_claims
    claims, meta = _claims_from_xlsm(xlsm, source_root)
    report = check_claims(claims, clang=clang)
    report["skipped_reasons"] = meta.get("skipped_reasons", {})
    out = {}
    # rows the check never saw say why (the file changed, or it is not in the given roots)
    derived = _derived_outputs(xlsm)
    for key, why in (meta.get("skipped_sequences") or {}).items():
        for name in derived.get(key, ()):
            out[(*key, name)] = "unchecked:" + why
    for claim in claims:
        for name in claim["outputs"]:
            v = (claim.get("verdicts") or {}).get(name, "unchecked:no_verdict")
            out[(*claim["key"], name)] = _LABEL.get(v, v)
    return out, report


def _derived_outputs(xlsm: str) -> dict[tuple, list]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsm, read_only=True)
    try:
        rows = wb["Test Evidence"].iter_rows(values_only=True)
        header = list(next(rows))
        out: dict[tuple, list] = {}
        for row in rows:
            d = dict(zip(header, row, strict=False))
            if d.get("Status") == "derived":
                out.setdefault((d["Source path"], d["Function"], d["Test Case ID"], d["Sequence"]), []).append(
                    d["Observable"])
        return out
    finally:
        wb.close()


LEGEND = [("clang_constexpr_agree", "clang 상수 평가(16비트 int 대상)가 같은 값을 냈다 — 벡터가 열어 둔 상태의 모든 채움 값에서"),
          ("clang_constexpr_agree_stubbed_callees", "같은 값 — 다만 이 함수가 부르는 함수는 하네스가 stub 했다(피호출 함수의 "
           "실제 동작은 확인하지 않았다)"),
          ("clang_mismatch / clang_eval_error", "모순 — 기대값을 낸 모델이 그 자리에서 틀렸다"),
          ("unchecked:<사유>", "하네스가 평가하지 못했다(판정 아님). source_changed = 행을 만든 소스와 지금 소스가 다르다"),
          ("—", "확정값이 아닌 행 — 확인할 값이 없다"),
          ("(공통)", "타깃 실행이 아니다. Execution 열은 not_run 그대로다.")]


def annotate(xlsm: str, out_path: str, labels: dict[tuple, str]) -> Counter:
    import os
    import tempfile

    import openpyxl
    wb = openpyxl.load_workbook(xlsm, keep_vba=xlsm.lower().endswith(".xlsm"))
    ws = wb["Test Evidence"]
    header = [c.value for c in ws[1]]
    col = header.index(COLUMN) + 1 if COLUMN in header else len(header) + 1
    ws.cell(1, col, COLUMN)
    idx = {h: i for i, h in enumerate(header)}
    counts: Counter = Counter()
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[idx["Status"]] != "derived":
            label = "—"
        else:
            key = (row[idx["Source path"]], row[idx["Function"]], row[idx["Test Case ID"]], row[idx["Sequence"]],
                   row[idx["Observable"]])
            label = labels.get(key, "unchecked:not_in_check")
        ws.cell(r, col, label)
        counts[label.split(":")[0] if label.startswith("unchecked") else label] += 1
    if "Check Legend" in wb.sheetnames:
        del wb["Check Legend"]
    legend = wb.create_sheet("Check Legend")
    legend.append(["Label", "Meaning"])
    for row in LEGEND:
        legend.append(list(row))
    # (review I6) written next to the target and moved into place: a failed save never leaves half a workbook
    fd, tmp = tempfile.mkstemp(suffix=Path(out_path).suffix, dir=str(Path(out_path).resolve().parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, out_path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return counts


def guard_output(out: str, xlsm: str, source_root: str) -> None:
    """(review W5) never the input (also through a hard link), never inside a read-only source tree, and the same
    workbook kind (a ``.xlsm`` copy keeps its macros; saved as ``.xlsx`` Excel would refuse it)."""
    import os
    import re as _re
    roots = [Path(r.strip()) for r in _re.split(r"[,;]", source_root) if r.strip()]
    target = Path(out)
    if Path(xlsm).resolve() == target.resolve() or (target.exists() and os.path.samefile(xlsm, target)):
        raise SystemExit("refusing to overwrite the input workbook: give --out another path")
    # the input's own folder may hold the copy (unlike `_guard_output`, which protects a reference's folder)
    for protected in (r.resolve() for r in roots):
        if target.resolve() == protected or protected in target.resolve().parents:
            raise SystemExit(f"refusing to write {target.resolve()}: inside read-only source {protected}")
    if target.suffix.lower() != Path(xlsm).suffix.lower():
        raise SystemExit(f"--out must keep the input's kind ({Path(xlsm).suffix}), not {target.suffix}")
    # (round 3 I5) before the clang run, not after it: the save would fail only at the end
    if not target.resolve().parent.is_dir():
        raise SystemExit(f"--out folder does not exist: {target.resolve().parent}")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--xlsm", required=True, help="a generated SUTS (not modified)")
    ap.add_argument("--source-root", required=True)
    ap.add_argument("--out", required=True, help="the annotated copy")
    ap.add_argument("--clang", default="clang")
    args = ap.parse_args(argv)
    guard_output(args.out, args.xlsm, args.source_root)
    labels, report = verdicts(args.xlsm, args.source_root, args.clang)
    counts = annotate(args.xlsm, args.out, labels)
    print(json.dumps({"rows": dict(counts), "check": {k: report[k] for k in ("claims", "checked", "agree", "mismatch",
                                                                            "eval_error", "unchecked",
                                                                            "skipped_reasons")}},
                     ensure_ascii=False))
    return 1 if report["mismatch"] or report["eval_error"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
