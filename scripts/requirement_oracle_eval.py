"""Requirement oracle evaluation (R5, G4) against a reference STS — independent of the extractor.

Two measures, both scored on what a *reference* (or generated) STS workbook says, never on the extractor's own output:

* **recall of reference thresholds** (G4 a): every threshold the reference STS uses in a test action or judgement
  (``15도 미만``, ``0.8m/s 이상``, ``5도( 0x05 ) 이하``) for an SRS ID is a gold item. It is recalled when the
  extractor found a fact with that value (unit-compatible) in the same requirement. The operator agreement is reported
  separately (the STS may state the other side of a threshold: ``15도 이상`` tests ``15도 미만``'s complement).
  Tolerance checks (``7 ± 5%``) are output judgements, not requirement thresholds — counted, not scored.
  Precision has no independent gold here: the facts are exported for a reviewer (``--review-csv``) and precision is
  reported ``None`` until labelled — a model grading its own extraction is not a measurement.

* **requirement-mutant discrimination** (G4 b): each parsed threshold / duration fact yields mutants — the boundary
  inclusion flipped (``>=`` ↔ ``>``) and the value moved one written step (``8.5`` → ``8.4``/``8.6``; ``4.00`` → ±0.01).
  A suite discriminates a mutant when one of its stimulus *points* for that requirement falls where the **line's**
  verdict (every condition on that subject, joined as the line joins them) **holds for the original and not for the
  mutant** — the step there expects the behaviour, and the mutant would not produce it. A point where the original
  verdict is false does not count (review r4 C-1): outside its condition a sentence claims nothing (another sentence of
  the requirement may act there), so a test at that point has no expectation the mutant could contradict. Separations
  on either side are reported as ``killed_either_side`` — an upper bound that holds only if the suite also asserts
  non-action.
  A point is a number in a test action: unnamed with a unit (``전원을 7V로 설정한다``) for any subject of that unit, or
  named (``u16g_ApiIn_Vsup = 1605``) for that subject only; a requirement quote after ``— 근거:`` is not a point. A
  comparator region in a test action (``0.8m/s 이상 빠르기``) counts only in the ``optimistic`` rate, as if the tester
  had picked the boundary value itself. Response constraints (``100ms 이내``) are measured, not stimulated — no mutant.

  ⚠ **Not independent for a generated STS built by `generators/sts_requirement_tc.py`**: it writes its points from
  the same extractor and the same written-precision step, so every fact it uses is discriminated by construction. The
  number measures whether it wrote every fact it could, and the reference's number measures what the reference tests —
  an independent mutant source (human-labelled facts, `--review-csv`) is needed before either is called superior.

* **cross-source discrimination** (R18, gap ③): mutants made from the **reference STS's own thresholds** — boundaries
  human testers wrote, independent of the extractor — scored on both suites with the same point rule. The verdict is
  the single comparison the tester wrote (``15도 미만`` → ``x < 15``), and a point is any stimulus of that requirement
  in a compatible unit: the threshold names no subject, so a point on another quantity of the same unit could be
  credited — kills are split by whether the requirement has one subject in that unit (``killed_single_subject``) or
  several (``killed_multi_subject``, an upper bound). This source favours what the reference chose to test, the
  extractor's favours what the SRS states — the generated suite is scored here on mutants it did not produce (the
  reference's own row is self-sourced: its regions are the thresholds, so only its points are reported). Thresholds
  the SRS block does not state are kept — a suite built from the SRS alone cannot know them (they may come from a
  system specification, a scenario precondition, or an SRS revision the reference was not written against).
  Left out and counted (``excluded``): thresholds only in the expected-result column, ones every row writes in both
  its action and its expected cell (``judgement_copied_to_action``), ones the SRS states only as a response constraint,
  unitless ones (``no_unit``). ``killable`` marks mutants a point can kill at all (a widening mutant never can); kills
  are split into ``killed_single_subject`` / ``killed_multi_subject`` / ``killed_unnamed_subject``, and a suite with no
  point in a compatible unit is not measured (``measurable`` false — the same holds for the R5 ``discrimination``).

Usage:
    .venv/Scripts/python.exe scripts/requirement_oracle_eval.py --srs SRS.txt --sts STS.xlsm [--sts-generated GEN.xlsm]
        --out r5.json [--review-csv facts.csv]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from decimal import Decimal
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from generators.requirement_oracle import (  # noqa: E402
    exact_value,
    is_outcome_section,
    line_holds,
    subject_of,
    written_step,
)

# longest first (``5step`` is not s); ``8v 이하``/``Km/h`` as the documents write them (R18 review W5)
_UNIT = r"(?:[Kk]m/h|KM/H|KPH|m/s|step|sec|ms|deg|℃|°C|mV|mv|mA|uA|µA|Hz|V|v|s|초|분|도|%|A)"
_NUM = r"\d+(?:\.\d+)?"
# a written threshold: its own number — not the tail of ``0x1FF0``, with its sign (``-4도``) and without its thousands
# separators (``1,000ms``) — R18 review W5
# (not ``\w``: Hangul is a word character, and ``열림각15도 이상`` / ``전압을9V 이상`` are thresholds — R18 review R2 W-B;
#  after a comma only a thousands group is refused — ``1,000`` is one number, ``Task(5,10,50ms)`` ends in ``50ms`` — R3 I-1)
_CMP_NUM = r"(?<![0-9A-Za-z_.])(?!(?<=\d,)\d{3}(?!\d))-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"
_OPS = {"이상": ">=", "이하": "<=", "미만": "<", "초과": ">"}
_COMPARATOR = re.compile(rf"(?P<num>{_CMP_NUM})\s*(?P<unit>{_UNIT})?\s*(?:\(\s*0x[0-9A-Fa-f]+\s*\))?\s*"
                         rf"(?P<op>이상|이하|미만|초과)")
_TOLERANCE = re.compile(rf"(?P<num>{_CMP_NUM})\s*(?P<unit>{_UNIT})?\s*±\s*(?P<tol>{_NUM})\s*(?P<tunit>{_UNIT})?")
_POINT = re.compile(rf"(?P<num>{_CMP_NUM})\s*(?P<unit>{_UNIT})\s*(?:\(\s*0x[0-9A-Fa-f]+\s*\))?\s*(?:로|으로|에서|를|을|만큼)?\s*"
                    rf"(?:설정|인가|유지|입력|회전|이동|열|닫|대기|경과)")
# ``u16g_ApiIn_Vsup = 1605 설정`` — a named stimulus; its unit may be absent (a raw ADC count)
_NAMED_POINT = re.compile(rf"(?P<sig>[A-Za-z_][A-Za-z0-9_]*|[가-힣A-Za-z][가-힣A-Za-z0-9_ ]{{0,30}}?)\s*=\s*"
                          rf"(?P<num>-?(?:0[xX][0-9A-Fa-f]+|{_NUM}))\s*(?P<unit>{_UNIT})?(?=\s*(?:설정|유지|$))")
_BASIS_MARK = " — 근거:"   # a quoted requirement after it is not a stimulus (``generators/sts_requirement_tc.BASIS_MARK``)
# ``입력 설정 (요구 경계): <subject> = <value><unit> …`` — the generated step names its subject verbatim (``B+``,
# ``Load(Average)``, ``LIN 통신 지속 시간``): read it by position, not by a name pattern (review r2 W2)
_BOUNDARY_STEP = re.compile(rf"^입력 설정 \(요구 경계\):\s*(?P<sig>.+?)\s+=\s+(?P<num>-?(?:0[xX][0-9A-Fa-f]+|{_NUM}))"
                            rf"\s*(?P<unit>{_UNIT})?")
_DURATION_SUFFIX = " 지속 시간"


def _num(text: str) -> float:
    return float(int(text, 16)) if text.lower().lstrip("-").startswith("0x") else float(text)
_SHEETS = ("3.SW Integration Test Spec", "3.SW Test Spec", "2.SW Test Spec")
_SRS_ID = re.compile(r"Sw[A-Za-z]+_[A-Za-z0-9_]+")


def _unit(u: str | None) -> str:
    u = (u or "").strip()
    return {"sec": "s", "초": "s", "deg": "도", "°C": "℃", "KPH": "km/h", "Km/h": "km/h", "KM/H": "km/h", "v": "V",
            "mv": "mV", "µA": "uA"}.get(u, u)


_TIME = {"ms": 0.001, "s": 1.0, "분": 60.0}


def _seconds(value: float, unit: str) -> float | None:
    return value * _TIME[unit] if unit in _TIME else None


def _same_quantity(a: float, ua: str, b: float, ub: str) -> bool:
    """Equal values in compatible units: same unit, either side unitless, or two time units."""
    ua, ub = _unit(ua), _unit(ub)
    if ua == ub or not ua or not ub:
        return abs(a - b) < 1e-9
    sa, sb = _seconds(a, ua), _seconds(b, ub)
    return sa is not None and sb is not None and abs(sa - sb) < 1e-9


def read_sts(path: str) -> dict[str, dict[str, list]]:
    """SRS ID → {"thresholds", "tolerances", "points", "regions"} from the test action (K) and expected result (L)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        name = next((s for s in _SHEETS if s in wb.sheetnames), None)
        if name is None:
            raise ValueError(f"no STS spec sheet in {Path(path).name}: {wb.sheetnames}")
        out: dict[str, dict[str, list]] = {}
        srs = None
        for row_no, row in enumerate(wb[name].iter_rows(min_row=4, values_only=True), start=4):
            cells = list(row) + [None] * 13
            ids = _SRS_ID.findall(str(cells[12] or ""))
            if ids:
                srs = ids[0]
            if srs is None:
                continue
            slot = out.setdefault(srs, {"thresholds": [], "tolerances": [], "points": [], "regions": []})
            for col, role in ((10, "action"), (11, "expected")):
                text = str(cells[col] or "").split(_BASIS_MARK)[0]
                taken = []
                for m in _TOLERANCE.finditer(text):
                    taken.append((m.start(), m.end()))
                    slot["tolerances"].append({"value": float(m.group("num").replace(",", "")), "unit": _unit(m.group("unit")),
                                               "raw": m.group(0), "role": role})
                for m in _COMPARATOR.finditer(text):
                    if any(s <= m.start() < e for s, e in taken):
                        continue
                    item = {"value": float(m.group("num").replace(",", "")), "unit": _unit(m.group("unit")),
                            "op": _OPS[m.group("op")], "raw": m.group(0), "role": role, "row": row_no,
                            "text": m.group("num").replace(",", "")}   # as written: sign and decimal places
                    slot["thresholds"].append(item)
                    if role == "action":
                        slot["regions"].append(item)
                if role == "action" and (b := _BOUNDARY_STEP.match(text)):
                    slot["points"].append({"value": _num(b.group("num")), "unit": _unit(b.group("unit")),
                                           "signal": b.group("sig").strip(), "raw": b.group(0)})
                elif role == "action":
                    named = []
                    for m in _NAMED_POINT.finditer(text):
                        named.append((m.start("num"), m.end()))
                        slot["points"].append({"value": _num(m.group("num")), "unit": _unit(m.group("unit")),
                                               "signal": m.group("sig").strip(), "raw": m.group(0)})
                    for m in _POINT.finditer(text):
                        if any(s <= m.start() < e for s, e in named):
                            continue
                        slot["points"].append({"value": float(m.group("num").replace(",", "")),
                                               "unit": _unit(m.group("unit")), "signal": None, "raw": m.group(0)})
        return out
    finally:
        wb.close()


def _fact_values(fact: dict) -> list[float]:
    v = fact.get("value")
    if isinstance(v, list):
        return [float(x) for x in v if isinstance(x, (int, float))]
    return [float(v)] if isinstance(v, (int, float)) else []


def _stated_in(text: str, value: float, unit: str) -> bool:
    """The quantity is written in the requirement's own text: the number followed by a compatible unit (``3도``,
    ``300ms`` for ``0.3 s``). A bare number is not enough — list numbers (``15. …``), IDs (``SyEI_03``) and other
    quantities (``50스텝`` for ``50도``) would make gold items the text never states."""
    for m in re.finditer(rf"(?P<num>{_CMP_NUM})\s*(?P<unit>{_UNIT})", text):
        if _same_quantity(float(m.group("num").replace(",", "")), m.group("unit"), value, unit):
            return True
    return False


def recall(requirements: list[dict], sts: dict, blocks: dict[str, str]) -> dict[str, Any]:
    """Only gold items whose number the SRS block itself writes are scored: a threshold the reference STS took from
    elsewhere (system spec, design) is not one the requirement text states — counted as ``not_stated_in_srs``."""
    facts_by_req = {r["req_id"]: [f for line in r["lines"] for f in line["facts"]] for r in requirements}
    items, found, op_agree, missing, not_stated = 0, 0, 0, [], []
    for req, slot in sorted(sts.items()):
        gold = {(t["value"], t["unit"]): t for t in slot["thresholds"]}  # distinct per requirement
        for (value, unit), t in sorted(gold.items()):
            if not _stated_in(blocks.get(req, ""), value, unit):
                not_stated.append({"req_id": req, "raw": t["raw"], "srs_block_found": req in blocks})
                continue
            items += 1
            hits = [f for f in facts_by_req.get(req, []) if any(_same_quantity(value, unit, x, f.get("unit") or "")
                                                                 for x in _fact_values(f))]
            if hits:
                found += 1
                complement = {">=": "<", "<": ">=", "<=": ">", ">": "<="}
                if any(f.get("op") in (t["op"], complement[t["op"]]) for f in hits):
                    op_agree += 1
            else:
                missing.append({"req_id": req, "raw": t["raw"], "in_srs_model": req in facts_by_req})
    return {"gold_items": items, "recalled": found, "recall": round(found / items, 4) if items else None,
            "operator_consistent": op_agree, "missing": missing, "not_stated_in_srs": not_stated}


def _tokens(x: str) -> list[str]:
    return re.sub(r"[^0-9A-Za-z가-힣_]+", " ", x).split()


def _same_signal(a: str, b: str) -> bool:
    """The same subject written with or without leading words / punctuation (``High -> Low Power mode 천이 시간`` read back
    from a test action as ``Low Power mode 천이 시간``): one normalized name ends the other. The shorter must still name
    something — an identifier or at least two words (``전압`` alone ends every voltage)."""
    na, nb = _tokens(a), _tokens(b)
    if not na or not nb:
        return False
    short, long_ = (na, nb) if len(na) <= len(nb) else (nb, na)
    if len(short) < 2 and "_" not in short[0] and not re.search(r"[a-z][A-Z]", short[0]) and short != long_:
        return False
    return long_[-len(short):] == short


def requirement_mutants(requirements: list[dict]) -> tuple[list[dict], dict[str, int]]:
    """Mutants of every parsed order comparison, and the facts left out with why."""
    flip = {">=": ">", ">": ">=", "<=": "<", "<": "<="}
    out, excluded = [], Counter()
    for r in requirements:
        for line in r["lines"]:
            for f in line["facts"]:
                if f["status"] != "parsed" or f["kind"] not in {"threshold", "duration"} or f.get("op") not in flip:
                    continue
                if not isinstance(f.get("value"), (int, float)) or isinstance(f.get("value"), bool):
                    continue
                if f.get("role") == "response_constraint":
                    excluded["response_constraint"] += 1
                    continue
                if f.get("negated"):
                    excluded["negated_condition"] += 1   # ``8.5V 이상을 만족하지 못하면``: the sense is reversed
                    continue
                if is_outcome_section(line.get("section") or ""):
                    excluded["outcome_section"] += 1   # ``<완료조건>`` / ``<Output>``: produced, not stimulated
                    continue
                value, step = exact_value(f), written_step(f)
                # the subject a stimulus sets: a parameter threshold's monitored quantity (review r2 C-B)
                base = {"req_id": r["req_id"], "signal": subject_of(f), "kind": f["kind"], "unit": _unit(f.get("unit")),
                        "op": f["op"], "value": float(value), "value_text": f.get("value_text"),
                        "raw": f.get("raw", ""), "_fact": f, "_line": line}
                out.append({**base, "mutant": "boundary_inclusion", "m_op": flip[f["op"]], "m_value": value})
                for d in (-step, step):
                    out.append({**base, "mutant": "value_shift", "m_op": f["op"], "m_value": value + d})
    return out, dict(excluded)


def discrimination(mutants: list[dict], sts: dict) -> dict[str, Any]:
    """``killed``: a stimulus point separates the mutant (strict). ``killed_optimistic`` also credits a comparator region
    in a test action at the fact's own value and unit (``15도 이상`` for ``15도 미만``) as if the tester had chosen the
    boundary value — an upper bound for a suite that states regions, not points."""
    killed, either, optimistic, by_kind, total = 0, 0, 0, Counter(), Counter()
    unjudgeable = 0
    rows = []

    def compatible(a, b):
        # same unit, or two time units (``300ms`` against ``0.3 s`` — the verdict scales it, review r2 W3)
        a, b = _unit(a), _unit(b)
        return a == b or (_seconds(1.0, a) is not None and _seconds(1.0, b) is not None)

    def usable(p, m):
        if m["kind"] == "duration":
            # a hold time is stimulated by how long the condition lasts: a ``… 지속 시간`` point, or an unnamed time
            named_ok = not p.get("signal") or p["signal"].endswith(_DURATION_SUFFIX.strip())
            return named_ok and bool(p["unit"]) and compatible(p["unit"], m["unit"] or "s")
        if p.get("signal"):
            # a named point stimulates its own subject only — a fact with no subject is not "any subject" (r3 W4)
            same_unit = not m["unit"] or not p["unit"] or compatible(p["unit"], m["unit"])
            return same_unit and bool(m.get("signal")) and _same_signal(p["signal"], m["signal"])
        # an unnamed point has a unit; a unit-less fact (a raw count) needs a named point
        return bool(m["unit"]) and bool(p["unit"]) and compatible(p["unit"], m["unit"])

    def line_verdict(p, m, override=None):
        # the point in the fact's unit (``300ms`` against ``0.3 s``), exact decimals
        return line_holds(Decimal(str(p["value"])) * _scale(p["unit"], m["unit"]), m["_fact"], m["_line"], override)

    for m in mutants:
        total[m["mutant"]] += 1
        slot = sts.get(m["req_id"]) or {}
        candidates = [p for p in slot.get("points", []) if usable(p, m)]
        points = [p["value"] for p in candidates]
        hit = side = None
        judged = True
        for p in candidates:
            a, b = line_verdict(p, m), line_verdict(p, m, (m["m_op"], m["m_value"]))
            if a is None or b is None:
                judged = False
                break
            if a and not b:
                hit = p["value"]      # the step expects the behaviour here; the mutant does not produce it
                break
            if a != b and side is None:
                side = p["value"]     # separated only where the original claims nothing
        if not judged:
            unjudgeable += 1
        if hit is not None or side is not None:
            either += 1
        region = next((r["raw"] for r in slot.get("regions", []) if _same_quantity(r["value"], r["unit"], m["value"],
                                                                                   m["unit"])), None)
        if hit is not None:
            killed += 1
            by_kind[m["mutant"]] += 1
        if hit is not None or region is not None:
            optimistic += 1
        rows.append({**{k: v for k, v in m.items() if not k.startswith("_")}, "m_value": float(m["m_value"]),
                     "points": points, "killed_by_point": hit, "separated_where_original_false": side,
                     "boundary_region": region})
    n = len(mutants)
    # (R18 review R3 W-3) a suite with no usable point for any mutant is not measured by points: its 0 is no score
    measurable = any(r["points"] for r in rows)
    return {"mutants": n, "measurable": measurable, "killed": killed,
            "rate": round(killed / n, 4) if n and measurable else None,
            "killed_either_side": either, "rate_either_side": round(either / n, 4) if n and measurable else None,
            "killed_optimistic": optimistic, "rate_optimistic": round(optimistic / n, 4) if n else None,
            "unjudgeable_combination": unjudgeable,
            "by_kind": {k: {"mutants": total[k], "killed": by_kind[k]} for k in total}, "rows": rows}


_COMPARE = {">=": lambda x, v: x >= v, ">": lambda x, v: x > v, "<=": lambda x, v: x <= v, "<": lambda x, v: x < v}
_FLIP = {">=": ">", ">": ">=", "<=": "<", "<": "<="}


def _units_compatible(a: str | None, b: str | None) -> bool:
    """Both units written and the same quantity: the same unit, or two time units (``300ms`` / ``0.3 s``). A threshold
    or a point without a unit cannot be placed without a subject (points and regions judged alike — R18 review W4)."""
    a, b = _unit(a), _unit(b)
    if not a or not b:
        return False
    return a == b or (_seconds(1.0, a) is not None and _seconds(1.0, b) is not None)


def _killable(op: str, value: Decimal, m_op: str, m_value: Decimal) -> bool:
    """Is there a point where the written comparison holds and the mutant's does not? A mutant that only widens the
    condition (``<`` → ``<=``, the value moved outward) can never be killed by a point that asserts the behaviour."""
    if op in {">=", ">"}:
        return m_value > value or (m_value == value and op == ">=" and m_op == ">")
    return m_value < value or (m_value == value and op == "<=" and m_op == "<")


def reference_mutants(reference: dict, blocks: dict[str, str] | None = None,
                      facts_by_req: dict[str, list] | None = None) -> tuple[list[dict], dict[str, int]]:
    """(R18) Mutants of every distinct order threshold the reference STS writes **as a stimulus** for a requirement:
    the boundary inclusion flipped and the value moved one **written** step (``0.8`` → ±0.1, ``15`` → ±1, the finest
    when a value is written twice). Left out, counted (R18 review W2, I4, R2 W-A): a threshold only in the
    expected-result column (an output judgement); one every row writes in **both** its action and its expected cell
    (acceptance criteria copied, ``CPU 부하 70%이하``); one the SRS states only as a response constraint
    (``100ms 이내``); a unitless one. Each mutant carries whether the SRS block states the quantity, whether the block
    exists, and whether the extractor recalled it."""
    out, excluded = [], Counter()
    for req, slot in sorted(reference.items()):
        groups: dict[tuple, dict] = {}
        for t in slot["thresholds"]:
            if t["op"] not in _FLIP:
                excluded["not_an_order_comparison"] += 1
                continue
            text = t.get("text") or str(t["value"])
            key = (Decimal(text), _unit(t["unit"]), t["op"])
            g = groups.setdefault(key, {"text": text, "roles": set(), "raw": t["raw"], "rows": {}})
            g["roles"].add(t["role"])
            g["rows"].setdefault(t.get("row"), set()).add(t["role"])
            if len(text.split(".")[1] if "." in text else "") > len(g["text"].split(".")[1] if "." in g["text"] else ""):
                g["text"] = text   # the finer written precision (``1.30`` over ``1.3`` — R18 review I2)
        for (value, unit, op), g in sorted(groups.items(), key=lambda x: (x[0][0], x[0][1], x[0][2])):
            if "action" not in g["roles"]:
                excluded["expected_result_only"] += 1
                continue
            if all(roles == {"action", "expected"} for r, roles in g["rows"].items() if "action" in roles and r is not None) \
                    and any(r is not None for r in g["rows"]):
                # every row that writes it as an action also writes it as that row's judgement: acceptance criteria
                # copied into both cells (KJPDS02 SwNTR_0301 ``CPU 부하 70%이하``) — R18 review R2 W-A
                excluded["judgement_copied_to_action"] += 1
                continue
            if not unit:
                excluded["no_unit"] += 1   # a point cannot be placed on a unitless threshold without its subject (I-d)
                continue
            block = (blocks or {}).get(req)
            stated = _stated_in(block or "", float(value), unit)
            same = [(f, line) for f, line in (facts_by_req or {}).get(req, [])
                    if any(_same_quantity(float(value), unit, x, f.get("unit") or "") for x in _fact_values(f))]
            # a response constraint (``100ms 이내``) is measured, not stimulated. (Not by section: the extractor's
            # section runs on past the last ``<Output>`` heading into the attribute table, so HDPDM01's verification
            # criteria ``Param( 3도 ) 초과한 열림각에서 끼임 발생한 경우`` — an input — read as Output; R2 I-a)
            if same and all(f.get("role") == "response_constraint" for f, _line in same):
                excluded["srs_states_it_as_output"] += 1
                continue
            step = Decimal(1).scaleb(-len(g["text"].split(".")[1]) if "." in g["text"] else 0)
            base = {"req_id": req, "unit": unit, "op": op, "value": value, "raw": g["raw"], "stated_in_srs": stated,
                    "srs_block_found": block is not None, "recalled_by_extractor": bool(same)}
            for mutant, m_op, m_value in (("boundary_inclusion", _FLIP[op], value), ("value_shift", op, value - step),
                                          ("value_shift", op, value + step)):
                out.append({**base, "mutant": mutant, "m_op": m_op, "m_value": m_value,
                            "killable": _killable(op, value, m_op, m_value)})
    return out, dict(excluded)


def _subject_groups(points: list[dict]) -> list[str]:
    """Distinct subjects among named points: a name that ends a longer one (``Low Power mode 천이 시간`` /
    ``High -> Low Power mode 천이 시간``) is that one; two long names a short one ends both of stay two
    (``저 전압 … 기준 전압`` / ``과 전압 … 기준 전압``). Longest first, so the result does not depend on the order
    the points come in (R18 review R3 W-1)."""
    names = sorted({p.get("signal") for p in points if p.get("signal")}, key=lambda x: (-len(_tokens(x)), x))
    groups: list[str] = []
    for name in names:
        if not any(_same_signal(name, g) for g in groups):
            groups.append(name)
    return groups


def cross_discrimination(mutants: list[dict], sts: dict, self_sourced: bool = False) -> dict[str, Any]:
    """(R18) ``killed``: a point of the requirement, in a compatible unit, where the written comparison holds and the
    mutant's does not (the step expects the behaviour there). ``killed_either_side`` counts a separation on either
    side; ``killed_optimistic`` also credits a comparator region at the threshold's own value (``15도 이상``), as if
    the tester had chosen the boundary value itself — **not reported for the suite the thresholds came from**: its
    regions *are* those thresholds, so the number would be fixed by construction (R18 review C1). A suite with no point
    in a compatible unit for any mutant is not measured (``measurable`` false, rates None — R2 W-D). A point names its
    subject only sometimes: kills are split into ``killed_single_subject`` (the requirement's same-unit points name one
    subject), ``killed_multi_subject`` (several — another quantity of that unit could have supplied it, an upper bound)
    and ``killed_unnamed_subject`` (an unnamed point: any quantity of that unit — R2 W-E)."""
    killed = either = optimistic = with_points = killable = single = multi = unnamed = 0
    by_stated = Counter()
    recall = Counter()
    rows = []
    for m in mutants:
        slot = sts.get(m["req_id"]) or {}
        points = [p for p in slot.get("points", []) if _units_compatible(p["unit"], m["unit"])]
        with_points += bool(points)
        hit = side = None
        for p in points:
            x = Decimal(str(p["value"])) * _scale(p["unit"], m["unit"])
            a, b = _COMPARE[m["op"]](x, m["value"]), _COMPARE[m["m_op"]](x, m["m_value"])
            if a and not b and hit is None:
                hit = p
            if a != b and side is None:
                side = p
        region = next((r["raw"] for r in slot.get("regions", []) if _units_compatible(r["unit"], m["unit"])
                       and _same_quantity(r["value"], r["unit"], float(m["value"]), m["unit"])), None)
        subjects = _subject_groups(points)
        killed += hit is not None
        either += hit is not None or side is not None
        optimistic += hit is not None or region is not None
        killable += m["killable"]
        if hit is not None:
            if not hit.get("signal"):
                unnamed += 1
            elif len(subjects) == 1 and all(p.get("signal") for p in points):
                single += 1
            else:
                multi += 1
        by_stated[(m["stated_in_srs"], hit is not None)] += 1
        recall[(m["recalled_by_extractor"], hit is not None)] += 1
        rows.append({**m, "value": float(m["value"]), "m_value": float(m["m_value"]),
                     "points": [p["value"] for p in points], "killed_by_point": None if hit is None else hit["value"],
                     "killed_by_signal": None if hit is None else (hit.get("signal") or ""),
                     "subjects_in_unit": subjects, "boundary_region": None if self_sourced else region})
    n = len(mutants)
    measurable = bool(with_points)

    def rate(k, d):
        return round(k / d, 4) if measurable and d else None
    stated = sum(v for (st, _k), v in by_stated.items() if st)
    return {"mutants": n, "measurable": measurable, "killed": killed, "rate": rate(killed, n),
            "killed_either_side": either, "rate_either_side": rate(either, n),
            "killed_optimistic": None if self_sourced else optimistic,
            "rate_optimistic": None if self_sourced or not n else round(optimistic / n, 4),
            "self_sourced": self_sourced,
            "self_sourced_note": "the suite the thresholds came from: its regions are those thresholds, so the "
                                 "optimistic rate would be fixed by construction — not reported" if self_sourced else "",
            # an unkillable mutant never satisfies "holds and the mutant does not": every kill is of a killable one
            "killable_mutants": killable, "rate_of_killable": rate(killed, killable),
            "killed_single_subject": single, "killed_multi_subject": multi, "killed_unnamed_subject": unnamed,
            "mutants_with_a_point_in_unit": with_points,
            "stated_in_srs": {"mutants": stated, "killed": by_stated[(True, True)]},
            "not_stated_in_srs": {"mutants": n - stated, "killed": by_stated[(False, True)]},
            "recalled_by_extractor": {"killed": recall[(True, True)], "not_killed": recall[(True, False)]},
            "not_recalled": {"killed": recall[(False, True)], "not_killed": recall[(False, False)]},
            "rows": rows}


def _scale(point_unit: str, fact_unit: str) -> Decimal:
    """Factor turning a point in ``point_unit`` into ``fact_unit`` (``300ms`` against ``0.3 s``); 1 otherwise."""
    a, b = _seconds(1.0, _unit(point_unit)), _seconds(1.0, _unit(fact_unit))
    if a is None or b is None or a == b:
        return Decimal(1)
    return Decimal(str(a)) / Decimal(str(b))


def evaluate(srs_text: str, sts_path: str, generated_path: str | None = None) -> dict[str, Any]:
    from generators.requirement_oracle import model, split_requirements
    m = model(srs_text)
    blocks = {b["req_id"]: b["text"] for b in split_requirements(srs_text)}
    reference = read_sts(sts_path)
    mutants, excluded = requirement_mutants(m["requirements"])
    report = {"srs_model": m["summary"], "reference_sts": sts_path,
              "reference_items": {"thresholds": sum(len(s["thresholds"]) for s in reference.values()),
                                  "tolerances": sum(len(s["tolerances"]) for s in reference.values()),
                                  "points": sum(len(s["points"]) for s in reference.values()),
                                  "regions": sum(len(s["regions"]) for s in reference.values())},
              "recall": recall(m["requirements"], reference, blocks),
              "precision": None, "precision_note": "no independent gold — see --review-csv",
              "mutants_excluded": excluded,
              "independence_note": "a generated STS built from the same extractor is discriminated by construction",
              "discrimination": {"reference": discrimination(mutants, reference)}}
    generated = read_sts(generated_path) if generated_path else None
    if generated is not None:
        report["discrimination"]["generated"] = discrimination(mutants, generated)
    # (R18) the other direction: mutants from the reference's own thresholds (independent of the extractor)
    facts_by_req = {r["req_id"]: [(f, line) for line in r["lines"] for f in line["facts"]] for r in m["requirements"]}
    cross, cross_excluded = reference_mutants(reference, blocks, facts_by_req)
    report["cross_source"] = {
        "source": "reference STS thresholds (human-written), stimuli only", "thresholds": len(cross) // 3,
        "boundaries": len({(x["req_id"], x["value"], x["unit"]) for x in cross}), "excluded": cross_excluded,
        "note": "the generated suite is scored on mutants it did not produce (the reference's thresholds); the "
                "reference's own row shows its points only — its regions are these thresholds (self-sourced). A point "
                "counts for any subject of the requirement in a compatible unit (the threshold names none): see "
                "killed_single_subject / killed_multi_subject",
        "reference": cross_discrimination(cross, reference, self_sourced=True)}
    if generated is not None:
        report["cross_source"]["generated"] = cross_discrimination(cross, generated)
    return report


def write_review_csv(srs_text: str, path: str) -> int:
    from generators.requirement_oracle import model
    rows = 0
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["req_id", "line", "kind", "signal", "op", "value", "unit", "status", "raw", "verdict(correct/wrong)"])
        for r in model(srs_text)["requirements"]:
            for line in r["lines"]:
                for f in line["facts"]:
                    w.writerow([r["req_id"], line["text"][:160], f["kind"], f.get("signal") or "", f.get("op"),
                                json.dumps(f.get("value")), f.get("unit") or "", f["status"], f.get("raw", ""), ""])
                    rows += 1
    return rows


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--srs", required=True, help="SRS text export (read-only)")
    ap.add_argument("--sts", required=True, help="reference STS workbook (read-only)")
    ap.add_argument("--sts-generated", default="")
    ap.add_argument("--out", required=True)
    ap.add_argument("--review-csv", default="")
    args = ap.parse_args(argv)
    text = Path(args.srs).read_text(encoding="utf-8")
    report = evaluate(text, args.sts, args.sts_generated or None)
    if args.review_csv:
        report["review_csv_rows"] = write_review_csv(text, args.review_csv)
    Path(args.out).write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    d = report["discrimination"]
    print(json.dumps({"recall": report["recall"]["recall"], "gold": report["recall"]["gold_items"],
                      "operator_consistent": report["recall"]["operator_consistent"],
                      "not_stated_in_srs": len(report["recall"]["not_stated_in_srs"]),
                      **{k: (v["killed"] if v["measurable"] else None, v["killed_optimistic"], v["mutants"])
                         for k, v in d.items()},
                      **{f"cross_{k}": (v["killed"] if v["measurable"] else None, v["killed_optimistic"], v["mutants"])
                         for k, v in report["cross_source"].items() if isinstance(v, dict) and "killed" in v}},
                     ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
